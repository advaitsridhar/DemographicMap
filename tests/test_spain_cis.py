"""Spain's religion from the CIS's pre-electoral surveys: reading, checks, binding, records.

No network: the workbook, the HTML and PDF marginals and the ficha técnica are
built here in the shapes the CIS publishes them in -- the PDFs as the text
pypdf reads from them, broken lines and split figures included.
"""

import io
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from scripts.fetch_census import spain_cis as s  # noqa: E402

ANSWERS_ES = ["Católico/a practicante", "Católico/a no practicante", "Creyente de otra religión",
              "Agnóstico/a (no niegan la existencia de Dios pero tampoco la descartan)",
              "Indiferente, no creyente", "Ateo/a (niegan la existencia de Dios)", "N.C."]

FICHA_ROWS = [
    [None, "3545/0 PREELECTORAL ELECCIONES AUTONÓMICAS 2026. COMUNIDAD AUTÓNOMA DE CASTILLA Y LEÓN"],
    [None, "FICHA TÉCNICA DEL ESTUDIO"],
    ["FICHA TÉCNICA"],
    ["Ámbito", "Castilla y León   (aut.)."],
    ["Universo", "Población española con derecho a voto en elecciones autonómicas y residente "
                 "en la comunidad autónoma de ambos sexos de 18 y más años."],
    ["Tamaño de la muestra:"],
    ["Diseñada", "8.000 entrevistas."],
    ["Realizada", "{total} entrevistas."],
    ["Afijación", "No proporcional"],
    [None, "Partiendo de una distribución proporcional por provincias, ..._x000D_"],
    ["Ponderación", "Para tratar la muestra en su conjunto es necesario aplicar coeficientes "
                    "de ponderación (PESO)._x000D_ Para la estimación a nivel de cada provincia, "
                    "en el fichero de microdatos se incluye la ponderación para cada una de ellas "
                    "(variable PESOPROV)."],
    ["Procedimiento de muestreo", "Se ha procedido a la selección aleatoria de teléfonos fijos y "
                                  "móviles. La selección de las personas se ha llevado a cabo "
                                  "mediante la aplicación de cuotas de sexo y edad._x000D_ Los "
                                  "cuestionarios se han aplicado mediante entrevista "
                                  "telefónica asistida por ordenador (CATI)."],
    ["Fecha de realización", "Del 6 al 13 de febrero de 2026."],
]

# Castilla y León's nine provinces: code, name, realised sample, mean weight.
CYL = [("05", "Ávila", 587, 0.9766), ("09", "Burgos", 1044, 1.0058), ("24", "León", 1377, 1.1183),
       ("34", "Palencia", 616, 0.9092), ("37", "Salamanca", 1012, 1.1053),
       ("40", "Segovia", 612, 0.8421), ("42", "Soria", 542, 0.5217),
       ("47", "Valladolid", 1637, 1.0439), ("49", "Zamora", 612, 1.1279)]


def province_shares(i):
    """Seven shares summing to 100, different for each province."""
    base = [26.0 + i, 36.0 - i, 2.0, 9.5, 10.5, 13.0, 3.0]
    return base


def region_of(provinces):
    mass = [n * w for _, _, n, w in provinces]
    whole = sum(mass)
    shares = [province_shares(i) for i in range(len(provinces))]
    return [sum(m * sh[k] for m, sh in zip(mass, shares)) / whole for k in range(7)]


def marginal_rows(title, shares, n):
    rows = [[None, "3545/0 PREELECTORAL"], [None, title], [None, "INFORME DE MARGINALES"],
            ["Pregunta A1"], ["Sexo:"], ["Hombre", 48.9], ["Mujer", 51.1], ["(N)", n], [None],
            ["Pregunta C6"], [s.QUESTION_ES]]
    rows += [[label, share] for label, share in zip(ANSWERS_ES, shares)]
    rows += [["(N)", n], [None], ["Pregunta C6a"],
             ["SÓLO A QUIENES SE DEFINEN EN MATERIA RELIGIOSA COMO CATÓLICOS/AS"],
             ["Casi nunca", 40.0], ["(N)", n - 10]]
    return rows


def workbook(provinces=CYL, region_shares=None, header="3545/0 PREELECTORAL ELECCIONES",
             total=None, drop=None, copy=None, sample=None, ficha_rows=FICHA_ROWS):
    """A study workbook; ``copy`` = (province, source) writes source's sheet under province's
    name, and ``sample`` = (province, n) gives a province's sheet another (N)."""
    import openpyxl
    book = openpyxl.Workbook()
    first = book.active
    first.title = "Contenidos"
    first.append([None, header])
    total = total if total is not None else sum(n for _, _, n, _ in provinces)
    ficha = book.create_sheet("Ficha técnica")
    for row in ficha_rows:
        ficha.append([c.format(total=f"{total:,}".replace(",", ".")) if isinstance(c, str) else c
                      for c in row])
    table = book.create_sheet("Tabla ponderaciones")
    table.append([None, "Muestra diseñada, realizada, coeficientes de ponderación y error (%)"])
    table.append([None, None, "Muestra"])
    table.append(["cpr", "Provincia", "Diseñada", "Realizada", "Ponderación", "Error (%)"])
    for code, name, n, w in provinces:
        table.append([code, name, 600, n, w, 4.0])
    table.append([None, "Total", 8000, total, None, 1.1])
    region = region_shares or region_of(provinces)
    book.create_sheet("Resultados_Castilla y León")
    for row in marginal_rows("Total", region, sum(n for _, _, n, _ in provinces)):
        book["Resultados_Castilla y León"].append(row)
    index = {name: i for i, (_, name, _, _) in enumerate(provinces)}
    for i, (code, name, n, _) in enumerate(provinces):
        if name == drop:
            continue
        source = copy[1] if copy and copy[0] == name else name
        j = index[source]
        n_shown = provinces[j][2]
        if sample and sample[0] == name:
            n_shown = sample[1]
        sheet = book.create_sheet(f"Resultados_{name}")
        for row in marginal_rows(f"Provincia de {name}", province_shares(j), n_shown):
            sheet.append(row)
    book.create_sheet("Sexo")
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def pdf(study, unit, shares, n, part=1):
    """PDF marginals as pypdf reads them: a cover, a page break inside the table, its foot."""
    printed = [f"{x:.1f}".replace(".", ",") for x in shares]
    lines = ["PREELECTORAL ELECCIONES", "AUTONÓMICAS 2026. COMUNIDAD", f"{unit}",
             f"Estudio nº {study}_{part}", "Febrero 2026", "Pregunta A1", "Sexo:", "Hombre 48,9",
             "(N) (1.000)", "CIS", f"Estudio nº{study}_{part}. PREELECTORAL", "Pág 3",
             "Pregunta C6",
             "¿Cómo se define Ud. en materia religiosa: católico/a practicante, católico/a no "
             "practicante, creyente de otra religión,",
             "agnóstico/a, indiferente o no creyente, o ateo/a?"]
    answers = [f"{a} {p}" for a, p in zip(ANSWERS_ES, printed)]
    lines += answers[:3] + ["CIS", f"Estudio nº{study}_{part}. PREELECTORAL ELECCIONES",
                            "AUTONÓMICAS 2026. COMUNIDAD", "Febrero 2026", "Pág 12"]
    lines += answers[3:] + [f"(N) ({n:,})".replace(",", "."),
                            "SÓLO A QUIENES SE DEFINEN EN MATERIA RELIGIOSA", "Casi nunca 40,0"]
    return "\n".join(lines)


STUDY = s.Study(3545, "Preelectoral elecciones autonómicas 2026. Comunidad autónoma de Castilla y León",
                "the 2026 election to the Cortes of Castile and León", "07",
                "https://www.cis.es/documents/20117/13765729/3545-multi.xlsx", "https://www.cis.es/x")

# Study 3286's ficha técnica as pypdf reads it, a few characters at a time.
FICHA_3286 = """ESTUDIO CIS nº 3
286
PREELECTORAL DEL PAÍS VA
SCO. ELECCIONES AUTONÓMICAS JULIO
2020
FICHA TÉCNICA
Ámbito
:
Comunidad
a
utónoma del País Vasco.
Universo
:
Población con derecho a voto en elecciones autonómic
as y residente en la c
omunidad a
utónoma del
País
Vasco.
Tamaño de la muestra
:
Diseñada:
3.
450
entrevistas, con la si
guiente distribución provincial:
Araba/Álava
1.
15
0
Bizkaia
1.
15
0
Gipu
zkoa
1.
15
0
Realizada:
3.
{total}
entrevistas, con la siguiente distribución provincial
:
Araba/Álava
1
.
{alava}
Bizkaia
1
.
125
Gip
u
zkoa
1
.
127
Afijación
:
No proporcional
.
Ponderación
:
Para tratar la muestra
para el conjunto de la comunidad autónoma es necesaria la aplicación de los
siguientes coeficientes de ponderación (PESO)
:
Araba/Álava
0
,
446
Bizkaia
1
,
579
Gipuzkoa
0
,
964
Puntos de muestreo
:
191
municipios y 3 provincias.
Procedimiento de muestreo
:
Se ha procedido a la selección aleatoria de teléfonos fijos y móviles con un porcentaje del
66
,
6
% y del 33,4
%,
respectivamente. La selección de los individuos se ha llevado a cabo mediante la aplicación de
cuotas de sexo y edad.
Los
cuestionarios se han aplicado mediante entrevista telefónic
a asistida por ordenador (CATI).
Fecha de realización
:
Del 1
0 al 1
6 de junio de 20
20."""

# Study 3263's ficha técnica as pypdf reads it, Cuadro 1 a figure to a line.
FICHA_3263 = """ESTUDIO CIS nº
3
2
63
MACROBARÓMETRO DE
OCTUBRE
2019. P
REELECTORAL ELECCIONES GENERALES 20
1
9
FICHA TÉCNICA
Ámbito:
Nacional
.
Universo:
Población
con derecho a voto en elecciones generales y residente en España
.
Tamaño de la muestra:
Diseñada:
1
8
.
00
0
entrevistas.
Realizada:
1
7
.
650
entrevistas.
Ponderación:
Para tratar la muestra en su conjunto es necesaria la aplicación de los
coeficientes de ponderación que figuran en el
Cuadro 1, al final de esta ficha técnica.
Procedimiento de muestreo:
Polietápico, estratificado por conglomerados, y de las unidades últ
imas (individuos) por rutas
aleatorias y cuotas de sexo y edad.
Los cuestionarios se han aplicado mediante entrevista personal en los domicilios.
Fecha
de realización:
Del
2
1
de septiembre
al
1
3
de
octubre
de 20
1
9
.
Cuadro 1.
-
Puntos de muestreo, entrevistas y coeficientes de ponderación
17
Rioja (La)
26
Rioja (La)
15
15
300
{rioja}
5,9
0,4166
18
Ceuta
(Ciudad Autónoma de)
51
Ceuta
1
1
60
60
12,9
0,5047
Total
1.093
1.0
85
18.000
17.650
0,75"""

RIOJA = [34.5, 40.1, 1.1, 6.3, 1.1, 15.5, 1.4]

BASQUE = {"01": ("ARABA/ÁLAVA", 1102, 0.446, [15.7, 35.4, 1.5, 14.2, 13.7, 18.1, 1.4]),
          "48": ("BIZKAIA", 1125, 1.579, [15.4, 39.4, 1.2, 12.3, 12.6, 18.0, 1.1]),
          "20": ("GIPUZKOA", 1127, 0.964, [16.1, 37.7, 1.3, 12.4, 13.5, 16.8, 2.2])}


def basque_region():
    mass = {c: n * w for c, (_, n, w, _) in BASQUE.items()}
    whole = sum(mass.values())
    return [sum(mass[c] * BASQUE[c][3][k] for c in BASQUE) / whole for k in range(7)]


def basque_texts(alava_n=1102, region=None):
    total = 1102 + 1125 + 1127
    study = s.PROVINCE_PDFS[0]
    texts = {c: pdf(3286, name, sh, alava_n if c == "01" else n, part=i)
             for i, (c, (name, n, _, sh)) in enumerate(sorted(BASQUE.items()), 1)}
    region_text = pdf(3286, "PREELECTORAL DEL PAÍS VASCO.", region or basque_region(), total)
    ficha = FICHA_3286.format(total=str(total)[1:], alava="102")
    return study, region_text, ficha, texts


class TestAnswers(unittest.TestCase):
    def test_every_answer_the_cis_offers(self):
        self.assertEqual([s.answer_label(a) for a in ANSWERS_ES],
                         ["Practising Catholic", "Non-practising Catholic",
                          "Believer in another religion", "Agnostic",
                          "Indifferent / non-believer", "Atheist", "Not stated"])
        # The HTML marginals put a zero-width space after "Católico/".
        self.assertEqual(s.answer_label("Católico/&#8203;a no practicante"),
                         "Non-practising Catholic")
        self.assertEqual(s.answer_label("N.S."), "Not stated")

    def test_an_answer_it_does_not_offer_stops(self):
        with self.assertRaises(SystemExit):
            s.answer_label("Musulmán/a")

    def test_numbers(self):
        self.assertEqual(s.number("15,2"), 15.2)
        self.assertEqual(s.number("(4.998)"), 4998)
        self.assertEqual(s.number("(998)"), 998)
        self.assertEqual(s.number(26.4), 26.4)
        self.assertEqual(s.number("-"), 0.0)


class TestText(unittest.TestCase):
    def test_figures_pypdf_splits_are_joined(self):
        self.assertEqual(s.flatten("Realizada: 3.\n354\nentrevistas"), "Realizada: 3.354 entrevistas")
        self.assertEqual(s.flatten("Del 1\n9 al 28 de marzo de 20\n2\n1."),
                         "Del 19 al 28 de marzo de 2021.")
        self.assertEqual(s.flatten("Araba/Álava 0 , 446 Bizkaia 1 , 579"),
                         "Araba/Álava 0,446 Bizkaia 1,579")
        self.assertEqual(s.squeeze("Gip u zkoa"), "gipuzkoa")

    def test_the_question_as_printed(self):
        self.assertEqual(s.tidy_question("Pregunta C6 &iquest;Cómo se define Ud. en materia\n"
                                         "religiosa: católico/&#8203;a practicante, o ateo/a? x"),
                         "¿Cómo se define Ud. en materia religiosa: católico/a practicante, o "
                         "ateo/a?")


class TestReading(unittest.TestCase):
    def test_religion_block(self):
        shares, n, question = s.religion_block(
            marginal_rows("Provincia de Ávila", province_shares(0), 587), "test")
        self.assertEqual(n, 587)
        self.assertEqual(shares["Practising Catholic"], 26.0)
        self.assertEqual(shares["Not stated"], 3.0)
        self.assertEqual(len(shares), 7)
        self.assertEqual(question, s.QUESTION_ES)

    def test_a_missing_answer_or_a_bad_sum_stops(self):
        rows = marginal_rows("x", province_shares(0), 587)
        bad = [r for r in rows if not (r and r[0] == "Ateo/a (niegan la existencia de Dios)")]
        with self.assertRaises(SystemExit):
            s.religion_block(bad, "test")
        skew = [[r[0], 30.0] if r and r[0] == "N.C." else r for r in rows]
        with self.assertRaises(SystemExit):
            s.religion_block(skew, "test")

    def test_no_question_stops(self):
        with self.assertRaises(SystemExit):
            s.religion_block([["Pregunta A1"], ["Sexo:"], ["(N)", 5]], "test")

    def test_weights_table(self):
        rows = [[None, "x"], ["cpr", "Provincia", "Diseñada", "Realizada", "Ponderación", "Error (%)"],
                ["05", "Ávila", 600, 587, 0.9766, 4.1], [9, "Burgos", 1039, 1044, 1.0058, 3.1],
                [None, "Total", 1639, 1631, None, 1.1]]
        table, total = s.weights_table(rows, "test")
        self.assertEqual(total, 1631)
        self.assertEqual(table["09"]["realised"], 1044)
        self.assertEqual(table["05"]["weight"], 0.9766)
        rows[-1][3] = 1700
        with self.assertRaises(SystemExit):
            s.weights_table(rows, "test")

    def test_ficha(self):
        text = " ".join(f"{c}:" if j == 0 and isinstance(c, str) else str(c)
                        for row in FICHA_ROWS for j, c in enumerate(row) if c).format(total="8.039")
        got = s.ficha(text, "test")
        self.assertEqual(got["realised"], 8039)
        self.assertEqual(got["year"], 2026)
        self.assertEqual(got["fieldwork"], "Del 6 al 13 de febrero de 2026")
        self.assertTrue(got["voters"] and got["weighted"])
        self.assertTrue(got["province_weight"])
        self.assertEqual(got["mode"], "telephone")
        self.assertTrue(got["quotas"])
        self.assertEqual(s.method_of(got), "telephone interviews (CATI) with quotas of sex and age, ")
        # The quotas are said only where the ficha says them.
        bare = s.ficha(text.replace("cuotas de sexo y edad", "rutas aleatorias"), "test")
        self.assertFalse(bare["quotas"])
        self.assertEqual(s.method_of(bare), "telephone interviews (CATI), ")
        # A barometer's universe names no election, and it has no provincial weight.
        barometer = ("Ámbito: Nacional. Universo: Población española de ambos sexos de 18 años "
                     "y más. Tamaño de la muestra: Diseñada: 4.000 entrevistas. Realizada: 4.042 "
                     "entrevistas. Fecha de realización: Del 1 al 4 de septiembre de 2026.")
        plain = s.ficha(barometer, "test")
        self.assertFalse(plain["voters"] or plain["province_weight"] or plain["weighted"])
        self.assertIsNone(plain["mode"])
        self.assertEqual(s.method_of(plain), "")
        with self.assertRaises(SystemExit):
            s.ficha(barometer.replace("española", "residente"), "test")
        # A municipal electorate takes in other EU citizens: not the note's universe.
        with self.assertRaises(SystemExit):
            s.ficha(barometer.replace("española de ambos sexos de 18 años y más",
                                      "con derecho a voto en elecciones locales"), "test")

    def test_a_ficha_pypdf_reads_in_pieces(self):
        got = s.ficha(FICHA_3286.format(total="354", alava="102"), "test")
        self.assertEqual(got["realised"], 3354)
        self.assertEqual(got["fieldwork"], "Del 10 al 16 de junio de 2020")
        self.assertEqual(s.english_dates(got["fieldwork"]), "10-16 June 2020")
        self.assertEqual(got["year"], 2020)
        self.assertTrue(got["voters"] and got["weighted"] and got["quotas"])
        self.assertFalse(got["province_weight"] or got["self_weighting"])
        self.assertEqual(got["mode"], "telephone")
        provinces = s.ficha_provinces(FICHA_3286.format(total="354", alava="102"),
                                      ["01", "20", "48"], "test")
        self.assertEqual(provinces, {"01": {"realised": 1102, "weight": 0.446},
                                     "20": {"realised": 1127, "weight": 0.964},
                                     "48": {"realised": 1125, "weight": 1.579}})

    def test_a_self_weighting_sample(self):
        madrid = ("Universo: Población con derecho a voto en elecciones autonómicas y residentes "
                  "en la Comunidad de Madrid. Tamaño de la muestra: Diseñada: 4.300 entrevistas. "
                  "Realizada: 4\n.\n124 entrevistas. Afijación: Proporcional. Ponderación: No "
                  "procede. Fecha de realización: Del 1\n9 al 28 de marzo de 20\n2\n1.")
        got = s.ficha(madrid, "test")
        self.assertTrue(got["self_weighting"] and got["voters"])
        self.assertEqual(got["realised"], 4124)
        self.assertIn("self-weighting", s.weighting_of(got, "comunidad"))

    def test_dates_and_sheet_names(self):
        self.assertEqual(s.english_dates("Del 6 al 13 de febrero de 2026."), "6-13 February 2026")
        self.assertEqual(s.english_dates("Del 28 de abril al 5 de mayo de 2026"),
                         "28 April - 5 May 2026")
        self.assertEqual(s.sheet_unit("Resultados_Ávila"), "avila")
        self.assertEqual(s.sheet_unit("Resultados Extremadura "), "extremadura")
        self.assertEqual(s.sheet_unit("Resultados_provincia de Soria"), "soria")
        self.assertEqual(s.sheet_unit("Almería"), "almeria")
        self.assertTrue(s.names_study("Estudio nº3437. PREELECTORAL", 3437))
        self.assertTrue(s.names_study("Estudio nº 3538_2", 3538))
        self.assertFalse(s.names_study("34370", 3437))


class TestPdfMarginals(unittest.TestCase):
    def test_reads_the_table_across_a_page_break(self):
        text = pdf(3538, "Provincia de Cáceres", [27.6, 40.4, 0.9, 7.4, 9.6, 11.9, 2.2], 822, 2)
        shares, n, question = s.pdf_marginals(text, s.WORKBOOKS[3], ("Cáceres",), "test")
        self.assertEqual(n, 822)
        self.assertEqual(shares["Non-practising Catholic"], 40.4)
        self.assertEqual(shares["Atheist"], 11.9)
        self.assertEqual(len(shares), 7)
        self.assertTrue(question.startswith("¿Cómo se define Ud. en materia religiosa"))
        self.assertTrue(question.endswith("o ateo/a?"))

    def test_another_studys_or_places_pdf_stops(self):
        text = pdf(3538, "Provincia de Cáceres", [27.6, 40.4, 0.9, 7.4, 9.6, 11.9, 2.2], 822, 2)
        with self.assertRaises(SystemExit):
            s.pdf_marginals(text, s.WORKBOOKS[3], ("Badajoz",), "test")
        with self.assertRaises(SystemExit):
            s.pdf_marginals(text.replace("3538", "3537"), s.WORKBOOKS[3], ("Cáceres",), "test")


class TestWorkbook(unittest.TestCase):
    def test_reads_every_province_and_the_comunidad(self):
        got = s.read_workbook(workbook(), STUDY)
        self.assertEqual(sorted(got["provinces"]), [c for c, *_ in CYL])
        self.assertEqual(got["provinces"]["42"]["n"], 542)
        self.assertEqual(got["region"]["n"], sum(n for _, _, n, _ in CYL))
        self.assertEqual(got["design"]["year"], 2026)
        self.assertEqual(got["region"]["question"], s.QUESTION_ES)

    def test_another_studys_workbook_stops(self):
        with self.assertRaises(SystemExit):
            s.read_workbook(workbook(header="3453/0 PREELECTORAL DE CATALUÑA"), STUDY)

    def test_a_comunidad_its_provinces_do_not_make_stops(self):
        # Separate weights put the comunidad a few tenths away: that passes.
        region = region_of(CYL)
        region[0] += 0.6
        region[1] -= 0.6
        s.read_workbook(workbook(region_shares=region), STUDY)
        # A point and a half is not weighting.
        region[0] += 0.9
        region[1] -= 0.9
        with self.assertRaises(SystemExit):
            s.read_workbook(workbook(region_shares=region), STUDY)

    def test_a_missing_province_or_a_wrong_total_stops(self):
        with self.assertRaises(SystemExit):
            s.read_workbook(workbook(drop="Soria"), STUDY)
        with self.assertRaises(SystemExit):
            s.read_workbook(workbook(total=8040), STUDY)

    def test_a_province_answered_by_other_than_its_sample_stops(self):
        with self.assertRaises(SystemExit):
            s.read_workbook(workbook(sample=("Soria", 600)), STUDY)

    def test_a_sheet_that_is_another_provinces_copy_is_withheld(self):
        # Study 3538: the Cáceres sheet is Badajoz's, figure for figure.
        got = s.read_workbook(workbook(copy=("Soria", "Burgos")), STUDY)
        self.assertNotIn("42", got["provinces"])
        self.assertIn("repeats Burgos's", got["withheld"]["42"])
        self.assertIn("realised sample is 542", got["withheld"]["42"])
        self.assertEqual(got["provinces"]["09"]["n"], 1044)
        self.assertEqual(len(got["provinces"]), 8)

    def test_two_provinces_with_one_set_of_shares_stop(self):
        # The same copy, but with the province's own sample: nothing tells
        # which sheet is whose, so nothing is written.
        with self.assertRaises(SystemExit):
            s.read_workbook(workbook(copy=("Soria", "Burgos"), sample=("Soria", 542)), STUDY)


class TestCopiedSheetFromItsPdf(unittest.TestCase):
    """Study 3538: Cáceres's sheet is Badajoz's, and its PDF marginals are its own."""

    def pdfs(self, soria_n=542, soria_shares=None, burgos_shares=None):
        index = {name: i for i, (_, name, _, _) in enumerate(CYL)}
        region = [round(x, 1) for x in region_of(CYL)]
        return {"42": pdf(3545, "Provincia de Soria", soria_shares or province_shares(index["Soria"]),
                          soria_n, 7),
                "09": pdf(3545, "Provincia de Burgos",
                          burgos_shares or province_shares(index["Burgos"]), 1044, 2),
                "CA": pdf(3545, "CASTILLA Y LEÓN", region, sum(n for _, _, n, _ in CYL), 0)}

    def test_the_province_is_read_from_its_pdf(self):
        got = s.read_workbook(workbook(copy=("Soria", "Burgos")), STUDY, self.pdfs())
        self.assertEqual(got["withheld"], {})
        soria = got["provinces"]["42"]
        self.assertTrue(soria["pdf"])
        self.assertEqual(soria["n"], 542)
        self.assertEqual(soria["shares"]["Practising Catholic"], 26.0 + 6)
        self.assertIn("repeats Burgos's", soria["via"])
        self.assertEqual(len(got["provinces"]), 9)

    def test_a_pdf_without_the_realised_sample_or_repeating_another_stops(self):
        with self.assertRaises(SystemExit):
            s.read_workbook(workbook(copy=("Soria", "Burgos")), STUDY, self.pdfs(soria_n=1044))
        with self.assertRaises(SystemExit):
            s.read_workbook(workbook(copy=("Soria", "Burgos")), STUDY,
                            self.pdfs(soria_shares=province_shares(1)))

    def test_a_pdf_disagreeing_with_its_sheet_stops(self):
        with self.assertRaises(SystemExit):
            s.read_workbook(workbook(copy=("Soria", "Burgos")), STUDY,
                            self.pdfs(burgos_shares=[28.0, 34.0, 2.0, 9.5, 10.5, 13.0, 3.0]))


class TestPdfStudy(unittest.TestCase):
    """Study 3286: the Basque Country's July 2020 survey, a PDF for each province."""

    def test_reads_the_provinces(self):
        study, region, ficha, texts = basque_texts()
        got = s.read_pdf_study(region, ficha, texts, study)
        self.assertEqual(sorted(got["provinces"]), ["01", "20", "48"])
        self.assertEqual(got["provinces"]["01"]["n"], 1102)
        self.assertEqual(got["provinces"]["48"]["weight"], 1.579)
        self.assertEqual(got["provinces"]["20"]["shares"]["Not stated"], 2.2)
        self.assertEqual(got["region"]["n"], 3354)
        self.assertEqual(got["design"]["year"], 2020)

    def test_a_province_off_its_sample_or_a_comunidad_off_its_provinces_stops(self):
        study, region, ficha, texts = basque_texts(alava_n=1101)
        with self.assertRaises(SystemExit):
            s.read_pdf_study(region, ficha, texts, study)
        skewed = basque_region()
        skewed[0] += 1.5
        skewed[1] -= 1.5
        study, region, ficha, texts = basque_texts(region=skewed)
        with self.assertRaises(SystemExit):
            s.read_pdf_study(region, ficha, texts, study)


class TestArchived(unittest.TestCase):
    """Study 3263: a macro-survey's comunidad PDFs, as the Internet Archive captured them."""

    def test_cuadro_and_capture(self):
        ficha = FICHA_3263.format(rioja="284")
        self.assertEqual(s.cuadro_realised(ficha, "26", "test"), 284)
        self.assertEqual(s.cuadro_realised(ficha, "51", "test"), 60)
        with self.assertRaises(SystemExit):
            s.cuadro_realised(ficha, "52", "test")            # Melilla is not in this one
        self.assertEqual(s.archive_date(s.ARCHIVED[0].url), "4 February 2020")
        design = s.ficha(ficha, "test")
        self.assertEqual(design["realised"], 17650)
        self.assertEqual(s.english_dates(design["fieldwork"]), "21 September - 13 October 2019")
        self.assertEqual(design["mode"], "face")
        self.assertTrue(design["quotas"] and design["voters"])

    def test_reads_la_rioja(self):
        rioja = s.ARCHIVED[0]
        got = s.read_archived(pdf(3263, "LA RIOJA", RIOJA, 284, part=1700),
                              FICHA_3263.format(rioja="284"), rioja)
        self.assertEqual(got["region"]["n"], 284)
        self.assertEqual(got["region"]["shares"]["Atheist"], 15.5)
        self.assertEqual(got["archived"], "4 February 2020")
        with self.assertRaises(SystemExit):                    # the PDF is not Cuadro 1's sample
            s.read_archived(pdf(3263, "LA RIOJA", RIOJA, 284, part=1700),
                            FICHA_3263.format(rioja="285"), rioja)
        with self.assertRaises(SystemExit):                    # another comunidad's PDF
            s.read_archived(pdf(3263, "CIUDAD AUTÓNOMA DE CEUTA", RIOJA, 284, part=1800),
                            FICHA_3263.format(rioja="284"), rioja)


HTML = """<html><head><title>3448 Preelectoral del País Vasco</title></head><body>
<p><b>Pregunta C5 </b></p><table><tr><td>Primaria</td><td>9,3</td></tr>
<tr><td>(N)</td><td>(4.998)</td></tr></table>
<p style='x'><b>Pregunta C6 </b><br><p>&iquest;Cómo se define Ud. en materia religiosa:
católico/a practicante, católico/a no practicante, creyente de otra religión, agnóstico/a,
indiferente o no creyente, o ateo/a?</p> <p></p><br></p><table border=0>
<col width=432.0><tr><td style='a'>Católico/&#8203;a practicante</td><td>15,2</td></tr>
<tr><td>Católico/&#8203;a no practicante</td><td>31,8</td></tr>
<tr><td>Creyente de otra religión</td><td>2,6</td></tr>
<tr><td>Agnóstico/&#8203;a (no niegan la existencia de Dios pero tampoco la descartan)</td><td>12,0</td></tr>
<tr><td>Indiferente, no creyente</td><td>16,3</td></tr>
<tr><td>Ateo/&#8203;a (niegan la existencia de Dios)</td><td>20,9</td></tr>
<tr><td>N.C.</td><td>1,2</td></tr>
<tr><td style='b'>(N)</td><td>(4.998)</td></tr></table></body></html>"""


class TestHtmlTotal(unittest.TestCase):
    def test_reads_the_comunidads_marginals(self):
        study = s.Study(3448, "t", "e", "16", "u", "p", "f")
        shares, n, question = s.read_html_total(HTML, study)
        self.assertEqual(n, 4998)
        self.assertEqual(shares["Atheist"], 20.9)
        self.assertAlmostEqual(sum(shares.values()), 100.0)
        self.assertEqual(question, s.QUESTION_ES)

    def test_marginals_of_another_study_stop(self):
        with self.assertRaises(SystemExit):
            s.read_html_total(HTML.replace("3448", "3437"), s.Study(3448, "t", "e", "16", "u", "p"))


class TestBinding(unittest.TestCase):
    def test_the_map_binds_by_code(self):
        provinces, comunidades = s.bind(*s.load_units())
        self.assertEqual(len(provinces), 52)
        self.assertEqual(len(comunidades), 19)
        self.assertEqual(provinces["35"]["name"], "Palmas, Las")
        self.assertEqual(provinces["42"]["name"], "Soria")
        self.assertEqual(comunidades["07"]["name"], "Castilla y León")
        self.assertEqual(provinces["42"]["parent"], comunidades["07"]["id"])

    def test_a_code_on_the_wrong_label_stops(self):
        admin1, admin2 = s.load_units()
        moved = [dict(u, name="Segovia") if u["name"] == "Soria" else u for u in admin2]
        with self.assertRaises(SystemExit):
            s.bind(admin1, moved)


class TestRecords(unittest.TestCase):
    def setUp(self):
        self.provinces, self.comunidades = s.bind(*s.load_units())

    def test_records(self):
        got = s.read_workbook(workbook(), STUDY)
        got["provinces"]["42"]["n"] = 250                    # low precision
        got["provinces"]["40"]["n"] = 90                     # under the minimum
        records, skipped = s.build([(STUDY, got)], self.provinces, self.comunidades)
        self.assertEqual(len(records), 9)                   # the comunidad and 8 provinces
        self.assertEqual(skipped, ["Segovia: 90 respondents (study 3545)"])
        by = {r["name"]: r for r in records}
        soria = by["Soria"]
        self.assertEqual(soria["level"], "admin2")
        self.assertEqual(soria["shape_id"], self.provinces["42"]["id"])
        self.assertTrue(soria["religion_basis"].startswith("survey estimate"))
        self.assertIn("Low precision", soria["religion_note"])
        self.assertIn("250 respondents in the province of Soria", soria["religion_note"])
        self.assertTrue(all("count" not in row for row in soria["religion"]))
        pcts = [row["pct"] for row in soria["religion"]]
        self.assertEqual(pcts, sorted(pcts, reverse=True))
        self.assertEqual(soria["religion_year"], 2026)
        self.assertIn("PESOPROV", soria["religion_note"])
        self.assertIn("telephone interviews (CATI) with quotas of sex and age", soria["religion_note"])
        self.assertIn(f"\"{s.QUESTION_ES}\"", soria["religion_note"])
        # Since when the census has not asked, and what the wording does.
        self.assertIn("Constitution of 1978", soria["religion_note"])
        self.assertIn("European Social Survey", soria["religion_note"])
        region = by["Castilla y León"]
        self.assertEqual(region["level"], "admin1")
        self.assertEqual(region["shape_id"], self.comunidades["07"]["id"])
        self.assertNotIn("Low precision", region["religion_note"])
        self.assertIn("(PESO)", region["religion_note"])
        self.assertEqual(region["sources"][0]["field"], "religion")
        self.assertEqual(region["sources"][0]["url"], STUDY.url)

    def test_a_province_read_from_its_pdf_cites_it(self):
        pdfs = TestCopiedSheetFromItsPdf().pdfs()
        got = s.read_workbook(workbook(copy=("Soria", "Burgos")), STUDY, pdfs)
        study = s.Study(**{**STUDY.__dict__, "pdfs": (("42", "https://x/es3545mar_07soria-pdf"),)})
        records, skipped = s.build([(study, got)], self.provinces, self.comunidades)
        self.assertEqual(skipped, [])
        soria = next(r for r in records if r["name"] == "Soria")
        self.assertEqual(soria["sources"][0]["url"], "https://x/es3545mar_07soria-pdf")
        self.assertIn("PDF marginals", soria["sources"][0]["name"])
        self.assertIn("repeats Burgos's", soria["religion_note"])
        burgos = next(r for r in records if r["name"] == "Burgos")
        self.assertEqual(burgos["sources"][0]["url"], STUDY.url)

    def test_a_comunidad_of_one_province_writes_both(self):
        madrid = s.TOTALS[2]
        design = s.ficha("Universo: Población con derecho a voto en elecciones autonómicas y "
                         "residentes en la Comunidad de Madrid. Tamaño de la muestra: Realizada: "
                         "4.124 entrevistas. Ponderación: No procede. Fecha de realización: Del "
                         "19 al 28 de marzo de 2021.", "test")
        shares = dict(zip(s.SUBSTANTIVE + ("Not stated",), [18.2, 38.5, 2.9, 12.5, 11.2, 15.0, 1.7]))
        got = {"design": design, "region": {"shares": shares, "n": 4124,
                                            "question": s.QUESTION_ES}}
        records, _ = s.build([(madrid, got)], self.provinces, self.comunidades)
        self.assertEqual(sorted(r["level"] for r in records), ["admin1", "admin2"])
        province = next(r for r in records if r["level"] == "admin2")
        region = next(r for r in records if r["level"] == "admin1")
        self.assertEqual(province["shape_id"], self.provinces["28"]["id"])
        self.assertEqual(region["shape_id"], self.comunidades["13"]["id"])
        self.assertEqual(province["religion"], region["religion"])
        self.assertIn("is the whole of", province["religion_note"])
        self.assertIn("self-weighting", region["religion_note"])
        self.assertEqual(region["religion_year"], 2021)

    def test_an_archived_comunidad_of_one_province(self):
        rioja, ceuta = s.ARCHIVED[0], s.ARCHIVED[1]
        ficha = FICHA_3263.format(rioja="284")
        got = s.read_archived(pdf(3263, "LA RIOJA", RIOJA, 284, part=1700), ficha, rioja)
        small = s.read_archived(pdf(3263, "CIUDAD AUTÓNOMA DE CEUTA", RIOJA, 60, part=1800),
                                ficha, ceuta)
        records, skipped = s.build([(rioja, got), (ceuta, small)], self.provinces,
                                   self.comunidades)
        self.assertEqual(sorted((r["level"], r["shape_id"]) for r in records),
                         sorted([("admin1", self.comunidades["17"]["id"]),
                                 ("admin2", self.provinces["26"]["id"])]))
        self.assertEqual(skipped, [f"{self.comunidades['18']['name']} and the province of "
                                   f"{self.provinces['51']['name']}: 60 respondents (study 3263)"])
        note = records[0]["religion_note"]
        self.assertIn("captured by the Internet Archive on 4 February 2020", note)
        self.assertIn("entitled to vote in general elections", note)
        self.assertIn("face-to-face interviews at home with quotas of sex and age", note)
        self.assertIn("21 September - 13 October 2019", note)
        self.assertIn("Low precision", note)
        self.assertEqual(records[0]["religion_year"], 2019)
        self.assertEqual(records[0]["sources"][0]["url"], rioja.url)
        self.assertIn("PDF marginals", records[0]["sources"][0]["name"])
        self.assertTrue(records[0]["religion_basis"].endswith("entitled to vote in general "
                                                              "elections"))
        self.assertIn("is the whole of", records[1]["religion_note"])

    def test_a_pdf_study_writes_its_provinces_only(self):
        study, region, ficha, texts = basque_texts()
        got = s.read_pdf_study(region, ficha, texts, study)
        records, skipped = s.build([(study, got)], self.provinces, self.comunidades)
        self.assertEqual(skipped, [])
        self.assertEqual(sorted(r["codes"]["ine_province"] for r in records), ["01", "20", "48"])
        self.assertTrue(all(r["level"] == "admin2" for r in records))
        alava = next(r for r in records if r["codes"]["ine_province"] == "01")
        self.assertEqual(alava["shape_id"], self.provinces["01"]["id"])
        self.assertEqual(alava["sources"][0]["url"], dict(study.pdfs)["01"])
        self.assertIn("unweighted", alava["religion_note"])
        self.assertIn("newest by province", alava["religion_note"])
        self.assertIn("10-16 June 2020", alava["religion_note"])
        self.assertEqual(alava["religion_year"], 2020)


if __name__ == "__main__":
    unittest.main()
