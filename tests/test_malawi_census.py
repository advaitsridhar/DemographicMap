"""Malawi's 2018 census by district and region, from the National Statistical
Office's Series A and E workbooks (scripts/fetch_census/malawi_census)."""

import io
import unittest
from unittest import mock

import openpyxl

from scripts.fetch_census import malawi_census as mw

# region: [(district or city, men, women)]
AREAS = {"Northern": [("Mzimba", 450, 490), ("Likoma", 7, 8), ("Mzuzu City", 110, 112)],
         "Central": [("Nkhotakota", 190, 200)],
         "Southern": [("Zomba", 350, 390), ("Zomba City", 51, 53)]}
AGE_LABELS = ["Less than 1 Year", "1-4"] + [f"{a}-{a + 4}" for a in range(5, 90, 5)] + ["90+"]
SHARES = [3, 12, 15, 14, 12, 10, 8, 6, 5, 4, 3, 2, 2, 1, 1, 1, 1, 0, 0, 0]  # of every 100


def ages(total):
    """Twenty age groups making ``total``."""
    out = [total * s // 100 for s in SHARES]
    out[1] += total - sum(out)
    return out


def workbooks(*, spell_nkhotakota="Nklhotakota", break_ages=False):
    a, e = openpyxl.Workbook(), openpyxl.Workbook()
    a1 = a.active
    a1.title = "A1"
    a1.append(["Table A1: Population of Malawi by Sex"])
    a1.append(["Region and District", 2018, None, None, 2008])
    a1.append([None, "Total", "Male", "Female", "Total", "Male", "Female"])
    regions = {r: (sum(m for _, m, _ in v), sum(f for _, _, f in v)) for r, v in AREAS.items()}
    tm, tf = sum(m for m, _ in regions.values()), sum(f for _, f in regions.values())
    a1.append(["Malawi", tm + tf, tm, tf, 1, 1, 0])
    for region, rows in AREAS.items():
        m, f = regions[region]
        a1.append(["", None])
        a1.append([region, m + f, m, f, 1, 1, 0])
        for name, dm, df in rows:
            a1.append([name, dm + df, dm, df, 1, 1, 0])
    a6 = a.create_sheet("A6")
    a6.append(["Table A6: Population by Age Group"])
    for region, rows in AREAS.items():
        m, f = regions[region]
        names = [spell_nkhotakota if n == "Nkhotakota" else n for n, _, _ in rows]
        cols = [ages(m + f)] + [ages(dm + df) for _, dm, df in rows]
        if break_ages and region == "Southern":
            cols[1][5] += 1
        a6.append(["Age Group", region])
        a6.append([None, f"{region} Total"] + names)
        a6.append(["Total", m + f] + [dm + df for _, dm, df in rows])
        for i, lab in enumerate(AGE_LABELS):
            a6.append([lab] + [c[i] for c in cols])
        a6.append([None])
    e4 = e.active
    e4.title = "E4"
    e4.append(["Table E4: Malawian Population by Tribe"])
    e4.append(["Region and District", None, "Tribe"])
    e4.append([None, "Total"] + list(mw.TRIBES))
    e5 = e.create_sheet("E5")
    e5.append(["Table E5. Population of Malawi by Denomination"])
    e5.append(["Area", "Denomination"])
    heads = ["Catholic", "CCAP", "Seventh Day Adventist/Baptist/Apostolic", "Anglican",
             "Pentecostal", "Other Christian Denominations", "Islam", "Traditional",
             "Other Denomination", "No Religion"]
    e5.append([None, "Total"] + heads)

    def line(sheet, name, total, n):
        parts = [total // n] * n
        parts[0] += total - sum(parts)
        sheet.append([name, total] + parts)

    line(e4, "Malawi", tm + tf - 3, len(mw.TRIBES))
    line(e5, "Malawi", tm + tf, len(heads))
    for region, rows in AREAS.items():
        m, f = regions[region]
        line(e4, region, m + f - 3 if region == "Northern" else m + f, len(mw.TRIBES))
        line(e5, region, m + f, len(heads))
        for name, dm, df in rows:
            line(e4, name, dm + df - (3 if name == "Likoma" else 0), len(mw.TRIBES))
            line(e5, name, dm + df, len(heads))
    e5.append(["Males", 1, 1])
    blobs = []
    for wb in (a, e):
        buf = io.BytesIO()
        wb.save(buf)
        blobs.append(buf.getvalue())
    return blobs, tm + tf


ADMIN1 = [{"id": "N", "name": "Northern Region"}, {"id": "C", "name": "Central Region"},
          {"id": "S", "name": "Southern Region"}]
ADMIN2 = [{"id": "mz", "name": "Mzimba", "parent": "N"}, {"id": "lk", "name": "Likoma",
                                                           "parent": "N"},
          {"id": "nk", "name": "Nkhotakota", "parent": "C"},
          {"id": "zo", "name": "Zomba", "parent": "S"}]


class MalawiCensusTest(unittest.TestCase):
    def table(self, **kw):
        (a, e), national = workbooks(**kw)
        with mock.patch.object(mw, "NATIONAL", national):
            return mw.parse(a, e)

    def test_cities_are_added_to_the_district_they_lie_in(self):
        recs = {r["shape_id"]: r for r in mw.build(self.table(), ADMIN1, ADMIN2)}
        mzimba = recs["mz"]
        self.assertEqual(mzimba["population"]["value"], 450 + 490 + 110 + 112)
        self.assertEqual(mzimba["sex_ratio"]["value"], round(1000 * 560 / 602))
        self.assertIn("Mzimba District and Mzuzu City, which lies within it",
                      mzimba["population_note"])
        self.assertEqual(recs["zo"]["population"]["value"], 350 + 390 + 51 + 53)

    def test_medians_religion_and_tribe(self):
        recs = {r["shape_id"]: r for r in mw.build(self.table(), ADMIN1, ADMIN2)}
        likoma = recs["lk"]
        groups15 = [(lo, w, n) for (lo, w), n in
                    zip([(0, 1), (1, 4)] + [(a, 5) for a in range(5, 90, 5)] + [(90, None)],
                        ages(15))]
        self.assertEqual(likoma["median_age"]["value"], mw.grouped_median(groups15))
        groups = {g["group"] for g in likoma["religion"]}
        self.assertIn("Church of Central Africa Presbyterian", groups)
        self.assertIn("Seventh Day Adventist/Baptist/Apostolic", groups)
        self.assertIn("Tribe of the 12 Malawians", likoma["ethnicity_note"])
        self.assertEqual(likoma["language"]["status"], "not_available")
        self.assertIn("no language", likoma["language"]["note"])

    def test_regions_take_no_population_of_their_own(self):
        recs = [r for r in mw.build(self.table(), ADMIN1, ADMIN2) if r["level"] == "admin1"]
        self.assertEqual(len(recs), 3)
        for r in recs:
            self.assertEqual(r["population"]["status"], "not_available")
            self.assertIsInstance(r["religion"], list)
            self.assertIn("value", r["median_age"])

    def test_age_groups_must_make_the_count(self):
        with self.assertRaises(SystemExit):
            self.table(break_ages=True)

    def test_a_polygon_without_a_census_district_is_refused(self):
        with self.assertRaises(SystemExit):
            mw.build(self.table(), ADMIN1, ADMIN2 + [{"id": "x", "name": "Nowhere",
                                                       "parent": "S"}])

    def test_a_census_district_without_a_polygon_is_refused(self):
        with self.assertRaises(SystemExit):
            mw.build(self.table(), ADMIN1, ADMIN2[:3])


if __name__ == "__main__":
    unittest.main()
