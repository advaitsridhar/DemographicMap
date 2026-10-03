"""Offline tests for the Balkan census readers: parsing, labels, sums, medians."""

import unittest
from collections import Counter

from scripts.fetch_census import balkans_common as common
from scripts.fetch_census import (albania_census, bosnia_age, bucharest_sectors, bulgaria_census, croatia,
                                  cyprus_census, greece_age, kosovo, montenegro, moldova_age, north_macedonia,
                                  north_macedonia_2002, romania_census, serbia_census)


class CommonTest(unittest.TestCase):
    def test_grouped_median_interpolates_within_the_group(self):
        # 10 people 0-4, 10 people 5-9: the middle person is at the 5-9 boundary.
        self.assertEqual(common.grouped_median([(0, 5, 10), (5, 5, 10)]), 5.0)
        self.assertEqual(common.grouped_median([(0, 5, 10), (5, 5, 30), (10, None, 1)]), 6.8)

    def test_median_in_the_open_group_is_refused(self):
        with self.assertRaises(SystemExit):
            common.grouped_median([(0, 5, 1), (5, None, 10)])

    def test_single_year_median(self):
        self.assertEqual(common.median_age(Counter({30: 1, 31: 1})), 31.0)

    def test_fold_drops_diacritics_and_dj(self):
        self.assertEqual(common.fold("Mađari Čair"), "madaricair")

    def test_match_names_is_one_to_one_and_reports_both_sides(self):
        shapes = [{"id": "a", "name": "Bar Municipality"}, {"id": "b", "name": "Budva Municipality"},
                  {"id": "c", "name": "Kotor Municipality"}]
        bound, left, spare = common.match_names({"1": "Bar", "2": "Budva", "3": "Tivat"}, shapes,
                                                strip=["Municipality"], who="test")
        self.assertEqual(bound, {"1": "a", "2": "b"})
        self.assertEqual(left, ["Tivat (3)"])
        self.assertEqual(spare, ["Kotor Municipality"])

    def test_check_sum_refuses_a_mismatch(self):
        common.check_sum(100, 100, "same")
        with self.assertRaises(SystemExit):
            common.check_sum(99, 100, "short")


class NorthMacedoniaTest(unittest.TestCase):
    def test_labels(self):
        self.assertEqual(north_macedonia.label_of("ethnicity", "Albanians"), "Albanian")
        self.assertEqual(north_macedonia.label_of(
            "religion", "Persons for whom data are taken from administrative sources"),
            "No religion data")

    def test_single_year_labels(self):
        self.assertEqual(north_macedonia.age_of("37"), 37)
        self.assertEqual(north_macedonia.age_of("100+"), 100)
        self.assertIsNone(north_macedonia.age_of("Age - TOTAL"))

    def test_five_year_groups_take_sexes_from_the_all_ages_row(self):
        groups = {"Age group - TOTAL": Counter({"male": 6, "female": 5, "sex - total": 11}),
                  "0-4": Counter({"male": 3, "female": 2}), "5-9": Counter({"male": 3, "female": 3}),
                  "90+": Counter()}
        rows, men, women = north_macedonia.five_year(groups)
        self.assertEqual((men, women), (6, 5))
        self.assertIn((0.0, 5.0, 5), rows)
        self.assertIn((90.0, None, 0), rows)

    def test_ts_and_c_are_one_letter(self):
        self.assertEqual(north_macedonia.key("Vraneshtitsa"), north_macedonia.key("Vraneshtica"))


class RomaniaTest(unittest.TestCase):
    def test_prefixes_and_i_hat(self):
        self.assertEqual(romania_census.bare("ORAȘ TÂRGU FRUMOS"), "TÂRGU FRUMOS")
        self.assertEqual(romania_census.fold("COVĂSÂNȚ", True), romania_census.fold("COVASINT"))
        self.assertEqual(romania_census.fold("CÂMPENI"), romania_census.fold("CAMPENI"))

    def test_labels_take_the_longest_prefix(self):
        self.assertEqual(romania_census.label_for("religion", "Ortodoxa Sarba"), "Serbian Orthodox")
        self.assertEqual(romania_census.label_for("religion", "Ortodoxa (Biserica Ortodoxa Romana)"),
                         "Orthodox")
        self.assertEqual(romania_census.label_for("ethnicity", "Romi"), "Roma")
        self.assertEqual(romania_census.label_for("ethnicity", "Români"), "Romanian")

    def test_suppressed_cells(self):
        self.assertIsNone(romania_census.number("*"))
        self.assertIsNone(romania_census.number("**"))
        self.assertEqual(romania_census.number("-"), 0.0)

    def test_what_the_stars_hide_is_one_bar(self):
        row = {"name": "ALBAC", "total": 1846.0, "groups": {"Romanian": 1693.0, "Roma": 64.0},
               "stars": 2}
        romania_census.check_row(row, "ethnicity")
        self.assertEqual(row["groups"][romania_census.SUPPRESSED], 89.0)
        with self.assertRaises(SystemExit):
            romania_census.check_row({"name": "X", "total": 10.0, "groups": {"a": 9.0},
                                      "stars": 0}, "ethnicity")
        with self.assertRaises(SystemExit):
            romania_census.check_row({"name": "X", "total": 10.0, "groups": {"a": 11.0},
                                      "stars": 1}, "ethnicity")

    def test_uat_rows_open_counties_and_read_ages(self):
        rows = [
            ["JUDET", "POPULATIA", "G R U P A", ""],
            ["", "", "0 - 4", "5 ani si peste"],
            ["", "", "ani", ""],
            ["A", "1.0", "2.0", "3.0"],
            ["ALBA", 200000.0, 50000.0, 150000.0],
            ["MUNICIPIUL ALBA IULIA", 150000.0, 40000.0, 110000.0],
            ["ALBAC", 50000.0, 10000.0, 40000.0],
        ]
        labels, data = romania_census.uat_rows(rows, {"alba"}, None)
        self.assertEqual(labels, [(0.0, 5.0), (5.0, None)])
        self.assertEqual([u["name"] for u in data["alba"]["uats"]],
                         ["MUNICIPIUL ALBA IULIA", "ALBAC"])
        self.assertEqual(data[""]["total"], 200000.0)


class MontenegroTest(unittest.TestCase):
    def test_protected_cells(self):
        self.assertIsNone(montenegro.number("z"))
        self.assertEqual(montenegro.number("-"), 0.0)
        self.assertEqual(montenegro.number(1234), 1234.0)

    def test_settle_keeps_the_hidden_people_apart(self):
        out = montenegro.settle("ethnicity", "Andrijevica", 100, {"Serbian": 90}, 1)
        self.assertEqual(out[montenegro.SUPPRESSED], 10)
        with self.assertRaises(SystemExit):
            montenegro.settle("ethnicity", "Andrijevica", 100, {"Serbian": 90}, 0)

    def test_labels(self):
        self.assertEqual(montenegro.label("ethnicity", "Crnogorci-Srbi"), "Montenegrin-Serbian")
        self.assertEqual(montenegro.label("language", "Srpsko-Hrvatski"), "Serbo-Croatian")
        self.assertEqual(montenegro.label("religion", "Ne želi da se izjasni"), "Not declared")


class MoldovaTest(unittest.TestCase):
    def test_floating_sums_round(self):
        self.assertEqual(moldova_age.persons(39562.9999999999), 39563.0)
        self.assertEqual(moldova_age.persons("-"), 0.0)


class CroatiaAgeTest(unittest.TestCase):
    def test_sheet_20_reads_the_all_settlements_rows(self):
        header = ("Županija", "", "", "", "Grad/općina", "Tip naselja", "", "Spol", "Sex",
                  "Ukupno\nTotal", "0 – 4", "5 i više")
        rows = [("20.",), header,
                ("Istarska", "Grad", "Istria", "Town", "Pula", "Ukupno", "Total", "sv.", "All", 10, 6, 4),
                ("Istarska", "Grad", "Istria", "Town", "Pula", "Ukupno", "Total", "m", "M", 5, 2, 3),
                ("Istarska", "Grad", "Istria", "Town", "Pula", "Ukupno", "Total", "ž", "W", 5, 2, 3),
                ("Istarska", "Grad", "Istria", "Town", "Pula", "U gradskim naseljima", "", "sv.", "All",
                 9, 4, 5)]
        ages = croatia.parse_ages(rows)
        unit = ages[("Istarska", "Grad", "Pula")]
        self.assertEqual((unit["total"], unit["men"], unit["women"]), (10, 5, 5))
        fields = croatia.age_fields(unit)
        self.assertEqual(fields["sex_ratio"]["value"], 100.0)
        self.assertEqual(fields["median_age"]["value"], 4.2)


class BosniaAgeTest(unittest.TestCase):
    def test_fr_t2_reads_the_sex_letters(self):
        rows = [("Područje", "Spol", "Ukupno", "0-4", "5-9", "85+", "Prosječna starost"),
                ("UNSKO-SANSKI KANTON", "Ukupno", 30, 10, 15, 5, 36.5),
                ("UNSKO-SANSKI KANTON", "M", 14, 5, 7, 2, 35.0),
                ("UNSKO-SANSKI KANTON", "Ž", 16, 5, 8, 3, 37.6),
                ("BIHAĆ", "Ukupno", 3, 1, 1, 1, 40.0)]
        out = bosnia_age.five_years(rows)
        unit = out["UNSKO-SANSKI KANTON"]
        self.assertEqual((unit["total"], unit["men"], unit["women"]), (30, 14, 16))
        self.assertEqual(unit["groups"], [(0.0, 5.0, 10), (5.0, 5.0, 15), (85.0, None, 5)])


class SerbiaTest(unittest.TestCase):
    def test_labels_and_totals(self):
        self.assertEqual(serbia_census.label_of("ethnicity", "Mađari"), "Hungarian")
        self.assertEqual(serbia_census.label_of("ethnicity",
                                                "Izjasnili se u smislu regionalne  pripadnosti"),
                         "Regional affiliation")
        self.assertIsNone(serbia_census.label_of("religion", "Svega hrišćanskа"))
        self.assertIsNone(serbia_census.label_of("religion", "Ukupno"))
        self.assertEqual(serbia_census.label_of("religion", "Nisu vernici (ateisti)"), "Atheism")
        self.assertEqual(serbia_census.label_of("language", "Klingonski"), "")

    def test_names_fold_dj_and_suffixes(self):
        self.assertEqual(serbia_census.key("Arandjelovac Municipality"), serbia_census.key("Aranđelovac"))
        self.assertEqual(serbia_census.key("Smederevska Palanka Municipal*"),
                         serbia_census.key("Smederevska Palanka"))
        self.assertEqual(serbia_census.key("Petrovac-na-Mlavi Municipality"),
                         serbia_census.key("Petrovac na Mlavi"))

    def test_age_groups_and_adults(self):
        def row(band, sex, n):
            return {"god": "2022", "nTipNaselja": "Ukupno", "nTer": "Ada", "nStarGrupa": band,
                    "nPol": sex, "vrednost": n}
        rows = [row("Ukupno", "Ukupno", 10), row("Ukupno", "Muško", 4), row("Ukupno", "Žensko", 6),
                row("0–4", "Ukupno", 4), row("85 i više godina", "Ukupno", 6),
                row("Punoletni (18+)", "Ukupno", 7), row("0–4", "Muško", 2)]
        saved = serbia_census.load
        serbia_census.load = lambda dataset: rows
        try:
            out = serbia_census.ages()
        finally:
            serbia_census.load = saved
        self.assertEqual(out["Ada"]["groups"], [(0.0, 5.0, 4), (85.0, None, 6)])
        self.assertEqual((out["Ada"]["men"], out["Ada"]["women"]), (4, 6))


class CyprusTest(unittest.TestCase):
    def test_district_spellings(self):
        for text in ("LEFKOSIA DISTRICT", "Lekfosia", "Lefkosia"):
            self.assertEqual(cyprus_census.district_key(text), "lefkosia")

    def test_community_keys(self):
        self.assertIn("aglantzia", cyprus_census.keys("Aglantzia or Aglangia"))
        self.assertIn("aglangia", cyprus_census.keys("Aglantzia or Aglangia"))
        self.assertIn("latsia", cyprus_census.keys("Latsia Municipality"))
        self.assertIn("agiavarvara", cyprus_census.keys("Agia Varvara Lefkosias"))
        self.assertEqual(cyprus_census.keys("Strovolos")[0], "strovolos")

    def test_parentheses_and_aliases(self):
        self.assertIn("voroklini", cyprus_census.keys("Voroklini (Oroklini)"))
        self.assertIn("paphos", cyprus_census.keys("Pafos"))

    def test_languages_named_after_countries_are_other(self):
        self.assertEqual(cyprus_census.LANGUAGES["indian"], "Other")
        self.assertEqual(cyprus_census.LANGUAGES["ukranian"], "Ukrainian")


class SpreadsheetMLTest(unittest.TestCase):
    def test_skipped_columns_and_numbers(self):
        blob = b"""<?xml version="1.0"?>
<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet"
 xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">
 <Worksheet ss:Name="T1"><Table>
  <Row><Cell><Data ss:Type="String">Berat</Data></Cell><Cell ss:Index="3"><Data ss:Type="Number">12</Data></Cell></Row>
  <Row ss:Index="3"><Cell><Data ss:Type="String">x</Data></Cell></Row>
 </Table></Worksheet></Workbook>"""
        rows = common.spreadsheetml(blob)["T1"]
        self.assertEqual(rows[0], ["Berat", None, 12.0])
        self.assertEqual(rows[1], [])
        self.assertEqual(rows[2], ["x"])

    def test_json_stat_one_is_read_as_two(self):
        v1 = {"dataset": {"dimension": {"id": ["A"], "size": [2],
                                        "A": {"category": {"index": {"a": 0, "b": 1},
                                                           "label": {"a": "x", "b": "y"}}}},
                          "value": [1, 2]}}
        out = common.unstack(common.json_stat1(v1, "u"))
        self.assertEqual(out, [({"A": ("a", "x")}, 1.0), ({"A": ("b", "y")}, 2.0)])


class AlbaniaTest(unittest.TestCase):
    def test_labels_take_the_longest_beginning(self):
        self.assertEqual(albania_census.label_of("religion", "Mysliman - Bektashi  Muslim - Bektashi"),
                         "Bektashi")
        self.assertEqual(albania_census.label_of("religion", "Mysliman Muslim"), "Islam")
        self.assertEqual(albania_census.label_of("religion", "Besimtarë të pacilësuar  Believers"),
                         "Unaffiliated believer")
        self.assertIsNone(albania_census.label_of("ethnicity", "Gjithsej  Total"))
        self.assertEqual(albania_census.label_of("ethnicity", "Grup etno-kulturor i përzier  Mixed"),
                         "Mixed")

    def test_hidden_cells_become_one_bar(self):
        rows = [["Tab. 1.12"], ["Qarku Dibër  Prefecture Dibër"],
                ["Gjithsej  Total", 100.0, 50.0, 50.0],
                ["Shqiptare  Albanian", 90.0, 45.0, 45.0],
                ["Greke  Greek", "..", "..", 1.0],
                ["Nuk disponohet  Not available", 8.0, 4.0, 4.0],
                ["Shënim ( .. nënkupton", None, None, None]]
        saved = (albania_census.http_get, albania_census.spreadsheetml, albania_census.NATIONAL)
        albania_census.http_get = lambda *a, **k: b""
        albania_census.spreadsheetml = lambda blob: {"Diber": rows}
        albania_census.NATIONAL = 100
        try:
            out = albania_census.read("ethnicity")
        finally:
            albania_census.http_get, albania_census.spreadsheetml, albania_census.NATIONAL = saved
        unit = out["diber"]
        self.assertEqual(unit["groups"][albania_census.SUPPRESSED], 2.0)
        self.assertEqual(unit["groups"]["Not stated"], 8.0)


class BucharestTest(unittest.TestCase):
    def test_age_bands(self):
        self.assertEqual(bucharest_sectors.band("0 - 4 ani"), (0.0, 5.0))
        self.assertEqual(bucharest_sectors.band("10-14 ani"), (10.0, 5.0))
        self.assertEqual(bucharest_sectors.band("85 ani si peste"), (85.0, None))
        self.assertIsNone(bucharest_sectors.band("TOTAL BUCURESTI"))


class NorthMacedonia2002Test(unittest.TestCase):
    def test_the_municipality_row_comes_before_its_settlement(self):
        pages = ["Tabela 2. x\nKi~evo 10 6 4 Ki~evo\nma`i 5 3 2 male\n`eni 5 3 2 female",
                 "Tabela 3. z\nKi~evo 10 10 - - - - - - - Ki~evo",
                 "Tabela 4. y\nDrugovo 3249 2790 448 2 - 9 Drugovo\nma`i 1682 1449 229 1 - 3 male\n"
                 "`eni 1567 1341 219 1 - 6 female\nDrugovo 1492 1250 237 2 - 3 Drugovo"]
        tables = north_macedonia_2002.sections(pages)
        total, men, women = north_macedonia_2002.municipality_rows(tables["religion"], "Drugovo")
        self.assertEqual(total, [3249, 2790, 448, 2, 0, 9])
        self.assertEqual(men[0] + women[0], 3249)


class CyprusBindingTest(unittest.TestCase):
    """build() against the map's own Cypriot polygons, with CYSTAT's tables
    replaced by a few communities that exercise each rule."""

    COMMUNITIES = [  # code, name, district, men, women
        ("1100", "Sia", "lefkosia", 430, 410),
        ("4215", "Kornos", "larnaka", 1040, 1041),
        ("1328", "Pano Koutrafas", "lefkosia", 18, 1),
        ("1329", "Kato Koutrafas", "lefkosia", 30, 34),
        ("1416", "Katydata", "lefkosia", 50, 50),
        ("6023", "Tremithousa", "pafos", 580, 591),
        ("4106", "Ormideia", "larnaka", 2000, 2100),
        ("1022", "Synoikismos Anthoupolis", "lefkosia", 800, 890),
        ("1024", "Lakatameia", "lefkosia", 19000, 21000),
        ("5356", "Troodos", "lemesos", 10, 7),
        ("4014", "Dromolaxia - Meneou", "larnaka", 3400, 3438),
        ("6340", "Karamoullides", "pafos", 1, 3),
    ]
    DISTRICT_CODES = {"lefkosia": "1000", "larnaka": "4000", "lemesos": "5000", "pafos": "6000",
                      "ammochostos": "3000"}

    def run_build(self):
        from collections import Counter
        comms, dist = {}, {}
        for code, name, d, men, women in self.COMMUNITIES:
            total = men + women
            comms[code] = {"name": name, "district": d, "men": float(men), "women": float(women),
                           "total": float(total),
                           "groups": Counter({(30.0, 5.0): total / 2, (35.0, 5.0): total / 2})}
        for d in self.DISTRICT_CODES:
            members = [u for u in comms.values() if u["district"] == d] or [
                {"men": 10.0, "women": 10.0, "total": 20.0}]
            dist[d] = {k: sum(u[k] for u in members) for k in ("men", "women", "total")}
        dist["total"] = {k: sum(dist[d][k] for d in self.DISTRICT_CODES) for k in ("men", "women", "total")}
        single = {d: {"ages": Counter({40: v["total"]}), "men": v["men"], "women": v["women"],
                      "total": v["total"]} for d, v in dist.items()}
        langs = {d: {None: v["total"], "Greek": v["total"]} for d, v in dist.items()}
        saved = (cyprus_census.single_years, cyprus_census.communities, cyprus_census.languages)
        cyprus_census.single_years = lambda: single
        cyprus_census.communities = lambda: (comms, dist)
        cyprus_census.languages = lambda: langs
        try:
            return cyprus_census.build()
        finally:
            cyprus_census.single_years, cyprus_census.communities, cyprus_census.languages = saved

    def test_rules(self):
        records = self.run_build()
        admin2 = [r for r in records if r["level"] == "admin2"]
        by_code = {str(r.get("codes", {}).get("cystat")): r for r in admin2 if r.get("codes")}
        # Sia and Kornos summed on the polygon labelled Sia.
        sia = next(r for r in admin2 if r["name"] == "Sia")
        self.assertEqual(sia["population"]["value"], 840 + 2081)
        self.assertIn("Kornos", sia["population_note"])
        self.assertIn("Nicosia district", sia["population_note"])
        # The swapped pair is on no polygon, and both polygons say why.
        self.assertNotIn("1328", by_code)
        self.assertNotIn("1329", by_code)
        kout = [r for r in admin2 if "Koutrafas" in r["name"]]
        self.assertEqual(len(kout), 2)
        self.assertTrue(all("swapped" in r["population"]["note"] for r in kout))
        # Pinned: Katydata and Tremithousa on the polygons that hold them.
        self.assertEqual(by_code["1416"]["shape_id"], "46923920B11980221307495")
        self.assertEqual(by_code["6023"]["shape_id"], "46923920B17841566161623")
        other = next(r for r in admin2 if r["shape_id"] == "46923920B20193016295799")
        self.assertIn("Agios Georgios", other["population"]["note"])
        # Ormideia is bound across the district line and says so, and so do
        # the two districts.
        self.assertIn("CYSTAT counts Ormideia in Larnaca", by_code["4106"]["population_note"])
        larnaca = next(r for r in records if r["level"] == "admin1" and r["name"] == "Larnaca")
        self.assertIn("Ormideia", larnaca["population_note"])
        self.assertIn("leaves out Sia", larnaca["population_note"])
        nicosia = next(r for r in records if r["level"] == "admin1" and r["name"] == "Nicosia")
        self.assertIn("includes Sia", nicosia["population_note"])
        self.assertIn("effective control", nicosia["population_note"])
        # Below 50 residents, no median or ratio; the counts are said in words.
        tiny = by_code["6340"]
        self.assertEqual(tiny["sex_ratio"]["status"], "not_available")
        self.assertIn("1 man and 3 women", tiny["sex_ratio"]["note"])
        # Lakatameia says it holds the undrawn Anthoupolis.
        self.assertIn("Synoikismos Anthoupolis", by_code["1024"]["population_note"])
        # A polygon with no community says why, and its marker displaces a
        # pre-1974 encyclopaedic figure.
        none = [r for r in admin2 if r["id"].startswith("CYP-2021-none-")]
        self.assertTrue(none)
        self.assertTrue(all(r["population"].get("displaces_before") == 1974 for r in none))
        kyrenia = [r for r in admin2 if r["id"].startswith("CYP-2021-outside-")]
        self.assertTrue(kyrenia and all(r["population"]["displaces_before"] == 1974 for r in kyrenia))
        # Every polygon of the map has exactly one record.
        ids = [r["shape_id"] for r in admin2]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertEqual(set(ids), {s["id"] for s in cyprus_census.shapes("CYP", "admin2")})

    def test_plural(self):
        self.assertEqual(cyprus_census.count_of(1, "man", "men"), "1 man")
        self.assertEqual(cyprus_census.count_of(2, "man", "men"), "2 men")
        self.assertEqual(cyprus_census.count_of(1_690, "resident"), "1,690 residents")


class ExactSharesTest(unittest.TestCase):
    def test_shares_add_to_a_hundred(self):
        counts = {f"g{i}": 3 for i in range(17)}
        counts["big"] = 1840
        out = common.exact_shares(counts, sum(counts.values()))
        self.assertAlmostEqual(sum(r["pct"] for r in out), 100.0, places=6)
        plain = sum(round(100 * v / 1891, 1) for v in counts.values())
        self.assertNotAlmostEqual(plain, 100.0, places=1)       # what shares() gives

    def test_order_and_counts(self):
        out = common.exact_shares({"b": 1, "a": 1, "c": 2}, 4)
        self.assertEqual([r["group"] for r in out], ["c", "a", "b"])
        self.assertEqual([r["count"] for r in out], [2, 1, 1])


class CroatiaClusterTest(unittest.TestCase):
    """A polygon the boundary file draws for several municipalities carries
    their sum; the Pirovac cluster waits for its merge."""

    PARTS = {("Zadarska", "Grad", "Benkovac"): (9680, 8000),
             ("Zadarska", "Općina", "Stankovci"): (1831, 1800),
             ("Zadarska", "Općina", "Lišane Ostrovičke"): (593, 500)}

    def tables(self, parts):
        by_key = {f: {} for f in croatia.SHEETS}
        ages = {}
        for (county, kind, name), (total, croats) in parts.items():
            key = (county, kind, name)
            counts = {"ethnicity": {"Croats": croats, "Serbs": total - croats},
                      "religion": {"Catholics": croats, "Orthodox": total - croats},
                      "language": {"Croatian": croats, "Other languages": total - croats}}
            for f in croatia.SHEETS:
                by_key[f][key] = {"county": county, "type": kind, "name": name, "total": total,
                                  "counts": counts[f]}
            half = total - total // 3                 # the median in the first, closed group
            ages[key] = {"groups": [(40.0, 5.0, half), (45.0, None, total - half)],
                         "men": half, "women": total - half, "total": total}
        return by_key, ages

    def test_benkovac_is_the_sum_of_three(self):
        by_key, ages = self.tables(self.PARTS)
        shape = {"id": "S1", "name": "Grad Benkovac"}
        rec = croatia.cluster("Grad Benkovac", "Zadarska",
                              (("Grad", "Benkovac"), ("Općina", "Stankovci"), ("Općina", "Lišane Ostrovičke")),
                              shape, by_key, ages)
        self.assertEqual(rec["population"]["value"], 9680 + 1831 + 593)
        self.assertEqual(rec["shape_id"], "S1")
        croat = next(r for r in rec["ethnicity"] if r["group"] == "Croatian")
        self.assertEqual(croat["count"], 10300)
        self.assertIn("Stankovci", rec["population_note"])
        self.assertIn("Lišane Ostrovičke", rec["ethnicity_note"])
        self.assertAlmostEqual(sum(r["pct"] for r in rec["ethnicity"]), 100.0, places=6)

    def test_undrawn_takes_its_parts_and_waits_for_the_merge(self):
        parts = dict(self.PARTS)
        for label, county, members in croatia.UNDRAWN:
            for kind, name in members:
                parts.setdefault((county, kind, name), (1000, 900))
        label, county, members, apart = croatia.REDRAWN_CLUSTER
        for kind, name in members:
            parts[(county, kind, name)] = (1500, 1400)
        by_key, ages = self.tables(parts)
        polys = [{"id": f"P{i}", "name": lab} for i, (lab, _, _) in enumerate(croatia.UNDRAWN)]
        polys += [{"id": "PIR", "name": label}] + [{"id": f"A{i}", "name": a} for i, a in enumerate(apart)]
        saved = croatia.log
        croatia.log = lambda *a: None
        try:
            records, taken = croatia.undrawn(by_key, ages, polys)
            self.assertEqual(len(taken), sum(len(m) for _, _, m in croatia.UNDRAWN) + len(members))
            # Still drawn apart: the three polygons say why, and carry no figure.
            waiting = [r for r in records if r["shape_id"] in ("PIR", "A0", "A1")]
            self.assertEqual(len(waiting), 3)
            self.assertTrue(all(r["sex_ratio"]["status"] == "not_available" for r in waiting))
            self.assertNotIn("population", {k for r in waiting for k, v in r.items()
                                            if k == "population" and isinstance(v, dict) and "value" in v})
            # Merged: one record on the kept polygon with the four's sum.
            merged, _ = croatia.undrawn(by_key, ages, polys[:-2])
            pir = [r for r in merged if r["shape_id"] == "PIR"]
            self.assertEqual(len(pir), 1)
            self.assertEqual(pir[0]["population"]["value"], 4 * 1500)
            # A part drawn in its own right stops the run.
            with self.assertRaises(SystemExit):
                croatia.undrawn(by_key, ages, polys + [{"id": "X", "name": "Općina Stankovci"}])
        finally:
            croatia.log = saved

    def test_national_median_against_eurostat(self):
        ages = {("A", None, None): {"groups": [(40.0, 5.0, 50), (45.0, None, 50)], "men": 50, "women": 50},
                ("A", "Grad", "X"): {"groups": [(0.0, 5.0, 999)], "men": 1, "women": 1}}
        saved = (croatia.NATIONAL_MEDIAN, croatia.log)
        croatia.log = lambda *a: None
        try:
            croatia.NATIONAL_MEDIAN = 45.2
            self.assertEqual(croatia.national_median(ages), 45.0)
            croatia.NATIONAL_MEDIAN = 45.4
            with self.assertRaises(SystemExit):
                croatia.national_median(ages)
        finally:
            croatia.NATIONAL_MEDIAN, croatia.log = saved


class AlbaniaDistrictTest(unittest.TestCase):
    def test_every_commune_once_and_sums(self):
        seen = set()
        for sid, (name, sheet, listed) in albania_census.DISTRICTS.items():
            for part in listed.split("|"):
                self.assertNotIn((sheet, part), seen)
                seen.add((sheet, part))
        self.assertEqual(len(seen), 373)
        self.assertEqual(len(albania_census.DISTRICTS), 36)
        self.assertTrue(set(albania_census.SAME_2023) <= set(albania_census.DISTRICTS))

    def test_rows_and_names(self):
        self.assertEqual(albania_census.unit_name("QENDËR    ."), "QENDËR")
        self.assertEqual(albania_census.unit_name("MOLLAS."), "MOLLAS")
        rows = [["2.1.2"], ["Bashkia/Komuna", "Gjithsej"],
                ["Gjithsej Total", 30.0, 10.0, 15.0, 5.0, 16.0, 5.0, 8.0, 3.0, 14.0, 5.0, 7.0, 2.0],
                ["BERAT", 20.0, 6.0, 10.0, 4.0, 11.0, 3.0, 6.0, 2.0, 9.0, 3.0, 4.0, 2.0],
                ["OTLLAK.", 10.0, 4.0, 5.0, 1.0, 5.0, 2.0, 2.0, 1.0, 5.0, 2.0, 3.0, 0.0]]
        units, whole = albania_census.sex_rows(rows, "test")
        self.assertEqual(units["OTLLAK"], (10.0, 5.0, 5.0))
        self.assertEqual(whole, (30.0, 16.0, 14.0))
        rows[3][5] = 12.0                                     # men and women no longer add up
        with self.assertRaises(SystemExit):
            albania_census.sex_rows(rows, "test")

    def test_districts_from_the_tables(self):
        """Each commune one person-count, so a district's figure is its count
        of communes; the twelve 2023 districts take Tab. 7's."""
        sheets = {}
        for i, sheet in enumerate(albania_census.SHEETS_2011):
            names = [p for n, s, listed in albania_census.DISTRICTS.values() if s == i
                     for p in listed.split("|")]
            k = float(len(names))
            sheets[sheet] = [["t"], ["Gjithsej Total", 2 * k, 2 * k, 0, 0, k, k, 0, 0, k, k, 0, 0]] + [
                [n, 2.0, 2.0, 0, 0, 1.0, 1.0, 0, 0, 1.0, 1.0, 0, 0] for n in names]
        munis = [["Gjithsej Total", float(albania_census.NATIONAL), float(albania_census.NATIONAL), 0, 0,
                  1_190_448.0, 1_190_448.0, 0, 0, 1_211_665.0, 1_211_665.0, 0, 0]]
        munis += [[m, 100.0, 100.0, 0, 0, 60.0, 60.0, 0, 0, 40.0, 40.0, 0, 0]
                  for m in albania_census.SAME_2023.values()]
        rest = albania_census.NATIONAL - 100 * len(munis[1:])
        munis.append(["Tiranë", float(rest), float(rest), 0, 0, 1_190_448.0 - 60 * 12, 1_190_448.0 - 60 * 12,
                      0, 0, 1_211_665.0 - 40 * 12, 1_211_665.0 - 40 * 12, 0, 0])
        saved = (albania_census.http_get, albania_census.sheets_of, albania_census.NATIONAL_2011,
                 albania_census.log)
        albania_census.http_get = lambda url, **k: url.encode()
        albania_census.sheets_of = lambda blob: sheets if b"3140" in blob else {"T": munis}
        albania_census.NATIONAL_2011 = 2 * 373
        albania_census.log = lambda *a: None
        try:
            admin1 = {s["id"]: s["name"] for s in albania_census.shapes("ALB", "admin1")}
            recs = albania_census.districts(admin1)
        finally:
            (albania_census.http_get, albania_census.sheets_of, albania_census.NATIONAL_2011,
             albania_census.log) = saved
        self.assertEqual(len(recs), 36)
        by = {r["shape_id"]: r for r in recs}
        kukes = by["67620474B92464717887955"]
        self.assertEqual((kukes["population"]["value"], kukes["population"]["year"]), (100, 2023))
        self.assertEqual(kukes["sex_ratio"]["value"], 150.0)
        berat = by["67620474B3220488890929"]
        self.assertEqual(berat["population"]["year"], 2011)
        self.assertEqual(berat["population"]["value"], 2 * len(albania_census.DISTRICTS[
            "67620474B3220488890929"][2].split("|")))
        korce = by["67620474B90792074310011"]
        self.assertEqual(korce["name"], "Korçë")
        self.assertIn("labels this polygon", korce["population_note"])
        self.assertEqual(berat["median_age"]["status"], "not_available")
        self.assertIn("three broad groups", berat["median_age"]["note"])


class BulgariaTest(unittest.TestCase):
    def test_transliteration(self):
        self.assertEqual(bulgaria_census.latin("Велинград"), "Velingrad")
        self.assertEqual(bulgaria_census.latin("Долна Митрополия"), "Dolna Mitropoliya")
        self.assertEqual(bulgaria_census.latin("Царево"), "Tsarevo")

    def test_codes(self):
        self.assertTrue(bulgaria_census.is_district("SOF"))
        self.assertTrue(bulgaria_census.is_municipality("SOF46"))
        self.assertFalse(bulgaria_census.is_municipality("SOF46-001"))
        self.assertFalse(bulgaria_census.is_municipality("XXX01"))

    def test_age_table(self):
        sexes = ("", "", "общо", "общо", "общо", "мъже", "мъже", "мъже", "жени", "жени", "жени")
        head = ("Код", "Име", "Общо", "0 - 4", "5+", "Общо", "0 - 4", "5+", "Общо", "0 - 4", "5+")
        rows = [("Таблица",), sexes, head,
                ("BG", "България", 10, 4, 6, 5, 2, 3, 5, 2, 3),
                ("VID", "Видин", 10, 4, 6, 5, 2, 3, 5, 2, 3),
                ("VID01", "Белоградчик", 10, 4, 6, 5, 2, 3, 5, 2, 3),
                ("VID01-00001", "Белоградчик", 1, 1, 0, 1, 1, 0, 0, 0, 0)]
        out = bulgaria_census.ages(rows)
        self.assertEqual(set(out), {"BG", "VID", "VID01"})
        self.assertEqual(out["VID01"]["groups"], [(0.0, 5.0, 4.0), (5.0, None, 6.0)])
        bad = list(rows)
        bad[5] = ("VID01", "Белоградчик", 10, 4, 5, 5, 2, 3, 5, 2, 3)
        with self.assertRaises(SystemExit):
            bulgaria_census.ages(bad)

    def test_composition_labels_and_footnotes(self):
        rows = [("Таблица 2",), ("Код", "Име", "Общо", "българска", "турска1", "непоказана", "друга"),
                ("BG", "България", 10, 6, 2, 1, 1), ("VID01", "Белоградчик", 5, 5, 0, 0, 0)]
        out = bulgaria_census.composition("ethnicity", rows)
        self.assertEqual(out["BG"]["groups"], {"Bulgarian": 6, "Turkish": 2, "Not stated": 1, "Other": 1})
        rows[1] = rows[1][:-1] + ("влашка",)
        with self.assertRaises(SystemExit):
            bulgaria_census.composition("ethnicity", rows)


class KosovoTest(unittest.TestCase):
    META = {
        "Komuna": {"code": "Komuna", "text": "Komuna", "values": ["0", "1"],
                   "valueTexts": ["KOSOVA", "Gllogoc"]},
        "Gjinia": {"code": "Gjinia", "text": "Sex", "values": ["0", "1", "2"],
                   "valueTexts": ["Total", "Male", "Female"]},
        "Viti": {"code": "Viti", "text": "Year", "values": ["0"], "valueTexts": ["2024"]},
    }

    def setUp(self):
        self.saved = (kosovo.px_meta, kosovo.px_table, dict(kosovo._SHAPE_KEYS))
        kosovo._SHAPE_KEYS.clear()
        kosovo._SHAPE_KEYS.update({"drenas": "Drenas", "kosova": "KOSOVA"})

    def tearDown(self):
        kosovo.px_meta, kosovo.px_table = self.saved[:2]
        kosovo._SHAPE_KEYS.clear()
        kosovo._SHAPE_KEYS.update(self.saved[2])

    def test_canon_reads_asks_spellings(self):
        self.assertEqual(kosovo.canon("Gllogoc"), "Drenas")
        self.assertEqual(kosovo.canon("Gllogovc"), "Drenas")
        self.assertEqual(kosovo.canon("KOSOVA"), "KOSOVA")

    def test_grouped_ages(self):
        meta = dict(self.META, Mosha={"code": "Mosha", "text": "Age", "values": ["0", "1", "2"],
                                      "valueTexts": ["Total", "0-4", "5+"]})
        cells = []
        for place, (total, men, women, young) in {"Gllogoc": (100, 48, 52, 60)}.items():
            muni = ("1", place)
            cells += [({"Komuna": muni, "Gjinia": ("0", "Total"), "Mosha": ("0", "Total")}, total),
                      ({"Komuna": muni, "Gjinia": ("1", "Male"), "Mosha": ("0", "Total")}, men),
                      ({"Komuna": muni, "Gjinia": ("2", "Female"), "Mosha": ("0", "Total")}, women),
                      ({"Komuna": muni, "Gjinia": ("0", "Total"), "Mosha": ("1", "0-4")}, young),
                      ({"Komuna": muni, "Gjinia": ("0", "Total"), "Mosha": ("2", "5+")}, total - young)]
        kosovo.px_meta = lambda url: meta
        kosovo.px_table = lambda url, select: cells
        out = kosovo.ages("u", grouped=True)
        self.assertEqual(out["Drenas"]["groups"], {(0.0, 5.0): 60, (5.0, None): 40})
        self.assertEqual((out["Drenas"]["men"], out["Drenas"]["women"]), (48, 52))

    def test_composition_labels(self):
        meta = dict(self.META, Etnia={"code": "Etnia", "text": "Ethnicity", "values": ["0", "1", "2"],
                                      "valueTexts": ["Total", "Albanian", "Egyptian"]})
        cells = [({"Komuna": ("1", "Gllogoc"), "Etnia": ("0", "Total")}, 10),
                 ({"Komuna": ("1", "Gllogoc"), "Etnia": ("1", "Albanian")}, 9),
                 ({"Komuna": ("1", "Gllogoc"), "Etnia": ("2", "Egyptian")}, 1)]
        kosovo.px_meta = lambda url: meta
        kosovo.px_table = lambda url, select: cells
        out = kosovo.composition("ethnicity", "u")
        self.assertEqual(out["Drenas"], {None: 10, "Albanian": 9, "Balkan Egyptian": 1})
        cells.append(({"Komuna": ("1", "Gllogoc"), "Etnia": ("3", "Martian")}, 0))
        with self.assertRaises(SystemExit):
            kosovo.composition("ethnicity", "u")


class GreeceAgeTest(unittest.TestCase):
    def test_single_years_and_the_open_group(self):
        payload = {"id": ["age", "geo", "sex", "time"], "dimension": {"time": {"category": {"index": {"2025": 0}}}}}
        cells = {("TOTAL", "EL51", "M", "2025"): 3.0, ("TOTAL", "EL51", "F", "2025"): 3.0,
                 ("Y_LT1", "EL51", "M", "2025"): 1.0, ("Y1", "EL51", "M", "2025"): 1.0,
                 ("Y_OPEN", "EL51", "M", "2025"): 1.0, ("Y2", "EL51", "F", "2025"): 3.0,
                 ("UNK", "EL51", "F", "2025"): 0.0}
        saved = (greece_age.http_json, greece_age.unpack, greece_age.log)
        greece_age.http_json = lambda url, **k: payload
        greece_age.unpack = lambda p: cells
        greece_age.log = lambda *a: None
        try:
            year, by_geo, totals = greece_age.ages(["EL51"])
        finally:
            greece_age.http_json, greece_age.unpack, greece_age.log = saved
        self.assertEqual(year, 2025)
        self.assertEqual(by_geo["EL51"]["M"], {0: 1.0, 1: 1.0, 100: 1.0})
        self.assertEqual(totals["EL51"], {"M": 3.0, "F": 3.0})

    def test_athos_row_and_the_sex_ratio_reason(self):
        import io
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Επίπεδο", "Κωδ", "Περιγραφή", "Μόνιμος Πληθυσμός 2021", "Άρρενες 2021", "Θήλεις 2021",
                   "Ποσοστό αρρένων 2021"])
        ws.append([4, "x", "ΑΓΙΟ ΟΡΟΣ", 1746, 1746, 0, 100])
        blob = io.BytesIO()
        wb.save(blob)
        saved = (greece_age.http_get, greece_age.log)
        greece_age.http_get = lambda url, **k: blob.getvalue()
        greece_age.log = lambda *a: None
        try:
            self.assertEqual(greece_age.athos(), (1746, 1746, 0))
        finally:
            greece_age.http_get, greece_age.log = saved

    def test_macedonia_thrace_says_athos_is_inside(self):
        self.assertIn("EL52", greece_age.ADMINISTRATIONS["Macedonia-Thrace"])
        self.assertIn("Mount Athos", greece_age.ATHOS_INSIDE)


if __name__ == "__main__":
    unittest.main()
