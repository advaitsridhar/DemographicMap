"""The two Main Results sheets as the C&SD workbook lays them out."""

import io
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT))

from scripts.fetch_census import hongkong_census as hk  # noqa: E402

# Table 3.9 (3), the both-sexes block, as the runner's probe printed the
# sheet: each row is its Chinese label with the figures, then its English
# label alone beneath, the share column last. The nesting is the column
# the label sits in.
E = None
ETHNICITY = [
    ["表3.9", E, E, E, E, "2021年按性別、種族及年齡組別劃分的人口（續）", E, E, E, E, E, E, E, E,
     "索引 Index"],
    ["Table 3.9", E, E, E, E, "Population by sex, ethnicity and age group, 2021 (cont'd)"],
    [],
    [E, E, E, E, E, E, "年齡組別"],
    [E, E, E, E, E, E, "Age group"],
    [E, E, E, E, E, E, "0 – 14", "15 – 24", "25 – 44", "45 – 64", "65+", "總計 Total"],
    [E, "性別", "種族", E, E, E, "數目", "數目", "數目", "數目", "數目", "數目", "百分比[1]"],
    [E, "Sex", "Ethnicity", E, E, E, "Number", "Number", "Number", "Number", "Number",
     "Number", "%[1]"],
    [],
    [E, "合計"],
    [E, "Both sexes"],
    [E, E, "華人", E, E, E, 747743, 550899, 1815457, 2254330, 1425073, 6793502, 91.6],
    [E, E, "Chinese"],
    [E, E, "非華人"],
    [E, E, "Non-Chinese"],
    [E, E, E, "菲律賓人", E, E, 3242, 4145, 134026, 56647, 3231, 201291, 2.7],
    [E, E, E, "Filipino"],
    [E, E, E, "印尼人", E, E, 355, 4983, 112851, 22711, 1165, 142065, 1.9],
    [E, E, E, "Indonesian"],
    [E, E, E, "南亞裔人士", E, E, 19935, 14566, 40163, 21351, 5954, 101969, 1.4],
    [E, E, E, "South Asian"],
    [E, E, E, E, "印度人", E, 6921, 5038, 18776, 9070, 2764, 42569, 0.6],
    [E, E, E, E, "Indian"],
    [E, E, E, E, "尼泊爾人", E, 5802, 3612, 11880, 6627, 1780, 29701, 0.4],
    [E, E, E, E, "Nepalese"],
    [E, E, E, E, "巴基斯坦人", E, 6372, 5524, 7341, 4095, 1053, 24385, 0.3],
    [E, E, E, E, "Pakistani"],
    [E, E, E, E, "其他南亞裔人士[2]", E, 840, 392, 2166, 1559, 357, 5314, 0.1],
    [E, E, E, E, "Other South Asian[2]"],
    [E, E, E, "泰國人", E, E, 372, 695, 2229, 7337, 2339, 12972, 0.2],
    [E, E, E, "Thai"],
    [E, E, E, "日本人", E, E, 1765, 440, 3321, 3769, 996, 10291, 0.1],
    [E, E, E, "Japanese"],
    [E, E, E, "韓國人", E, E, 1154, 756, 3638, 2611, 541, 8700, 0.1],
    [E, E, E, "Korean"],
    [E, E, E, "其他亞洲人", E, E, 1036, 974, 4904, 2584, 1076, 10574, 0.1],
    [E, E, E, "Other Asian"],
    [E, E, E, "白人", E, E, 9490, 4042, 22973, 20256, 4821, 61582, 0.8],
    [E, E, E, "White"],
    [E, E, E, "其他[3]", E, E, 22746, 9614, 20227, 11219, 6318, 70124, 0.9],
    [E, E, E, "Others[3]"],
    [E, E, E, "小計", E, E, 60095, 40215, 344332, 148485, 26441, 619568, 8.4],
    [E, E, E, "Sub-total"],
    [E, E, "總計", E, E, E, 807838, 591114, 2159789, 2402815, 1451514, 7413070, 100],
    [E, E, "Total"],
    [],
    ["註釋："],
    ["[1]", E, E, E, E, "數字顯示在性別組別的總人數中所佔的百分比。"],
    ["[3]", E, E, E, E, "數字包括報稱有多於一個種族的人士。"],
    ["Notes :"],
    ["[1]", E, E, E, E, "Figures represent the percentages in respect of the sex group totals."],
    ["[3]", E, E, E, E, "Figures include persons who reported more than one ethnicity."],
]

# Table 3.13: three places of birth and the total, each a number and a share.
LANGUAGE = [
    ["表3.13", E, E, "2021年按慣用交談語言及出生地點劃分的5歲及以上人口[1]"],
    ["Table 3.13", E, E, "Population[1] aged 5 and over by usual spoken language and place of birth, 2021"],
    [],
    [E, E, E, E, "出生地點"],
    [E, E, E, E, "Place of birth"],
    [E, E, E, E, "香港", E, "中國內地／澳門／台灣", E, "其他地方", E, "總計"],
    [E, E, E, E, "Hong Kong", E, "The mainland of China / Macao / Taiwan", E, "Elsewhere", E, "Total"],
    [E, "慣用交談語言", E, E, "數目", "百分比", "數目", "百分比", "數目", "百分比", "數目", "百分比"],
    [E, "Usual spoken language", E, E, "Number", "%", "Number", "%", "Number", "%", "Number", "%"],
    [],
    [E, "廣州話", E, E, 4222634, 97, 1890477, 85.7, 215836, 34.8, 6328947, 88.2],
    [E, "Cantonese"],
    [E, "英語", E, E, 62678, 1.4, 8675, 0.4, 259429, 41.8, 330782, 4.6],
    [E, "English"],
    [E, "普通話", E, E, 28436, 0.7, 123220, 5.6, 13795, 2.2, 165451, 2.3],
    [E, "Putonghua"],
    [E, "福建話", E, E, 5332, 0.1, 54274, 2.5, 1258, 0.2, 60864, 0.8],
    [E, "Fukien"],
    [E, "客家話", E, E, 5628, 0.1, 34882, 1.6, 1004, 0.2, 41514, 0.6],
    [E, "Hakka"],
    [E, "潮州話", E, E, 3534, 0.1, 33475, 1.5, 612, 0.1, 37621, 0.5],
    [E, "Chiu Chau"],
    [E, "其他中國方言", E, E, 4114, 0.1, 59871, 2.7, 587, 0.1, 64572, 0.9],
    [E, "Other Chinese dialects"],
    [E, "菲律賓語", E, E, 997, 0, 78, 0, 28338, 4.6, 29413, 0.4],
    [E, "Filipino (Tagalog)"],
    [E, "印尼語", E, E, 129, 0, 481, 0, 23634, 3.8, 24244, 0.3],
    [E, "Indonesian (Bahasa Indonesia)"],
    [E, "日本語", E, E, 1170, 0, 150, 0, 7384, 1.2, 8704, 0.1],
    [E, "Japanese"],
    [E, "其他語言", E, E, 17880, 0.4, 626, 0, 68509, 11, 87015, 1.2],
    [E, "Others"],
    [E, "總計", E, E, 4352532, 100, 2206209, 100, 620386, 100, 7179127, 100],
    [E, "Total"],
    [],
    ["註釋："],
    ["[1]", E, E, "數字不包括失去語言能力的人士。"],
    ["Note :"],
    ["[1]", E, E, "Figures exclude mute persons."],
]


def workbook(ethnicity=ETHNICITY, language=LANGUAGE) -> bytes:
    """The two sheets as an .xlsx in memory, named as the Department names them."""
    import openpyxl
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in ((hk.ETHNICITY_SHEET, ethnicity), (hk.LANGUAGE_SHEET, language)):
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def altered(rows, label, **cells):
    """The rows with one Chinese-label row's cells changed, found by its English label."""
    out = [list(r) for r in rows]
    for i, row in enumerate(out[1:], 1):
        if any(isinstance(c, str) and c == label for c in row):
            for index, value in cells.items():
                out[i - 1][int(index[1:])] = value
            return out
    raise KeyError(label)


class Reading(unittest.TestCase):
    def test_ethnicity_leaves_and_their_printed_shares(self):
        counts, printed = hk.read_ethnicity(ETHNICITY)
        self.assertEqual(counts["Chinese"], 6793502)
        self.assertEqual(counts["Filipino"], 201291)
        self.assertEqual(counts["Other South Asian"], 5314)
        self.assertEqual(counts["Other ethnic groups"], 70124)
        self.assertNotIn("South Asian", counts)
        self.assertNotIn("Non-Chinese", counts)
        self.assertNotIn("Sub-total", counts)
        self.assertEqual(sum(counts.values()), 7413070)
        self.assertEqual(printed["Chinese"], 91.6)
        self.assertEqual(printed["White"], 0.8)

    def test_language_total_column(self):
        counts, printed = hk.read_language(LANGUAGE)
        self.assertEqual(counts["Cantonese"], 6328947)
        self.assertEqual(counts["Filipino (Tagalog)"], 29413)
        self.assertEqual(counts["Other languages"], 87015)
        self.assertNotIn("Total", counts)
        self.assertEqual(sum(counts.values()), 7179127)
        self.assertEqual(printed["Cantonese"], 88.2)

    def test_the_record_as_the_report_states_it(self):
        records = hk.build(workbook())
        self.assertEqual(len(records), 1)
        rec = records[0]
        self.assertEqual(rec["name"], "Hong Kong Special Administrative Region")
        self.assertEqual(rec["id"], "CHN-hong-kong-special-administrative-region")
        self.assertEqual(rec["country"], "CHN")
        self.assertEqual(rec["ethnicity"][0], {"group": "Chinese", "pct": 91.6, "count": 6793502})
        self.assertEqual(rec["ethnicity"][1], {"group": "Filipino", "pct": 2.7, "count": 201291})
        self.assertEqual(len(rec["ethnicity"]), 13)
        self.assertAlmostEqual(sum(r["pct"] for r in rec["ethnicity"]), 100.0, delta=0.3)
        self.assertEqual(rec["ethnicity_year"], 2021)
        self.assertEqual(rec["language"][0], {"group": "Cantonese", "pct": 88.2, "count": 6328947})
        self.assertEqual(rec["language"][1], {"group": "English", "pct": 4.6, "count": 330782})
        self.assertEqual(len(rec["language"]), 11)
        self.assertAlmostEqual(sum(r["pct"] for r in rec["language"]), 100.0, delta=0.3)
        self.assertEqual(rec["language_year"], 2021)
        self.assertEqual(rec["language_basis"], "usual spoken language, population aged 5 and over")
        self.assertIn("aged 5 and over", rec["language_note"])
        self.assertIn("Table 3.9", rec["ethnicity_note"])
        # Religion is the China policy's business, not this file's.
        self.assertEqual(rec["religion"], {"status": "not_available"})
        self.assertEqual({s["field"] for s in rec["sources"]}, {"ethnicity", "language"})
        self.assertTrue(all(s["url"] == hk.URL for s in rec["sources"]))

    def test_a_footnote_mark_is_not_part_of_a_label(self):
        counts, _ = hk.read_ethnicity(ETHNICITY)
        self.assertNotIn("Others[3]", counts)
        self.assertNotIn("Other South Asian[2]", counts)


class Refusals(unittest.TestCase):
    def test_a_moved_total_column_is_refused(self):
        # Chinese's total shifted by one column: the count read is an age group.
        rows = altered(ETHNICITY, "Chinese", c11=1425073, c12=6793502)
        with self.assertRaises(SystemExit):
            hk.read_ethnicity(rows)

    def test_south_asian_must_be_its_detail_rows(self):
        rows = altered(ETHNICITY, "Indian", c11=42568)
        with self.assertRaises(SystemExit):
            hk.read_ethnicity(rows)

    def test_the_leaves_must_sum_to_the_printed_total(self):
        rows = altered(ETHNICITY, "Total", c11=7413071)
        with self.assertRaises(SystemExit):
            hk.read_ethnicity(rows)

    def test_language_rows_must_sum_to_the_printed_total(self):
        rows = altered(LANGUAGE, "Total", c10=7179128)
        with self.assertRaises(SystemExit):
            hk.read_language(rows)

    def test_a_share_that_disagrees_with_its_count_is_refused(self):
        rows = altered(LANGUAGE, "English", c11=5.6)
        with self.assertRaises(SystemExit):
            hk.build(workbook(language=rows))

    def test_cantonese_must_be_the_stated_share(self):
        counts, printed = hk.read_language(LANGUAGE)
        counts["Cantonese"], counts["English"] = 6000000, 659729
        printed["Cantonese"], printed["English"] = 83.6, 9.2
        with self.assertRaises(SystemExit):
            hk.checked("language", counts, "Cantonese", hk.CANTONESE_SHARE, printed)

    def test_a_missing_sheet_is_refused(self):
        import openpyxl
        book = openpyxl.Workbook()
        out = io.BytesIO()
        book.save(out)
        with self.assertRaises(SystemExit):
            hk.build(out.getvalue())

    def test_a_row_without_its_english_label_is_refused(self):
        rows = [list(r) for r in ETHNICITY]
        rows = [r for r in rows if not (r and any(c == "Nepalese" for c in r))]
        with self.assertRaises(SystemExit):
            hk.read_ethnicity(rows)


if __name__ == "__main__":
    unittest.main()
