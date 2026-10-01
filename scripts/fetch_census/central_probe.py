#!/usr/bin/env python3
"""Read-only probes for the central European census readers.

Germany, Austria, Switzerland, Liechtenstein, Poland, Czechia, Slovakia,
Hungary, Slovenia, the Netherlands and Luxembourg each publish population by
age and sex, and several of them census compositions, through their own
portals. Which table, which layout and which geography is not guessable, so
each named probe below asks one office what it serves and prints only what is
needed to write a reader against it: a table's variables, a catalogue's
matching entries, a file's first lines. Nothing is written; the output is the
log.

Usage:
    python -m scripts.fetch_census.central_probe aut che pol
    python -m scripts.fetch_census.central_probe get URL [REGEX]
"""

from __future__ import annotations

import gzip
import json
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
from typing import Any, Callable

from ._shared import log
from common import USER_AGENT  # noqa: E402  (_shared puts scripts/ on the path)

TIMEOUT = 90


def fetch(url: str, *, data: bytes | None = None, headers: dict[str, str] | None = None,
          timeout: int = TIMEOUT) -> tuple[int, dict[str, str], bytes]:
    req = urllib.request.Request(url, data=data, headers={
        "User-Agent": USER_AGENT, "Accept-Encoding": "gzip", **(headers or {})})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            body = resp.read()
            if resp.headers.get("Content-Encoding") == "gzip":
                body = gzip.decompress(body)
            return resp.status, dict(resp.headers), body
    except urllib.error.HTTPError as err:
        body = err.read() or b""
        try:
            if err.headers.get("Content-Encoding") == "gzip":
                body = gzip.decompress(body)
        except OSError:
            pass
        return err.code, dict(err.headers or {}), body
    except Exception as exc:  # noqa: BLE001 - a probe reports, it does not stop
        return -1, {"error": f"{exc.__class__.__name__}: {exc}"}, b""


def text(body: bytes) -> str:
    for enc in ("utf-8-sig", "cp1250", "latin-1"):
        try:
            return body.decode(enc)
        except UnicodeDecodeError:
            continue
    return body.decode("utf-8", "replace")


def show(url: str, find: str | None = None, *, limit: int = 40, raw: int = 0,
         headers: dict[str, str] | None = None, data: bytes | None = None) -> str:
    status, head, body = fetch(url, headers=headers, data=data)
    kind = head.get("Content-Type") or head.get("content-type") or head.get("error", "")
    log(f"\n## {url}\n   HTTP {status}, {len(body):,} bytes, {kind}")
    page = text(body)
    if find:
        seen: list[str] = []
        for m in re.finditer(find, page):
            hit = m.group(0)
            if hit not in seen:
                seen.append(hit)
        log(f"   {len(seen)} distinct matches of {find!r}")
        for hit in seen[:limit]:
            log("   - " + " ".join(hit.split())[:300])
    if raw:
        log("   " + page[:raw].replace("\n", "\n   "))
    return page


def pxweb_meta(url: str, values: int = 8) -> dict[str, Any] | None:
    status, _, body = fetch(url)
    log(f"\n## PxWeb {url}\n   HTTP {status}, {len(body):,} bytes")
    if status != 200:
        log("   " + text(body)[:300])
        return None
    try:
        meta = json.loads(text(body))
    except json.JSONDecodeError:
        log("   not JSON: " + text(body)[:300])
        return None
    if isinstance(meta, list):
        for row in meta[:values * 10]:
            log(f"   - {row.get('id')} [{row.get('type')}] {row.get('text')}")
        return None
    log(f"   title: {meta.get('title')}")
    for var in meta.get("variables", []):
        vals, texts = var.get("values", []), var.get("valueTexts", [])
        sample = ", ".join(f"{v}={t}" for v, t in list(zip(vals, texts))[:values])
        log(f"   {var.get('code')} ({var.get('text')}): {len(vals)} values; {sample}"
            + (f" ... last {vals[-1]}={texts[-1]}" if len(vals) > values else ""))
    return meta


def jpath(obj: Any, keys: str) -> Any:
    for key in keys.split("."):
        if isinstance(obj, dict):
            obj = obj.get(key)
        elif isinstance(obj, list) and key.isdigit():
            obj = obj[int(key)] if int(key) < len(obj) else None
        else:
            return None
    return obj


# ---------------------------------------------------------------------------
# One probe per office
# ---------------------------------------------------------------------------

def aut() -> None:
    for year in (2026, 2025, 2024):
        show(f"https://data.statistik.gv.at/web/meta.jsp?dataset=OGD_bevstandjbab2002_BevStand_{year}",
             r"OGD_[A-Za-z0-9_\-]+\.csv|<title>[^<]*")
    show("https://data.statistik.gv.at/web/catalog.jsp", r"dataset=OGD_bev[^\"&'<> ]+", limit=60)


def che() -> None:
    pxweb_meta("https://www.pxweb.bfs.admin.ch/api/v1/de/px-x-0102010000_101/px-x-0102010000_101.px")
    status, _, body = fetch("https://www.pxweb.bfs.admin.ch/api/v1/de/")
    log(f"\n## BFS PxWeb root: HTTP {status}, {len(body):,} bytes")
    try:
        rows = json.loads(text(body))
        log(f"   {len(rows)} databases")
        for row in rows:
            label = f"{row.get('dbid') or row.get('id')} {row.get('text')}"
            if re.search(r"Religion|Konfession|Sprache|Hauptsprache", label):
                log("   - " + label[:200])
    except Exception as exc:  # noqa: BLE001
        log(f"   {exc}: {text(body)[:200]}")
    page = show("https://www.agvchapp.bfs.admin.ch/api/communes/snapshot?date=01-01-2009", raw=600)
    lines = page.splitlines()
    log(f"   {len(lines)} lines")


def lie() -> None:
    show("https://www.statistikportal.li/de", r"href=\"[^\"]*(?:bev|volks|etab|Bev|Volks)[^\"]*\"", limit=60)
    show("https://etab.llv.li/", r"href=\"[^\"]+\"|<title>[^<]*", limit=40)


def pol() -> None:
    for query in ("mediana", "wiek"):
        page = show(f"https://bdl.stat.gov.pl/api/v1/variables/search?name={query}&format=json"
                    f"&page-size=100&lang=pl")
        try:
            data = json.loads(page)
            log(f"   totalRecords {data.get('totalRecords')}")
            for row in data.get("results", [])[:100]:
                log(f"   - {row.get('id')} subj {row.get('subjectId')} lvl {row.get('level')} "
                    f"years {row.get('years', [])[-2:] if row.get('years') else ''} | "
                    f"{row.get('n1')} | {row.get('n2')} | {row.get('n3')} | {row.get('n4')}")
        except Exception as exc:  # noqa: BLE001
            log(f"   {exc}")


def cze() -> None:
    show("https://csu.gov.cz/otevrena-data", r"href=\"[^\"]*(?:produkty|opendata|otevren)[^\"]*\"", limit=80)
    show("https://data.csu.gov.cz/", r"href=\"[^\"]+\"|<title>[^<]*", limit=30)


def svk() -> None:
    status, _, body = fetch("https://data.statistics.sk/api/v2/collection?lang=en")
    log(f"\n## DATAcube collection: HTTP {status}, {len(body):,} bytes")
    try:
        coll = json.loads(text(body))
        items = jpath(coll, "link.item") or []
        log(f"   {len(items)} datasets")
        for item in items:
            label = item.get("label", "")
            if re.search(r"(?i)age|nationalit|religio|mother tongue|census", label) and \
                    re.search(r"(?i)district|okres|municip", label):
                log(f"   - {item.get('href')} | {label[:160]}")
    except Exception as exc:  # noqa: BLE001
        log(f"   {exc}: {text(body)[:300]}")


def hun() -> None:
    show("https://www.ksh.hu/stadat?lang=en&theme=nep", r"href=\"[^\"]*stadat_files[^\"]*\"[^<]*<", limit=80)
    show("https://nepszamlalas2022.ksh.hu/en/", r"href=\"[^\"]+\"", limit=60)


def svn() -> None:
    pxweb_meta("https://pxweb.stat.si/SiStatData/api/v1/en/Data/05C4002S.px")
    status, _, body = fetch("https://pxweb.stat.si/SiStatData/api/v1/en/Data?query=municipalities&filter=*")
    log(f"\n## SiStat search: HTTP {status}, {len(body):,} bytes")
    try:
        for row in json.loads(text(body))[:80]:
            log(f"   - {row.get('id')} {row.get('title') or row.get('text')}"[:200])
    except Exception as exc:  # noqa: BLE001
        log(f"   {exc}: {text(body)[:300]}")


def nld() -> None:
    page = show("https://opendata.cbs.nl/ODataApi/odata/03759ned/DataProperties?$format=json")
    try:
        for row in json.loads(page).get("value", []):
            log(f"   - {row.get('Key')} [{row.get('Type')}] {row.get('Title')}")
    except Exception as exc:  # noqa: BLE001
        log(f"   {exc}")
    flt = urllib.parse.quote("substringof('ezindte',Title) or substringof('eligi',Title)")
    page = show(f"https://opendata.cbs.nl/ODataCatalog/Tables?$format=json&$filter={flt}")
    try:
        for row in json.loads(page).get("value", []):
            log(f"   - {row.get('Identifier')} | {row.get('Title')} | {row.get('Period')} | "
                f"{row.get('Frequency')}")
    except Exception as exc:  # noqa: BLE001
        log(f"   {exc}")


def lux() -> None:
    status, _, body = fetch("https://lustat.statec.lu/rest/dataflow/all/all/latest?detail=allstubs",
                            headers={"Accept": "application/vnd.sdmx.structure+json;version=1.0"})
    log(f"\n## LUSTAT dataflows: HTTP {status}, {len(body):,} bytes")
    page = text(body)
    try:
        flows = jpath(json.loads(page), "data.dataflows") or []
        log(f"   {len(flows)} dataflows")
        for flow in flows:
            name = flow.get("name") or jpath(flow, "names.en") or ""
            if re.search(r"(?i)commun|munic|langu|âge|age\b|age group", name):
                log(f"   - {flow.get('agencyID')}:{flow.get('id')}({flow.get('version')}) {name[:150]}")
    except Exception as exc:  # noqa: BLE001
        log(f"   {exc}: {page[:300]}")


def head_lines(url: str, n: int = 6, *, grep: str | None = None, limit: int = 40) -> list[str]:
    status, head, body = fetch(url)
    page = text(body)
    lines = page.splitlines()
    log(f"\n## {url}\n   HTTP {status}, {len(body):,} bytes, {len(lines):,} lines, "
        f"{head.get('Content-Type') or head.get('error', '')}")
    for line in lines[:n]:
        log("   | " + line[:240])
    if grep:
        hits = [line for line in lines if re.search(grep, line)]
        log(f"   {len(hits)} lines match {grep!r}")
        for line in hits[:limit]:
            log("   > " + line[:240])
    return lines


def aut2() -> None:
    stem = "https://data.statistik.gv.at/data/OGD_bevstandjbab2002_BevStand_2026"
    head_lines(stem + "_HEADER.csv", 12)
    head_lines(stem + ".csv", 5)
    head_lines(stem + "_C-GRGEMAKT-0.csv", 4, grep=r"GRGEMAKT-9|GRGEMAKT-10[12]", limit=8)
    head_lines(stem + "_C-GALTEJ112-0.csv", 3, grep=r"GALTEJ112-1(0\d|1\d)\b", limit=20)
    head_lines(stem + "_C-C11-0.csv", 4)
    head_lines("https://www.statistik.at/verzeichnis/reglisten/polbezirke.csv", 8)


def che2() -> None:
    lines = head_lines("https://www.agvchapp.bfs.admin.ch/api/communes/snapshot?date=01-01-2009", 1)
    rows = [line.split(",") for line in lines[1:]]
    cantons = {r[0]: r[7] for r in rows if len(r) > 7 and r[4] == "1"}
    districts = [r for r in rows if len(r) > 7 and r[4] == "2"]
    log(f"   2009: {len(cantons)} cantons, {len(districts)} districts, "
        f"{sum(1 for r in rows if len(r) > 7 and r[4] == '3')} communes")
    for r in districts:
        log(f"   D {cantons.get(r[5], r[5])} {r[1]} {r[6]}")
    for q in ("snapshot?date=01-01-2025", "correspondances?startPeriod=01-01-2009&endPeriod=01-01-2025",
              "correspondances?startPeriod=01-01-2009&endPeriod=01-01-2025&includeUnmodified=true",
              "levels?date=01-01-2025"):
        head_lines("https://www.agvchapp.bfs.admin.ch/api/communes/" + q, 4)
    pxweb_meta("https://www.pxweb.bfs.admin.ch/api/v1/de/px-x-0104010000_101/px-x-0104010000_101.px")
    status, _, body = fetch("https://www.pxweb.bfs.admin.ch/api/v1/de/")
    try:
        rows = json.loads(text(body))
        for row in rows:
            label = f"{row.get('dbid') or row.get('id')} {row.get('text')}"
            if re.search(r"(?i)relig|konfess|sprach", label):
                log("   - " + label[:200])
        log(f"   sample root rows: {rows[:3]}")
    except Exception as exc:  # noqa: BLE001
        log(f"   {exc}")


def lie2() -> None:
    for url in ("https://etab.llv.li/PXWeb/api/v1/de/", "https://etab.llv.li/api/v1/de/",
                "https://etab.llv.li/PXWeb/api/v1/de/eTab/"):
        pxweb_meta(url)
    show("https://www.statistikportal.li/de/themen/bevoelkerung/bevoelkerungsstand",
         r"href=\"[^\"]+\.(?:xlsx?|csv|pdf)\"|href=\"[^\"]*etab[^\"]*\"", limit=40)
    show("https://www.statistikportal.li/de/erhebungen-register/volkszaehlung",
         r"href=\"[^\"]+\.(?:xlsx?|csv|pdf)\"|href=\"[^\"]*etab[^\"]*\"", limit=40)


def pol2() -> None:
    base = "https://bdl.stat.gov.pl/api/v1"
    show(f"{base}/subjects/P4280?format=json&lang=pl", raw=1200)
    for q in ("roczniki", "mediana", "struktura ludno", "ludność według"):
        page = show(f"{base}/subjects/search?name={urllib.parse.quote(q)}&format=json&page-size=50&lang=pl")
        try:
            for row in json.loads(page).get("results", [])[:30]:
                log(f"   - {row.get('id')} lvls {row.get('levels')} vars {row.get('hasVariables')} "
                    f"{row.get('name')} | parent {row.get('parentId')}")
        except Exception as exc:  # noqa: BLE001
            log(f"   {exc}")
    show(f"{base}/data/by-variable/1650633?format=json&unit-level=5&page-size=3&lang=pl", raw=900)


def cze2() -> None:
    show("https://csu.gov.cz/otevrena_data", r"href=\"[^\"]+\"[^>]*>[^<]*(?:[Vv]ěk|[Oo]byvatel)[^<]*<",
         limit=40)
    show("https://csu.gov.cz/otevrena_data", r"href=\"[^\"]*(?:produkty|katalog|sady)[^\"]*\"", limit=30)


def svk2() -> None:
    for dim in ("om7009rr_vuc", "om7009rr_obd", "om7009rr_ukaz", "om7009rr_poh", "om7009rr_vek",
                "om7005rr_ukaz"):
        cube = dim.split("_")[0]
        page = show(f"https://data.statistics.sk/api/v2/dimension/{cube}/{dim}?lang=en")
        try:
            data = json.loads(page)
            cat = jpath(data, f"dimension.{dim}.category") or jpath(data, "category") or {}
            labels = cat.get("label", {})
            keys = list(labels)
            log(f"   {len(keys)} values: " + "; ".join(f"{k}={labels[k]}" for k in keys[:12])
                + (f" ... {keys[-1]}={labels[keys[-1]]}" if len(keys) > 12 else ""))
        except Exception as exc:  # noqa: BLE001
            log(f"   {exc}: {page[:300]}")
    show("https://www.scitanie.sk/", r"href=\"[^\"]*(?:otvoren|open|data|csv|xlsx)[^\"]*\"", limit=40)


def hun2() -> None:
    show("https://nepszamlalas2022.ksh.hu/en/results/tables", r"href=\"[^\"]+\"[^>]*>[^<]{3,}<", limit=80)
    show("https://www.ksh.hu/stadat?lang=en&theme=nep", r"<a [^>]*href=\"[^\"]+\"[^>]*>[^<]{3,}</a>", limit=80)


def svn2() -> None:
    for path in ("", "Data/", "Data/05C4010S.px", "Data/05C4004S.px", "Data/05C5002S.px",
                 "Data/05C2002S.px"):
        pxweb_meta(f"https://pxweb.stat.si/SiStatData/api/v1/en/{path}")


def nld2() -> None:
    for url in ("https://datasets.cbs.nl/odata/v1/CBS/03759ned",
                "https://odata4.cbs.nl/CBS/03759ned",
                "https://opendata.cbs.nl/ODataApi/odata/03759ned",
                "http://opendata.cbs.nl/ODataApi/odata/03759ned"):
        show(url, raw=500)


def lux2() -> None:
    base = "https://lustat.statec.lu/rest"
    for flow in ("DF_B1607", "DF_X021", "DF_B1102"):
        show(f"{base}/dataflow/LU1/{flow}/latest?references=all&detail=referencepartial",
             r"<(?:str|structure):Name[^>]*xml:lang=\"en\"[^>]*>[^<]+|id=\"[A-Z_0-9]+\"", limit=40,
             headers={"Accept": "application/vnd.sdmx.structure+xml;version=2.1"})
    show(f"{base}/data/LU1,DF_B1607,1.0/all?startPeriod=2021&dimensionAtObservation=AllDimensions",
         raw=1500, headers={"Accept": "application/vnd.sdmx.data+csv;version=1.0.0"})
    show("https://statistiques.public.lu/fr/recensement.html", r"href=\"[^\"]*(?:langu|xls|recensement)[^\"]*\"",
         limit=40)


def che3() -> None:
    """Which 2025 communes straddle a 2009 district, and the BFS tables on religion."""
    lines = fetch("https://www.agvchapp.bfs.admin.ch/api/communes/correspondances?"
                  "startPeriod=01-01-2009&endPeriod=01-01-2025&includeUnmodified=true")[2]
    rows = list(csv_rows(text(lines)))
    spans: dict[str, set[str]] = {}
    for r in rows:
        spans.setdefault(f"{r['TerminalCode']} {r['TerminalName']}", set()).add(r["InitialParentName"])
    log(f"   {len(rows)} correspondences, {len(spans)} terminal communes")
    for commune, parents in sorted(spans.items()):
        if len(parents) > 1:
            log(f"   ! {commune}: {sorted(parents)}")
    import time
    status, _, body = fetch("https://www.pxweb.bfs.admin.ch/api/v1/de/")
    dbs = [row.get("dbid") for row in json.loads(text(body))]
    log(f"   titles of {len(dbs)} BFS tables matching religion or language:")
    for dbid in dbs:
        if not dbid or not dbid.startswith("px-x-01"):
            continue
        status, _, body = fetch(f"https://www.pxweb.bfs.admin.ch/api/v1/de/{dbid}/{dbid}.px", timeout=30)
        try:
            title = json.loads(text(body)).get("title", "")
        except Exception:  # noqa: BLE001
            title = f"HTTP {status}"
        if re.search(r"(?i)relig|konfess|sprach", title):
            log(f"   - {dbid}: {title[:220]}")
        time.sleep(0.2)


def csv_rows(page: str, delimiter: str = ","):
    import csv
    import io
    return csv.DictReader(io.StringIO(page), delimiter=delimiter)


def lie3() -> None:
    base = "https://etab.llv.li/PXWeb/api/v1/de/eTab/"

    def walk(path: str, depth: int) -> None:
        status, _, body = fetch(base + urllib.parse.quote(path))
        try:
            rows = json.loads(text(body))
        except Exception:  # noqa: BLE001
            log(f"   {path}: HTTP {status} {text(body)[:120]}")
            return
        for row in rows:
            rid = row.get("id")
            log(f"   {'  ' * depth}{row.get('type')} {path}{rid} | {row.get('text')}"[:220])
            if row.get("type") == "l" and depth < 3:
                walk(f"{path}{rid}/", depth + 1)
    walk("Bevölkerung/", 0)


def pol3() -> None:
    base = "https://bdl.stat.gov.pl/api/v1"
    for url in (f"{base}/variables?subject-id=P3814&format=json&page-size=100&lang=pl",
                f"{base}/subjects?parent-id=G7&format=json&page-size=100&lang=pl"):
        page = show(url)
        try:
            for row in json.loads(page).get("results", []):
                log(f"   - {row.get('id')} {row.get('name') or ''} {row.get('n1') or ''} | "
                    f"{row.get('n2') or ''} | {row.get('n3') or ''} lvls {row.get('levels')} "
                    f"years {str(row.get('years'))[-40:]}")
        except Exception as exc:  # noqa: BLE001
            log(f"   {exc}")


def svk3() -> None:
    status, _, body = fetch("https://data.statistics.sk/api/v2/collection?lang=en")
    items = jpath(json.loads(text(body)), "link.item") or []
    for item in items:
        label = item.get("label", "")
        if re.search(r"(?i)2021|census|sodb|nationalit|religio|tongue|language", label):
            log(f"   - {item.get('href')} | {label[:160]}")
    show("https://data.statistics.sk/api/v2/dataset/om7005rr/SK0101/2024/IN010088/SPOLU?lang=en&type=json",
         raw=700)
    show("https://statdata.statistics.sk/public/ui/dc/domov/data-view", r"https?://[^\"' ]+", limit=20)


def hun3() -> None:
    import openpyxl
    import io as _io
    for code in ("1.1.2", "1.1.6", "1.1.7"):
        url = f"https://nepszamlalas2022.ksh.hu/en/results/final-data/tables/nsz2022-{code}-eng.xlsx"
        status, _, body = fetch(url)
        log(f"\n## {url}: HTTP {status}, {len(body):,} bytes")
        if status != 200:
            continue
        book = openpyxl.load_workbook(_io.BytesIO(body), read_only=True, data_only=True)
        for sheet in book.worksheets[:6]:
            rows = list(sheet.iter_rows(values_only=True, max_row=14))
            log(f"   sheet {sheet.title!r}: {sheet.max_row} rows x {sheet.max_column} cols")
            for row in rows:
                cells = [str(c)[:28] for c in row[:14] if c is not None]
                if cells:
                    log("     " + " | ".join(cells))


def svn3() -> None:
    status, _, body = fetch("https://pxweb.stat.si/SiStatData/api/v1/en/Data/")
    rows = json.loads(text(body))
    for row in rows:
        title = row.get("text", "")
        if re.search(r"(?i)municipal", title) and re.search(r"(?i)\bage\b|census|ethnic|religio|tongue",
                                                              title):
            log(f"   - {row.get('id')} {title[:200]}")


def nld3() -> None:
    for url in ("https://www.cbs.nl/", "https://archive.org/wayback/available?url=opendata.cbs.nl/ODataApi/odata/03759ned",
                "https://service.pdok.nl/cbs/wijkenbuurten/2022/wfs/v1_0?request=GetCapabilities&service=WFS"):
        show(url, raw=400)


def lux3() -> None:
    base = "https://lustat.statec.lu/rest/data"
    for flow in ("LU1,DF_X021,1.1/all?startPeriod=2015&endPeriod=2017",
                 "LU1,DSD_CENSUS_GROUP1_3@DF_B1607,1.0/all"):
        lines = head_lines(f"{base}/{flow}", 3)
        log(f"   ({len(lines)} lines)")
        # which municipalities and periods, briefly
        head = lines[0].split(",") if lines else []
        log(f"   columns: {head}")
    show("https://statistiques.public.lu/fr.html", r"href=\"[^\"]*(?:recens|rp20|langu)[^\"]*\"", limit=30)


def cze3() -> None:
    query = """PREFIX dct: <http://purl.org/dc/terms/>
PREFIX dcat: <http://www.w3.org/ns/dcat#>
SELECT ?d ?title ?url WHERE {
  ?d a dcat:Dataset ; dct:title ?title ; dcat:distribution ?dist .
  ?dist dcat:downloadURL ?url .
  FILTER(CONTAINS(LCASE(STR(?title)), "věk"))
  FILTER(CONTAINS(STR(?url), "csu.gov.cz") || CONTAINS(STR(?url), "czso.cz"))
} LIMIT 60"""
    url = "https://data.gov.cz/sparql?query=" + urllib.parse.quote(query)
    page = show(url, headers={"Accept": "application/sparql-results+json"})
    try:
        for b in json.loads(page)["results"]["bindings"]:
            log(f"   - {b['title']['value'][:110]} | {b['url']['value']}")
    except Exception as exc:  # noqa: BLE001
        log(f"   {exc}: {page[:300]}")


def aut3() -> None:
    page = text(fetch("https://data.statistik.gv.at/web/catalog.jsp")[2])
    ids = sorted(set(re.findall(r"dataset=(OGD_[^\"&'<> ]+)", page)))
    log(f"   {len(ids)} OGD datasets")
    for i in ids:
        if re.search(r"(?i)vz|volksz|relig|sprach|2001|census|reg", i):
            log(f"   - {i}")


def deu3() -> None:
    base = "https://www-genesis.destatis.de/genesisWS/rest/2020"
    for term in ("Sprache", "gesprochene Sprache"):
        body = urllib.parse.urlencode({"term": term, "category": "tables", "pagelength": "50",
                                       "language": "de"}).encode()
        status, _, reply = fetch(f"{base}/find/find", data=body,
                                 headers={"username": "GAST", "password": "GAST",
                                          "Content-Type": "application/x-www-form-urlencoded"})
        log(f"\n## GENESIS find {term!r}: HTTP {status}, {len(reply):,} bytes")
        try:
            data = json.loads(text(reply))
            log(f"   status: {data.get('Status')}")
            for row in (data.get("Tables") or [])[:50]:
                log(f"   - {row.get('Code')} | {row.get('Content')[:160]} | {row.get('Time')}")
        except Exception as exc:  # noqa: BLE001
            log(f"   {exc}: {text(reply)[:300]}")


def cze4() -> None:
    lines = head_lines("https://data.csu.gov.cz/opendata/sady/OBY02BHVD/distribuce/csv", 4)
    if lines:
        rows = list(csv_rows("\n".join(lines)))
        kinds = Counter_(r.get("uzemi_cis") or r.get("UZEMI_CIS") or "?" for r in rows)
        years = Counter_(r.get("casref_do") or r.get("rok") or "?" for r in rows)
        log(f"   territory kinds {dict(kinds)}; periods {dict(list(years.items())[-6:])}")
        log(f"   columns {list(rows[0].keys()) if rows else []}")


def Counter_(items):  # noqa: N802 - a local alias keeps the import where it is used
    from collections import Counter
    return Counter(items)


def pol4() -> None:
    base = "https://bdl.stat.gov.pl/api/v1"
    show(f"{base}/data/by-variable/746289?format=json&unit-level=5&page-size=2&lang=pl", raw=700)
    show(f"{base}/data/by-variable/746289?format=json&unit-level=0&lang=pl", raw=500)
    for subject in ("P1336", "P2137"):
        page = show(f"{base}/variables?subject-id={subject}&format=json&page-size=12&lang=pl")
        try:
            for row in json.loads(page).get("results", []):
                log(f"   - {row.get('id')} {row.get('n1')} | {row.get('n2')} | {row.get('n3')}")
        except Exception as exc:  # noqa: BLE001
            log(f"   {exc}")


def svk4() -> None:
    show("https://data.statistics.sk/api/v2/dataset/om7009rr/SK0101/2024/IN010053/1/all?lang=en&type=json",
         raw=600)
    page = show("https://data.statistics.sk/api/v2/dataset/om7005rr/all/2024/IN010088/SPOLU?lang=en&type=json")
    try:
        data = json.loads(page)
        log(f"   ids {data.get('id')} sizes {data.get('size')} values {str(data.get('value'))[:300]}")
    except Exception as exc:  # noqa: BLE001
        log(f"   {exc}")
    show("https://www.scitanie.sk/", r"href=\"[^\"]+\"", limit=80)


def hun4() -> None:
    show("https://nepszamlalas2022.ksh.hu/", r"href=\"[^\"]*(?:tabl|xlsx|terulet|kiadvany|adatbazis)[^\"]*\"",
         limit=60)
    show("https://nepszamlalas2022.ksh.hu/eredmenyek/vegleges-adatok/tablak/",
         r"href=\"[^\"]+\.xlsx\"", limit=80)
    show("https://nepszamlalas2022.ksh.hu/adatbazis/", raw=1500)


def nld4() -> None:
    for prefix in ("opendata.cbs.nl/ODataFeed/odata/03759ned", "opendata.cbs.nl/ODataApi/odata/03759ned",
                   "opendata.cbs.nl/CsvDownload/csv/03759ned", "opendata.cbs.nl/ODataApi/OData/03759ned"):
        head_lines(f"http://web.archive.org/cdx/search/cdx?url={prefix}*&limit=25&fl=timestamp,original,length",
                   25)


def lux4() -> None:
    csv_accept = {"Accept": "application/vnd.sdmx.data+csv;version=1.0.0"}
    for flow in ("LU1,DF_X021,1.1/all?startPeriod=2017&endPeriod=2017",
                 "LU1,DSD_CENSUS_GROUP1_3@DF_B1607,1.0/all"):
        status, _, body = fetch(f"https://lustat.statec.lu/rest/data/{flow}", headers=csv_accept)
        lines = text(body).splitlines()
        log(f"\n## {flow}: HTTP {status}, {len(lines)} lines")
        for line in lines[:4]:
            log("   | " + line[:300])
        if len(lines) > 1:
            rows = list(csv_rows("\n".join(lines)))
            for col in rows[0]:
                vals = Counter_(r[col] for r in rows)
                if len(vals) < 400:
                    log(f"   {col}: {len(vals)} values, e.g. {list(vals)[:12]}")
                else:
                    log(f"   {col}: {len(vals)} values")
    show("https://statistiques.public.lu/fr/publications/recensement.html", r"href=\"[^\"]+\"[^>]*>[^<]{4,}<",
         limit=80)


def aut4() -> None:
    for i in (1, 2, 3, 4):
        stem = f"https://data.statistik.gv.at/data/OGD_f0743_VZ_HIS_GEM_{i}"
        head_lines(stem + "_HEADER.csv", 12)
    show("https://data.statistik.gv.at/web/meta.jsp?dataset=OGD_f0743_VZ_HIS_GEM_1",
         r"OGD_[A-Za-z0-9_\-]+\.csv|<title>[^<]*|Religion[^<]{0,80}|Umgangssprache[^<]{0,80}")


def deu4() -> None:
    base = "https://genesis.destatis.de/genesisWS/rest/2020"
    for term in ("Sprache", "Religion"):
        body = urllib.parse.urlencode({"term": term, "category": "tables", "pagelength": "60",
                                       "language": "de"}).encode()
        status, _, reply = fetch(f"{base}/find/find", data=body,
                                 headers={"username": "GAST", "password": "GAST",
                                          "Content-Type": "application/x-www-form-urlencoded"})
        log(f"\n## GENESIS find {term!r}: HTTP {status}, {len(reply):,} bytes")
        try:
            data = json.loads(text(reply))
            log(f"   status: {data.get('Status')}")
            for row in (data.get("Tables") or [])[:60]:
                log(f"   - {row.get('Code')} | {row.get('Content')[:170]} | {row.get('Time')}")
        except Exception as exc:  # noqa: BLE001
            log(f"   {exc}: {text(reply)[:300]}")


def che4() -> None:
    import time
    status, _, body = fetch("https://www.pxweb.bfs.admin.ch/api/v1/de/")
    dbs = [row.get("dbid") for row in json.loads(text(body))]
    for dbid in dbs:
        if not dbid or not re.match(r"px-x-010[2-6]", dbid):
            continue
        status, _, body = fetch(f"https://www.pxweb.bfs.admin.ch/api/v1/de/{dbid}/{dbid}.px", timeout=30)
        try:
            title = json.loads(text(body)).get("title", "")
        except Exception:  # noqa: BLE001
            title = f"HTTP {status} {text(body)[:60]}"
        log(f"   - {dbid}: {title[:230]}")
        time.sleep(0.15)
    show("https://www.bfs.admin.ch/bfs/de/home/statistiken/bevoelkerung/sprachen-religionen/religionen.html",
         r"href=\"[^\"]*asset[^\"]*\"[^>]*>[^<]*<|href=\"[^\"]*dam-api[^\"]*\"", limit=40)


def cze5() -> None:
    lines = head_lines("https://csu.gov.cz/docs/107508/bc8f2d41-4d3a-a8f4-02fa-800d9cd27266/"
                       "130181-24data2023.csv", 4)
    rows = list(csv_rows("\n".join(lines)))
    if rows:
        for col in rows[0]:
            vals = Counter_(r[col] for r in rows)
            log(f"   {col}: {len(vals)} values, e.g. {list(vals)[:8]}")
    query = """PREFIX dct: <http://purl.org/dc/terms/>
PREFIX dcat: <http://www.w3.org/ns/dcat#>
SELECT ?title ?url WHERE {
  ?d a dcat:Dataset ; dct:title ?title ; dcat:distribution ?dist .
  ?dist dcat:downloadURL ?url .
  FILTER(CONTAINS(STR(?url), "130181") || CONTAINS(LCASE(STR(?title)), "jednotek věku"))
} LIMIT 80"""
    page = show("https://data.gov.cz/sparql?query=" + urllib.parse.quote(query),
                headers={"Accept": "application/sparql-results+json"})
    try:
        for b in json.loads(page)["results"]["bindings"]:
            log(f"   - {b['title']['value'][:110]} | {b['url']['value']}")
    except Exception as exc:  # noqa: BLE001
        log(f"   {exc}")


def hun5() -> None:
    show("https://nepszamlalas2022.ksh.hu/eredmenyek/statikus-tablak", r"href=\"[^\"]+\"[^>]*>[^<]{3,}<",
         limit=120)
    show("https://nepszamlalas2022.ksh.hu/adatbazis/app.js?v1", r"(?:https?:)?//[^\"' ]+|/api/[^\"' ]+|\.php[^\"' ]*",
         limit=40)


def che5() -> None:
    for table in ("px-x-0103030000_211", "px-x-0103030000_220", "px-x-0103030000_223"):
        pxweb_meta(f"https://www.pxweb.bfs.admin.ch/api/v1/de/{table}/{table}.px")
    head_lines("https://www.agvchapp.bfs.admin.ch/api/communes/correspondances?"
               "startPeriod=01-01-2009&endPeriod=01-01-2011", 1)


def svk5() -> None:
    show("https://www.scitanie.sk/obyvatelia/rozsirene-vysledky", r"href=\"[^\"]+\"[^>]*>[^<]{3,}<", limit=80)
    show("https://www.scitanie.sk/obyvatelia/zakladne-vysledky/pocet-obyvatelov/SR/SK0/SR",
         r"href=\"[^\"]*(?:zakladne-vysledky|xlsx|csv|stiahn)[^\"]*\"", limit=60)


def nld5() -> None:
    show("https://service.pdok.nl/cbs/wijkenbuurten/2022/wfs/v1_0?request=DescribeFeatureType&service=WFS"
         "&version=2.0.0&typeNames=wijkenbuurten:gemeenten", r"name=\"[a-z_0-9]+\"", limit=120)
    show("https://www.cbs.nl/nl-nl/zoeken?q=religie%20provincie", r"href=\"[^\"]*(?:relig|kerk)[^\"]*\"",
         limit=30)


def lux5() -> None:
    show("https://statistiques.public.lu/fr/publications/recensement/diversite-linguistique.html",
         r"href=\"[^\"]+\.(?:xlsx?|csv|pdf)\"|href=\"[^\"]*lustat[^\"]*\"", limit=40)
    status, _, body = fetch("https://lustat.statec.lu/rest/dataflow/all/all/latest?detail=allstubs",
                            headers={"Accept": "application/vnd.sdmx.structure+json;version=1.0"})
    flows = jpath(json.loads(text(body)), "data.dataflows") or []
    for flow in flows:
        name = flow.get("name") or ""
        if re.search(r"(?i)2011|census|recens|RP20|langu|main language|religi", name) and \
                not re.search(r"(?i)working status|occupation|economic|education", name):
            log(f"   - {flow.get('id')}({flow.get('version')}) {name[:150]}")


def hun6() -> None:
    page = text(fetch("https://nepszamlalas2022.ksh.hu/adatbazis/app.js?v1")[2])
    for pattern in (r"fetch\([^)]{0,120}\)", r"[\"'`][^\"'`]{0,60}(?:api|\.json|data/)[^\"'`]{0,80}[\"'`]"):
        hits = sorted(set(re.findall(pattern, page)))
        log(f"   {len(hits)} matches of {pattern!r}")
        for hit in hits[:40]:
            log("   - " + hit[:200])
    show("https://map.ksh.hu/nepszamlalas/", r"src=\"[^\"]+\"|href=\"[^\"]+\"", limit=30)


def nld6() -> None:
    show("https://www.cbs.nl/nl-nl/maatwerk/2026/11/religie-naar-regio-2021-2025",
         r"href=\"[^\"]+\.(?:xlsx|csv|zip|ods)\"|<title>[^<]*|<p>[^<]{20,400}</p>", limit=20)
    show("https://www.cbs.nl/nl-nl/zoeken?q=leeftijd%20gemeente%20geslacht%20maatwerk",
         r"href=\"[^\"]*(?:maatwerk|cijfers/detail)[^\"]*\"", limit=30)


def svk6() -> None:
    page = show("https://gis.scitanie.sk/portal/sharing/rest/search?q=n%C3%A1rodnos%C5%A5&f=json&num=40")
    try:
        for row in json.loads(page).get("results", []):
            log(f"   - {row.get('type')} | {row.get('title')} | {row.get('url')}")
    except Exception as exc:  # noqa: BLE001
        log(f"   {exc}")
    for q in ("okres", "vierovyznanie", "jazyk"):
        page = text(fetch(f"https://gis.scitanie.sk/portal/sharing/rest/search?q={q}&f=json&num=40")[2])
        try:
            for row in json.loads(page).get("results", []):
                log(f"   [{q}] {row.get('type')} | {row.get('title')} | {row.get('url')}")
        except Exception as exc:  # noqa: BLE001
            log(f"   {exc}")
    show("https://www.scitanie.sk/mapa-stranok", r"href=\"[^\"]*obyvatelia/zakladne-vysledky[^\"]*\"", limit=60)


def lie7() -> None:
    url = ("https://etab.llv.li/PXWeb/api/v1/de/eTab/" +
           urllib.parse.quote("Bevölkerung/Bevölkerungsstand/Stichtag 31 Dezember/211.004.px"))
    meta = json.loads(text(fetch(url)[2]))
    for v in meta["variables"]:
        log(f"   {v['code']!r}: {list(zip(v['values'], v['valueTexts']))[:16]}")
    query = {"query": [
        {"code": "Jahr", "selection": {"filter": "item", "values": ["0"]}},
        {"code": "Altersjahr", "selection": {"filter": "item", "values": ["0"]}},
        {"code": "Geschlecht", "selection": {"filter": "item", "values": ["0"]}},
        {"code": "Heimat", "selection": {"filter": "item", "values": ["0"]}},
        {"code": "Wohnort", "selection": {"filter": "item",
                                          "values": [str(i) for i in range(14)]}}],
        "response": {"format": "json-stat2"}}
    for fmt in ("json-stat2", "json-stat", "csv"):
        query["response"]["format"] = fmt
        status, _, body = fetch(url, data=json.dumps(query).encode(),
                                headers={"Content-Type": "application/json"})
        log(f"\n## POST {fmt}: HTTP {status}, {len(body):,} bytes")
        log("   " + text(body)[:1500].replace("\n", "\n   "))


def xlsx_dump(url: str, rows: int = 12, sheets: int = 8, cols: int = 14) -> None:
    import io as _io
    import openpyxl
    status, _, body = fetch(url)
    log(f"\n## {url}: HTTP {status}, {len(body):,} bytes")
    if status != 200:
        return
    book = openpyxl.load_workbook(_io.BytesIO(body), read_only=True, data_only=True)
    log(f"   sheets: {book.sheetnames}")
    for sheet in book.worksheets[:sheets]:
        log(f"   sheet {sheet.title!r}: {sheet.max_row} rows x {sheet.max_column} cols")
        for row in sheet.iter_rows(values_only=True, max_row=rows):
            cells = [str(c)[:30] for c in row[:cols] if c is not None]
            if cells:
                log("     " + " | ".join(cells))


def hun7() -> None:
    page = text(fetch("https://nepszamlalas2022.ksh.hu/adatbazis/app.js?v1")[2])
    for pattern in (r"\bal\s*=\s*[^,;]{0,100}", r"apiBaseUrl\s*[:=]\s*[^,;]{0,100}", r"new ol\([^)]{0,100}\)",
                    r"https://[a-z0-9.\-]*ksh\.hu[^\"'` ]*"):
        hits = sorted(set(re.findall(pattern, page)))
        log(f"   {pattern!r}: {hits[:12]}")
    for url in ("https://nepszamlalas2022.ksh.hu/adatbazis/api/version",
                "https://nepszamlalas2022.ksh.hu/api/version",
                "https://nepszamlalas2022.ksh.hu/adatbazis/api/index/hu/0"):
        show(url, raw=600)


def svk7() -> None:
    for q in ("obyv_", "Hosted narodnost", "okres narod", "materinsk", "nabozen"):
        page = text(fetch("https://gis.scitanie.sk/portal/sharing/rest/search?q="
                          + urllib.parse.quote(q) + "&f=json&num=100")[2])
        try:
            rows = json.loads(page).get("results", [])
            log(f"   [{q}] {len(rows)} results")
            for row in rows:
                if row.get("url") and "server/rest" in (row.get("url") or ""):
                    log(f"   - {row.get('type')} | {row.get('title')} | {row.get('url')}")
        except Exception as exc:  # noqa: BLE001
            log(f"   {exc}: {page[:200]}")
    show("https://gis.scitanie.sk/server/rest/services/Hosted?f=json", raw=3000)


def nld7() -> None:
    xlsx_dump("https://www.cbs.nl/-/media/_excel/2026/11/religie_2025_tabellen.xlsx", rows=16)


def hun8() -> None:
    base = "https://nepszamlalas2022.ksh.hu/api"
    for path in ("index/hu/V67", "index/en/V67", "index/hu", "dataflows", "structure"):
        show(f"{base}/{path}", raw=1500)


def svk8() -> None:
    base = "https://gis.scitanie.sk/server/rest/services/Hosted/"
    for service in ("obyv_ekchar_nar_dlnar", "obyv_ekchar_nar_matjaz", "obyv_ekchar_nabo_vekskup",
                    "obyv_ekchar_matjaz_vekskup", "obyv_demo_vek_zloz"):
        page = text(fetch(base + service + "/FeatureServer/layers?f=json")[2])
        try:
            layers = json.loads(page).get("layers", [])
        except Exception as exc:  # noqa: BLE001
            log(f"   {service}: {exc}: {page[:200]}")
            continue
        log(f"\n## {service}: {len(layers)} layers")
        for layer in layers:
            fields = [f["name"] for f in layer.get("fields", [])]
            log(f"   layer {layer.get('id')} {layer.get('name')!r}: {len(fields)} fields: "
                f"{fields[:60]}")


def nld8() -> None:
    import io as _io
    import openpyxl
    body = fetch("https://www.cbs.nl/-/media/_excel/2026/11/religie_2025_tabellen.xlsx")[2]
    book = openpyxl.load_workbook(_io.BytesIO(body), read_only=True, data_only=True)
    for name, cols in (("Tabel 2", 28), ("Toelichting", 3), ("Tabel 1", 6)):
        log(f"\n## sheet {name}")
        for row in book[name].iter_rows(values_only=True):
            cells = [str(c)[:60] for c in row[:cols] if c is not None]
            if cells:
                log("   " + " | ".join(cells))


def hun9() -> None:
    page = text(fetch("https://nepszamlalas2022.ksh.hu/adatbazis/app.js?v1")[2])
    for word in ("/api/index/", "/api/structure/", "/api/dataflows/"):
        for m in list(re.finditer(re.escape(word), page))[:2]:
            log(f"   ...{page[max(0, m.start() - 700):m.start() + 300]}...".replace("\n", " "))


def svk9() -> None:
    base = "https://gis.scitanie.sk/server/rest/services/Hosted/"
    for service, layer in (("obyv_ekchar_nabo_vekskup", 5), ("obyv_ekchar_matjaz_vekskup", 5),
                           ("obyv_ekchar_nar_vekskup", None), ("obyv_ekchar_nar_dlnar", 2)):
        if layer is None:
            page = text(fetch(base + service + "/FeatureServer/layers?f=json")[2])
            try:
                for lay in json.loads(page).get("layers", []):
                    log(f"   {service} layer {lay['id']} {lay['name']!r}: "
                        f"{[(f['name'], f.get('alias')) for f in lay.get('fields', [])][:40]}")
            except Exception as exc:  # noqa: BLE001
                log(f"   {exc}")
            continue
        page = text(fetch(f"{base}{service}/FeatureServer/{layer}?f=json")[2])
        try:
            info = json.loads(page)
            log(f"\n## {service}/{layer} {info.get('name')!r}")
            log(f"   fields: {[(f['name'], f.get('alias')) for f in info.get('fields', [])]}")
        except Exception as exc:  # noqa: BLE001
            log(f"   {exc}: {page[:200]}")
        q = (f"{base}{service}/FeatureServer/{layer}/query?where=1%3D1&outFields=*&returnGeometry=false"
             f"&resultRecordCount=4&f=json")
        page = text(fetch(q)[2])
        try:
            for feat in json.loads(page).get("features", []):
                log(f"   row: {json.dumps(feat.get('attributes'), ensure_ascii=False)[:700]}")
        except Exception as exc:  # noqa: BLE001
            log(f"   {exc}: {page[:200]}")
        count = text(fetch(f"{base}{service}/FeatureServer/{layer}/query?where=1%3D1&returnCountOnly=true&f=json")[2])
        log(f"   count: {count[:100]}")


def hun10() -> None:
    base = "https://nepszamlalas2022.ksh.hu/api"
    version = json.loads(text(fetch(f"{base}/version")[2]))["version"]
    for lang in ("en", "hu"):
        page = text(fetch(f"{base}/index/{version}/{lang}")[2])
        try:
            data = json.loads(page)
        except Exception as exc:  # noqa: BLE001
            log(f"   {lang}: {exc}: {page[:300]}")
            continue
        flows = data.get("dataflows", [])
        log(f"\n## index {version}/{lang}: keys {list(data)[:10]}, {len(flows)} dataflows")
        for flow in flows:
            label = json.dumps(flow, ensure_ascii=False)
            if re.search(r"(?i)district|járás|religi|vallás|nation|nemzetis|tongue|anyanyelv|age|kor\b|sex|nem\b",
                         label):
                log(f"   - {label[:400]}")
        if flows:
            log(f"   first: {json.dumps(flows[0], ensure_ascii=False)[:600]}")


def hun11() -> None:
    base = "https://nepszamlalas2022.ksh.hu/api"
    version = json.loads(text(fetch(f"{base}/version")[2]))["version"]
    index = json.loads(text(fetch(f"{base}/index/{version}/en")[2]))
    for flow in index["dataflows"]:
        if flow["id"] in ("WBS001", "WBS003", "WBS008", "WBS009", "WBS010"):
            log(f"   {flow['id']}: {json.dumps(flow['dimensions'], ensure_ascii=False)[:1500]}")
    dims = index.get("dimensions")
    log(f"   dimensions entry: {json.dumps(dims, ensure_ascii=False)[:1500]}")
    for flow in ("WBS001", "WBS003"):
        page = text(fetch(f"{base}/structure/{flow}/{version}")[2])
        log(f"\n## structure {flow}: {len(page):,} chars: {page[:1500]}")
        try:
            data = json.loads(page)
            cls = jpath(data, "data.codelists") or []
            for cl in cls:
                codes = cl.get("codes", [])
                log(f"   codelist {cl.get('id')}: {len(codes)} codes, e.g. "
                    f"{[(c.get('id'), jpath(c, 'names.en') or c.get('name')) for c in codes[:8]]}")
        except Exception as exc:  # noqa: BLE001
            log(f"   {exc}")
    for path in (f"dataflows/WBS001/{version}", f"dataflows/WBS001/{version}/d/CL_TIME_PERIOD,CL_TERUL_GEO4,CL_NEME_SEX,CL_KEV_AGE",
                 f"dataflows/WBS001/latest"):
        show(f"{base}/{path}", raw=800)


def hun12() -> None:
    base = "https://nepszamlalas2022.ksh.hu/api"
    version = json.loads(text(fetch(f"{base}/version")[2]))["version"]
    page = text(fetch("https://nepszamlalas2022.ksh.hu/adatbazis/app.js?v1")[2])
    for m in list(re.finditer(r"OBS_VALUE", page))[:3]:
        log(f"   ...{page[max(0, m.start() - 600):m.start() + 400]}...".replace("\n", " "))
    for flow in ("WBS001", "WBS003"):
        data = json.loads(text(fetch(f"{base}/structure/{flow}/{version}")[2]))
        dsds = jpath(data, "data.dataStructures") or []
        for dsd in dsds:
            dims = jpath(dsd, "dataStructureComponents.dimensionList.dimensions") or []
            log(f"   {flow} dimensions: {[(d.get('id'), jpath(d, 'localRepresentation.enumeration')) for d in dims]}")
        for cl in jpath(data, "data.codelists") or []:
            codes = cl.get("codes", [])
            if cl["id"] in ("CL_TERUL_GEO4", "CL_TERUL_GEO5"):
                named = [(c["id"], jpath(c, "names.en") or c.get("name"), c.get("parent")) for c in codes]
                log(f"   {cl['id']}: {len(codes)}; first 30: {named[:30]}")
                log(f"   {cl['id']}: codes by length: {Counter_(len(c['id']) for c in codes)}")
                log(f"   {cl['id']}: sample 5-char: {[n for n in named if len(n[0]) == 5][:10]}")
                log(f"   {cl['id']}: sample 3-char: {[n for n in named if len(n[0]) == 3][:40]}")
            if cl["id"] in ("CL_TEL_SZ_ADAT", "CL_KEV_AGE"):
                log(f"   {cl['id']}: {[(c['id'], jpath(c, 'names.en') or c.get('name'), c.get('parent')) for c in codes]}"[:6000])
    tries = [f"dataflows/WBS001/{version}/d/TIME_PERIOD:2022,TERUL_GEO4:HU110,NEME_SEX:M,KEV_AGE",
             f"dataflows/WBS001/{version}/d/TIME_PERIOD:2022,TERUL_GEO4:HU110,NEME_SEX:M+F,KEV_AGE:TOTAL+Y1+Y2",
             f"dataflows/WBS001/{version}/d/TERUL_GEO4:HU110"]
    for path in tries:
        show(f"{base}/{path}", raw=900)


def deu5() -> None:
    base = "https://genesis.destatis.de/genesisWS/rest/2020"
    for term in ("vorwiegend gesprochene Sprache", "Haushaltssprache", "Einwanderungsgeschichte",
                 "Migrationshintergrund"):
        body = urllib.parse.urlencode({"term": term, "category": "tables", "pagelength": "40",
                                       "language": "de"}).encode()
        status, _, reply = fetch(f"{base}/find/find", data=body,
                                 headers={"username": "GAST", "password": "GAST",
                                          "Content-Type": "application/x-www-form-urlencoded"})
        try:
            data = json.loads(text(reply))
            log(f"\n## GENESIS {term!r}: {len(data.get('Tables') or [])} tables")
            for row in (data.get("Tables") or [])[:40]:
                log(f"   - {row.get('Code')} | {' '.join(row.get('Content', '').split())[:170]}")
        except Exception as exc:  # noqa: BLE001
            log(f"   {exc}: {text(reply)[:200]}")
    show("https://www.destatis.de/SiteGlobals/Forms/Suche/Expertensuche_Formular.html?"
         "templateQueryString=vorwiegend+gesprochene+Sprache", r"href=\"[^\"]+\"[^>]*>[^<]{10,}<", limit=40)


def hun13() -> None:
    base = "https://nepszamlalas2022.ksh.hu/api"
    version = json.loads(text(fetch(f"{base}/version")[2]))["version"]
    data = json.loads(text(fetch(f"{base}/structure/WBS003/{version}")[2]))
    for cl in jpath(data, "data.codelists") or []:
        if cl["id"] == "CL_TEL_SZ_ADAT":
            codes = [(c["id"], jpath(c, "names.en"), c.get("parent")) for c in cl["codes"]]
            log(f"   CL_TEL_SZ_ADAT from 60: {codes[60:]}")
    for path in (f"dataflows/WBS003/{version}/d/TIME_PERIOD:2022,TERUL_GEO5:122,TEL_SZ_ADAT",
                 f"dataflows/WBS003/{version}/d/TIME_PERIOD:2022,TERUL_GEO5:HU120,TEL_SZ_ADAT",
                 f"dataflows/WBS001/{version}/d/TIME_PERIOD:2022,TERUL_GEO4:122,NEME_SEX:M,KEV_AGE"):
        page = text(fetch(f"{base}/{path}")[2])
        try:
            rows = json.loads(page)
            log(f"\n## {path}: {len(rows)} rows")
            for r in rows[:200]:
                key = r.get("TEL_SZ_ADAT") or r.get("KEV_AGE")
                log(f"   {key}={r.get('OBS_VALUE')}")
        except Exception as exc:  # noqa: BLE001
            log(f"   {exc}: {page[:300]}")


def deu6() -> None:
    """The Zensus 2022 database's answers to a few forms of the same request.

    The account reaches this probe through the environment and goes only into
    request headers; nothing here prints it.
    """
    import os
    auth = {"username": os.environ.get("ZENSUS_USER", ""), "password": os.environ.get("ZENSUS_PASSWORD", "")}
    base = "https://ergebnisse.zensus2022.de/api/rest/2020"
    for path in ("helloworld/logincheck", "catalogue/tables?selection=1000A-1018&language=de"):
        status, _, body = fetch(f"{base}/{path}", data=b"", headers={**auth, "Content-Type":
                                                                     "application/x-www-form-urlencoded"})
        log(f"\n## {path}: HTTP {status}: {text(body)[:400]}")
    variants = {
        "as written": {"name": "1000A-1018", "area": "all", "format": "ffcsv", "compress": "false",
                       "language": "de", "regionalvariable": "GEORB1", "regionalschluessel": ""},
        "no key": {"name": "1000A-1018", "area": "all", "format": "ffcsv", "compress": "false",
                   "language": "de", "regionalvariable": "GEORB1"},
        "Land cut": {"name": "1000A-1018", "area": "all", "format": "ffcsv", "compress": "false",
                     "language": "de", "regionalvariable": "GEOBL1"},
        "default": {"name": "1000A-1018", "area": "all", "format": "ffcsv", "compress": "false",
                    "language": "de"},
        "Kreise": {"name": "1000A-1018", "area": "all", "format": "ffcsv", "compress": "false",
                   "language": "de", "regionalvariable": "GEOLK4"},
    }
    for label, params in variants.items():
        status, head, body = fetch(f"{base}/data/tablefile", data=urllib.parse.urlencode(params).encode(),
                                   headers={**auth, "Content-Type": "application/x-www-form-urlencoded"})
        kind = head.get("Content-Type") or head.get("content-type")
        log(f"\n## tablefile {label}: HTTP {status}, {len(body):,} bytes, {kind}")
        if not body.startswith(b"PK"):
            log("   " + text(body)[:300].replace("\n", " "))


def che7() -> None:
    for url in ("https://www.bfs.admin.ch/bfs/de/home/statistiken/bevoelkerung/sprachen-religionen/religionen.html",
                "https://www.bfs.admin.ch/bfs/de/home/statistiken/bevoelkerung/sprachen-religionen/sprachen.html",
                "https://www.bfs.admin.ch/bfs/de/home/statistiken/bevoelkerung/sprachen-religionen.assetdetail.html"):
        show(url, r"assetdetail\.\d+\.html|dam-api[^\"' ]+|href=\"[^\"]*(?:xlsx|je-d-01\.08)[^\"]*\"", limit=40)


def nld9() -> None:
    url = ("https://service.pdok.nl/cbs/wijkenbuurten/2022/wfs/v1_0?request=GetFeature&service=WFS&version=2.0.0"
           "&typeNames=wijkenbuurten:gemeenten&propertyName=gemeentecode,gemeentenaam,water,mannen,vrouwen,jaar"
           "&outputFormat=application/json&count=5")
    show(url, raw=1500)
    for year in (2022, 2023):
        show(f"https://service.pdok.nl/cbs/wijkenbuurten/{year}/wfs/v1_0?request=GetFeature&service=WFS"
             f"&version=2.0.0&typeNames=wijkenbuurten:gemeenten&resultType=hits", raw=400)
    show("https://www.cbs.nl/nl-nl/maatwerk/2023/14/kerncijfers-wijken-en-buurten-2022",
         r"href=\"[^\"]+\.(?:xlsx|zip|csv)\"", limit=10)


CHE_RELIGION_ASSETS = (36343816, 36429858, 35627011, 35607916, 35607913, 35607915, 35607929,
                       35627019, 35607932, 36429856, 36429857)
CHE_LANGUAGE_ASSETS = (36434442, 36434449, 36630112, 36063525, 36434435, 36063516, 36063463,
                       36063461)


def che8() -> None:
    """What each BFS asset on the religion and language pages is, and the xlsx ones' heads."""
    for asset in CHE_RELIGION_ASSETS + CHE_LANGUAGE_ASSETS:
        status, head, body = fetch(f"https://dam-api.bfs.admin.ch/hub/api/dam/assets/{asset}")
        page = text(body)
        titles = re.findall(r'"(?:title|name|filename|fileName|mimeType|language)"\s*:\s*"([^"]{1,160})"',
                            page)
        log(f"\n# asset {asset}: meta HTTP {status}: {' | '.join(dict.fromkeys(titles))[:600]}")
        if status != 200:
            log("   " + page[:300].replace("\n", " "))
        status, head, body = fetch(f"https://dam-api.bfs.admin.ch/hub/api/dam/assets/{asset}/master")
        kind = head.get("Content-Type") or head.get("content-type") or head.get("error", "")
        disp = head.get("Content-Disposition") or head.get("content-disposition") or ""
        log(f"   master HTTP {status}, {len(body):,} bytes, {kind}, {disp[:120]}")
        if status == 200 and body[:2] == b"PK":
            import io as _io
            import openpyxl
            try:
                book = openpyxl.load_workbook(_io.BytesIO(body), read_only=True, data_only=True)
            except Exception as exc:  # noqa: BLE001
                log(f"   not a workbook: {exc}")
                continue
            log(f"   sheets: {book.sheetnames[:12]}")
            for sheet in book.worksheets[:2]:
                log(f"   sheet {sheet.title!r}")
                for row in sheet.iter_rows(values_only=True, max_row=16):
                    cells = [str(c)[:28] for c in row[:12] if c is not None]
                    if cells:
                        log("     " + " | ".join(cells))


def svn4() -> None:
    """Census 2002 by municipality: religion, mother tongue and ethnic affiliation."""
    show("https://www.stat.si/popis2002/si/", r"href=\"[^\"]+\"[^>]*>[^<]{3,80}<", limit=60)
    show("https://www.stat.si/popis2002/si/rezultati/rezultati_red.asp", r"href=\"[^\"]+\"[^>]*>[^<]{3,80}<",
         limit=80)
    for q in ("ter=OBC&st=7", "ter=OBC&st=8", "ter=OBC&st=9"):
        show(f"https://www.stat.si/popis2002/si/rezultati/rezultati_red.asp?{q}", raw=1200)


def aut5() -> None:
    """Census 2001 religion and Umgangssprache by Bezirk; the 2021 religion survey."""
    show("https://www.statistik.at/statistiken/bevoelkerung-und-soziales/bevoelkerung/"
         "bevoelkerungsstand/bevoelkerung-nach-religion",
         r"href=\"[^\"]+\.(?:xlsx?|ods|csv|pdf)\"|href=\"[^\"]*(?:religion|Religion)[^\"]*\"", limit=40)
    show("https://www.statistik.at/statistiken/bevoelkerung-und-soziales/bevoelkerung/"
         "bevoelkerungsstand/bevoelkerung-nach-sprache",
         r"href=\"[^\"]+\.(?:xlsx?|ods|csv|pdf)\"|href=\"[^\"]*(?:sprach|Sprach)[^\"]*\"", limit=40)
    for url in ("https://www.statistik.at/blickgem/vz1/g10101.pdf",
                "https://www.statistik.at/blickgem/G0101/g10101.pdf",
                "https://www.statistik.at/fileadmin/pages/405/VZ2001_Hauptergebnisse_I_Oesterreich.pdf"):
        status, head, body = fetch(url)
        log(f"\n## {url}: HTTP {status}, {len(body):,} bytes, "
            f"{head.get('Content-Type') or head.get('content-type') or head.get('error', '')}")


# ---------------------------------------------------------------------------
# Round ten: older censuses and official surveys for the fields today's
# censuses and registers do not ask.
# ---------------------------------------------------------------------------

def pdf_text(url: str, pages: int = 2, limit: int = 2500) -> None:
    """A PDF's page count and the text of its first pages."""
    status, head, body = fetch(url)
    log(f"\n## PDF {url}: HTTP {status}, {len(body):,} bytes")
    if status != 200 or not body.startswith(b"%PDF"):
        log("   " + text(body)[:300].replace("\n", " "))
        return
    import io as _io
    from pypdf import PdfReader
    reader = PdfReader(_io.BytesIO(body))
    log(f"   {len(reader.pages)} pages")
    for i, page in enumerate(reader.pages[:pages]):
        words = page.extract_text() or ""
        log(f"   --- page {i + 1}\n   " + words[:limit].replace("\n", "\n   "))


def aut6() -> None:
    """What a 2001-census Gemeinde sheet of "Ein Blick auf die Gemeinde" holds, and where the index is."""
    pdf_text("https://www.statistik.at/blickgem/vz1/g10101.pdf", pages=3)
    for url in ("https://www.statistik.at/blickgem/", "https://www.statistik.at/blickgem/index.jsp",
                "https://www.statistik.at/atlas/blick/"):
        show(url, r"href=\"[^\"]+\"[^>]*>[^<]{2,60}<", limit=30)
    show("https://www.statistik.at/statistiken/bevoelkerung-und-soziales/bevoelkerung/bevoelkerungsstruktur",
         r"href=\"[^\"]*(?:religi|Religi|sprach|Sprach|umgang)[^\"]*\"", limit=30)


def svn5() -> None:
    """SiStat's census tables of 2002 and 1991, and the 2002 census site in the Wayback Machine."""
    status, _, body = fetch("https://pxweb.stat.si/SiStatData/api/v1/en/Data/")
    try:
        rows = json.loads(text(body))
    except json.JSONDecodeError:
        log(f"   SiStat list: HTTP {status}, not JSON")
        rows = []
    hits = [r for r in rows if re.search(r"(?i)census|ethnic|religio|mother tongue|language",
                                         str(r.get("text")))]
    log(f"\n## SiStat: {len(rows)} tables, {len(hits)} about a census or ethnicity/language/religion")
    for r in hits:
        log(f"   - {r.get('id')} {str(r.get('text'))[:200]}")
    head_lines("http://web.archive.org/cdx/search/cdx?url=stat.si/popis2002/si/rezultati/*"
               "&limit=80&fl=timestamp,original,statuscode&filter=statuscode:200&collapse=urlkey", 80)


def che9() -> None:
    """BFS's asset search for the structural survey's religion and language tables by canton and district."""
    for query in ("orderNr=je-d-01.08.03.02", "orderNr=je-d-01.08.03.01", "orderNr=je-d-01.08.02.02",
                  "language=de&title=Religionszugeh%C3%B6rigkeit",
                  "language=de&q=Religionszugeh%C3%B6rigkeit%20Kanton",
                  "language=de&prodima=900046&articleModel=900052"):
        head_lines(f"https://dam-api.bfs.admin.ch/hub/api/dam/assets?{query}", 1)
        status, _, body = fetch(f"https://dam-api.bfs.admin.ch/hub/api/dam/assets?{query}")
        page = text(body)
        found = re.findall(r'"(?:title|orderNr|id|shortTextGnp)"\s*:\s*"?([^",]{1,140})', page)
        log(f"   {len(found)} fields: {' | '.join(found[:40])}")
    status, _, body = fetch("https://www.pxweb.bfs.admin.ch/api/v1/de/")
    dbs = [row.get("dbid") for row in json.loads(text(body))]
    import time
    for dbid in dbs:
        if not dbid or re.match(r"px-x-010[2-4]", dbid) or not re.match(r"px-x-01", dbid):
            continue
        status, _, body = fetch(f"https://www.pxweb.bfs.admin.ch/api/v1/de/{dbid}/{dbid}.px", timeout=30)
        try:
            title = json.loads(text(body)).get("title", "")
        except Exception:  # noqa: BLE001
            title = f"HTTP {status}"
        log(f"   - {dbid}: {title[:200]}")
        time.sleep(0.1)


def nld10() -> None:
    """StatLine's open-data hosts once more, and every attribute PDOK's 2022 gemeenten carry."""
    for url in ("https://opendata.cbs.nl/ODataApi/odata/03759ned/Perioden?$format=json",
                "https://datasets.cbs.nl/odata/v1/CBS/03759ned/PeriodenCodes",
                "https://dataderden.cbs.nl/ODataApi/odata/",
                "https://opendata.cbs.nl/statline/portal.html"):
        show(url, raw=300)
    show("https://service.pdok.nl/cbs/wijkenbuurten/2022/wfs/v1_0?request=DescribeFeatureType&service=WFS"
         "&version=2.0.0&typeNames=wijkenbuurten:gemeenten", r"name=\"[A-Za-z_0-9]+\"", limit=200)


def lux6() -> None:
    """The pre-2018 communes in LUSTAT's population series, and the 2021 language tables."""
    show("https://lustat.statec.lu/rest/codelist/LU1/CL_CANTON_COMMUNE/latest",
         r"id=\"[0-9C]+\"[^>]*>\s*<(?:com|common):Name xml:lang=\"en\">[^<]+", limit=130)
    csv_accept = {"Accept": "application/vnd.sdmx.data+csv;version=1.0.0"}
    status, _, body = fetch("https://lustat.statec.lu/rest/data/LU1,DF_X021,1.1/all?startPeriod=2011&endPeriod=2018",
                            headers=csv_accept)
    rows = list(csv_rows(text(body)))
    log(f"\n## DF_X021 2011-2018: HTTP {status}, {len(rows)} rows")
    per_year = Counter_(r.get("TIME_PERIOD") for r in rows)
    log(f"   units per year: {dict(sorted(per_year.items()))}")
    for url in ("https://statistiques.public.lu/fr/recensement/diversite-linguistique.html",
                "https://statistiques.public.lu/en/recensement.html"):
        show(url, r"href=\"[^\"]+\.(?:xlsx?|csv|zip|pdf)\"|href=\"[^\"]*(?:lingu|langu)[^\"]*\"", limit=30)


def deu7() -> None:
    """The Zensus database's refusal, asked several ways; the Regionaldatenbank's religion tables."""
    import os
    auth = {"username": os.environ.get("ZENSUS_USER", ""), "password": os.environ.get("ZENSUS_PASSWORD", "")}
    log(f"   credential present: {bool(auth['username'])}, {bool(auth['password'])}")
    base = "https://ergebnisse.zensus2022.de/api/rest/2020"
    forms = {
        "POST form, headers": dict(data=b"", headers={**auth, "Content-Type": "application/x-www-form-urlencoded"}),
        "POST, headers, no type": dict(data=b"", headers=auth),
        "GET, headers": dict(headers=auth),
        "GET, no credential": dict(),
    }
    for label, kwargs in forms.items():
        status, head, body = fetch(f"{base}/helloworld/logincheck", **kwargs)
        log(f"   logincheck {label}: HTTP {status}: {' '.join(text(body)[:160].split())}")
    status, head, body = fetch("https://ergebnisse.zensus2022.de/api/rest/2020/helloworld/whoami")
    log(f"   whoami: HTTP {status}: {' '.join(text(body)[:160].split())}")
    show("https://ergebnisse.zensus2022.de/datenbank/online/", raw=400)
    rdb = "https://www.regionalstatistik.de/genesisws/rest/2020"
    for term in ("Religion", "Zensus 2022"):
        body = urllib.parse.urlencode({"term": term, "category": "tables", "pagelength": "40",
                                       "language": "de"}).encode()
        status, _, reply = fetch(f"{rdb}/find/find", data=body,
                                 headers={"username": "GAST", "password": "GAST",
                                          "Content-Type": "application/x-www-form-urlencoded"})
        try:
            data = json.loads(text(reply))
            log(f"\n## Regionaldatenbank {term!r}: {len(data.get('Tables') or [])} tables")
            for row in (data.get("Tables") or [])[:40]:
                log(f"   - {row.get('Code')} | {' '.join(row.get('Content', '').split())[:170]}")
        except Exception as exc:  # noqa: BLE001
            log(f"   Regionaldatenbank {term!r}: HTTP {status} {exc}: {text(reply)[:200]}")


def px_all(url: str) -> None:
    """Every value of every variable of one PxWeb table."""
    status, _, body = fetch(url)
    log(f"\n## PxWeb {url}: HTTP {status}")
    try:
        meta = json.loads(text(body))
    except json.JSONDecodeError:
        log("   " + text(body)[:300])
        return
    log(f"   title: {meta.get('title')}")
    for var in meta.get("variables", []):
        pairs = list(zip(var.get("values", []), var.get("valueTexts", [])))
        shown = pairs if len(pairs) <= 80 else pairs[:12] + [("...", "...")] + pairs[-6:]
        log(f"   {var.get('code')} ({var.get('text')}): {len(pairs)} values: "
            + "; ".join(f"{v}={t}" for v, t in shown))


def svn6() -> None:
    """The 2002 census tables by municipality and by statistical region, in full."""
    for table in ("05W1002S", "05W1006S", "05W1007S", "05W1001S", "0558601S", "05W2202S"):
        px_all(f"https://pxweb.stat.si/SiStatData/api/v1/en/Data/{table}.px")


def walk_strings(obj: Any, keys: tuple[str, ...], out: list[str], depth: int = 0) -> None:
    if depth > 6:
        return
    if isinstance(obj, dict):
        for k, v in obj.items():
            if k in keys and isinstance(v, (str, int)):
                out.append(f"{k}={v}")
            else:
                walk_strings(v, keys, out, depth + 1)
    elif isinstance(obj, list):
        for v in obj:
            walk_strings(v, keys, out, depth + 1)


def che10() -> None:
    """The shape of BFS's asset search, and the religion and language tables it finds."""
    for query in ("language=de&title=Religionszugeh%C3%B6rigkeit",
                  "language=de&title=Hauptsprachen",
                  "language=de&title=Religion%20Kanton"):
        status, _, body = fetch(f"https://dam-api.bfs.admin.ch/hub/api/dam/assets?{query}")
        log(f"\n## DAM {query}: HTTP {status}, {len(body):,} bytes")
        try:
            data = json.loads(text(body))
        except json.JSONDecodeError:
            log("   " + text(body)[:300])
            continue
        if isinstance(data, dict):
            log(f"   top-level keys: {list(data)[:20]}")
            items = next((v for v in data.values() if isinstance(v, list)), [])
        else:
            items = data
        log(f"   {len(items)} items")
        if items and query.endswith("keit"):
            log("   first item: " + json.dumps(items[0], ensure_ascii=False)[:2500])
        for item in items[:60]:
            found: list[str] = []
            walk_strings(item, ("ident", "id", "title", "orderNr", "periodFrom", "periodTo",
                                "mimetype", "mimeType", "articleModel"), found)
            log("   - " + " | ".join(found[:12])[:300])


def aut7() -> None:
    """The other sheets of the 2001 Gemeinde series, and the 2021 religion survey's page."""
    for section in ("vz2", "vz3", "vz4", "vz5", "vz6", "vz7"):
        pdf_text(f"https://www.statistik.at/blickgem/{section}/g10101.pdf", pages=1, limit=1800)
    show("https://www.statistik.at/blickgem/", r"src=\"[^\"]+\"|[\"'][^\"']*\.(?:json|do|jsp)[^\"']*[\"']",
         limit=30)
    for url in ("https://www.statistik.at/statistiken/bevoelkerung-und-soziales/bevoelkerung",
                "https://www.statistik.at/statistiken/bevoelkerung-und-soziales/bevoelkerung/"
                "bevoelkerungsstand"):
        show(url, r"href=\"[^\"]*(?:religi|Religi|sprach|Sprach|umgang|volkszaehl|registerz)[^\"]*\"",
             limit=40)


def deu8() -> None:
    """The Regionaldatenbank's Zensus tables, asked by code rather than by word."""
    rdb = "https://www.regionalstatistik.de/genesisws/rest/2020"
    for path, params in (("helloworld/logincheck", {}),
                         ("catalogue/tables", {"selection": "1000A*", "pagelength": "50"}),
                         ("catalogue/tables", {"selection": "12111*", "pagelength": "60"})):
        body = urllib.parse.urlencode({**params, "language": "de"}).encode()
        status, _, reply = fetch(f"{rdb}/{path}", data=body,
                                 headers={"username": "GAST", "password": "GAST",
                                          "Content-Type": "application/x-www-form-urlencoded"})
        page = text(reply)
        log(f"\n## Regionaldatenbank {path} {params}: HTTP {status}, {len(reply):,} bytes")
        try:
            data = json.loads(page)
        except json.JSONDecodeError:
            log("   " + " ".join(page[:300].split()))
            continue
        rows = data.get("List") or data.get("Tables") or []
        log(f"   status {data.get('Status')}; {len(rows)} rows")
        for row in rows[:60]:
            log(f"   - {row.get('Code')} | {' '.join(str(row.get('Content', '')).split())[:170]}")


def dam_items(query: str, pages: int = 3) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for page in range(pages):
        status, _, body = fetch(f"https://dam-api.bfs.admin.ch/hub/api/dam/assets?{query}&skip={20 * page}")
        try:
            data = json.loads(text(body))
        except json.JSONDecodeError:
            break
        items = data.get("data") or []
        out += items
        if len(items) < 20:
            break
    return out


def che11() -> None:
    """BFS's structural-survey tables on religion and main language, with their geography."""
    for query in ("language=de&title=Religionszugeh%C3%B6rigkeit",
                  "language=de&title=Religion",
                  "language=de&title=Hauptsprachen",
                  "language=de&title=Volksz%C3%A4hlung%202000",
                  "language=de&title=Bezirk%20Religion"):
        items = dam_items(query)
        log(f"\n## DAM {query}: {len(items)} items")
        for item in items:
            bfs, desc = item.get("bfs") or {}, item.get("description") or {}
            cat = desc.get("categorization") or {}
            space = ", ".join(s.get("name", "") for s in cat.get("spatialdivision") or [])
            log(f"   - {jpath(item, 'ids.damId')} | {jpath(bfs, 'articleModel.name')} | "
                f"{jpath(item, 'shop.orderNr')} | {jpath(desc, 'bibliography.period')} | {space} | "
                f"{jpath(desc, 'titles.main')}"[:300])


def aut8() -> None:
    """Which Gemeinde codes the 2001 sheets answer to: today's, or those of 2001."""
    lines = head_lines("https://www.statistik.at/verzeichnis/reglisten/gemliste_knz.csv", 4)
    codes = [m.group(1) for line in lines for m in [re.match(r"^\s*\d;[^;]*;(\d{5});", line)] if m]
    if not codes:
        codes = [m.group(0) for line in lines for m in [re.search(r"\b[1-9]\d{4}\b", line)] if m]
    log(f"   {len(codes)} codes read from today's list")
    picks = ([c for c in codes if c[:3] in ("620", "621", "622", "623")][:4]
             + [c for c in codes if c.startswith("307")][-4:]
             + ["32419", "60901", "61101", "62006", "10101"])
    for code in picks:
        status, head, body = fetch(f"https://www.statistik.at/blickgem/vz7/g{code}.pdf")
        first = ""
        if status == 200 and body.startswith(b"%PDF"):
            import io as _io
            from pypdf import PdfReader
            first = " ".join((PdfReader(_io.BytesIO(body)).pages[0].extract_text() or "").split()[:6])
        log(f"   vz7/g{code}: HTTP {status}, {len(body):,} bytes; {first}")


def book_dump(url: str, rows: int = 14, cols: int = 16, grep: str | None = None,
              limit: int = 12) -> None:
    """An .xlsx or .xls workbook's sheets, first rows, and rows matching ``grep``."""
    import io as _io
    status, head, body = fetch(url)
    disp = head.get("Content-Disposition") or head.get("content-disposition") or ""
    log(f"\n## {url}: HTTP {status}, {len(body):,} bytes, {disp[:100]}")
    if status != 200:
        log("   " + text(body)[:200])
        return
    sheets: list[tuple[str, list[list[Any]]]] = []
    if body[:2] == b"PK":
        import openpyxl
        book = openpyxl.load_workbook(_io.BytesIO(body), read_only=True, data_only=True)
        for sh in book.worksheets:
            sheets.append((sh.title, [list(r) for r in sh.iter_rows(values_only=True)]))
    elif body[:4] == b"\xd0\xcf\x11\xe0":
        import xlrd
        book = xlrd.open_workbook(file_contents=body)
        for sh in book.sheets():
            sheets.append((sh.name, [sh.row_values(i) for i in range(sh.nrows)]))
    else:
        log("   neither xlsx nor xls: " + text(body)[:200])
        return
    log(f"   sheets: {[name for name, _ in sheets][:20]} ({len(sheets)})")
    for name, table_rows in sheets[:3]:
        log(f"   sheet {name!r}: {len(table_rows)} rows")
        for row in table_rows[:rows]:
            cells = [str(c)[:26] for c in row[:cols] if c not in (None, "")]
            if cells:
                log("     " + " | ".join(cells))
        if grep:
            hits = [r for r in table_rows if any(re.search(grep, str(c)) for c in r[:4])]
            log(f"   {len(hits)} rows match {grep!r}")
            for row in hits[:limit]:
                log("     > " + " | ".join(str(c)[:22] for c in row[:cols] if c not in (None, "")))


def che12() -> None:
    """The 2000 census's religion and main language by commune, and religion by canton since 2010."""
    dam = "https://dam-api.bfs.admin.ch/hub/api/dam/assets/{}/master"
    book_dump(dam.format(193515), grep=r"Bezirk|Amt|District|Distretto|^\s*\.{3,}|^\d{1,4}\s", limit=10)
    book_dump(dam.format(147501), grep=r"Bezirk|Amt|District|Distretto", limit=6)
    book_dump(dam.format(36347568), rows=30)


def round13() -> None:
    """Luxembourg's 2021 language tables, Zensus 2022 religion off the blocked database host,
    and Statistik Austria's 2021 religion survey."""
    links = r"href=\"[^\"]*(?:xlsx?|csv|ods|pdf|lingu|langu|religi|Religi|rp08|zensus|Zensus)[^\"]*\""
    for url in ("https://statistiques.public.lu/fr/recherche.html?q=langue%20principale%20commune",
                "https://statistiques.public.lu/fr/recensement.html",
                "https://statistiques.public.lu/fr/publications/series/recensement-population.html",
                "https://www.zensus2022.de/DE/Ergebnisse-des-Zensus/_inhalt.html",
                "https://www.zensus2022.de/DE/Aktuelles/Religion.html",
                "https://www.statistik.rlp.de/gesellschaft-staat/bevoelkerung-und-gebiet/zensus-2022",
                "https://www.statistik.at/suche?tx_solr%5Bq%5D=Religionszugeh%C3%B6rigkeit",
                "https://www.statistik.at/statistiken/bevoelkerung-und-soziales/bevoelkerung/"
                "bevoelkerungsstand/historische-volkszaehlungen"):
        show(url, links, limit=30)


def ods_dump(url: str, rows: int = 30, cols: int = 12) -> None:
    """An OpenDocument spreadsheet's sheets and first rows (content.xml read directly)."""
    import io as _io
    import zipfile
    import xml.etree.ElementTree as ET
    status, _, body = fetch(url)
    log(f"\n## {url}: HTTP {status}, {len(body):,} bytes")
    if status != 200 or body[:2] != b"PK":
        return
    ns = {"table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
          "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0"}
    root = ET.fromstring(zipfile.ZipFile(_io.BytesIO(body)).read("content.xml"))
    for sheet in root.iter(f"{{{ns['table']}}}table"):
        name_key = "{%s}name" % ns["table"]
        log(f"   sheet {sheet.get(name_key)!r}")
        shown = 0
        for row in sheet.iter(f"{{{ns['table']}}}table-row"):
            cells = []
            for cell in row.iter(f"{{{ns['table']}}}table-cell"):
                words = " ".join("".join(p.itertext()) for p in cell.iter(f"{{{ns['text']}}}p"))
                repeat = int(cell.get(f"{{{ns['table']}}}number-columns-repeated", "1"))
                cells += [words] * min(repeat, 3)
            cells = [c[:24] for c in cells if c][:cols]
            if cells:
                log("     " + " | ".join(cells))
                shown += 1
            if shown >= rows:
                break


def aut9() -> None:
    """Statistik Austria's page on religious affiliation (the 2021 survey) and its tables."""
    page = show("https://www.statistik.at/statistiken/bevoelkerung-und-soziales/bevoelkerung/"
                "weiterfuehrende-bevoelkerungsstatistiken/religionsbekenntnis",
                r"href=\"[^\"]+\.(?:ods|xlsx?|pdf|csv)\"|<p>[^<]{40,400}</p>", limit=30)
    for link in re.findall(r"href=\"([^\"]+\.(?:ods|xlsx))\"", page)[:6]:
        ods_dump(urllib.parse.urljoin("https://www.statistik.at/", link), rows=40)
    ods_dump("https://www.statistik.at/fileadmin/pages/402/Religion.ods", rows=24)


def deu9() -> None:
    """Rhineland-Palatinate's own Zensus 2022 results, for its three statistical regions."""
    for url in ("https://www.statistik.rlp.de/themen/zensus/ergebnisse",
                "https://www.statistik.rlp.de/themen/zensus"):
        page = show(url, r"href=\"[^\"]*(?:\.xlsx?|\.pdf|\.csv|religi|Religi|kreis|Kreis)[^\"]*\"",
                    limit=40)
        for link in re.findall(r"href=\"([^\"]*(?:ergebnisse|zensus)[^\"]*)\"", page)[:12]:
            full = urllib.parse.urljoin(url, link)
            if full.rstrip("/") != url.rstrip("/"):
                show(full, r"href=\"[^\"]*(?:\.xlsx?|\.pdf|\.csv)\"|[^>]{0,80}Religion[^<]{0,80}",
                     limit=15)


def deu10() -> None:
    """Rhineland-Palatinate's Zensus 2022 regional population table: is religion in it?"""
    base = "https://www.statistik.rlp.de/fileadmin/statistik.rlp.de/Dokumente_und_Bilder/1_Themen/3_Zensus/"
    for name in ("07_RP_Regionaltabelle_Bevoelkerung_Z22.xlsx", "07_RP_Regionaltabelle_Demografie_Z22.xlsx",
                 "Eckzahlen.xlsx"):
        book_dump(base + name, rows=12, cols=14, grep=r"(?i)religi|kathol|evangel|Koblenz|Trier|Rheinhessen",
                  limit=8)


def deu11() -> None:
    """The header of the RLP Zensus population sheet, and its Kreis rows."""
    import io as _io
    import openpyxl
    url = ("https://www.statistik.rlp.de/fileadmin/statistik.rlp.de/Dokumente_und_Bilder/1_Themen/"
           "3_Zensus/07_RP_Regionaltabelle_Bevoelkerung_Z22.xlsx")
    status, _, body = fetch(url)
    log(f"\n## {url}: HTTP {status}, {len(body):,} bytes")
    book = openpyxl.load_workbook(_io.BytesIO(body), read_only=True, data_only=True)
    for name in ("Bevölkerung", "Bevölkerung_nachrichtlich"):
        rows = [list(r) for r in book[name].iter_rows(values_only=True)]
        log(f"   sheet {name!r}: {len(rows)} rows x {max(len(r) for r in rows)} cols")
        for r in rows[:9]:
            log("     " + " | ".join(str(c)[:30].replace("\n", " ") for c in r if c is not None))
        for r in rows:
            if r and str(r[0]).strip() in ("07", "071", "072", "073", "07111", "07211", "07311"):
                log("     > " + " | ".join(str(c)[:14] for c in r[:60] if c is not None))


def aut10() -> None:
    """Municipality 62280: what it is, and which Ein-Blick sheets exist for it."""
    head_lines("https://www.statistik.at/verzeichnis/reglisten/gemliste_knz.csv", 1, grep=r"^6228|^622[0-9]{2};",
               limit=60)
    for section in ("vz1", "vz2", "vz3", "vz7", "blick1", "blick2", "blick3", "ae1", "ae2"):
        status, head, body = fetch(f"https://www.statistik.at/blickgem/{section}/g62280.pdf")
        first = ""
        if body.startswith(b"%PDF"):
            import io as _io
            from pypdf import PdfReader
            first = " | ".join((PdfReader(_io.BytesIO(body)).pages[0].extract_text() or "").split("\n")[:14])
        log(f"   {section}/g62280: HTTP {status}, {len(body):,} bytes, "
            f"{head.get('Content-Type') or head.get('content-type')}; {first[:700]}")


def hun14() -> None:
    """Fejér's settlements as KSH's 2022 codelist files them, and where the 2013 division is published."""
    base = "https://nepszamlalas2022.ksh.hu/api"
    version = json.loads(text(fetch(f"{base}/version")[2]))["version"]
    data = json.loads(text(fetch(f"{base}/structure/WBS003/{version}")[2]))
    codes = next(cl["codes"] for cl in data["data"]["codelists"] if cl["id"] == "CL_TERUL_GEO5")
    name = {c["id"]: jpath(c, "names.hu") or jpath(c, "names.en") or c.get("name") for c in codes}
    fejer = [c["id"] for c in codes if c.get("parent") == "HU211"]
    log(f"   Fejér's districts: {[(d, name[d]) for d in fejer]}")
    named = [c["id"] for c in codes if len(c["id"]) == 3
             and re.search(r"(?i)enying|fehérvár|polgárd", name[c["id"]] or "")]
    for d in named:
        kids = [(c["id"], name[c["id"]]) for c in codes if c.get("parent") == d]
        log(f"   {d} {name[d]}: {len(kids)} settlements: {kids}")
    for url in ("https://www.ksh.hu/docs/helysegnevtar/hnt_letoltes_2014.xls",
                "https://www.ksh.hu/docs/helysegnevtar/hnt_letoltes_2013.xls",
                "https://www.ksh.hu/docs/helysegnevtar/hnt_letoltes_2014.xlsx"):
        status, head, body = fetch(url)
        log(f"   {url}: HTTP {status}, {len(body):,} bytes, {head.get('Content-Type') or head.get('error', '')}")
        if status == 200 and body[:4] in (b"\xd0\xcf\x11\xe0", b"PK\x03\x04"):
            book_dump(url, rows=6, grep=r"(?i)polg[aá]rd", limit=20)
            return
    show("https://web.archive.org/cdx/search/cdx?url=ksh.hu/docs/helysegnevtar/hnt_letoltes_201*"
         "&output=txt&limit=40", r"\S+ \d{14} \S+ \S+ \d{3}", limit=40)
    for url in ("https://njt.hu/jogszabaly/2012-218-20-22", "https://net.jogtar.hu/jogszabaly?docid=a1200218.kor"):
        page = show(url)
        for m in list(re.finditer(r"Polgárdi", page))[:3]:
            log("   ..." + " ".join(re.sub(r"<[^>]+>", " ", page[max(0, m.start() - 200):m.start() + 700]).split()))


def che13() -> None:
    """BFS tables of the resident population by commune on an older commune state; Höfe's sexes."""
    for query in ("language=de&title=St%C3%A4ndige%20Wohnbev%C3%B6lkerung%20Gemeinden",
                  "language=de&title=Bilanz%20der%20st%C3%A4ndigen%20Wohnbev%C3%B6lkerung",
                  "language=de&title=Regionalportr%C3%A4ts",
                  "language=de&title=Bev%C3%B6lkerung%20Gemeinden%202009",
                  "language=de&title=Strukturerhebung%20Stichprobe"):
        items = dam_items(query, pages=5)
        log(f"\n## DAM {query}: {len(items)} items")
        for item in items:
            desc = item.get("description") or {}
            period = jpath(desc, "bibliography.period") or ""
            title = jpath(desc, "titles.main") or ""
            if "Stichprobe" not in query and not re.search(r"200[89]|2010|2011", f"{period} {title}"):
                continue
            log(f"   - {jpath(item, 'ids.damId')} | {jpath(item, 'shop.orderNr')} | {period} | "
                f"{jpath(item, 'bfs.articleModel.name')} | {title}"[:260])
    table = "https://www.pxweb.bfs.admin.ch/api/v1/de/px-x-0102010000_101/px-x-0102010000_101.px"
    meta = {v["code"]: v for v in json.loads(text(fetch(table)[2]))["variables"]}
    code = {k: next(c for c in meta if c.startswith(k)) for k in
            ("Jahr", "Kanton", "Bevölkerungstyp", "Staatsangehörigkeit", "Geschlecht", "Alter")}
    places = dict(zip(meta[code["Kanton"]]["values"], meta[code["Kanton"]]["valueTexts"]))
    hoefe = [k for k, v in places.items() if re.search(r"Höfe|Wollerau|Freienbach|Feusisberg", v)]
    log(f"   Höfe entries: {[(k, places[k]) for k in hoefe]}")
    query = [{"code": code["Jahr"], "selection": {"filter": "item", "values": ["2025"]}},
             {"code": code["Kanton"], "selection": {"filter": "item", "values": hoefe}},
             {"code": code["Bevölkerungstyp"], "selection": {"filter": "item", "values": ["1"]}},
             {"code": code["Staatsangehörigkeit"], "selection": {"filter": "all", "values": ["*"]}},
             {"code": code["Geschlecht"], "selection": {"filter": "all", "values": ["*"]}},
             {"code": code["Alter"], "selection": {"filter": "item", "values": ["-99999"]}}]
    body = json.dumps({"query": query, "response": {"format": "json-stat2"}}).encode()
    status, _, reply = fetch(table, data=body, headers={"Content-Type": "application/json"})
    try:
        from .pxweb import unstack
        for key, value in unstack(json.loads(text(reply))):
            log(f"   {places[key[code['Kanton']][0]]} | nat {key[code['Staatsangehörigkeit']][0]} | "
                f"sex {key[code['Geschlecht']][0]}: {value:,.0f}")
    except Exception as exc:  # noqa: BLE001
        log(f"   HTTP {status}: {exc}: {text(reply)[:300]}")


def lux7() -> None:
    """Official population figures for the six communes merged in 2018, before the merger."""
    page = show("https://lustat.statec.lu/rest/dataflow/LU1/all/latest")
    flows = re.findall(r'<structure:Dataflow id="([^"]+)"[^>]*>.*?<common:Name xml:lang="en">([^<]+)',
                       page, flags=re.S)
    log(f"   {len(flows)} dataflows; those naming communes, localities or the 2011 census:")
    for fid, title in flows:
        if re.search(r"(?i)commun|municip|localit|2011|canton", title):
            log(f"   - {fid}: {title[:160]}")
    for q in ("population+commune", "population+localite", "recensement+2011"):
        status, _, body = fetch(f"https://data.public.lu/api/1/datasets/?q={q}&page_size=12")
        try:
            data = json.loads(text(body))
        except json.JSONDecodeError:
            log(f"   data.public.lu {q}: HTTP {status}: {text(body)[:200]}")
            continue
        log(f"\n## data.public.lu {q}: {data.get('total')} datasets")
        for ds in data.get("data", []):
            org = jpath(ds, "organization.name") or ""
            log(f"   - {ds.get('slug')} | {org} | {ds.get('title')}"[:220])
            for res in (ds.get("resources") or [])[:6]:
                log(f"       {res.get('format')} {res.get('title')} {res.get('url')}"[:220])
    show("https://web.archive.org/cdx/search/cdx?url=statistiques.public.lu/stat/TableViewer/tableView*"
         "&matchType=prefix&filter=original:.*ReportId=12859.*&output=txt&limit=12",
         r"\S+ \d{14} \S+", limit=12)


def pol5() -> None:
    """BDL variables for the usual-resident population ('ludność rezydująca'), and at which levels."""
    for term in ("rezyduj", "rezydenci"):
        for kind in ("variables", "subjects"):
            status, _, body = fetch(f"https://bdl.stat.gov.pl/api/v1/{kind}/search?name={term}&format=json"
                                    f"&page-size=50&lang=pl")
            try:
                data = json.loads(text(body))
            except json.JSONDecodeError:
                log(f"   {kind} {term}: HTTP {status}: {text(body)[:200]}")
                continue
            rows = data.get("results") or []
            log(f"\n## BDL {kind} search {term!r}: {data.get('totalRecords')} ({len(rows)} shown)")
            for r in rows[:40]:
                log(f"   - {r.get('id')} | {r.get('subjectId') or r.get('parentId')} | levels "
                    f"{r.get('levels')} | {r.get('n1') or r.get('name')} | {r.get('n2', '')} | "
                    f"{r.get('n3', '')}"[:240])


def nld11() -> None:
    """How many respondents stand behind each province in CBS's religion workbook."""
    book_dump("https://www.cbs.nl/-/media/_excel/2026/11/religie_2025_tabellen.xlsx", rows=40,
              grep=r"(?i)respondent|waarnem|steekproef|ongewogen|aantal|onderzoek|enquête|bron",
              limit=30)
    show("https://www.cbs.nl/nl-nl/maatwerk/2026/11/religie-naar-regio-2021-2025",
         r"(?i)[^.<>]{0,200}(?:respondent|steekproef|enquête|waarneming|ondervraagd)[^.<>]{0,200}", limit=12)


def aut11() -> None:
    """What sample Statistik Austria's 2021 religion estimate rests on."""
    url = ("https://www.statistik.at/statistiken/bevoelkerung-und-soziales/bevoelkerung/"
           "weiterfuehrende-bevoelkerungsstatistiken/religionsbekenntnis")
    page = show(url, r"href=\"[^\"]*(?:\.pdf|\.ods|\.xlsx|presse|Religion)[^\"]*\"", limit=30)
    plain = " ".join(re.sub(r"<[^>]+>", " ", page).split())
    for m in list(re.finditer(r"(?i)stichprob|befragt|erhebung|mikrozensus|personen", plain))[:10]:
        log("   ..." + plain[max(0, m.start() - 250):m.start() + 250])
    status, _, body = fetch("https://www.statistik.at/fileadmin/pages/439/neu__Religion_2021_Bundesland.ods")
    if status == 200:
        import zipfile
        import io as _io
        content = zipfile.ZipFile(_io.BytesIO(body)).read("content.xml").decode("utf-8")
        words = " ".join(re.sub(r"<[^>]+>", " ", content).split())
        for m in list(re.finditer(r"(?i)quelle|stichprob|befragt|erhebung|hochgerechnet|anmerkung|n\s*=",
                                  words))[:8]:
            log("   ods ..." + words[max(0, m.start() - 120):m.start() + 380])


def svn7() -> None:
    """SiStat's tables by settlement, for Šentrupert's sex ratio (the prison at Dob)."""
    status, _, body = fetch("https://pxweb.stat.si/SiStatData/api/v1/sl/Data/")
    try:
        rows = json.loads(text(body))
    except json.JSONDecodeError:
        log(f"   HTTP {status}: {text(body)[:200]}")
        return
    hits = [r for r in rows if re.search(r"(?i)naselj", str(r.get("text")))]
    log(f"   {len(rows)} tables; {len(hits)} by settlement:")
    for r in hits[:25]:
        log(f"   - {r.get('id')} | {r.get('text')}"[:200])


def pdf_grep(url: str, pattern: str, *, width: int = 300, limit: int = 20) -> None:
    """Every match of ``pattern`` in a PDF's text, with the words around it."""
    status, _, body = fetch(url)
    log(f"\n## PDF {url}: HTTP {status}, {len(body):,} bytes")
    if status != 200 or not body.startswith(b"%PDF"):
        return
    import io as _io
    from pypdf import PdfReader
    reader = PdfReader(_io.BytesIO(body))
    words = " ".join(" ".join((p.extract_text() or "").split()) for p in reader.pages)
    hits = list(re.finditer(pattern, words))
    log(f"   {len(reader.pages)} pages, {len(hits)} matches of {pattern!r}")
    for m in hits[:limit]:
        log("   ..." + words[max(0, m.start() - width):m.start() + width])


def sheet_text(url: str, sheets: tuple[str, ...], limit: int = 60) -> None:
    """Every non-empty row of the named sheets, cells with their column index."""
    import io as _io
    import openpyxl
    status, _, body = fetch(url)
    log(f"\n## {url}: HTTP {status}, {len(body):,} bytes")
    if status != 200:
        return
    book = openpyxl.load_workbook(_io.BytesIO(body), read_only=True, data_only=True)
    for name in sheets:
        if name not in book.sheetnames:
            log(f"   no sheet {name!r}; {book.sheetnames}")
            continue
        rows = [r for r in book[name].iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
        log(f"   sheet {name!r}: {len(rows)} rows")
        for r in rows[:limit]:
            log("     " + " | ".join(f"[{j}] {' '.join(str(c).split())[:400]}" for j, c in enumerate(r)
                                   if c not in (None, "")))


def aut12() -> None:
    """The sample behind the 2021 religion questions, from Statistik Austria's standard documentation."""
    pdf_grep("https://www.statistik.at/fileadmin/shared/QM/Standarddokumentationen/B_2/"
             "std_b_religionzugehoerigkeit.pdf",
             r"(?i)stichprob|befragt|respondent|bundesl|haushalte|personen|ausschöpf|rücklauf|n\s*=",
             width=260, limit=30)


def nld12() -> None:
    """CBS's workbook's own explanation: the survey, its sample, and what 'Ander geloof' holds."""
    sheet_text("https://www.cbs.nl/-/media/_excel/2026/11/religie_2025_tabellen.xlsx",
               ("Introductie", "Toelichting", "Begrippen"), limit=60)


def che14() -> None:
    """The structural survey's sample by canton, and BFS population tables on an older commune state."""
    for asset in (36432820, 35167595, 31425811):
        book_dump(f"https://dam-api.bfs.admin.ch/hub/api/dam/assets/{asset}/master", rows=45, cols=12)
    for order in ("su-d-01.02.04.08", "su-d-01.02.04.10", "hs-d-01.02.02.02.09.20",
                  "hs-d-01.02.02.02.09.21", "hs-d-01.02.02.02.09.22", "hs-d-01.02.02.02.09.26",
                  "hs-d-01.02.02.01.20", "je-d-01.02.03.01"):
        for item in dam_items(f"orderNr={order}", pages=1):
            desc = item.get("description") or {}
            log(f"   {order}: {jpath(item, 'ids.damId')} | {jpath(desc, 'bibliography.period')} | "
                f"{jpath(desc, 'titles.main')}"[:240])
    for item in dam_items("language=de&title=Wohnbev%C3%B6lkerung%20nach%20Bezirken%20und%20Gemeinden",
                          pages=3):
        desc = item.get("description") or {}
        log(f"   - {jpath(item, 'ids.damId')} | {jpath(item, 'shop.orderNr')} | "
            f"{jpath(desc, 'bibliography.period')} | {jpath(desc, 'titles.main')}"[:240])


def lux8() -> None:
    """STATEC's population by commune 2000-2017 and the 2011 census by commune, age and sex."""
    grep = r"(?i)rosport|mompach|hobscheid|septfontaines|boevange|tuntange|total|luxembourg$|^commune"
    for url in ("https://download.data.public.lu/resources/population-par-commune-et-code-lau2-depuis-2000/"
                "20170511-141118/popcom2000-2017_LAU2.xlsx",
                "https://download.data.public.lu/resources/population-de-residence-habituelle-par-commune-et-age-"
                "au-1er-fevrier-2011/20160711-131312/Population_par_commune_et_age_au_1er_fevrier_2011.xlsx",
                "https://download.data.public.lu/resources/population-de-residence-habituelle-par-commune-et-sexe-"
                "au-1er-fevrier-2011/20160711-131552/Population_par_commune_et_sexe_au_1er_fevrier_2011.xlsx"):
        book_dump(url, rows=8, cols=26, grep=grep, limit=14)


def svn8() -> None:
    """Šentrupert and the settlement of Dob pri Mirni by sex (the prison at Dob)."""
    url = "https://pxweb.stat.si/SiStatData/api/v1/sl/Data/05C5003S.px"
    meta = pxweb_meta(url, values=6)
    if not meta:
        return
    var = {v["code"]: v for v in meta["variables"]}
    place = next(c for c, v in var.items() if re.search(r"(?i)obč|nasel", v["text"]))
    hits = [(v, t) for v, t in zip(var[place]["values"], var[place]["valueTexts"])
            if re.search(r"(?i)šentrupert|dob pri mirni|^\s*dob\b", t)]
    log(f"   {place}: {hits}")
    query = []
    for code, v in var.items():
        if code == place:
            query.append({"code": code, "selection": {"filter": "item", "values": [h[0] for h in hits]}})
        elif re.search(r"(?i)spol", v["text"]):
            query.append({"code": code, "selection": {"filter": "all", "values": ["*"]}})
        elif re.search(r"(?i)leto|obdobje|čas", v["text"]) or v.get("time"):
            query.append({"code": code, "selection": {"filter": "top", "values": ["1"]}})
        else:
            query.append({"code": code, "selection": {"filter": "item", "values": [v["values"][0]]}})
    body = json.dumps({"query": query, "response": {"format": "json-stat2"}}).encode()
    status, _, reply = fetch(url, data=body, headers={"Content-Type": "application/json"})
    try:
        from .pxweb import unstack
        for key, value in unstack(json.loads(text(reply))):
            log(f"   {key}: {value}")
    except Exception as exc:  # noqa: BLE001
        log(f"   HTTP {status}: {exc}: {text(reply)[:300]}")


def hun15() -> None:
    """The 2014 gazetteer's header with column numbers, and one settlement's WBS003 cells."""
    import xlrd
    status, _, body = fetch("https://www.ksh.hu/docs/helysegnevtar/hnt_letoltes_2014.xls")
    book = xlrd.open_workbook(file_contents=body)
    sheet = book.sheet_by_index(0)
    for i in range(4):
        log(f"   row {i}: " + " | ".join(f"[{j}] {' '.join(str(c).split())}" for j, c in
                                        enumerate(sheet.row_values(i)) if c not in (None, "")))
    counts = Counter_(str(sheet.cell_value(i, 0)).strip() != "" for i in range(sheet.nrows))
    log(f"   {sheet.nrows} rows; non-empty first cells {counts}")
    base = "https://nepszamlalas2022.ksh.hu/api"
    version = json.loads(text(fetch(f"{base}/version")[2]))["version"]
    rows = json.loads(text(fetch(f"{base}/dataflows/WBS003/{version}/d/TIME_PERIOD:2022,"
                                 f"TERUL_GEO5:17525+02802,TEL_SZ_ADAT")[2]))
    log(f"   WBS003 for Polgárdi and Enying: {len(rows)} rows")
    for r in rows:
        if r["TERUL_GEO5"] == "17525":
            log(f"     {r['TEL_SZ_ADAT']}={r.get('OBS_VALUE')}")


def che15() -> None:
    """BFS's commune balance 1991-2025: one sheet a year, and whose communes each sheet lists."""
    import io as _io
    import openpyxl
    status, head, body = fetch("https://dam-api.bfs.admin.ch/hub/api/dam/assets/36681837/master", timeout=300)
    log(f"   HTTP {status}, {len(body):,} bytes, {head.get('Content-Disposition') or ''}")
    book = openpyxl.load_workbook(_io.BytesIO(body), read_only=True, data_only=True)
    log(f"   sheets: {book.sheetnames}")
    grep = re.compile(r"(?i)peseux|corcelles|^\W*neuchâtel|giubiasco|bellinzona|haldenstein|thun|"
                      r"busswil|lyss|ederswiler|clavaleyres|bezirk thun|amtsbezirk")
    for name in [n for n in book.sheetnames if re.search(r"2009|2010|2016|2020|2025", n)][:5]:
        rows = list(book[name].iter_rows(values_only=True))
        log(f"\n   sheet {name!r}: {len(rows)} rows")
        for r in rows[:7]:
            log("     " + " | ".join(f"[{j}] {' '.join(str(c).split())[:40]}" for j, c in enumerate(r)
                                   if c not in (None, "")))
        codes = [r for r in rows if r and re.match(r"^\s*\.*\s*\d{1,4}\b", str(r[0] or ""))]
        log(f"   rows starting with a number: {len(codes)}")
        for r in [r for r in rows if r and any(grep.search(str(c or "")) for c in r[:2])][:14]:
            log("     > " + " | ".join(f"[{j}] {' '.join(str(c).split())[:30]}" for j, c in enumerate(r[:8])
                                     if c not in (None, "")))


def nld13() -> None:
    """The survey question's categories in full, from the workbook's 'Begrippen'."""
    import io as _io
    import openpyxl
    status, _, body = fetch("https://www.cbs.nl/-/media/_excel/2026/11/religie_2025_tabellen.xlsx")
    book = openpyxl.load_workbook(_io.BytesIO(body), read_only=True, data_only=True)
    for r in book["Begrippen"].iter_rows(values_only=True):
        if r and str(r[0] or "").startswith("Gelovigen"):
            log("   " + " ".join(str(r[1]).split()))
    for r in book["Tabel 2"].iter_rows(values_only=True):
        cells = [str(c) for c in r if c not in (None, "")]
        if cells and not re.match(r"^\d", cells[-1] if cells else ""):
            log("   T2: " + " | ".join(" ".join(c.split())[:160] for c in cells))


def svn9() -> None:
    """Šentrupert's settlements by sex in 2026: where its men outnumber its women."""
    url = "https://pxweb.stat.si/SiStatData/api/v1/sl/Data/05C5003S.px"
    meta = json.loads(text(fetch(url)[2]))
    var = {v["code"]: v for v in meta["variables"]}
    place = next(c for c, v in var.items() if re.search(r"(?i)obč|nasel", v["text"]))
    year = next(c for c, v in var.items() if re.search(r"(?i)leto", v["text"]))
    measure = next(c for c, v in var.items() if re.search(r"(?i)meritve", v["text"]))
    wanted = [v for v in var[place]["values"] if v == "211" or (v.startswith("211") and len(v) == 6)]
    query = [{"code": place, "selection": {"filter": "item", "values": wanted}},
             {"code": year, "selection": {"filter": "item", "values": ["2026"]}},
             {"code": measure, "selection": {"filter": "item", "values": ["0", "1", "2"]}}]
    body = json.dumps({"query": query, "response": {"format": "json-stat2"}}).encode()
    reply = json.loads(text(fetch(url, data=body, headers={"Content-Type": "application/json"})[2]))
    from .pxweb import unstack
    table: dict[str, dict[str, float]] = {}
    for key, value in unstack(reply):
        table.setdefault(key[place][1], {})[key[measure][0]] = value
    for name, v in sorted(table.items(), key=lambda kv: -(kv[1].get("1", 0) - kv[1].get("2", 0)))[:8]:
        log(f"   {name}: total {v.get('0')}, men {v.get('1')}, women {v.get('2')}")


def che16() -> None:
    """The 2009 balance sheet's rows the register lacks, and the 2010 mergers' predecessors in it."""
    import io as _io
    import openpyxl
    status, _, body = fetch("https://dam-api.bfs.admin.ch/hub/api/dam/assets/36681837/master", timeout=300)
    book = openpyxl.load_workbook(_io.BytesIO(body), read_only=True, data_only=True)
    rows = list(book["2009"].iter_rows(values_only=True))
    want = re.compile(r"^\.*\s*0*(4505|4670|4695|5269|5397|4161|4163|4162|4164|4165|4166|5171|5173|5174|"
                      r"5175|5176|5003|5004|5001)\b")
    context = None
    for r in rows:
        first = str(r[0] or "").strip()
        if not re.match(r"^\.*\s*\d", first):
            context = first
        if want.match(first):
            log(f"   {first} | 1 Jan {r[1]} | 31 Dec {r[9]} | under {context}")


# ---------------------------------------------------------------------------
# Generic probes, driven by options rather than by a named function, so that a
# question about one more file costs a dispatch and not a commit. Every
# argument is one whitespace-free token, because the workflow splits its
# command on whitespace.
#
#   fetch URL [URL ...] [--grep RE] [--rows N] [--bytes N] [--context N]
#         [--links RE] [--sheets N] [--sheet-match RE] [--start N] [--cols N]
#         [--width N] [--member RE] [--delim C] [--wayback TS|latest]
#         [--post JSON] [--form K=V&K=V] [--accept TYPE] [--jpath A.B.0]
#       Fetch each URL and describe it: a workbook's sheets and rows, a zip's
#       members (and the CSV or workbook member matching --member), CSV's
#       header and the lines matching --grep, JSON's shape, or a page's
#       links and text.
#   cdx PREFIX [--match RE] [--rows N] [--cdx-filter F]
#       The Internet Archive's captures under PREFIX, one per URL.
# ---------------------------------------------------------------------------

CDX = "https://web.archive.org/cdx/search/cdx"


def _clip(value: Any, width: int) -> str:
    return ("" if value is None else " ".join(str(value).split()))[:width]


def _wayback(url: str, stamp: str) -> str | None:
    if stamp != "latest":
        return f"https://web.archive.org/web/{stamp}id_/{url}"
    query = urllib.parse.urlencode({"url": url, "output": "json", "filter": "statuscode:200",
                                    "fl": "timestamp", "limit": "-1"})
    status, _, body = fetch(f"{CDX}?{query}", timeout=120)
    try:
        rows = json.loads(text(body) or "[]")
    except json.JSONDecodeError:
        rows = []
    if status != 200 or len(rows) < 2:
        log(f"   no capture of {url} (HTTP {status})")
        return None
    return f"https://web.archive.org/web/{rows[-1][0]}id_/{url}"


def _book_rows(blob: bytes) -> list[tuple[str, list[list[Any]]]]:
    import io as _io
    if blob[:2] == b"PK":
        import openpyxl
        book = openpyxl.load_workbook(_io.BytesIO(blob), read_only=True, data_only=True)
        return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)])
                for ws in book.worksheets]
    import xlrd
    book = xlrd.open_workbook(file_contents=blob)
    return [(ws.name, [ws.row_values(i) for i in range(ws.nrows)]) for ws in book.sheets()]


def _show_book(blob: bytes, args: Any) -> None:
    sheets = _book_rows(blob)
    log(f"   {len(sheets)} sheets: " + " | ".join(t for t, _ in sheets[:60]))
    if args.sheet_match:
        sheets = [s for s in sheets if re.search(args.sheet_match, s[0], re.I)]
    grep = re.compile(args.grep, re.I) if args.grep else None
    for title, rows in sheets[:args.sheets]:
        width = max((len(r) for r in rows), default=0)
        log(f"   -- sheet {title!r}: {len(rows)} rows x {width} cols")
        shown = 0
        for i, row in enumerate(rows):
            if i < args.start:
                continue
            line = " | ".join(_clip(c, args.width) for c in row[:args.cols])
            if grep and not grep.search(line) and i >= args.start + args.head:
                continue
            log(f"     r{i}: {line}")
            shown += 1
            if shown >= args.rows:
                break


def _show_csv(page: str, args: Any) -> None:
    lines = page.splitlines()
    log(f"   {len(lines):,} lines")
    for line in lines[:args.head or 2]:
        log("   | " + line[:args.bytes])
    if args.grep:
        grep = re.compile(args.grep, re.I)
        hits = [line for line in lines[1:] if grep.search(line)]
        log(f"   {len(hits)} lines match {args.grep!r}")
        for line in hits[:args.rows]:
            log("   > " + line[:args.bytes])


def _show_zip(blob: bytes, args: Any) -> None:
    import io as _io
    import zipfile
    with zipfile.ZipFile(_io.BytesIO(blob)) as archive:
        infos = archive.infolist()
        log(f"   zip, {len(infos)} members")
        for info in infos[:args.rows]:
            log(f"     {info.filename} ({info.file_size:,})")
        if not args.member:
            return
        for info in infos:
            if re.search(args.member, info.filename, re.I):
                data = archive.read(info)
                log(f"   member {info.filename}:")
                _describe(data, info.filename, args)
                return
        log(f"   no member matches {args.member!r}")


def _describe(blob: bytes, url: str, args: Any) -> None:
    lower = url.lower().split("?")[0]
    if blob[:2] == b"PK":
        import io as _io
        import zipfile
        try:
            names = zipfile.ZipFile(_io.BytesIO(blob)).namelist()
        except zipfile.BadZipFile:
            names = []
        if "[Content_Types].xml" in names and not lower.endswith(".ods"):
            _show_book(blob, args)
        else:
            _show_zip(blob, args)
        return
    if blob[:8] == bytes.fromhex("d0cf11e0a1b11ae1"):
        _show_book(blob, args)
        return
    if blob[:4] == b"%PDF":
        log(f"   a PDF of {len(blob):,} bytes")
        return
    page = text(blob)
    stripped = page.lstrip("﻿").lstrip()
    if stripped.startswith(("{", "[")) and not args.grep:
        try:
            data = json.loads(stripped)
        except ValueError:
            data = None
        if data is not None:
            if args.jpath:
                data = jpath(data, args.jpath)
            log("   json: " + json.dumps(data, ensure_ascii=False)[:args.bytes])
            return
    if lower.endswith((".csv", ".txt", ".tsv")) or (args.delim and "<html" not in page[:400].lower()):
        _show_csv(page, args)
        return
    if args.grep:
        hits = list(re.finditer(args.grep, page, re.I))
        log(f"   {len(hits)} matches of {args.grep!r}")
        for m in hits[:args.rows]:
            s = max(0, m.start() - args.context)
            log("   ~ " + " ".join(page[s:m.end() + args.context].split())[:args.bytes])
        return
    links = re.findall(r"""(?is)<a\b[^>]*href\s*=\s*["']([^"']+)["'][^>]*>(.*?)</a>""", page)
    if links:
        match = re.compile(args.links, re.I) if args.links else None
        shown = 0
        for href, label in links:
            label = " ".join(re.sub(r"<[^>]+>", " ", label).split())
            if match and not (match.search(href) or match.search(label)):
                continue
            log(f"   link: {label[:90]!r} -> {urllib.parse.urljoin(url, href)}")
            shown += 1
            if shown >= args.rows:
                break
        log(f"   ({len(links)} links in all)")
    words = " ".join(re.sub(r"(?is)<(script|style).*?</\1>|<[^>]+>", " ", page).split())
    log("   text: " + words[:args.bytes])


def generic(argv: list[str]) -> int:
    import argparse
    ap = argparse.ArgumentParser(prog="central_probe")
    ap.add_argument("cmd", choices=["fetch", "cdx"])
    ap.add_argument("target", nargs="+")
    for flag, default in (("--rows", 40), ("--bytes", 1500), ("--context", 160),
                          ("--sheets", 2), ("--start", 0), ("--cols", 14), ("--width", 24),
                          ("--head", 2), ("--timeout", TIMEOUT)):
        ap.add_argument(flag, type=int, default=default)
    for flag in ("--grep", "--links", "--sheet-match", "--member", "--delim", "--wayback",
                 "--post", "--form", "--accept", "--jpath", "--match", "--cdx-filter"):
        ap.add_argument(flag)
    args = ap.parse_args(argv)
    if args.cmd == "cdx":
        for prefix in args.target:
            params = [("url", prefix), ("matchType", "prefix"), ("output", "json"),
                      ("filter", "statuscode:200"), ("fl", "original,timestamp,mimetype,length"),
                      ("collapse", "urlkey"), ("limit", "20000")]
            if args.cdx_filter:
                params.append(("filter", args.cdx_filter))
            status, _, body = fetch(f"{CDX}?{urllib.parse.urlencode(params)}", timeout=180)
            try:
                rows = json.loads(text(body) or "[]")
            except json.JSONDecodeError:
                rows = []
            match = re.compile(args.match, re.I) if args.match else None
            shown = 0
            for row in rows[1:]:
                original = urllib.parse.unquote(row[0])
                if match and not match.search(original):
                    continue
                log(f"   {row[1]} {row[2][:24]:24} {row[3]:>9} {original}")
                shown += 1
                if shown >= args.rows:
                    break
            log(f"   cdx {prefix}: HTTP {status}, {shown} shown of {max(len(rows) - 1, 0)}")
        return 0
    for url in args.target:
        target = _wayback(url, args.wayback) if args.wayback else url
        if target is None:
            continue
        data, headers = None, {}
        if args.post:
            data, headers = args.post.encode(), {"Content-Type": "application/json"}
        elif args.form:
            data = args.form.encode()
            headers = {"Content-Type": "application/x-www-form-urlencoded"}
        if args.accept:
            headers["Accept"] = args.accept
        status, head, body = fetch(target, data=data, headers=headers, timeout=args.timeout)
        kind = head.get("Content-Type") or head.get("content-type") or head.get("error", "")
        disp = head.get("Content-Disposition") or head.get("content-disposition") or ""
        log(f"\n## {target}\n   HTTP {status}, {len(body):,} bytes, {kind} {disp[:80]}")
        if body:
            try:
                _describe(body, url, args)
            except Exception as exc:  # noqa: BLE001 - a probe reports, it does not stop
                log(f"   could not describe it: {exc.__class__.__name__}: {exc}")
    return 0


def deu_zensus() -> None:
    """The Zensus 2022 database with the account: who it thinks we are, which tables carry
    citizenship, and which form of the tablefile request it accepts.

    The credentials come from the environment and are scrubbed out of every
    line printed; the run's product is a log that gets committed.
    """
    import os
    base = "https://ergebnisse.zensus2022.de/api/rest/2020"
    user = os.environ.get("ZENSUS_USER", "").strip()
    password = os.environ.get("ZENSUS_PASSWORD", "").strip()
    secrets = [s for s in (user, password) if s]

    def scrub(value: str) -> str:
        for secret in secrets:
            value = value.replace(secret, "***")
        return value

    log(f"   account set: {bool(user)}; password set: {bool(password)}")

    def post(path: str, params: dict[str, str], *, how: str = "headers", show_bytes: int = 600,
             grep: str | None = None) -> tuple[int, bytes]:
        body = dict(params)
        headers = {"Accept": "*/*", "Content-Type": "application/x-www-form-urlencoded"}
        if how == "headers":
            headers.update({"username": user, "password": password})
        elif how == "body":
            body.update({"username": user, "password": password})
        status, head, blob = fetch(f"{base}/{path}", data=urllib.parse.urlencode(body).encode(),
                                   headers=headers, timeout=180)
        if blob[:2] == b"PK":
            import io as _io
            import zipfile
            with zipfile.ZipFile(_io.BytesIO(blob)) as archive:
                blob = archive.read(archive.namelist()[0])
        shown = scrub(text(blob))
        log(f"\n## POST {path} {sorted(params)} ({how}): HTTP {status}, {len(blob):,} bytes, "
            f"{head.get('Content-Type') or head.get('error', '')}")
        if grep:
            hits = re.findall(grep, shown)
            log(f"   {len(hits)} matches of {grep!r}")
            for hit in hits[:80]:
                log("   - " + " ".join(str(hit).split())[:240])
        else:
            log("   " + " ".join(shown[:show_bytes].split()))
        return status, blob

    post("helloworld/logincheck", {}, how="headers")
    post("helloworld/logincheck", {}, how="body")
    title = r'"Code"\s*:\s*"[^"]+"\s*,\s*"Content"\s*:\s*"[^"]{0,160}"'
    for term in ("Staatsangehörigkeit", "Geburtsland", "Migration"):
        post("find/find", {"term": term, "category": "tables", "pagelength": "200",
                           "language": "de"}, grep=title)
    post("catalogue/tables", {"selection": "1000A-*", "area": "all", "pagelength": "500",
                              "language": "de"}, grep=title)
    for variant in ({"regionalvariable": "GEOBL1", "regionalschluessel": ""},
                    {"regionalvariable": "GEOBL1"}, {}):
        params = {"name": "1000A-1018", "area": "all", "format": "ffcsv", "compress": "false",
                  "language": "de", **variant}
        post("data/tablefile", params, show_bytes=400)


PROBES: dict[str, Callable[[], None]] = {
    "deu_zensus": deu_zensus,
    "che16": che16,
    "che15": che15, "nld13": nld13, "svn9": svn9,
    "aut12": aut12, "nld12": nld12, "che14": che14, "lux8": lux8, "svn8": svn8, "hun15": hun15,
    "hun14": hun14, "che13": che13, "lux7": lux7, "pol5": pol5, "nld11": nld11, "aut11": aut11,
    "svn7": svn7,
    "aut10": aut10,
    "deu11": deu11,
    "deu10": deu10,
    "deu9": deu9,
    "aut9": aut9,
    "round13": round13,
    "che12": che12,
    "che11": che11, "aut8": aut8,
    "svn6": svn6, "che10": che10, "aut7": aut7, "deu8": deu8,
    "aut6": aut6, "svn5": svn5, "che9": che9, "nld10": nld10, "lux6": lux6, "deu7": deu7,
    "che8": che8, "svn4": svn4, "aut5": aut5,
    "che7": che7, "nld9": nld9,
    "deu6": deu6,
    "hun13": hun13,
    "deu5": deu5,
    "hun12": hun12,
    "hun11": hun11,
    "hun10": hun10,
    "hun9": hun9, "svk9": svk9,
    "hun8": hun8, "svk8": svk8, "nld8": nld8,
    "hun7": hun7, "svk7": svk7, "nld7": nld7,
    "lie7": lie7,
    "hun6": hun6, "nld6": nld6, "svk6": svk6,
    "cze5": cze5, "hun5": hun5, "che5": che5, "svk5": svk5, "nld5": nld5, "lux5": lux5,
    "cze4": cze4, "pol4": pol4, "svk4": svk4, "hun4": hun4, "nld4": nld4, "lux4": lux4,
    "aut4": aut4, "deu4": deu4, "che4": che4,
    "che3": che3, "lie3": lie3, "pol3": pol3, "svk3": svk3, "hun3": hun3, "svn3": svn3,
    "nld3": nld3, "lux3": lux3, "cze3": cze3, "aut3": aut3, "deu3": deu3,
    "aut": aut, "che": che, "lie": lie, "pol": pol, "cze": cze, "svk": svk,
    "hun": hun, "svn": svn, "nld": nld, "lux": lux,
    "aut2": aut2, "che2": che2, "lie2": lie2, "pol2": pol2, "cze2": cze2, "svk2": svk2,
    "hun2": hun2, "svn2": svn2, "nld2": nld2, "lux2": lux2,
}


def main(argv: list[str]) -> int:
    if not argv:
        print(__doc__)
        return 0
    if argv[0] == "get":
        show(argv[1], argv[2] if len(argv) > 2 else None, raw=0 if len(argv) > 2 else 1500)
        return 0
    if argv[0] in ("fetch", "cdx"):
        return generic(argv)
    if argv[0] == "px":
        for url in argv[1:]:
            pxweb_meta(url)
        return 0
    for name in argv:
        log(f"\n=== {name}")
        try:
            PROBES[name]()
        except Exception as exc:  # noqa: BLE001
            log(f"   probe {name} failed: {exc.__class__.__name__}: {exc}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
