#!/usr/bin/env python3
"""Central Asia and Iran reconnaissance: what each route serves, before a reader.

Read-only; nothing is written and the output is the log. Each subcommand
answers one question and prints only what decides it.

* ``get URL [URL...]`` -- status, type and length of each body; with
  ``--links RE`` the links whose href or text match, made absolute; with
  ``--grep RE`` the passages that match, with ``--context`` characters
  around each; otherwise the opening ``--chars`` characters of the text.
  ``--data`` POSTs a form or JSON body instead of a GET.
* ``spa URL`` -- a single-page application's own API paths: the page, its
  scripts, and every ``api/...`` path or absolute URL those scripts name.
* ``cdx PATTERN`` -- the Wayback Machine's captures under a URL pattern.
* ``hdx ISO[,ISO...]`` -- every HDX dataset for the country that carries a
  population table, with its licence and resources.
* ``res DATASET`` -- an HDX dataset's resources with their URLs; with
  ``--peek`` each CSV's or workbook sheet's header and first rows.
* ``json URL`` -- a JSON document's shape: the keys at ``--path`` and the
  first ``--rows`` items of a list there.
* ``tree URL --match RE`` -- the nodes of a nested JSON catalogue (SIAT's)
  whose names match, with their ids and codes.
* ``siat ID [ID...]`` -- what each SIAT indicator's data file holds: rows by
  code length (region or district) and the years covered.
* ``xlsx URL`` -- a big workbook's structure: every row whose first cell is a
  label rather than a number, with that row's non-empty cells.
* ``pages URL [URL...] --match RE`` -- the whole text of every PDF page that
  matches, up to ``--most`` pages per file.

Usage:
    python -m scripts.fetch_census.central_asia_probe get https://stat.gov.kg/ru/ --links census
    python -m scripts.fetch_census.central_asia_probe cdx amar.org.ir/* --match xlsx
    python -m scripts.fetch_census.central_asia_probe hdx IRN,KAZ
"""

from __future__ import annotations

import argparse
import csv
import html
import io
import json
import re
import urllib.error
import urllib.parse
import urllib.request
from typing import Any

from ._shared import log

HDX = "https://data.humdata.org/api/3/action"
UA = "DemographicMap/1.0 (+https://github.com/advaitsridhar/DemographicMap)"
TIMEOUT = 90
AGE = re.compile(r"^[TFM]_?\d{1,2}_\d{1,2}$|^[TFM]_?\d{1,2}_?plus$", re.I)


def fetch(url: str, *, timeout: int = TIMEOUT, accept: str | None = None,
          data: bytes | None = None, ctype: str | None = None,
          limit: int | None = None) -> tuple[int, str, bytes, str]:
    """(status, content type, body, final URL); an HTTP error is an answer."""
    headers = {"User-Agent": UA}
    if accept:
        headers["Accept"] = accept
    if ctype:
        headers["Content-Type"] = ctype
    req = urllib.request.Request(url, headers=headers, data=data)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as fh:
            body = fh.read(limit) if limit else fh.read()
            return fh.status, fh.headers.get("Content-Type", ""), body, fh.geturl()
    except urllib.error.HTTPError as err:
        body = b""
        try:
            body = err.read()[:4000]
        except Exception:  # noqa: BLE001 -- the status is the answer
            pass
        return err.code, (err.headers.get("Content-Type", "") if err.headers else ""), body, url


def decode(body: bytes, charset: str | None, ctype: str) -> str:
    if charset:
        return body.decode(charset, "replace")
    found = re.search(r"charset=([\w-]+)", ctype or "", re.I)
    if not found:
        found = re.search(rb"<meta[^>]+charset=[\"']?([\w-]+)", body[:4000], re.I)
        enc = found.group(1).decode() if found else "utf-8"
    else:
        enc = found.group(1)
    try:
        return body.decode(enc, "replace")
    except LookupError:
        return body.decode("utf-8", "replace")


def text_of(page: str) -> str:
    page = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", page)
    page = re.sub(r"(?s)<[^>]+>", " ", page)
    return " ".join(html.unescape(page).split())


def cmd_get(urls: list[str], *, links: str, grep: str, chars: int, context: int,
            rows: int, charset: str | None, data: str | None, accept: str | None,
            ctype: str | None, raw: bool) -> None:
    for url in urls:
        body_data = data.encode("utf-8") if data is not None else None
        try:
            status, kind, body, final = fetch(url, timeout=180, data=body_data,
                                              accept=accept, ctype=ctype)
        except Exception as err:  # noqa: BLE001 -- a refusal is an answer
            log(f"{url}: {type(err).__name__}: {str(err)[:300]}")
            continue
        page = decode(body, charset, kind)
        moved = f" -> {final}" if final != url else ""
        log(f"{url}{moved}: HTTP {status} {kind} {len(body):,} B")
        if links:
            seen: set[str] = set()
            found = 0
            for href, label in re.findall(
                    r"""(?is)<a\b[^>]*?href\s*=\s*["']([^"']+)["'][^>]*>(.*?)</a>""", page):
                full = urllib.parse.urljoin(final, html.unescape(href.strip()))
                label = text_of(label)[:120]
                if full in seen or not re.search(links, f"{full} {label}", re.I):
                    continue
                seen.add(full)
                found += 1
                if found <= rows:
                    log(f"  {label} -> {full}")
            log(f"  {found} links match {links!r}")
        if grep:
            plain = page if raw else text_of(page)
            hits = list(re.finditer(grep, plain, re.I))
            log(f"  {len(hits)} passages match {grep!r}")
            for m in hits[:rows]:
                a, b = max(0, m.start() - context), min(len(plain), m.end() + context)
                log(f"  ... {plain[a:b]} ...")
        if not links and not grep:
            plain = page if raw else text_of(page)
            log("  " + plain[:chars])


def cmd_spa(url: str, rows: int) -> None:
    try:
        status, kind, body, final = fetch(url)
    except Exception as err:  # noqa: BLE001 -- a DNS failure is an answer too
        log(f"{url}: {type(err).__name__}: {err}")
        return
    text = body.decode("utf-8", "replace")
    log(f"{url}: HTTP {status} {kind} {len(body):,} B")
    scripts = re.findall(r"<script[^>]+src=[\"']([^\"']+)", text)
    pattern = r"""["'`]([^"'`\s]*api/[^"'`\s]{2,160})["'`]"""
    paths: set[str] = set(re.findall(pattern, text))
    for src in scripts[:30]:
        full = urllib.parse.urljoin(final, src)
        try:
            s, _, js, _ = fetch(full, timeout=120)
        except Exception as err:  # noqa: BLE001 -- reported
            log(f"  script {full}: {type(err).__name__}: {err}")
            continue
        jtext = js.decode("utf-8", "replace")
        found = set(re.findall(pattern, jtext))
        found |= set(re.findall(r"""["'`](https?://[^"'`\s]{6,160})["'`]""", jtext))
        log(f"  script {full}: HTTP {s} {len(js):,} B, {len(found)} paths")
        paths |= found
    for p in sorted(paths)[:rows]:
        log(f"    {p}")


def cmd_cdx(pattern: str, limit: int, match: str, since: str, depth: int = 0,
            every: bool = False) -> None:
    params = {"url": pattern, "output": "json", "limit": str(limit),
              "filter": "statuscode:200"}
    if not every:
        params["collapse"] = "urlkey"
    if since:
        params["from"] = since
    url = "http://web.archive.org/cdx/search/cdx?" + urllib.parse.urlencode(params)
    try:
        status, _, body, _ = fetch(url, timeout=240)
    except Exception as err:  # noqa: BLE001 -- reported
        log(f"cdx {pattern}: {type(err).__name__}: {err}")
        return
    if status != 200:
        log(f"cdx {pattern}: HTTP {status} {body[:200]!r}")
        return
    rows = json.loads(body or b"[]")
    keep = [r for r in rows[1:] if not match or re.search(match, r[2], re.I)]
    log(f"cdx {pattern}: {max(len(rows) - 1, 0)} captures, {len(keep)} matching {match!r}")
    if depth:
        # A folder tree of thousands of files is read by its folders: the
        # captures counted under each path prefix of ``depth`` segments.
        folders: dict[str, int] = {}
        for r in keep:
            path = urllib.parse.urlsplit(r[2]).path.lower()
            folder = "/".join(path.split("/")[:depth])
            folders[folder] = folders.get(folder, 0) + 1
        for folder, n in sorted(folders.items()):
            log(f"  {n:5d}  {folder}")
        return
    for r in keep[:400]:
        size = r[6] if len(r) > 6 else "?"
        log(f"  {r[1]} {r[4]} {r[3][:28]} {size:>9} {r[2]}")


def hdx(path: str, **params: Any) -> Any:
    url = f"{HDX}/{path}?" + urllib.parse.urlencode(params)
    status, _, body, _ = fetch(url, accept="application/json")
    if status != 200:
        raise SystemExit(f"HDX {path} answered {status}")
    return json.loads(body)["result"]


def cmd_hdx(codes: list[str], query: str) -> None:
    for code in codes:
        found = hdx("package_search", fq=f"groups:{code.lower()}", q=query, rows=80)
        log(f"== {code}: {found.get('count')} datasets match {query!r}")
        for p in found.get("results") or []:
            org = (p.get("organization") or {}).get("name", "")
            log(f"  {p.get('name')} [{org}] licence={p.get('license_id')} "
                f"ref={p.get('dataset_date')}")
            for r in p.get("resources") or ():
                log(f"      - {str(r.get('name'))[:100]} ({r.get('format')})")


def cmd_res(datasets: list[str], peek: bool, rows: int) -> None:
    for name in datasets:
        p = hdx("package_show", id=name)
        log(f"== {name}: licence={p.get('license_id')} {p.get('license_other') or ''}")
        meth = p.get("methodology_other") or p.get("methodology") or ""
        if meth:
            log(f"   method: {str(meth)[:300]}")
        for r in p.get("resources") or ():
            log(f"  - {r.get('name')} ({r.get('format')}, {r.get('size') or '?'} B)")
            log(f"    {r.get('url')}")
            low = str(r.get("name") or "").lower()
            if not peek or not low.endswith((".xlsx", ".csv")):
                continue
            status, _, body, _ = fetch(str(r.get("url")), timeout=600)
            if status != 200:
                log(f"    HTTP {status}")
                continue
            if low.endswith(".csv"):
                lines = body.decode("utf-8-sig", "replace").splitlines()
                header = next(csv.reader(io.StringIO(lines[0]))) if lines else []
                ages = [c for c in header if AGE.match(c.strip())]
                log(f"    {len(lines)} lines; {len(header)} cols; {len(ages)} age cols")
                for line in lines[:rows]:
                    log(f"    | {line[:400]}")
                continue
            import openpyxl
            book = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
            for sheet in book.worksheets:
                log(f"    sheet {sheet.title!r}: {sheet.max_row} rows x {sheet.max_column} cols")
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    if i >= rows:
                        break
                    cells = ["" if c is None else str(c) for c in row]
                    log(f"    | {' | '.join(cells[:30])[:500]}")


def walk(doc: Any, path: str) -> Any:
    for part in [p for p in path.split(".") if p]:
        if isinstance(doc, list):
            doc = doc[int(part)]
        else:
            doc = doc[part]
    return doc


def cmd_json(urls: list[str], path: str, rows: int, data: str | None, chars: int) -> None:
    for url in urls:
        body_data = data.encode("utf-8") if data is not None else None
        try:
            status, kind, body, _ = fetch(url, timeout=180, data=body_data,
                                          accept="application/json",
                                          ctype="application/json" if data else None)
        except Exception as err:  # noqa: BLE001 -- reported
            log(f"{url}: {type(err).__name__}: {str(err)[:300]}")
            continue
        log(f"{url}: HTTP {status} {kind} {len(body):,} B")
        try:
            doc = walk(json.loads(body), path)
        except Exception as err:  # noqa: BLE001 -- reported with the opening bytes
            log(f"  not JSON at {path!r}: {type(err).__name__}: "
                f"{' '.join(body[:300].decode('utf-8', 'replace').split())}")
            continue
        if isinstance(doc, dict):
            log(f"  dict with {len(doc)} keys: {list(doc)[:40]}")
            for k in list(doc)[:rows]:
                log(f"   {k}: {json.dumps(doc[k], ensure_ascii=False)[:chars]}")
        elif isinstance(doc, list):
            log(f"  list of {len(doc)}")
            for item in doc[:rows]:
                log(f"   {json.dumps(item, ensure_ascii=False)[:chars]}")
        else:
            log(f"  {json.dumps(doc, ensure_ascii=False)[:chars]}")


def cmd_tree(url: str, match: str, keys: list[str], rows: int) -> None:
    """Every node of a nested JSON catalogue whose name matches, with its id.

    SIAT's catalogue is one 3 MB tree of sections, groups and indicators;
    this walks it and prints the nodes a pattern picks out, so an indicator
    can be found by what it measures rather than by paging the site.
    """
    status, _, body, _ = fetch(url, timeout=240, accept="application/json")
    log(f"{url}: HTTP {status} {len(body):,} B")
    found = 0

    def visit(node: Any, path: list[str]) -> None:
        nonlocal found
        if isinstance(node, list):
            for item in node:
                visit(item, path)
            return
        if not isinstance(node, dict):
            return
        names = [str(node.get(k) or "") for k in keys]
        label = " | ".join(n for n in names if n)
        if label and re.search(match, label, re.I):
            found += 1
            if found <= rows:
                log(f"  id={node.get('id')} code={node.get('code')} {label[:220]}"
                    f"  <- {' / '.join(path[-2:])[:120]}")
        here = path + ([names[0][:40]] if names and names[0] else [])
        for value in node.values():
            if isinstance(value, (list, dict)):
                visit(value, here)

    visit(json.loads(body), [])
    log(f"  {found} nodes match {match!r}")


def cmd_siat(ids: list[str], rows: int) -> None:
    """What each SIAT indicator's data file holds: units by code length, years."""
    for ident in ids:
        pointer_url = (f"https://api.siat.stat.uz/sdmx/{ident}/table/download/"
                       f"?download_format=json")
        status, _, body, _ = fetch(pointer_url, accept="application/json")
        if status != 200:
            log(f"siat {ident}: pointer HTTP {status}")
            continue
        pointer = json.loads(body)
        url = pointer.get("file") or pointer.get("file_2")
        status, _, body, _ = fetch(str(url), timeout=240)
        if status != 200:
            log(f"siat {ident}: data HTTP {status} {url}")
            continue
        doc = json.loads(body)
        head = doc[0] if isinstance(doc, list) else doc
        data = head.get("data") or []
        meta = {k: v for k, v in head.items() if k != "data"}
        lengths: dict[int, int] = {}
        for row in data:
            lengths[len(str(row.get("Code") or ""))] = lengths.get(
                len(str(row.get("Code") or "")), 0) + 1
        years = sorted(k for k in (data[0] if data else {}) if re.fullmatch(r"\d{4}", k))
        log(f"siat {ident}: {len(data)} rows; code lengths {lengths}; years "
            f"{years[:2]}..{years[-2:]}; updated {pointer.get('updated_at')}")
        log(f"  meta: {json.dumps(meta, ensure_ascii=False)[:600]}")
        for row in data[:rows]:
            log(f"   {json.dumps(row, ensure_ascii=False)[:300]}")


def workbook_rows(body: bytes) -> list[tuple[str, Any]]:
    """[(sheet name, row iterator)] for an .xlsx (openpyxl) or a legacy .xls (xlrd)."""
    if body[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        import xlrd
        book = xlrd.open_workbook(file_contents=body)

        def rows_of(sheet: Any) -> Any:
            # A function, not a generator expression in the comprehension:
            # that one would close over the loop variable and read every
            # sheet's rows from the last sheet.
            return (sheet.row_values(i) for i in range(sheet.nrows))
        return [(s.name, rows_of(s)) for s in book.sheets()]
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    return [(name, book[name].iter_rows(values_only=True)) for name in book.sheetnames]


def cmd_xlsx(url: str, sheets: list[str], rows: int, width: int, skip: str,
             grep: str = "", sheet_match: str = "", cells: int = 0) -> None:
    """A big workbook's structure: the rows whose first cell is a label.

    A census workbook that stacks one block per region down a sheet is read
    by where its blocks start; printing every row whose first cell is not a
    plain number (an age, a code) shows that, with each such row's non-empty
    cells, without printing the thousands of figures between them.

    ``grep`` prints instead only the rows one of whose cells matches it (a
    settlement table's county, district and city rows); ``sheet_match`` reads
    only the sheets whose name matches. Legacy ``.xls`` is read too.
    """
    status, _, body, _ = fetch(url, timeout=600)
    log(f"{url}: HTTP {status} {len(body):,} B")
    if status != 200:
        return
    book = workbook_rows(body)
    names = [name for name, _ in book]
    log(f"  sheets ({len(names)}): {names[:60]}")
    pattern = re.compile(grep) if grep else None
    for name, it in book:
        if sheets and name not in sheets:
            continue
        if sheet_match and not re.search(sheet_match, name):
            continue
        log(f"  -- {name!r}")
        shown = 0
        for i, row in enumerate(it, start=1):
            row = list(row or [])
            if cells:
                # The first rows cell by cell, with each cell's column: where
                # a merged heading sits is what a block reader depends on.
                if i > cells:
                    break
                log(f"  r{i}: " + " | ".join(f"c{j}={str(c).strip()[:width]!r}"
                                            for j, c in enumerate(row)
                                            if c not in (None, "")))
                continue
            first = "" if not row or row[0] is None else str(row[0]).strip()
            if not pattern and skip and re.fullmatch(skip, first):
                continue
            cells = [str(c).strip()[:width] for c in row if c not in (None, "")]
            if not cells:
                continue
            if pattern and not any(pattern.search(str(c)) for c in row if c not in (None, "")):
                continue
            shown += 1
            if shown > rows:
                log(f"  ... more rows; stopped at {rows}")
                break
            log(f"  r{i}: {' | '.join(cells)[:900]}")


def cmd_abadi(url: str, kinds: set[str], counties: set[str], width: int, notes: bool) -> None:
    """Iran's 2016 settlement table (CN95_HouseholdPopulationVillage_NN.xlsx).

    One sheet lists the province, every county (record code 2), district
    (3), rural district (4), city (5) and village (6, 8) with households,
    persons, men and women. This prints the rows of the record codes asked
    for (``kinds``), in the counties asked for (their two-digit codes, or
    all), and the workbook's notes sheet whole.
    """
    status, _, body, _ = fetch(url, timeout=600)
    log(f"{url}: HTTP {status} {len(body):,} B")
    if status != 200:
        return
    book = workbook_rows(body)
    for name, it in book:
        rows = [list(r or []) for r in it]
        if name != book[0][0]:
            if notes:
                log(f"  -- notes sheet {name!r}")
                for row in rows:
                    cells = [str(c).strip() for c in row if c not in (None, "")]
                    if cells:
                        log("  " + " | ".join(cells))
            continue
        head = [str(c or "").strip() for c in rows[1]] if len(rows) > 1 else []
        try:
            kind_col = next(i for i, h in enumerate(head) if "ركورد" in h or "رکورد" in h)
        except StopIteration:
            log(f"  no record-code column in {head}")
            return
        log(f"  {name!r}: {len(rows)} rows; record code in column {kind_col}")
        counts: dict[str, int] = {}
        for i, row in enumerate(rows[2:], start=3):
            kind = str(row[kind_col] if kind_col < len(row) else "").strip().split(".")[0]
            counts[kind] = counts.get(kind, 0) + 1
            county = str(row[2] if len(row) > 2 and row[2] is not None else "").strip()
            if re.fullmatch(r"\d+(?:\.0)?", county):
                county = f"{int(float(county)):02d}"
            if kinds and kind not in kinds:
                continue
            if counties and county not in counties:
                continue
            cells = [str(c).strip()[:width] for c in row if c not in (None, "")]
            log(f"  r{i}: {' | '.join(cells)}")
        log(f"  rows by record code: {counts}")


def cmd_pages(urls: list[str], match: str, most: int, nfkc: bool = False,
              only: str = "") -> None:
    """The text of every page of each PDF whose text matches ``match``.

    A census book's tables run over pages whose numbers differ from book to
    book; their titles and "Продолжение табл. N" lines do not. Printing those
    pages whole gives a parser its real layout, wrapped labels and all.

    ``nfkc`` folds the text first (Persian and Arabic PDFs often carry their
    letters as presentation forms, which no pattern in ordinary letters
    matches); ``only`` ("3,7-9") prints those pages whatever they hold.
    """
    import logging
    import unicodedata

    from pypdf import PdfReader
    logging.getLogger("pypdf").setLevel(logging.ERROR)
    pattern = re.compile(match)
    wanted: set[int] = set()
    for part in only.split(","):
        if part.strip():
            low, _, high = part.partition("-")
            wanted.update(range(int(low), int(high or low) + 1))
    for url in urls:
        status, _, body, _ = fetch(url, timeout=300)
        if status != 200:
            log(f"{url}: HTTP {status}")
            continue
        pages = PdfReader(io.BytesIO(body)).pages
        hits = 0
        log(f"{url}: {len(pages)} pages")
        for number, page in enumerate(pages, start=1):
            text = page.extract_text() or ""
            if nfkc:
                text = unicodedata.normalize("NFKC", text)
            if wanted:
                if number not in wanted:
                    continue
            elif not pattern.search(text):
                continue
            hits += 1
            if hits > most:
                log(f"  ... more pages match; stopped at {most}")
                break
            log(f"--- page {number} ---")
            for line in text.splitlines():
                log(line.rstrip())


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    pg = sub.add_parser("pages")
    pg.add_argument("urls", nargs="+")
    pg.add_argument("--match", required=True)
    pg.add_argument("--most", type=int, default=60)
    pg.add_argument("--nfkc", action="store_true", help="fold presentation forms first")
    pg.add_argument("--only", default="", help="print these pages whatever they hold")
    x2 = sub.add_parser("xlsx")
    x2.add_argument("url")
    x2.add_argument("--sheets", default="")
    x2.add_argument("--rows", type=int, default=80)
    x2.add_argument("--width", type=int, default=40)
    x2.add_argument("--skip", default=r"\d+|\d+\D{0,3}", help="first-cell pattern to leave out")
    x2.add_argument("--grep", default="", help="print only rows with a cell matching this")
    x2.add_argument("--sheet-match", default="", help="read only sheets whose name matches")
    x2.add_argument("--cells", type=int, default=0,
                    help="print this many first rows cell by cell, with column numbers")
    ab = sub.add_parser("abadi")
    ab.add_argument("url")
    ab.add_argument("--kinds", default="2,3,5", help="record codes to print")
    ab.add_argument("--counties", default="", help="two-digit county codes (all if empty)")
    ab.add_argument("--width", type=int, default=30)
    ab.add_argument("--notes", action="store_true")
    s2 = sub.add_parser("siat")
    s2.add_argument("ids", nargs="+")
    s2.add_argument("--rows", type=int, default=3)
    t = sub.add_parser("tree")
    t.add_argument("url")
    t.add_argument("--match", required=True)
    t.add_argument("--keys", default="name_en,name_ru")
    t.add_argument("--rows", type=int, default=150)
    g = sub.add_parser("get")
    g.add_argument("urls", nargs="+")
    g.add_argument("--links", default="")
    g.add_argument("--grep", default="")
    g.add_argument("--chars", type=int, default=600)
    g.add_argument("--context", type=int, default=160)
    g.add_argument("--rows", type=int, default=60)
    g.add_argument("--charset", default=None)
    g.add_argument("--data", default=None)
    g.add_argument("--accept", default=None)
    g.add_argument("--ctype", default=None)
    g.add_argument("--raw", action="store_true", help="grep the markup, not the text")
    s = sub.add_parser("spa")
    s.add_argument("url")
    s.add_argument("--rows", type=int, default=300)
    c = sub.add_parser("cdx")
    c.add_argument("pattern")
    c.add_argument("--limit", type=int, default=5000)
    c.add_argument("--match", default="")
    c.add_argument("--since", default="")
    c.add_argument("--depth", type=int, default=0,
                   help="count captures by folder, this many path segments deep")
    c.add_argument("--every", action="store_true",
                   help="every capture of each URL, not only the first")
    h = sub.add_parser("hdx")
    h.add_argument("codes")
    h.add_argument("--query", default="population")
    r = sub.add_parser("res")
    r.add_argument("datasets", nargs="+")
    r.add_argument("--peek", action="store_true")
    r.add_argument("--rows", type=int, default=5)
    j = sub.add_parser("json")
    j.add_argument("urls", nargs="+")
    j.add_argument("--path", default="")
    j.add_argument("--rows", type=int, default=10)
    j.add_argument("--chars", type=int, default=400)
    j.add_argument("--data", default=None)
    args = ap.parse_args()
    if args.cmd == "pages":
        cmd_pages(args.urls, args.match, args.most, args.nfkc, args.only)
    elif args.cmd == "xlsx":
        cmd_xlsx(args.url, [s for s in args.sheets.split(",") if s], args.rows,
                 args.width, args.skip, args.grep, args.sheet_match, args.cells)
    elif args.cmd == "abadi":
        cmd_abadi(args.url, {k for k in args.kinds.split(",") if k},
                  {c for c in args.counties.split(",") if c}, args.width, args.notes)
    elif args.cmd == "siat":
        cmd_siat(args.ids, args.rows)
    elif args.cmd == "tree":
        cmd_tree(args.url, args.match, [k for k in args.keys.split(",") if k], args.rows)
    elif args.cmd == "get":
        cmd_get(args.urls, links=args.links, grep=args.grep, chars=args.chars,
                context=args.context, rows=args.rows, charset=args.charset,
                data=args.data, accept=args.accept, ctype=args.ctype, raw=args.raw)
    elif args.cmd == "spa":
        cmd_spa(args.url, args.rows)
    elif args.cmd == "cdx":
        cmd_cdx(args.pattern, args.limit, args.match, args.since, args.depth, args.every)
    elif args.cmd == "hdx":
        cmd_hdx([x.strip().upper() for x in args.codes.split(",") if x.strip()], args.query)
    elif args.cmd == "res":
        cmd_res(args.datasets, args.peek, args.rows)
    else:
        cmd_json(args.urls, args.path, args.rows, args.data, args.chars)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
