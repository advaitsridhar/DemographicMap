#!/usr/bin/env python3
"""Tunisia -- the 2024 census by governorate and delegation: people, ages, sexes.

The Institut National de la Statistique took the Recensement Général de la
Population et de l'Habitat on 6 November 2024 and published, with its first
results in May 2025, a workbook of the whole population by secteur (imada),
sex and five-year age group: every sector with its delegation and governorate,
by code and by name in French and Arabic (sheet ``Population``, one row per
sector, sex and age group). Added up, that is every delegation and every
governorate: the head count, men per 1,000 women, and the median age
interpolated within the five-year group holding the middle person.

**Checks before anything is written.** Every row is one sector, one sex and
one age group, the groups running from 0 without a gap to an open last group;
the sectors make the national count INS announced, 11,972,169; each
delegation sits in one governorate.

**Binding.** The map's first level is the 24 governorates, labelled in French;
its second the 264 delegations of the 2014 census, labelled mostly in Arabic
and a few in French, some with stray direction marks and Persian letters. A
delegation is matched within its governorate on its Arabic or its French name,
both folded (hamza forms, ta marbuta, alef maqsura and Persian letters read as
the Arabic letter, marks and spaces dropped), then on the closest spelling
where each is the other's best (``DELEGATIONS`` declares the few that need a
label). INS counts 279 delegations in 2024 where the map draws the 264 of
2014: each of the fifteen made since is put back on the polygon of the 2014
delegation its sectors belonged to, sector by sector, from INS's 2014 table
of the population by sector (``SECTORS_2014``), and the polygon's figures are
the two delegations added together, age group by age group, which its note
says. Everything left over is logged, both sides.

Usage:
    python -m scripts.fetch_census.tunisia_census
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import unicodedata
from collections import defaultdict
from difflib import SequenceMatcher
from typing import Any

from ._shared import PROCESSED, http_get, log, measure, record, write_json

URL = ("https://www.ins.tn/sites/default/files-ftp3/files/2025-05/"
       "Population_Secteur_Age_Sexe_Dynamique.xlsx")
LANDING = ("https://www.ins.tn/enquetes/"
           "recensement-general-de-la-population-et-de-lhabitat-2024")
SOURCE = ("Institut National de la Statistique (INS), Recensement Général de la Population "
          "et de l'Habitat 2024, répartition de la population totale par secteur, âge et sexe")
LICENCE = "INS publication (reuse with attribution)"
YEAR = 2024
NATIONAL = 11_972_169
OUT = "tunisia_census.json"
SITE = PROCESSED.parent.parent / "site" / "data"

# INS's French governorate name, folded -> the map's label, where folding
# alone does not join them.
GOVERNORATES = {"lekef": "El Kef", "mannouba": "Manouba", "jandouba": "Jendouba",
                "zaghouane": "Zaghouan", "mednine": "Médenine"}
# The 2014 census's population by sector (secteur, imada), each under its
# delegation as drawn in 2014: what places the sectors of a delegation made
# since then on the polygon of the delegation they were part of.
SECTORS_2014 = ("https://www.ins.tn/sites/default/files-ftp3/files/2021-03/"
                "fr_population_logement_menages_par_secteurs_2014.xlsx")
REGIONS = {"tunisie", "nordest", "nordouest", "centreest", "centreouest", "sudest",
           "sudouest", "grandtunis", "districttunis"}
# INS's delegation (governorate code, delegation code) -> the map's label,
# where neither name folds to the label and no spelling is close enough. The
# label is looked for in the governorate first and then in the whole country:
# Ouled Chamekh is Mahdia's, and its polygon is filed under Sousse.
DELEGATIONS: dict[tuple[str, str], str] = {
    ("11", "1153"): "باب بحر",                 # Bab el Bhar
    ("11", "1162"): "الزهور",                  # Ezzouhour (Tunis)
    ("22", "2251"): "Jendouba Sud",            # Jendouba, the delegation of the town
    ("22", "2259"): "Balta Bouawene",
    ("31", "3166"): "الزاوية قصيبة الثريات",    # Zaouia-Ksiba-Thrayet
    ("33", "3353"): "والد شامخ",               # Ouled Chamekh
    ("42", "4253"): "الزهور",                  # Ezzouhour (Kasserine)
}
# A pair left over is bound on its closest spelling at this similarity or more.
SPELLING = 0.82
# A sector of a new delegation is found in 2014 under another spelling at this.
SECTOR_SPELLING = 0.8
# A delegation made since 2014 none of whose sectors INS's 2014 table names
# -> the 2014 delegation (as ``place_key`` writes it) whose outline holds it.
# None is needed today: all fifteen are found sector by sector.
HOMES_2014: dict[str, str] = {}
AGE = re.compile(r"^\s*(\d{1,3})\s*-\s*(\d{1,3})\s*$")
OPEN = re.compile(r"^\s*(\d{1,3})\s*(?:\+|ans?\s*et\s*plus|et\s*\+|et\s*plus)\s*$", re.I)
ARABIC = str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ٱ": "ا", "ة": "ه", "ى": "ي",
                        "ی": "ي", "ک": "ك", "ؤ": "و", "ئ": "ي", "ـ": None})


def fold(text: Any) -> str:
    """A name reduced to its letters: accents, marks, direction controls and
    spaces dropped; Arabic hamza forms, ta marbuta, alef maqsura and Persian
    letters read as the plain Arabic letter."""
    raw = unicodedata.normalize("NFKD", str(text or "")).translate(ARABIC)
    out = "".join(c for c in raw.lower()
                  if c.isalnum() and not unicodedata.combining(c)
                  and unicodedata.category(c) != "Cf")
    return re.sub(r"(.)\1+", r"\1", out)


def latin(text: Any) -> str:
    """The French name without 'Délégation de' and its kind words."""
    name = re.sub(r"(?i)^\s*d[ée]l[ée]gation\s+d[e']\s*", "", str(text or ""))
    return fold(name)


def place_key(text: Any) -> str:
    """A governorate's or delegation's French name as the 2014 and 2024
    tables can be joined on: without 'Gouvernorat' or 'Délégation', 'de',
    'du' or 'd'', and a leading article -- 2014's 'Gouvernorat du Kef' and
    2024's 'Le Kef' are both 'kef'."""
    name = re.sub(r"(?i)^\s*(gouvernorat|d[ée]l[ée]gation)\s+", "", str(text or ""))
    name = re.sub(r"(?i)^(d[eu]s?\s+|d['’]\s*)", "", name)
    name = re.sub(r"(?i)^(l[ea]s?\s+|l['’]\s*|el\s+)", "", name)
    return fold(name)


def number(cell: Any) -> int:
    if cell is None or cell == "":
        return 0
    if isinstance(cell, (int, float)):
        return int(round(cell))
    text = re.sub(r"[\s  ,]", "", str(cell))
    return int(float(text)) if text else 0


def bounds(label: str) -> tuple[int, int | None]:
    if m := AGE.match(label):
        return int(m.group(1)), int(m.group(2))
    if m := OPEN.match(label):
        return int(m.group(1)), None
    raise SystemExit(f"tunisia_census: an age group reads {label!r}")


def grouped_median(groups: list[tuple[int, int | None, int]]) -> float | None:
    base = sum(n for _, _, n in groups)
    half, before = base / 2, 0
    for low, high, people in groups:
        if before + people >= half:
            if high is None or people <= 0:
                return None
            return round(low + (half - before) / people * (high - low + 1), 1)
        before += people
    return None


def read(rows: list[list[Any]]) -> dict[str, Any]:
    """{"governorates": {code: unit}, "delegations": {code: unit}, "groups": [...]}."""
    header = [str(c or "").strip() for c in rows[0]]
    at = {h: i for i, h in enumerate(header)}
    need = ("Code_Gouvernorat", "Gouvernorat", "الولاية", "Code_Delegation", "Délégations",
            "المعتمدية", "ID_Sexe", "Tranche d'âge", "Population")
    missing = [n for n in need if n not in at]
    if missing:
        raise SystemExit(f"tunisia_census: the sheet has no {missing}; it has {header}")
    govs: dict[str, dict[str, Any]] = {}
    dels: dict[str, dict[str, Any]] = {}
    secs: dict[str, dict[str, Any]] = {}
    labels: set[str] = set()
    total = 0
    sector_code = at.get("Code_secteur_D")
    sector_name = at.get("Secteurs")
    for row in rows[1:]:
        cells = list(row) + [None] * len(header)
        gcode = str(cells[at["Code_Gouvernorat"]] or "").strip()
        dcode = str(cells[at["Code_Delegation"]] or "").strip()
        if not gcode or not dcode:
            continue
        sex = str(cells[at["ID_Sexe"]] or "").strip().upper()
        if sex not in ("M", "F"):
            raise SystemExit(f"tunisia_census: a row's sex reads {sex!r}")
        label = str(cells[at["Tranche d'âge"]] or "").strip()
        labels.add(label)
        people = number(cells[at["Population"]])
        total += people
        scode = str(cells[sector_code] or "").strip() if sector_code is not None else ""
        levels = [(govs, gcode, "Gouvernorat", "الولاية"),
                  (dels, dcode, "Délégations", "المعتمدية")]
        for table, code, fr, ar in levels:
            unit = table.setdefault(code, {"code": code, "fr": str(cells[at[fr]] or "").strip(),
                                           "ar": str(cells[at[ar]] or "").strip(),
                                           "gov": gcode, "M": 0, "F": 0,
                                           "ages": defaultdict(int)})
            if table is dels and unit["gov"] != gcode:
                raise SystemExit(f"tunisia_census: delegation {code} sits in governorates "
                                 f"{unit['gov']} and {gcode}")
            unit[sex] += people
            unit["ages"][label] += people
        if scode:
            sec = secs.setdefault(scode, {"code": scode, "gov": gcode, "del": dcode,
                                          "fr": str(cells[sector_name] or "").strip()
                                          if sector_name is not None else scode,
                                          "M": 0, "F": 0, "ages": defaultdict(int)})
            sec[sex] += people
            sec["ages"][label] += people
    groups = sorted({bounds(lab) + (lab,) for lab in labels}, key=lambda g: g[0])
    edge = 0
    for low, high, lab in groups:
        if low != edge:
            raise SystemExit(f"tunisia_census: the age groups do not run from 0 without a "
                             f"gap: {[g[2] for g in groups]}")
        edge = high + 1 if high is not None else None
    if groups[-1][1] is not None:
        raise SystemExit("tunisia_census: no open-ended last age group")
    if total != NATIONAL:
        raise SystemExit(f"tunisia_census: the sectors make {total:,}, not INS's {NATIONAL:,}")
    for table in (govs, dels, secs):
        for unit in table.values():
            unit["total"] = unit["M"] + unit["F"]
            unit["groups"] = [(lo, hi, unit["ages"].get(lab, 0)) for lo, hi, lab in groups]
    return {"governorates": govs, "delegations": dels, "sectors": secs, "groups": groups}


def read_2014(rows: list[list[Any]]) -> dict[str, dict[str, str]]:
    """{governorate (folded): {sector (folded): its 2014 delegation (folded)}}.

    The sheet lists the country, its regions, each governorate ("Gouvernorat
    de ..."), each delegation ("Délégation de ...") and the delegation's
    sectors, in that order, one row each; the whole population is in the
    'Ensemble milieux' block's first column. Each delegation's sectors must
    make it (a row that is a region heading would not)."""
    out: dict[str, dict[str, str]] = defaultdict(dict)
    gov = deleg = None
    totals: dict[tuple[str, str], int] = {}
    made: dict[tuple[str, str], int] = defaultdict(int)
    for row in rows:
        cells = list(row) + [None] * 16
        name = str(cells[0] or "").strip()
        if not name:
            continue
        people = cells[11]
        if not isinstance(people, (int, float)):
            continue
        key = fold(name)
        if re.match(r"(?i)^gouvernorat\s+d", name):
            gov, deleg = place_key(name), None
            continue
        if re.match(r"(?i)^d[ée]l[ée]gation\s+d", name):
            deleg = place_key(name)
            totals[(gov or "", deleg)] = int(people)
            continue
        if key in REGIONS or gov is None or deleg is None:
            continue
        if int(people) > totals[(gov, deleg)]:
            # A region's heading between two governorates, not a sector.
            log(f"  2014: {name!r} ({int(people):,}) is more than its delegation; skipped")
            continue
        out[gov][key] = deleg
        made[(gov, deleg)] += int(people)
    off = [f"{g}/{d}: sectors {made[(g, d)]:,} against {t:,}" for (g, d), t in totals.items()
           if made[(g, d)] != t]
    if off:
        log(f"  2014 sectors that do not make their delegation: {off[:12]}")
    return out


def carved(table: dict[str, Any], held: dict[str, list[str]], sectors_2014: dict[str, dict[str, str]],
           report: list[str]) -> dict[str, list[str]]:
    """Polygon id -> the 2024 sectors of delegations it holds that the map does
    not draw, each placed by the 2014 delegation its sector belonged to."""
    bound_codes = {c for codes in held.values() for c in codes}
    by_name: dict[tuple[str, str], str] = {}
    for sid, codes in held.items():
        for code in codes:
            unit = table["delegations"][code]
            by_name[(unit["gov"], place_key(unit["fr"]))] = sid
    gov_fold = {code: place_key(u["fr"]) for code, u in table["governorates"].items()}
    extra: dict[str, list[str]] = defaultdict(list)
    for code, unit in sorted(table["delegations"].items()):
        if code in bound_codes:
            continue
        mine = [s for s in table["sectors"].values() if s["del"] == code]
        known = sectors_2014.get(gov_fold[unit["gov"]], {})
        placed: dict[str, str] = {}
        undrawn: set[str] = set()
        for sec in mine:
            old = known.get(fold(sec["fr"]))
            if old is None and known:
                # A sector spelt differently in the two censuses ("Zanouch",
                # "Zannouch"): the closest name of the governorate's 2014
                # sectors, where it is close.
                best = max(known, key=lambda k: SequenceMatcher(None, fold(sec["fr"]), k).ratio())
                score = SequenceMatcher(None, fold(sec["fr"]), best).ratio()
                if score >= SECTOR_SPELLING:
                    old = known[best]
                    report.append(f"    sector {sec['fr']} read as 2014's {best} ({score:.2f})")
            sid = by_name.get((unit["gov"], old)) if old else None
            if sid:
                placed[sec["code"]] = sid
            elif old:
                undrawn.add(old)
        homes = set(placed.values())
        if not placed and code in HOMES_2014:
            home = by_name.get((unit["gov"], HOMES_2014[code]))
            if home is None:
                raise SystemExit(f"tunisia_census: HOMES_2014 puts {code} in "
                                 f"{HOMES_2014[code]!r}, which is drawn nowhere")
            report.append(f"    {unit['fr']}: none of its sectors in the 2014 table; put in "
                          f"{HOMES_2014[code]!r}, whose outline holds it")
            placed = {sec["code"]: home for sec in mine}
            homes = {home}
        if len(homes) == 1 and len(placed) < len(mine):
            # A sector renamed or made since 2014 goes where all its
            # delegation's other sectors went.
            home = next(iter(homes))
            for sec in mine:
                placed.setdefault(sec["code"], home)
        if len(placed) != len(mine) or not mine:
            where = (f"; found under 2014 delegations bound to no polygon: {sorted(undrawn)}"
                     if undrawn else "")
            report.append(f"  delegation {code} {unit['fr']} ({unit['total']:,}): "
                          f"{len(mine) - len(placed)} of its {len(mine)} sectors not placed: "
                          f"{[s['fr'] for s in mine if s['code'] not in placed]}{where}")
            continue
        for sec_code, sid in placed.items():
            extra[sid].append(sec_code)
        names = sorted({next(iter(n for (g, n), s in by_name.items() if s == sid))
                        for sid in set(placed.values())})
        report.append(f"  delegation {code} {unit['fr']} ({unit['total']:,}), made since 2014: "
                      f"its {len(mine)} sectors were part of {names}")
    return extra


def figures(unit: dict[str, Any]) -> dict[str, Any]:
    return {"median": grouped_median(unit["groups"]),
            "ratio": round(1000 * unit["M"] / unit["F"]) if unit["F"] else None}


def merge(units: list[dict[str, Any]]) -> dict[str, Any]:
    out = {"M": sum(u["M"] for u in units), "F": sum(u["F"] for u in units)}
    out["total"] = out["M"] + out["F"]
    out["groups"] = [(lo, hi, sum(u["groups"][i][2] for u in units))
                     for i, (lo, hi, _) in enumerate(units[0]["groups"])]
    return out


def bind_governorates(govs: dict[str, dict[str, Any]], admin1: list[dict[str, Any]]
                      ) -> dict[str, dict[str, Any]]:
    by_fold = {fold(s["name"]): s for s in admin1}
    bound: dict[str, dict[str, Any]] = {}
    for code, unit in govs.items():
        key = latin(unit["fr"])
        shape = by_fold.get(key) or by_fold.get(fold(GOVERNORATES.get(key, "")))
        if shape is None:
            raise SystemExit(f"tunisia_census: no governorate polygon for {unit['fr']!r}")
        if shape["id"] in {s["id"] for s in bound.values()}:
            raise SystemExit(f"tunisia_census: two governorates on {shape['name']!r}")
        bound[code] = shape
    if len(bound) != len(admin1):
        raise SystemExit(f"tunisia_census: {len(bound)} governorates for {len(admin1)} polygons")
    return bound


def bind_delegations(dels: dict[str, dict[str, Any]], admin2: list[dict[str, Any]],
                     govs: dict[str, dict[str, Any]], report: list[str]
                     ) -> dict[str, list[str]]:
    """Polygon id -> the INS delegation code it holds. A delegation INS
    counts and the map does not draw is left out here and logged."""
    held: dict[str, list[str]] = {}
    taken: set[str] = set()
    # Declared labels first, the governorate's polygons before the country's.
    for (gcode, code), label in DELEGATIONS.items():
        if code not in dels:
            raise SystemExit(f"tunisia_census: DELEGATIONS names {code}, which INS does not count")
        gshape = govs[gcode]
        hits = [s for s in admin2 if s["parent"] == gshape["id"] and fold(s["name"]) == fold(label)]
        if not hits:
            hits = [s for s in admin2 if fold(s["name"]) == fold(label)]
        if len(hits) != 1 or hits[0]["id"] in taken:
            raise SystemExit(f"tunisia_census: {len(hits)} polygons labelled {label!r} for {code}")
        held[hits[0]["id"]] = [code]
        taken.add(hits[0]["id"])
    declared = {code for (_, code) in DELEGATIONS}
    for gcode, gshape in govs.items():
        free = {s["id"]: s for s in admin2
                if s["parent"] == gshape["id"] and s["id"] not in taken}
        mine = {c: u for c, u in dels.items() if u["gov"] == gcode and c not in declared}
        left: dict[str, dict[str, Any]] = {}
        for code, unit in sorted(mine.items()):
            keys = {fold(unit["ar"]), latin(unit["fr"])}
            hits = [s for s in free.values() if fold(s["name"]) in keys]
            if len(hits) == 1:
                held[hits[0]["id"]] = [code]
                taken.add(hits[0]["id"])
                del free[hits[0]["id"]]
            else:
                left[code] = unit
        # Leftovers on both sides of one governorate: the closest spellings,
        # each the other's best.
        def ratio(unit: dict[str, Any], shape: dict[str, Any]) -> float:
            name = fold(shape["name"])
            return max(SequenceMatcher(None, fold(unit["ar"]), name).ratio(),
                       SequenceMatcher(None, latin(unit["fr"]), name).ratio())
        for code, unit in sorted(left.items()):
            best = max(free.values(), key=lambda s: ratio(unit, s), default=None)
            back = (max(left.values(), key=lambda u: ratio(u, best))
                    if best is not None else None)
            if best is not None and ratio(unit, best) >= SPELLING and back is unit:
                report.append(f"  {gshape['name']}: {unit['fr']} / {unit['ar']} bound to "
                              f"{best['name']!r} by spelling ({ratio(unit, best):.2f})")
                held[best["id"]] = [code]
                taken.add(best["id"])
                del free[best["id"]]
            else:
                report.append(f"  {gshape['name']}: delegation {code} {unit['fr']} / "
                              f"{unit['ar']} ({unit['M'] + unit['F']:,}) has no polygon of "
                              f"its own")
        for shape in free.values():
            report.append(f"  {gshape['name']}: polygon {shape['name']!r} has no delegation")
    return held


def display(name: str) -> str:
    """INS writes its French names in lower case after the first letter."""
    return " ".join(w if any(c.isupper() for c in w) else w.capitalize()
                    for w in str(name).split())


def build(table: dict[str, Any], admin1: list[dict[str, Any]], admin2: list[dict[str, Any]],
          sectors_2014: dict[str, dict[str, str]] | None = None
          ) -> tuple[list[dict[str, Any]], list[str]]:
    report: list[str] = []
    govs = bind_governorates(table["governorates"], admin1)
    records = []
    sources = [{"field": "population/median_age/sex_ratio", "name": SOURCE, "url": LANDING,
                "license": LICENCE, "year": YEAR}]
    for code, shape in sorted(govs.items()):
        unit = table["governorates"][code]
        records.append(unit_record(f"TUN-RGPH2024-{code}", shape, "admin1", unit, sources,
                                   f"the governorate of {display(unit['fr'])}"))
    held = bind_delegations(table["delegations"], admin2, govs, report)
    extra = carved(table, held, sectors_2014 or {}, report)
    # A delegation made since 2014 that is neither drawn nor put back on the
    # polygon it was carved from leaves that polygon short, and which polygon
    # it is cannot be told: nothing is written until it is placed.
    unplaced = {code for code, unit in table["delegations"].items()
                if code not in {c for codes in held.values() for c in codes}
                and not any(s["del"] == code for sid in extra for s in
                            (table["sectors"][x] for x in extra[sid]))}
    if unplaced:
        for line in report:
            log(line)
        raise SystemExit(f"tunisia_census: delegations neither drawn nor placed: "
                         f"{sorted(unplaced)}")
    labels = {s["id"]: s for s in admin2}
    for sid, codes in sorted(held.items()):
        units = [table["delegations"][c] for c in codes]
        added = [table["sectors"][x] for x in extra.get(sid, [])]
        unit = units[0] if not added else merge(units + added)
        name = display(units[0]["fr"])
        if added:
            new = sorted({display(table["delegations"][s["del"]]["fr"]) for s in added})
            what = (f"the delegation of {name} as drawn in 2014; the 2024 census counts its "
                    f"ground as {name} and {' and '.join(new)}, made a delegation since, "
                    f"which are added together here")
        else:
            what = f"the delegation of {name}"
        records.append(unit_record(f"TUN-RGPH2024-{'-'.join(codes)}", labels[sid], "admin2",
                                   unit, sources, what))
    return records, report


def unit_record(rid: str, shape: dict[str, Any], level: str, unit: dict[str, Any],
                sources: list[dict[str, Any]], what: str) -> dict[str, Any]:
    f = figures(unit)
    return record(
        rid, shape["name"], level=level, parent="TUN", country="TUN",
        match_by="shape_id", shape_id=shape["id"],
        population=measure(unit["total"], year=YEAR, source=SOURCE),
        population_note=(f"The 2024 census count of {what}: {unit['M']:,} men and "
                         f"{unit['F']:,} women."),
        median_age=measure(f["median"], unit="years", year=YEAR, source=SOURCE)
        if f["median"] is not None else None,
        median_age_note=("Interpolated within the five-year age group that holds the middle "
                         "person, from the 2024 census's count of everyone by five-year age "
                         "group."),
        sex_ratio=measure(f["ratio"], unit="males_per_1000_females", year=YEAR, source=SOURCE)
        if f["ratio"] is not None else None,
        sources=sources)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--run", help="a label that tells one dispatch from another; not used")
    parser.parse_args()
    import openpyxl                                 # noqa: PLC0415
    blob = http_get(URL, binary=True, cache=False, timeout=300)
    if blob[:2] != b"PK":
        raise SystemExit(f"tunisia_census: {URL} is not a workbook: {blob[:80]!r}")
    log(f"  read {URL}: {len(blob):,} bytes, sha256 {hashlib.sha256(blob).hexdigest()}")
    book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    if "Population" not in book.sheetnames:
        raise SystemExit(f"tunisia_census: no sheet 'Population': {book.sheetnames}")
    rows = [list(r) for r in book["Population"].iter_rows(values_only=True)]
    table = read(rows)
    log(f"  {len(table['governorates'])} governorates, {len(table['delegations'])} "
        f"delegations; age groups {[g[2] for g in table['groups']]}")
    for code, unit in sorted(table["delegations"].items()):
        f = figures(unit)
        log(f"  {unit['gov']} {code} {unit['fr']} | {unit['ar']} | {unit['total']:,} | "
            f"median {f['median']} | men/1000 women {f['ratio']}")
    old = http_get(SECTORS_2014, binary=True, cache=False, timeout=180)
    if old[:2] != b"PK":
        raise SystemExit(f"tunisia_census: {SECTORS_2014} is not a workbook: {old[:80]!r}")
    sheet = openpyxl.load_workbook(io.BytesIO(old), read_only=True, data_only=True).worksheets[0]
    sectors_2014 = read_2014([list(r) for r in sheet.iter_rows(values_only=True)])
    log(f"  2014: {sum(len(v) for v in sectors_2014.values()):,} sectors in "
        f"{len(sectors_2014)} governorates")
    admin1 = json.loads((SITE / "admin1" / "TUN.units.json").read_text())
    admin2 = json.loads((SITE / "admin2" / "TUN.units.json").read_text())
    records, report = build(table, admin1, admin2, sectors_2014)
    for line in report:
        log(line)
    bound = sum(1 for r in records if r["level"] == "admin2")
    log(f"  {sum(1 for r in records if r['level'] == 'admin1')} governorates and {bound} of "
        f"{len(admin2)} delegation polygons bound")
    write_json(PROCESSED / OUT, records)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
