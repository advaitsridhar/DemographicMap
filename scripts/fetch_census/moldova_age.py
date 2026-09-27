#!/usr/bin/env python3
"""Moldova: median age, sex ratio and head count by district, from the 2024 census.

The National Bureau of Statistics published the 2024 census's final results
with annex workbooks. Two of their tables give every district, municipality
and Gagauzia -- the 35 units the census counted:

* ``Anexa_Caracteristici_demografice_RPL2024.xlsx``, table 2.3 -- the usually
  resident population by five-year age group (0-4 ... 75-79, 80+);
* ``Anexa_Localitati_RPL2024.xlsx``, table 8.3 -- the same population by sex.

Median age is interpolated within the five-year group that holds the middle
person: the bureau publishes nothing finer by district. The country's
median from the same groups is logged beside the check that the units add up
to the published 2,409,207. Sex ratio is males per 100 females.

Bender and Transnistria, which the census did not reach, carry a stated gap;
the map draws each unit at both levels, as moldova_census does.

Usage:
    python -m scripts.fetch_census.moldova_age
"""

from __future__ import annotations

import argparse
import io
import re
from typing import Any

from ._shared import NOT_AVAILABLE, PROCESSED, gap, http_get, log, measure, record, write_json
from .balkans_common import check_sum, grouped_median
from .moldova_census import count, display, fold, key, shapes

OUT = "moldova_age.json"
YEAR = 2024
BASE = "https://statistica.gov.md/files/files/ComPresa/Recensamant/2024/Ro/"
AGES = BASE + "Anexa_Caracteristici_demografice_RPL2024.xlsx"
SEXES = BASE + "Anexa_Localitati_RPL2024.xlsx"
PAGE = ("https://statistica.gov.md/ro/rezultatele-finale-ale-recensamantului-populatiei-si-"
        "locuintelor-2024-10118.html")
SOURCE = ("Biroul Național de Statistică, Recensământul Populației și Locuințelor 2024, "
          "final results, annex table {table}")
LICENCE = "Official statistics of the Republic of Moldova (reuse with attribution)"
NATIONAL = 2_409_207
UNITS = 35
UNIT_NAME = re.compile(r"^(mun|raionul|r|uta)\b", re.I)
NOT_COUNTED = ("Moldova's 2024 census was not taken on the left bank of the Dniester or in "
               "Bender, so the bureau publishes no age or sex for this unit; Transnistria's "
               "own 2015 census published its age structure for the republic as a whole only.")


def rows_of(url: str, sheet: str) -> list[tuple[Any, ...]]:
    import openpyxl
    blob = http_get(url, binary=True, timeout=600)
    log(f"  {url.rsplit('/', 1)[-1]}: {len(blob):,} bytes")
    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"moldova_age: no sheet {sheet!r} in {wb.sheetnames}")
    return [tuple(r) for r in wb[sheet].iter_rows(values_only=True)]


def persons(cell: Any) -> float:
    """A count, rounded: the workbook stores some sums as 39562.9999999999,
    which truncating would read as 39,562."""
    if isinstance(cell, float):
        return float(round(cell))
    return float(count(cell))


def text(cell: Any) -> str:
    return " ".join(str(cell or "").split())


def ages(rows: list[tuple[Any, ...]]) -> tuple[dict[str, dict[str, Any]], list[tuple[float, float | None, float]]]:
    """({unit key: {"name", "total", "groups"}}, the country's groups) from table 2.3."""
    head = next(i for i, r in enumerate(rows) if any(text(c) == "0-4" for c in r))
    bands: list[tuple[int, float, float | None]] = []
    total_col = None
    for j, c in enumerate(rows[head]):
        label = text(c)
        m = re.match(r"^(\d+)-(\d+)$", label)
        top = re.match(r"^(\d+)\+$", label)
        if label.lower() == "total":
            total_col = j
        elif m:
            bands.append((j, float(m.group(1)), float(int(m.group(2)) - int(m.group(1)) + 1)))
        elif top:
            bands.append((j, float(top.group(1)), None))
    if total_col is None or not bands or bands[-1][2] is not None:
        raise SystemExit(f"moldova_age: table 2.3's header is not what was read: {rows[head]}")
    units: dict[str, dict[str, Any]] = {}
    country = None
    for r in rows[head + 1:]:
        cells = [text(c) for c in r]
        name = next((c for c in cells[:3] if c and not re.fullmatch(r"[\dA-Z]+", c)), "")
        if not name or total_col >= len(r) or r[total_col] in (None, ""):
            continue
        groups = [(lo, w, persons(r[j])) for j, lo, w in bands]
        total = persons(r[total_col])
        check_sum(sum(n for _, _, n in groups), total, f"moldova_age: ages of {name}")
        if name.lower() == "total":
            country = groups
            check_sum(total, NATIONAL, "moldova_age: the country's total")
            continue
        if UNIT_NAME.match(name):
            units[unit_key(name)] = {"name": name, "total": total, "groups": groups}
    if country is None or len(units) != UNITS:
        raise SystemExit(f"moldova_age: table 2.3 gave {len(units)} units and "
                         f"{'a' if country else 'no'} country row")
    check_sum(sum(u["total"] for u in units.values()), NATIONAL, "moldova_age: units against the country")
    return units, country


def unit_key(name: str) -> str:
    """Gagauzia is "UTA Găgăuzia" in one table and spelled out in the other."""
    if fold(name).startswith("unitatea teritoriala autonoma"):
        name = "UTA " + name.split()[-1]
    return key(name)


def sexes(rows: list[tuple[Any, ...]]) -> dict[str, tuple[float, float, float]]:
    """{unit key: (total, men, women)} from table 8.3.

    Every row says what it is in its "Tip dezagregare" cell -- "Raioane" for
    a district, a municipality or Gagauzia, "Comune" and "Localitati" for
    what lies inside -- and names itself in the cell after.
    """
    head = next(i for i, r in enumerate(rows) if any("Masculin" == text(c) for c in r))
    men_col = next(j for j, c in enumerate(rows[head]) if text(c) == "Masculin")
    women_col = next(j for j, c in enumerate(rows[head]) if text(c) == "Feminin")
    out: dict[str, tuple[float, float, float]] = {}
    for r in rows[head + 1:]:
        cells = [text(c) for c in r]
        if "Raioane" not in cells:
            continue
        name = cells[cells.index("Raioane") + 1]
        men, women = persons(r[men_col]), persons(r[women_col])
        total = persons(r[men_col - 1])
        check_sum(men + women, total, f"moldova_age: sexes of {name}")
        if unit_key(name) in out:
            raise SystemExit(f"moldova_age: table 8.3 names {name} twice")
        out[unit_key(name)] = (total, men, women)
    if len(out) != UNITS:
        raise SystemExit(f"moldova_age: table 8.3 gave {len(out)} units, not {UNITS}")
    check_sum(sum(v[0] for v in out.values()), NATIONAL, "moldova_age: table 8.3's units")
    return out


def build() -> list[dict[str, Any]]:
    units, country = ages(rows_of(AGES, "2.3"))
    sex = sexes(rows_of(SEXES, "8.3"))
    log(f"  country: median {grouped_median(country)} from five-year groups")
    drawn = shapes()
    records = []
    for k, unit in sorted(units.items()):
        if k not in sex:
            raise SystemExit(f"moldova_age: {unit['name']} has ages and no sexes")
        total, men, women = sex[k]
        check_sum(total, unit["total"], f"moldova_age: {unit['name']}'s two tables")
        median = grouped_median(unit["groups"])
        for level in ("admin1", "admin2"):
            shape = drawn.get(level, {}).get(k)
            if not shape:
                raise SystemExit(f"moldova_age: {unit['name']!r} is not a {level} shape")
            records.append(record(
                f"MDA-age-{level}-{k}", display(re.sub(r"^Raionul\s+", "", unit["name"])), level=level, parent="MDA",
                country="MDA", match_by="shape_id", shape_id=shape,
                population=measure(int(total), year=YEAR, source=SOURCE.format(table="8.3")),
                median_age=measure(median, unit="years", year=YEAR, source=SOURCE.format(table="2.3")),
                median_age_note=("Interpolated within the five-year age group that holds the "
                                 "middle person, from the census's usually resident population "
                                 "by age group (annex table 2.3): the bureau publishes nothing "
                                 "finer by district."),
                sex_ratio=measure(round(100 * men / women, 1), unit="males_per_100_females",
                                  year=YEAR, source=SOURCE.format(table="8.3")),
                sources=[{"field": "median_age", "name": SOURCE.format(table="2.3"), "url": AGES,
                          "page": PAGE, "year": YEAR, "license": LICENCE},
                         {"field": "population/sex_ratio", "name": SOURCE.format(table="8.3"),
                          "url": SEXES, "page": PAGE, "year": YEAR, "license": LICENCE}]))
        log(f"  {display(unit['name'])}: {total:,.0f}, median {median}, "
            f"{round(100 * men / women, 1)} men per 100 women")
    for name in ("Bender", "Transnistria"):
        k = key(name)
        for level in ("admin1", "admin2"):
            shape = drawn.get(level, {}).get(k)
            if shape:
                records.append(record(
                    f"MDA-age-{level}-{k}", name, level=level, parent="MDA", country="MDA",
                    match_by="shape_id", shape_id=shape,
                    median_age=gap(NOT_AVAILABLE, NOT_COUNTED),
                    sex_ratio=gap(NOT_AVAILABLE, NOT_COUNTED)))
    return records


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.parse_args()
    log("moldova_age: 2024 census, age and sex by district")
    records = build()
    write_json(PROCESSED / OUT, records)
    log(f"  wrote {OUT}: {len(records)} records")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
