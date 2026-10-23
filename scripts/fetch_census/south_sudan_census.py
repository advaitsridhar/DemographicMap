#!/usr/bin/env python3
"""South Sudan -- the 2008 census by state and county: median age and sex ratio.

South Sudan has held no census of its own. The Fifth Sudan Population and
Housing Census of 21 April 2008, enumerated in the south by the Southern Sudan
Centre for Census, Statistics and Evaluation (now the National Bureau of
Statistics), counted everyone by five-year age group and sex down to the boma;
the US Census Bureau's subnational tables for HDX carry those counts for the
10 states and 79 counties of 2008 (``south-sudan-subnational-boundaries-and-
tabular-data``, sheet ``Age-Sex``).

**What is written**: for each state and each county the map draws, the median
age interpolated within the five-year group holding the middle person, and
men per 1,000 women -- the census's own, which stand in front of the
projections OCHA's population tables give for 2022. The census's head count
is not the population written: for each county that is OCHA's estimate for
2025, the most recent figure, from the estimates workbook of its population
dataset (``cod-ps-ssd``, sheet ``Pop_Stats_Summary``, column ``Population -
2025``), bound by the county's P-code (the Abyei Area's row is never read);
the 2008 count is kept in the note.

**Checks before anything is written.** Every row's age groups make its total,
its men and women make it too, and the counties of each state make the state.
The age groups are read in the finest partition the sheet holds (five-year
groups to 90-94 and 95 and over, the sheet's 65-and-over sum set aside).
Counties bind to the map's within their state by name (``COUNTIES`` holds the
spellings that differ); everything left over is logged, both sides. The
Abyei Area, which Sudan and South Sudan both claim, is never written.

Usage:
    python -m scripts.fetch_census.south_sudan_census
"""

from __future__ import annotations

import argparse
import io
import json
import re
import unicodedata
from typing import Any

from . import uscb
from ._shared import PROCESSED, http_get, log, measure, record, write_json

DATASET = "south-sudan-subnational-boundaries-and-tabular-data"
YEAR = 2008
OUT = "south_sudan_census.json"
SITE = PROCESSED.parent.parent / "site" / "data"
# The publisher's name has no comma in it: the figure card shows the text before
# the first one.
SOURCE = ("Southern Sudan Centre for Census Statistics and Evaluation (SSCCSE), 5th Sudan "
          "Population and Housing Census 2008, population by five-year age group and sex, as the "
          "US Census Bureau tabulates it for HDX")
LICENCE = "Creative Commons Attribution International (CC BY), published via HDX"
# OCHA's county estimates for 2025, from its population dataset's estimates
# workbook (the build's own OCHA tables are of 2022).
ESTIMATES = "cod-ps-ssd"
ESTIMATE_YEAR = 2025
ESTIMATE_FILE = re.compile(r"population_estimates_data\.xlsx$", re.I)
ESTIMATE_SHEET = "Pop_Stats_Summary"
ESTIMATE_COLUMN = f"Population - {ESTIMATE_YEAR}"
SOURCE_2025 = ("OCHA, Common Operational Dataset -- population statistics (cod-ps-ssd), South "
               "Sudan population estimates by county for 2025")
LICENCE_2025 = "Creative Commons Attribution for Intergovernmental Organisations (CC BY-IGO)"
# The census's county (folded) -> the map's label, where the two differ by
# more than spelling folds away. Filled from the run's log.
COUNTIES: dict[str, str] = {
    "pijicanal": "Canal/Pigi",
    "lopa": "Lafon",            # Lopa/Lafon county, seat Lafon
}
CLAIMED = ("abyei",)
GROUP = re.compile(r"^([BMF])(\d{2})(\d{2})$")
OPEN = re.compile(r"^([BMF])(\d{2})PL$")


def fold(text: Any) -> str:
    stripped = unicodedata.normalize("NFKD", str(text or ""))
    return "".join(c for c in stripped.lower() if c.isalpha() and not unicodedata.combining(c))


def number(value: Any) -> int | None:
    got = uscb.number(value)
    return None if got is None else int(round(got))


def partition(names: list[str], sex: str, finest: bool = True
              ) -> list[tuple[int, int | None, int]]:
    """A run of this sex's age groups from 0 to an open last group, as (low,
    high, column): five-year groups where the sheet has them, and the open
    group with the highest start (``finest``) or the lowest, so a coarser sum
    beside the detail (65 and over beside 65-69 ... 95 and over) is never
    counted with it."""
    closed: dict[int, tuple[int, int]] = {}
    opens: dict[int, int] = {}
    for i, name in enumerate(names):
        if (m := GROUP.match(name)) and m.group(1) == sex:
            low, high = int(m.group(2)), int(m.group(3))
            if low not in closed or high < closed[low][0]:
                closed[low] = (high, i)
        elif (m := OPEN.match(name)) and m.group(1) == sex:
            opens[int(m.group(2))] = i
    run: list[tuple[int, int | None, int]] = []
    edge = 0
    top = (max(opens) if finest else min(opens)) if opens else None
    while top is not None and edge < top:
        if edge not in closed:
            raise SystemExit(f"south_sudan_census: no {sex} age group starts at {edge}")
        high, i = closed[edge]
        run.append((edge, high, i))
        edge = high + 1
    if top is None or edge != top:
        raise SystemExit(f"south_sudan_census: the {sex} age groups do not reach an open group")
    run.append((top, None, opens[top]))
    return run


def grouped_median(groups: list[tuple[int, int | None, int]]) -> float | None:
    base = sum(n for _, _, n in groups)
    half, before = base / 2, 0
    for low, high, people in groups:
        if before + people >= half:
            if high is None or people <= 0:
                return None
            return round(low + (half - before) / people * (high - low + 1), 1)
        before += people
    return None


def read(rows: list[list[Any]]) -> dict[str, Any]:
    names, _ = uscb.columns(rows)
    at = {n: i for i, n in enumerate(names) if n}
    runs = [tuple(partition(names, s, finest) for s in "BMF") for finest in (True, False)]
    out: dict[str, Any] = {"national": None, "states": {}, "counties": [], "notes": []}
    for row in rows[2:]:
        level = number(row[at["ADM_LEVEL"]])
        if level not in (0, 1, 2):
            continue
        name = str(row[at["AREA_NAME"]] or "").strip()
        total, men, women = (number(row[at[k]]) for k in ("BTOTL", "MTOTL", "FTOTL"))
        if total is None or men is None or women is None:
            raise SystemExit(f"south_sudan_census: {name}: no total, male or female count")
        if men + women != total:
            raise SystemExit(f"south_sudan_census: {name}: {men:,} men and {women:,} women "
                             f"do not make {total:,}")
        # The finest groups where they add up, else the coarser run (to 65
        # and over): a row whose detail above 65 is incomplete still gives a
        # median, which lies far below 65.
        groups = None
        for both, men_cols, women_cols in runs:
            trial = [(lo, hi, number(row[i]) or 0) for lo, hi, i in both]
            made = (sum(n for _, _, n in trial),
                    sum(number(row[i]) or 0 for _, _, i in men_cols),
                    sum(number(row[i]) or 0 for _, _, i in women_cols))
            if made == (total, men, women):
                groups = trial
                break
        if groups is None:
            raise SystemExit(f"south_sudan_census: {name}: no run of age groups makes its "
                             f"{total:,} people")
        note = str(row[at["AS_CMNT"]] or "").strip() if "AS_CMNT" in at else ""
        if note:
            out["notes"].append(f"{name}: {note}")
        unit = {"name": name, "total": total, "men": men, "women": women, "groups": groups,
                "state": str(row[at["ADM1_NAME"]] or "").strip(),
                "code": str(row[at["GEO_MATCH"]] or "").strip() if "GEO_MATCH" in at else name}
        if level == 0:
            out["national"] = unit
        elif level == 1:
            out["states"][fold(name)] = unit
        else:
            out["counties"].append(unit)
    for key, state in out["states"].items():
        made = sum(c["total"] for c in out["counties"] if fold(c["state"]) == key)
        if made != state["total"]:
            raise SystemExit(f"south_sudan_census: {state['name']}'s counties make {made:,}, "
                             f"not {state['total']:,}")
    if out["national"] and sum(s["total"] for s in out["states"].values()) != out["national"]["total"]:
        raise SystemExit("south_sudan_census: the states do not make the national total")
    return out


def figures(unit: dict[str, Any]) -> dict[str, Any]:
    return {"median": grouped_median(unit["groups"]),
            "ratio": round(1000 * unit["men"] / unit["women"])}


def read_estimates(rows: list[list[Any]]) -> dict[str, dict[str, Any]]:
    """County P-code -> {"name", "other", "state", "population"} from OCHA's
    estimates sheet, the Abyei Area's row left out."""
    head = next((i for i, r in enumerate(rows)
                 if {"Admin2_Pcode", ESTIMATE_COLUMN} <= {str(c or "").strip() for c in r}), None)
    if head is None:
        raise SystemExit(f"south_sudan_census: no column {ESTIMATE_COLUMN!r} by Admin2_Pcode")
    at = {str(c or "").strip(): j for j, c in enumerate(rows[head])}
    out: dict[str, dict[str, Any]] = {}
    for r in rows[head + 1:]:
        code = str(r[at["Admin2_Pcode"]] or "").strip()
        value = r[at[ESTIMATE_COLUMN]]
        if not code or value in (None, ""):
            continue
        name = str(r[at["Admin2"]] or "").strip()
        if any(c in fold(name) or c in fold(r[at["Admin1"]]) for c in CLAIMED):
            continue
        try:
            people = round(float(value))
        except (TypeError, ValueError):
            raise SystemExit(f"south_sudan_census: {code} {name}: {value!r} is no number")
        if code in out or people <= 0:
            raise SystemExit(f"south_sudan_census: {code} {name}: twice, or {people}")
        other = str(r[at["Admin2_Alternate_Name"]] or "").strip() \
            if "Admin2_Alternate_Name" in at else ""
        out[code] = {"name": name, "other": other, "population": people,
                     "state": str(r[at["Admin1"]] or "").strip()}
    return out


def bind_estimates(estimates: dict[str, dict[str, Any]], admin2: list[dict[str, Any]],
                   report: list[str]) -> dict[str, dict[str, Any]]:
    """The map's county id -> its estimate: by P-code, else by the one county
    of that name. Every county drawn must have one, or the run stops."""
    out: dict[str, dict[str, Any]] = {}
    left = dict(estimates)
    for shape in admin2:
        code = (shape.get("codes") or {}).get("pcode")
        if code in left:
            out[shape["id"]] = left.pop(code)
    for shape in admin2:
        if shape["id"] in out:
            continue
        hits = [c for c, e in left.items()
                if fold(shape["name"]) in (fold(e["name"]), fold(e["other"]))]
        if len(hits) == 1:
            out[shape["id"]] = left.pop(hits[0])
            report.append(f"estimate for {shape['name']!r} bound by name")
    missing = [s["name"] for s in admin2 if s["id"] not in out]
    for code, e in sorted(left.items()):
        report.append(f"estimate {code} {e['name']} ({e['population']:,}): no county drawn")
    if missing:
        raise SystemExit(f"south_sudan_census: no {ESTIMATE_YEAR} estimate for {missing}")
    return out


def unit_record(rid: str, shape: dict[str, Any], level: str, unit: dict[str, Any],
                what: str, estimate: dict[str, Any] | None = None) -> dict[str, Any]:
    f = figures(unit)
    extra: dict[str, Any] = {}
    sources = [{"field": "median_age/sex_ratio", "name": SOURCE,
                "url": uscb.dataset_url(DATASET), "year": YEAR, "license": LICENCE}]
    if estimate:
        extra = {"population": measure(estimate["population"], year=ESTIMATE_YEAR,
                                       source=SOURCE_2025),
                 "population_note": (f"OCHA's estimate of the population of {what} for "
                                     f"{ESTIMATE_YEAR}; the 2008 census counted "
                                     f"{unit['total']:,}.")}
        sources.append({"field": "population", "name": SOURCE_2025,
                        "url": uscb.dataset_url(ESTIMATES), "year": ESTIMATE_YEAR,
                        "license": LICENCE_2025})
    return record(
        rid, shape["name"], level=level, parent="SSD", country="SSD",
        match_by="shape_id", shape_id=shape["id"], **extra,
        median_age=measure(f["median"], unit="years", year=YEAR, source=SOURCE)
        if f["median"] is not None else None,
        median_age_note=(f"Interpolated within the five-year age group that holds the middle "
                         f"person, from the 2008 census's count by five-year age group of "
                         f"everyone in {what}: {unit['total']:,} people."),
        sex_ratio=measure(f["ratio"], unit="males_per_1000_females", year=YEAR, source=SOURCE),
        sex_ratio_note=(f"Men per 1,000 women at the 2008 census in {what}: {unit['men']:,} "
                        f"men and {unit['women']:,} women."),
        codes={"uscb": unit["code"]},
        sources=sources)


def build(table: dict[str, Any], admin1: list[dict[str, Any]], admin2: list[dict[str, Any]],
          estimates: dict[str, dict[str, Any]] | None = None
          ) -> tuple[list[dict[str, Any]], list[str]]:
    report: list[str] = []
    records: list[dict[str, Any]] = []
    bound = bind_estimates(estimates, admin2, report) if estimates else {}
    states = {fold(s["name"]): s for s in admin1}
    state_of: dict[str, dict[str, Any]] = {}
    for key, unit in sorted(table["states"].items()):
        shape = states.get(key)
        if shape is None:
            report.append(f"state {unit['name']!r}: no polygon")
            continue
        state_of[key] = shape
        # A state that counted a county of the claimed Abyei Area (Warrap)
        # is written without it: the map draws the Abyei Area apart.
        claimed = [c for c in table["counties"] if fold(c["state"]) == key
                   and any(w in fold(c["name"]) for w in CLAIMED)]
        what = unit["name"].title()
        if claimed:
            kept = [c for c in table["counties"] if fold(c["state"]) == key and c not in claimed]
            if len({tuple(g[:2] for g in c["groups"]) for c in kept}) != 1:
                raise SystemExit(f"south_sudan_census: {unit['name']}'s counties are not read "
                                 "in one set of age groups")
            unit = {**unit, "total": sum(c["total"] for c in kept),
                    "men": sum(c["men"] for c in kept), "women": sum(c["women"] for c in kept),
                    "groups": [(lo, hi, sum(c["groups"][i][2] for c in kept))
                               for i, (lo, hi, _) in enumerate(kept[0]["groups"])]}
            what = (f"{what} without "
                    + " and ".join(f"{c['name'].title()} County ({c['total']:,} people)"
                                   for c in claimed)
                    + ", in the Abyei Area, which Sudan and South Sudan both claim")
            report.append(f"state {unit['name']!r}: written without {[c['name'] for c in claimed]}")
        records.append(unit_record(f"SSD-2008-{key}", shape, "admin1", unit, what))
    free = {s["id"]: s for s in admin2}
    for unit in table["counties"]:
        key = fold(unit["name"])
        if any(c in key for c in CLAIMED):
            report.append(f"county {unit['name']!r}: in the claimed Abyei Area; not written")
            continue
        parent = state_of.get(fold(unit["state"]))
        wanted = fold(COUNTIES.get(key, unit["name"]))
        hits = [s for s in free.values() if fold(s["name"]) == wanted
                and (parent is None or s["parent"] == parent["id"])]
        if len(hits) != 1:
            report.append(f"county {unit['name']!r} ({unit['state'].title()}, "
                          f"{unit['total']:,}): {len(hits)} polygons")
            continue
        shape = hits[0]
        del free[shape["id"]]
        records.append(unit_record(f"SSD-2008-{unit['code']}", shape, "admin2", unit,
                                   f"{unit['name'].title()} County", bound.get(shape["id"])))
    for shape in free.values():
        report.append(f"polygon {shape['name']!r}: no county")
    return records, report


def estimate_sheets() -> dict[str, list[list[Any]]]:
    """Every sheet's rows in the newest estimates workbook OCHA's population
    dataset carries."""
    import openpyxl                                 # noqa: PLC0415
    found = sorted((r for r in uscb.package(ESTIMATES).get("resources", [])
                    if ESTIMATE_FILE.search(r.get("name") or "")),
                   key=lambda r: r.get("name") or "", reverse=True)
    if not found:
        raise SystemExit(f"south_sudan_census: {ESTIMATES} carries no estimates workbook")
    log(f"  estimates: {found[0].get('name')} ({found[0]['url']})")
    blob = http_get(found[0]["url"], binary=True, cache=False, timeout=180)
    book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    if ESTIMATE_SHEET not in book.sheetnames:
        raise SystemExit(f"south_sudan_census: no sheet {ESTIMATE_SHEET} in {found[0].get('name')}")
    return {name: [list(r) for r in book[name].iter_rows(values_only=True)]
            for name in book.sheetnames}


def check_states(estimates: dict[str, dict[str, Any]], rows: list[list[Any]]) -> int:
    """Each state's counties must make the total the workbook's summary gives
    it, to within rounding; the number of states checked."""
    totals = {fold(r[0]): float(r[1]) for r in rows[1:]
              if len(r) > 1 and r[0] and isinstance(r[1], (int, float))
              and "total" not in fold(r[0])}
    checked = 0
    for key, total in totals.items():
        mine = [e["population"] for e in estimates.values() if fold(e["state"]) == key]
        made = sum(mine)
        if any(c in key for c in CLAIMED):
            continue
        # Each county's figure rounded to a whole person.
        if abs(made - total) > 1 + len(mine) / 2:
            raise SystemExit(f"south_sudan_census: the counties of {key} make {made:,}, the "
                             f"summary {total:,.0f}")
        checked += 1
    return checked


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--run", help="a label that tells one dispatch from another; not used")
    parser.parse_args()
    import openpyxl                                 # noqa: PLC0415
    blob = http_get(uscb.workbook_url(DATASET), binary=True, cache=False, timeout=180)
    book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    meta = " ".join(str(c) for r in uscb.sheet_rows(book, "Metadata") for c in r if c)
    if str(YEAR) not in meta:
        raise SystemExit("south_sudan_census: the workbook's metadata never names 2008")
    table = read(uscb.sheet_rows(book, "Age-Sex"))
    for line in table["notes"][:40]:
        log("  note: " + line[:300])
    for unit in [*table["states"].values(), *table["counties"]]:
        f = figures(unit)
        log(f"  {unit['state'][:24]:24} {unit['name'][:28]:28} {unit['total']:>9,}  median "
            f"{f['median']}  men/1000 women {f['ratio']}")
    sheets = estimate_sheets()
    estimates = read_estimates(sheets[ESTIMATE_SHEET])
    states = next((rows for name, rows in sheets.items() if "state" in name.lower()), [])
    log(f"  {len(estimates)} counties' estimates for {ESTIMATE_YEAR}, making "
        f"{sum(e['population'] for e in estimates.values()):,} without the Abyei Area; "
        f"{check_states(estimates, states)} states' totals made by their counties")
    admin1 = json.loads((SITE / "admin1" / "SSD.units.json").read_text())
    admin2 = json.loads((SITE / "admin2" / "SSD.units.json").read_text())
    records, report = build(table, admin1, admin2, estimates)
    for line in report:
        log("  " + line)
    log(f"  {sum(r['level'] == 'admin1' for r in records)} states and "
        f"{sum(r['level'] == 'admin2' for r in records)} of {len(admin2)} counties bound")
    write_json(PROCESSED / OUT, records)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
