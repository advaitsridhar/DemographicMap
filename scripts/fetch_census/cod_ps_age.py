#!/usr/bin/env python3
"""Median age and sex ratio from OCHA's COD-PS, where no census gives them.

Most of Latin America's first- and second-level units had a population and
nothing else: Peru's 26 departments and 196 provinces, Venezuela's states and
municipios, Paraguay's, El Salvador's, Uruguay's. OCHA's Common Operational
Dataset for each of those countries is the national statistical office's own
table (cod_ps.py already reads its totals) and splits every unit by sex and by
five-year age group. That is enough for two of the map's questions:

- **sex ratio**: males per 1,000 females, from the table's F and M totals;
- **median age**: interpolated within the five-year group that holds the
  middle person, the standard way to take a median from grouped ages.

**What a figure is.** For a reference year after the country's last census the
table is the office's projection from that census, and the note on every unit
says which, quoting the dataset's own methodology line where it has one.
Venezuela's is the 2011 census itself. So these fill gaps and never replace a
census's own median or ratio: the build lists this file in FILL_ONLY, whatever
year either figure is for.

**Checks, per unit, before anything is written:**
- the age groups start at 0 and run without a gap to an open-ended last group;
- they add up to the unit's total, allowing only for a column of people of
  unstated age, which is left out of the median;
- women and men add up to the total;
- the median falls inside a closed group, never the open-ended one.
A unit failing any of these is left out and the log says why. Which of the
map's levels a table describes is measured by its names against the map's own
(cod_ps.which_level), not taken from the file's "adm2".

Usage:
    python -m scripts.fetch_census.cod_ps_age [--only PER,VEN]
"""

from __future__ import annotations

import argparse
import csv
import io
import re
import urllib.request
from typing import Any

from ._shared import PROCESSED, log, record, write_json
from .cod_ps import (DATASET_PAGE, HEADERS, TIMEOUT, catalogue, dataset_year, is_usable,
                     iso3, licence, name_column, reference_year, which_level)

OUT = "cod_ps_age.json"

# The Americas, where the map's regions were emptiest (the owner's ask of
# 25 September 2026). The reader is not specific to them: --only widens it.
AMERICAS = frozenset({
    "ARG", "ATG", "BHS", "BLZ", "BOL", "BRA", "BRB", "CHL", "COL", "CRI", "CUB", "DMA",
    "DOM", "ECU", "GRD", "GTM", "GUY", "HND", "HTI", "JAM", "KNA", "LCA", "MEX", "NIC",
    "PAN", "PER", "PRY", "SLV", "SUR", "TTO", "URY", "VCT", "VEN"})

# A table, by the file's own name for its level: "per_admpop_adm1_2022",
# "CMR_admpop2_2025.csv".
TABLE_LEVEL = re.compile(r"adm(?:pop)?_?([123])(?!\d)")
AGE = re.compile(r"^([TFM])_?(\d{1,2})_(\d{1,2})$", re.I)
OPEN = re.compile(r"^([TFM])_?(\d{1,2})_?plus$", re.I)
TOTAL = re.compile(r"^([TFM])(?:_?TL)?$", re.I)
UNSTATED = re.compile(r"^T_?(?:unidentified|unknown|unstated|ns|nd)$", re.I)
PCODE = re.compile(r"^adm(?:in)?_?([123])_?pcode$", re.I)
# Rounding in a projection's cells, not a disagreement.
TOLERANCE = 0.005
# Units the boundary file names differently: renamed since the table's year,
# or abbreviated in it.
ALIASES = {
    ("VEN", "Distrito Federal"): ["Distrito Capital"],   # renamed 1999
    ("VEN", "Vargas"): ["La Guaira"],                      # renamed 2019
    ("HND", "Islas de La Bahia"): ["Bay Islands"],
    ("PRY", "Pdte. Hayes"): ["Presidente Hayes"],
}
# Rows whose unit is not the shape of the same name. Peru's table has the
# department of Lima whole, Metropolitan Lima's ten million inside it; the map
# draws Metropolitan Lima and Lima Region apart, and on Lima Region the
# department's figures would be a different place's.
EXCLUDE = {("PER", "admin1", "Lima"): "the whole department, Metropolitan Lima inside it"}


def number(cell: Any) -> float | None:
    text = str(cell if cell is not None else "").strip().replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def age_columns(columns: list[str]) -> dict[str, Any] | None:
    """Which columns hold what: totals by sex, age groups by sex, the open group.

    None when the table has no age breakdown at all, or one that does not run
    from 0 without a gap -- such a table cannot give a median.
    """
    totals: dict[str, str] = {}
    groups: dict[str, dict[tuple[int, int], str]] = {}
    opens: dict[str, tuple[int, str]] = {}
    unstated = None
    for column in columns:
        c = column.strip()
        if m := AGE.match(c):
            groups.setdefault(m.group(1).upper(), {})[(int(m.group(2)), int(m.group(3)))] = c
        elif m := OPEN.match(c):
            opens[m.group(1).upper()] = (int(m.group(2)), c)
        elif m := TOTAL.match(c):
            totals[m.group(1).upper()] = c
        elif UNSTATED.match(c):
            unstated = c
    sex = "T" if "T" in groups and "T" in opens else None
    if sex is None and all(s in groups and s in opens for s in "FM"):
        sex = "FM"
    if sex is None or "F" not in totals or "M" not in totals:
        return None
    for s in sex:
        bounds = sorted(groups[s])
        edge = 0
        for low, high in bounds:
            if low != edge or high < low:
                return None
            edge = high + 1
        if opens[s][0] != edge:
            return None
    if sex == "FM" and sorted(groups["F"]) != sorted(groups["M"]):
        return None
    return {"sexes": sex, "totals": totals, "groups": groups, "opens": opens,
            "unstated": unstated}


def grouped_median(counts: list[tuple[int, int | None, float]]) -> float | None:
    """The median of ``(low, high, people)`` groups; None if it is in the open one."""
    base = sum(n for _, _, n in counts)
    if base <= 0:
        return None
    half, before = base / 2, 0.0
    for low, high, people in counts:
        if before + people >= half:
            if high is None or people <= 0:
                return None
            return round(low + (half - before) / people * (high - low + 1), 1)
        before += people
    return None


def unit_figures(row: dict[str, Any], cols: dict[str, Any]) -> tuple[dict[str, Any] | None, str]:
    """A row's median and ratio, or None with the reason."""
    women, men = number(row.get(cols["totals"]["F"])), number(row.get(cols["totals"]["M"]))
    if not women or not men:
        return None, "no female or male total"
    total = number(row.get(cols["totals"]["T"])) if "T" in cols["totals"] else women + men
    if not total:
        return None, "no total"
    if abs(women + men - total) > max(2, TOLERANCE * total):
        return None, f"women {women:,.0f} and men {men:,.0f} do not add up to {total:,.0f}"
    first = cols["sexes"][0]
    bands = sorted(cols["groups"][first])
    counts: list[tuple[int, int | None, float]] = []
    for low, high in bands:
        cells = [number(row.get(cols["groups"][s][(low, high)])) for s in cols["sexes"]]
        if any(c is None for c in cells):
            return None, f"no count for ages {low}-{high}"
        counts.append((low, high, sum(cells)))
    cells = [number(row.get(cols["opens"][s][1])) for s in cols["sexes"]]
    if any(c is None for c in cells):
        return None, "no count for the open-ended age group"
    counts.append((cols["opens"][first][0], None, sum(cells)))
    unstated = number(row.get(cols["unstated"])) or 0 if cols["unstated"] else 0
    aged = sum(n for _, _, n in counts)
    if abs(aged + unstated - total) > max(2, TOLERANCE * total):
        return None, f"its age groups add up to {aged:,.0f}, not {total:,.0f}"
    median = grouped_median(counts)
    if median is None:
        return None, "the middle person is in the open-ended age group"
    return {"median": median, "ratio": round(1000 * men / women), "unstated": unstated,
            "aged": aged}, ""


def tables(package: dict[str, Any]) -> list[dict[str, Any]]:
    """Every table in the dataset that names its level: CSVs, and workbook sheets."""
    out: list[dict[str, Any]] = []
    for resource in package.get("resources") or ():
        name = str(resource.get("name") or "")
        low = name.lower()
        if re.search(r"gazetteer|admgz|admgaz|boundar", low):
            continue
        is_csv, is_book = low.endswith(".csv"), low.endswith((".xlsx", ".xlsm"))
        if not (is_book or (is_csv and TABLE_LEVEL.search(low))):
            continue
        try:
            with urllib.request.urlopen(urllib.request.Request(
                    str(resource.get("url")), headers=HEADERS), timeout=TIMEOUT) as fh:
                body = fh.read()
        except Exception as err:                     # noqa: BLE001 -- reported
            log(f"    unreadable: {name}: {type(err).__name__}: {err}")
            continue
        if is_csv:
            reader = csv.DictReader(io.StringIO(body.decode("utf-8-sig", "replace")))
            out.append({"label": name, "level": TABLE_LEVEL.search(low).group(1),
                        "columns": [c.strip() for c in reader.fieldnames or []],
                        "rows": [{(k or "").strip(): v for k, v in r.items()} for r in reader],
                        "method": ""})
            continue
        import openpyxl
        try:
            book = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        except Exception as err:                     # noqa: BLE001 -- reported
            log(f"    unreadable workbook: {name}: {type(err).__name__}: {err}")
            continue
        method = ""
        for sheet in book.worksheets:
            if "method" in sheet.title.lower():
                for row in sheet.iter_rows(values_only=True, max_row=20):
                    if (row and len(row) > 1
                            and str(row[0] or "").strip().lower().startswith("methodology used")):
                        method = str(row[1] or "").strip()
        for sheet in book.worksheets:
            found = TABLE_LEVEL.search(sheet.title.lower().replace(" ", ""))
            if not found:
                continue
            rows = sheet.iter_rows(values_only=True)
            header = [str(c).strip() if c is not None else "" for c in next(rows, [])]
            body_rows = [{h: v for h, v in zip(header, r) if h} for r in rows
                         if any(v not in (None, "") for v in r)]
            out.append({"label": f"{name} {sheet.title}_", "level": found.group(1),
                        "columns": [h for h in header if h], "rows": body_rows,
                        "method": ""})
        for table in out:
            if table["label"].startswith(name):
                table["method"] = method
    return out


def country_records(package: dict[str, Any]) -> list[dict[str, Any]]:
    code = iso3(package)
    stub = str(package.get("name") or "")
    log(f"  {stub} ({code})")
    if not is_usable(package):
        log(f"    refused: licence {licence(package)}")
        return []
    method = ""
    newest: dict[str, tuple[int, dict[str, Any], dict[str, Any]]] = {}
    for table in tables(package):
        method = method or table["method"]
        cols = age_columns(table["columns"])
        if cols is None:
            continue
        year = reference_year(table["columns"], [{k: str(v) for k, v in r.items()}
                                                  for r in table["rows"]], table["label"])
        year = year or dataset_year(package)
        if year is None:
            log(f"    refused {table['label']}: no reference year")
            continue
        level = table["level"]
        if level not in newest or year > newest[level][0]:
            newest[level] = (year, table, cols)
    if not newest:
        log("    no table with a full breakdown by sex and five-year age group")
        return []
    out: list[dict[str, Any]] = []
    claimed: set[str] = set()
    for cod_level in sorted(newest):
        year, table, cols = newest[cod_level]
        unit_col = name_column(table["columns"], cod_level)
        if unit_col is None:
            log(f"    refused {table['label']}: no level-{cod_level} name column")
            continue
        names = [str(r.get(unit_col) or "").strip() for r in table["rows"]]
        level, why = which_level(code, names)
        if level is None or level in claimed:
            log(f"    refused {table['label']}: {why if level is None else level + ' already read'}")
            continue
        claimed.add(level)
        parent_col = name_column(table["columns"], "1")
        pcode_col = next((c for c in table["columns"]
                          if (m := PCODE.match(c.strip())) and m.group(1) == cod_level), None)
        source = (f"OCHA, Common Operational Dataset -- population statistics ({stub}), "
                  f"reference year {year}")
        basis = (f" The office's method: \"{method[:220].rstrip()}"
                 f"{'...' if len(method) > 220 else ''}\"" if method else "")
        written = refused = 0
        seen: set[str] = set()
        # A name is ambiguous only within one parent: Mexico has a Benito
        # Juarez in several states, and the build tells those apart by state.
        def place(row: dict[str, Any]) -> tuple[str, str]:
            parent = str(row.get(parent_col) or "").strip() \
                if level == "admin2" and parent_col else ""
            return parent, str(row.get(unit_col) or "").strip()
        named: dict[tuple[str, str], int] = {}
        for row in table["rows"]:
            named[place(row)] = named.get(place(row), 0) + 1
        for row in table["rows"]:
            name = str(row.get(unit_col) or "").strip()
            pcode = str(row.get(pcode_col) or name).strip() if pcode_col else name
            if not name or pcode in seen:
                continue
            seen.add(pcode)
            # Two rows under one name at one level cannot both be the shape of
            # that name, and nothing says which is: Nicaragua's table calls
            # two first-level units "Las Minas".
            if named[place(row)] > 1:
                log(f"    left out {name}: {named[place(row)]} rows carry the name")
                continue
            if (code, level, name) in EXCLUDE:
                log(f"    left out {name}: {EXCLUDE[(code, level, name)]}")
                continue
            figures, why_not = unit_figures(row, cols)
            if figures is None:
                refused += 1
                log(f"    left out {name}: {why_not}")
                continue
            unstated = (f" {figures['unstated']:,.0f} people of unstated age are left out "
                        f"of the median." if figures["unstated"] else "")
            out.append(record(
                f"{code}-CODPSAGE-{pcode}", name, level=level, parent=code, country=code,
                parent_name=(str(row.get(parent_col) or "").strip() or None)
                if level == "admin2" and parent_col else None,
                codes={"pcode": pcode}, aliases=ALIASES.get((code, name)),
                median_age={"value": figures["median"], "unit": "years", "year": year,
                            "source": source},
                median_age_note=(
                    f"Interpolated within the five-year age group that holds the middle "
                    f"person, from the age breakdown in OCHA's COD-PS for this country, "
                    f"reference year {year}, which is the national statistical office's "
                    f"table. For a year after the last census it is a projection, not a "
                    f"count.{basis}{unstated}"),
                sex_ratio={"value": figures["ratio"], "unit": "males_per_1000_females",
                           "year": year, "source": source},
                sex_ratio_note=(f"Males per 1,000 females in OCHA's COD-PS for this country, "
                                f"reference year {year}; for a year after the last census, "
                                f"the office's projection.{basis}"),
                sources=[{"field": "median_age/sex_ratio", "name": source,
                          "url": DATASET_PAGE.format(stub=stub), "year": year,
                          "license": licence(package)}]))
            written += 1
        log(f"    {level}: {written} unit(s) from {table['label']} ({why}); "
            f"{refused} left out")
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--only", default="",
                    help="comma-separated ISO3s; default the Americas")
    args = ap.parse_args()
    only = {x.upper() for x in args.only.split(",") if x} or AMERICAS
    packages = [p for p in catalogue() if iso3(p) in only]
    log(f"cod_ps_age: {len(packages)} dataset(s) for {len(only)} countries")
    out: list[dict[str, Any]] = []
    for package in sorted(packages, key=lambda p: str(p.get("name"))):
        out.extend(country_records(package))
    if not out:
        raise SystemExit("cod_ps_age: nothing read")
    by_country: dict[str, int] = {}
    for rec in out:
        by_country[rec["country"]] = by_country.get(rec["country"], 0) + 1
    log("  written: " + ", ".join(f"{c} {n}" for c, n in sorted(by_country.items())))
    write_json(PROCESSED / OUT, out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
