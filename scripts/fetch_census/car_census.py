#!/usr/bin/env python3
"""Central African Republic -- the 2021 census mapping's count, and 2016 ages.

The map draws the 17 prefectures and 72 sub-prefectures of the 2003 census
(RGPH03), whose religion, language and ethnicity car_prefecture.json reads.
Their populations were OCHA's 2015 projection (69 of 72), and nothing gave a
median age or a sex ratio at either level.

**Population, 2021.** The Institut Centrafricain des Statistiques et des
Etudes Economiques et Sociales (ICASEES) counted the population at the mapping
stage of the fourth census (RGPH-4 cartography, reference date 1 March 2021)
and published it by commune of the 2020 division, as the U.S. Census Bureau's
workbook for the CAR tabulates it (sheet ``Population``: 20 prefectures, 79
sub-prefectures, 181 communes). The polygons are the 2003 division's, so each
2021 commune is found among the 2003 communes of the same name inside the
prefecture it descends from (Mambéré, Lim-Pendé and Ouham-Fafa were made out
of Mambéré-Kadéï, Ouham-Pendé and Ouham in 2020), a commune with no 2003
namesake goes with its 2021 sub-prefecture's other communes, and Bégoua --
made a sub-prefecture of Bangui in 2020, with no 2003 namesake and no
siblings -- goes to Bimbo, whose 2003 polygon surrounds Bangui and holds it.
A polygon takes the sum of its communes, a prefecture the sum of its polygons.

OCHA's edition of the same table (``icasees_cartocensitaire2021_formatocha_
public_V2.xlsx`` in cod-ps-caf) assigns the count to the 2003 communes itself
(sheet ``PrevCommune``), and the two assignments are compared polygon by
polygon on every run. They differ only where OCHA files a commune under a
namesake in another prefecture -- the town of Nola (Sangha-Mbaéré) under
Lobaye's commune of Nola -- and any other difference stops the run.

**No sex ratio from it.** The mapping table splits every unit's count into
men and women in one fixed proportion (49 to 51 to the person, urban and rural
alike), so it describes no unit's sexes; the run checks that this is still so
and says it on every record.

**Prefectures' ages and sexes, 2016.** ICASEES's *Annuaire statistique de la
République Centrafricaine 2016* (2018), tables 110-116, gives the
administrative population of each of the 17 prefectures by sex and five-year
age group at 31 December 2016 -- an estimate, not a count -- as the Census
Bureau's workbook tabulates it (sheet ``Age-Sex``). A median is taken by
interpolation within the five-year group holding the middle person, and only
for a prefecture whose age proportions are its own: one that shares them with
the country or with another prefecture to within COPY_TOLERANCE points in
every group is left out (cod_ps_age's test).

Usage:
    python -m scripts.fetch_census.car_census [--dump]
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
from collections import defaultdict
from typing import Any

from ._shared import NOT_AVAILABLE, PROCESSED, gap, http_get, http_json, log, measure, record, write_json
from .binding import bind, fold
from .cod_ps_age import COPY_TOLERANCE, grouped_median, same_structure
from .uscb import dataset_url, number, sheet_rows, workbook_url


def series(names: list[str]) -> str:
    """'A', 'A and B', 'A, B and C'."""
    names = list(names)
    return names[0] if len(names) == 1 else ", ".join(names[:-1]) + " and " + names[-1]

USCB_DATASET = "car-subnational-boundaries-and-tabular-data"
OCHA_DATASET = "cod-ps-caf"
HDX_API = "https://data.humdata.org/api/3/action"
OUT = "car_census.json"
SITE = PROCESSED.parent.parent / "site" / "data"
COUNT_YEAR = 2021
AGE_YEAR = 2016
COUNT_SOURCE = ("Institut Centrafricain des Statistiques et des Etudes Economiques et Sociales "
                "(ICASEES), RGPH-4: Principaux résultats de la cartographie 2021 (2022), "
                "population by commune, as OCHA publishes it")
COUNT_SOURCE_USCB = ("Institut Centrafricain des Statistiques et des Etudes Economiques et "
                     "Sociales (ICASEES), RGPH-4: Principaux résultats de la cartographie 2021 "
                     "(2022), population by prefecture, sub-prefecture and commune, as tabulated "
                     "by the U.S. Census Bureau")
AGE_SOURCE = ("Institut Centrafricain des Statistiques et des Etudes Economiques et Sociales "
              "(ICASEES), Annuaire statistique de la République Centrafricaine 2016 (2018), "
              "tables 110-116, as tabulated by the U.S. Census Bureau")
LICENCE = "CC BY-IGO, published via HDX"

# The 2021 prefectures made in 2020 out of 2003 ones (checked on every run).
NEW_PREFECTURES = {"mambere": "mamberekadei", "limpende": "ouhampende", "ouhamfafa": "ouham"}
# A 2021 commune with no 2003 namesake and no siblings to place it by, and the
# 2003 sub-prefecture whose polygon holds it: Bégoua, 12 km north of Bangui's
# centre, inside the 3,037 km2 polygon drawn as Bimbo, which surrounds the
# 71 km2 city (measured on the CGAZ polygons).
PLACED = {("bangui", "begoua"): ("ombellampoko", "bimbo")}

# Sub-prefecture names the boundary file spells otherwise, as the 2003 census
# writes them (car_prefecture's list) and as OCHA's table does -- Djemah,
# Dede-Makouba, and Bantangafo, a doubled letter. Keyed by the folded name.
SP_ALIASES = {fold(k): v for k, v in {
    "Aba": "Abba", "Bossemptele": "Bossemtélé", "Nanga-Boguila": "Nagha Boguila",
    "Ndjoukou": "Djoukou", "Djemah": "Djéma", "Dede-Makouba": "Dédé-Mokouba",
    "Bantangafo": "Batangafo"}.items()}


def spelled(name: str) -> str:
    """A sub-prefecture's name as the boundary file writes it."""
    return SP_ALIASES.get(fold(name), name)


# A polygon whose whole count is a figure the table also prints for another
# unit, so it cannot be taken for this ground's own: Dédé-Mokouba's one
# commune, Haute-Kadéi, has exactly Mobaye commune's 23,708 (Basse-Kotto), in
# both the Census Bureau's tabulation and OCHA's. Which of the two is the copy
# cannot be told (the table splits every unit 49 to 51, so that the men and
# the women agree too says nothing more). The sentence goes on the population
# as its reason, and beside it as its note: OCHA's 2015 projection for
# Dédé-Mokouba (the shared patch's alias) is shown in its place, and the note
# says why the 2021 count is not.
SUSPECT = {fold("Dédé-Mokouba"): (
    "The 2021 census mapping's table gives Haute-Kadéi, this sub-prefecture's one commune, "
    "exactly the count it gives Mobaye commune in Basse-Kotto (23,708 people), and which of "
    "the two is a copy of the other cannot be told, so that count cannot be taken for this "
    "sub-prefecture's own and is not shown here.")}

# The fixed proportion of men in every unit of the mapping table, and how far
# a unit may sit from it (rounding of whole people) before the table is taken
# to describe sexes after all.
FIXED_MALE_SHARE = 0.49
FIXED_SLACK = 0.0005

SEX_GAP = ("The 2021 census mapping's table splits every unit's count into men and women in "
           "one fixed proportion (49 to 51), so it gives no sex ratio of this sub-prefecture's "
           "own; ICASEES's 2016 estimates by sex (Annuaire statistique de la République "
           "Centrafricaine 2016) stop at the prefecture.")
# Added where the 2025 table gives every sub-prefecture of a prefecture one
# share of men (checked on each run, one_sex_split).
SEX_GAP_2025 = (" ICASEES's projection for 2025 gives every sub-prefecture of a prefecture the "
                "same proportion of men, so it gives none of this one's own either.")
AGE_GAP = ("ICASEES's 2016 estimates by age (Annuaire statistique de la République "
           "Centrafricaine 2016) stop at the 17 prefectures and its 2021 census mapping counts "
           "carry no ages, so no median age is published for the sub-prefecture.")
SEAT_GAP = ("No published list of the sub-prefectures' seats has been found, so none is "
            "shown; the census tables name the sub-prefectures and their communes, not their "
            "seats.")


def table(rows: list[list[Any]]) -> list[dict[str, Any]]:
    """A USCB sheet as dicts keyed by its field-name row (row 0); row 1 is aliases."""
    names = ["" if v is None else str(v).strip() for v in rows[0]]
    return [{n: v for n, v in zip(names, row) if n} for row in rows[2:]
            if any(v not in (None, "") for v in row)]


def plain(rows: list[list[Any]]) -> list[dict[str, Any]]:
    """An OCHA sheet: one header row."""
    names = ["" if v is None else str(v).strip() for v in rows[0]]
    return [{n: v for n, v in zip(names, row) if n} for row in rows[1:]
            if any(v not in (None, "") for v in row)]


def level(row: dict[str, Any]) -> int | None:
    value = number(row.get("ADM_LEVEL"))
    return int(value) if value is not None else None


def text(row: dict[str, Any], key: str) -> str:
    return " ".join(str(row.get(key) or "").split())


def prev_communes(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """OCHA's PrevCommune sheet: the 2021 count on each 2003 commune."""
    out = []
    for row in rows:
        total = number(row.get("T_TL"))
        if total is None or not text(row, "Pcode3"):
            continue
        out.append({"pref_code": text(row, "Pcode1"), "pref": text(row, "PrevPréfecture"),
                    "sp_code": text(row, "Pcode2"), "sp": text(row, "PrevSous-préfecture"),
                    "code": text(row, "Pcode3"), "name": text(row, "PrevCommune"),
                    "total": int(round(total)),
                    "men": number(row.get("M_TL")) or 0.0, "women": number(row.get("F_TL")) or 0.0})
    return out


def counts_2021(rows: list[dict[str, Any]]) -> dict[str, Any]:
    """The USCB Population sheet: the same count on the 2021 frame."""
    out: dict[str, Any] = {"country": None, "prefectures": {}, "sps": {}, "communes": []}
    for row in rows:
        figures = tuple(number(row.get(k)) for k in ("POP_BTOTL", "POP_MTOTL", "POP_FTOTL"))
        if any(f is None for f in figures):
            continue
        total, men, women = (int(f) for f in figures)
        unit = {"name": text(row, "AREA_NAME"), "pref": text(row, "ADM1_NAME"),
                "sp": text(row, "ADM2_NAME"), "total": total, "men": men, "women": women}
        lvl = level(row)
        if lvl == 0:
            out["country"] = unit
        elif lvl == 1:
            out["prefectures"][fold(unit["name"])] = unit
        elif lvl == 2:
            out["sps"][(fold(unit["pref"]), fold(unit["name"]))] = unit
        elif lvl == 3:
            out["communes"].append(unit)
    return out


def communes_2003(rows: list[dict[str, Any]]) -> list[tuple[str, str, str]]:
    """(prefecture, sub-prefecture, commune), folded, from a 2003-vintage sheet."""
    return [(fold(text(r, "ADM1_NAME")), fold(text(r, "ADM2_NAME")), fold(text(r, "AREA_NAME")))
            for r in rows if level(r) == 3]


def own_crosswalk(c21: dict[str, Any], g03: list[tuple[str, str, str]]
                  ) -> tuple[dict[tuple[str, str], list[dict[str, Any]]], list[str]]:
    """2003 (pref, sp) -> the 2021 communes it holds, by commune names; and what failed."""
    by_name: dict[tuple[str, str], list[tuple[str, str]]] = defaultdict(list)
    anywhere: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for pref, sp, commune in g03:
        by_name[(pref, commune)].append((pref, sp))
        anywhere[commune].append((pref, sp))
    placed: dict[int, tuple[str, str]] = {}
    for i, unit in enumerate(c21["communes"]):
        pref, name = fold(unit["pref"]), fold(unit["name"])
        if (pref, name) in PLACED:
            placed[i] = PLACED[(pref, name)]
            continue
        hits = by_name.get((NEW_PREFECTURES.get(pref, pref), name), [])
        if len(hits) != 1:
            hits = anywhere.get(name, []) if not hits else hits
        if len(hits) == 1:
            placed[i] = hits[0]
    problems: list[str] = []
    for i, unit in enumerate(c21["communes"]):
        if i in placed or not unit["total"]:
            continue
        key = (fold(unit["pref"]), fold(unit["sp"]))
        siblings = {placed[j] for j, u in enumerate(c21["communes"])
                    if j in placed and (fold(u["pref"]), fold(u["sp"])) == key}
        if len(siblings) == 1:
            placed[i] = siblings.pop()
        else:
            problems.append(f"{unit['name']} ({unit['sp']}, {unit['pref']}): no 2003 namesake "
                            f"and no one place for its siblings")
    held: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for i, key in sorted(placed.items()):
        held[key].append(c21["communes"][i])
    return dict(held), problems


def fixed_split(c21: dict[str, Any]) -> int:
    """How many units of the mapping table depart from the fixed 49:51 split."""
    off = 0
    for unit in [*c21["prefectures"].values(), *c21["sps"].values(), *c21["communes"]]:
        if unit["total"] and abs(unit["men"] / unit["total"] - FIXED_MALE_SHARE) > max(
                FIXED_SLACK, 1.5 / unit["total"]):
            off += 1
    return off


def twin_counts(groups: dict[str, list[tuple[str, str]]],
                placement: dict[tuple[str, str], list[dict[str, Any]]]
                ) -> dict[int, list[tuple[str, dict[str, Any]]]]:
    """Counts of 1,000 or more the table gives two communes or more: {count: [(polygon,
    commune)]}. Mobaye's and Haute-Kadéi's 23,708, Lissa's and Pouyamba's 8,825."""
    by_count: dict[int, list[tuple[str, dict[str, Any]]]] = defaultdict(list)
    for sid, keys in sorted(groups.items()):
        for key in keys:
            for unit in placement[key]:
                if unit["total"] >= 1000:
                    by_count[unit["total"]].append((sid, unit))
    return {count: found for count, found in by_count.items() if len(found) > 1}


def twin_sentence(sid: str, twins: dict[int, list[tuple[str, dict[str, Any]]]],
                  shapes: dict[str, dict[str, Any]], prefectures: dict[str, dict[str, Any]],
                  pref_of: dict[str, str]) -> str:
    """What a sub-prefecture's note says of its communes that share a count with another."""
    out = []
    for count, found in sorted(twins.items()):
        mine = [u for s, u in found if s == sid]
        others = [(s, u) for s, u in found if s != sid]
        if not mine:
            continue
        if not others:
            names = " and ".join(u["name"].title() for u in mine)
            out.append(f"The 2021 mapping's table gives two of its communes, {names}, the same "
                       f"count ({count:,} people each); one may be a copy of the other, and both "
                       f"are counted here.")
            continue
        s, other = others[0]
        pref = prefectures.get(pref_of.get(s, ""), {}).get("name", "")
        out.append(f"The 2021 mapping's table gives its commune of {mine[0]['name'].title()} "
                   f"exactly the count it gives {other['name'].title()}, a commune of "
                   f"{shapes[s]['name']} in "
                   f"{pref} ({count:,} people); which of the two is a copy of the other cannot "
                   f"be told, and {mine[0]['name'].title()}'s is counted here.")
    return " ".join(out)


def twin_sentence_prefecture(key: str, twins: dict[int, list[tuple[str, dict[str, Any]]]],
                             shapes: dict[str, dict[str, Any]],
                             prefectures: dict[str, dict[str, Any]],
                             pref_of: dict[str, str]) -> str:
    """What a prefecture's note says of the shared counts its sum includes."""
    out = []
    for count, found in sorted(twins.items()):
        mine = [(s, u) for s, u in found if pref_of.get(s) == key]
        others = [(s, u) for s, u in found if pref_of.get(s) != key]
        if not mine:
            continue
        if not others:
            names = " and ".join(f"{u['name'].title()} ({shapes[s]['name']})" for s, u in mine)
            out.append(f"The sum includes two communes the 2021 mapping's table gives the same "
                       f"count, {names}, {count:,} people each; one may be a copy of the other.")
            continue
        s, unit = mine[0]
        o, other = others[0]
        where = prefectures.get(pref_of.get(o, ""), {}).get("name", "")
        withheld = SUSPECT.get(fold(shapes[s]["name"]))
        out.append(f"The sum includes the {count:,} people the 2021 mapping's table gives "
                   f"{unit['name'].title()} ({shapes[s]['name']}), the same count it gives "
                   f"{other['name'].title()} ({shapes[o]['name']}, {where}); which of the two "
                   f"is a copy of the other cannot be told"
                   + (f", and {shapes[s]['name']} itself shows no 2021 count for that reason."
                      if withheld else "."))
    return " ".join(out)


def age_table(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    """{prefecture (folded): {"groups": [[low, high, people]...], "men", "women"}}."""
    out: dict[str, dict[str, Any]] = {}
    for row in rows:
        if level(row) != 1:
            continue
        groups: list[list[Any]] = []
        for low in range(0, 80, 5):
            cell = number(row.get(f"B{low:02d}{low + 4:02d}"))
            if cell is None:
                break
            groups.append([low, low + 4, cell])
        last = number(row.get("B80PL"))
        men, women, total = (number(row.get(k)) for k in ("MTOTL", "FTOTL", "BTOTL"))
        if len(groups) != 16 or last is None or not men or not women or not total:
            continue
        groups.append([80, None, last])
        out[fold(text(row, "AREA_NAME"))] = {"name": text(row, "AREA_NAME"), "groups": groups,
                                             "men": men, "women": women, "total": total}
    return out


def own_structures(ages: dict[str, dict[str, Any]]) -> tuple[set[str], list[str]]:
    """The prefectures whose age proportions are their own, and the report."""
    whole: dict[tuple[int, Any], float] = defaultdict(float)
    for entry in ages.values():
        for low, high, people in entry["groups"]:
            whole[(low, high)] += people
    national = [[lo, hi, n] for (lo, hi), n in whole.items()]
    report: list[str] = []
    copied: set[str] = set()
    keys = sorted(ages)
    for k in keys:
        if same_structure(ages[k]["groups"], national):
            copied.add(k)
            report.append(f"  ages: {ages[k]['name']} carries the national proportions")
    for i, a in enumerate(keys):
        for b in keys[i + 1:]:
            if same_structure(ages[a]["groups"], ages[b]["groups"]):
                copied |= {a, b}
                report.append(f"  ages: {ages[a]['name']} and {ages[b]['name']} share one "
                              f"structure (within {COPY_TOLERANCE} points)")
    return set(keys) - copied, report


def ocha_workbook_url(key: str = "cartocensitaire") -> str:
    body = http_json(f"{HDX_API}/package_show?id={OCHA_DATASET}")
    # By the resource's name or its file's: the 2025 table's name is not its
    # file's (ocha-car-projection-population-2025-icasees.xlsx).
    found = [r for r in body["result"].get("resources", [])
             if key in str(r.get("name") or "").lower()
             or key in str(r.get("url") or "").lower().rsplit("/", 1)[-1]]
    if not found:
        raise SystemExit(f"car_census: {OCHA_DATASET} no longer carries a resource named {key}")
    return found[0]["url"]


# ICASEES's projection for 2025 by sub-prefecture of the 2020 division, as OCHA
# publishes it for humanitarian planning (ocha-car-projection-population-2025-
# icasees.xlsx, sheet "Population_HNRP 2026", column Projection_2025_ic).
PROJECTION_KEY = "projection-population-2025"
PROJECTION_YEAR = 2025
PROJECTION_OUT = "car_projection.json"
PROJECTION_SOURCE = ("Institut Centrafricain des Statistiques et des Etudes Economiques et "
                     "Sociales (ICASEES), projection of the population for 2025 by "
                     "sub-prefecture from the 2003 census (RGPH 2003), as OCHA publishes it for "
                     "humanitarian planning (cod-ps-caf)")
# The 2025 table's spellings of 2020 sub-prefectures the mapping table writes
# otherwise (folded): (prefecture, sub-prefecture) -> the mapping table's.
PROJECTION_ALIASES: dict[tuple[str, str], tuple[str, str]] = {}


def projection_rows(rows: list[list[Any]]) -> list[dict[str, Any]]:
    """The 2025 table's sub-prefectures: prefecture, sub-prefecture, code, figure,
    and its men and women where the table gives them."""
    head = [str(c or "").strip() for c in rows[0]]
    at = {fold(h): i for i, h in enumerate(head)}
    value = next((i for k, i in at.items() if k.startswith("projection2025")), None)
    if value is None or "prefecture" not in at or "sousprefecture" not in at:
        raise SystemExit(f"car_census: the 2025 table's columns are not as expected: {head}")
    out = []
    for row in rows[1:]:
        figure = number(row[value]) if len(row) > value else None
        name = str(row[at["sousprefecture"]] or "").strip()
        if figure is None or not name:
            continue
        men = number(row[at["h2025"]]) if "h2025" in at and len(row) > at["h2025"] else None
        women = number(row[at["f2025"]]) if "f2025" in at and len(row) > at["f2025"] else None
        out.append({"pref": str(row[at["prefecture"]] or "").strip(), "sp": name,
                    "code": str(row[at.get("admin2pcode", value)] or "").strip(),
                    "value": float(figure), "men": men, "women": women})
    return out


# Two sub-prefectures whose shares of men in the 2025 table agree this closely
# have one proportion applied to both.
SHARE_TOLERANCE = 5e-4


def one_sex_split(rows: list[dict[str, Any]]) -> tuple[bool, list[str]]:
    """Whether the 2025 table gives every sub-prefecture of each prefecture the
    same share of men (one proportion applied to all), and the report."""
    by_pref: dict[str, list[float]] = defaultdict(list)
    for row in rows:
        if row.get("men") and row.get("women"):
            by_pref[row["pref"]].append(row["men"] / (row["men"] + row["women"]))
    report = []
    copied = bool(by_pref)
    for pref, shares in sorted(by_pref.items()):
        spread = max(shares) - min(shares)
        if len(shares) > 1 and spread > SHARE_TOLERANCE:
            copied = False
            report.append(f"  2025 sexes: {pref}: shares of men {min(shares):.4f}-"
                          f"{max(shares):.4f} over {len(shares)} sub-prefectures")
        else:
            report.append(f"  2025 sexes: {pref}: one share of men, {shares[0]:.4f}, for "
                          f"{len(shares)} sub-prefecture{'s' if len(shares) > 1 else ''}")
    return copied, report


# How alike two folded names must be to be one unit spelled two ways: Bantangafo
# and Batangafo 0.95, Dede-Makouba and Dédé-Mokouba 0.91; Bangui Kaga, one of
# Bangui's four 2025 units, and Bangui, the 2020 sub-prefecture, only 0.75.
SPELLING_LIKENESS = 0.85


def close_names(wanted: list[str], offered: list[str], least: float
                ) -> tuple[dict[str, str], list[str]]:
    """Each wanted name (folded) -> the one offered name (folded) it is closest to,
    where that is at least ``least`` alike and no other wanted name is closer to
    it: the 2025 table's spellings of the mapping table's units ("Bantangafo",
    "Dede-Makouba", "Nana-Gribizi"). Returns the matches and the report."""
    import difflib                                   # noqa: PLC0415
    scores = sorted(((difflib.SequenceMatcher(None, w, o).ratio(), w, o)
                     for w in wanted for o in offered), reverse=True)
    out: dict[str, str] = {}
    taken: set[str] = set()
    report = []
    for score, w, o in scores:
        if score < least or w in out or o in taken:
            continue
        out[w] = o
        taken.add(o)
        report.append(f"{w} -> {o} ({score:.2f})")
    return out, report


def projection(rows: list[dict[str, Any]], c21: dict[str, Any],
               placement: dict[tuple[str, str], list[dict[str, Any]]],
               groups: dict[str, list[tuple[str, str]]], shapes: dict[str, dict[str, Any]],
               admin1: list[dict[str, Any]], totals: dict[str, int], url: str,
               withheld: dict[str, str] | None = None, twinned: dict[str, str] | None = None
               ) -> tuple[list[dict[str, Any]], dict[str, str], dict[str, str], list[str]]:
    """The 2025 figures on the 2003 polygons, where the 2020 sub-prefectures nest in them.

    The 2025 table's prefectures and sub-prefectures are those of the 2020
    division as the mapping table names them, some spelled otherwise: a name
    not found is matched to the closest one of its prefecture not otherwise
    taken, and every such match is logged. Each 2020 sub-prefecture lies where
    the commune crosswalk puts its communes; a polygon takes the 2025 figures of
    the 2020 sub-prefectures it holds whole, and only where it holds nothing of
    any other. ``withheld`` gives the polygons whose 2021 count is not shown
    and why, which their 2025 note repeats in place of that count; ``twinned``
    what a polygon's 2021 count says of a commune in it that the table gives
    another commune's count (Mobaye, Grimari), which the 2025 note repeats
    beside that count. Returns the records (held behind the 2021 counts), the
    sentence for each polygon and prefecture that gets none, and the report."""
    withheld = withheld or {}
    twinned = twinned or {}
    report: list[str] = []
    # Each 2020 sub-prefecture's 2021 count and the polygons its communes are on.
    sp_polys: dict[tuple[str, str], set[str]] = defaultdict(set)
    sp_count: dict[tuple[str, str], int] = defaultdict(int)
    sp_name: dict[tuple[str, str], str] = {}
    for sid, keys in groups.items():
        for key in keys:
            for unit in placement[key]:
                sp21 = (fold(unit["pref"]), fold(unit["sp"]))
                sp_polys[sp21].add(sid)
                sp_count[sp21] += unit["total"]
                sp_name[sp21] = unit["sp"].title() if unit["sp"].isupper() else unit["sp"]
    known = set(c21["sps"])
    prefs = sorted({p for p, _ in known})
    # The 2025 table's prefectures, spelled as the mapping table spells them.
    asked = sorted({fold(r["pref"]) for r in rows} - set(prefs))
    pref_as, lines = close_names(asked, prefs, 0.8)
    report += [f"  2025: prefecture {line}" for line in lines]
    by_pref: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        p = fold(row["pref"])
        by_pref[pref_as.get(p, p)].append(row)
    figures: dict[tuple[str, str], float] = {}
    unmatched: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for pref, found in sorted(by_pref.items()):
        mine = sorted(s for p, s in known if p == pref)
        exact = {fold(r["sp"]) for r in found} & set(mine)
        spelled, lines = close_names(sorted({fold(r["sp"]) for r in found} - exact),
                                     sorted(set(mine) - exact), SPELLING_LIKENESS)
        report += [f"  2025: {pref}: {line}" for line in lines]
        for row in found:
            sp = fold(row["sp"])
            sp = sp if sp in exact else spelled.get(sp, "")
            key = PROJECTION_ALIASES.get((pref, sp), (pref, sp))
            if key in known:
                figures[key] = figures.get(key, 0.0) + row["value"]
            else:
                unmatched[pref].append(row)
        left = sorted(set(mine) - {s for p, s in figures if p == pref})
        if unmatched.get(pref) or left:
            report.append(f"  2025: {pref}: 2025 rows not matched "
                          f"{[r['sp'] for r in unmatched.get(pref, [])]}; mapping-table "
                          f"sub-prefectures without a 2025 row {left}")
    # A 2020 prefecture any of whose 2025 rows names no 2020 sub-prefecture is
    # divided otherwise in 2025: none of its sub-prefectures is placed.
    divided = set(unmatched)
    reasons: dict[str, str] = {}
    records: list[dict[str, Any]] = []
    have: dict[str, float] = {}
    for sid, keys in sorted(groups.items()):
        members = sorted({(fold(u["pref"]), fold(u["sp"])) for key in keys
                          for u in placement[key]})
        name = shapes[sid]["name"]
        whole = all(sp_polys[m] == {sid} for m in members)
        cut = [m for m in members if m[0] in divided]
        missing = [m for m in members if m not in figures and m[0] not in divided]
        if cut:
            found = unmatched[cut[0][0]]
            units = ", ".join(r["sp"] for r in found)
            reasons[sid] = (
                f"ICASEES's projection for {PROJECTION_YEAR} divides {found[0]['pref']} "
                f"prefecture into units of its own ({units}; "
                f"{sum(r['value'] for r in found):,.0f} people) that are not the sub-prefectures "
                f"of the 2020 division whose communes make up this one as the 2003 census drew "
                f"it, so no {PROJECTION_YEAR} figure is shown here.")
            report.append(f"  2025: {name}: {cut[0][0]} divided otherwise in 2025")
            continue
        if not whole:
            spread = [f"{sp_name[m]} ({', '.join(sorted(shapes[s]['name'] for s in sp_polys[m]))})"
                      for m in members if sp_polys[m] != {sid}]
            reasons[sid] = (
                f"ICASEES's projection for {PROJECTION_YEAR} is by sub-prefecture of the 2020 "
                f"division, and the communes of "
                + ", ".join(spread)
                + f" lie partly in this sub-prefecture as the 2003 census drew it and partly in "
                f"others, so no {PROJECTION_YEAR} figure describes it.")
            report.append(f"  2025: {name}: not whole: {'; '.join(spread)}")
            continue
        if missing or not members:
            report.append(f"  2025: {name}: no 2025 row for {missing or 'any of its units'}")
            continue
        value = sum(figures[m] for m in members)
        made = sum(sp_count[m] for m in members)
        parts = [sp_name[m] for m in members]
        have[sid] = value
        same = len(parts) == 1 and fold(parts[0]) == fold(name)
        count = (f" {withheld[sid]}" if sid in withheld else
                 f" The census mapping counted {totals.get(sid, made):,} people here on 1 March "
                 f"2021." + (f" {twinned[sid]}" if twinned.get(sid) else ""))
        note = (f"ICASEES's projection of the population of {name} for {PROJECTION_YEAR}, from "
                f"the 2003 census, as OCHA publishes it for humanitarian planning"
                + ("" if same else
                   f": the 2020 sub-prefectures of {series(parts)} added up, which lie in "
                   f"{name} as the 2003 census drew it")
                + "." + count)
        records.append(record(
            f"CAF-PROJ{PROJECTION_YEAR}-{sid}", name, level="admin2", parent="CAF",
            country="CAF", match_by="shape_id", shape_id=sid,
            population=measure(round(value), year=PROJECTION_YEAR, source=PROJECTION_SOURCE),
            population_note=note,
            sources=[{"field": "population", "name": PROJECTION_SOURCE, "url": url,
                      "year": PROJECTION_YEAR, "license": LICENCE},
                     {"field": "population note", "name": COUNT_SOURCE_USCB,
                      "url": dataset_url(USCB_DATASET), "year": COUNT_YEAR,
                      "license": LICENCE}]))
        rate = (value / made) ** (1 / (PROJECTION_YEAR - COUNT_YEAR)) - 1 if made else 0.0
        report.append(f"  2025: {name}: {value:,.0f} ({' + '.join(parts)}; 2021 {made:,}, "
                      f"{rate:+.1%} a year)" + (" -- 2021 count withheld" if sid in withheld
                                               else ""))
    # A prefecture all of whose polygons have a 2025 figure takes their sum.
    pref_of: dict[str, str] = {sid: keys[0][0] for sid, keys in groups.items()}
    pref_reasons: dict[str, str] = {}
    for poly in admin1:
        key = fold(poly["name"])
        sids = [sid for sid, p in pref_of.items() if p == key]
        if not sids:
            continue
        if all(sid in have for sid in sids):
            value = sum(have[sid] for sid in sids)
            records.append(record(
                f"CAF-PROJ{PROJECTION_YEAR}-{poly['id']}", poly["name"], level="admin1",
                parent="CAF", country="CAF", match_by="shape_id", shape_id=poly["id"],
                population=measure(round(value), year=PROJECTION_YEAR, source=PROJECTION_SOURCE),
                population_note=(
                    f"ICASEES's projection of the population of {poly['name']} for "
                    f"{PROJECTION_YEAR}, from the 2003 census, as OCHA publishes it for "
                    f"humanitarian planning: its sub-prefectures' figures added up."),
                sources=[{"field": "population", "name": PROJECTION_SOURCE, "url": url,
                          "year": PROJECTION_YEAR, "license": LICENCE}]))
            report.append(f"  2025: prefecture {poly['name']}: {value:,.0f}")
        else:
            lacking = sorted(shapes[sid]["name"] for sid in sids if sid not in have)
            pref_reasons[poly["id"]] = (
                f"ICASEES's projection for {PROJECTION_YEAR} gives no figure for "
                f"{series(lacking)} as the 2003 census drew "
                f"{'it' if len(lacking) == 1 else 'them'}, so none is shown for the prefecture.")
            report.append(f"  2025: prefecture {poly['name']}: none ({', '.join(lacking)})")
    return records, reasons, pref_reasons, report


def build(prev: list[dict[str, Any]], c21: dict[str, Any], g03: list[tuple[str, str, str]],
          ages: dict[str, dict[str, Any]], admin1: list[dict[str, Any]],
          admin2: list[dict[str, Any]], ocha_url: str, uscb_url: str,
          proj_rows: list[dict[str, Any]] | None = None, proj_url: str = "",
          held: list[dict[str, Any]] | None = None
          ) -> tuple[list[dict[str, Any]], list[str]]:
    """The 2021 counts and 2016 ages on the 2003 polygons; with ``proj_rows``,
    the 2025 projection's records are added to ``held`` (car_projection.json)
    and the polygons and prefectures it gives no figure say why."""
    report: list[str] = []

    def fail(message: str) -> None:
        """Stop, after printing what was found: the log is how a refusal is read."""
        for line in report:
            log(line)
        raise SystemExit(f"car_census: {message}")

    country = c21["country"]["total"] if c21.get("country") else None
    made = sum(c["total"] for c in prev)
    report.append(f"PrevCommune: {len(prev)} communes of the 2003 frame, {made:,} people; "
                  f"the national line {country:,}" if country else f"PrevCommune: {made:,}")
    if country is not None and made != country:
        fail(f"the 2003-frame communes make {made:,}, the country {country:,}")
    off = fixed_split(c21)
    units = len(c21["prefectures"]) + len(c21["sps"]) + len(c21["communes"])
    report.append(f"mapping table: {off} of {units} units depart from the fixed "
                  f"{FIXED_MALE_SHARE:.0%} share of men")
    if off > 0.05 * units:
        fail("the mapping table now gives units their own sexes; read its sex ratios instead "
             "of withholding them")
    # Duplicate figures, named for anyone reading the run.
    seen: dict[int, list[str]] = defaultdict(list)
    for c in prev:
        if c["total"] >= 1000:
            seen[c["total"]].append(f"{c['name']} ({c['sp']})")
    for total, names in sorted(seen.items()):
        if len(names) > 1:
            report.append(f"  the same count, {total:,}, for {' and '.join(names)}")

    # The assignment by commune names, from the Census Bureau's 2021 frame.
    placement, problems = own_crosswalk(c21, g03)
    report += [f"  crosswalk: {p}" for p in problems]
    if problems:
        fail(f"{len(problems)} communes of 2021 could not be placed on the 2003 frame")
    parents = {u["id"]: u["name"] for u in admin1}
    pref_poly = {fold(u["name"]): u for u in admin1}
    bangui1 = pref_poly.get("bangui")
    bangui2 = next((u for u in admin2 if bangui1 and u["parent"] == bangui1["id"]), None)

    def parent_name(pref: str) -> str:
        return pref_poly[fold(pref)]["name"] if fold(pref) in pref_poly else ""

    departments = {f"{p}/{s}": (spelled(s), parent_name(p)) for (p, s) in placement
                   if p != "bangui"}
    bound, missing = bind(departments, admin2, parents)
    report += [f"  not bound to a polygon: {m}" for m in missing]
    if missing:
        fail("sub-prefectures left without a polygon: " + ", ".join(missing))
    if bangui2:
        bound["bangui/bangui"] = bangui2["id"]
    groups: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for key, sid in bound.items():
        groups[sid].append(tuple(key.split("/")))
    names_by_sid: dict[str, set[str]] = defaultdict(set)
    for pref, sp, commune in g03:
        if f"{pref}/{sp}" in bound:
            names_by_sid[bound[f"{pref}/{sp}"]].add(commune)

    # OCHA's own assignment of the same count to the 2003 communes, as a check.
    by_sp: dict[str, dict[str, Any]] = {}
    for c in prev:
        entry = by_sp.setdefault(c["sp_code"], {"name": c["sp"], "pref": c["pref"], "total": 0})
        entry["total"] += c["total"]
    ocha_departments = {code: (spelled(e["name"]), parent_name(e["pref"]))
                        for code, e in by_sp.items() if fold(e["pref"]) != "bangui"}
    ocha_bound, ocha_missing = bind(ocha_departments, admin2, parents)
    report += [f"  OCHA's table, not bound to a polygon: {m}" for m in ocha_missing]
    ocha_by_poly: dict[str, int] = defaultdict(int)
    for code, sid in ocha_bound.items():
        ocha_by_poly[sid] += by_sp[code]["total"]
    if bangui2:
        ocha_by_poly[bangui2["id"]] = sum(e["total"] for e in by_sp.values()
                                          if fold(e["pref"]) == "bangui")
    shapes2 = {u["id"]: u for u in admin2}
    totals = {sid: sum(u["total"] for key in keys for u in placement[key])
              for sid, keys in groups.items()}
    diff = {sid: ocha_by_poly.get(sid, 0) - totals[sid] for sid in totals
            if ocha_by_poly.get(sid, 0) != totals[sid]}
    unexplained = []
    for sid, d in diff.items():
        if d > 0:
            continue
        # OCHA lacks -d here: a commune of that count, filed by OCHA under a
        # polygon that has a commune of the same name and exactly d too many.
        found = None
        for key in groups[sid]:
            for u in placement[key]:
                if u["total"] != -d:
                    continue
                for other, e in diff.items():
                    if e == -d and fold(u["name"]) in names_by_sid.get(other, set()):
                        found = (u, other)
        if found:
            u, other = found
            report.append(f"  OCHA's table files {u['name'].title()} ({shapes2[sid]['name']}, "
                          f"{u['total']:,}) under {shapes2[other]['name']}'s commune of the same "
                          f"name; the names place it here")
        else:
            unexplained.append(sid)
    unexplained += [sid for sid, d in diff.items() if d > 0 and not any(
        e == -d for e in diff.values())]
    for sid in unexplained:
        report.append(f"  disagreement: {shapes2[sid]['name']}: OCHA's 2003 communes "
                      f"{ocha_by_poly.get(sid, 0):,}, the 2021 communes by name {totals[sid]:,}")
    if unexplained:
        fail(f"OCHA's assignment and the commune names disagree on {len(unexplained)} polygons "
             f"for no reason found")

    cite_count = [{"field": "population", "name": COUNT_SOURCE,
                   "url": f"https://data.humdata.org/dataset/{OCHA_DATASET}", "year": COUNT_YEAR,
                   "license": LICENCE},
                  {"field": "population", "name": COUNT_SOURCE_USCB,
                   "url": dataset_url(USCB_DATASET), "year": COUNT_YEAR, "license": LICENCE}]
    # The sub-prefectures' notes on their sexes, ages and seats rest on the
    # mapping table (its fixed split, no ages, no seats), on the 2016 yearbook
    # (whose ages and sexes stop at the prefecture) and, for 2025, on ICASEES's
    # projection.
    cite_notes = [{"field": "sex_ratio note/median_age note/capital note",
                   "name": COUNT_SOURCE_USCB, "url": dataset_url(USCB_DATASET),
                   "year": COUNT_YEAR, "license": LICENCE},
                  {"field": "sex_ratio note/median_age note", "name": AGE_SOURCE,
                   "url": dataset_url(USCB_DATASET), "year": AGE_YEAR, "license": LICENCE}]
    records: list[dict[str, Any]] = []
    pref_totals: dict[str, int] = defaultdict(int)
    reasons: dict[str, str] = {}
    pref_reasons: dict[str, str] = {}
    sexes_copied = False
    twins = twin_counts(groups, placement)
    for count, found in sorted(twins.items()):
        report.append(f"  twins: {count:,} for " + " and ".join(
            f"{u['name'].title()} ({shapes2[sid]['name']})" for sid, u in found))
    pref_of = {sid: keys[0][0] for sid, keys in groups.items()}
    if proj_rows:
        # A polygon whose 2021 count is withheld (SUSPECT) takes the 2025 figure,
        # which rests on the 2003 census, and its note says why the count is not
        # shown; one whose count holds a commune the table gives another's count
        # says so beside the count, as the count's own note does.
        withheld = {sid: SUSPECT[fold(shapes2[sid]["name"])] for sid in groups
                    if fold(shapes2[sid]["name"]) in SUSPECT}
        twinned = {sid: twin_sentence(sid, twins, shapes2, pref_poly, pref_of)
                   for sid in groups}
        made, reasons, pref_reasons, lines = projection(
            proj_rows, c21, placement, groups, shapes2, admin1, totals, proj_url, withheld,
            twinned)
        report += lines
        sexes_copied, lines = one_sex_split(proj_rows)
        report += lines
        if held is not None:
            held += made
    for sid, keys in sorted(groups.items()):
        shape = shapes2[sid]
        units = [u for key in keys for u in placement[key]]
        total = totals[sid]
        pref_totals[keys[0][0]] += total
        sps21 = sorted({u["sp"].title() for u in units})
        whole = (len(sps21) == 1 and fold(sps21[0]) == keys[0][1]
                 and len(units) == sum(1 for u in c21["communes"]
                                       if fold(u["sp"]) == keys[0][1]
                                       and fold(u["pref"]) == fold(units[0]["pref"])))
        suspect = SUSPECT.get(fold(shape["name"]))
        sources = cite_count
        seat: dict[str, Any] = {"capital": gap(NOT_AVAILABLE, SEAT_GAP)}
        if bangui2 and sid == bangui2["id"]:
            # Bangui's one division is the city, which is also the prefecture.
            seat = {"capital": "Bangui",
                    "capital_note": ("This polygon is the city of Bangui, which is also the "
                                     "prefecture drawn one level up; its seat is the city "
                                     "itself.")}
        if suspect:
            pop = gap(NOT_AVAILABLE, suspect)
            note = suspect
            sources = [{**src, "field": "population note"} for src in cite_count]
        elif not total:
            pop = gap(NOT_AVAILABLE, f"The 2021 census mapping counted no one in the communes "
                                     f"of {shape['name']}.")
            note = None
        else:
            pop = measure(total, year=COUNT_YEAR, source=COUNT_SOURCE_USCB)
            if whole:
                note = (f"The people counted in {shape['name']} at the mapping stage of the "
                        f"fourth general census (RGPH-4), 1 March 2021.")
            else:
                listed = ", ".join(f"{u['name'].title()} {u['total']:,}" for u in units)
                note = (f"The people counted at the mapping stage of the fourth general census "
                        f"(RGPH-4), 1 March 2021, on {shape['name']}'s ground as the 2003 census "
                        f"drew it, its communes added up ({listed})"
                        + (f"; since 2020 this ground is the sub-prefectures of "
                           f"{series(sps21)}." if len(sps21) > 1 else "."))
            twin = twin_sentence(sid, twins, shapes2, pref_poly, pref_of)
            if twin:
                note = f"{note} {twin}"
        if sid in reasons and note:
            note = f"{note} {reasons[sid]}"
        noted = list(cite_notes)
        said = (["population note"] if sid in reasons else []) + (
            ["sex_ratio note"] if sexes_copied else [])
        if said and proj_rows:
            noted.append({"field": "/".join(said), "name": PROJECTION_SOURCE,
                          "url": proj_url or f"https://data.humdata.org/dataset/{OCHA_DATASET}",
                          "year": PROJECTION_YEAR, "license": LICENCE})
        if bangui2 and sid == bangui2["id"]:
            noted = [{**src, "field": src["field"].replace("capital note", "capital")}
                     for src in noted]
        records.append(record(
            f"CAF-RGPH4-{sid}", shape["name"], level="admin2", parent="CAF", country="CAF",
            match_by="shape_id", shape_id=sid, population=pop,
            **({"population_note": note} if note else {}),
            sex_ratio=gap(NOT_AVAILABLE, SEX_GAP + (SEX_GAP_2025 if sexes_copied else "")),
            median_age=gap(NOT_AVAILABLE, AGE_GAP), sources=sources + noted, **seat))
        report.append(f"  {shape['name']}: {total:,} ({len(units)} communes of 2021"
                      + ("" if whole else f"; 2021 sub-prefectures {', '.join(sps21)}") + ")"
                      + (" -- withheld" if suspect else ""))


    own, why = own_structures(ages)
    report += why
    cite_age = [{"field": "median_age/sex_ratio", "name": AGE_SOURCE,
                 "url": dataset_url(USCB_DATASET), "year": AGE_YEAR, "license": LICENCE}]
    for key, poly in sorted(pref_poly.items()):
        total = pref_totals.get(key)
        if not total:
            report.append(f"  prefecture {poly['name']}: no 2021 count")
            continue
        twin = twin_sentence_prefecture(key, twins, shapes2, pref_poly, pref_of)
        fields: dict[str, Any] = {
            "population": measure(total, year=COUNT_YEAR, source=COUNT_SOURCE_USCB),
            "population_note": (f"The people counted in {poly['name']} at the mapping stage of "
                                f"the fourth general census (RGPH-4), 1 March 2021, on its ground "
                                f"as the 2003 census drew it: the sum of its sub-prefectures' "
                                f"counts in the table." + (f" {twin}" if twin else "")
                                + (f" {pref_reasons[poly['id']]}"
                                   if poly["id"] in pref_reasons else "")),
            "sources": list(cite_count) + (
                [{"field": "population note", "name": PROJECTION_SOURCE,
                  "url": proj_url or f"https://data.humdata.org/dataset/{OCHA_DATASET}",
                  "year": PROJECTION_YEAR, "license": LICENCE}]
                if poly["id"] in pref_reasons else [])}
        entry = ages.get(key)
        if entry:
            fields["sex_ratio"] = measure(round(1000 * entry["men"] / entry["women"]),
                                          unit="males_per_1000_females", year=AGE_YEAR,
                                          source=AGE_SOURCE)
            fields["sex_ratio_note"] = ("Men per 1,000 women in ICASEES's administrative "
                                        "population at 31 December 2016 (Annuaire statistique "
                                        "de la République Centrafricaine 2016), an estimate "
                                        "rather than a count.")
            median = grouped_median([tuple(g) for g in entry["groups"]]) if key in own else None
            if median is not None:
                fields["median_age"] = {"value": median, "unit": "years", "year": AGE_YEAR,
                                        "source": AGE_SOURCE}
                fields["median_age_note"] = (
                    "Interpolated within the five-year age group that holds the middle person, "
                    "from ICASEES's administrative population by age at 31 December 2016 "
                    "(Annuaire statistique de la République Centrafricaine 2016), an estimate "
                    "rather than a count.")
            else:
                fields["median_age"] = gap(NOT_AVAILABLE, (
                    "ICASEES's 2016 estimates by age (Annuaire statistique de la République "
                    "Centrafricaine 2016) give this prefecture the same age proportions as "
                    "another prefecture or the country, so no median age of its own can be "
                    "taken from them."))
            fields["sources"] += cite_age
        records.append(record(f"CAF-RGPH4-{poly['id']}", poly["name"], level="admin1",
                              parent="CAF", country="CAF", match_by="shape_id",
                              shape_id=poly["id"], **fields))
        report.append(f"  prefecture {poly['name']}: {total:,}"
                      + (f"; 2016: {entry['men'] / entry['women'] * 1000:.0f} men per 1,000 "
                         f"women, median {fields['median_age'].get('value')}" if entry else ""))
    return records, report


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dump", action="store_true", help="print the OCHA sheet's rows, write nothing")
    ap.add_argument("--run", help="a label that tells one dispatch from another; not used")
    args = ap.parse_args()
    import openpyxl                                  # noqa: PLC0415
    ocha_url = ocha_workbook_url()
    blob = http_get(ocha_url, binary=True, cache=False, timeout=300)
    log(f"car_census: {ocha_url}: {len(blob):,} bytes, sha256 {hashlib.sha256(blob).hexdigest()}")
    book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    prev = prev_communes(plain(sheet_rows(book, "PrevCommune")))
    uscb_url = workbook_url(USCB_DATASET)
    ublob = http_get(uscb_url, binary=True, cache=False, timeout=300)
    log(f"car_census: {uscb_url}: {len(ublob):,} bytes, sha256 {hashlib.sha256(ublob).hexdigest()}")
    ubook = openpyxl.load_workbook(io.BytesIO(ublob), read_only=True, data_only=True)
    c21 = counts_2021(table(sheet_rows(ubook, "Population")))
    g03 = communes_2003(table(sheet_rows(ubook, "Ethnicity")))
    ages = age_table(table(sheet_rows(ubook, "Age-Sex")))
    if args.dump:
        for c in prev:
            log(f"  {c['pref_code']} {c['pref']} | {c['sp_code']} {c['sp']} | {c['code']} "
                f"{c['name']} | {c['total']}")
        for k, e in ages.items():
            log(f"  {e['name']}: {e['groups']}")
        return 0
    admin1 = json.loads((SITE / "admin1" / "CAF.units.json").read_text())
    admin2 = json.loads((SITE / "admin2" / "CAF.units.json").read_text())
    proj_url = ocha_workbook_url(PROJECTION_KEY)
    # What the dataset says of the projection (its base, its method), for the notes.
    meta = http_json(f"{HDX_API}/package_show?id={OCHA_DATASET}")["result"]
    for res in meta.get("resources", []):
        if res.get("url") == proj_url:
            log(f"  2025 resource: {res.get('name')}: "
                + " ".join(str(res.get('description') or '').split())[:600])
    log("  dataset notes: " + " ".join(str(meta.get("notes") or "").split())[:900])
    log("  dataset caveats: " + " ".join(str(meta.get("caveats") or "").split())[:900])
    pblob = http_get(proj_url, binary=True, cache=False, timeout=300)
    log(f"car_census: {proj_url.rsplit('/', 1)[-1]}: {len(pblob):,} bytes, sha256 "
        f"{hashlib.sha256(pblob).hexdigest()}")
    pbook = openpyxl.load_workbook(io.BytesIO(pblob), read_only=True, data_only=True)
    sheet = next(ws for ws in pbook.worksheets
                 if any("projection_2025" in str(c or "").lower()
                        for c in next(ws.iter_rows(values_only=True), ())))
    proj_rows = projection_rows([list(r) for r in sheet.iter_rows(values_only=True)])
    log(f"  {sheet.title}: {len(proj_rows)} sub-prefectures, "
        f"{sum(r['value'] for r in proj_rows):,.0f} people for {PROJECTION_YEAR}")
    held: list[dict[str, Any]] = []
    records, report = build(prev, c21, g03, ages, admin1, admin2, ocha_url, uscb_url,
                            proj_rows, f"https://data.humdata.org/dataset/{OCHA_DATASET}", held)
    for line in report:
        log(line)
    log(f"  {sum(r['level'] == 'admin1' for r in records)} prefectures and "
        f"{sum(r['level'] == 'admin2' for r in records)} sub-prefecture polygons written; "
        f"{PROJECTION_YEAR}: {sum(r['level'] == 'admin1' for r in held)} prefectures and "
        f"{sum(r['level'] == 'admin2' for r in held)} sub-prefectures")
    write_json(PROCESSED / OUT, records)
    write_json(PROCESSED / PROJECTION_OUT, held)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
