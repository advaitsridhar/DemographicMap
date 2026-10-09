#!/usr/bin/env python3
"""Iran: the 2016 census by province and shahrestan -- age, sex and citizenship.

**Where the tables are.** The Statistical Centre of Iran published the 2016
(1395) census's detailed tables as Excel workbooks under
``amar.org.ir/Portals/0/census/1395/results/tables/``. Its host ends the TLS
handshake before a standard client reads a byte (the reason this map's
Iranian provinces carried no age or composition), so the workbooks are read
from the Internet Archive's captures of those same files, raw (the ``id_``
modifier), as ``probe_xls --wayback`` reads them. Four families are used,
all from the jamiat (population) tafsili (detailed) series:

* ``ostani/1-jamiat_ostani.xls`` -- Table 1, population by single year of age
  and sex, one sheet per province (31), urban and rural;
* ``shahrestan/jamiatNN.xls`` -- the same Table 1 for every shahrestan, one
  workbook per province (NN, the Centre's province code 00-30) and one sheet
  per shahrestan, named with its four-digit code;
* ``3-jamiat-k.xls`` and ``shahrestan/3/3-jamiat-NN-k.xls`` -- Table 3,
  population by sex and country of citizenship (Iran, Afghanistan, Iraq,
  Pakistan, Turkey, other countries, not stated), by province and by
  shahrestan.

**What is written.** For every province and every shahrestan the map draws:
the census population; the median age, interpolated within the single year
of age that holds the middle person; men per hundred women; and the
citizenship composition as the ethnicity field under ``ethnicity_basis:
"citizenship"``, by the owner's decision of 19 September 2026 (the model is
Japan's and Korea's nationality files). Iran's census asks no ethnicity; it
counts citizenship, and that count is what this is -- a person's country of
citizenship, not their ethnic group -- so each is labelled as a nationality
("Iranian national", "Afghan national"; see ``SHOWN``) and never as a
people's name, which the group tree would file under an ethnic family.
Religion was asked in 2016 too: the
Centre's Statistical Yearbook 1395 (English, chapter 3, table 3.18) counts
it by province, and that is written on every first-level polygon that is a
whole province (its Christian and Zoroastrian headings are swapped in the
English edition, and are restored -- see RELIGION_COLUMNS). No table of
religion below the province is among the archived census files.

**Binding.** The map's shahrestans are geoBoundaries' and their English
labels follow no one transliteration ("Kahlkhal", "Savdkuh", "Sharekurd").
The Centre's sheets are named in Persian. The bridge is OCHA's Common
Operational Dataset (cod-ab-irn), which names the same 2016 shahrestans in
Persian and in English: a sheet is matched to OCHA's row by its Persian name
inside its province, and OCHA's English name to the drawn polygon inside the
drawn province by a folded comparison, or by ``DRAWN`` where the boundary
file spells it differently. Every pair is one-to-one or the run stops. One
sheet is spelt otherwise by OCHA (``CENSUS_FA``: Khuzestan's Haftgel, "هفتگل"
in the census and "هفتکل" in cod-ab-irn).

**What the map draws differently from 2016, measured on the site's own
tiles.**

* Three 2016 shahrestans are each drawn as two polygons: Tehran ("City of
  Tehran", the city, and "Tehran", the rest of the county in three pieces
  around it), Isfahan ("Isfahan", the city, and "Isfahan County", the rest)
  and Mehdishahr ("Shahmirzad", its Shahmirzad district, made a county of
  its own after the census, and "Mehdishahr", its central district). The
  shahrestan's figure describes neither half, so each half takes the
  census's own count of what it is (``SPLIT``), from the province's
  settlement table (``CN95_HouseholdPopulationVillage_NN.xlsx``: every
  county, district, city and village with its persons, men and women).
  Tehran is its shahrestan's only city, so Table 1's urban block is the city
  and its rural and unsettled blocks the rest, and both halves take ages
  too; Isfahan's shahrestan holds thirteen other cities and Mehdishahr's
  ages are not tabulated by district, so those four halves take persons and
  sex only. Table 3 counts citizenship by shahrestan alone, so no half takes
  a citizenship share. Each half says why it lacks what it lacks.
* Abu Musa, an island county, is not drawn.
* Two polygons carry another county's label: the one labelled "Shahariar"
  in Alborz is Fardis (Fardis town, 35.72 N 50.98 E, lies in it; Shahriar is
  in Tehran province and drawn there), and the one labelled "Paveh" in
  Tehran province is Pardis (Pardis, 35.74 N 51.77 E; Paveh is in
  Kermanshah and drawn there). They are bound by id under the census's name.
* Bandar-e Gaz county (Golestan) is drawn under a first-level polygon of its
  own labelled "Mazandaran" (Bandar-e Gaz town, 36.77 N 53.95 E, lies in it).
  That small polygon takes Bandar-e Gaz's figures, and the drawn Golestan,
  which therefore lacks it, takes Golestan's less Bandar-e Gaz's -- exact,
  because these are counts by single year of age.

**Checks, each a stop.** Every sheet's single years add up to its total, for
both sexes; the five-year subtotals the sheet prints equal the single years
they cover; men and women add up to both sexes; every province's
shahrestans add up to the province; the provinces to the country's
79,926,270; every citizenship row's categories to its total, and that total
to the age table's.

Usage:
    python -m scripts.fetch_census.iran_census
"""

from __future__ import annotations

import argparse
import io
import re
import sys
import time
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from ._shared import PROCESSED, http_get, log, measure, record, shares, write_json
from .redatam import median_age

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from common import NOT_AVAILABLE, as_drawn, gap, read_json, shard_name  # noqa: E402

ISO3 = "IRN"
YEAR = 2016
OUT = "iran_census.json"
ROOT = Path(__file__).resolve().parent.parent.parent
SITE = ROOT / "site" / "data"
BASE = "https://www.amar.org.ir/Portals/0/census/1395/results/tables/jamiat/tafsili"
WAYBACK = "https://web.archive.org/web/2019id_/{url}"
PROVINCE_AGES = f"{BASE}/ostani/1-jamiat_ostani.xls"
PROVINCE_CITIZENSHIP = f"{BASE}/3-jamiat-k.xls"
COUNTY_AGES = BASE + "/shahrestan/jamiat{nn:02d}.xls"
COUNTY_CITIZENSHIP = BASE + "/shahrestan/3/3-jamiat-{nn:02d}-k.xls"
CODAB = ("https://data.humdata.org/dataset/07f4ec78-42c7-4606-ae62-4f1bff918c45/resource/"
         "6536032d-54c1-43cd-b4a0-0ecb28dc18b6/download/irn_adminboundaries_tabulardata.xlsx")
SOURCE = ("Statistical Centre of Iran, National Population and Housing Census 2016 (1395), "
          "detailed tables (Table 1: population by age and sex; Table 3: population by sex "
          "and citizenship), as the Internet Archive captured them from amar.org.ir")
PAGE = "https://www.amar.org.ir/"
LICENCE = "Official statistics of the Statistical Centre of Iran"
NATIONAL = 79_926_270

# The Centre's province codes, with its Persian names (checked against every
# sheet) and the map's first-level labels.
PROVINCES: dict[int, tuple[str, str]] = {
    0: ("مرکزی", "Markazi"), 1: ("گیلان", "Gilan"), 2: ("مازندران", "Mazandaran"),
    3: ("آذربایجان شرقی", "East Azerbaijan"), 4: ("آذربایجان غربی", "West Azerbaijan"),
    5: ("کرمانشاه", "Kermanshah"), 6: ("خوزستان", "Khuzestan"), 7: ("فارس", "Fars"),
    8: ("کرمان", "Kerman"), 9: ("خراسان رضوی", "Razavi Khorasan"), 10: ("اصفهان", "Isfahan"),
    11: ("سیستان و بلوچستان", "Sistan and Baluchestan"), 12: ("کردستان", "Kurdistan"),
    13: ("همدان", "Hamadan"), 14: ("چهارمحال و بختیاری", "Chaharmahal and Bakhtiari"),
    15: ("لرستان", "Lorestan"), 16: ("ایلام", "Ilam"),
    17: ("کهگیلویه و بویراحمد", "Kohgiluyeh and Boyer-Ahmad"), 18: ("بوشهر", "Bushehr"),
    19: ("زنجان", "Zanjan"), 20: ("سمنان", "Semnan"), 21: ("یزد", "Yazd"),
    22: ("هرمزگان", "Hormozgan"), 23: ("تهران", "Tehran"), 24: ("اردبیل", "Ardabil"),
    25: ("قم", "Qom"), 26: ("قزوین", "Qazvin"), 27: ("گلستان", "Golestan"),
    28: ("خراسان شمالی", "North Khorasan"), 29: ("خراسان جنوبی", "South Khorasan"),
    30: ("البرز", "Alborz"),
}
# OCHA's first-level name for a shahrestan where it is not the province's:
# cod-ab-irn files Some'e Sara under a first level of its own, "Somee-Sara".
CODAB_PROVINCE = {"Somee-Sara": "Gilan"}

# OCHA's English name -> the boundary file's label, inside the same province.
DRAWN: dict[str, str] = {
    "Ardabil": "Ardebil", "Bilasavar": "Bilehsavar", "Germi": "Garmi", "Khalkhal": "Kahlkhal",
    "Kowsar": "Kosar", "Meshginshahr": "Meshkinshahr", "Dayyer": "Deyr", "Ganaveh": "Genaveh",
    "Kiaar": "Kiar", "Shahrekord": "Sharekurd", "Kalibar": "Kaleybar",
    "Kherameh": "Kharameh", "Khorrambid": "Khorambid", "Qir-o-Karzin": "Qir and Karzin",
    "Zarrindasht": "Zarindasht", "Astane-ye-Ashrafiyeh": "Astan-e-Ashrafieh",
    "Langrud": "Langerud", "Some'e-Sara": "Sowme'eh Sara", "Kolaleh": "Kalaleh",
    "Bandar-e-Torkaman": "Turkman", "Hamadan": "Hamedan", "Kabudarahang": "Kabutarahang",
    "Bandar-Abbas": "Bandar-e-Abbas", "Bandar-Lengeh": "Bandar-e-Lengeh",
    "Abdanan": "Abdanan County", "Badre": "Badreh", "Shirvan-o-Chardavol": "Chardavol",
    "Darreh Shahr": "Darreh Shahr County", "Dehloran": "Dehloran County",
    "Eyvan": "Eyvan County", "Mehran": "Mehran County", "Aran-o-Bidgol": "Aran and Bidgol",
    "Booeino Miyandasht": "Buin and Miandasht", "Borkhar": "Barkhar", "Dehaghan": "Dehaqan",
    "Falavarjan": "Falaverjan", "Faridan": "Fereydan", "Khorobiyabanak": "Khur and Biabanak",
    "Shahinshahro Meymeh": "Shahin Shahr and Meymeh", "Shahreza": "Sahreza",
    "Tiran-o-Korun": "Tiran and Karvan", "Ghaleye-Ganj": "Ghalehganj",
    "Kuhbonan": "Kuhbanan", "Normashir": "Narmashir", "Rabar": "Rabor",
    "Roudbar-e-Jonub": "Rudbar-e-Jonub", "Dalaho": "Dalahu",
    "Islamabad-e-Gharb": "Eslamabad-e Gharb", "Solas-e-Babajani": "Salas-e Babajani",
    "Guotvand": "Gotvand", "Hoveizeh": "Hoveyzeh", "Karoun": "Karun",
    "Khorramshahr": "Khoramshahr", "Mahshahr": "Bandar-e-Mahshahr", "Omidiyeh": "Omidieh",
    "Kohgeluyeh": "Kohkiluyeh", "Yasooj": "Buyerahmad", "Divandarreh": "Divandareh",
    "Qorveh": "Ghorveh", "Sannandaj": "Sanandaj", "Doureh": "Doreh",
    "Khorramabad": "Khoramabad", "Poldokhtar": "Pol-e-Dokhtar", "Romeshkan": "Rumeshkhan",
    "Komeijan": "Komijan", "Feredunkenar": "Fereydunkenar", "Qaemshahr": "Ghaemshahr",
    "Northern Savadkooh": "Savadkuh-e-Shomali", "Savadkuh": "Savdkuh",
    "Esfarayen": "Esfarayen County", "Faroj": "Faruj", "Germeh": "Garmeh",
    "Maneh-o-Samalqan": "Maneh and Samalqan", "Razo Jalgelan": "Raz and Jargalan",
    "Shirvan": "Shirwan", "Nishapur": "Neyshabur", "Rashtkhar": "Roshtkhar",
    "Torghabe-o-Shandiz": "Binalud", "Zave": "Zaveh", "Meyami": "Miami",
    "Dalgan": "Dalagan", "Fanouj": "Fanuj", "Ghasre Ghand": "Qasr-e-Ghand",
    "Hamoun": "Hamun", "Nimrouz": "Nimruz", "Sibo Soran": "Sib and Suran", "Zehak": "Zahak",
    "Qaen": "Ghaenat", "Zirkouh": "Zirkuh", "Qarchak": "Gharchak", "Shahr-e Qods": "Qods",
    "Oshnaviyeh": "Ashnavieh", "Chaldoran": "Chaldoran County", "Naqadeh": "Naghadeh",
    "Shahindej": "Shahindezh", "Eejrud": "Ijrud", "Khorramdarreh": "Khoramdareh",
    # The boundary file's label is another county's; the town is in the polygon.
    "Fardis": "Shahariar", "Pardis": "Paveh",
}
# Labels the boundary file puts on another county's polygon: the record is
# named by the census and the label kept as an alias.
MISLABELLED = {"Fardis": "Shahariar", "Pardis": "Paveh"}
# 2016 shahrestans the map draws as two polygons, and what each polygon is in
# the census's own units: the province, the county's two-digit code in the
# settlement table and its sheet code in Table 1; the part one polygon is --
# a city (the settlement table's four-digit code) or a district (two digits)
# -- and the polygon that is it; and the polygon that is the rest of the
# county. Measured on the boundary file at full resolution: "City of Tehran"
# (627 km2) holds Tajrish and Rey, regions 1 and 20 of the city's 22, and
# "Tehran" (752 km2) is three pieces around it -- north-west (Soleqan,
# Vardij and Keshar, the Kan district's villages), east and south; "Isfahan"
# (187 km2) holds the city's centre and Rehnan, while Baharestan, Qahjavarestan
# and Khatunabad lie in "Isfahan County"; Shahmirzad town and Fulad Mahalleh
# lie in "Shahmirzad", and Mehdishahr and Darjazin in "Mehdishahr".
SPLIT: dict[str, dict[str, Any]] = {
    "Tehran": {
        "province": 23, "county": "01", "code": "2301", "part": ("city", "1576"),
        "part_label": "City of Tehran", "rest_label": "Tehran",
        "part_what": "the city of Tehran, its 22 municipal regions",
        "rest_what": ("Tehran shahrestan less the city of Tehran: its rural population (the "
                      "Kan and Aftab districts and the central district's Siahrud rural "
                      "district) and the people the census counts as unsettled"),
    },
    "Isfahan": {
        "province": 10, "county": "02", "code": "1002", "part": ("city", "1406"),
        "part_label": "Isfahan", "rest_label": "Isfahan County",
        "part_what": "the city of Isfahan, its 15 municipal regions",
        "rest_what": ("Isfahan shahrestan less the city of Isfahan: its {others} other "
                      "cities, Baharestan the largest, and its rural districts"),
    },
    "Mehdishahr": {
        "province": 20, "county": "05", "code": "2005", "part": ("district", "01"),
        "part_label": "Shahmirzad", "rest_label": "Mehdishahr",
        "part_what": ("the Shahmirzad district of Mehdishahr shahrestan (the town of "
                      "Shahmirzad and the Poshtkuh and Chashm rural districts), made a "
                      "county of its own after the census"),
        "rest_what": ("Mehdishahr shahrestan less its Shahmirzad district: the central "
                      "district, with the towns of Mehdishahr and Darjazin"),
    },
}
# The settlement table of each province ("abadi"): every county (record code
# 2), district (3), rural district (4), city (5) and village (6, 8), with its
# households, persons, men and women. A zoned city's regions are rows of
# their own whose ShahrTop column names the city.
SETTLEMENTS = ("https://www.amar.org.ir/Portals/0/census/1395/results/abadi/"
               "CN95_HouseholdPopulationVillage_{nn:02d}.xlsx")
SETTLEMENT_SOURCE = ("Statistical Centre of Iran, National Population and Housing Census 2016 "
                     "(1395), population and households by settlement (county, district, city "
                     "and village), as the Internet Archive captured it from amar.org.ir")
# The settlement table's header row, where the reader expects each column.
SETTLEMENT_COLUMNS = {2: "کد", 4: "کد", 6: "کد", 11: "کدرکورد", 12: "ShahrTop",
                      13: "SwDiv", 15: "جمعیت", 16: "مرد", 17: "زن"}
# Table 1's four blocks of columns, as their headings read once folded.
BLOCKS = ("جمع", "نقاطشهری", "نقاطروستایی", "غیرساکن")
MEDIAN_SPLIT_GAP = {
    "Isfahan": ("Table 1 of the 2016 census counts ages by shahrestan, with its urban and "
                "rural population apart, and Isfahan shahrestan's urban population is the "
                "city of Isfahan and {others} other cities; no age table describes either "
                "polygon of the split shahrestan, so no median age is written."),
    "Mehdishahr": ("Table 1 of the 2016 census counts ages by shahrestan, and no table of "
                   "ages by district is among the Centre's archived census tables, so no "
                   "median age is written for either district's polygon."),
}
CITIZENSHIP_SPLIT_GAP = ("Table 3 of the 2016 census counts citizenship by shahrestan only, "
                         "with no city or district row, and the map draws {county} shahrestan "
                         "as two polygons; no share is written for either.")
# Drawn nowhere: an island county.
NOT_DRAWN = {"Abumusa"}
# A shahrestan drawn under another first-level polygon than its province's:
# OCHA's English name -> the drawn province label, and which drawn polygon of
# that label (the one holding exactly this child).
ELSEWHERE = {"Bandar-e-Gaz": "Mazandaran"}
# A sheet whose Persian name OCHA spells otherwise, inside the same province:
# the Centre's sheet name -> OCHA's. Khuzestan's 0622 is "هفتگل" (Haftgel) in
# the census and "شهرستان هفتکل" (Haftkel, IR015013, English "Haftgol") in
# cod-ab-irn -- the two usual spellings of the one county, and the only sheet
# and the only OCHA row of Khuzestan's 27 left over when the others are paired.
CENSUS_FA = {"هفتگل": "هفتکل"}

LABELS = {"ایران": "Iranian", "افغانستان": "Afghan", "عراق": "Iraqi", "پاکستان": "Pakistani",
          "ترکیه": "Turkish", "سایر کشورها": "Other nationalities",
          "اظهار نشده": "Not stated"}
# How each country of citizenship is written on the record. Table 3 counts
# passports, and a bare adjective is what the group tree reads as a people:
# "Iranian" files under the Iranian peoples, which would make Iran's
# Azerbaijani, Kurdish, Turkmen and Arab citizens Persians, Kurds and Baloch
# on the map, and "Afghan", "Iraqi" and "Turkish" would do the same to
# Afghanistan's, Iraq's and Turkey's citizens. Each is written as the
# nationality it is, which the tree files with the other national identities.
SHOWN = {"Iranian": "Iranian national", "Afghan": "Afghan national",
         "Iraqi": "Iraqi national", "Pakistani": "Pakistani national",
         "Turkish": "Turkish national", "Other nationalities": "Other nationalities",
         "Not stated": "Not stated"}

PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")

# Religion by province: the Statistical Centre's English Statistical Yearbook
# 1395 (2016-17), chapter 3, table 3.18 "Population by religion and ostan,
# 1395 census". No table of religion below the province is among the
# archived census tables (jamiat 1, 2, 3 and 9; faaliat, khanevar, maskan,
# mohajerat, tahsilat, zanashui), so the shahrestans carry none.
YEARBOOK = "https://www.amar.org.ir/Portals/1/yearbook/1395/03.pdf"
YEARBOOK_SOURCE = ("Statistical Centre of Iran, Iran Statistical Yearbook 1395 (2016-17), "
                   "chapter 3 (Population), table 3.18: population by religion and ostan, "
                   "1395 census")
# The yearbook's English province names -> the map's labels.
YEARBOOK_OSTAN = {
    "East Azarbayejan": "East Azerbaijan", "West Azarbayejan": "West Azerbaijan",
    "Ardebil": "Ardabil", "Esfahan": "Isfahan", "Alborz": "Alborz", "Ilam": "Ilam",
    "Bushehr": "Bushehr", "Tehran": "Tehran",
    "Chaharmahal & Bakhtiyari": "Chaharmahal and Bakhtiari",
    "South Khorasan": "South Khorasan", "Khorasan-e-Razavi": "Razavi Khorasan",
    "North Khorasan": "North Khorasan", "Khuzestan": "Khuzestan", "Zanjan": "Zanjan",
    "Semnan": "Semnan", "Sistan & Baluchestan": "Sistan and Baluchestan", "Fars": "Fars",
    "Qazvin": "Qazvin", "Qom": "Qom", "Kordestan": "Kurdistan", "Kerman": "Kerman",
    "Kermanshah": "Kermanshah", "Kohgiluyeh & Boyerahmad": "Kohgiluyeh and Boyer-Ahmad",
    "Golestan": "Golestan", "Gilan": "Gilan", "Lorestan": "Lorestan",
    "Mazandaran": "Mazandaran", "Markazi": "Markazi", "Hormozgan": "Hormozgan",
    "Hamedan": "Hamadan", "Yazd": "Yazd",
}
# Table 3.18's columns after the total, as printed, and the map's labels.
# The English yearbook swaps two headings: its table 3.17 prints the 2011
# census's 117,704 under "Zoroastrian" and 25,271 under "Christian", where
# the Centre's own Persian selected results of that census (census-90-
# results.pdf, table 3) count 117,704 Christians (مسیحی) and 25,271
# Zoroastrians (زرتشتی). Table 3.18 is set the same way: its "Christian"
# column has Yazd's 3,600 and Tehran's 8,579, its "Zoroastrian" column
# Tehran's 43,987, Isfahan's 8,628 and West Azerbaijan's 7,647 -- the
# Armenian and Assyrian churches. read_religion checks table 3.17 before
# trusting this.
RELIGION_COLUMNS = ("Islam", "Zoroastrianism", "Christianity", "Judaism", "Other religions",
                    "Not stated")
CENSUS_2011_PERSIAN = {"Christianity": 117_704, "Zoroastrianism": 25_271}


def fa(text: Any) -> str:
    """Persian text with the Arabic letter forms, digits and spacing folded."""
    s = str(text or "").translate(PERSIAN_DIGITS)
    s = (s.replace("ي", "ی").replace("ى", "ی").replace("ئ", "ی").replace("ك", "ک")
         .replace("ة", "ه").replace("ؤ", "و")
         .replace("‌", "").replace("‏", "").replace("‎", "")
         .replace("أ", "ا").replace("إ", "ا").replace("آ", "ا"))
    return re.sub(r"\s+", "", s.replace("شهرستان", ""))


def fold(name: str) -> str:
    text = unicodedata.normalize("NFKD", str(name or "").lower())
    return "".join(c for c in text if c.isalnum() and not unicodedata.combining(c))


def number(cell: Any) -> int:
    text = str(cell if cell is not None else "").translate(PERSIAN_DIGITS).replace(",", "").strip()
    if text in ("", "-", "–", "*"):
        return 0
    return int(round(float(text)))


# --------------------------------------------------------------------------
# Table 1: population by single year of age and sex.

AGE_SINGLE = re.compile(r"^(\d+)ساله$")
AGE_GROUP = re.compile(r"^(\d+)-(\d+)ساله$")
AGE_OPEN = re.compile(r"^(\d+)ساله.*(?:بیشتر|بالا|بالاتر)$")


def age_of(label: Any) -> tuple[str, Any]:
    """('single', age) / ('group', (lo, hi)) / ('open', age) / ('total', None) / ('unstated', None)."""
    s = fa(label)
    if s in ("تمامیسنین", "جمع", "کل"):
        return "total", None
    if s == "کمترازیکساله":
        return "single", 0
    if s in ("اظهارنشده", "نامشخص"):
        return "unstated", None
    if m := AGE_OPEN.match(s):
        return "open", int(m.group(1))
    if m := AGE_SINGLE.match(s):
        return "single", int(m.group(1))
    if m := AGE_GROUP.match(s):
        return "group", (int(m.group(1)), int(m.group(2)))
    return "", None


def parse_ages(rows: list[list[Any]], label: str, block: int = 0) -> dict[str, Any]:
    """One Table 1 sheet: {"total": (both, men, women), "men": Counter, "women": Counter}.

    The columns are found by the header row that reads both sexes, men,
    women ("مردوزن", "مرد", "زن") under the "all" heading; the first such
    triple is the whole population, and the next three are the urban, rural
    and unsettled population (``block`` 1, 2 and 3), each under its heading
    in the row above, which is checked.
    """
    head = None
    for i, row in enumerate(rows[:8]):
        cells = [fa(c) for c in row]
        for j in range(len(cells) - 2):
            if cells[j] in ("مردوزن", "مردوزن") and cells[j + 1] == "مرد" and cells[j + 2] == "زن":
                head = (i, j)
                break
        if head:
            break
    if head is None:
        raise SystemExit(f"iran_census: {label}: no both/men/women header row")
    both_col = head[1] + 3 * block
    if block:
        sub = [fa(c) for c in rows[head[0]]]
        above = [fa(c) for c in rows[head[0] - 1]] if head[0] else []
        if (sub[both_col:both_col + 3] != ["مردوزن", "مرد", "زن"]
                or both_col >= len(above) or above[both_col] != BLOCKS[block]):
            raise SystemExit(f"iran_census: {label}: block {block} is not headed "
                             f"{BLOCKS[block]!r} over both sexes, men and women")
    label_col = None
    out = {"men": Counter(), "women": Counter(), "groups": {}, "unstated": [0, 0, 0]}
    total = None
    for row in rows[head[0] + 1:]:
        if label_col is None:
            for j in range(both_col):
                if age_of(row[j])[0] == "total":
                    label_col = j
                    break
            if label_col is None:
                continue
        kind, age = age_of(row[label_col])
        if not kind:
            continue
        both, men, women = (number(row[both_col + k]) if both_col + k < len(row) else 0
                            for k in range(3))
        if men + women != both:
            raise SystemExit(f"iran_census: {label}, {row[label_col]!r}: men {men:,} and "
                             f"women {women:,} do not make {both:,}")
        if kind == "total":
            if total is None:
                total = (both, men, women)
        elif kind in ("single", "open"):
            if age in out["men"] or age in out["women"]:
                raise SystemExit(f"iran_census: {label}: age {age} twice")
            out["men"][age] += men
            out["women"][age] += women
            if kind == "open":
                out["open"] = age
        elif kind == "group":
            out["groups"][age] = (both, men, women)
        elif kind == "unstated":
            out["unstated"] = [both, men, women]
    if total is None:
        raise SystemExit(f"iran_census: {label}: no all-ages row")
    out["total"] = total
    check_ages(out, label)
    return out


def check_ages(table: dict[str, Any], label: str) -> None:
    both, men, women = table["total"]
    s_men = sum(table["men"].values()) + table["unstated"][1]
    s_women = sum(table["women"].values()) + table["unstated"][2]
    if (s_men, s_women) != (men, women):
        raise SystemExit(f"iran_census: {label}: single years make {s_men:,} men and "
                         f"{s_women:,} women against {men:,} and {women:,}")
    for (lo, hi), (g_both, g_men, g_women) in table["groups"].items():
        m = sum(n for a, n in table["men"].items() if lo <= a <= hi)
        w = sum(n for a, n in table["women"].items() if lo <= a <= hi)
        if (m, w) != (g_men, g_women):
            raise SystemExit(f"iran_census: {label}: ages {lo}-{hi} make {m:,} men and "
                             f"{w:,} women against the sheet's {g_men:,} and {g_women:,}")


def add_ages(tables: list[dict[str, Any]]) -> dict[str, Any]:
    out = {"men": Counter(), "women": Counter(), "groups": {}, "unstated": [0, 0, 0],
           "total": (0, 0, 0)}
    for t in tables:
        out["men"].update(t["men"])
        out["women"].update(t["women"])
        out["unstated"] = [a + b for a, b in zip(out["unstated"], t["unstated"])]
        out["total"] = tuple(a + b for a, b in zip(out["total"], t["total"]))
    return out


def less(table: dict[str, Any], part: dict[str, Any]) -> dict[str, Any]:
    out = {"men": Counter(table["men"]), "women": Counter(table["women"]), "groups": {},
           "unstated": [a - b for a, b in zip(table["unstated"], part["unstated"])],
           "total": tuple(a - b for a, b in zip(table["total"], part["total"]))}
    out["men"].subtract(part["men"])
    out["women"].subtract(part["women"])
    if any(v < 0 for c in (out["men"], out["women"]) for v in c.values()):
        raise SystemExit("iran_census: a part has more people at some age than its whole")
    return out


# --------------------------------------------------------------------------
# Table 3: population by sex and citizenship.

def citizenship_columns(rows: list[list[Any]]) -> tuple[int, dict[int, str], int]:
    """(header row index, {column: label}, total column)."""
    for i, row in enumerate(rows[:6]):
        cells = [fa(c) for c in row]
        if fa("ایران") in cells and fa("افغانستان") in cells:
            columns = {}
            total = None
            for j, c in enumerate(cells):
                for persian, english in LABELS.items():
                    if c == fa(persian):
                        columns[j] = english
                if c == fa("جمع"):
                    total = j
            if total is None or len(columns) != len(LABELS):
                raise SystemExit(f"iran_census: citizenship header {row!r} lacks a column")
            return i, columns, total
    raise SystemExit("iran_census: no citizenship header row")


def parse_citizenship_block(rows: list[list[Any]], label: str) -> dict[str, Any]:
    """{name: (total, {label: count})} for the both-sexes block of a Table 3 sheet."""
    head, columns, total_col = citizenship_columns(rows)
    out: dict[str, Any] = {}
    in_block = False
    for row in rows[head + 1:]:
        name = fa(row[0]) if row else ""
        if name in ("مردوزن",):
            in_block = True
            out["*"] = read_citizenship(row, columns, total_col, label)
            continue
        if name in ("مرد", "زن"):
            if in_block:
                break
            continue
        if in_block and name:
            out[name] = read_citizenship(row, columns, total_col, f"{label} {row[0]}")
    if "*" not in out:
        raise SystemExit(f"iran_census: {label}: no both-sexes row")
    return out


def read_citizenship(row: list[Any], columns: dict[int, str], total_col: int,
                     label: str) -> tuple[int, dict[str, int]]:
    counts = {english: number(row[j]) for j, english in columns.items()}
    total = number(row[total_col])
    if sum(counts.values()) != total:
        raise SystemExit(f"iran_census: {label}: citizenships make {sum(counts.values()):,} "
                         f"against the total {total:,}")
    return total, counts


# --------------------------------------------------------------------------
# Reading the workbooks.

PAUSE = 10.0                          # seconds between two Wayback requests
COOLDOWNS = (120, 240, 300, 300)      # waits after the Archive refuses us
_last_request = [0.0]


OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"      # a legacy .xls
ZIP = b"PK\x03\x04"                             # an .xlsx


def fetch(url: str, magic: bytes = OLE) -> bytes:
    """One workbook from the Wayback Machine, spaced out and waited out.

    The Archive refuses connections for a minute or more when one address asks
    too often, and each capture costs two requests (the redirect to the exact
    timestamp, then the file). The first run was refused after six or seven
    workbooks in a row, so requests are spaced PAUSE seconds apart and a refusal
    is waited out for minutes rather than retried at once.
    """
    target = WAYBACK.format(url=url)
    blob: bytes | str = b""
    for cooldown in (*COOLDOWNS, None):
        wait = PAUSE - (time.monotonic() - _last_request[0])
        if wait > 0:
            time.sleep(wait)
        try:
            blob = http_get(target, binary=True, timeout=180, retries=1)
            break
        except RuntimeError as err:
            if cooldown is None:
                raise
            log(f"  the Archive refused {url.rsplit('/', 1)[-1]} ({err.__cause__}); "
                f"waiting {cooldown}s")
            time.sleep(cooldown)
        finally:
            _last_request[0] = time.monotonic()
    assert isinstance(blob, bytes)
    if not blob.startswith(magic):
        raise SystemExit(f"iran_census: {url} came back as something other than a "
                         f"workbook: {blob[:80]!r}")
    return blob


def sheets(blob: bytes) -> list[tuple[str, list[list[Any]]]]:
    import xlrd
    book = xlrd.open_workbook(file_contents=blob)
    return [(s.name, [s.row_values(i) for i in range(s.nrows)]) for s in book.sheets()]


SHEET_CODE = re.compile(r"^(.*?)(\d{2,4})$")


def split_sheet_name(name: str) -> tuple[str, str]:
    """('اراک', '0001') from 'اراک0001'."""
    s = name.strip().translate(PERSIAN_DIGITS)
    m = SHEET_CODE.match(s)
    if not m:
        raise SystemExit(f"iran_census: sheet {name!r} carries no code")
    return m.group(1).strip(), m.group(2)


# --------------------------------------------------------------------------
# Binding.

def load_codab(blob: bytes) -> dict[str, list[dict[str, str]]]:
    """OCHA's shahrestans by English province name: [{en, fa, pcode}]."""
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    rows = book["ADM2"].iter_rows(values_only=True)
    header = [str(c or "").strip() for c in next(rows)]
    col = {h: i for i, h in enumerate(header)}
    out: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        if not row or not row[col["ADM2_EN"]]:
            continue
        en = str(row[col["ADM2_EN"]]).replace("_x000D_", "").strip()
        province = str(row[col["ADM1_EN"]]).strip()
        out[CODAB_PROVINCE.get(province, province)].append(
            {"en": en, "fa": str(row[col["ADM2_FA"]]), "pcode": str(row[col["ADM2_PCODE"]])})
    return out


def ocha_fa(name: str) -> str:
    """A census sheet's Persian name, folded, as OCHA spells it (CENSUS_FA)."""
    folded = fa(name)
    for census, ocha in CENSUS_FA.items():
        if folded == fa(census):
            return fa(ocha)
    return folded


def drawn_units() -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    a1 = as_drawn(read_json(SITE / "admin1" / shard_name(ISO3), []))
    a2 = as_drawn(read_json(SITE / "admin2" / shard_name(ISO3), []))
    return a1, a2


def bind_counties(counties: list[dict[str, Any]], codab: dict[str, list[dict[str, str]]],
                  a1: list[dict[str, Any]], a2: list[dict[str, Any]]
                  ) -> tuple[dict[str, dict[str, Any]], list[str]]:
    """Centre code -> drawn polygon, through OCHA's Persian and English names.

    Returns ({code: shape}, notes). A county in SPLIT or NOT_DRAWN is left
    out with a note; anything else unmatched stops the run.
    """
    name1 = {u["id"]: u["name"] for u in a1}
    by_province: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for u in a2:
        by_province[name1.get(u.get("parent"), "")].append(u)
    bound: dict[str, dict[str, Any]] = {}
    notes: list[str] = []
    used: set[str] = set()
    problems: list[str] = []
    for county in counties:
        province = PROVINCES[county["province"]][1]
        wanted = ocha_fa(county["name"])
        rows = [r for r in codab.get(province, []) if fa(r["fa"]) == wanted]
        if len(rows) != 1:
            # Excel cuts a sheet name at 31 characters; the code takes four.
            rows = [r for r in codab.get(province, [])
                    if fa(r["fa"]).startswith(wanted) and len(wanted) >= 4]
        if len(rows) != 1:
            unclaimed = [f"{r['en']}={r['fa']}" for r in codab.get(province, [])
                         if not any(fa(r["fa"]) == ocha_fa(c["name"]) for c in counties
                                    if c["province"] == county["province"])]
            problems.append(f"{province} {county['name']} ({county['code']}) matches "
                            f"{len(rows)} of OCHA's shahrestans; OCHA's unmatched there: "
                            f"{', '.join(unclaimed)}")
            continue
        en = rows[0]["en"]
        county["en"] = en
        if en in NOT_DRAWN:
            notes.append(f"{en} ({county['code']}): not drawn; left out")
            continue
        if en in SPLIT:
            notes.append(f"{en} ({county['code']}): drawn as {SPLIT[en]['part_label']!r} and "
                         f"{SPLIT[en]['rest_label']!r}; the shahrestan's figure goes on "
                         f"neither, each takes its own part's")
            continue
        drawn_province = ELSEWHERE.get(en, province)
        label = DRAWN.get(en, en)
        pool = by_province.get(drawn_province, [])
        hits = [u for u in pool if fold(u["name"]) == fold(label)]
        if en in ELSEWHERE:
            hits = [u for u in hits
                    if sum(1 for v in a2 if v.get("parent") == u.get("parent")) == 1]
        if len(hits) != 1 or hits[0]["id"] in used:
            problems.append(f"{province} {en} ({county['code']}) finds {len(hits)} drawn "
                            f"polygons labelled {label!r}")
            continue
        used.add(hits[0]["id"])
        bound[county["code"]] = hits[0]
    if problems:
        for line in problems:
            log(f"  unbound: {line}")
        raise SystemExit(f"iran_census: {len(problems)} shahrestans not bound (listed above)")
    return bound, notes


# --------------------------------------------------------------------------
# Records.

def fields(table: dict[str, Any], citizenship: tuple[int, dict[str, int]] | None,
           note_extra: str = "") -> dict[str, Any]:
    both, men, women = table["total"]
    ages = Counter({a: table["men"][a] + table["women"][a]
                    for a in set(table["men"]) | set(table["women"])})
    median = median_age(ages)
    out: dict[str, Any] = {
        "population": measure(both, year=YEAR, source=SOURCE),
        "median_age": measure(median, unit="years", year=YEAR, source=SOURCE),
        "median_age_note": ("Interpolated within the single year of age that holds the middle "
                            "person, from Table 1 of the 2016 census (population by single year "
                            "of age and sex)."
                            + (f" {table['unstated'][0]:,} people of unstated age are left "
                               f"out." if table["unstated"][0] else "") + note_extra),
        "sex_ratio": measure(round(100 * men / women, 1), unit="males_per_100_females",
                             year=YEAR, source=SOURCE),
        "sex_ratio_note": f"{men:,} men and {women:,} women, 2016 census." + note_extra,
        "sources": [{"field": "population/median_age/sex_ratio", "name": SOURCE,
                     "url": PAGE, "year": YEAR, "license": LICENCE}],
    }
    if citizenship is not None:
        total, counts = citizenship
        if total != both:
            raise SystemExit(f"iran_census: citizenship total {total:,} against the age "
                             f"table's {both:,}")
        unknown = set(counts) - set(SHOWN)
        if unknown:
            raise SystemExit(f"iran_census: no label for citizenship {sorted(unknown)}")
        out["ethnicity"] = shares({SHOWN[k]: v for k, v in counts.items() if v}, total=total)
        out["ethnicity_year"] = YEAR
        out["ethnicity_basis"] = "citizenship"
        out["ethnicity_note"] = (
            "Country of citizenship, from Table 3 of the 2016 census (population by sex and "
            "citizenship): Iranian nationals and the nationals of each other country the "
            "Centre tabulates. Iran's census does not ask ethnicity; by the owner's decision of "
            "19 September 2026 the citizenship it counts is shown in its place, as Japan's and "
            "Korea's nationality is. It is citizenship, not ethnic group: every Iranian "
            "citizen, whatever their ethnicity, is an 'Iranian national' here."
            + note_extra)
        out["sources"].append({"field": "ethnicity", "name": SOURCE, "url": PAGE,
                               "year": YEAR, "license": LICENCE})
    return out


def build(province_ages: dict[int, dict[str, Any]],
          province_citizenship: dict[str, tuple[int, dict[str, int]]],
          counties: list[dict[str, Any]], codab: dict[str, list[dict[str, str]]],
          a1: list[dict[str, Any]], a2: list[dict[str, Any]]
          ) -> tuple[list[dict[str, Any]], list[str]]:
    # The provinces add up to the country, and each province's shahrestans to it.
    national = sum(t["total"][0] for t in province_ages.values())
    if national != NATIONAL:
        raise SystemExit(f"iran_census: the provinces make {national:,} against {NATIONAL:,}")
    for nn, table in province_ages.items():
        parts = [c["ages"] for c in counties if c["province"] == nn]
        summed = add_ages(parts)
        if summed["total"] != table["total"]:
            raise SystemExit(f"iran_census: {PROVINCES[nn][1]}'s shahrestans make "
                             f"{summed['total']} against the province's {table['total']}")
    bound, notes = bind_counties(counties, codab, a1, a2)
    rows: list[dict[str, Any]] = []
    for county in counties:
        shape = bound.get(county["code"])
        if shape is None:
            continue
        en = county["en"]
        name = en if en in MISLABELLED else shape["name"]
        extra = (f" The boundary file labels this polygon {MISLABELLED[en]!r}; "
                 f"it is {en} ({county['name']}), whose town lies in it."
                 if en in MISLABELLED else "")
        rows.append(record(
            f"IRN-CENSUS-{county['code']}", name, level="admin2", parent=ISO3, country=ISO3,
            match_by="shape_id", shape_id=shape["id"],
            aliases=[county["name"], en] + ([shape["name"]] if en in MISLABELLED else []),
            codes={"sci": county["code"]},
            **fields(county["ages"], county["citizenship"], extra)))
    # The first level: every province whole, except where the map draws
    # Bandar-e Gaz under a first-level polygon of its own.
    gaz = next((c for c in counties if c.get("en") == "Bandar-e-Gaz"), None)
    by_label: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for u in a1:
        by_label[u["name"]].append(u)
    children = Counter(u.get("parent") for u in a2)
    for nn, (persian, english) in PROVINCES.items():
        polygons = sorted(by_label.get(english, []), key=lambda u: -children[u["id"]])
        if not polygons:
            raise SystemExit(f"iran_census: the map draws no {english}")
        table = province_ages[nn]
        citizen = province_citizenship.get(fa(persian))
        if citizen is None:
            raise SystemExit(f"iran_census: Table 3 has no row for {persian}")
        extra = ""
        if english == "Golestan" and gaz is not None:
            table = less(table, gaz["ages"])
            total, counts = citizen
            g_total, g_counts = gaz["citizenship"]
            citizen = (total - g_total, {k: counts[k] - g_counts[k] for k in counts})
            extra = (" The map draws Bandar-e Gaz county under a first-level polygon of its "
                     "own, so this is Golestan less Bandar-e Gaz, from the same tables.")
        rows.append(record(
            f"IRN-CENSUS-P{nn:02d}", english, level="admin1", parent=ISO3, country=ISO3,
            match_by="shape_id", shape_id=polygons[0]["id"], aliases=[persian],
            codes={"sci": f"{nn:02d}"}, **fields(table, citizen, extra)))
        if english == "Mazandaran" and len(polygons) > 1 and gaz is not None:
            small = polygons[-1]
            if children[small["id"]] != 1:
                raise SystemExit("iran_census: the second Mazandaran polygon does not hold "
                                 "exactly one shahrestan")
            rows.append(record(
                f"IRN-CENSUS-P{nn:02d}-GAZ", "Mazandaran", level="admin1", parent=ISO3,
                country=ISO3, match_by="shape_id", shape_id=small["id"],
                aliases=["Bandar-e Gaz"], codes={"sci": gaz["code"]},
                **fields(gaz["ages"], gaz["citizenship"],
                         " This first-level polygon holds Bandar-e Gaz county alone (a county "
                         "of Golestan province), so it carries Bandar-e Gaz's figures.")))
    return rows, notes


def parse_religion(pages: list[str]) -> dict[str, dict[str, int]]:
    """{map label: {religion: count}} from table 3.18, with table 3.17's check.

    Every row's religions make its total, the provinces make the country's
    row, and table 3.17's 2011 column must carry the swapped headings that
    RELIGION_COLUMNS undoes (see there) -- otherwise the run stops.
    """
    text = "\n".join(pages)
    head = re.search(r"3\.17\.\s*POPULATION BY SEX AND RELIGION(.*?)Source", text, re.S)
    if not head:
        raise SystemExit("iran_census: the yearbook has no table 3.17")
    printed = {}
    for label in ("Christian", "Zoroastrian"):
        m = re.search(rf"^{label}\s*\.*\s*((?:\d+\s+){{8}}\d+)\s*$", head.group(1), re.M)
        if not m:
            raise SystemExit(f"iran_census: table 3.17 has no {label} row")
        printed[label] = int(m.group(1).split()[3])        # 1390, both sexes
    if (printed["Zoroastrian"], printed["Christian"]) != (
            CENSUS_2011_PERSIAN["Christianity"], CENSUS_2011_PERSIAN["Zoroastrianism"]):
        raise SystemExit(f"iran_census: table 3.17's 2011 column reads {printed}; the "
                         f"heading swap RELIGION_COLUMNS undoes is not there")
    body = re.search(r"3\.18\.\s*POPULATION BY RELIGION AND OSTAN(.*?)Source", text, re.S)
    if not body:
        raise SystemExit("iran_census: the yearbook has no table 3.18")
    out: dict[str, dict[str, int]] = {}
    for line in body.group(1).splitlines():
        m = re.match(r"^\s*(.+?)\s*\.{2,}\s*((?:\d+\s+){6}\d+)\s*$", line)
        if not m:
            continue
        name = " ".join(m.group(1).split())
        total, *counts = (int(x) for x in m.group(2).split())
        if sum(counts) != total:
            raise SystemExit(f"iran_census: table 3.18's {name}: religions make "
                             f"{sum(counts):,} against {total:,}")
        if name == "Total country":
            out["*"] = dict(zip(RELIGION_COLUMNS, counts))
            continue
        if name not in YEARBOOK_OSTAN:
            raise SystemExit(f"iran_census: table 3.18's {name!r} is not a known ostan")
        out[YEARBOOK_OSTAN[name]] = dict(zip(RELIGION_COLUMNS, counts))
    if "*" not in out or len(out) != len(PROVINCES) + 1:
        raise SystemExit(f"iran_census: table 3.18 has {len(out) - ('*' in out)} ostans")
    made = {k: sum(v[k] for name, v in out.items() if name != "*") for k in RELIGION_COLUMNS}
    if made != out["*"]:
        raise SystemExit(f"iran_census: table 3.18's ostans make {made} against {out['*']}")
    return out


def read_religion() -> dict[str, dict[str, int]]:
    kept = cached("religion")
    if kept is None:
        import logging

        from pypdf import PdfReader
        logging.getLogger("pypdf").setLevel(logging.ERROR)
        blob = http_get(WAYBACK.replace("2019id_", "2023id_").format(url=YEARBOOK),
                        binary=True, timeout=300)
        assert isinstance(blob, bytes)
        if blob[:5] != b"%PDF-":
            raise SystemExit(f"iran_census: {YEARBOOK} came back as something other than a "
                             f"PDF: {blob[:80]!r}")
        pages = [(p.extract_text() or "") for p in PdfReader(io.BytesIO(blob)).pages]
        kept = parse_religion(pages)
        keep("religion", kept)
    return kept


def add_religion(rows: list[dict[str, Any]], religion: dict[str, dict[str, int]],
                 province_ages: dict[int, dict[str, Any]]) -> None:
    """Religion on every first-level record that is a whole province.

    Golestan is drawn without Bandar-e Gaz, which has a first-level polygon
    of its own; the yearbook counts the province whole and publishes nothing
    for the county, so neither polygon takes a share, and each says why.
    """
    note = ("Religion as each person declared it in the 2016 census (the Centre's "
            "categories: Muslim, Christian, Zoroastrian, Jewish, other, not stated), from "
            "table 3.18 of the Statistical Yearbook 1395, read from the Internet Archive's "
            "capture of amar.org.ir. The English yearbook prints the Christian and "
            "Zoroastrian headings swapped; they are restored from the Centre's Persian "
            "results (see the adapter).")
    for r in rows:
        if r["level"] == "admin2":
            if isinstance(r.get("religion"), list):
                continue
            r["religion"] = gap(NOT_AVAILABLE, (
                "The 2016 census's religion is published by province only (Statistical "
                "Yearbook 1395, table 3.18). No table of religion by shahrestan is among the "
                "Centre's archived detailed census tables (population tables 1, 2, 3 and 9; "
                "activity, households, housing, migration, education, marriage), so no "
                "share is written below the province."))
            continue
        if r["level"] != "admin1":
            continue
        name = r["name"]
        if r["id"].endswith("-GAZ") or name == "Golestan":
            r["religion"] = gap(NOT_AVAILABLE, (
                "The 2016 census's religion is published by province only (Statistical "
                "Yearbook 1395, table 3.18). The map draws Bandar-e Gaz county under a "
                "first-level polygon of its own, so neither it nor Golestan without it is a "
                "unit the table counts, and no share is written for either."))
            continue
        counts = religion[name]
        nn = next(k for k, v in PROVINCES.items() if v[1] == name)
        if sum(counts.values()) != province_ages[nn]["total"][0]:
            raise SystemExit(f"iran_census: {name}'s religions make {sum(counts.values()):,} "
                             f"against the census's {province_ages[nn]['total'][0]:,}")
        r["religion"] = shares({k: v for k, v in counts.items() if v})
        r["religion_year"] = YEAR
        r["religion_note"] = note
        r["sources"].append({"field": "religion", "name": YEARBOOK_SOURCE, "url": YEARBOOK,
                             "year": YEAR, "license": LICENCE})


def code_of(cell: Any, width: int) -> str:
    """A settlement-table code as text of its printed width ('01', '1576')."""
    text = str(cell if cell is not None else "").strip()
    if re.fullmatch(r"\d+(?:\.0)?", text):
        return f"{int(float(text)):0{width}d}"
    return text


def read_settlements(rows: list[list[Any]], county: str, label: str) -> dict[str, Any]:
    """One county's rows of a province's settlement table.

    {"county": (persons, men, women), "districts": {code: (p, m, w)},
    "cities": {code: (p, m, w)}, "regions": {city code: (p, m, w) summed}}.
    Every row's men and women make its persons, a zoned city's regions make
    the city, and the districts make no more than the county (the rest, if
    any, are the unsettled, whom the table counts at the county alone).
    """
    if len(rows) < 3:
        raise SystemExit(f"iran_census: {label}: the settlement table is empty")
    head = [fa(c) if not str(c or "").isascii() else str(c or "").strip() for c in rows[1]]
    for col, name in SETTLEMENT_COLUMNS.items():
        if col >= len(head) or head[col] != name:
            raise SystemExit(f"iran_census: {label}: the settlement table's column {col} reads "
                             f"{head[col] if col < len(head) else None!r}, not {name!r}")
    out: dict[str, Any] = {"county": None, "districts": {}, "cities": {}, "regions": {}}
    for row in rows[2:]:
        row = list(row) + [None] * (18 - len(row))
        if code_of(row[2], 2) != county:
            continue
        kind = code_of(row[11], 1)
        if kind not in ("2", "3", "5"):
            continue
        p, m, w = (number(row[c]) for c in (15, 16, 17))
        if m + w != p:
            raise SystemExit(f"iran_census: {label}: {row[3]} {row[5]} {row[7]}: men {m:,} and "
                             f"women {w:,} do not make {p:,}")
        if kind == "2":
            out["county"] = (p, m, w)
        elif kind == "3":
            out["districts"][code_of(row[4], 2)] = (p, m, w)
        else:
            top = code_of(row[12], 4)
            if top:
                old = out["regions"].get(top, (0, 0, 0))
                out["regions"][top] = (old[0] + p, old[1] + m, old[2] + w)
            else:
                out["cities"][code_of(row[6], 4)] = (p, m, w)
    if out["county"] is None:
        raise SystemExit(f"iran_census: {label}: no row for county {county}")
    for city, summed in out["regions"].items():
        if out["cities"].get(city) != summed:
            raise SystemExit(f"iran_census: {label}: city {city}'s regions make {summed}, "
                             f"against the city's {out['cities'].get(city)}")
    made = tuple(sum(d[k] for d in out["districts"].values()) for k in range(3))
    if any(made[k] > out["county"][k] for k in range(3)):
        raise SystemExit(f"iran_census: {label}: the districts make {made}, more than the "
                         f"county's {out['county']}")
    return out


def less3(a: tuple[int, int, int], b: tuple[int, int, int]) -> tuple[int, int, int]:
    return (a[0] - b[0], a[1] - b[1], a[2] - b[2])


def split_figures(en: str, settlements: dict[str, Any], county_ages: dict[str, Any],
                  sheet: list[list[Any]] | None) -> dict[str, Any]:
    """What each half of a split shahrestan is, checked against Table 1.

    ``county_ages`` is the shahrestan's Table 1 (all ages, both sexes); the
    settlement table must count the shahrestan at exactly its total. The part
    one polygon is (a city, or a district) is read from the settlement table
    and the rest is the shahrestan less it. Where the part is the
    shahrestan's only city, Table 1's urban block is that city -- checked to
    the person, men and women -- and its rural and unsettled blocks are the
    rest, so ``sheet`` (the shahrestan's Table 1 sheet) gives both halves'
    ages.
    """
    spec = SPLIT[en]
    county = settlements["county"]
    if county != tuple(county_ages["total"]):
        raise SystemExit(f"iran_census: {en}: the settlement table counts {county}, Table 1 "
                         f"{tuple(county_ages['total'])}")
    kind, code = spec["part"]
    pool = settlements["cities"] if kind == "city" else settlements["districts"]
    if code not in pool:
        raise SystemExit(f"iran_census: {en}: the settlement table has no {kind} {code}")
    part = pool[code]
    rest = less3(county, part)
    if min(rest) <= 0:
        raise SystemExit(f"iran_census: {en}: the rest of the county is {rest}")
    others = sorted(c for c in settlements["cities"] if c != code)
    out: dict[str, Any] = {"county": list(county), "part": list(part), "rest": list(rest),
                           "other_cities": len(others), "ages": None}
    if kind == "city" and not others:
        if sheet is None:
            raise SystemExit(f"iran_census: {en}: its only city's ages need Table 1's sheet")
        blocks = [parse_ages(sheet, f"{en} block {b}", block=b) for b in range(4)]
        if tuple(blocks[0]["total"]) != county:
            raise SystemExit(f"iran_census: {en}: Table 1's sheet counts {blocks[0]['total']}")
        if tuple(blocks[1]["total"]) != part:
            raise SystemExit(f"iran_census: {en}: Table 1's urban block counts "
                             f"{blocks[1]['total']}, the settlement table's only city {part}")
        remainder = add_ages([blocks[2], blocks[3]])
        whole = add_ages([blocks[1], remainder])
        if (whole["men"] != blocks[0]["men"] or whole["women"] != blocks[0]["women"]
                or tuple(whole["total"]) != county):
            raise SystemExit(f"iran_census: {en}: the urban, rural and unsettled blocks do "
                             f"not make the whole, age by age")
        out["ages"] = {"part": table_json(blocks[1]), "rest": table_json(remainder)}
        out["unsettled"] = list(blocks[3]["total"])
    return out


def split_records(a1: list[dict[str, Any]], a2: list[dict[str, Any]],
                  figures: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    """Each half of every 2016 shahrestan the map draws as two, with its own count.

    The census counts the shahrestan whole, and its figure describes neither
    polygon. Each half takes what the census counts for the part it is
    (``SPLIT``, ``split_figures``): persons and sex from the settlement table,
    ages where Table 1 isolates the part, and a stated gap for the rest.
    """
    name1 = {u["id"]: u["name"] for u in a1}
    rows = []
    for en, spec in SPLIT.items():
        province = PROVINCES[spec["province"]][1]
        made = figures[en]
        others = made["other_cities"]
        for half in ("part", "rest"):
            label = spec[f"{half}_label"]
            hits = [u for u in a2 if u["name"] == label and name1.get(u.get("parent")) == province]
            if len(hits) != 1:
                raise SystemExit(f"iran_census: {len(hits)} drawn polygons labelled {label!r} "
                                 f"in {province}")
            other = spec["rest_label" if half == "part" else "part_label"]
            what = spec[f"{half}_what"].format(others=others)
            persons, men, women = made[half]
            note = (f"The 2016 census counts {en} shahrestan whole ({made['county'][0]:,} "
                    f"people); the map draws it as two polygons, and this one is {what}, "
                    f"which the census counts at {persons:,}. The other polygon, "
                    f"{other!r}, is the rest.")
            if made.get("ages"):
                table = table_from(made["ages"][half])
                extra = (" Read from Table 1's urban block, which is the city: the city is "
                         "the shahrestan's only one, and the settlement table counts it at the "
                         "same persons, men and women." if half == "part" else
                         f" Read from Table 1's rural and unsettled blocks; the "
                         f"{made['unsettled'][0]:,} unsettled people are counted at the "
                         f"shahrestan alone, so they are with its rural ground here.")
                body = fields(table, None, extra)
                sources = body.pop("sources")
            else:
                body = {
                    "population": measure(persons, year=YEAR, source=SETTLEMENT_SOURCE),
                    "sex_ratio": measure(round(100 * men / women, 1),
                                         unit="males_per_100_females", year=YEAR,
                                         source=SETTLEMENT_SOURCE),
                    "sex_ratio_note": (f"{men:,} men and {women:,} women, from the 2016 "
                                       f"census's settlement table."),
                    "median_age": gap(NOT_AVAILABLE, MEDIAN_SPLIT_GAP[en].format(others=others)),
                }
                sources = [{"field": "population/sex_ratio", "name": SETTLEMENT_SOURCE,
                            "url": PAGE, "year": YEAR, "license": LICENCE}]
            rows.append(record(
                f"IRN-CENSUS-SPLIT-{fold(label)}", label, level="admin2", parent=ISO3,
                country=ISO3, match_by="shape_id", shape_id=hits[0]["id"],
                codes={"sci": spec["code"]}, population_note=note,
                ethnicity=gap(NOT_AVAILABLE, CITIZENSHIP_SPLIT_GAP.format(county=en)),
                sources=sources, **body))
    return rows


def read_split(counties: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    """The split shahrestans' halves, from data/raw/iran or the Archive."""
    import openpyxl
    kept = cached("split")
    if kept is not None and set(kept) == set(SPLIT):
        return kept
    by_code = {c["code"]: c for c in counties}
    out: dict[str, dict[str, Any]] = {}
    for en, spec in SPLIT.items():
        nn = spec["province"]
        book = openpyxl.load_workbook(io.BytesIO(fetch(SETTLEMENTS.format(nn=nn), magic=ZIP)),
                                      read_only=True, data_only=True)
        rows = [list(r) for r in book[book.sheetnames[0]].iter_rows(values_only=True)]
        settlements = read_settlements(rows, spec["county"], f"settlements {nn:02d}")
        county = by_code.get(spec["code"])
        if county is None:
            raise SystemExit(f"iran_census: no Table 1 sheet {spec['code']} for {en}")
        kind, code = spec["part"]
        table1 = None
        if kind == "city" and not [c for c in settlements["cities"] if c != code]:
            table1 = {f"{nn:02d}{split_sheet_name(n)[1][-2:]}": r
                      for n, r in sheets(fetch(COUNTY_AGES.format(nn=nn)))}.get(spec["code"])
        out[en] = split_figures(en, settlements, county["ages"], table1)
    keep("split", out)
    return out


# What has been read is kept, as the tables' counts, under data/raw/iran (which
# the repository tracks): the Archive refuses a run that asks for all 64
# workbooks at once often enough that a run reads what it can within BUDGET,
# keeps it, and stops; the next run starts where it stopped. The kept files
# are the parsed tables -- counts by single year of age and sex, and by
# citizenship -- not the workbooks.
CACHE = PROCESSED.parent / "raw" / "iran"
BUDGET = 30 * 60


def table_json(table: dict[str, Any]) -> dict[str, Any]:
    return {"men": {str(a): n for a, n in sorted(table["men"].items())},
            "women": {str(a): n for a, n in sorted(table["women"].items())},
            "groups": [[lo, hi, *v] for (lo, hi), v in sorted(table["groups"].items())],
            "unstated": list(table["unstated"]), "total": list(table["total"]),
            **({"open": table["open"]} if "open" in table else {})}


def table_from(data: dict[str, Any]) -> dict[str, Any]:
    out = {"men": Counter({int(a): n for a, n in data["men"].items()}),
           "women": Counter({int(a): n for a, n in data["women"].items()}),
           "groups": {(g[0], g[1]): tuple(g[2:]) for g in data["groups"]},
           "unstated": list(data["unstated"]), "total": tuple(data["total"])}
    if "open" in data:
        out["open"] = data["open"]
    return out


def cached(name: str) -> Any:
    path = CACHE / f"census2016_{name}.json"
    return read_json(path, None) if path.exists() else None


def keep(name: str, payload: Any) -> None:
    write_json(CACHE / f"census2016_{name}.json", payload)


def read_provinces() -> tuple[dict[int, Any], dict[str, Any]]:
    kept = cached("provinces")
    if kept is None:
        province_ages: dict[int, dict[str, Any]] = {}
        for name, rows in sheets(fetch(PROVINCE_AGES)):
            persian, code = split_sheet_name(name)
            nn = int(code)
            if fa(persian) != fa(PROVINCES[nn][0]):
                raise SystemExit(f"iran_census: province sheet {name!r} is not "
                                 f"{PROVINCES[nn][0]}")
            province_ages[nn] = parse_ages(rows, name)
        if sorted(province_ages) != sorted(PROVINCES):
            raise SystemExit(f"iran_census: province sheets {sorted(province_ages)}")
        (_, rows), = sheets(fetch(PROVINCE_CITIZENSHIP))
        block = parse_citizenship_block(rows, "3-jamiat-k")
        kept = {"ages": {str(nn): table_json(t) for nn, t in province_ages.items()},
                "citizenship": {k: [v[0], v[1]] for k, v in block.items()}}
        keep("provinces", kept)
    ages = {int(nn): table_from(t) for nn, t in kept["ages"].items()}
    citizenship = {k: (v[0], v[1]) for k, v in kept["citizenship"].items()}
    for nn, table in ages.items():
        check_ages(table, f"province {nn:02d}")
    return ages, citizenship


def read_province(nn: int) -> list[dict[str, Any]]:
    """One province's shahrestans: [{province, code, name, ages, citizenship}]."""
    # Keyed by the sheet's code: the two workbooks may cut a long name at
    # different lengths, never the code.
    ages = {split_sheet_name(n)[1]: (split_sheet_name(n)[0], r)
            for n, r in sheets(fetch(COUNTY_AGES.format(nn=nn)))}
    cit = {split_sheet_name(n)[1]: r
           for n, r in sheets(fetch(COUNTY_CITIZENSHIP.format(nn=nn)))}
    if set(ages) != set(cit):
        raise SystemExit(f"iran_census: province {nn:02d}'s two tables carry different "
                         f"codes: {sorted(set(ages) ^ set(cit))}")
    counties = []
    for code, (persian, rows) in sorted(ages.items()):
        full = f"{nn:02d}{code[-2:]}"
        if len(code) == 4 and int(code[:2]) != nn:
            raise SystemExit(f"iran_census: sheet {persian}{code} is filed under "
                             f"province {nn:02d}")
        block = parse_citizenship_block(cit[code], f"3-jamiat-{nn:02d} {persian}")
        counties.append({"province": nn, "code": full, "name": persian,
                         "ages": parse_ages(rows, f"jamiat{nn:02d} {persian}"),
                         "citizenship": block["*"]})
    return counties


def read_all(budget: float = BUDGET, only: set[int] | None = None
             ) -> tuple[dict[int, Any], dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    """Every province's shahrestans, from data/raw/iran or the Archive.

    ``only`` limits the fetching to those provinces (the rest are read from
    data/raw/iran or left for another run), so that several runs on several
    machines can share the work the Archive rations.
    """
    started = time.monotonic()
    province_ages, province_citizenship = read_provinces()
    counties: list[dict[str, Any]] = []
    waiting = []
    for nn in sorted(PROVINCES):
        kept = cached(f"{nn:02d}")
        if kept is None:
            if time.monotonic() - started > budget or (only is not None and nn not in only):
                waiting.append(nn)
                continue
            fresh = read_province(nn)
            kept = [{**c, "ages": table_json(c["ages"]),
                     "citizenship": [c["citizenship"][0], c["citizenship"][1]]} for c in fresh]
            keep(f"{nn:02d}", kept)
        mine = [{**c, "ages": table_from(c["ages"]),
                 "citizenship": (c["citizenship"][0], c["citizenship"][1])} for c in kept]
        for c in mine:
            check_ages(c["ages"], f"{c['code']} {c['name']}")
        counties += mine
        log(f"  province {nn:02d} {PROVINCES[nn][1]}: {len(mine)} shahrestans")
    if waiting:
        raise SystemExit(f"iran_census: {len(waiting)} provinces not yet read ({waiting}); "
                         f"the {len(PROVINCES) - len(waiting)} read are kept in data/raw/iran "
                         f"and a re-run carries on from them")
    codab = load_codab(http_get(CODAB, binary=True))  # type: ignore[arg-type]
    return province_ages, province_citizenship, counties, codab


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--provinces", default="",
                    help="fetch only these provinces' tables, e.g. 5-9,12 (SCI codes)")
    args = ap.parse_args()
    only = None
    if args.provinces:
        only = set()
        for part in args.provinces.split(","):
            low, _, high = part.partition("-")
            only.update(range(int(low), int(high or low) + 1))
    log("iran_census: 2016 census tables from the Internet Archive's captures of amar.org.ir")
    province_ages, province_citizenship, counties, codab = read_all(only=only)
    a1, a2 = drawn_units()
    rows, notes = build(province_ages, province_citizenship, counties, codab, a1, a2)
    religion = read_religion()
    halves = split_records(a1, a2, read_split(counties))
    add_religion(rows + halves, religion, province_ages)
    for line in notes:
        log(f"  {line}")
    national = add_ages(list(province_ages.values()))
    ages = Counter({a: national["men"][a] + national["women"][a]
                    for a in set(national["men"]) | set(national["women"])})
    log(f"  Iran: median age {median_age(ages)}, {national['total'][1]:,} men and "
        f"{national['total'][2]:,} women")
    for level in ("admin1", "admin2"):
        mine = [r for r in rows if r["level"] == level]
        log(f"  {level}: {len(mine)} records")
    for r in rows:
        eth = ", ".join(f"{s['group']} {s['pct']}" for s in r["ethnicity"][:3])
        log(f"    {r['level']} {r['name']} [{r['shape_id']}]: {r['population']['value']:,}, "
            f"median {r['median_age']['value']}, ratio {r['sex_ratio']['value']}; {eth}")
    for r in halves:
        median = r["median_age"].get("value", "a stated gap")
        log(f"    admin2 {r['name']} [{r['shape_id']}] (half of a split shahrestan): "
            f"{r['population']['value']:,}, median {median}, ratio {r['sex_ratio']['value']}")
    for r in rows:
        if r["level"] == "admin1" and isinstance(r.get("religion"), list):
            log(f"    religion {r['name']}: " + ", ".join(
                f"{s['group']} {s['pct']}" for s in r["religion"][:3]))
    write_json(PROCESSED / OUT, rows + halves)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
