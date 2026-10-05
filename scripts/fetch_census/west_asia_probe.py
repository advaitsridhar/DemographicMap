#!/usr/bin/env python3
"""Read-only probes for the statistical offices of the Arab states and Israel.

The Gulf offices publish through open-data portals of three kinds --
OpenDataSoft (Bahrain's data.gov.bh, Qatar's data.gov.qa), CKAN (Saudi
Arabia's open data) and single-page apps with a JSON API behind them -- and
the Levant's through WordPress sites and SharePoint libraries. Each of the
subcommands below answers one question about one of those and prints only
what decides it. Nothing is written and nothing is committed but the log.
Every argument is one whitespace-free token, as the workflow splits its
command on whitespace.

Subcommands:

    ods BASE [--search WORDS] [--rows N]
        An OpenDataSoft catalogue: each dataset's id, title, record count and
        field names. WORDS is a '+'-joined search.
    odsrows BASE DATASET [--where W] [--select S] [--group-by G] [--rows N]
        Records of one OpenDataSoft dataset (Explore API v2.1).
    get URL [URL ...] [--grep REGEX] [--links REGEX] [--chars N]
        Status, type and length of each answer; the text around each REGEX
        match, the links whose address or label match, or the opening.
    wayback URL [--rows N] [--grep REGEX]
        The Internet Archive's captures of a URL pattern (CDX); with REGEX,
        only those whose address matches.
    xlsx URL [URL ...] [--rows N] [--cols N] [--sheets N] [--width N]
        A workbook's sheets and first rows, read as .xlsx whatever the URL.
    pdf URL [URL ...] [--pages P] [--chars N] [--grep REGEX]
        A PDF's page count and, laid out row by row, the opening of each page
        in P (comma-separated, 1-based; default the first two), or of each
        page matching REGEX; with --chars 0, only the lines matching REGEX.
    uscb DATASET SHEET [--grep REGEX] [--rows N] [--level L] [--dictionary REGEX]
        One sheet of the US Census Bureau's workbook in an HDX dataset: its
        field names and aliases, and the rows (of ADM_LEVEL L, if given) with
        the geography columns and the fields whose name matches REGEX; with
        --dictionary, the workbook's data-dictionary rows (definition, the
        office's original field name, its table) of the fields matching it.
    geonames CC [--iso ISO3] [--unit-level L] [--classes P,L] [--grep REGEX] [--rows N]
        GeoNames' dump of one country (``CC.zip``): every feature of the given
        classes, with its Arabic names, its point, and the drawn unit of
        ``ISO3`` at level ``L`` (default admin2) whose polygon, read from the
        map's own tiles, holds the point. REGEX filters on the feature's names.
    wdpoints QID [--iso ISO3] [--unit-level L] [--grep REGEX] [--rows N]
        Wikidata's settlements, neighbourhoods and districts in the country
        whose item is QID, with their Arabic and English labels, their points
        and the drawn unit holding each.

Usage:
    python -m scripts.fetch_census.west_asia_probe ods https://www.data.gov.bh --search population
    python -m scripts.fetch_census.west_asia_probe get https://example.org --links xlsx
"""

from __future__ import annotations

import argparse
import html
import json
import re
import urllib.error
import urllib.parse
import urllib.request
from typing import Any

from ._shared import log

UA = ("DemographicMap/1.0 (+https://github.com/advaitsridhar/DemographicMap) "
      "python-urllib")
TIMEOUT = 90
CDX = "https://web.archive.org/cdx/search/cdx"


def fetch(url: str, *, accept: str | None = None, limit: int | None = None
          ) -> tuple[int, str, bytes]:
    headers = {"User-Agent": UA}
    if accept:
        headers["Accept"] = accept
    req = urllib.request.Request(url, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=TIMEOUT) as fh:
            body = fh.read(limit) if limit else fh.read()
            return fh.status, fh.headers.get("Content-Type", ""), body
    except urllib.error.HTTPError as err:
        body = b""
        try:
            body = err.read()[:3000]
        except Exception:  # noqa: BLE001
            pass
        return err.code, (err.headers.get("Content-Type", "") if err.headers else ""), body
    except Exception as err:  # noqa: BLE001 -- a probe reports, it does not raise
        return -1, type(err).__name__, str(err).encode()


def text_of(body: bytes, ctype: str) -> str:
    m = re.search(r"charset=([\w-]+)", ctype or "")
    enc = m.group(1) if m else "utf-8"
    try:
        return body.decode(enc, "replace")
    except LookupError:
        return body.decode("utf-8", "replace")


def plain(page: str) -> str:
    page = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", page)
    page = re.sub(r"(?i)<br\s*/?>|</(p|div|tr|li|h\d|td|th)>", "\n", page)
    page = re.sub(r"<[^>]+>", " ", page)
    page = html.unescape(page)
    return "\n".join(" ".join(line.split()) for line in page.split("\n") if line.strip())


def cmd_ods(base: str, search: str | None, rows: int) -> None:
    base = base.rstrip("/")
    params = {"limit": str(min(rows, 100))}
    if search:
        params["where"] = " OR ".join(f'search("{w}")' for w in search.split("+"))
    offset = 0
    shown = 0
    while shown < rows:
        params["offset"] = str(offset)
        url = f"{base}/api/explore/v2.1/catalog/datasets?" + urllib.parse.urlencode(params)
        status, ctype, body = fetch(url, accept="application/json")
        if status != 200:
            log(f"{url}: {status} {ctype} {body[:300]!r}")
            return
        data = json.loads(body)
        results = data.get("results") or []
        if offset == 0:
            log(f"{base}: {data.get('total_count')} datasets")
        for ds in results:
            meta = (ds.get("metas") or {}).get("default") or {}
            fields = [f.get("name") for f in ds.get("fields") or []]
            log(f"  {ds.get('dataset_id')} | {meta.get('title')} | records={meta.get('records_count')}"
                f" | modified={str(meta.get('modified'))[:10]}")
            log(f"      fields: {', '.join(str(f) for f in fields)[:400]}")
            shown += 1
        if not results or len(results) < int(params["limit"]):
            break
        offset += len(results)


def cmd_odsrows(base: str, dataset: str, where: str | None, select: str | None,
                group_by: str | None, rows: int) -> None:
    """Records, a hundred a page, until ``rows`` are shown.

    The arguments are percent-decoded first: a quote cannot be passed through
    the workflow, so ``%22`` stands for it (``--where zone=%2213%22``).
    """
    base = base.rstrip("/")
    params = {"limit": str(min(rows, 100))}
    if where:
        params["where"] = urllib.parse.unquote(where)
    if select:
        params["select"] = urllib.parse.unquote(select)
    if group_by:
        params["group_by"] = urllib.parse.unquote(group_by)
    shown = offset = 0
    while shown < rows:
        params["offset"] = str(offset)
        url = (f"{base}/api/explore/v2.1/catalog/datasets/{dataset}/records?"
               + urllib.parse.urlencode(params))
        status, ctype, body = fetch(url, accept="application/json")
        if status != 200:
            log(f"{url}: {status} {ctype} {body[:500]!r}")
            return
        data = json.loads(body)
        results = data.get("results") or []
        if offset == 0:
            log(f"{dataset}: total_count={data.get('total_count')}")
        for rec in results:
            rec = {k: v for k, v in rec.items() if k not in ("geo_shape", "geo_point")}
            log("  " + json.dumps(rec, ensure_ascii=False)[:600])
            shown += 1
        if group_by or len(results) < int(params["limit"]):
            break
        offset += len(results)


def cmd_get(urls: list[str], grep: str | None, links: str | None, chars: int,
            context: int, most: int = 40, raw: bool = False) -> None:
    for url in urls:
        status, ctype, body = fetch(url)
        log(f"== {url}\n   {status} {ctype} {len(body):,} bytes")
        page = text_of(body, ctype)
        if links:
            pat = re.compile(links, re.I)
            seen = set()
            for m in re.finditer(r'(?is)<a\b[^>]*?href\s*=\s*["\']([^"\']+)["\'][^>]*>(.*?)</a>',
                                 page):
                href, label = html.unescape(m.group(1)), plain(m.group(2))[:120]
                full = urllib.parse.urljoin(url, href)
                if (pat.search(href) or pat.search(label)) and full not in seen:
                    seen.add(full)
                    log(f"   - {label!r} -> {full}")
            log(f"   {len(seen)} matching links")
        if grep:
            text = plain(page) if "html" in ctype.lower() and not raw else page
            hits = 0
            for m in re.finditer(grep, text, flags=re.I):
                a, b = max(0, m.start() - context), min(len(text), m.end() + context)
                log("   ~ " + " ".join(text[a:b].split()))
                hits += 1
                if hits >= most:
                    break
            log(f"   {hits} matches for {grep!r}")
        if not links and not grep:
            text = plain(page) if "html" in ctype.lower() else page
            log("   " + text[:chars].replace("\n", "\n   "))


def cmd_wayback(pattern: str, rows: int, grep: str | None = None) -> None:
    """Captures of a URL pattern; with ``grep``, only those whose address matches.

    A single-page app's captures are mostly its scripts and styles, so the
    filter runs over a larger listing than is printed.
    """
    url = CDX + "?" + urllib.parse.urlencode({
        "url": pattern, "output": "json", "limit": str(rows * 20 if grep else rows),
        "collapse": "urlkey", "filter": "statuscode:200"})
    status, ctype, body = fetch(url)
    if status != 200:
        log(f"{url}: {status} {body[:300]!r}")
        return
    rows_ = json.loads(body or b"[]")
    pat = re.compile(grep, re.I) if grep else None
    picked = [r for r in rows_[1:] if pat is None or pat.search(str(r[2]))]
    log(f"{pattern}: {max(0, len(rows_) - 1)} captures, {len(picked)} shown")
    for row in picked[:rows]:
        log("  " + " ".join(str(c) for c in row[1:5]))


def cmd_xlsx(urls: list[str], rows: int, cols: int, sheets: int, width: int) -> None:
    """A workbook's sheets and first rows, whatever its address says it is.

    Some offices serve .xlsx from an address with no extension (Kuwait's
    census tables are ``CensusData?st_id=4&handler=ExportExcel``).
    """
    import io

    import openpyxl
    for url in urls:
        status, ctype, body = fetch(url)
        log(f"== {url}\n   {status} {ctype} {len(body):,} bytes")
        if status != 200:
            log("   " + text_of(body, ctype)[:300])
            continue
        try:
            book = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        except Exception as err:  # noqa: BLE001 -- a probe reports
            log(f"   cannot open: {type(err).__name__}: {err}")
            continue
        for name in book.sheetnames[:sheets]:
            sheet = book[name]
            log(f"   [{name}] {sheet.max_row} x {sheet.max_column}")
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i >= rows:
                    break
                cells = ["" if c is None else str(c)[:width] for c in row[:cols]]
                if any(cells):
                    log(f"   {i + 1:>3} | " + " | ".join(cells))


def cmd_pdf(urls: list[str], pages: str | None, chars: int, grep: str | None,
            rows: int) -> None:
    import sys
    from pathlib import Path
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from probe_pdf import PAGE_BREAK, laid_out  # noqa: E402
    wanted = [int(p) for p in (pages or "1,2").split(",") if p.strip()]
    pat = re.compile(grep) if grep else None
    for url in urls:
        status, ctype, body = fetch(url)
        log(f"== {url}\n   {status} {ctype} {len(body):,} bytes")
        if status != 200 or not body.startswith(b"%PDF"):
            log("   " + text_of(body, ctype)[:300])
            continue
        try:
            text = laid_out(body)
        except Exception as err:  # noqa: BLE001 -- a probe reports
            log(f"   cannot read: {type(err).__name__}: {err}")
            continue
        sheets = text.split(PAGE_BREAK)
        log(f"   {len(sheets)} pages")
        if pat and chars == 0:
            hits = [(i, line) for i, t in enumerate(sheets) for line in t.splitlines()
                    if pat.search(line)]
            for i, line in hits[:rows]:
                log(f"   p{i + 1}: {line[:200]}")
            log(f"   {len(hits)} matching lines")
            continue
        picked = ([i for i, t in enumerate(sheets) if pat.search(t)][:rows] if pat
                  else [n - 1 for n in wanted if 0 < n <= len(sheets)])
        for i in picked:
            log(f"   -- page {i + 1}\n" + sheets[i][:chars])


def cmd_uscb(dataset: str, sheet: str, grep: str | None, rows: int,
             level: int | None, dictionary: str | None = None) -> None:
    import io

    import openpyxl

    from . import uscb
    url = uscb.workbook_url(dataset)
    status, ctype, body = fetch(url)
    log(f"{dataset}: {url} {status} {len(body):,} bytes")
    book = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    if dictionary:
        # The data dictionary says what each field is: its definition, the
        # office's own field name and the table it was read from.
        pat = re.compile(dictionary)
        for name in book.sheetnames:
            if name.strip().lower() != "data dictionary":
                continue
            for row in book[name].iter_rows(values_only=True):
                cells = ["" if v is None else " ".join(str(v).split()) for v in row]
                if cells and cells[0] and pat.search(cells[0]):
                    log(f"  {cells[0]}: " + " | ".join(c[:300] for c in cells[1:] if c))
    # A sheet name with a space cannot be one argument; '-' or '_' stands for it.
    wanted = re.sub(r"[-_]", " ", sheet).strip().lower()
    actual = next((s for s in book.sheetnames
                   if re.sub(r"[-_]", " ", s).strip().lower() == wanted), sheet)
    table = uscb.sheet_rows(book, actual)
    names, aliases = uscb.columns(table)
    pat = re.compile(grep or "^$")
    geo = [i for i, n in enumerate(names) if n in uscb.GEOGRAPHY and n not in
           ("GEO_CONCAT", "CNTRY_NAME", "USCBCMNT", "GENC_CODE", "FIPS_CODE")]
    picked = [i for i, n in enumerate(names) if n and n not in uscb.GEOGRAPHY and pat.search(n)]
    log(f"  {sheet}: {len(table) - 2} rows; fields: {', '.join(n for n in names if n)}"[:3000])
    for i in picked[:40]:
        log(f"    {names[i]} = {aliases[i]}")
    at = {n: i for i, n in enumerate(names)}
    shown = 0
    for row in table[2:]:
        if level is not None and uscb.number(row[at["ADM_LEVEL"]]) != level:
            continue
        cells = [str(row[i]) for i in geo] + [str(row[i]) for i in picked[:30]]
        log("  | " + " | ".join(cells)[:700])
        shown += 1
        if shown >= rows:
            break


ARABIC_LETTERS = re.compile(r"[؀-ۿ]")


def cmd_geonames(cc: str, iso3: str | None, level: str, classes: str, grep: str | None,
                 rows: int, zoom: int) -> None:
    """A country's GeoNames features, each with the drawn unit holding its point."""
    import io
    import zipfile

    status, ctype, body = fetch(f"https://download.geonames.org/export/dump/{cc}.zip")
    log(f"{cc}.zip: {status} {len(body):,} bytes")
    if status != 200:
        return
    with zipfile.ZipFile(io.BytesIO(body)) as zf:
        text = zf.read(f"{cc}.txt").decode("utf-8")
    wanted = set(classes.split(","))
    pat = re.compile(grep, re.I) if grep else None
    feats = []
    for line in text.splitlines():
        f = line.split("\t")
        if len(f) < 15 or f[6] not in wanted:
            continue
        names = [f[1], f[2]] + [n for n in f[3].split(",") if n]
        if pat and not any(pat.search(n) for n in names):
            continue
        arabic = [n for n in f[3].split(",") if ARABIC_LETTERS.search(n)]
        feats.append((f[0], f[1], arabic, float(f[4]), float(f[5]), f"{f[6]}.{f[7]}", f[14]))
    holder: dict[str, str] = {}
    if iso3:
        from shapely.geometry import Point

        from .sea_common import drawn, polygons
        shapes = polygons(level, iso3, zoom)
        label = {u["id"]: u["name"] for u in drawn(iso3, level)}
        for gid, _n, _a, lat, lon, _c, _p in feats:
            hits = [label.get(sid, sid) for sid, g in shapes.items() if g.contains(Point(lon, lat))]
            holder[gid] = " + ".join(hits) or "-"
    log(f"{len(feats)} features of classes {classes}" + (f" matching {grep!r}" if grep else ""))
    for gid, name, arabic, lat, lon, code, pop in sorted(feats, key=lambda r: r[1])[:rows]:
        log(f"  {gid} | {name} | {' / '.join(arabic[:3])} | {lat:.5f},{lon:.5f} | {code} | "
            f"pop {pop} | in {holder.get(gid, '?')}")


WDQS = "https://query.wikidata.org/sparql"
# Places a person lives in: settlements, neighbourhoods, suburbs, districts.
SETTLEMENT_CLASSES = ("wd:Q486972", "wd:Q123705", "wd:Q188509", "wd:Q2983893",
                      "wd:Q1549591", "wd:Q15284")


def cmd_wdpoints(country: str, iso3: str | None, level: str, grep: str | None, rows: int,
                 zoom: int) -> None:
    """Wikidata's settlements and neighbourhoods in one country, with the drawn unit
    holding each point. ``country`` is the country's item (Q817 for Kuwait)."""
    query = f"""
SELECT ?item ?ar ?en ?coord WHERE {{
  ?item wdt:P17 wd:{country} ; wdt:P625 ?coord ; wdt:P31/wdt:P279* ?cls .
  VALUES ?cls {{ {' '.join(SETTLEMENT_CLASSES)} }}
  OPTIONAL {{ ?item rdfs:label ?ar FILTER(LANG(?ar) = "ar") }}
  OPTIONAL {{ ?item rdfs:label ?en FILTER(LANG(?en) = "en") }}
}}"""
    url = WDQS + "?" + urllib.parse.urlencode({"format": "json", "query": query})
    status, ctype, body = fetch(url, accept="application/sparql-results+json")
    log(f"WDQS {country}: {status} {len(body):,} bytes")
    if status != 200:
        log("  " + text_of(body, ctype)[:400])
        return
    seen: dict[str, tuple[str, str, float, float]] = {}
    for b in json.loads(body)["results"]["bindings"]:
        m = re.match(r"Point\(([-\d.]+) ([-\d.]+)\)", b["coord"]["value"])
        if not m:
            continue
        qid = b["item"]["value"].rsplit("/", 1)[-1]
        seen.setdefault(qid, (b.get("ar", {}).get("value", ""), b.get("en", {}).get("value", ""),
                              float(m.group(2)), float(m.group(1))))
    pat = re.compile(grep, re.I) if grep else None
    items = {q: v for q, v in seen.items() if not pat or pat.search(v[0]) or pat.search(v[1])}
    holder: dict[str, str] = {}
    if iso3:
        from shapely.geometry import Point

        from .sea_common import drawn, polygons
        shapes = polygons(level, iso3, zoom)
        label = {u["id"]: u["name"] for u in drawn(iso3, level)}
        for qid, (_ar, _en, lat, lon) in items.items():
            hits = [label.get(sid, sid) for sid, g in shapes.items() if g.contains(Point(lon, lat))]
            holder[qid] = " + ".join(hits) or "-"
    log(f"{len(items)} items")
    for qid, (ar, en, lat, lon) in sorted(items.items(), key=lambda kv: kv[1][1] or kv[1][0])[:rows]:
        log(f"  {qid} | {en} | {ar} | {lat:.5f},{lon:.5f} | in {holder.get(qid, '?')}")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("cmd", choices=["ods", "odsrows", "get", "wayback", "uscb", "xlsx", "pdf",
                                    "geonames", "wdpoints"])
    ap.add_argument("target", nargs="+")
    ap.add_argument("--search")
    ap.add_argument("--where")
    ap.add_argument("--select")
    ap.add_argument("--group-by")
    ap.add_argument("--rows", type=int, default=60)
    ap.add_argument("--grep")
    ap.add_argument("--links")
    ap.add_argument("--chars", type=int, default=1500)
    ap.add_argument("--context", type=int, default=150)
    ap.add_argument("--level", type=int)
    ap.add_argument("--cols", type=int, default=16)
    ap.add_argument("--sheets", type=int, default=3)
    ap.add_argument("--width", type=int, default=24)
    ap.add_argument("--raw", action="store_true", help="grep the page as sent, tags and all")
    ap.add_argument("--pages")
    ap.add_argument("--dictionary")
    ap.add_argument("--iso")
    ap.add_argument("--unit-level", default="admin2")
    ap.add_argument("--classes", default="P")
    ap.add_argument("--zoom", type=int, default=8)
    args = ap.parse_args()
    if args.cmd == "uscb":
        cmd_uscb(args.target[0], args.target[1], args.grep, args.rows, args.level,
                 args.dictionary)
        return 0
    if args.cmd == "geonames":
        cmd_geonames(args.target[0], args.iso, args.unit_level, args.classes, args.grep,
                     args.rows, args.zoom)
        return 0
    if args.cmd == "wdpoints":
        cmd_wdpoints(args.target[0], args.iso, args.unit_level, args.grep, args.rows, args.zoom)
        return 0
    if args.cmd == "ods":
        for base in args.target:
            cmd_ods(base, args.search, args.rows)
    elif args.cmd == "odsrows":
        cmd_odsrows(args.target[0], args.target[1], args.where, args.select,
                    args.group_by, args.rows)
    elif args.cmd == "get":
        cmd_get(args.target, args.grep, args.links, args.chars, args.context, max(40, args.rows),
                args.raw)
    elif args.cmd == "wayback":
        for pattern in args.target:
            cmd_wayback(pattern, args.rows, args.grep)
    elif args.cmd == "xlsx":
        cmd_xlsx(args.target, args.rows, args.cols, args.sheets, args.width)
    elif args.cmd == "pdf":
        cmd_pdf(args.target, args.pages, args.chars, args.grep, args.rows)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
