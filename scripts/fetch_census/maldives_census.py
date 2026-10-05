#!/usr/bin/env python3
"""Maldives -- Census 2022 by atoll: head count, sex ratio, median age, nationality.

The Maldives Bureau of Statistics publishes its 2022 census atoll by atoll in
one workbook, the *Atoll Level Indicator Sheet - Population*: for Malé and each
of the twenty administrative atolls, the resident population by sex, split into
Maldivians and foreigners, and the median age of the resident population, of
its Maldivians and of its foreigners. That is three of the map's fields for
every atoll, where the map had a head count from Wikidata for some atolls and
nothing else.

**Which block.** The sheet prints the atolls twice. First under
"Administrative Islands & Non-administrative Islands" -- every locality,
the resorts and industrial islands among them -- and then under
"Administrative Islands" alone, the inhabited islands. The first is everyone
counted in the atoll and is what this reads: a resort's staff live on the
resort, inside the atoll the map draws. The second is read only to check the
sheet against itself: the inhabited islands of each atoll, plus the resorts
and the industrial islands the sheet totals nationally, make the first block.

**Checks**, all exact, before anything is written:

* every row's males and females make its total, and its Maldivians and
  foreigners make its residents, for both sexes;
* the twenty atolls of each block are the twenty codes, once each, and add up
  to the sheet's own Atolls and Administrative Islands rows;
* Malé and the Atolls make the Republic; the Administrative and the
  Non-administrative Islands make the Atolls; resorts and industrial islands
  make the Non-administrative;
* every median is a number of years a census could print.

**The median age is the Bureau's own**, the resident population's, as it
prints it: in whole years, or half years where the middle falls between two.
Nothing here recomputes it.

**Binding.** Each row carries its atoll's code -- "North Thiladhunmathi (HA)".
The map's first level names atolls by their letter names (Haa Alif), its
second by their formal names (North Thiladhunmathe), and both are matched
through the code, never by a similar spelling. Where a second-level polygon
lies inside a first-level one, the two must be the same atoll, or the run
stops. The map draws thirteen atolls at the first level and nineteen at the
second; an atoll without a polygon is named in the log and placed nowhere.
Gnaviyani (Fuvahmulah) has none at either level, so its 9,177 people are not
put on Addu, whose polygon lies 37 km away.

**Malé.** The boundary file has two second-level polygons labelled Male'. One
covers the whole of Kaafu -- from South Malé Atoll's southern islands to
Kaashidhoo -- and Malé island with it, as GeoNames' point for the capital
shows. Its people are Kaafu's and Malé's together, so it takes the sum of the
two rows, and a sex ratio from the summed sexes; the census prints no median
for the two together, and two medians do not make a third, so that one field
is left with its reason. The other polygon is a fragment some thirty metres
across on Malé's north shore; the city's people are on the first polygon, and
writing them here as well would count them twice. At the first level, Kaafu
is the administrative atoll, which does not include Malé: it takes Kaafu's
own row, and says so.

**Nationality** -- Maldivian or foreigner, the census's only question about
who a person is -- is written to a file of its own,
``maldives_nationality.json``, as the ethnicity field under
``ethnicity_basis: "nationality"`` (the owner's decision of 19 September
2026, as for Japan, Korea and Pakistan). It needs the Maldives' ethnicity line
in ``common.NOT_COLLECTED_POLICY`` to go, which is proposed rather than done
here; until then the country's declaration stands beside it.

Usage:
    python -m scripts.fetch_census.maldives_census
"""

from __future__ import annotations

import argparse
import io
import re
from typing import Any

from ._shared import (NOT_AVAILABLE, PROCESSED, gap, http_get, log, measure,
                      record, shares, write_json)
from .south_asia_common import fold, load_units, males_per_100

URL = ("https://statisticsmaldives.gov.mv/mbs/wp-content/uploads/2023/09/"
       "Atoll-Level-Indicator-Sheet-Population.xlsx")
YEAR = 2022
SOURCE = ("Maldives Bureau of Statistics, Population and Housing Census 2022, "
          "Atoll Level Indicator Sheet - Population")
LICENCE = "Maldives Bureau of Statistics publication, reuse with attribution"
OUT = "maldives_census.json"
NATIONALITY_OUT = "maldives_nationality.json"

# The census's atoll code -> the atoll's letter name, which the map's first
# level uses, and the formal name the map's second level draws it under (as
# the boundary file spells it). Gnaviyani has no second-level polygon.
ATOLLS: dict[str, tuple[str, str | None]] = {
    "HA": ("Haa Alif", "North Thiladhunmathe"),
    "HDh": ("Haa Dhaalu", "South Thiladhunmathe"),
    "Sh": ("Shaviyani", "North Miladhunmadulu"),
    "N": ("Noonu", "South Miladhunmadulu"),
    "R": ("Raa", "North Maalhosmadulu"),
    "B": ("Baa", "South Maalhosmadulu"),
    "Lh": ("Lhaviyani", "Faadhippolhu"),
    "K": ("Kaafu", "Male'"),
    "AA": ("Alif Alif", "North Ari"),
    "ADh": ("Alif Dhaal", "South Ari"),
    "V": ("Vaavu", "Felidhoo"),
    "M": ("Meemu", "Mulaku"),
    "F": ("Faafu", "North Nilandhoo"),
    "Dh": ("Dhaalu", "South Nilandhoo"),
    "Th": ("Thaa", "Kolhumadulu"),
    "L": ("Laamu", "Hadhdhunmathee"),
    "GA": ("Gaafu Alif", "North Huvadhoo"),
    "GDh": ("Gaafu Dhaalu", "South Huvadhoo"),
    "Gn": ("Gnaviyani", None),
    "S": ("Addu", "Addoo"),
}
# Where the two Male' polygons are told apart: the one that is Kaafu covers
# both of these points; the other must be a fragment at Malé itself.
MALE_CITY = (73.5092, 4.1752)       # GeoNames 1282027, Malé (PPLC)
KAASHIDHOO = (73.4630, 4.9570)      # Kaafu's northernmost inhabited island
FRAGMENT = 0.02                     # degrees: how far from Malé a fragment may lie

ATOLL_ROW = re.compile(r"^(?P<name>.+?)\s*\((?P<code>[A-Za-z]+)\)$")
FIELDS = ("residents", "maldivians", "foreigners")


# ---------------------------------------------------------------------------
# Reading the sheet
# ---------------------------------------------------------------------------

def sheet_rows(blob: bytes) -> list[list[Any]]:
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    sheet = book["Atoll"] if "Atoll" in book.sheetnames else book.worksheets[0]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def text(cell: Any) -> str:
    return re.sub(r"\s+", " ", str(cell)).strip() if cell is not None else ""


def columns(rows: list[list[Any]]) -> tuple[int, dict[str, int]]:
    """The header row and the first column of each group this reads.

    Found by the sheet's own words, so a column added or moved shows as a
    refusal rather than as figures read from the wrong place.
    """
    for i, row in enumerate(rows[:12]):
        labels = [text(c) for c in row]
        if "Resident Population" not in labels:
            continue
        sub = [text(c) for c in rows[i + 1]] if i + 1 < len(rows) else []
        found = {
            "residents": labels.index("Resident Population"),
            "maldivians": labels.index("Resident Maldivians")
            if "Resident Maldivians" in labels else -1,
            "foreigners": labels.index("Resident Foreigners")
            if "Resident Foreigners" in labels else -1,
            "median": labels.index("Median Age") if "Median Age" in labels else -1,
        }
        if -1 in found.values():
            raise SystemExit(f"maldives: the header row has {labels}; one of "
                             "Resident Population, Resident Maldivians, Resident "
                             "Foreigners or Median Age is missing")
        for key in FIELDS:
            at = found[key]
            if sub[at:at + 3] != ["Both sexes", "Male", "Female"]:
                raise SystemExit(f"maldives: {key} is not Both sexes, Male, Female "
                                 f"but {sub[at:at + 3]}")
        if not sub[found["median"]].startswith("Resident Population"):
            raise SystemExit("maldives: the first Median Age column is "
                             f"{sub[found['median']]!r}, not the resident population's")
        return i, found
    raise SystemExit("maldives: no header row naming Resident Population")


def figures(row: list[Any], cols: dict[str, int], label: str) -> dict[str, Any]:
    out: dict[str, Any] = {"label": label}
    for key in FIELDS:
        at = cols[key]
        values = row[at:at + 3]
        if not all(isinstance(v, (int, float)) and float(v).is_integer() for v in values):
            raise SystemExit(f"maldives: {label}: {key} is {values}, not three counts")
        out[key] = tuple(int(v) for v in values)
    median = row[cols["median"]]
    out["median"] = float(median) if isinstance(median, (int, float)) else None
    return out


def read(rows: list[list[Any]]) -> dict[str, Any]:
    """The sheet's rows by what they are: totals, and the atolls of each block."""
    head, cols = columns(rows)
    totals: dict[str, dict[str, Any]] = {}
    blocks: dict[str, dict[str, dict[str, Any]]] = {"all": {}, "inhabited": {}}
    block: str | None = None
    width = max(len(rows[head]), len(rows[head + 1]))
    for row in rows[head + 2:]:
        row = list(row) + [None] * (width - len(row))
        label = text(row[1])
        if not label:
            continue
        has_figures = isinstance(row[cols["residents"]], (int, float))
        if label.startswith("Administrative Islands & Non"):
            block = "all"
            continue
        if label == "Administrative Islands":
            block = "inhabited"
            if has_figures:
                totals["Administrative Islands"] = figures(row, cols, label)
            continue
        if label in ("Non Administrative Islands", "Resorts", "Industrial Islands"):
            block = None
            if has_figures and label not in totals:
                totals[label] = figures(row, cols, label)
            continue
        if label == "Republic" or label == "Atolls" or label.startswith("Male'"):
            if has_figures:
                totals["Male'" if label.startswith("Male'") else label] = (
                    figures(row, cols, label))
            continue
        match = ATOLL_ROW.match(label)
        if match and block and has_figures:
            code = match["code"]
            if code not in ATOLLS:
                raise SystemExit(f"maldives: {label!r} carries a code this reader "
                                 "does not know")
            if code in blocks[block]:
                raise SystemExit(f"maldives: {code} appears twice in the {block} block")
            blocks[block][code] = figures(row, cols, label)
            blocks[block][code]["name"] = match["name"]
    return {"totals": totals, "all": blocks["all"], "inhabited": blocks["inhabited"]}


def add(rows: list[dict[str, Any]]) -> dict[str, tuple[int, int, int]]:
    return {key: tuple(sum(r[key][i] for r in rows) for i in range(3)) for key in FIELDS}


def check(sheet: dict[str, Any]) -> None:
    totals, every, inhabited = sheet["totals"], sheet["all"], sheet["inhabited"]
    wanted = {"Republic", "Male'", "Atolls", "Administrative Islands",
              "Non Administrative Islands", "Resorts", "Industrial Islands"}
    if set(totals) != wanted:
        raise SystemExit(f"maldives: totals rows {sorted(totals)}, not {sorted(wanted)}")
    for name, block in (("all localities", every), ("inhabited islands", inhabited)):
        if set(block) != set(ATOLLS):
            raise SystemExit(f"maldives: the {name} block has {sorted(block)}, not "
                             f"the twenty atolls")
    for row in [*totals.values(), *every.values(), *inhabited.values()]:
        for key in FIELDS:
            both, male, female = row[key]
            if male + female != both:
                raise SystemExit(f"maldives: {row['label']}: {key} {male:,} + "
                                 f"{female:,} is not {both:,}")
        for i in range(3):
            if row["maldivians"][i] + row["foreigners"][i] != row["residents"][i]:
                raise SystemExit(f"maldives: {row['label']}: Maldivians and "
                                 "foreigners do not make the residents")
    def same(label: str, got: dict[str, tuple], want: dict[str, Any]) -> None:
        for key in FIELDS:
            if tuple(got[key]) != tuple(want[key]):
                raise SystemExit(f"maldives: {label}: {key} {got[key]} against "
                                 f"{want[key]}")
    same("the twenty atolls, all localities, against Atolls",
         add(list(every.values())), totals["Atolls"])
    same("the twenty atolls, inhabited islands, against Administrative Islands",
         add(list(inhabited.values())), totals["Administrative Islands"])
    same("Male' and the Atolls against the Republic",
         add([totals["Male'"], totals["Atolls"]]), totals["Republic"])
    same("the inhabited and the other islands against the Atolls",
         add([totals["Administrative Islands"], totals["Non Administrative Islands"]]),
         totals["Atolls"])
    same("resorts and industrial islands against the Non-administrative Islands",
         add([totals["Resorts"], totals["Industrial Islands"]]),
         totals["Non Administrative Islands"])
    for code in ATOLLS:
        if inhabited[code]["residents"][0] > every[code]["residents"][0]:
            raise SystemExit(f"maldives: {code}'s inhabited islands count more "
                             "people than all its localities")
    for row in [totals["Male'"], *every.values()]:
        median = row["median"]
        if median is None or not 10 <= median <= 60 or (2 * median) % 1:
            raise SystemExit(f"maldives: {row['label']}: median age {median!r} is "
                             "not a whole or half year a census would print")


# ---------------------------------------------------------------------------
# The map's units
# ---------------------------------------------------------------------------

def contains(bbox: list[float], point: tuple[float, float]) -> bool:
    return bbox[0] <= point[0] <= bbox[2] and bbox[1] <= point[1] <= bbox[3]


def bind_units(units1: list[dict[str, Any]], units2: list[dict[str, Any]]
               ) -> tuple[dict[str, dict], dict[str, dict], dict[str, Any] | None]:
    """Code -> first-level unit, code -> second-level unit, and Malé's fragment."""
    by_name1: dict[str, list[dict]] = {}
    for unit in units1:
        by_name1.setdefault(fold(unit.get("name")), []).append(unit)
    by_name2: dict[str, list[dict]] = {}
    for unit in units2:
        by_name2.setdefault(fold(unit.get("name")), []).append(unit)
    first: dict[str, dict] = {}
    second: dict[str, dict] = {}
    fragment = None
    for code, (letter, formal) in ATOLLS.items():
        hits = by_name1.get(fold(letter), [])
        if len(hits) > 1:
            raise SystemExit(f"maldives: {len(hits)} first-level units are called {letter}")
        if hits:
            first[code] = hits[0]
        if formal is None:
            continue
        hits = by_name2.get(fold(formal), [])
        if code == "K":
            kaafu = [u for u in hits if contains(u["bbox"], MALE_CITY)
                     and contains(u["bbox"], KAASHIDHOO)]
            rest = [u for u in hits if u not in kaafu]
            if len(kaafu) != 1 or len(rest) > 1:
                raise SystemExit(f"maldives: the Male' polygons are not one Kaafu "
                                 f"and at most one fragment: {[u['bbox'] for u in hits]}")
            if rest:
                b = rest[0]["bbox"]
                if (b[2] - b[0] > FRAGMENT or b[3] - b[1] > FRAGMENT
                        or abs(b[0] - MALE_CITY[0]) > FRAGMENT
                        or abs(b[1] - MALE_CITY[1]) > FRAGMENT):
                    raise SystemExit(f"maldives: the second Male' polygon {b} is not "
                                     "a fragment at Malé; it needs a decision")
                fragment = rest[0]
            hits = kaafu
        if len(hits) != 1:
            raise SystemExit(f"maldives: {len(hits)} second-level units are called "
                             f"{formal}")
        second[code] = hits[0]
    # A second-level polygon drawn inside a first-level one must be the same atoll.
    ids1 = {u["id"]: code for code, u in first.items()}
    for code, unit in second.items():
        parent = unit.get("parent")
        if parent in ids1 and ids1[parent] != code:
            raise SystemExit(f"maldives: {unit['name']} ({code}) is drawn inside "
                             f"{first[ids1[parent]]['name']} ({ids1[parent]})")
        if parent not in ids1 and parent != "MDV" and parent in {
                u["id"] for u in units1}:
            raise SystemExit(f"maldives: {unit['name']} has parent {parent}")
    return first, second, fragment


# ---------------------------------------------------------------------------
# Fields
# ---------------------------------------------------------------------------

def cite(fields: list[str]) -> list[dict[str, Any]]:
    return [{"field": field, "name": SOURCE, "url": URL, "year": YEAR,
             "license": LICENCE} for field in fields]


def fields(row: dict[str, Any], *, extra: str = "", median: bool = True
           ) -> dict[str, Any]:
    both, male, female = row["residents"]
    f_both, f_male, f_female = row["foreigners"]
    out: dict[str, Any] = {
        "population": measure(both, year=YEAR, source=SOURCE),
        "population_note": (
            f"Census 2022: everyone resident here on census night -- on the "
            f"inhabited islands, the resorts and the industrial islands alike -- "
            f"{f_both:,} of them foreign residents.{extra}"),
        "sex_ratio": measure(males_per_100(male, female), unit="males_per_100_females",
                             year=YEAR, source=SOURCE),
        "sex_ratio_note": (
            f"Census 2022: {male:,} males and {female:,} females resident. Foreign "
            f"residents are counted where they live: {f_male:,} of the males and "
            f"{f_female:,} of the females here.{extra}"),
    }
    kept = ["population", "sex_ratio"]
    if median:
        out["median_age"] = measure(row["median"], unit="years", year=YEAR, source=SOURCE)
        out["median_age_note"] = (
            "Census 2022: the median age of the resident population, Maldivians "
            "and foreigners together, as the Bureau prints it -- in whole years, "
            "or half years where the middle falls between two." + extra)
        kept.append("median_age")
    out["sources"] = cite(kept)
    return out


def nationality(row: dict[str, Any], *, extra: str = "") -> dict[str, Any]:
    maldivians, foreigners = row["maldivians"][0], row["foreigners"][0]
    return {
        "ethnicity": shares({"Maldivian": maldivians, "Foreigner": foreigners}),
        "ethnicity_year": YEAR,
        "ethnicity_basis": "nationality",
        "ethnicity_note": (
            "Nationality, not ethnicity: the Maldives census asks whether a person "
            "is a Maldivian or a foreigner and nothing about ethnicity, and "
            "publishes no breakdown of foreigners by country in its atoll sheet. "
            f"Census 2022 counts {maldivians:,} Maldivians and {foreigners:,} "
            f"foreign residents here.{extra}"),
        "sources": cite(["ethnicity"]),
    }


def summed(a: dict[str, Any], b: dict[str, Any], label: str) -> dict[str, Any]:
    out = {key: tuple(x + y for x, y in zip(a[key], b[key])) for key in FIELDS}
    out["label"] = label
    out["median"] = None
    return out


# ---------------------------------------------------------------------------

def build(sheet: dict[str, Any], units1: list[dict[str, Any]],
          units2: list[dict[str, Any]]
          ) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    every, male_city = sheet["all"], sheet["totals"]["Male'"]
    first, second, fragment = bind_units(units1, units2)
    main: list[dict[str, Any]] = []
    nat: list[dict[str, Any]] = []
    kaafu_note = (f" Kaafu is the administrative atoll, which does not include "
                  f"Malé: the city's {male_city['residents'][0]:,} people are "
                  "counted apart and are not in this figure.")
    for code, unit in sorted(first.items()):
        extra = kaafu_note if code == "K" else ""
        common = dict(level="admin1", parent="MDV", country="MDV",
                      match_by="shape_id", shape_id=unit["id"])
        main.append(record(f"MDV-{code}", unit["name"], **common,
                           **fields(every[code], extra=extra)))
        nat.append(record(f"MDV-NAT-{code}", unit["name"], **common,
                          **nationality(every[code], extra=extra)))
    for code, unit in sorted(second.items()):
        common = dict(level="admin2", parent="MDV", country="MDV",
                      match_by="shape_id", shape_id=unit["id"])
        if code == "K":
            both = summed(every["K"], male_city, "Male' and Kaafu")
            extra = (f" The boundary file draws Kaafu and Malé island as this one "
                     f"polygon, so the figure is Kaafu's "
                     f"{every['K']['residents'][0]:,} and Malé's "
                     f"{male_city['residents'][0]:,} (with Villimalé and "
                     f"Hulhumalé) together.")
            body = fields(both, extra=extra, median=False)
            body["median_age"] = gap(NOT_AVAILABLE, (
                f"The census prints Malé's median age ({male_city['median']:g}) "
                f"and Kaafu's ({every['K']['median']:g}) apart and none for the "
                "two together, which is what this polygon draws; two medians do "
                "not make a third."))
            main.append(record(f"MDV-ATOLL-{code}", unit["name"], **common, **body))
            nat.append(record(f"MDV-NAT-ATOLL-{code}", unit["name"], **common,
                              **nationality(both, extra=extra)))
            continue
        main.append(record(f"MDV-ATOLL-{code}", unit["name"], **common,
                           **fields(every[code])))
        nat.append(record(f"MDV-NAT-ATOLL-{code}", unit["name"], **common,
                          **nationality(every[code])))
    if fragment is not None:
        why = ("This second polygon labelled Male' is a fragment some thirty "
               "metres across on Malé's north shore. The city's people are counted "
               "on the other Male' polygon, which covers the island, and putting "
               "them here as well would count them twice.")
        main.append(record("MDV-ATOLL-male-fragment", fragment["name"],
                           level="admin2", parent="MDV", country="MDV",
                           match_by="shape_id", shape_id=fragment["id"],
                           population=gap(NOT_AVAILABLE, why),
                           median_age=gap(NOT_AVAILABLE, why),
                           sex_ratio=gap(NOT_AVAILABLE, why)))
    log(f"  first level: {len(first)} atolls bound -- "
        + ", ".join(f"{c} {u['name']}" for c, u in sorted(first.items())))
    log("    not drawn at the first level: "
        + ", ".join(f"{c} {ATOLLS[c][0]}" for c in ATOLLS if c not in first)
        + "; and Malé")
    log(f"  second level: {len(second)} polygons bound"
        + (", and Malé's fragment given its reason" if fragment else ""))
    log("    not drawn at the second level: "
        + ", ".join(f"{c} {ATOLLS[c][0]}" for c in ATOLLS if c not in second))
    return main, nat


def main() -> int:
    argparse.ArgumentParser(description=__doc__,
                            formatter_class=argparse.RawDescriptionHelpFormatter
                            ).parse_args()
    log(f"maldives_census: {SOURCE}")
    blob = http_get(URL, binary=True)
    log(f"  {len(blob):,} bytes from {URL}")
    sheet = read(sheet_rows(blob))
    check(sheet)
    totals = sheet["totals"]
    republic, city, atolls = (totals[k]["residents"][0]
                              for k in ("Republic", "Male'", "Atolls"))
    log(f"  Republic {republic:,} = Malé {city:,} + Atolls {atolls:,}; the twenty "
        "atolls of both blocks, resorts and industrial islands all reconcile")
    main_records, nat_records = build(sheet, load_units("MDV", "admin1"),
                                      load_units("MDV", "admin2"))
    write_json(PROCESSED / OUT, main_records)
    write_json(PROCESSED / NATIONALITY_OUT, nat_records)
    log(f"  wrote {len(main_records)} records to {OUT} and {len(nat_records)} "
        f"to {NATIONALITY_OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
