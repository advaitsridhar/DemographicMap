#!/usr/bin/env python3
"""The Netherlands: population and sex ratio by gemeente, 1 January 2022.

CBS's "Kerncijfers wijken en buurten 2022" (KWB 2022) counts the registered
population (BRP) of every gemeente, wijk and buurt on 1 January 2022, with
men and women, and CBS publishes it as a workbook on download.cbs.nl. Its
gemeenten are those of 1 January 2022: 345.

**Vintage.** The map draws 344 gemeenten: the 2022 division after Weesp
joined Amsterdam on 24 March 2022, and before Brielle, Hellevoetsluis and
Westvoorne became Voorne aan Zee in 2023. So Amsterdam's polygon is given
Amsterdam and Weesp summed, which is what it holds, and every other gemeente
is its own polygon.

**Median age.** Not written. CBS publishes age by gemeente in single years in
StatLine table 03759ned, but StatLine's open-data service (opendata.cbs.nl)
refused every connection from the fetch runner, and the key-figures workbook
gives five broad age groups only (0-14, 15-24, 25-44, 45-64, 65+), from which a
median would be a guess across a twenty-year band.

Usage:
    python -m scripts.fetch_census.netherlands_gemeente
"""

from __future__ import annotations

import argparse
import io
from typing import Any

from ._shared import NOT_AVAILABLE, PROCESSED, gap, http_get, log, measure, record, write_json
from .central_ages import SEX_RATIO_UNIT, check_sum, fold, report_unbound, sex_ratio, units

WORKBOOK = "https://download.cbs.nl/regionale-kaarten/kwb-2022.xlsx"
PAGE = "https://www.cbs.nl/nl-nl/maatwerk/2023/14/kerncijfers-wijken-en-buurten-2022"
SOURCE = "CBS, Kerncijfers wijken en buurten 2022 (population register, 1 January 2022)"
LICENCE = "CC BY 4.0 (CBS)"
OUT = PROCESSED / "netherlands_gemeente.json"
YEAR = 2022
NATIONAL = 17_590_672             # CBS, population of the Netherlands on 1 January 2022
EXPECTED = 345
MEDIAN_GAP = ("CBS publishes age by gemeente in single years (StatLine 03759ned), but its open-data "
              "service refused every connection from the fetch runner, and its key figures by "
              "gemeente give five broad age groups only, too coarse for a median.")
MERGED_INTO = {"Weesp": "Amsterdam"}          # 24 March 2022, before the map's division
# CBS's spelling -> the boundary file's, where folding does not reach it.
MAP_NAMES = {"Hengelo": "Hengelo (O)", "Bergen (L.)": "Bergen (L)", "Bergen (NH.)": "Bergen (NH)"}


def number(cell: Any) -> float | None:
    if cell is None:
        return None
    if isinstance(cell, (int, float)):
        return float(cell)
    text = str(cell).strip().replace(" ", "")
    if text in ("", ".", "-", "x"):
        return None
    try:
        return float(text)
    except ValueError:
        return None


def read() -> tuple[dict[str, dict[str, Any]], float | None]:
    """{gemeente code: {name, total, men, women}} and the national total."""
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(http_get(WORKBOOK, binary=True, timeout=300)),
                                  read_only=True, data_only=True)
    sheet = book[book.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header: list[str] = []
    for row in rows:
        cells = [str(c).strip() if c is not None else "" for c in row]
        if "a_inw" in cells and "recs" in cells:
            header = cells
            break
    if not header:
        raise SystemExit(f"netherlands_gemeente: no header with a_inw and recs in {book.sheetnames}")
    code_col = next(i for i, c in enumerate(header) if c.startswith("gwb_code"))
    at = {k: header.index(k) for k in ("regio", "recs", "a_inw", "a_man", "a_vrouw")}
    log(f"  columns: {header[:12]} ...")
    out: dict[str, dict[str, Any]] = {}
    national = None
    for row in rows:
        if not row or len(row) <= max(at.values()) or row[at["recs"]] is None:
            continue
        kind = str(row[at["recs"]]).strip()
        if kind == "Land":
            national = number(row[at["a_inw"]])
        if kind != "Gemeente":
            continue
        code = str(row[code_col]).strip()
        out[code] = {"name": str(row[at["regio"]]).strip(), "total": number(row[at["a_inw"]]),
                     "men": number(row[at["a_man"]]), "women": number(row[at["a_vrouw"]])}
    return out, national


def build() -> list[dict[str, Any]]:
    log("netherlands_gemeente: CBS Kerncijfers wijken en buurten 2022, by gemeente")
    gemeenten, national = read()
    log(f"  {len(gemeenten)} gemeenten; Nederland {national or 0:,.0f}")
    if len(gemeenten) != EXPECTED:
        raise SystemExit(f"netherlands_gemeente: {len(gemeenten)} gemeenten, not {EXPECTED}")
    missing = [g["name"] for g in gemeenten.values() if None in (g["total"], g["men"], g["women"])]
    if missing:
        raise SystemExit(f"netherlands_gemeente: no count for {missing}")
    check_sum((g["total"] for g in gemeenten.values()), NATIONAL, "gemeenten against the Netherlands")
    if national is not None:
        check_sum([national], NATIONAL, "the workbook's own national row")
    worst = max(abs(g["men"] + g["women"] - g["total"]) for g in gemeenten.values())
    log(f"  men + women against the total: largest difference {worst:,.0f}")
    if worst > 0:
        raise SystemExit("netherlands_gemeente: men and women do not make a gemeente's total")

    by_name = {g["name"]: g for g in gemeenten.values()}
    for part, whole in MERGED_INTO.items():
        into = by_name[whole]
        for key in ("total", "men", "women"):
            into[key] += by_name[part][key]
        into["merged"] = part
        log(f"  {whole}: {part} added, {into['total']:,.0f}")

    shapes = units("NLD", "admin2")
    by_key: dict[str, list[dict[str, Any]]] = {}
    for shape in shapes:
        by_key.setdefault(fold(shape["name"]), []).append(shape)
    records, unbound, used = [], [], set()
    for code, g in sorted(gemeenten.items()):
        if g["name"] in MERGED_INTO:
            continue
        drawn = MAP_NAMES.get(g["name"], g["name"])
        hits = [s for s in by_key.get(fold(drawn), []) if s["id"] not in used]
        if len(hits) != 1:
            unbound.append(f"{g['name']} ({code})")
            continue
        shape = hits[0]
        used.add(shape["id"])
        merged = g.get("merged")
        what = f"{g['name']} and {merged}, which joined it on 24 March 2022," if merged else g["name"]
        records.append(record(
            f"NLD-KWB2022-{code}", g["name"], level="admin2", parent=shape["parent"], country="NLD",
            codes={"cbs_gemeente": code}, aliases=[shape["name"]] if shape["name"] != g["name"] else [],
            match_by="shape_id", shape_id=shape["id"],
            sources=[{"field": "population/sex_ratio", "name": SOURCE, "url": PAGE,
                      "license": LICENCE, "year": YEAR}],
            population=measure(int(g["total"]), year=YEAR, source=SOURCE),
            sex_ratio=measure(sex_ratio(g["men"], g["women"]), unit=SEX_RATIO_UNIT, year=YEAR,
                              source=SOURCE),
            sex_ratio_note=(f"Males per 100 females registered in {what} on 1 January 2022 "
                            f"(CBS key figures by gemeente)."),
            median_age=gap(NOT_AVAILABLE, MEDIAN_GAP)))
    left = [s["name"] for s in shapes if s["id"] not in used]
    report_unbound("netherlands_gemeente", unbound, left)
    if unbound or left:
        raise SystemExit(f"netherlands_gemeente: {len(unbound)} gemeenten and {len(left)} polygons "
                         f"unpaired")
    return records


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.parse_args()
    records = build()
    write_json(OUT, records)
    log(f"  wrote {OUT.name}: {len(records)} records")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
