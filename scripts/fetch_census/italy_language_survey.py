#!/usr/bin/env python3
"""Italy: the language used in the family, by region and by macro-region --
ISTAT's 2015 survey on the use of Italian, dialects and other languages.

Italy's census has never asked a language question of the whole country (the
South Tyrolean language-group declaration is a separate register), so the
language field of the twenty regions and five macro-regions carried only the
European Social Survey where a region reached 100 respondents -- eleven of the
twenty -- and nothing at all in the other nine.

ISTAT measures it itself. Its multipurpose household survey *I cittadini e il
tempo libero* (2015) asked every member of the sampled households aged 6 and
over which language they use mostly in the family, with friends and with
strangers; the results were published in December 2017 as *L'uso della lingua
italiana, dei dialetti e delle altre lingue in Italia*, with a workbook of
tables. Tavola 1 gives, for each region, the persons aged 6 and over by the
language used in the family -- as percentages (``tav1``) and as weighted
estimates in thousands (``tav1 (segue)``):

* solo o prevalentemente italiano -- written "Italian";
* solo o prevalentemente dialetto -- "Italian dialects". In Trentino-Alto
  Adige (and so in the Nord-Est that contains it) the dialect of half the
  answers is South Tyrol's German one, so there it is written "Local dialect
  (Italian or German)";
* sia italiano che dialetto -- "Italian and local dialect";
* altra lingua -- "Other language": German and Ladin in South Tyrol, Friulian
  and Slovene in Friuli, Sardinian, French and Franco-Provençal in the Aosta
  Valley, and the languages of immigrants everywhere;
* altro -- "Other";
* and the rest of the persons aged 6 and over, who gave no answer, as "Not
  stated".

A survey estimate, so the file's name ends in ``_survey``: it fills a region no
count describes and never replaces one, and as a national office's survey it
stands in front of the European Social Survey.

**Macro-regions** are ISTAT's ripartizioni, the map's first level, and each is
its regions' weighted estimates added: the survey's estimator is calibrated to
regional population totals, so the regions' estimates add to the ripartizione's
own. Nothing is copied from a region onto another unit.

Checks, each of which stops the run:

* every region's categories and its not-stated remainder make its published
  total of persons aged 6 and over, and the remainder is under 2%;
* the published percentages and the estimates in thousands agree within the
  rounding of a thousand;
* Bolzano and Trento make Trentino-Alto Adige, and the regions make Italy, in
  every column;
* all twenty drawn regions and five macro-regions are bound one to one.

Usage:
    python -m scripts.fetch_census.italy_language_survey
"""

from __future__ import annotations

import argparse
import io
import json
import re
from pathlib import Path
from typing import Any

from ._shared import PROCESSED, http_get, log, record, write_json
from .binding import fold

OUT = "italy_language_survey.json"
YEAR = 2015
URL = "https://www.istat.it/wp-content/uploads/2017/12/Lingue-e-dialetti_2015_Tavole.xlsx"
PAGE = "https://www.istat.it/it/archivio/207961"
SOURCE = ("ISTAT, L'uso della lingua italiana, dei dialetti e delle altre lingue in Italia "
          "(survey 'I cittadini e il tempo libero', 2015), Tavola 1")
LICENCE = "CC BY 3.0 IT (ISTAT)"
SITE = PROCESSED.parent.parent / "site" / "data"

# Tavola 1's columns for the family, in the order ISTAT prints them.
CATEGORIES = [
    ("Solo o prevalentemente italiano", "Italian"),
    ("Solo o prevalentemente dialetto", "Italian dialects"),
    ("Sia italiano che dialetto", "Italian and local dialect"),
    ("Altra lingua", "Other language"),
    ("Altro", "Other"),
]
MIXED_DIALECT = "Local dialect (Italian or German)"
NOT_STATED = "Not stated"

# ISTAT's region name -> the map's, by the name's start (the Aosta Valley is
# written bilingually, with whichever apostrophe the workbook was typed with).
REGION_NAMES = {
    "valled": "Valle d'Aosta",
    "friuli": "Friuli Venezia Giulia",
}
# Rows Tavola 1 may carry beside the regions, which are not drawn at the
# second level and are rebuilt from the regions for the first.
AGGREGATES = {"nordovest", "nordest", "centro", "sud", "isole", "mezzogiorno", "nord"}


def canonical(name: str) -> str:
    for start, drawn in REGION_NAMES.items():
        if fold(name).startswith(start):
            return drawn
    return name
# ISTAT's ripartizioni, the map's first level.
RIPARTIZIONI = {
    "Nord-Ovest": ["Piemonte", "Valle d'Aosta", "Liguria", "Lombardia"],
    "Nord-Est": ["Trentino-Alto Adige", "Veneto", "Friuli Venezia Giulia", "Emilia-Romagna"],
    "Centro": ["Toscana", "Umbria", "Marche", "Lazio"],
    "Sud": ["Abruzzo", "Molise", "Campania", "Puglia", "Basilicata", "Calabria"],
    "Isole": ["Sicilia", "Sardegna"],
}
# Regions and ripartizioni whose dialects are not all Italo-Romance.
GERMAN_DIALECTS = {"Trentino-Alto Adige", "Nord-Est"}
PROVINCES = ("Bolzano/Bozen", "Trento")


def cell(value: Any) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def number(value: Any) -> float | None:
    """A table value: '-' is none of them, '..' or blank is not published."""
    text = cell(value).replace(",", ".")
    if text in ("-", "--"):
        return 0.0
    try:
        return float(text)
    except ValueError:
        return None


def read_sheet(rows: list[list[Any]], *, with_total: bool) -> dict[str, dict[str, float]]:
    """Tavola 1 or its (segue): region -> {column label: value}."""
    header_at = None
    for i, row in enumerate(rows):
        labels = [cell(v) for v in row]
        if any(label.startswith("Solo o prevalentemente italiano") for label in labels):
            header_at = i
            break
    if header_at is None:
        raise SystemExit("italy_language_survey: no header row in Tavola 1")
    header = [cell(v) for v in rows[header_at]]
    # The family block is the first run of the five categories; "Con amici"
    # repeats them further right.
    columns: dict[str, int] = {}
    for label, _ in CATEGORIES:
        at = next((j for j, h in enumerate(header) if h == label or
                   (label != "Altro" and h.startswith(label)) or
                   (label == "Altro" and re.match(r"^Altro\b", h))), None)
        if at is None:
            raise SystemExit(f"italy_language_survey: no column {label!r} in {header}")
        columns[label] = at
    if with_total:
        at = next((j for j, h in enumerate(header) if h.startswith("Persone di 6")), None)
        if at is None:
            raise SystemExit(f"italy_language_survey: no total column in {header}")
        columns["total"] = at
    out: dict[str, dict[str, float]] = {}
    for row in rows[header_at + 1:]:
        name = canonical(cell(row[0]) if row else "")
        if (not name or name.startswith("Fonte") or name.startswith("(")
                or fold(name) in AGGREGATES):
            continue
        values = {k: number(row[j]) if j < len(row) else None for k, j in columns.items()}
        if any(v is None for v in values.values()):
            continue
        out[name] = values           # type: ignore[assignment]
    return out


def workbook_rows() -> tuple[list[list[Any]], list[list[Any]]]:
    import openpyxl
    blob = http_get(URL, binary=True, cache=True, timeout=300)
    assert isinstance(blob, bytes)
    book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    pct = [list(r) for r in book["tav1"].iter_rows(values_only=True)]
    thousands = [list(r) for r in book["tav1 (segue)"].iter_rows(values_only=True)]
    return pct, thousands


def check(pct: dict[str, dict[str, float]], kilo: dict[str, dict[str, float]]) -> None:
    missing = sorted(set(kilo) ^ set(pct))
    if missing:
        raise SystemExit(f"italy_language_survey: rows in one sheet only: {missing}")
    for name, row in kilo.items():
        made = sum(row[label] for label, _ in CATEGORIES)
        rest = row["total"] - made
        if rest < -2 or rest > 0.02 * row["total"] + 1:
            raise SystemExit(f"italy_language_survey: {name}: the categories make {made:,.0f} "
                             f"thousand against {row['total']:,.0f}")
        slack = 100.0 * 1.0 / row["total"] + 0.15
        for label, _ in CATEGORIES:
            share = 100.0 * row[label] / row["total"]
            if abs(share - pct[name][label]) > slack:
                raise SystemExit(f"italy_language_survey: {name} {label}: {pct[name][label]}% "
                                 f"published, {share:.2f}% from the thousands")
    for key in ["total"] + [label for label, _ in CATEGORIES]:
        parts = sum(kilo[p][key] for p in PROVINCES)
        if abs(parts - kilo["Trentino-Alto Adige"][key]) > 2:
            raise SystemExit(f"italy_language_survey: Bolzano and Trento make {parts:,.0f} "
                             f"thousand {key}, Trentino-Alto Adige {kilo['Trentino-Alto Adige'][key]}")
        regions = [n for n in kilo if n not in PROVINCES and n != "Italia"]
        made = sum(kilo[n][key] for n in regions)
        if abs(made - kilo["Italia"][key]) > max(12, 0.002 * kilo["Italia"][key]):
            raise SystemExit(f"italy_language_survey: {len(regions)} regions make {made:,.0f} "
                             f"thousand {key}, Italy {kilo['Italia'][key]:,.0f}")


def composition(people: float, shares: dict[str, float], name: str) -> list[dict[str, Any]]:
    """Published shares -> the map's rows, with the unanswered remainder."""
    rows = []
    for label, english in CATEGORIES:
        if english == "Italian dialects" and name in GERMAN_DIALECTS:
            english = MIXED_DIALECT
        share = shares[label]
        rows.append({"group": english, "pct": round(share, 1),
                     "count": int(round(people * share / 100.0))})
    rest = round(100.0 - sum(shares[label] for label, _ in CATEGORIES), 1)
    if rest >= 0.1:
        rows.append({"group": NOT_STATED, "pct": rest, "count": int(round(people * rest / 100.0))})
    rows = [r for r in rows if r["pct"] > 0]
    rows.sort(key=lambda r: (-r["pct"], r["group"]))
    return rows


def note(name: str, people: float, parts: list[str] | None = None) -> str:
    text = (
        "ISTAT's multipurpose household survey 'I cittadini e il tempo libero' (2015), published "
        "in 2017 as 'L'uso della lingua italiana, dei dialetti e delle altre lingue in Italia', "
        "Tavola 1: persons aged 6 and over by the language they use mostly in the family -- only "
        "or mainly Italian, only or mainly dialect, both Italian and dialect, another language, "
        "or another answer; those who gave none are Not stated. A survey estimate, weighted by "
        f"ISTAT to the population ({people / 1e6:,.2f} million persons aged 6 and over here), "
        "not a count. The sample is about 24,000 households in about 850 municipalities, every "
        "member interviewed (a parent answers for a child under 14), and each region -- in "
        "Trentino-Alto Adige each province -- is a domain the design is drawn to estimate; "
        "ISTAT does not publish the respondents per region. Universe: residents of private "
        "households aged 6 and over. Italy's census does not ask language.")
    if name in GERMAN_DIALECTS:
        text += (" In South Tyrol the dialect most answers name is the German one, so here the "
                 "dialect row is Italian or German.")
    if parts:
        text += (" The ripartizione's figure is its regions' estimates added: "
                 + ", ".join(parts) + ".")
    return text


def fields(name: str, people: float, shares: dict[str, float],
           parts: list[str] | None = None) -> dict[str, Any]:
    return {
        "language": composition(people, shares, name),
        "language_year": YEAR,
        "language_basis": "survey estimate: language used mostly in the family, persons aged 6+",
        "language_note": note(name, people, parts),
        "sources": [{"field": "language", "name": SOURCE, "url": PAGE, "year": YEAR,
                     "license": LICENCE}],
    }


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    pct_rows, kilo_rows = workbook_rows()
    pct = read_sheet(pct_rows, with_total=False)
    kilo = read_sheet(kilo_rows, with_total=True)
    check(pct, kilo)
    log(f"italy_language_survey: {len(kilo)} rows; Italy "
        f"{kilo['Italia']['total']:,.0f} thousand persons aged 6 and over")

    admin1 = {fold(u["name"]): u for u in json.loads((SITE / "admin1" / "ITA.units.json").read_text())}
    admin2 = {fold(u["name"]): u for u in json.loads((SITE / "admin2" / "ITA.units.json").read_text())}
    regions = {n: n for n in kilo if n not in PROVINCES and n != "Italia"}
    if sorted(map(fold, regions)) != sorted(admin2):
        raise SystemExit(f"italy_language_survey: ISTAT regions {sorted(regions)} against drawn "
                         f"{sorted(u['name'] for u in admin2.values())}")
    records: list[dict[str, Any]] = []
    for name, istat in sorted(regions.items()):
        shape = admin2[fold(name)]
        people = kilo[istat]["total"] * 1000
        records.append(record(
            f"ITA-ISTAT-LANG-{fold(name)}", shape["name"], level="admin2", parent="ITA",
            country="ITA", match_by="shape_id", shape_id=shape["id"],
            **fields(name, people, pct[istat])))
    if sorted(map(fold, RIPARTIZIONI)) != sorted(admin1):
        raise SystemExit(f"italy_language_survey: ripartizioni against drawn {sorted(admin1)}")
    for macro, members in RIPARTIZIONI.items():
        shape = admin1[fold(macro)]
        for member in members:
            if admin2[fold(member)].get("parent") != shape["id"]:
                raise SystemExit(f"italy_language_survey: {member} is not drawn inside {macro}")
        people = sum(kilo[regions[m]]["total"] for m in members) * 1000
        shares = {label: sum(pct[regions[m]][label] * kilo[regions[m]]["total"] for m in members)
                  / (people / 1000) for label, _ in CATEGORIES}
        records.append(record(
            f"ITA-ISTAT-LANG-{fold(macro)}", shape["name"], level="admin1", parent="ITA",
            country="ITA", match_by="shape_id", shape_id=shape["id"],
            **fields(macro, people, shares, parts=members)))
    for r in records:
        top = r["language"][0]
        log(f"  {r['level']} {r['name']}: {top['group']} {top['pct']}%, "
            + ", ".join(f"{g['group']} {g['pct']}" for g in r["language"][1:]))
    write_json(Path(args.out) if args.out else PROCESSED / OUT, records)
    log(f"  {len(records)} records (20 regions, 5 ripartizioni)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
