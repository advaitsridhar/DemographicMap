#!/usr/bin/env python3
"""Switzerland: nationality as ethnicity, by canton and by the districts the map draws.

Switzerland's population register counts nationality, not ethnicity. By the
owner's decision of 19 September 2026 that count is carried on the ethnicity
field under ``ethnicity_basis: "nationality"`` (see ``central_nationality``).

**What is read.** BFS table px-x-0102010000_104: the permanent and
non-permanent resident population on 31 December by commune, place of birth
and nationality (some 200 countries), from STATPOP. The permanent resident
population is read, every place of birth together. The country's own row
decides which nationalities are named (``NAMED_SHARE`` of the country or
more); the communes are then asked for those and their total, and the rest of
each unit is "Other nationalities".

**Where.** The same ground as ``switzerland_ages``: the districts of 2009
that the map draws, each the sum of today's communes whose every 2009
predecessor lay inside it, through BFS's historicised register (AGVCH). A
district straddled by a commune merged across its 2009 line since takes no
figure -- BFS publishes no split of a merged commune -- and says so. The
cantons are the sums of today's communes by the canton of their 2009
predecessors, so Moutier, which left Bern for Jura on 1 January 2026, counts
in the Bern the map draws. One commune of today crosses a 2009 cantonal line:
Murten (Fribourg) took in Clavaleyres (Bern) in 2022; it is counted in
Fribourg, and both cantons' notes say so. It is counted whole in Fribourg's
Lac district too, the rest of it lying there, and Lac's note says whose
people that adds and how many (Clavaleyres' last count on its own, from BFS's
balance by commune). Fribourg's seven districts then make the canton, which is
checked. Bern's Laupen, which held Clavaleyres in 2009, keeps its gap.

Usage:
    python -m scripts.fetch_census.switzerland_nationality [--year 2025]
"""

from __future__ import annotations

import argparse
import re
import time
from collections import Counter
from typing import Any

from ._shared import NOT_AVAILABLE, PROCESSED, gap, http_get, log, record, write_json
from .central_ages import check_sum, fold, units
from .central_nationality import GERMAN, OTHER, composition, label_for, named, note
from .pxweb import http_json, unstack
from .switzerland_ages import LAKE, VINTAGE, WATER, agvch, drawn_units

TABLE = "https://www.pxweb.bfs.admin.ch/api/v1/de/px-x-0102010000_104/px-x-0102010000_104.px"
PORTAL = "https://www.pxweb.bfs.admin.ch/pxweb/de/px-x-0102010000_104/"
SOURCE = ("Federal Statistical Office (BFS/OFS), STATPOP, px-x-0102010000_104 (resident "
          "population by commune, place of birth and nationality), and the historicised register "
          "of communes (AGVCH)")
LICENCE = "BFS open data (free reuse with attribution)"
OUT = PROCESSED / "switzerland_nationality.json"
NATIONAL = "8100"
TOTAL = "-99999"
PERMANENT = "1"
# A nationality is named when its people are at least this share of the country's.
NAMED_SHARE = 0.004
BATCH = 80


def variables() -> tuple[dict[str, dict[str, Any]], dict[str, str]]:
    meta = {v["code"]: v for v in http_json(TABLE)["variables"]}
    code = {k: next(c for c in meta if c.startswith(k)) for k in
            ("Jahr", "Kanton", "Bevölkerungstyp", "Geburtsort", "Staatsangehörigkeit")}
    return meta, code


def query(code: dict[str, str], year: int, places: list[str], nationalities: list[str]
          ) -> list[tuple[dict[str, tuple[str, str]], float]]:
    q = [
        {"code": code["Jahr"], "selection": {"filter": "item", "values": [str(year)]}},
        {"code": code["Kanton"], "selection": {"filter": "item", "values": places}},
        {"code": code["Bevölkerungstyp"], "selection": {"filter": "item", "values": [PERMANENT]}},
        {"code": code["Geburtsort"], "selection": {"filter": "item", "values": [TOTAL]}},
        {"code": code["Staatsangehörigkeit"], "selection": {"filter": "item", "values": nationalities}},
    ]
    return unstack(http_json(TABLE, {"query": q, "response": {"format": "json-stat2"}}))


def todays_communes(meta: dict[str, dict[str, Any]], place: str, year: int
                    ) -> tuple[list[str], dict[str, set[str]], set[str], dict[str, str]]:
    """The table's communes, today's-commune -> its 2009 communes, the lake
    areas the register lists and the table leaves out, and today's names.

    The register date is the one whose correspondence lands on exactly the
    table's communes, as switzerland_ages finds it.
    """
    px = [v for v in meta[place]["values"] if re.fullmatch(r"\d{4}", v) and v != NATIONAL]
    for end in (f"01-01-{year + 1}", f"31-12-{year}", f"01-01-{year}"):
        rows = agvch(f"correspondances?startPeriod={VINTAGE}&endPeriod={end}&includeUnmodified=true")
        names = {str(int(r["TerminalCode"])).zfill(4): r["TerminalName"] for r in rows}
        only_px = set(px) - set(names)
        only_register = set(names) - set(px)
        if not only_px and all(LAKE.search(names[c]) for c in only_register):
            initial_of: dict[str, set[str]] = {}
            for r in rows:
                initial_of.setdefault(str(int(r["TerminalCode"])).zfill(4), set()).add(r["InitialCode"])
            log(f"  the table's {len(px)} communes are the register's of {end}")
            return px, initial_of, only_register, names
    raise SystemExit("switzerland_nationality: no register date matches the table's communes")


def build(year: int) -> list[dict[str, Any]]:
    log(f"switzerland_nationality: BFS px-x-0102010000_104, 31 December {year}, read against "
        f"the districts of {VINTAGE}")
    meta, code = variables()
    if str(year) not in meta[code["Jahr"]]["values"]:
        raise SystemExit(f"switzerland_nationality: the table has no {year}")
    nat_values = meta[code["Staatsangehörigkeit"]]["values"]
    nat_names = dict(zip(nat_values, meta[code["Staatsangehörigkeit"]]["valueTexts"]))

    # The country decides which nationalities are named.
    national: dict[str, float] = {}
    for key, value in query(code, year, [NATIONAL], nat_values):
        national[key[code["Staatsangehörigkeit"]][0]] = value
    total = national.pop(TOTAL)
    check_sum(national.values(), total, "nationalities against Switzerland")
    labels = {c: label_for(nat_names[c], GERMAN) for c in national}
    unlabelled = [f"{c} {nat_names[c]!r} ({n:,.0f})" for c, n in national.items()
                  if labels[c] is None and n / total >= NAMED_SHARE]
    if unlabelled:
        raise SystemExit(f"switzerland_nationality: nationalities above the naming threshold "
                         f"with no label: {unlabelled}")
    names = named(national, total, labels, share=NAMED_SHARE, always=["8100"])
    log(f"  named: {', '.join(str(labels[n]) for n in names)}")

    # Every commune, for the named nationalities and its total.
    px, initial_of, lakes, todays = todays_communes(meta, code["Kanton"], year)
    asked = [TOTAL, *names]
    counts: dict[str, dict[str, float]] = {}
    for start in range(0, len(px), BATCH):
        for key, value in query(code, year, px[start:start + BATCH], asked):
            counts.setdefault(key[code["Kanton"]][0], {})[key[code["Staatsangehörigkeit"]][0]] = value
        time.sleep(0.5)
    check_sum((c.get(TOTAL, 0) for c in counts.values()), total, "communes against Switzerland")

    def fields(communes: list[str], where: str, extra: str = "") -> dict[str, Any]:
        unit: Counter = Counter()
        for c in communes:
            unit.update(counts.get(c, {}))
        whole = unit.pop(TOTAL, 0)
        unit["__rest__"] = whole - sum(unit.values())
        return {"ethnicity": composition(dict(unit), whole, names,
                                         {**labels, "__rest__": None}, where=where),
                "ethnicity_year": year, "ethnicity_basis": "nationality",
                "ethnicity_note": note("Permanent resident population by nationality",
                                       f"31 December {year}",
                                       "BFS's STATPOP count from the population registers "
                                       "(px-x-0102010000_104)", extra=extra)}

    sources = [{"field": "ethnicity", "name": SOURCE, "url": PORTAL, "license": LICENCE,
                "year": year}]
    snapshot = agvch(f"snapshot?date={VINTAGE}")
    cantons = {r["HistoricalCode"]: r for r in snapshot if r["Level"] == "1"}
    districts = {r["HistoricalCode"]: r for r in snapshot if r["Level"] == "2"}
    old_communes = {r["BfsCode"]: r for r in snapshot if r["Level"] == "3"}
    canton_of_2009 = {c: districts[r["Parent"]]["Parent"] for c, r in old_communes.items()}

    records: list[dict[str, Any]] = []
    bound, unbound, left = drawn_units(snapshot)
    written = gaps = 0
    district_totals: dict[str, list[float]] = {}
    for name, canton_name, members, shape in bound:
        todays_here = sorted(t for t, ini in initial_of.items() if ini & members and t not in lakes)
        straddling = [t for t in todays_here if not initial_of[t] <= members]
        # A commune merged across a cantonal line is counted whole in the
        # district that holds the rest of it, as its canton counts it whole.
        across = [t for t in straddling
                  if crosses_canton_only(initial_of[t], members, canton_of_2009)]
        straddling = [t for t in straddling if t not in across]
        extra = ""
        if across and not straddling:
            extra = took_in_note(across, initial_of, members, old_communes, todays, name)
        if straddling:
            why = (f"Not written: the map draws this district as it was in 2009, and "
                   f"{', '.join(todays[t] for t in straddling)} -- today's commune(s) -- merged "
                   f"2009 communes from inside and outside it; BFS publishes nationality by "
                   f"today's communes only and no split of a merged commune.")
            unit_fields = {"ethnicity": gap(NOT_AVAILABLE, why)}
            gaps += 1
        else:
            unit_fields = fields(todays_here, name, extra)
            district_totals.setdefault(next(iter(canton_of_2009[c] for c in members)), []).append(
                sum(counts.get(c, {}).get(TOTAL, 0) for c in todays_here))
            written += 1
        records.append(record(
            f"CHE-2009-{fold(name)}-nat", shape["name"], level="admin2", parent=shape["parent"],
            country="CHE", codes={"bfs_2009": name}, match_by="shape_id", shape_id=shape["id"],
            sources=sources, **unit_fields))
    log(f"  districts: {written} written, {gaps} straddled; unbound {unbound}; "
        f"polygons with no unit {left}")
    if unbound or set(left) - WATER:
        raise SystemExit("switzerland_nationality: units and polygons do not pair up one to one")

    # Cantons, by the canton of each commune's 2009 predecessors.
    by_canton: dict[str, list[str]] = {}
    crossing: dict[str, list[str]] = {}
    for t, ini in initial_of.items():
        if t in lakes:
            continue
        homes = Counter(canton_of_2009[c] for c in ini if c in canton_of_2009)
        if not homes:
            raise SystemExit(f"switzerland_nationality: {todays[t]} has no 2009 canton")
        home = homes.most_common(1)[0][0]
        by_canton.setdefault(home, []).append(t)
        if len(homes) > 1:
            for other in homes:
                if other != home:
                    crossing.setdefault(home, []).append(
                        f"{todays[t]} (counted here) took in communes the map draws in "
                        f"{cantons[other]['Name'].split('/')[0].strip()}")
                    crossing.setdefault(other, []).append(
                        f"{todays[t]} is counted in {cantons[home]['Name'].split('/')[0].strip()}, "
                        f"though it took in communes the map draws here")
    shapes = units("CHE", "admin1")
    used = set()
    for hist, members in sorted(by_canton.items()):
        label = cantons[hist]["Name"]
        hits = [s for s in shapes if any(fold(p) == fold(s["name"]) for p in label.split("/"))]
        if len(hits) != 1 or hits[0]["id"] in used:
            raise SystemExit(f"switzerland_nationality: canton {label!r} meets {[h['name'] for h in hits]}")
        used.add(hits[0]["id"])
        extra = ("; ".join(crossing[hist]) + ".") if hist in crossing else ""
        records.append(record(
            f"CHE-canton-{fold(label.split('/')[0])}-nat", hits[0]["name"], level="admin1",
            parent="CHE", country="CHE", match_by="shape_id", shape_id=hits[0]["id"],
            sources=sources, **fields(members, label, extra)))
    if len(used) != len(shapes):
        raise SystemExit(f"switzerland_nationality: {len(used)} cantons bound of {len(shapes)}")
    # A canton whose every district was written is the sum of them: the
    # communes counted whole across a cantonal line are counted in the same
    # canton at both levels.
    drawn_in: Counter = Counter(next(iter(canton_of_2009[c] for c in members))
                                for _, _, members, _ in bound)
    for hist, members in by_canton.items():
        totals = district_totals.get(hist, [])
        if drawn_in.get(hist) and len(totals) == drawn_in[hist]:
            whole = sum(counts.get(c, {}).get(TOTAL, 0) for c in members)
            check_sum(totals, whole, f"{cantons[hist]['Name']}'s districts against the canton")
    return records


def crosses_canton_only(initial: set[str], members: set[str],
                        canton_of_2009: dict[str, str]) -> bool:
    """A commune of today whose 2009 predecessors outside this district all lay
    in another canton, and the rest inside it, so the canton the district
    belongs to counts it whole: Murten, which took in Clavaleyres (Bern)."""
    inside, outside = initial & members, initial - members
    if not inside or not outside:
        return False
    here = {canton_of_2009.get(c) for c in members}
    if len(here) != 1:
        return False
    canton = here.pop()
    homes = Counter(canton_of_2009.get(c) for c in initial)
    return (all(canton_of_2009.get(c) not in (None, canton) for c in outside)
            and homes.most_common(1)[0][0] == canton)


def took_in_note(across: list[str], initial_of: dict[str, set[str]], members: set[str],
                 old_communes: dict[str, dict[str, str]], todays: dict[str, str],
                 district: str) -> str:
    """What a district counts besides its 2009 ground, and how many people that is."""
    said = []
    for t in across:
        for code in sorted(initial_of[t] - members):
            name = old_communes.get(code, {}).get("Name", code)
            when, people = last_count(code)
            size = (f", which had {people:,.0f} residents on 31 December {when}"
                    if people is not None else "")
            said.append(f"{todays[t]} is counted whole: it took in {name}, a commune of "
                        f"another canton in 2009{size}, so this figure counts them too")
    return "; ".join(said) + "." if said else ""


def last_count(code: str) -> tuple[int | None, float | None]:
    """A commune's last population on 31 December in BFS's balance by commune,
    the latest year's sheet that still lists it."""
    import io as _io

    import openpyxl
    from .switzerland_ages import BALANCE, FIRST_BALANCE_YEAR, balance_sheet
    try:
        book = openpyxl.load_workbook(_io.BytesIO(http_get(BALANCE, binary=True, timeout=300)),
                                      read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 -- the note is then said without the count
        log(f"  the balance by commune could not be read ({exc}); no count for {code}")
        return None, None
    key = str(int(code)).zfill(4)
    years = sorted((int(s) for s in book.sheetnames if s.isdigit()), reverse=True)
    for when in years:
        if when < FIRST_BALANCE_YEAR:
            break
        sheet, _ = balance_sheet([list(r) for r in book[str(when)].iter_rows(values_only=True)])
        if key in sheet:
            return when, sheet[key]
    return None, None


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--year", type=int, default=2025)
    args = ap.parse_args()
    records = build(args.year)
    write_json(OUT, records)
    log(f"  wrote {OUT.name}: {len(records)} records")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
