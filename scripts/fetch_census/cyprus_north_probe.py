#!/usr/bin/env python3
"""Read-only probes for the census taken in the north of Cyprus.

The Turkish Cypriot administration in the north (not recognised
internationally) took a population and housing census on 4 December 2011
("2011 Nüfus ve Konut Sayımı"). Its results were published by the State
Planning Organisation (Devlet Planlama Örgütü, ``devplan.org``) and are now
kept by the statistics institute (``istatistik.gov.ct.tr``). Neither host's
layout is guessable from here, and the sandbox that writes the adapter cannot
reach either, so this asks the runner and prints only what a decision needs:
which files the hosts and the Internet Archive hold, and what is in them.

Each probe is named, and several run in one call:

    python -m scripts.fetch_census.cyprus_north_probe --run cdx_devplan,home

A probe that takes a URL reads it from ``--url`` (repeatable), from the live
host or, with ``--wayback TS``, from the Internet Archive's raw capture.

Nothing is written; the output is the log.
"""

from __future__ import annotations

import argparse
import html
import io
import json
import re
import urllib.error
import urllib.parse
import urllib.request
from typing import Any, Callable

UA = "DemographicMap/1.0 (+https://github.com/advaitsridhar/DemographicMap) python-urllib"
TIMEOUT = 150
CDX = "https://web.archive.org/cdx/search/cdx"
TERMS = r"(?i).*(n%C3%BCfus|nufus|n%FCfus|say%C4%B1m|sayim|census|population).*"


def log(*args: Any) -> None:
    print(*args, flush=True)


def as_uri(url: str) -> str:
    """The institute's pages are named in Turkish ('NÜFUS-SAYIMLARI'), and
    urllib sends a request line in ASCII: escape what is not, keep what is."""
    return urllib.parse.quote(url, safe=":/?&=%#+,;@!$'()*[]~")


def fetch(url: str, timeout: int = TIMEOUT) -> tuple[int, str, bytes]:
    """(status, content type, body); an HTTP error is returned, not raised."""
    req = urllib.request.Request(as_uri(url), headers={"User-Agent": UA})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return resp.status, resp.headers.get("Content-Type", ""), resp.read()
    except urllib.error.HTTPError as exc:
        return exc.code, exc.headers.get("Content-Type", "") if exc.headers else "", b""


def raw_url(url: str, wayback: str | None) -> str:
    return f"https://web.archive.org/web/{wayback}id_/{url}" if wayback else url


def cdx(query: dict[str, str], limit: int = 400) -> list[list[str]]:
    params = {"output": "json", "fl": "timestamp,original,mimetype,statuscode,length",
              "collapse": "urlkey", "limit": str(limit), **query}
    status, _, body = fetch(CDX + "?" + urllib.parse.urlencode(params), timeout=240)
    if status != 200:
        log(f"  CDX {query} -> HTTP {status}")
        return []
    rows = json.loads(body.decode("utf-8", "replace") or "[]")
    return rows[1:] if rows else []


def show_cdx(label: str, query: dict[str, str], limit: int = 400, keep: str | None = None) -> None:
    log(f"== CDX {label}: {query}")
    try:
        rows = cdx(query, limit)
    except Exception as exc:                              # noqa: BLE001 -- reported
        log(f"  {type(exc).__name__}: {str(exc)[:200]}")
        return
    if keep:
        rows = [r for r in rows if re.search(keep, urllib.parse.unquote(r[1]), re.I)]
    log(f"  {len(rows)} rows")
    for ts, orig, mime, status, length in rows:
        log(f"  {ts} {status} {mime[:28]:28} {length:>9} {urllib.parse.unquote(orig)}")


def links_of(body: str, base: str) -> list[tuple[str, str]]:
    out = []
    for m in re.finditer(r"""(?is)<a\b[^>]*?href\s*=\s*["']?([^"' >]+)["']?[^>]*>(.*?)</a>""", body):
        href = html.unescape(m.group(1))
        text = re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", m.group(2)))).strip()
        out.append((urllib.parse.urljoin(base, href), text))
    return out


def show_links(url: str, grep: str | None, wayback: str | None = None, limit: int = 120) -> None:
    target = raw_url(url, wayback)
    log(f"== links {target}")
    try:
        status, ctype, body = fetch(target)
    except Exception as exc:                              # noqa: BLE001 -- reported
        log(f"  {type(exc).__name__}: {str(exc)[:200]}")
        return
    log(f"  HTTP {status} {ctype} {len(body):,} bytes")
    if not body:
        return
    text = body.decode("utf-8", "replace")
    seen = set()
    shown = 0
    for href, label in links_of(text, url):
        if href in seen:
            continue
        seen.add(href)
        if grep and not re.search(grep, urllib.parse.unquote(href) + " " + label, re.I):
            continue
        log(f"  {label[:70]!r} -> {urllib.parse.unquote(href)}")
        shown += 1
        if shown >= limit:
            log("  ...")
            break
    if not shown:
        title = re.search(r"(?is)<title>(.*?)</title>", text)
        log(f"  no matching links; title {title.group(1).strip()[:120] if title else None!r}; "
            f"starts {' '.join(text[:300].split())!r}")


def show_head(url: str, wayback: str | None = None) -> None:
    target = raw_url(url, wayback)
    try:
        status, ctype, body = fetch(target)
    except Exception as exc:                              # noqa: BLE001 -- reported
        log(f"== {target}\n  {type(exc).__name__}: {str(exc)[:200]}")
        return
    log(f"== {target}\n  HTTP {status} {ctype} {len(body):,} bytes; starts {body[:8]!r}")


def show_pdf(url: str, wayback: str | None, pages: str | None, grep: str | None,
             chars: int = 4000) -> None:
    import pdfplumber
    target = raw_url(url, wayback)
    log(f"== pdf {target}")
    status, ctype, body = fetch(target, timeout=300)
    log(f"  HTTP {status} {ctype} {len(body):,} bytes")
    if status != 200 or not body.startswith(b"%PDF"):
        return
    with pdfplumber.open(io.BytesIO(body)) as pdf:
        n = len(pdf.pages)
        log(f"  {n} pages")
        wanted: list[int] = []
        for part in (pages or "").split(","):
            if "-" in part:
                a, b = part.split("-")
                wanted += list(range(int(a), int(b) + 1))
            elif part.strip():
                wanted.append(int(part))
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if wanted and i not in wanted:
                continue
            if not wanted and grep and not re.search(grep, text, re.I):
                continue
            if not wanted and not grep:
                first = " | ".join(line for line in text.splitlines()[:3])
                log(f"  p{i}: {first[:200]}")
                continue
            log(f"  --- page {i} ---")
            log(text[:chars])


def show_xls(url: str, wayback: str | None, rows: int = 40, start: int = 0, width: int = 24,
             cols: int = 16) -> None:
    target = raw_url(url, wayback)
    log(f"== workbook {target}")
    status, ctype, body = fetch(target, timeout=300)
    log(f"  HTTP {status} {ctype} {len(body):,} bytes; starts {body[:8]!r}")
    if status != 200:
        return
    if body[:2] == b"PK":
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        for ws in wb.worksheets:
            log(f"  sheet {ws.title!r} {ws.max_row}x{ws.max_column}")
            for r, row in enumerate(ws.iter_rows(values_only=True)):
                if r >= start + rows:
                    break
                if r >= start:
                    log(f"   {r}: " + " | ".join("" if v is None else str(v)[:width] for v in row[:cols]))
    elif body[:4] == b"\xd0\xcf\x11\xe0":
        import xlrd
        wb = xlrd.open_workbook(file_contents=body)
        for ws in wb.sheets():
            log(f"  sheet {ws.name!r} {ws.nrows}x{ws.ncols}")
            for r in range(start, min(start + rows, ws.nrows)):
                log(f"   {r}: " + " | ".join(str(v)[:width] for v in ws.row_values(r)[:cols]))


# --- named probes -----------------------------------------------------------

def p_cdx_devplan(a: argparse.Namespace) -> None:
    show_cdx("devplan Nufus-2011 prefix", {"url": "devplan.org/Nufus-2011/", "matchType": "prefix"})
    show_cdx("devplan census terms", {"url": "devplan.org", "matchType": "domain",
                                      "filter": "original:" + TERMS}, limit=600)


def p_cdx_istatistik(a: argparse.Namespace) -> None:
    show_cdx("istatistik.gov.ct.tr census terms",
             {"url": "istatistik.gov.ct.tr", "matchType": "domain", "filter": "original:" + TERMS},
             limit=800)


def p_cdx_istatistik_files(a: argparse.Namespace) -> None:
    show_cdx("istatistik.gov.ct.tr files",
             {"url": "istatistik.gov.ct.tr", "matchType": "domain",
              "filter": "original:(?i).*\\.(pdf|xls|xlsx|zip)$"}, limit=1500)


def p_home(a: argparse.Namespace) -> None:
    grep = r"n[uü]fus|say[iı]m|census|population|yay[iı]n|istatistik"
    for url in ("http://www.devplan.org/", "https://www.devplan.org/",
                "https://istatistik.gov.ct.tr/", "http://www.devplan.org/Nufus-2011/"):
        show_links(url, grep)


def p_links(a: argparse.Namespace) -> None:
    for url in a.url:
        show_links(url, a.grep, a.wayback)


def p_head(a: argparse.Namespace) -> None:
    for url in a.url:
        show_head(url, a.wayback)


def p_pdf(a: argparse.Namespace) -> None:
    for url in a.url:
        show_pdf(url, a.wayback, a.pages, a.grep, a.chars)


def p_xls(a: argparse.Namespace) -> None:
    for url in a.url:
        show_xls(url, a.wayback, a.rows, a.start, a.width, a.cols)


def p_geonames(a: argparse.Namespace) -> None:
    """GeoNames' Cyprus features that are places (class P) or areas (A), each
    with its point and its Latin-script alternate names: binding evidence for
    the census's Turkish names, never data. One line each, tab-separated:
    GN, id, name, lat, lon, feature code, alternate names joined by '|'."""
    import zipfile
    status, ctype, body = fetch("https://download.geonames.org/export/dump/CY.zip", timeout=300)
    log(f"== GeoNames CY.zip HTTP {status} {len(body):,} bytes")
    if status != 200:
        return
    with zipfile.ZipFile(io.BytesIO(body)) as zf:
        text = zf.read("CY.txt").decode("utf-8")
    n = 0
    for line in text.splitlines():
        f = line.split("\t")
        if len(f) < 15 or f[6] not in ("P", "A"):
            continue
        alts = sorted({x for x in f[3].split(",")
                       if x and not re.search(r"[\u0370-\u03ff\u1f00-\u1fff]", x) and not x.startswith("http")})
        log("\t".join(["GN", f[0], f[1], f[4], f[5], f[7], "|".join(alts)]))
        n += 1
    log(f"== GeoNames: {n} features of class P or A")


WIKIDATA_QUERY = """
SELECT ?item ?lat ?lon ?tr ?el ?en (GROUP_CONCAT(DISTINCT ?alt; separator="|") AS ?alts)
       (GROUP_CONCAT(DISTINCT STRAFTER(STR(?cls), "entity/"); separator="|") AS ?classes) WHERE {
  SERVICE wikibase:box {
    ?item wdt:P625 ?loc .
    bd:serviceParam wikibase:cornerSouthWest "Point(32.2 34.5)"^^geo:wktLiteral .
    bd:serviceParam wikibase:cornerNorthEast "Point(34.7 35.75)"^^geo:wktLiteral .
  }
  ?item wdt:P31 ?cls .
  ?item p:P625/psv:P625 [wikibase:geoLatitude ?lat; wikibase:geoLongitude ?lon] .
  ?item rdfs:label ?tr FILTER(LANG(?tr) = "tr")
  OPTIONAL { ?item rdfs:label ?el FILTER(LANG(?el) = "el") }
  OPTIONAL { ?item rdfs:label ?en FILTER(LANG(?en) = "en") }
  OPTIONAL { ?item skos:altLabel ?alt FILTER(LANG(?alt) IN ("tr", "en")) }
} GROUP BY ?item ?lat ?lon ?tr ?el ?en
"""


def p_wikidata(a: argparse.Namespace) -> None:
    """Wikidata's items on the island with a Turkish label and a point, with
    their Greek and English labels and classes: binding evidence only. One
    line each: WD, item, lat, lon, tr, el, en, alternate labels, classes."""
    url = "https://query.wikidata.org/sparql?" + urllib.parse.urlencode(
        {"query": WIKIDATA_QUERY, "format": "json"})
    status, ctype, body = fetch(url, timeout=300)
    log(f"== Wikidata HTTP {status} {len(body):,} bytes")
    if status != 200:
        log(body[:300].decode("utf-8", "replace"))
        return
    rows = json.loads(body)["results"]["bindings"]
    for r in rows:
        v = {k: r[k]["value"] for k in r}
        log("\t".join(["WD", v["item"].rsplit("/", 1)[-1], v.get("lat", ""), v.get("lon", ""),
                       v.get("tr", ""), v.get("el", ""), v.get("en", ""), v.get("alts", ""),
                       v.get("classes", "")]))
    log(f"== Wikidata: {len(rows)} items")


OVERPASS = "https://overpass-api.de/api/interpreter"
OSM_PLACES = """[out:json][timeout:180];
node["place"](34.95,32.5,35.75,34.65);
out body;"""
OSM_ADMIN = """[out:json][timeout:180];
relation["boundary"="administrative"]["admin_level"~"^(6|7|8|9|10)$"](34.95,32.5,35.75,34.65);
out tags;"""


OVERPASS_MIRRORS = (OVERPASS, "https://overpass.private.coffee/api/interpreter",
                    "https://maps.mail.ru/osm/tools/overpass/api/interpreter")


def overpass(query: str) -> list[dict[str, Any]]:
    """The query's elements, from the main instance or a mirror: one run had
    the main instance answer 'Network is unreachable' where the run before
    had read it."""
    for host in OVERPASS_MIRRORS:
        url = host + "?" + urllib.parse.urlencode({"data": query})
        try:
            status, ctype, body = fetch(url, timeout=300)
        except Exception as exc:                          # noqa: BLE001 -- reported
            log(f"== Overpass {host}: {type(exc).__name__}: {str(exc)[:150]}")
            continue
        log(f"== Overpass {host} HTTP {status} {len(body):,} bytes")
        if status == 200:
            return json.loads(body).get("elements", [])
        log(body[:300].decode("utf-8", "replace"))
    return []


def p_osm_places(a: argparse.Namespace) -> None:
    """OpenStreetMap's place nodes in the north half of the island, with their
    Turkish, Greek and other names: binding evidence only. One line each:
    OSM, node id, lat, lon, place, name, name:tr, name:el, name:en, alt/old names."""
    for el in overpass(OSM_PLACES):
        t = el.get("tags", {})
        alts = "|".join(v for k, v in sorted(t.items())
                        if k in ("alt_name", "old_name", "alt_name:tr", "old_name:tr", "alt_name:el",
                                 "old_name:el", "official_name", "name:tr-CY", "short_name"))
        log("\t".join(["OSM", str(el["id"]), str(el.get("lat")), str(el.get("lon")),
                       t.get("place", ""), t.get("name", ""), t.get("name:tr", ""),
                       t.get("name:el", ""), t.get("name:en", ""), alts]))


def p_osm_admin(a: argparse.Namespace) -> None:
    """OpenStreetMap's administrative boundary relations on the island (levels
    6-10): which of the north's districts, municipalities and quarters it
    draws, before any geometry is asked for."""
    for el in overpass(OSM_ADMIN):
        t = el.get("tags", {})
        log("\t".join(["ADM", str(el["id"]), t.get("admin_level", ""), t.get("name", ""),
                       t.get("name:tr", ""), t.get("name:el", ""), t.get("name:en", ""),
                       t.get("ref", ""), t.get("is_in", "")]))


OSM_GEOM = """[out:json][timeout:240];
relation["boundary"="administrative"]["admin_level"~"^(8|9)$"](35.05,32.5,35.75,34.65);
out geom;"""


def p_osm_geom(a: argparse.Namespace) -> None:
    """The outlines of OpenStreetMap's quarters and villages in the north
    (admin levels 8 and 9), so the sandbox can measure which drawn polygon
    each census quarter lies in: binding evidence only. One line each: GEOM,
    relation id, admin level, name, WKT simplified to about 20 metres."""
    for el in overpass(OSM_GEOM):
        t = el.get("tags", {})
        log("\t".join(["GEOM", str(el["id"]), t.get("admin_level", ""), t.get("name", ""),
                       outline_wkt(el)]))


def outline_wkt(el: dict[str, Any]) -> str:
    """A relation's outer ways, polygonised and simplified, as WKT to five
    decimals (about a metre), or EMPTY where they do not close."""
    from shapely import to_wkt
    from shapely.geometry import LineString
    from shapely.ops import polygonize, unary_union
    lines = [LineString([(p["lon"], p["lat"]) for p in m["geometry"]])
             for m in el.get("members", [])
             if m.get("type") == "way" and m.get("role") in ("outer", "")
             and len(m.get("geometry") or []) >= 2]
    polys = list(polygonize(unary_union(lines))) if lines else []
    if not polys:
        return "EMPTY"
    shape_ = unary_union(polys).simplify(0.0002)
    wkt = to_wkt(shape_, rounding_precision=5, trim=True)
    if len(wkt) > 60000:
        wkt = to_wkt(shape_.simplify(0.001), rounding_precision=5, trim=True)
    return wkt


# Where a census quarter's outline crosses a drawn polygon's line, the share
# of its buildings on each side says more about where its people live than the
# share of its area. (south, west, north, east)
BUILDING_BOXES = {
    "north_nicosia": (35.165, 33.27, 35.255, 33.45),
    "kyrenia": (35.285, 33.22, 35.36, 33.41),
    "esentepe": (35.27, 33.40, 35.38, 33.63),
    "famagusta": (35.09, 33.84, 35.17, 33.98),
    "karavas": (35.31, 33.13, 35.37, 33.24),
    "mesarya": (35.15, 33.70, 35.25, 33.82),
    "morfou": (35.17, 32.93, 35.23, 33.08),
    "trikomo": (35.26, 33.86, 35.31, 33.93),
    # Geçitköy (Panagra) and Çamlıbel (Myrtou), whose shared edge runs close
    # to Geçitköy's houses.
    "panagra": (35.30, 33.04, 35.36, 33.11),
}


def p_osm_buildings(a: argparse.Namespace) -> None:
    """OpenStreetMap's buildings in the boxes above (or only those named by
    --boxes), counted per cell of a thousandth of a degree (about 90 by 110
    metres): one line per cell, BLD, box, lat, lon of the cell's corner,
    count. Binding evidence only."""
    from collections import Counter
    wanted = set(filter(None, (a.boxes or "").split(",")))
    unknown = wanted - set(BUILDING_BOXES)
    if unknown:
        raise SystemExit(f"cyprus_north_probe: no building box {sorted(unknown)}")
    for name, (s, w, n, e) in BUILDING_BOXES.items():
        if wanted and name not in wanted:
            continue
        query = f'[out:json][timeout:240];way["building"]({s},{w},{n},{e});out center;'
        cells: Counter = Counter()
        for el in overpass(query):
            c = el.get("center")
            if c:
                cells[(int(c["lat"] * 1000), int(c["lon"] * 1000))] += 1
        log(f"== buildings {name}: {sum(cells.values())} in {len(cells)} cells")
        for (la, lo), k in sorted(cells.items()):
            log(f"BLD\t{name}\t{la / 1000:.3f}\t{lo / 1000:.3f}\t{k}")


# Names OpenStreetMap may give the United Nations buffer zone between the two
# sides (the "Green Line"), in English, Turkish and Greek.
GREEN_NAMES = ("buffer|tampon|ara b[oö]lge|unficyp|green line|ye[sş]il hat|ate[sş]kes|ceasefire|"
               "νεκρή|ουδέτερ|πράσινη γραμμή")
ISLAND = (34.5, 32.2, 35.8, 34.7)


def relation_polygon(el: dict[str, Any]) -> Any:
    """A relation's (or closed way's) area as one shapely geometry: its outer
    rings polygonised, less its inner ones; None where nothing closes."""
    from shapely.geometry import LineString
    from shapely.ops import polygonize, unary_union

    def rings(roles: tuple[str, ...]) -> Any:
        if el.get("type") == "way":
            pts = [(p["lon"], p["lat"]) for p in el.get("geometry") or []]
            lines = [LineString(pts)] if roles != ("inner",) and len(pts) >= 4 else []
        else:
            lines = [LineString([(p["lon"], p["lat"]) for p in m["geometry"]])
                     for m in el.get("members", [])
                     if m.get("type") == "way" and m.get("role") in roles
                     and len(m.get("geometry") or []) >= 2]
        faces = list(polygonize(unary_union(lines))) if lines else []
        return unary_union(faces) if faces else None

    outer = rings(("outer", ""))
    if outer is None:
        return None
    inner = rings(("inner",))
    return outer.difference(inner) if inner is not None else outer


def p_green_line(a: argparse.Namespace) -> None:
    """Which side of the line between the two sides a place is on, by
    OpenStreetMap: binding evidence only.

    * ISIN: every OpenStreetMap area that holds each --point (label,lat,lon);
    * CAND: the relations and ways on the island named for the United Nations
      buffer zone, and the island's boundaries of admin level 3 to 4;
    * ZONE: the outline of each such area clipped to --box (south,west,north,
      east), to about three metres, and of the buffer zone island-wide, to
      about fifty; with each --point's distance to it, in metres (negative
      inside).
    """
    from shapely import to_wkt
    from shapely.geometry import Point, box
    points = []
    for spec in a.point:
        label, lat, lon = spec.split(",")
        points.append((label, float(lat), float(lon)))
        for el in overpass(f"[out:json][timeout:90];is_in({lat},{lon});out tags;"):
            t = el.get("tags", {})
            extra = "|".join(f"{k}={v}" for k, v in sorted(t.items())
                             if k in ("military", "landuse", "place", "disputed", "type", "border_type",
                                      "de_facto", "political_division", "protect_class"))
            log("\t".join(["ISIN", label, lat, lon, el.get("type", ""), str(el.get("id")),
                           t.get("name", ""), t.get("name:en", ""), t.get("admin_level", ""),
                           t.get("boundary", ""), extra]))
    s, w, n, e = ISLAND
    query = (f'[out:json][timeout:180];(relation["name"~"{GREEN_NAMES}",i]({s},{w},{n},{e});'
             f'way["name"~"{GREEN_NAMES}",i]({s},{w},{n},{e});'
             f'relation["name:en"~"{GREEN_NAMES}",i]({s},{w},{n},{e});'
             f'relation["boundary"]["admin_level"~"^[34]$"]({s},{w},{n},{e}););out tags;')
    found = overpass(query)
    zone_ids = set()
    for el in found:
        t = el.get("tags", {})
        green = bool(re.search(GREEN_NAMES, " ".join([t.get("name", ""), t.get("name:en", "")]), re.I))
        if green:
            zone_ids.add((el["type"], el["id"]))
        log("\t".join(["CAND", el["type"], str(el["id"]), t.get("name", ""), t.get("name:en", ""),
                       t.get("admin_level", ""), t.get("boundary", ""),
                       "|".join(f"{k}={v}" for k, v in sorted(t.items())
                                if k in ("military", "landuse", "type", "disputed", "border_type")),
                       "green" if green else ""]))
    clip = None
    if a.box:
        south, west, north, east = map(float, a.box.split(","))
        clip = box(west, south, east, north)
    for kind, ident in sorted(zone_ids | {("relation", el["id"]) for el in found
                                           if el["type"] == "relation"
                                           and el.get("tags", {}).get("admin_level") in ("3", "4")}):
        geo = overpass(f"[out:json][timeout:240];{kind}({ident});out geom;")
        if not geo:
            continue
        shape_ = relation_polygon(geo[0])
        name = geo[0].get("tags", {}).get("name", "")
        if shape_ is None or shape_.is_empty:
            log(f"ZONE\t{kind}\t{ident}\t{name}\tnone\tEMPTY")
            continue
        for label, lat, lon in points:
            pt = Point(lon, lat)
            d = shape_.boundary.distance(pt) * 111_000 * (-1 if shape_.contains(pt) else 1)
            log(f"DIST\t{kind}\t{ident}\t{name}\t{label}\t{d:.0f}")
        if clip is not None:
            part = shape_.intersection(clip).simplify(0.00003)
            log(f"ZONE\t{kind}\t{ident}\t{name}\tbox\t{to_wkt(part, rounding_precision=5, trim=True)}")
        if (kind, ident) in zone_ids:
            whole = shape_.simplify(0.0005)
            log(f"ZONE\t{kind}\t{ident}\t{name}\tisland\t{to_wkt(whole, rounding_precision=4, trim=True)}")


PROBES: dict[str, Callable[[argparse.Namespace], None]] = {
    "green_line": p_green_line,
    "osm_buildings": p_osm_buildings,
    "osm_places": p_osm_places, "osm_admin": p_osm_admin, "osm_geom": p_osm_geom,
    "cdx_devplan": p_cdx_devplan, "cdx_istatistik": p_cdx_istatistik,
    "cdx_istatistik_files": p_cdx_istatistik_files, "home": p_home, "links": p_links,
    "head": p_head, "pdf": p_pdf, "xls": p_xls, "geonames": p_geonames, "wikidata": p_wikidata,
}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--run", required=True, help="comma-separated probe names: " + ", ".join(PROBES))
    ap.add_argument("--url", action="append", default=[])
    ap.add_argument("--wayback", help="read --url from this Internet Archive capture, raw")
    ap.add_argument("--grep", help="regex a link or page must match")
    ap.add_argument("--pages", help="PDF pages to print, e.g. 1-3,9")
    ap.add_argument("--chars", type=int, default=4000)
    ap.add_argument("--rows", type=int, default=40)
    ap.add_argument("--start", type=int, default=0, help="first workbook row to print")
    ap.add_argument("--width", type=int, default=24, help="characters per cell")
    ap.add_argument("--cols", type=int, default=16, help="cells per row")
    ap.add_argument("--boxes", help="osm_buildings: only these boxes, comma-separated")
    ap.add_argument("--point", action="append", default=[],
                    help="green_line: label,lat,lon (repeatable; no spaces)")
    ap.add_argument("--box", help="green_line: south,west,north,east to clip outlines to")
    args = ap.parse_args()
    for name in args.run.split(","):
        if name not in PROBES:
            raise SystemExit(f"cyprus_north_probe: no probe {name!r}; have {sorted(PROBES)}")
        try:
            PROBES[name](args)
        except Exception as exc:                          # noqa: BLE001 -- reported
            log(f"!! {name}: {type(exc).__name__}: {str(exc)[:300]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
