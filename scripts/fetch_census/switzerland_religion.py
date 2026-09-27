#!/usr/bin/env python3
"""Switzerland: religion by canton, from BFS's structural survey (a survey estimate).

Since 2010 the Swiss census asks religion only in the structural survey
(Strukturerhebung), an annual sample of the permanent resident population
aged 15 and over, weighted by BFS. Its table je-d-01.08.02.02, "Religions-
zugehörigkeit nach Grossregion und Kanton" (DAM asset 36347568), gives one
sheet per year with the weighted number of people in each religious group
and each estimate's confidence interval, as a percentage of the estimate.

The records follow the rules the language table's reader
(``switzerland.py``) already applies to the same survey: an estimate whose
interval is wider than ±25% of itself is dropped rather than shown, and a
cell BFS suppressed ("X") is left out, so a canton's shares can fall short of
100. They are survey estimates and say so in ``religion_basis``; the output
is a ``*_survey.json`` file, which the build uses only where no count exists.

Usage:
    python -m scripts.fetch_census.switzerland_religion [--year 2024]
"""

from __future__ import annotations

import argparse
import io
from typing import Any

from ._shared import PROCESSED, dated, http_get, log, record, shares, write_json
from .central_ages import fold, units

DATA = "https://dam-api.bfs.admin.ch/hub/api/dam/assets/36347568/master"
PORTAL = "https://www.bfs.admin.ch/asset/de/36347568"
SOURCE = ("Federal Statistical Office (BFS/OFS), Strukturerhebung, je-d-01.08.02.02 "
          "(Religionszugehörigkeit nach Grossregion und Kanton)")
LICENCE = "BFS open data (free reuse with attribution)"
OUT = PROCESSED / "switzerland_religion_survey.json"
MAX_INTERVAL = 25.0
RELIGION = {
    "Evangelisch-reformiert": "Reformed", "Römisch-katholisch": "Roman Catholic",
    "Andere christliche Glaubensgemeinschaften": "Other Christian",
    "Jüdische Glaubensgemeinschaften": "Judaism", "Jüdische Glaubensgemeinschaft": "Judaism",
    "Islamische Glaubensgemeinschaften": "Islam", "Islamische Glaubensgemeinschaft": "Islam",
    "Andere Religionsgemeinschaften": "Other religion",
    "Ohne Religionszugehörigkeit": "No religion", "Religionszugehörigkeit unbekannt": "Not stated",
    "Unbekannt": "Not stated", "Ohne Angabe": "Not stated",
}
ALIASES = {"Appenzell A. Rh.": "Appenzell Ausserrhoden", "Appenzell I. Rh.": "Appenzell Innerrhoden"}


def clean(cell: Any) -> str:
    return " ".join(str(cell if cell is not None else "").split())


def read(year: int) -> dict[str, dict[str, tuple[float | None, float | None]]]:
    """{row label: {category: (estimate, interval %)}}; None where BFS printed X."""
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(http_get(DATA, binary=True, timeout=300)),
                                  read_only=True, data_only=True)
    if str(year) not in book.sheetnames:
        raise SystemExit(f"switzerland_religion: no sheet {year}; sheets {book.sheetnames}")
    rows = [list(r) for r in book[str(year)].iter_rows(values_only=True)]
    sub_at = next(i for i, r in enumerate(rows) if any(clean(c).startswith("Anzahl") for c in r))
    head, sub = [clean(c) for c in rows[sub_at - 1]], [clean(c) for c in rows[sub_at]]
    columns: dict[int, str] = {}
    label = ""
    for j, cell in enumerate(head):
        label = cell or label
        if sub[j].startswith("Anzahl") or (label == "Total" and head[j] == "Total"):
            columns[j] = label
    log(f"  {year}: columns {columns}")
    unknown = [c for c in columns.values() if c != "Total" and c not in RELIGION]
    if unknown:
        raise SystemExit(f"switzerland_religion: categories this reader does not know: {unknown}")

    def value(cell: Any) -> float | None:
        text = clean(cell)
        return None if text in ("X", "x", "") else float(cell)

    out: dict[str, dict[str, tuple[float | None, float | None]]] = {}
    for r in rows[sub_at + 1:]:
        name = clean(r[0])
        if not name or all(c is None for c in r[1:]):
            continue
        entry = {}
        for j, category in columns.items():
            ci = value(r[j + 1]) if category != "Total" and j + 1 < len(r) else None
            entry[category] = (value(r[j]), ci)
        if name in out and out[name] != entry:
            raise SystemExit(f"switzerland_religion: {name} is printed twice with different figures")
        out[name] = entry
    return out


def build(year: int) -> list[dict[str, Any]]:
    log(f"switzerland_religion: BFS structural survey, religion by canton, {year}")
    table = read(year)
    shapes = units("CHE", "admin1")
    by_name = {fold(s["name"]): s for s in shapes}
    records, used = [], set()
    note = (f"BFS structural survey {year} (Strukturerhebung): religious affiliation of the "
            f"permanent resident population aged 15 and over, weighted by BFS. It is a sample: "
            f"estimates whose confidence interval is wider than ±{MAX_INTERVAL:.0f}% of the "
            f"estimate are dropped, and cells BFS suppresses for too few respondents are left "
            f"out, so the shares can fall short of 100. BFS does not print the number of "
            f"respondents per canton in this table. The census has not asked every resident "
            f"since 2000.")
    for name, cells in table.items():
        parts = [p.strip() for p in ALIASES.get(name, name).split("/")]
        shape = next((by_name[fold(p)] for p in parts if fold(p) in by_name), None)
        if shape is None:
            continue                      # the country and its Grossregionen
        total = cells["Total"][0]
        if not total:
            raise SystemExit(f"switzerland_religion: {name} has no total")
        published = sum(v for k, (v, _) in cells.items() if k != "Total" and v is not None)
        suppressed = [k for k, (v, _) in cells.items() if v is None]
        if published > total * 1.001 or (not suppressed and published < total * 0.999):
            raise SystemExit(f"switzerland_religion: {name}: groups make {published:,.0f} of "
                             f"{total:,.0f} (suppressed {suppressed})")
        counts: dict[str, float] = {}
        dropped = []
        for german, (v, ci) in cells.items():
            if german == "Total" or v is None:
                continue
            if ci is not None and ci > MAX_INTERVAL:
                dropped.append(f"{RELIGION[german]} ±{ci:.0f}%")
                continue
            counts[RELIGION[german]] = counts.get(RELIGION[german], 0.0) + v
        rows = shares(counts, total=total)
        used.add(shape["id"])
        log(f"  {shape['name']}: {total:,.0f} aged 15+; dropped {dropped}; suppressed {suppressed}")
        records.append(record(
            f"CHE-SE{year}-{fold(shape['name'])}", shape["name"], level="admin1", parent="CHE",
            country="CHE", match_by="shape_id", shape_id=shape["id"],
            sources=[{"field": "religion", "name": SOURCE, "url": PORTAL, "license": LICENCE,
                      "year": year}],
            religion=rows, religion_year=dated(rows, year),
            religion_basis="survey estimate: structural survey, permanent resident population 15+",
            religion_note=note))
    left = sorted(s["name"] for s in shapes if s["id"] not in used)
    if left:
        raise SystemExit(f"switzerland_religion: cantons with no row: {left}")
    return records


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--year", type=int, default=2024)
    args = ap.parse_args()
    records = build(args.year)
    write_json(OUT, records)
    log(f"  wrote {OUT.name}: {len(records)} records")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
