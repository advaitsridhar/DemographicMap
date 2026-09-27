#!/usr/bin/env python3
"""Malta: the 2021 census by locality -- population, median age, sex ratio
and religion.

Malta's 68 localities are both of this map's levels (the boundary file draws
the same polygons at the first and the second), and they had a head count
from Wikipedia infoboxes and nothing else. The National Statistics Office's
Census of Population and Housing 2021, Volume 1, publishes its tables as
workbooks, one per chapter:

* Chapter 1, "Population": for every locality a sheet of its residents by sex
  and **single year of age** (less than 1 to 89, and over 89), Tables 1.5;
  and Table 1.2, each locality's total by sex, to check them against.
* Chapter 5, "Religious affiliation": Table 5.3, residents aged 15 and over
  by locality and religion -- Roman Catholicism, Islam, Orthodoxy, Hinduism,
  Church of England, Protestantism, Buddhism, Judaism, other religious groups
  and no religious affiliation.

**Read from the Internet Archive.** nso.gov.mt answers the runner 403
Forbidden, page and file alike. The archive holds the office's own
workbooks at their own addresses (captured 9 December 2023 and 2 July 2023),
and those captures are read byte for byte (``id_``), not a copy.

The census asked two more questions this map shows:

* "What is your racial origin?" (Q11): Volume 1, Table 4.3, every resident by
  locality. No workbook of that chapter was captured, so the table is read off
  the report's own pages (the archive's capture of the Volume 1 PDF), and
  written as the ethnicity field with ``ethnicity_basis`` "racial origin".
* "What is the main language that you grew up speaking from early
  childhood?" (aged 5+): tabulated by locality only for Maltese citizens
  (Volume 3, Table 3.6, read off the report's pages as Table 4.3 is), and for
  everyone only by district. The Maltese citizens' languages are written as
  shares of *all* the locality's residents aged 5 and over (from the single
  years of Chapter 1), so the composition adds to the citizens' share and the
  other residents -- whose language the NSO tabulates by district only -- are
  left visibly unaccounted for rather than assumed to speak like citizens.

**Names.** The census writes each locality's Maltese name with its article
("Il-Birgu", "Ħal Qormi", "Raħal Ġdid"); the boundary file writes a short
or English one ("Birgu", "Qormi", "Paola"). The 68 pairs are declared below
and checked one to one against both sides.

Checks, each of which stops the run:

* every locality's single years make its total for each sex, and its sexes
  make its total, and those totals are Table 1.2's;
* the localities make Malta's 519,562;
* the NSO publishes an average age, not a median (Table 1.2): Malta's mean
  recomputed from the localities' single years must be within 0.6 years of
  it (the half-year is how a completed age is read);
* every locality's religion categories make its published total aged 15 and
  over, and the localities make Malta's.

Usage:
    python -m scripts.fetch_census.malta_census
"""

from __future__ import annotations

import argparse
import io
import json
import re
from collections import Counter
from typing import Any

from ._shared import (
    PROCESSED, dated, http_get, log, measure, record, shares, write_json,
)

OUT = "malta_census.json"
YEAR = 2021
SOURCE = "NSO Malta, Census of Population and Housing 2021, Volume 1"
LICENCE = "NSO Malta (attribution)"
SITE = PROCESSED.parent.parent / "site" / "data"
ARCHIVE = "https://web.archive.org/web/{stamp}id_/https://nso.gov.mt/wp-content/uploads/{name}"
CHAPTER_1 = ARCHIVE.format(stamp="20231209013436", name="Census-Vol-1_Chapter-1.xlsx")
CHAPTER_5 = ARCHIVE.format(stamp="20230702041150", name="Census-Vol-1_Chapter-5.xlsx")
VOLUME_1 = ARCHIVE.format(stamp="20250628190541",
                          name="Census-of-Population-2021-volume1.pdf")
PAGE = "https://nso.gov.mt/census-of-population-and-housing-2021-final-report/"
NATIONAL = 519562

# Table 4.3 writes Gozo's Żebbuġ without the ", Għawdex" the workbooks give it.
SPELLINGS = {"Iż-Żebbuġ": "Iż-Żebbuġ, Għawdex"}

VOLUME_3 = ARCHIVE.format(stamp="20240129150900",
                          name="volume3-Census-of-Population-2021.pdf")
LANGUAGES = ("Maltese", "English", "Italian", "German", "French", "Arabic", "Other language")


def language_note(citizens: float, residents: float) -> str:
    return (
        f"Census {YEAR} question \"What is the main language that you grew up speaking from early "
        "childhood?\", asked of residents aged 5 and over. The NSO tabulates it by locality for "
        f"Maltese citizens only (Volume 3, Table 3.6): {citizens:,.0f} of this locality's "
        f"{residents:,.0f} residents aged 5 and over. Their languages are shown as shares of all "
        f"{residents:,.0f}, so the {residents - citizens:,.0f} residents who are not Maltese "
        "citizens -- tabulated by district only -- are not accounted for here.")


# The census's name -> the boundary file's.
LOCALITIES = {
    "Birkirkara": "Birkirkara", "Birżebbuġa": "Birżebbuġa", "Bormla": "Bormla",
    "Floriana": "Floriana", "Għajnsielem and Comino": "Ghajnsielem",
    "Ħad-Dingli": "Dingli", "Ħal Balzan": "Balzan", "Ħal Għargħur": "Għargħur",
    "Ħal Għaxaq": "Għaxaq", "Ħal Kirkop": "Kirkop", "Ħal Lija": "Lija", "Ħal Luqa": "Luqa",
    "Ħal Qormi": "Qormi", "Ħal Safi": "Safi", "Ħal Tarxien": "Tarxien", "Ħ'Attard": "Attard",
    "Ħaż-Żabbar": "Żabbar", "Ħaż-Żebbuġ": "Żebbuġ Malta", "Il-Birgu": "Birgu",
    "Il-Fgura": "Fgura", "Il-Fontana": "Fontana", "Il-Gudja": "Gudja", "Il-Gżira": "Gżira",
    "Il-Ħamrun": "Ħamrun", "Il-Kalkara": "Kalkara", "Il-Marsa": "Marsa",
    "Il-Mellieħa": "Mellieħa", "Il-Mosta": "Mosta", "Il-Munxar": "Munxar", "Il-Qala": "Qala",
    "Il-Qrendi": "Qrendi", "In-Nadur": "Nadur", "In-Naxxar": "Naxxar",
    "Ir-Rabat": "Rabat Malta", "Ir-Rabat, Għawdex": "Rabat Gozo",
    "Is-Siġġiewi": "Siġġiewi", "Is-Swieqi": "Swieqi", "Ix-Xagħra": "Xagħra",
    "Ix-Xewkija": "Xewkija", "Ix-Xgħajra": "Xgħajra", "Iż-Żebbuġ, Għawdex": "Żebbuġ Gozo",
    "Iż-Żejtun": "Żejtun", "Iż-Żurrieq": "Żurrieq", "L-Għarb": "Gharb", "L-Għasri": "Ghasri",
    "L-Iklin": "Iklin", "L-Imdina": "Mdina", "L-Imġarr": "Mġarr", "L-Imqabba": "Mqabba",
    "L-Imsida": "Msida", "L-Imtarfa": "Mtarfa", "L-Isla": "Isla", "Marsaskala": "Marsaskala",
    "Marsaxlokk": "Marsaxlokk", "Pembroke": "Pembroke", "Raħal Ġdid": "Paola",
    "San Ġiljan": "Saint Julian's", "San Ġwann": "Saint John", "San Lawrenz": "Saint Lawerence",
    "San Pawl Il-Baħar": "Saint Paul's Bay", "Santa Luċija": "Saint Lucia's",
    "Santa Venera": "Santa Venera", "Ta' Kerċem": "Kerċem", "Ta' Sannat": "Sannat",
    "Ta' Xbiex": "Ta' Xbiex", "Tal-Pieta'": "Pietà", "Tas-Sliema": "Sliema",
    "Valletta": "Valletta",
}

RELIGIONS = {
    "Roman Catholicism": "Roman Catholic",
    "Islam": "Muslim",
    "Orthodoxy": "Orthodox",
    "Hinduism": "Hindu",
    "Church of England": "Church of England",
    "Protestantism": "Protestant",
    "Buddhism": "Buddhist",
    "Judaism": "Jewish",
    "Other religious groups": "Other religion",
    "No religious affiliation": "No religion",
}


def workbook(url: str):
    import openpyxl
    blob = http_get(url, binary=True, cache=True, timeout=300)
    return openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)


def cell(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("‐", "-")).strip()


def number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def single_years(rows: list[list[Any]], where: str) -> dict[str, Any]:
    """A Table 1.5 sheet: two blocks of Age | Males | Females | Total."""
    men, women = Counter(), Counter()
    total = None
    for row in rows:
        for start in (0, 5):
            if len(row) < start + 4:
                continue
            label = cell(row[start])
            m, f, t = (number(v) for v in row[start + 1:start + 4])
            if label == "Total" and start == 0 and t is not None:
                total = (m, f, t)
                continue
            if t is None or "-" in label:
                continue          # a decade's subtotal, or a blank
            # A sex with nobody of that age is printed as a dash, not a 0.
            m, f = m or 0.0, f or 0.0
            if label == "Less than 1":
                age = 0
            elif label.startswith("Over "):
                age = int(label.split()[1]) + 1
            elif label.isdigit():
                age = int(label)
            else:
                continue
            if m + f != t:
                raise SystemExit(f"malta_census: {where} age {label}: {m:.0f} + {f:.0f} "
                                 f"is not {t:.0f}")
            men[age] += m
            women[age] += f
    if total is None:
        raise SystemExit(f"malta_census: {where}: no Total row")
    if sum(men.values()) != total[0] or sum(women.values()) != total[1]:
        raise SystemExit(f"malta_census: {where}: single years make {sum(men.values()):,.0f} "
                         f"men and {sum(women.values()):,.0f} women against {total}")
    # An age nobody in the locality has is printed as dashes and read as no
    # row at all; what must hold is that every age read is 0 to 89 or "over 89".
    if not set(men) | set(women) <= set(range(0, 91)):
        raise SystemExit(f"malta_census: {where}: ages {sorted(set(men) | set(women))}")
    return {"men": total[0], "women": total[1], "ages": men + women}


def read_chapter_1() -> dict[str, dict[str, Any]]:
    book = workbook(CHAPTER_1)
    table_12: dict[str, tuple[Any, ...]] = {}
    for row in book["T1.2_pop_sex_age_locality"].iter_rows(values_only=True):
        name = SPELLINGS.get(cell(row[0]), cell(row[0])) if row else ""
        m, f, t, _, _, _, mean = (number(v) for v in (list(row) + [None] * 8)[1:8])
        if name and t is not None:
            table_12[name] = (m, f, t, mean)
    out: dict[str, dict[str, Any]] = {}
    for sheet in book.sheetnames:
        if not re.match(r"^\d+-", sheet):
            continue
        rows = [list(r) for r in book[sheet].iter_rows(values_only=True)]
        name = next((cell(r[0]) for r in rows[1:4] if r and cell(r[0])), "")
        if name not in LOCALITIES:
            raise SystemExit(f"malta_census: sheet {sheet!r} is for {name!r}, which is not "
                             "a declared locality")
        entry = single_years(rows, name)
        published = table_12.get(name)
        if not published or published[:2] != (entry["men"], entry["women"]):
            raise SystemExit(f"malta_census: {name}: Table 1.5 says {entry['men']:,.0f} men "
                             f"and {entry['women']:,.0f} women, Table 1.2 {published}")
        out[name] = entry
    if sorted(out) != sorted(LOCALITIES):
        raise SystemExit(f"malta_census: localities with no sheet: "
                         f"{sorted(set(LOCALITIES) - set(out))}")
    whole = sum(e["men"] + e["women"] for e in out.values())
    if whole != NATIONAL or table_12.get("MALTA", (0, 0, 0))[2] != NATIONAL:
        raise SystemExit(f"malta_census: the localities make {whole:,.0f}, Malta is "
                         f"{NATIONAL:,}")
    log(f"  chapter 1: {len(out)} localities make {whole:,.0f}")
    out["_mean"] = table_12["MALTA"][3]
    return out


def read_religion() -> dict[str, dict[str, Any]]:
    book = workbook(CHAPTER_5)
    rows = [list(r) for r in book["5.3"].iter_rows(values_only=True)]
    header_at = next(i for i, r in enumerate(rows)
                     if r and cell(r[0]).startswith("District and locality"))
    header = [cell(v) for v in rows[header_at]]
    cols = {i: RELIGIONS[h] for i, h in enumerate(header) if h in RELIGIONS}
    if len(cols) != len(RELIGIONS) or "Total" not in header:
        raise SystemExit(f"malta_census: Table 5.3's columns are {header}")
    total_at = header.index("Total")
    out: dict[str, dict[str, Any]] = {}
    for row in rows[header_at + 1:]:
        name = SPELLINGS.get(cell(row[0]), cell(row[0])) if row else ""
        if name not in LOCALITIES and name != "MALTA":
            continue
        counts = {label: number(row[i]) or 0.0 for i, label in cols.items()}
        total = number(row[total_at])
        if total is None or sum(counts.values()) != total:
            raise SystemExit(f"malta_census: {name}: religions make {sum(counts.values()):,.0f} "
                             f"against {total}")
        out[name] = {"counts": counts, "total": total}
    missing = sorted(set(LOCALITIES) - set(out))
    if missing:
        raise SystemExit(f"malta_census: Table 5.3 has no row for {missing}")
    made = sum(out[n]["total"] for n in LOCALITIES)
    if made != out["MALTA"]["total"]:
        raise SystemExit(f"malta_census: localities make {made:,.0f} aged 15+, Malta "
                         f"{out['MALTA']['total']:,.0f}")
    log(f"  chapter 5: {len(LOCALITIES)} localities, {made:,.0f} aged 15 and over")
    return out


RACIAL_ORIGINS = ("Caucasian", "Asian", "Arab", "African", "Hispanic or Latino",
                  "More than one racial origin")
ROW = re.compile(r"^(?P<name>\D+?)\s+" + r"\s+".join([r"([\d,]+)"] * 7) + r"\s*$")


def read_racial_origin() -> dict[str, dict[str, Any]]:
    """Table 4.3 of the Volume 1 report, read off its two pages.

    No workbook of Chapter 4 was captured, so the report's own table is read:
    each line is a locality and seven counts -- Caucasian, Asian, Arab,
    African, Hispanic or Latino, more than one racial origin, and the total.
    """
    import pdfplumber
    blob = http_get(VOLUME_1, binary=True, cache=True, timeout=600)
    out: dict[str, dict[str, Any]] = {}
    with pdfplumber.open(io.BytesIO(blob)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if "TABLE 4.3. Total population by racial origin and locality" not in text:
                continue
            for line in text.splitlines():
                m = ROW.match(line.replace("‐", "-").strip())
                if not m:
                    continue
                name = SPELLINGS.get(m.group("name").strip(), m.group("name").strip())
                values = [float(v.replace(",", "")) for v in m.groups()[1:]]
                counts = dict(zip(RACIAL_ORIGINS, values[:6]))
                if sum(counts.values()) != values[6]:
                    raise SystemExit(f"malta_census: Table 4.3 {name}: origins make "
                                     f"{sum(counts.values()):,.0f} against {values[6]:,.0f}")
                out[name] = {"counts": counts, "total": values[6]}
    missing = sorted(set(LOCALITIES) - set(out))
    if missing:
        raise SystemExit(f"malta_census: Table 4.3 has no row for {missing}")
    if sum(out[n]["total"] for n in LOCALITIES) != NATIONAL or out["MALTA"]["total"] != NATIONAL:
        raise SystemExit("malta_census: Table 4.3's localities do not make Malta")
    log(f"  table 4.3: {len(LOCALITIES)} localities' racial origin make {NATIONAL:,}")
    return out


LANGUAGE_VALUE = re.compile(r"^(?:[\d,]+|-+)$")


def language_row(line: str) -> tuple[str, list[float]] | None:
    """One line of Table 3.6: a locality and seven languages and a total.

    The report prints an empty cell as a dash, and pdf text runs two or more
    adjacent dashes together ("--" is two empty cells), so a run of dashes is
    read as that many zeros; the row's own total then checks the reading.
    """
    tokens = line.replace("\u2010", "-").split()
    at = len(tokens)
    while at > 0 and LANGUAGE_VALUE.match(tokens[at - 1]):
        at -= 1
    name, cells = " ".join(tokens[:at]), tokens[at:]
    values: list[float] = []
    for token in cells:
        if set(token) == {"-"}:
            values.extend([0.0] * len(token))
        else:
            values.append(float(token.replace(",", "")))
    if not name or len(values) != len(LANGUAGES) + 1:
        return None
    return name, values


def read_language() -> dict[str, dict[str, Any]]:
    """Volume 3, Table 3.6: Maltese citizens aged 5+ by main language and locality."""
    import pdfplumber
    blob = http_get(VOLUME_3, binary=True, cache=True, timeout=600)
    out: dict[str, dict[str, Any]] = {}
    folded = {name.lower(): name for name in LOCALITIES}
    with pdfplumber.open(io.BytesIO(blob)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if "TABLE 3.6. Maltese population aged 5 and over" not in text:
                continue
            for line in text.splitlines():
                row = language_row(line)
                if row is None:
                    continue
                name, values = row
                name = SPELLINGS.get(name, name)
                name = folded.get(name.lower(), name)
                counts = dict(zip(LANGUAGES, values[:-1]))
                if sum(counts.values()) != values[-1]:
                    raise SystemExit(f"malta_census: Table 3.6 {name}: languages make "
                                     f"{sum(counts.values()):,.0f} against {values[-1]:,.0f}")
                if name in out and out[name]["total"] != values[-1]:
                    raise SystemExit(f"malta_census: Table 3.6 prints {name} twice, differently")
                out[name] = {"counts": counts, "total": values[-1]}
    missing = sorted(set(LOCALITIES) - set(out))
    if missing:
        raise SystemExit(f"malta_census: Table 3.6 has no row for {missing}")
    whole = out.get("Total")
    if whole is None:
        raise SystemExit("malta_census: Table 3.6 has no Total row")
    for key in LANGUAGES:
        made = sum(out[n]["counts"][key] for n in LOCALITIES)
        if made != whole["counts"][key]:
            raise SystemExit(f"malta_census: Table 3.6's localities make {made:,.0f} {key}, "
                             f"its total {whole['counts'][key]:,.0f}")
    log(f"  table 3.6: {len(LOCALITIES)} localities, {whole['total']:,.0f} Maltese citizens "
        "aged 5 and over")
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    ages = read_chapter_1()
    published_mean = ages.pop("_mean")
    religion = read_religion()
    origin = read_racial_origin()
    language = read_language()
    for name in LOCALITIES:
        if origin[name]["total"] != ages[name]["men"] + ages[name]["women"]:
            raise SystemExit(f"malta_census: {name}: Table 4.3 counts "
                             f"{origin[name]['total']:,.0f}, Table 1.2 "
                             f"{ages[name]['men'] + ages[name]['women']:,.0f}")
    national = sum((e["ages"] for e in ages.values()), Counter())
    people = sum(national.values())
    mean = sum((a + 0.5) * n for a, n in national.items()) / people
    log(f"  Malta: median {median(national)}, mean {mean:.2f} (Table 1.2 publishes "
        f"{published_mean:.2f})")
    if published_mean is None or abs(mean - published_mean) > 0.6:
        raise SystemExit("malta_census: the single years do not give the published mean age")

    admin1 = {u["name"]: u for u in json.loads((SITE / "admin1" / "MLT.units.json").read_text())}
    admin2 = {u["name"]: u for u in json.loads((SITE / "admin2" / "MLT.units.json").read_text())}
    if sorted(admin1) != sorted(LOCALITIES.values()) or sorted(admin2) != sorted(admin1):
        raise SystemExit("malta_census: the map's localities are not the 68 declared: "
                         f"{sorted(set(admin1) ^ set(LOCALITIES.values()))}")

    records: list[dict[str, Any]] = []
    for name, drawn in sorted(LOCALITIES.items(), key=lambda kv: kv[1]):
        a, r, o = ages[name], religion[name], origin[name]
        rows = shares(r["counts"], total=r["total"])
        origins = shares(o["counts"], total=o["total"])
        residents = sum(n for age, n in a["ages"].items() if age >= 5)
        citizens = language[name]["total"]
        if citizens > residents:
            raise SystemExit(f"malta_census: {name}: {citizens:,.0f} Maltese citizens aged 5+ "
                             f"against {residents:,.0f} residents aged 5+")
        spoken = shares({k: v for k, v in language[name]["counts"].items() if v > 0},
                        total=residents)
        for level, shapes in (("admin1", admin1), ("admin2", admin2)):
            records.append(record(
                f"MLT-NSO-{level}-{drawn}", drawn, level=level, parent="MLT", country="MLT",
                match_by="shape_id", shape_id=shapes[drawn]["id"],
                aliases=[name] if name != drawn else [],
                population=measure(int(a["men"] + a["women"]), year=YEAR, source=SOURCE),
                population_note=(f"Residents of {name} counted by the census of 21 November "
                                 f"{YEAR} (Table 1.2)."),
                median_age=measure(median(a["ages"]), unit="years", year=YEAR, source=SOURCE),
                median_age_note=("Interpolated within the single year of age holding the "
                                 "middle person, from Table 1.5's residents by sex and single "
                                 "year of age (over 89 as one open group)."),
                sex_ratio=measure(round(1000 * a["men"] / a["women"]),
                                  unit="males_per_1000_females", year=YEAR, source=SOURCE),
                religion=rows,
                religion_year=dated(rows, YEAR),
                religion_note=(f"Census {YEAR} question \"What religion, religious "
                               "denomination or body do you belong to?\", residents aged 15 "
                               "and over (Table 5.3). Every respondent is in one of the ten "
                               "categories; 'Other religion' is the NSO's 'other religious "
                               "groups'."),
                ethnicity=origins,
                ethnicity_year=dated(origins, YEAR),
                ethnicity_note=(f"Census {YEAR} question \"What is your racial origin?\" "
                                "(through one's biological parents), all residents (Volume 1, "
                                "Table 4.3). The questionnaire also offered 'Other'; the NSO "
                                "publishes none, having coded those answers into the six "
                                "groups shown. Citizenship is a separate question."),
                ethnicity_basis="racial origin",
                language=spoken,
                language_year=dated(spoken, YEAR),
                language_note=language_note(citizens, residents),
                sources=[{"field": "population/median age/sex ratio", "name": SOURCE,
                          "url": CHAPTER_1, "year": YEAR, "license": LICENCE},
                         {"field": "religion", "name": SOURCE, "url": CHAPTER_5,
                          "year": YEAR, "license": LICENCE},
                         {"field": "ethnicity", "name": SOURCE, "url": VOLUME_1,
                          "year": YEAR, "license": LICENCE},
                         {"field": "language", "name": SOURCE.replace("Volume 1", "Volume 3"),
                          "url": VOLUME_3, "year": YEAR, "license": LICENCE}]))
    write_json(args.out or PROCESSED / OUT, records)
    log(f"  {len(records)} records ({len(LOCALITIES)} localities at both levels)")
    return 0


def median(ages: Counter) -> float | None:
    from .redatam import median_age
    return median_age(ages)


if __name__ == "__main__":
    raise SystemExit(main())
