#!/usr/bin/env python3
"""Morocco -- median age at the 2024 census, by region and province, where HCP
publishes the population by five-year age group.

The Haut-Commissariat au Plan's results dashboard and provincial atlas give
every province's ages only as the shares under 15, 15-59 and 60 and over, from
which no median can be had. Its regional directorates publish more:

* a workbook of the region's demographic and socio-economic indicators, every
  region, province, prefecture and commune with the share of its municipal
  population in each five-year age group (``WORKBOOKS``: Rabat-Salé-Kénitra,
  Fès-Meknès), sheet ``Population``, rows by HCP's geographic code;
* provincial notes whose annex gives the same shares for the province and the
  region side by side (``ANNEXES``: five provinces of Tanger-Tétouan-Al
  Hoceima);
* notes on the region and each province of Souss-Massa, with a table of the
  population by five-year age group and sex (``SERIES``).

From those shares (to 0.1%) the median is interpolated within the five-year
group that holds the middle person.

**Binding** is the legal-population reader's (``morocco_rgph2024``), by HCP's
geographic code for the workbooks and by name for the notes. Guelmim-Oued
Noun's median, which HCP gives in a note on the southern provinces, counts
Al Mahbass, a commune across the Western Sahara line, and is stated, not
shown. Every other region and province bound says what HCP gives instead.

**Checks before anything is written.** Every unit's shares must run from 0-4
without a gap to an open last group and make 100% to within rounding; the
region's shares in each note must agree; and every unit's shares must make
the dashboard's own shares under 15 and of 60 and over, to within 0.6 points,
which ties each table to the unit it is bound to.

Usage:
    python -m scripts.fetch_census.morocco_rgph2024_ages
"""

from __future__ import annotations

import argparse
import io
import json
import re
from typing import Any

from . import morocco_rgph2024 as legal
from ._shared import NOT_AVAILABLE, PROCESSED, gap, http_get, log, measure, record, write_json

YEAR = 2024
OUT = "morocco_rgph2024_ages.json"
LICENCE = legal.LICENCE
HCP = "Haut-Commissariat au Plan (HCP)"
# Region (as HCP names it) -> (workbook, its landing page, its title).
WORKBOOKS = {
    "Rabat-Salé-Kénitra": (
        "https://www.hcp.ma/region-rabat/docs/docs/RGPH2024-Indicateurs%20demographiques%20et"
        "%20socioeconomiques%20de%20la%20region%20de%20Rabat.xlsx",
        "https://www.hcp.ma/region-rabat/",
        "RGPH 2024, indicateurs démographiques et socio-économiques de la région de "
        "Rabat-Salé-Kénitra"),
    "Fès-Meknès": (
        "https://www.hcp.ma/region-fes/attachment/2679745/",
        "https://www.hcp.ma/region-fes/Les-Principaux-Resultats-de-la-Region-Fes-Meknes-selon-"
        "le-RGPH-2024_a117.html",
        "Les principaux résultats de la région Fès-Meknès selon le RGPH 2024"),
}
# Province (as HCP's legal-population workbook names it) -> its note, whose
# annex gives the province and the Tanger-Tétouan-Al Hoceima region.
ANNEXES = {
    "Province de Tétouan": "https://www.hcp.ma/file/244701/",
    "Province d'Ouezzane": "https://www.hcp.ma/file/244700/",
    "Province de Larache": "https://www.hcp.ma/file/244699/",
    "Préfecture de Tanger-Assilah": "https://www.hcp.ma/file/244698/",
    "Préfecture de M'diq-Fnideq": "https://www.hcp.ma/file/244697/",
}
# Souss-Massa's regional directorate publishes one such note for the region and
# each province, its table "Répartition selon le groupe d'âges quinquennal"
# giving men, women and both: note -> the unit's HCP name, or names to tell two
# prefectures apart by the note's own words.
SERIES = {
    "https://www.hcp.ma/region-agadir/attachment/2802780/": ("Région de Souss-Massa",),
    "https://www.hcp.ma/region-agadir/attachment/2802781/": (
        "Préfecture d'Agadir-Ida-Ou-Tanane", "Préfecture d'Inezgane-Aït Melloul"),
    "https://www.hcp.ma/region-agadir/attachment/2802782/": (
        "Préfecture d'Agadir-Ida-Ou-Tanane", "Préfecture d'Inezgane-Aït Melloul"),
    "https://www.hcp.ma/region-agadir/attachment/2802783/": ("Province de Chtouka-Aït Baha",),
    "https://www.hcp.ma/region-agadir/attachment/2802784/": ("Province de Taroudannt",),
    "https://www.hcp.ma/region-agadir/attachment/2802862/": ("Province de Tata",),
    "https://www.hcp.ma/region-agadir/attachment/2802786/": ("Province de Tiznit",),
}
SERIES_TITLE = ("RGPH 2024, caractéristiques démographiques et socio-économiques de la "
                "population, notes de la région de Souss-Massa et de ses préfectures et provinces")
SERIES_LANDING = ("https://www.hcp.ma/region-agadir/br-Caracteristiques-demographiques-et-socio-"
                  "economiques-de-la-population-RGPH-2024-donnees-regionales-et-provinciales_a210"
                  ".html")
# The dashboard's shares under 15 and of 60 and over, against which every
# unit's five-year shares are checked (to within this many points).
BROAD = {"Part de la population de moins de 15 ans (%)": (0, 14),
         "Part de la population de 60 ans et plus (%)": (60, None)}
BROAD_SLACK = 0.6
ANNEX_REGION = "Tanger-Tétouan-Al Hoceima"
ANNEX_TITLE = ("RGPH 2024, caractéristiques démographiques et socioéconomiques de la population, "
               "notes provinciales de la région de Tanger-Tétouan-Al Hoceima, annexe")
LANDING_DOWNLOADS = "https://www.hcp.ma/downloads/RGPH-2024_t22752.html"
# A region whose published median counts ground across the Western Sahara line.
STATED = {
    "Guelmim-Oued Noun": (
        "HCP's note on the southern provinces of the Kingdom (Les provinces du Sud du Royaume: "
        "données démographiques, September 2025) puts the region's median age at 31.3 years in "
        "2024, against 27.0 in 2014; that median counts the commune of Al Mahbass, which lies "
        "south of the Western Sahara line and outside the region as drawn here, and no median "
        "for the region without it is published, so none is shown."),
}
ONLY_SHARES = ("HCP's results dashboard and provincial atlas give this unit's 2024 ages only as "
               "the shares under 15, 15-59 and 60 and over, from which no median can be had.")
AGE = re.compile(r"^\s*(\d{1,2})\s*-\s*(\d{1,2})\s*(?:ans)?\s*$")
OPEN = re.compile(r"^\s*(\d{1,2})\s*(?:ans)?\s*(?:et|ou)\s*(?:plus|\+)\s*$|^\s*(\d{1,2})\s*\+\s*$",
                  re.I)
ROW = re.compile(r"^\s*(\d{1,2}\s*-\s*\d{1,2}|\d{1,2}\s*(?:ans)?\s*(?:et|ou)\s*(?:plus|\+))"
                 r"(?:\s*ans)?\s+([\d,.\s]+?)\s*$", re.I)


def bounds(label: Any) -> tuple[int, int | None] | None:
    text = str(label or "").strip()
    if m := AGE.match(text):
        return int(m.group(1)), int(m.group(2))
    if m := OPEN.match(text):
        return int(m.group(1) or m.group(2)), None
    return None


def median(groups: list[tuple[int, int | None, float]]) -> float | None:
    """The median of (low, high, share) groups; None if it falls in the open one."""
    total = sum(s for _, _, s in groups)
    half, before = total / 2, 0.0
    for low, high, share in groups:
        if before + share >= half:
            if high is None or share <= 0:
                return None
            return round(low + (half - before) / share * (high - low + 1), 1)
        before += share
    return None


def check(name: str, groups: list[tuple[int, int | None, float]]) -> None:
    edge = 0
    for low, high, _ in groups:
        if low != edge:
            raise SystemExit(f"morocco_rgph2024_ages: {name}'s age groups break at {low}")
        edge = high + 1 if high is not None else -1
    if not groups or groups[-1][1] is not None:
        raise SystemExit(f"morocco_rgph2024_ages: {name} has no open last age group")
    total = sum(s for _, _, s in groups)
    if abs(total - 100) > 0.8:
        raise SystemExit(f"morocco_rgph2024_ages: {name}'s age shares make {total:.1f}%")


def read_workbook(rows: list[list[Any]]) -> dict[int, dict[str, Any]]:
    """HCP code -> {"name", "groups"} for the regions and provinces of a regional
    directorate's workbook (sheet Population)."""
    head = next((i for i, r in enumerate(rows) if any(bounds(c) == (0, 4) for c in r)), None)
    if head is None:
        raise SystemExit("morocco_rgph2024_ages: no row of age groups in the workbook")
    cols = [(j, bounds(c)) for j, c in enumerate(rows[head]) if bounds(c)]
    # The first run of age groups is the one of the whole population.
    run = []
    for j, b in cols:
        if run and b[0] == 0:
            break
        run.append((j, b))
    out: dict[int, dict[str, Any]] = {}
    for r in rows[head + 1:]:
        try:
            code = int(str(r[0]).strip())
        except (TypeError, ValueError):
            continue
        if code >= 100_000:                     # cercles and communes
            continue
        try:
            groups = [(lo, hi, float(r[j])) for j, (lo, hi) in run]
        except (TypeError, ValueError):
            continue
        name = str(r[1] or "").strip()
        check(name, groups)
        out[code] = {"name": name, "groups": groups}
    return out


def read_annex(text: str) -> tuple[list, list]:
    """(the province's groups, the region's groups) from a provincial note's
    annex: the block headed Ensemble, rows '0-4 ans 7,5 7,0 7,2 8,4 8,0 8,2' --
    the province's men, women and both, then the region's."""
    block, runs, run = None, [], None
    for line in text.splitlines():
        word = line.strip()
        if word in ("Urbain", "Rural", "Ensemble"):
            block, run = word, None
            continue
        m = ROW.match(line)
        b = bounds(m.group(1)) if m else None
        values = ([float(v.replace(",", ".")) for v in re.findall(r"\d+(?:,\d+)?", m.group(2))]
                  if m else [])
        if b is None or len(values) != 6:
            run = None
            continue
        if b[0] == 0:
            run = {"block": block, "rows": []}
            runs.append(run)
        if run is not None:
            run["rows"].append((b, values))
    whole = [r["rows"] for r in runs if r["block"] == "Ensemble"
             and r["rows"] and r["rows"][-1][0][1] is None]
    if not whole or any(w != whole[0] for w in whole):
        raise SystemExit(f"morocco_rgph2024_ages: {len(whole)} different tables of ages for "
                         "the whole population in the annex")
    return ([(b[0], b[1], v[2]) for b, v in whole[0]],
            [(b[0], b[1], v[5]) for b, v in whole[0]])


def read_series(text: str) -> list[tuple[int, int | None, float]]:
    """The first whole table of five-year groups with men, women and both
    ('0-4 ans 7,7 7,0 7,3' ... '75ans et + 3,9 4,1 4'): both sexes' shares."""
    run: list = []
    for line in text.splitlines():
        m = ROW.match(line)
        b = bounds(m.group(1)) if m else None
        values = ([float(v.replace(",", ".")) for v in re.findall(r"\d+(?:,\d+)?", m.group(2))]
                  if m else [])
        if b is None or len(values) != 3:
            run = []
            continue
        if b[0] == 0:
            run = []
        run.append((b[0], b[1], values[2]))
        if b[1] is None and run[0][0] == 0:
            return run
    raise SystemExit("morocco_rgph2024_ages: no whole table of five-year ages in the note")


def broad_check(name: str, groups: list[tuple[int, int | None, float]],
                broad: dict[str, float]) -> None:
    """The five-year shares must make the dashboard's under 15 and 60 and over."""
    for (low, high), share in broad.items():
        made = sum(s for lo, hi, s in groups
                   if lo >= low and (high is None or (hi is not None and hi <= high)))
        if abs(made - share) > BROAD_SLACK:
            raise SystemExit(f"morocco_rgph2024_ages: {name}'s five-year shares make {made:.1f}% "
                             f"aged {low}-{high}, the dashboard {share}%")


def build(table: dict[str, Any], admin1: list[dict[str, Any]], admin2: list[dict[str, Any]],
          ages: dict[str, dict[str, Any]]) -> tuple[list[dict[str, Any]], list[str]]:
    """``ages``: legal.key(name) -> {"groups", "source", "url", "how"}."""
    bound, report = legal.build(table, admin1, admin2)
    out = []
    for rec in bound:
        code = int((rec.get("codes") or {}).get("hcp") or 0)
        if not code or "value" not in (rec.get("population") or {}):
            continue                                # across the line: said there
        row = table["regions"].get(code) if rec["level"] == "admin1" else \
            table["provinces"].get(code)
        if row is None:
            continue
        found = ages.get(legal.key(row["name"]))
        stated = next((why for name, why in STATED.items()
                       if legal.key(name) == legal.key(row["name"])), None)
        if stated:
            fields = {"median_age": gap(NOT_AVAILABLE, stated)}
        elif found:
            value = median(found["groups"])
            fields = {
                "median_age": measure(value, unit="years", year=YEAR, source=found["source"]),
                "median_age_note": (
                    "Interpolated within the five-year age group that holds the middle person, "
                    f"from the shares of {row['name']}'s municipal population in each five-year "
                    f"age group at the 2024 census{found['how']}, given to 0.1%."),
                "sources": [{"field": "median_age", "name": found["source"],
                             "url": found["url"], "license": LICENCE, "year": YEAR}]}
        else:
            fields = {"median_age": gap(NOT_AVAILABLE, ONLY_SHARES)}
            report.append(f"{row['name']}: no five-year ages")
        out.append(record(f"MAR-RGPH2024-AGE-{code}", rec["name"], level=rec["level"],
                          parent="MAR", country="MAR", match_by="shape_id",
                          shape_id=rec["shape_id"], codes={"hcp": str(code)}, **fields))
    return out, report


def gather(table: dict[str, Any]) -> dict[str, dict[str, Any]]:
    """legal.key(name) -> the unit's five-year shares and their source."""
    import openpyxl                                 # noqa: PLC0415
    import pdfplumber                               # noqa: PLC0415
    ages: dict[str, dict[str, Any]] = {}
    for region, (url, landing, title) in WORKBOOKS.items():
        blob = http_get(url, binary=True, cache=False, timeout=180)
        if blob[:2] != b"PK":
            raise SystemExit(f"morocco_rgph2024_ages: {url} is not a workbook")
        book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
        if "Population" not in book.sheetnames:
            raise SystemExit(f"morocco_rgph2024_ages: {url} has no sheet Population")
        units = read_workbook([list(r) for r in book["Population"].iter_rows(values_only=True)])
        mine = {c: u for c, u in units.items()
                if c in table["regions"] or c in table["provinces"]}
        for code, unit in sorted(mine.items()):
            hcp = (table["regions"].get(code) or table["provinces"].get(code))["name"]
            if legal.key(hcp) != legal.key(unit["name"]):
                raise SystemExit(f"morocco_rgph2024_ages: code {code} is {unit['name']!r} in "
                                 f"{title} and {hcp!r} in the legal population")
            ages[legal.key(hcp)] = {"groups": unit["groups"], "url": landing,
                                    "source": f"{HCP}, {title}", "how": ""}
            log(f"  {hcp}: median {median(unit['groups'])} ({title})")
        if legal.key(region) not in ages:
            raise SystemExit(f"morocco_rgph2024_ages: {title} has no row for {region}")
    regions = []
    for province, url in ANNEXES.items():
        blob = http_get(url, binary=True, cache=False, timeout=180)
        with pdfplumber.open(io.BytesIO(blob)) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        mine, region = read_annex(text)
        check(province, mine)
        check(ANNEX_REGION, region)
        regions.append(region)
        ages[legal.key(province)] = {
            "groups": mine, "url": LANDING_DOWNLOADS, "source": f"{HCP}, {ANNEX_TITLE}",
            "how": ", in HCP's note on the province"}
        log(f"  {province}: median {median(mine)}; the region {median(region)}")
    if any(r != regions[0] for r in regions):
        raise SystemExit(f"morocco_rgph2024_ages: the notes give {ANNEX_REGION} different "
                         "shares")
    if regions:
        ages[legal.key(ANNEX_REGION)] = {
            "groups": regions[0], "url": LANDING_DOWNLOADS, "source": f"{HCP}, {ANNEX_TITLE}",
            "how": ", in the annex of HCP's provincial notes"}
    for url, names in SERIES.items():
        blob = http_get(url, binary=True, cache=False, timeout=180)
        with pdfplumber.open(io.BytesIO(blob)) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
        # The note's own name for its unit, where two are possible.
        opening = legal.latin(" ".join(pages[:2]))
        named = [n for n in names if legal.key(n) in opening] if len(names) > 1 else list(names)
        if len(named) != 1:
            raise SystemExit(f"morocco_rgph2024_ages: {url} names {named or 'none'} of {names}")
        groups = read_series("\n".join(pages))
        check(named[0], groups)
        ages[legal.key(named[0])] = {
            "groups": groups, "url": SERIES_LANDING, "source": f"{HCP}, {SERIES_TITLE}",
            "how": ", in HCP's note on it"}
        log(f"  {named[0]}: median {median(groups)} ({url})")
    return ages


def dashboard_broad() -> dict[str, dict[tuple[int, int | None], float]]:
    """legal.key(unit) -> the dashboard's shares under 15 and of 60 and over."""
    from .morocco_rgph2024_sexes import post       # noqa: PLC0415
    rows = post(84, ["libelle_fr", "libelle_prov", "Key(Nom indicateur)"],
                [{"col": "Milieu", "op": "==", "val": "Ensemble"},
                 {"col": "Sexe", "op": "==", "val": "Ensemble"},
                 {"col": "libelle_commune", "op": "IS NULL"},
                 {"col": "Key(Nom indicateur)", "op": "IN", "val": list(BROAD)}])
    out: dict[str, dict[tuple[int, int | None], float]] = {}
    for r in rows:
        name = r.get("libelle_prov") or r.get("libelle_fr")
        if r.get("v") is not None and name:
            out.setdefault(legal.key(name), {})[BROAD[r["Key(Nom indicateur)"]]] = float(r["v"])
    return out


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--run", help="a label that tells one dispatch from another; not used")
    parser.parse_args()
    blob = http_get(legal.URL, binary=True, cache=False, timeout=120)
    if blob[:2] != b"PK":
        raise SystemExit(f"morocco_rgph2024_ages: {legal.URL} is not a workbook")
    table = legal.read(blob)
    ages = gather(table)
    broad = dashboard_broad()
    for key, unit in sorted(ages.items()):
        if key not in broad:
            raise SystemExit(f"morocco_rgph2024_ages: the dashboard has no broad ages for {key}")
        broad_check(key, unit["groups"], broad[key])
    log(f"  {len(ages)} units' five-year shares agree with the dashboard's under 15 and 60+")
    admin1 = json.loads((legal.SITE / "admin1" / "MAR.units.json").read_text())
    admin2 = json.loads((legal.SITE / "admin2" / "MAR.units.json").read_text())
    records, report = build(table, admin1, admin2, ages)
    for line in report:
        log("  " + line)
    shown = [r for r in records if "value" in (r.get("median_age") or {})]
    log(f"  {len(records)} records: {len(shown)} medians "
        f"({sum(r['level'] == 'admin1' for r in shown)} regions, "
        f"{sum(r['level'] == 'admin2' for r in shown)} provinces); the rest say why")
    write_json(PROCESSED / OUT, records)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
