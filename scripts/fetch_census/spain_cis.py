#!/usr/bin/env python3
"""Spain: religion by province and by comunidad from the CIS's pre-electoral surveys.

Spain's census does not ask religion. The Centro de Investigaciones
Sociológicas (CIS), the state's survey institute, asks it in nearly every
study, in the same words and with the same answers since its barometers
adopted them: "¿Cómo se define Ud. en materia religiosa: católico/a
practicante, católico/a no practicante, creyente de otra religión,
agnóstico/a, indiferente o no creyente, o ateo/a?", plus N.C. for no answer.
Those are the categories written here, in English: Practising Catholic,
Non-practising Catholic, Believer in another religion, Agnostic, Indifferent /
non-believer, Atheist, Not stated.

**Why not the monthly barometers.** They interview about 4,000 people a month,
and pooled over a year they would reach most provinces. But the CIS publishes
each barometer's results crossed only by sex, age, education, size of
municipality, vote recall, occupation, class, ideology and religiosity (the
"Resultados y cruces" workbooks) -- never by comunidad or province -- and it
releases the microdata only through a request form that takes the requester's
name, organisation, e-mail address and purpose. This project does not fill in
that form, and does not fetch the files behind it. The CIS's other route to
its microdata, the Ficheros Integrados de Datos, needs an account.

**What is read instead.** Before a regional election the CIS surveys that
comunidad with a sample drawn province by province -- the provinces are the
electoral districts -- and publishes every question's weighted marginals for
the comunidad and for each province. The recent ones put them in the study's
tabulation workbook ("Tabulaciones XLSX"): one sheet of marginals for the
comunidad and one per province, the realised sample and the mean weight of
each province in "Tabla ponderaciones", and the design in "Ficha técnica".
Four are published that way (``WORKBOOKS``): Extremadura (study 3538,
fieldwork late 2025), Aragón (3543), Castilla y León (3545) and Andalucía
(3558), all 2026 -- 22 provinces and their four comunidades, of which
Cáceres is withheld because its sheet is a copy of Badajoz's (see below). For Galicia
(3437) and the Basque Country (3448) the 2024 surveys are on the CIS's site
for the comunidad only (``TOTALS``, HTML marginals with a ficha técnica PDF):
their study pages still link province tables, which answer 404, so only the
comunidad is written for those two. Catalonia's 2024 survey (3453) is linked
the same way and none of its tables answers; the May 2023 regional and July
2023 general-election surveys (3402, 3411) link marginals by comunidad and
province on the CIS's old site, which answers 404 and 410 and which the
Internet Archive did not capture.

These are survey estimates of the people each survey was drawn from -- Spanish
citizens aged 18 and over resident in the comunidad and entitled to vote in
its election -- not counts, so the output file's name ends in ``_survey``:
the build lets it fill a composition no count has written, and as a national
office's survey it stands in front of the European Social Survey. The CIS's
shares are written as it publishes them (weighted, without counts): a
province's respondents are no measure of its people, and a roll-up must
price the shares against population rather than add interviews.

Checks, each of which stops the run:

* every answer is one the question offers, and every offered answer is there;
* each unit's shares make 100 (to 0.5 points);
* each province's respondents are the realised sample the weights table gives
  it, and the provinces make the comunidad's respondents and the table's total
  -- with one exception, which is withheld rather than stopping the rest: a
  province whose sheet is, figure for figure, another province's (study
  3538's "Provincia de Cáceres" sheet repeats Badajoz's shares and its 1,215
  respondents; Cáceres's realised sample is 822), so Cáceres is left out;
* no two provinces carry the same shares;
* the comunidad's shares are, to 1 point in every answer, the provinces'
  shares weighted by their realised samples and mean weights (the CIS
  weights the comunidad and each province separately, PESO and PESOPROV, so
  the two agree to a few tenths of a point rather than exactly);
* the ficha técnica still describes Spanish citizens aged 18 and over, and the
  workbook is the study it is listed as;
* every unit has at least 100 respondents (100-299 is marked low precision);
* every province binds by its INE or NUTS code to exactly one drawn polygon
  whose label is that province's name, inside the polygon its comunidad binds
  to by name.

Usage:
    python -m scripts.fetch_census.spain_cis
"""

from __future__ import annotations

import argparse
import html
import io
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from ._shared import PROCESSED, http_get, log, record, write_json
from .binding import fold

OUT = "spain_cis_survey.json"
SITE = PROCESSED.parent.parent / "site" / "data"

MIN_N = 100
LOW_PRECISION = 300
SHARE_TOLERANCE = 0.5
# How far a comunidad's share may sit from its provinces' combined: see check_region.
REGION_TOLERANCE = 1.0
LICENCE = ("CIS reuse conditions (Ley 37/2007): source to be cited as "
           "'Origen de los datos: Centro de Investigaciones Sociológicas'")

QUESTION_ES = ("¿Cómo se define Ud. en materia religiosa: católico/a practicante, "
               "católico/a no practicante, creyente de otra religión, agnóstico/a, "
               "indiferente o no creyente, o ateo/a?")
QUESTION = re.compile(r"c[oó]mo\s+se\s+define\s+ud\.?\s+en\s+materia\s+religiosa", re.I)

# The CIS's answers, by how each begins once accents, case and the zero-width
# spaces its HTML puts after "Católico/" are folded away.
ANSWERS: tuple[tuple[str, str], ...] = (
    ("catolicoapracticante", "Practising Catholic"),
    ("catolicoanopracticante", "Non-practising Catholic"),
    ("creyentedeotrareligion", "Believer in another religion"),
    ("agnosticoa", "Agnostic"),
    ("indiferente", "Indifferent / non-believer"),
    ("ateoa", "Atheist"),
)
NOT_STATED = {"nc", "ns", "nsnc", "nocontesta", "nosabe"}
SUBSTANTIVE = tuple(label for _, label in ANSWERS)


@dataclass(frozen=True)
class Study:
    number: int
    title: str          # the CIS's own title
    election: str       # what the survey was for, in English
    ccaa: str           # INE code of the comunidad surveyed
    url: str            # the workbook, or the comunidad's HTML marginals
    page: str           # the study's page on cis.es
    ficha: str | None = None   # the ficha técnica PDF, where the workbook does not carry it


WORKBOOKS: tuple[Study, ...] = (
    Study(3558, "Preelectoral elecciones autonómicas 2026. Comunidad autónoma de Andalucía",
          "the 2026 election to the Parliament of Andalusia", "01",
          "https://www.cis.es/documents/20117/13957047/3558-multi.xlsx",
          "https://www.cis.es/es/estudios/preelectoral-elecciones-autonomicas-2026-comunidad-"
          "autonoma-de-andalucia"),
    Study(3545, "Preelectoral elecciones autonómicas 2026. Comunidad autónoma de Castilla y León",
          "the 2026 election to the Cortes of Castile and León", "07",
          "https://www.cis.es/documents/20117/13765729/3545-multi.xlsx",
          "https://www.cis.es/es/estudios/preelectoral-elecciones-autonomicas-2026-comunidad-"
          "autonoma-de-castilla-y-leon"),
    # The CIS files study 3543's workbook as "3453-multi.xlsx"; its own header
    # says 3543, which is checked.
    Study(3543, "Preelectoral elecciones autonómicas 2026. Comunidad autónoma de Aragón",
          "the 2026 election to the Cortes of Aragon", "02",
          "https://www.cis.es/documents/20117/13730886/3453-multi.xlsx",
          "https://www.cis.es/es/estudios/preelectoral-elecciones-autonomicas-2026.-comunidad-"
          "autonoma-de-aragon"),
    Study(3538, "Preelectoral elecciones autonómicas 2025. Comunidad autónoma de Extremadura",
          "the 2025 election to the Assembly of Extremadura", "11",
          "https://www.cis.es/documents/20117/13697506/3538-multi.xlsx",
          "https://www.cis.es/es/estudios/preelectoral-elecciones-autonomicas-2025.-comunidad-"
          "autonoma-de-extremadura"),
)

TOTALS: tuple[Study, ...] = (
    Study(3448, "Preelectoral del País Vasco. Elecciones autonómicas 2024",
          "the 2024 election to the Basque Parliament", "16",
          "https://www.cis.es/documents/20117/1559601/es3448mar-html.html",
          "https://www.cis.es/es/estudios/preelectoral-del-pais-vasco.-elecciones-autonomicas-2024",
          "https://www.cis.es/documents/20117/1559601/FT3448.pdf"),
    Study(3437, "Preelectoral de Galicia. Elecciones autonómicas 2024",
          "the 2024 election to the Parliament of Galicia", "12",
          "https://www.cis.es/documents/20117/1559046/es3437mar-htm.html",
          "https://www.cis.es/es/estudios/preelectoral-de-galicia.-elecciones-autonomicas-2024",
          "https://www.cis.es/documents/20117/1559046/FT3437.pdf"),
)

# INE's province codes: the names the CIS and the boundary file use, the
# comunidad (INE code) and the NUTS 2024 code where the province is one
# NUTS 3 region (the Balearic and Canary NUTS 3 regions are islands).
PROVINCES: dict[str, tuple[tuple[str, ...], str, str | None]] = {
    "01": (("Araba/Álava", "Álava", "Araba"), "16", "ES211"),
    "02": (("Albacete",), "08", "ES421"),
    "03": (("Alicante/Alacant", "Alicante", "Alacant"), "10", "ES521"),
    "04": (("Almería",), "01", "ES611"),
    "05": (("Ávila",), "07", "ES411"),
    "06": (("Badajoz",), "11", "ES431"),
    "07": (("Balears, Illes", "Illes Balears", "Baleares"), "04", None),
    "08": (("Barcelona",), "09", "ES511"),
    "09": (("Burgos",), "07", "ES412"),
    "10": (("Cáceres",), "11", "ES432"),
    "11": (("Cádiz",), "01", "ES612"),
    "12": (("Castellón/Castelló", "Castellón", "Castelló"), "10", "ES522"),
    "13": (("Ciudad Real",), "08", "ES422"),
    "14": (("Córdoba",), "01", "ES613"),
    "15": (("A Coruña", "Coruña, A", "Coruña", "La Coruña"), "12", "ES111"),
    "16": (("Cuenca",), "08", "ES423"),
    "17": (("Girona", "Gerona"), "09", "ES512"),
    "18": (("Granada",), "01", "ES614"),
    "19": (("Guadalajara",), "08", "ES424"),
    "20": (("Gipuzkoa", "Guipúzcoa"), "16", "ES212"),
    "21": (("Huelva",), "01", "ES615"),
    "22": (("Huesca",), "02", "ES241"),
    "23": (("Jaén",), "01", "ES616"),
    "24": (("León",), "07", "ES413"),
    "25": (("Lleida", "Lérida"), "09", "ES513"),
    "26": (("La Rioja", "Rioja, La"), "17", "ES230"),
    "27": (("Lugo",), "12", "ES112"),
    "28": (("Madrid",), "13", "ES300"),
    "29": (("Málaga",), "01", "ES617"),
    "30": (("Murcia",), "14", "ES620"),
    "31": (("Navarra",), "15", "ES220"),
    "32": (("Ourense", "Orense"), "12", "ES113"),
    "33": (("Asturias",), "03", "ES120"),
    "34": (("Palencia",), "07", "ES414"),
    "35": (("Las Palmas", "Palmas, Las"), "05", None),
    "36": (("Pontevedra",), "12", "ES114"),
    "37": (("Salamanca",), "07", "ES415"),
    "38": (("Santa Cruz de Tenerife",), "05", None),
    "39": (("Cantabria",), "06", "ES130"),
    "40": (("Segovia",), "07", "ES416"),
    "41": (("Sevilla",), "01", "ES618"),
    "42": (("Soria",), "07", "ES417"),
    "43": (("Tarragona",), "09", "ES514"),
    "44": (("Teruel",), "02", "ES242"),
    "45": (("Toledo",), "08", "ES425"),
    "46": (("Valencia/València", "Valencia", "València"), "10", "ES523"),
    "47": (("Valladolid",), "07", "ES418"),
    "48": (("Bizkaia", "Vizcaya"), "16", "ES213"),
    "49": (("Zamora",), "07", "ES419"),
    "50": (("Zaragoza",), "02", "ES243"),
    "51": (("Ceuta",), "18", "ES630"),
    "52": (("Melilla",), "19", "ES640"),
}

# INE's comunidad codes and the names the CIS and the boundary file give them.
COMUNIDADES: dict[str, tuple[str, ...]] = {
    "01": ("Andalucía",), "02": ("Aragón",),
    "03": ("Principado de Asturias", "Asturias", "Asturias (Principado de)"),
    "04": ("Illes Balears", "Balears (Illes)", "Islas Baleares"), "05": ("Canarias",),
    "06": ("Cantabria",), "07": ("Castilla y León",), "08": ("Castilla-La Mancha",),
    "09": ("Cataluña/Catalunya", "Cataluña", "Catalunya"),
    "10": ("Comunitat Valenciana", "Comunidad Valenciana"), "11": ("Extremadura",),
    "12": ("Galicia",),
    "13": ("Comunidad de Madrid", "Madrid (Comunidad de)", "Madrid"),
    "14": ("Región de Murcia", "Murcia (Región de)", "Murcia"),
    "15": ("Comunidad Foral de Navarra", "Navarra (Comunidad Foral de)", "Navarra"),
    "16": ("País Vasco/Euskadi", "País Vasco", "Euskadi"), "17": ("La Rioja",),
    "18": ("Ciudad Autónoma de Ceuta", "Ceuta"), "19": ("Ciudad Autónoma de Melilla", "Melilla"),
}

MONTHS = {"enero": "January", "febrero": "February", "marzo": "March", "abril": "April",
          "mayo": "May", "junio": "June", "julio": "July", "agosto": "August",
          "septiembre": "September", "setiembre": "September", "octubre": "October",
          "noviembre": "November", "diciembre": "December"}


# --- reading the CIS's tables ------------------------------------------------

def answer_label(text: str) -> str:
    """The CIS's answer -> the map's label; an answer it does not offer stops the run."""
    key = fold(html.unescape(str(text)).replace("​", ""))
    if key in NOT_STATED:
        return "Not stated"
    for start, label in ANSWERS:
        if key.startswith(start):
            return label
    raise SystemExit(f"spain_cis: an answer the religion question does not offer: {text!r}")


def number(value: Any) -> float:
    """A cell or an HTML figure as a number: 15,2 -> 15.2; (4.998) -> 4998."""
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().strip("()").strip()
    if re.fullmatch(r"\d{1,3}(\.\d{3})+", text):          # a thousands separator
        return float(text.replace(".", ""))
    return float(text.replace(",", "."))


def composition(pairs: list[tuple[str, Any]], where: str) -> dict[str, float]:
    """(answer, share) pairs -> {label: share}, every answer offered and none other."""
    shares: dict[str, float] = {}
    for text, value in pairs:
        label = answer_label(text)
        shares[label] = shares.get(label, 0.0) + number(value)
    missing = [label for label in SUBSTANTIVE if label not in shares]
    if missing:
        raise SystemExit(f"spain_cis: {where}: the religion question lacks {missing}")
    total = sum(shares.values())
    if abs(total - 100.0) > SHARE_TOLERANCE:
        raise SystemExit(f"spain_cis: {where}: the religion answers make {total:.2f}%")
    return shares


def religion_block(rows: list[list[Any]], where: str) -> tuple[dict[str, float], int]:
    """The religion question's shares and its (N) in one sheet of marginals."""
    for i, row in enumerate(rows):
        first = next((c for c in row if c not in (None, "")), None)
        if first is None or not QUESTION.search(str(first)):
            continue
        pairs: list[tuple[str, Any]] = []
        for follow in rows[i + 1:]:
            cells = [c for c in follow if c not in (None, "")]
            if not cells:
                continue
            label = str(cells[0]).strip()
            if label == "(N)":
                if len(cells) < 2:
                    raise SystemExit(f"spain_cis: {where}: an (N) row with no figure")
                return composition(pairs, where), int(number(cells[1]))
            if label.lower().startswith("pregunta") or len(cells) < 2:
                break
            pairs.append((label, cells[1]))
        raise SystemExit(f"spain_cis: {where}: the religion question ends without its (N)")
    raise SystemExit(f"spain_cis: {where}: no religion question")


def weights_table(rows: list[list[Any]], where: str) -> tuple[dict[str, dict[str, Any]], int]:
    """'Tabla ponderaciones' -> {INE province code: name, realised, weight}, and the total."""
    # The header row's own cells, not the sheet's title above it ("Muestra
    # diseñada, realizada, coeficientes de ponderación y error"), which names both.
    head = next((i for i, row in enumerate(rows)
                 if any(str(c).strip().lower() == "realizada" for c in row)
                 and any(str(c).strip().lower().startswith("ponderaci") for c in row)), None)
    if head is None:
        raise SystemExit(f"spain_cis: {where}: no weights table")
    table: dict[str, dict[str, Any]] = {}
    total = None
    for row in rows[head + 1:]:
        cells = [c for c in row if c not in (None, "")]
        if not cells:
            continue
        if any(str(c).strip().lower() == "total" for c in cells[:2]):
            numbers = [c for c in cells if not isinstance(c, str) or re.fullmatch(r"[\d.,]+", c)]
            if len(numbers) < 2:
                raise SystemExit(f"spain_cis: {where}: the weights table's total has {cells}")
            total = int(number(numbers[1]))
            break
        code = str(cells[0]).strip()
        if not re.fullmatch(r"\d{1,2}", code) or len(cells) < 5:
            raise SystemExit(f"spain_cis: {where}: a weights table row reads {cells}")
        code = code.zfill(2)
        table[code] = {"name": str(cells[1]).strip(), "designed": int(number(cells[2])),
                       "realised": int(number(cells[3])), "weight": number(cells[4])}
    if total is None or not table:
        raise SystemExit(f"spain_cis: {where}: the weights table has no total")
    if sum(t["realised"] for t in table.values()) != total:
        raise SystemExit(f"spain_cis: {where}: the provinces' realised samples make "
                         f"{sum(t['realised'] for t in table.values()):,}, the total {total:,}")
    return table, total


def ficha(text: str, where: str) -> dict[str, Any]:
    """The design as a ficha técnica states it: universe, realised sample, fieldwork."""
    flat = re.sub(r"\s+", " ", text.replace("_x000D_", " "))
    universe = re.search(r"Universo:?\s*(.+?)\s*(?:Tamaño de la muestra|Diseñada)", flat)
    realised = re.search(r"Realizada:?\s*([\d.]+)\s*entrevistas", flat)
    dates = re.search(r"Fecha de realizaci[oó]n:?\s*(.+?\d{4})", flat)
    if not (universe and realised and dates):
        raise SystemExit(f"spain_cis: {where}: the ficha técnica no longer states the universe, "
                         f"sample and fieldwork as it did")
    out = {"universe": universe.group(1).strip(), "realised": int(number(realised.group(1))),
           "fieldwork": dates.group(1).strip(), "year": int(dates.group(1)[-4:]),
           "telephone": bool(re.search(r"CATI|telef[oó]nic", flat, re.I)),
           "weighted": bool(re.search(r"\bPESO\b", flat)),
           "province_weight": "PESOPROV" in flat}
    universe_text = fold(out["universe"])
    if "espanola" not in universe_text or "18" not in out["universe"]:
        raise SystemExit(f"spain_cis: {where}: the universe is now {out['universe']!r}; the "
                         f"note says Spanish citizens aged 18 and over")
    out["voters"] = "voto" in universe_text
    return out


def names_study(text: str, number_: int) -> bool:
    """Whether a page or sheet names the study -- after "nº" too, which \\b misses."""
    return re.search(rf"(?<!\d){number_}(?!\d)", text) is not None


def english_dates(text: str) -> str:
    """'Del 6 al 13 de febrero de 2026' -> '6-13 February 2026', or the text as it is."""
    t = text.strip().rstrip(".")
    m = re.fullmatch(r"(?i)del (\d{1,2})(?: de (\w+))? al (\d{1,2}) de (\w+) de (\d{4})", t)
    if m:
        d1, m1, d2, m2, year = m.groups()
        m2e = MONTHS.get(m2.lower())
        m1e = MONTHS.get(m1.lower()) if m1 else None
        if m2e and (m1 is None or m1e):
            return (f"{d1} {m1e} - {d2} {m2e} {year}" if m1e and m1e != m2e
                    else f"{d1}-{d2} {m2e} {year}")
    return t


def sheet_unit(name: str) -> str:
    """A results sheet's name -> the unit it describes, folded."""
    name = re.sub(r"(?i)^\s*resultados[\s_]*", "", name)
    name = re.sub(r"(?i)^(?:de\s+)?(?:la\s+)?provincia\s+de\s+", "", name)
    return fold(name)


def read_workbook(raw: bytes, study: Study) -> dict[str, Any]:
    """A pre-electoral study's workbook -> its design, its provinces and its comunidad."""
    import openpyxl                                    # noqa: PLC0415 -- the runner has it
    book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    where = f"study {study.number}"

    def rows_of(title: str) -> list[list[Any]]:
        return [list(r) for r in book[title].iter_rows(values_only=True)]

    names = book.sheetnames
    header = " ".join(str(c) for row in rows_of(names[0])[:4] for c in row if c)
    if not re.search(rf"(?<!\d){study.number}/\d", header):
        raise SystemExit(f"spain_cis: the workbook listed as study {study.number} is headed "
                         f"{header[:120]!r}")
    sheet = next((n for n in names if fold(n).startswith("fichatecnica")), None)
    table = next((n for n in names if fold(n).startswith("tablaponderaciones")), None)
    if sheet is None or table is None:
        raise SystemExit(f"spain_cis: {where}: sheets are {names}")
    design = ficha(" ".join(f"{c}:" if j == 0 and isinstance(c, str) else str(c)
                            for row in rows_of(sheet) for j, c in enumerate(row) if c), where)
    weights, total = weights_table(rows_of(table), where)
    if total != design["realised"]:
        raise SystemExit(f"spain_cis: {where}: the weights table totals {total:,}, the ficha "
                         f"técnica {design['realised']:,}")
    by_unit = {sheet_unit(n): n for n in names}
    region_names = COMUNIDADES[study.ccaa]
    region_sheet = next((by_unit[fold(n)] for n in region_names if fold(n) in by_unit), None)
    if region_sheet is None:
        raise SystemExit(f"spain_cis: {where}: no sheet for {region_names[0]} among {names}")
    region, region_n = religion_block(rows_of(region_sheet), f"{where} {region_sheet!r}")
    read: dict[str, dict[str, Any]] = {}
    for code, entry in weights.items():
        if code not in PROVINCES or PROVINCES[code][1] != study.ccaa:
            raise SystemExit(f"spain_cis: {where}: province code {code} ({entry['name']}) is not "
                             f"in comunidad {study.ccaa}")
        known = {fold(n) for n in PROVINCES[code][0]}
        if fold(entry["name"]) not in known:
            raise SystemExit(f"spain_cis: {where}: province {code} is {entry['name']!r} in the "
                             f"weights table")
        title = next((by_unit[k] for k in known if k in by_unit), None)
        if title is None:
            raise SystemExit(f"spain_cis: {where}: no results sheet for {entry['name']}")
        shares, n = religion_block(rows_of(title), f"{where} {title!r}")
        read[code] = {"name": entry["name"], "shares": shares, "n": n,
                      "weight": entry["weight"], "sheet": title, "realised": entry["realised"]}
    provinces, withheld = sort_out(read, where)
    expected = {c for c, (_, ccaa, _) in PROVINCES.items() if ccaa == study.ccaa}
    if set(provinces) | set(withheld) != expected:
        raise SystemExit(f"spain_cis: {where}: provinces {sorted(read)}, expected "
                         f"{sorted(expected)}")
    if region_n != total:
        raise SystemExit(f"spain_cis: {where}: the comunidad has {region_n:,} respondents, the "
                         f"weights table {total:,}")
    if withheld:
        log(f"  {where}: the comunidad is not checked against its provinces: "
            f"{', '.join(read[c]['name'] for c in withheld)} withheld")
    else:
        check_region(region, provinces, where)
    modified = getattr(book.properties, "modified", None)
    return {"design": design, "region": {"shares": region, "n": region_n},
            "provinces": provinces, "withheld": withheld,
            "modified": modified.date().isoformat() if modified else None}


def sort_out(read: dict[str, dict[str, Any]], where: str
             ) -> tuple[dict[str, dict[str, Any]], dict[str, str]]:
    """Provinces whose sheet is theirs, and those whose sheet is another's copy.

    A province's respondents to the religion question are its realised sample;
    a sheet that disagrees stops the run -- except where it is, figure for
    figure, another province's sheet, whose own sample it carries. That is the
    CIS's workbook repeating one province under another's name, as study 3538
    does: its "Provincia de Cáceres" sheet is Badajoz's, every share to the
    fifteenth decimal and its 1,215 respondents (Cáceres's realised sample is
    822). Such a province is withheld and said so, never given the copy.
    """
    good = {c: p for c, p in read.items() if p["n"] == p["realised"]}
    withheld: dict[str, str] = {}
    for code, prov in read.items():
        if code in good:
            continue
        twin = next((c for c, p in good.items()
                     if p["n"] == prov["n"] and p["shares"] == prov["shares"]), None)
        if twin is None:
            raise SystemExit(f"spain_cis: {where}: {prov['name']} has {prov['n']:,} respondents "
                             f"to the religion question and a realised sample of "
                             f"{prov['realised']:,}")
        withheld[code] = (f"the CIS's sheet for {prov['name']} repeats {good[twin]['name']}'s, "
                          f"every share and its {prov['n']:,} respondents ({prov['name']}'s "
                          f"realised sample is {prov['realised']:,})")
        log(f"  {where}: {prov['name']} withheld: {withheld[code]}")
    for code, prov in good.items():
        twins = [p["name"] for c, p in good.items() if c != code and p["shares"] == prov["shares"]]
        if twins:
            raise SystemExit(f"spain_cis: {where}: {prov['name']} and {twins} carry the same "
                             f"shares")
    return good, withheld


def check_region(region: dict[str, float], provinces: dict[str, dict[str, Any]],
                 where: str) -> float:
    """The comunidad's shares against its provinces' weighted by sample and mean weight.

    A province's realised sample times its mean weight is its weighted size,
    so the comunidad's share is close to the provinces' shares in those
    proportions. Not exactly: the CIS weights the comunidad's marginals with
    PESO (calibrated to sex, age, province, education and size of
    municipality) and each province's with PESOPROV (calibrated within the
    province), so the two agree to a few tenths of a point -- 0.32 at most in
    Andalucía, 0.54 in Castilla y León. The log gives the agreement for each
    study, beside what the provinces' shares make added up unweighted.
    """
    mass = {c: p["n"] * p["weight"] for c, p in provinces.items()}
    whole = sum(mass.values())
    sample = sum(p["n"] for p in provinces.values())
    worst, plain = 0.0, 0.0
    for label in region:
        made = sum(mass[c] * p["shares"].get(label, 0.0) for c, p in provinces.items()) / whole
        flat = sum(p["n"] * p["shares"].get(label, 0.0) for p in provinces.values()) / sample
        worst = max(worst, abs(made - region[label]))
        plain = max(plain, abs(flat - region[label]))
        if abs(made - region[label]) > REGION_TOLERANCE:
            raise SystemExit(f"spain_cis: {where}: the provinces make {label} {made:.2f}%, the "
                             f"comunidad {region[label]:.2f}%")
    log(f"  {where}: the provinces make the comunidad's shares to within {worst:.2f} points "
        f"weighted by their samples and mean weights ({plain:.2f} unweighted)")
    return worst


def read_html_total(text: str, study: Study) -> tuple[dict[str, float], int]:
    """A comunidad's HTML marginals -> the religion question's shares and its (N)."""
    if not names_study(text, study.number):
        raise SystemExit(f"spain_cis: the marginals listed as study {study.number} do not name it")
    m = QUESTION.search(html.unescape(text))
    if not m:
        raise SystemExit(f"spain_cis: study {study.number}: no religion question")
    body = html.unescape(text)[m.end():]
    table = re.search(r"(?is)<table.*?</table>", body)
    if not table:
        raise SystemExit(f"spain_cis: study {study.number}: the religion question has no table")
    rows: list[list[Any]] = []
    for row in re.findall(r"(?is)<tr.*?</tr>", table.group(0)):
        cells = [re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", c)).strip()
                 for c in re.findall(r"(?is)<td.*?</td>", row)]
        rows.append(cells)
    return religion_block([[QUESTION_ES]] + rows, f"study {study.number}")


# --- the map -------------------------------------------------------------------

def load_units() -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    from common import as_drawn                       # noqa: PLC0415 -- on the path via _shared
    a1 = as_drawn(json.loads((SITE / "admin1" / "ESP.units.json").read_text()))
    a2 = as_drawn(json.loads((SITE / "admin2" / "ESP.units.json").read_text()))
    return a1, a2


def bind(admin1: list[dict[str, Any]], admin2: list[dict[str, Any]]
         ) -> tuple[dict[str, dict[str, Any]], dict[str, dict[str, Any]]]:
    """Every INE province -> its drawn polygon, and every comunidad -> the polygon around them.

    A province binds by its code -- INE's, or the NUTS 3 region it is -- to
    exactly one drawn polygon, and that polygon's own label must be the
    province's name. A comunidad is the one polygon all its provinces are
    drawn inside, and its label must be the comunidad's name.
    """
    provinces: dict[str, dict[str, Any]] = {}
    used: set[str] = set()
    for code, (names, _, nuts) in PROVINCES.items():
        hits = [u for u in admin2 if (u.get("codes") or {}).get("ine_province") == code
                or (nuts and (u.get("codes") or {}).get("nuts") == nuts)]
        if not hits:
            hits = [u for u in admin2 if fold(u["name"]) in {fold(n) for n in names}]
        if len(hits) != 1:
            raise SystemExit(f"spain_cis: province {code} ({names[0]}) matches "
                             f"{[u['name'] for u in hits]} on the map")
        unit = hits[0]
        if fold(unit["name"]) not in {fold(n) for n in names}:
            raise SystemExit(f"spain_cis: province {code} ({names[0]}) has code-matched the "
                             f"polygon labelled {unit['name']!r}")
        if unit["id"] in used:
            raise SystemExit(f"spain_cis: polygon {unit['name']} bound twice")
        used.add(unit["id"])
        provinces[code] = unit
    by_id = {u["id"]: u for u in admin1}
    comunidades: dict[str, dict[str, Any]] = {}
    for ccaa, names in COMUNIDADES.items():
        parents = {provinces[c]["parent"] for c, (_, cc, _) in PROVINCES.items() if cc == ccaa}
        if len(parents) != 1 or next(iter(parents)) not in by_id:
            raise SystemExit(f"spain_cis: comunidad {ccaa}'s provinces are drawn inside {parents}")
        unit = by_id[next(iter(parents))]
        if fold(unit["name"]) not in {fold(n) for n in names}:
            raise SystemExit(f"spain_cis: comunidad {ccaa} ({names[0]}) is drawn as "
                             f"{unit['name']!r}")
        comunidades[ccaa] = unit
    return provinces, comunidades


# --- records ---------------------------------------------------------------------

def rows_of(shares: dict[str, float]) -> list[dict[str, Any]]:
    rows = [{"group": label, "pct": round(share, 1)} for label, share in shares.items()]
    rows.sort(key=lambda r: (-r["pct"], r["group"]))
    return rows


def fields(study: Study, design: dict[str, Any], unit: str, n: int,
           shares: dict[str, float], what: str, updated: str | None,
           province: bool = False) -> dict[str, Any]:
    precision = (" Low precision: under 300 respondents." if n < LOW_PRECISION else "")
    method = ("telephone interviews (CATI) with quotas of sex and age, " if design["telephone"]
              else "")
    if province and design.get("province_weight"):
        weighting = ("weighted by the CIS's post-stratification coefficients for the province "
                     "(PESOPROV); these are its published weighted shares")
    elif not province and design.get("weighted"):
        weighting = ("weighted by the CIS's coefficients for the comunidad (PESO); these are "
                     "its published weighted shares")
    else:
        weighting = "shares as the CIS publishes them"
    comunidad = COMUNIDADES[study.ccaa][0]
    voters = design.get("voters", False)
    return {
        "religion": rows_of(shares),
        "religion_year": design["year"],
        "religion_basis": ("survey estimate: self-identification, Spanish citizens aged 18+"
                           + (" entitled to vote in the regional election" if voters else "")),
        "religion_note": (
            f"CIS study {study.number}, the pre-electoral survey for {study.election}, fieldwork "
            f"{english_dates(design['fieldwork'])}: \"{QUESTION_ES}\" Answers as the CIS offers "
            f"them; its N.C. (no answer) is Not stated. A survey estimate, not a count: "
            f"{n:,} respondents in {what} ({method}{weighting}).{precision} Universe: Spanish "
            f"citizens aged 18 and over resident in {comunidad}"
            + (" and entitled to vote in its election" if voters else "")
            + ", so foreign residents are not in it. Spain's census does not ask religion, and "
              "the CIS's monthly barometers publish it by no region."),
        "sources": [{"field": "religion",
                     "name": (f"CIS, study {study.number}: {study.title} -- weighted marginals "
                              f"for {unit}" + (f" (updated {updated})" if updated else "")),
                     "url": study.url, "year": design["year"], "license": LICENCE}],
    }


def build(results: list[tuple[Study, dict[str, Any]]], provinces: dict[str, dict[str, Any]],
          comunidades: dict[str, dict[str, Any]]) -> tuple[list[dict[str, Any]], list[str]]:
    records: list[dict[str, Any]] = []
    skipped: list[str] = []
    for study, got in results:
        design, updated = got["design"], got.get("modified")
        region = got["region"]
        unit = comunidades[study.ccaa]
        if region["n"] < MIN_N:
            skipped.append(f"{unit['name']}: {region['n']} respondents (study {study.number})")
        else:
            records.append(record(
                f"ESP-CIS{study.number}-CA{study.ccaa}", unit["name"], level="admin1",
                parent="ESP", country="ESP", match_by="shape_id", shape_id=unit["id"],
                codes={"ine_ccaa": study.ccaa, "cis_study": study.number, "cis_n": region["n"]},
                **fields(study, design, unit["name"], region["n"], region["shares"],
                         unit["name"], updated)))
        for code, why in sorted(got.get("withheld", {}).items()):
            skipped.append(f"{provinces[code]['name']}: {why} (study {study.number})")
        for code, prov in sorted(got.get("provinces", {}).items()):
            shape = provinces[code]
            if shape["parent"] != unit["id"]:
                raise SystemExit(f"spain_cis: {shape['name']} is not drawn inside {unit['name']}")
            if prov["n"] < MIN_N:
                skipped.append(f"{shape['name']}: {prov['n']} respondents (study {study.number})")
                continue
            records.append(record(
                f"ESP-CIS{study.number}-{code}", shape["name"], level="admin2",
                parent="ESP", country="ESP", match_by="shape_id", shape_id=shape["id"],
                codes={"ine_province": code, "cis_study": study.number, "cis_n": prov["n"]},
                **fields(study, design, prov["name"], prov["n"], prov["shares"],
                         f"the province of {prov['name']}", updated, province=True)))
    ids = [r["shape_id"] for r in records]
    if len(ids) != len(set(ids)):
        raise SystemExit("spain_cis: a polygon was written twice")
    return records, skipped


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    admin1, admin2 = load_units()
    provinces, comunidades = bind(admin1, admin2)
    log(f"  bound {len(provinces)} provinces and {len(comunidades)} comunidades by code")

    results: list[tuple[Study, dict[str, Any]]] = []
    for study in WORKBOOKS:
        raw = http_get(study.url, binary=True, cache=False, timeout=180)
        assert isinstance(raw, bytes)
        got = read_workbook(raw, study)
        log(f"  study {study.number}: {len(got['provinces'])} provinces, "
            f"{got['region']['n']:,} respondents, fieldwork {got['design']['fieldwork']}")
        results.append((study, got))
    for study in TOTALS:
        assert study.ficha
        page = http_get(study.url, cache=False, timeout=180)
        assert isinstance(page, str)
        shares, n = read_html_total(page, study)
        from pypdf import PdfReader                    # noqa: PLC0415 -- the runner has it
        pdf = http_get(study.ficha, binary=True, cache=False, timeout=180)
        assert isinstance(pdf, bytes)
        design = ficha(" ".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(pdf)).pages),
                       f"study {study.number}")
        if n != design["realised"]:
            raise SystemExit(f"spain_cis: study {study.number}: {n:,} answered the religion "
                             f"question of {design['realised']:,} interviewed")
        log(f"  study {study.number}: {n:,} respondents, fieldwork {design['fieldwork']}")
        results.append((study, {"design": design, "region": {"shares": shares, "n": n}}))

    records, skipped = build(results, provinces, comunidades)
    for r in records:
        log(f"  {r['level']} {r['name']} (n={r['codes']['cis_n']:,}, {r['religion_year']}): "
            + ", ".join(f"{g['group']} {g['pct']}" for g in r["religion"]))
    for line in skipped:
        log(f"  left out: {line}")
    write_json(Path(args.out) if args.out else PROCESSED / OUT, records)
    log(f"  {len(records)} records: "
        f"{sum(r['level'] == 'admin1' for r in records)} comunidades, "
        f"{sum(r['level'] == 'admin2' for r in records)} provinces")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
