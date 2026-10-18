#!/usr/bin/env python3
"""Burundi: men per 1,000 women for every province and commune, from the 2008 census.

Burundi's third general census of population and housing (RGPH, 16 August
2008) counted 8,053,574 residents, 3,964,906 men and 4,088,668 women; the
decree of 5 April 2010 that published its final results annexes the
population by province, commune and sex. The statistics institute (ISTEEBU,
now INSBU) republishes the census's first-chapter tables as a workbook,
Chapitre-1.xlsx, whose Tableau 1.6 is the resident population by province and
commune and by sex, and whose Tableau 1.4 the same with ordinary and
collective households apart. The institute's site no longer serves its older
links (the decree's PDF under wp-content answers with a page of about 1 KB),
so the Internet Archive's capture of the workbook is read.

The map shows OCHA's projection for 2022 (cod-ps-bdi) for every commune's
head count, median age and sex ratio. A census's sex ratio stands in front of
a projection's, so this writes the 2008 census's for every unit the map draws.
The census published no ages by province or commune (its age tables are by
residence, urban or rural, for the country), so the projection's medians stay.

**Which rows make which polygon.** The map draws the 116 rural communes of
2008 and Bujumbura Mairie as the three communes it has had since 2014, made
from the city's thirteen communes of 2008, now its zones: Muha (Kanyosha,
Kinindo, Musaga), Mukaza (Buyenzi, Bwiza, Nyakabiga, Rohero) and Ntahangwa
(Buterere, Cibitoke, Gihosha, Kamenge, Kinama, Ngagara). Rumonge province was
made in 2015 from Bururi's communes of Burambi, Buyengero and Rumonge and
Bujumbura Rural's of Bugarama and Muhuta; each province is the sum of the
communes that make it as drawn.

**Checks.** Every row's men and women make its total; each province's
communes make its row; the provinces make the national row (8,053,574); and
Tableau 1.6 and Tableau 1.4's totals agree row by row.

Usage:
    python -m scripts.fetch_census.burundi_census
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
from typing import Any

from ._shared import PROCESSED, http_get, log, measure, record, write_json
from .binding import fold

OUT = "burundi_census.json"
YEAR = 2008
SITE = PROCESSED.parent.parent / "site" / "data"
PAGE = "https://www.insbu.bi/documents/recensements/Chapitre-1.xlsx"
WORKBOOK = f"https://web.archive.org/web/20251225150927id_/{PAGE}"
SOURCE = ("Institut de Statistiques et d'Études Économiques du Burundi, Troisième Recensement "
          "Général de la Population et de l'Habitat 2008, Tableau 1.6: population résidante "
          "par province et par commune selon le sexe")
LICENCE = "Government statistics publication (reuse with attribution)"
NATIONAL = (3_964_906, 4_088_668, 8_053_574)
PROVINCES_2008 = 17
COMMUNES_2008 = 129
# Bujumbura Mairie's communes of today -> its communes of 2008, now their zones.
MAIRIE = {
    "Muha": ("KANYOSHA", "KININDO", "MUSAGA"),
    "Mukaza": ("BUYENZI", "BWIZA", "NYAKABIGA", "ROHERO"),
    "Ntahangwa": ("BUTERERE", "CIBITOKE", "GIHOSHA", "KAMENGE", "KINAMA", "NGAGARA"),
}
# Rumonge province (2015): its communes and the provinces of 2008 they were in.
RUMONGE = {"BUGARAMA": "BUJUMBURA RURAL", "MUHUTA": "BUJUMBURA RURAL",
           "BURAMBI": "BURURI", "BUYENGERO": "BURURI", "RUMONGE": "BURURI"}
# The boundary file's spelling of a commune -> the census's, where neither is
# the other's beginning (the boundary file cuts names at ten letters).
SPELLING = {"Kanyosha1": "KANYOSHA", "Nyanrusang": "NYARUSANGE", "Buhiga": "BUHUGA",
            "Mpinga-Kay": "MPINGA"}
# The boundary file's spelling of a province -> the census's.
PROVINCE_SPELLING = {"Karuzi": "KARUSI"}
# Rows the workbook mislabels, by (province, label, men, women): the name the
# decree of 5 April 2010 gives them in its annex. The workbook calls two of
# Makamba's communes "MAKAMBA"; the annex calls the first Mabanda, which also
# falls between Kibago and Makamba in the table's alphabetical order.
RELABELLED = {("MAKAMBA", "MAKAMBA", 22_910, 22_926): "MABANDA"}


def number(value: Any) -> int | None:
    if value in (None, "", "-"):
        return None
    try:
        return int(round(float(str(value).replace(" ", "").replace(",", ""))))
    except ValueError:
        return None


def read_rows(rows: list[list[Any]], columns: tuple[int, int, int] = (1, 2, 3)
              ) -> list[dict[str, Any]]:
    """[{name, male, female, total}] for every labelled row with three counts,
    from the sheet's rows, each checked: its men and women make its total."""
    out = []
    m, f, t = columns
    for row in rows:
        if not row or len(row) <= t:
            continue
        name = " ".join(str(row[0] or "").split())
        male, female, total = number(row[m]), number(row[f]), number(row[t])
        if not name or None in (male, female, total):
            continue
        if male + female != total:
            raise SystemExit(f"burundi_census: {name}: {male:,} men and {female:,} women "
                             f"against {total:,}")
        out.append({"name": name, "male": male, "female": female, "total": total})
    return out


def provinces(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    """{province: {row, communes: {name: row}}}: after the national row, each
    province's row is followed by its communes', which make it in men, women
    and total; the provinces make the national row."""
    if not rows or fold(rows[0]["name"]) != "burundi" \
            or (rows[0]["male"], rows[0]["female"], rows[0]["total"]) != NATIONAL:
        raise SystemExit(f"burundi_census: the table does not open with Burundi's "
                         f"{NATIONAL[2]:,} ({rows[:1]})")
    out: dict[str, dict[str, Any]] = {}
    i = 1
    while i < len(rows):
        head, j = rows[i], i + 1
        made = {"male": 0, "female": 0, "total": 0}
        while j < len(rows) and made["total"] < head["total"]:
            for k in made:
                made[k] += rows[j][k]
            j += 1
        if any(made[k] != head[k] for k in made) or j == i + 1:
            raise SystemExit(f"burundi_census: {head['name']}'s communes make {made}, not "
                             f"its row")
        if head["name"] in out:
            raise SystemExit(f"burundi_census: {head['name']} appears twice")
        for r in rows[i + 1:j]:
            label = RELABELLED.get((head["name"], r["name"], r["male"], r["female"]))
            if label:
                r["name"] = label
        communes = {r["name"]: r for r in rows[i + 1:j]}
        if len(communes) != j - i - 1:
            raise SystemExit(f"burundi_census: a commune of {head['name']} appears twice")
        out[head["name"]] = {"row": head, "communes": communes}
        i = j
    made = {k: sum(p["row"][k] for p in out.values()) for k in ("male", "female", "total")}
    ncom = sum(len(p["communes"]) for p in out.values())
    if (made["male"], made["female"], made["total"]) != NATIONAL \
            or len(out) != PROVINCES_2008 or ncom != COMMUNES_2008:
        raise SystemExit(f"burundi_census: {len(out)} provinces, {ncom} communes, making "
                         f"{made}")
    return out


def added(rows: list[dict[str, Any]]) -> dict[str, int]:
    return {k: sum(r[k] for r in rows) for k in ("male", "female", "total")}


def ratio(figures: dict[str, int]) -> int:
    return round(1000 * figures["male"] / figures["female"])


def title(name: str) -> str:
    return "-".join(part.capitalize() for part in name.lower().split("-"))


def census_commune(label: str, communes: dict[str, dict[str, Any]]) -> str | None:
    """The census's commune for a polygon's label among one province's communes:
    the same name, the census's name beginning with the label (cut at ten
    letters), or SPELLING; None where no one commune answers."""
    want = SPELLING.get(label.strip())
    if want is not None:
        return want if want in communes else None
    key = fold(label)
    exact = [c for c in communes if fold(c) == key]
    if len(exact) == 1:
        return exact[0]
    begins = [c for c in communes if fold(c).startswith(key)] if len(key) >= 8 else []
    return begins[0] if len(begins) == 1 else None


def build(table: dict[str, dict[str, Any]], admin1: list[dict[str, Any]],
          admin2: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[str]]:
    cite = [{"field": "sex_ratio", "name": SOURCE, "url": PAGE, "year": YEAR,
             "license": LICENCE}]
    by_fold = {fold(p): p for p in table}
    mairie = next((p for p in table if "mairie" in fold(p)), None)
    if mairie is None:
        raise SystemExit(f"burundi_census: no Bujumbura Mairie in {sorted(table)}")
    names1 = {s["id"]: s["name"].strip() for s in admin1}
    records, report = [], []
    used: set[tuple[str, str]] = set()

    def note(parts: list[str], whose: str) -> str:
        if len(parts) == 1:
            return (f"Men per 1,000 women among the {whose}'s residents counted by the 2008 "
                    "census.")
        listed = ", ".join(title(p) for p in parts[:-1]) + " and " + title(parts[-1])
        return (f"Men per 1,000 women among the residents of {listed}, counted by the 2008 "
                f"census and added up: {whose}.")

    for shape in admin2:
        label = shape["name"].strip()
        if label in MAIRIE:
            parts = list(MAIRIE[label])
            rows = [table[mairie]["communes"].get(p) for p in parts]
            if None in rows:
                raise SystemExit(f"burundi_census: {label}: {parts} not all in {mairie}")
            used.update((mairie, p) for p in parts)
            why = (f"the communes of 2008 that are now {label}'s zones, {label} having been "
                   "one of Bujumbura's three communes since 2014")
            text = note(parts, why)
        else:
            province = names1.get(shape["parent"], "")
            homes = ([RUMONGE.get(c) for c in RUMONGE] if fold(province) == "rumonge" else
                     [by_fold.get(fold(PROVINCE_SPELLING.get(province, province)))])
            found = None
            for home in dict.fromkeys(h for h in homes if h):
                hit = census_commune(label, table[home]["communes"])
                if hit and (fold(province) != "rumonge" or RUMONGE.get(hit) == home):
                    found = (home, hit)
                    break
            if found is None:
                report.append(f"{label} ({province}): no commune of the census answers; not "
                              "written")
                continue
            if found in used:
                raise SystemExit(f"burundi_census: {found} answers two polygons")
            used.add(found)
            rows = [table[found[0]]["communes"][found[1]]]
            text = note([found[1]], "commune")
            if fold(found[1]) != fold(label):
                report.append(f"{label} ({province}): the census's {found[1]}")
        figures = added(rows)
        records.append(record(
            f"BDI-2008-{shape['id']}", label, level="admin2", parent="BDI", country="BDI",
            match_by="shape_id", shape_id=shape["id"],
            sex_ratio=measure(ratio(figures), unit="males_per_1000_females", year=YEAR,
                              source=SOURCE),
            sex_ratio_note=text, sources=cite))
    left = sorted(f"{p}/{c}" for p in table for c in table[p]["communes"] if (p, c) not in used)
    if left:
        report.append(f"census communes on no polygon: {left}")
    for shape in admin1:
        label = shape["name"].strip()
        key = fold(label)
        if key == "rumonge":
            rows = [table[p]["communes"][c] for c, p in RUMONGE.items()]
            text = ("Men per 1,000 women among the residents of Bugarama, Burambi, Buyengero, "
                    "Muhuta and Rumonge, counted by the 2008 census and added up: the "
                    "province was made from these communes of Bururi and Bujumbura Rural in "
                    "2015.")
        else:
            home = (by_fold.get(fold(PROVINCE_SPELLING.get(label, label)))
                    or (mairie if "mairie" in key else None))
            if home is None:
                raise SystemExit(f"burundi_census: no province of the census for {label}")
            kept = {c: r for c, r in table[home]["communes"].items()
                    if RUMONGE.get(c) != home}
            rows = list(kept.values())
            if len(kept) == len(table[home]["communes"]):
                text = ("Men per 1,000 women among the province's residents counted by the "
                        "2008 census.")
            else:
                gone = sorted(title(c) for c in table[home]["communes"] if c not in kept)
                text = ("Men per 1,000 women among the residents of the province's communes "
                        "counted by the 2008 census, without " + " and ".join(gone)
                        + ", which have been Rumonge province's since 2015.")
        figures = added(rows)
        records.append(record(
            f"BDI-2008-{shape['id']}", label, level="admin1", parent="BDI", country="BDI",
            match_by="shape_id", shape_id=shape["id"],
            sex_ratio=measure(ratio(figures), unit="males_per_1000_females", year=YEAR,
                              source=SOURCE),
            sex_ratio_note=text, sources=cite))
    return records, report


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--run", default="", help="a label that tells two dispatches apart; ignored")
    ap.parse_args()
    import openpyxl                                 # noqa: PLC0415
    blob = http_get(WORKBOOK, binary=True, cache=False, timeout=300)
    log(f"  read {WORKBOOK}: {len(blob):,} bytes, sha256 {hashlib.sha256(blob).hexdigest()}")
    book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    t16 = read_rows([list(r) for r in book["ETA6"].iter_rows(values_only=True)])
    t14 = read_rows([list(r) for r in book["ETA4"].iter_rows(values_only=True)], (7, 8, 9))
    # Row by row the same people; the two sheets spell some names apart
    # ("KARUSI" and "KARUZI", "VYANDA" and "Vyanda").
    figures = [(r["male"], r["female"], r["total"]) for r in t16]
    if figures != [(r["male"], r["female"], r["total"]) for r in t14]:
        diff = [(a, b) for a, b in zip(t16, t14)
                if (a["male"], a["female"], a["total"]) != (b["male"], b["female"], b["total"])]
        raise SystemExit(f"burundi_census: Tableaux 1.6 and 1.4 disagree: {diff[:5]} "
                         f"({len(t16)} and {len(t14)} rows)")
    table = provinces(t16)
    log(f"  Tableau 1.6: {len(table)} provinces and "
        f"{sum(len(p['communes']) for p in table.values())} communes making "
        f"{NATIONAL[2]:,}, each province's communes making its row; Tableau 1.4 agrees")
    admin1 = json.loads((SITE / "admin1" / "BDI.units.json").read_text())
    admin2 = json.loads((SITE / "admin2" / "BDI.units.json").read_text())
    records, report = build(table, admin1, admin2)
    for line in report:
        log("  " + line)
    for r in records:
        log(f"  {r['level']} {r['name']}: {r['sex_ratio']['value']} men per 1,000 women")
    log(f"  {sum(1 for r in records if r['level'] == 'admin1')} of {len(admin1)} provinces and "
        f"{sum(1 for r in records if r['level'] == 'admin2')} of {len(admin2)} communes written")
    write_json(PROCESSED / OUT, records)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
