#!/usr/bin/env python3
"""Poland: religion, national-ethnic identification and home language by
voivodeship and powiat, from the 2021 census (NSP 2021) final tables.

Statistics Poland publishes the final results as three workbooks on one page
(stat.gov.pl, "Tablice z ostatecznymi danymi w zakresie przynależności
narodowo-etnicznej, języka używanego w domu oraz przynależności do wyznania
religijnego"), one per question, each with a sheet for the country, the 16
voivodeships, the 380 powiats and the gminas. The page and all three files
were requested on a runner and their sheets read before this was written.

**Two of the three questions allow two answers.** National-ethnic
identification and language used at home are each asked twice ("and a
second?"), and the tables count every declaration: Dolnośląskie's 2,904,894
people made 2,883,187 Polish declarations and 62,376 others, which is more
people than live there. So these two compositions are multi-response, as New
Zealand's are, and each share is the percentage of people who named that
identification or language, not a slice of a whole; the note on every record
says so. Religion is asked once and partitions the population -- including the
fifth of it that declined to answer, which is kept as its own bar, because a
composition that quietly renormalised away 20% of Poland would be a different
and worse fact.

**The religion table is a tree.** Its rows carry a classification level from 1
(the unit) to 7 (a single church), and a composition is one cut through that
tree: the Christian branches at level 5, the non-Christian religions at level
4, no religion at level 3 and the non-response at level 2. Every other row is
a subtotal or a detail of one of those and would double-count.

**Names.** GUS writes voivodeships in capitals and powiats as adjectives
("bolesławiecki"); geoBoundaries writes the English voivodeship name and
"powiat bolesławiecki", with eleven land powiats as "X County" and city
powiats by city. The record carries the boundary file's spellings as aliases
and each powiat names its voivodeship, so the two "bielski" powiats scope to
their own.

The long tails -- 190 identifications, 300 languages, 200 churches -- are
not all translated. A named group with an English name here is kept as its own
bar; the rest are summed into an "Other" bar, so nothing is dropped and no
Polish adjective reaches a translated chart.

Usage:
    python -m scripts.fetch_census.poland
"""

from __future__ import annotations

import argparse
import io
import re
from typing import Any

from ._shared import (
    NOT_AVAILABLE, PROCESSED, gap, http_get, log, measure, record, shares, write_json,
)

PAGE = ("https://stat.gov.pl/spisy-powszechne/nsp-2021/nsp-2021-wyniki-ostateczne/"
        "tablice-z-ostatecznymi-danymi-w-zakresie-przynaleznosci-narodowo-etnicznej-"
        "jezyka-uzywanego-w-domu-oraz-przynaleznosci-do-wyznania-religijnego,10,1.html")
BASE = "https://stat.gov.pl/download/gfx/portalinformacyjny/pl/defaultaktualnosci/6536/10/1/1/"
FILES = {
    "religion": BASE + "przynaleznosc_wyznaniowa_-_dane_nsp_2021_dla_kraju_i_jednostek_podzialu_terytorialnego_1.xlsx",
    "ethnicity": BASE + "przynaleznosc_narodowo-etniczna_-_dane_nsp_2021_dla_kraju_i_jednostek_podzialu_terytorialnego_v2.xlsx",
    "language": BASE + "jezyk_uzywany_w_domu_-_dane_nsp_2021_dla_kraju_i_jednostek_podzialu_terytorialnego.xlsx",
}
# Sheet names differ between the workbooks: the religion one has no space.
SHEETS = {
    "religion": {"admin1": "TABL.3", "admin2": "TABL.4"},
    "ethnicity": {"admin1": "TABL. 3", "admin2": "TABL. 4"},
    "language": {"admin1": "TABL. 3", "admin2": "TABL. 4"},
}
SOURCE = "Statistics Poland (GUS), National Population and Housing Census 2021, final results"
LICENCE = "GUS open data (free reuse with attribution)"
YEAR = 2021
EXPECTED = {"admin1": 16, "admin2": 380}

VOIVODESHIPS = {
    "DOLNOŚLĄSKIE": "Lower Silesian Voivodeship",
    "KUJAWSKO-POMORSKIE": "Kuyavian-Pomeranian Voivodeship",
    "LUBELSKIE": "Lublin Voivodeship",
    "LUBUSKIE": "Lubusz Voivodeship",
    "ŁÓDZKIE": "Łódź Voivodeship",
    "MAŁOPOLSKIE": "Lesser Poland Voivodeship",
    "MAZOWIECKIE": "Masovian Voivodeship",
    "OPOLSKIE": "Opole Voivodeship",
    "PODKARPACKIE": "Subcarpathian Voivodeship",
    "PODLASKIE": "Podlaskie Voivodeship",
    "POMORSKIE": "Pomeranian Voivodeship",
    "ŚLĄSKIE": "Silesian Voivodeship",
    "ŚWIĘTOKRZYSKIE": "Świętokrzyskie Voivodeship",
    "WARMIŃSKO-MAZURSKIE": "Warmian-Masurian Voivodeship",
    "WIELKOPOLSKIE": "Greater Poland Voivodeship",
    "ZACHODNIOPOMORSKIE": "West Pomeranian Voivodeship",
}
# geoBoundaries names eleven land powiats after their seat rather than by
# the adjective, keyed by voivodeship because "bielski" is two places.
COUNTY_ALIASES = {
    ("ŚLĄSKIE", "bielski"): "Bielsko County",
    ("ZACHODNIOPOMORSKIE", "kołobrzeski"): "Colberg County",
    ("ŚLĄSKIE", "gliwicki"): "Gliwice County",
    ("MAZOWIECKIE", "grodziski"): "Grodzisk Mazowiecki County",
    ("ŁÓDZKIE", "kutnowski"): "Kutno County",
    ("KUJAWSKO-POMORSKIE", "lipnowski"): "Lipno County",
    ("ŚLĄSKIE", "raciborski"): "Racibórz County",
    ("MAZOWIECKIE", "siedlecki"): "Siedlce County",
    ("ŚLĄSKIE", "tarnogórski"): "Tarnowskie Góry County",
    ("ŚLĄSKIE", "wodzisławski"): "Wodzisław County",
    ("MAZOWIECKIE", "żyrardowski"): "Żyrardów County",
}

ETHNICITY = {
    "Polska": "Polish", "śląska": "Silesian", "kaszubska": "Kashubian",
    "niemiecka": "German", "ukraińska": "Ukrainian", "białoruska": "Belarusian",
    "łemkowska": "Lemko", "romska": "Roma", "rosyjska": "Russian",
    "litewska": "Lithuanian", "żydowska": "Jewish", "czeska": "Czech",
    "słowacka": "Slovak", "ormiańska": "Armenian", "tatarska": "Tatar",
    "karaimska": "Karaim", "wietnamska": "Vietnamese", "angielska": "English",
    "włoska": "Italian", "francuska": "French", "amerykańska": "American",
    "hiszpańska": "Spanish", "grecka": "Greek", "irlandzka": "Irish",
    "holenderska": "Dutch", "brytyjska": "British", "bułgarska": "Bulgarian",
    "chińska": "Chinese", "hinduska": "Indian", "turecka": "Turkish",
    "mazurska": "Masurian", "góralska": "Goral", "kociewska": "Kociewian",
    "norweska": "Norwegian", "szwedzka": "Swedish", "kanadyjska": "Canadian",
    "gruzińska": "Georgian", "austriacka": "Austrian", "belgijska": "Belgian",
    "szkocka": "Scottish", "japońska": "Japanese", "koreańska": "Korean",
    "mołdawska": "Moldovan", "rumuńska": "Romanian", "węgierska": "Hungarian",
    "chorwacka": "Croatian", "serbska": "Serbian", "portugalska": "Portuguese",
    "duńska": "Danish", "fińska": "Finnish", "łotewska": "Latvian",
    "estońska": "Estonian", "macedońska": "Macedonian", "słoweńska": "Slovene",
    "islandzka": "Icelandic", "australijska": "Australian",
    "brazylijska": "Brazilian", "meksykańska": "Mexican", "egipska": "Egyptian",
    "nigeryjska": "Nigerian", "kazachska": "Kazakh", "azerska": "Azerbaijani",
    "uzbecka": "Uzbek", "syryjska": "Syrian", "irańska": "Iranian",
    "iracka": "Iraqi", "kurdyjska": "Kurdish", "palestyńska": "Palestinian",
    "libańska": "Lebanese", "tunezyjska": "Tunisian", "marokańska": "Moroccan",
    "algierska": "Algerian", "filipińska": "Filipino",
    "indonezyjska": "Indonesian", "tajska": "Thai", "pakistańska": "Pakistani",
    "nepalska": "Nepali", "mongolska": "Mongolian", "tybetańska": "Tibetan",
    "rusińska": "Rusyn", "bojkowska": "Boyko", "huculska": "Hutsul",
    "łużycka": "Sorbian", "śląsko-cieszyńska": "Cieszyn Silesian",
    "podlaska": "Podlasian", "pomorska": "Pomeranian", "słowiańska": "Slavic",
    "wielkopolska": "Greater Polish", "kurpiowska": "Kurpian",
    "szwajcarska": "Swiss", "afgańska": "Afghan", "czeczeńska": "Chechen",
    "Nieustalona": "Not stated", "nieustalona": "Not stated",
}
LANGUAGE = {
    "Polski": "Polish", "angielski": "English", "śląski": "Silesian",
    "niemiecki": "German", "kaszubski": "Kashubian", "rosyjski": "Russian",
    "ukraiński": "Ukrainian", "białoruski": "Belarusian", "łemkowski": "Lemko",
    "romski": "Romani", "litewski": "Lithuanian", "francuski": "French",
    "hiszpański": "Spanish", "włoski": "Italian", "niderlandzki": "Dutch",
    "czeski": "Czech", "słowacki": "Slovak", "ormiański": "Armenian",
    "wietnamski": "Vietnamese", "chiński": "Chinese", "grecki": "Greek",
    "hebrajski": "Hebrew", "jidysz": "Yiddish", "arabski": "Arabic",
    "turecki": "Turkish", "węgierski": "Hungarian", "portugalski": "Portuguese",
    "norweski": "Norwegian", "szwedzki": "Swedish", "duński": "Danish",
    "fiński": "Finnish", "japoński": "Japanese", "koreański": "Korean",
    "hindi": "Hindi", "polski język migowy": "Polish Sign Language",
    "ruski": "Ruthenian", "bułgarski": "Bulgarian", "rumuński": "Romanian",
    "chorwacki": "Croatian", "serbski": "Serbian", "gruziński": "Georgian",
    "kazachski": "Kazakh", "mongolski": "Mongolian", "tatarski": "Tatar",
    "łotewski": "Latvian", "estoński": "Estonian", "irlandzki": "Irish",
    "islandzki": "Icelandic", "esperanto": "Esperanto",
    "gwara góralska": "Goral dialect", "flamandzki": "Flemish",
    "macedoński": "Macedonian", "słoweński": "Slovene",
    "afrykanerski": "Afrikaans", "suahili": "Swahili", "perski": "Persian",
    "urdu": "Urdu", "tamilski": "Tamil", "bengalski": "Bengali",
    "nepali": "Nepali", "tajski": "Thai", "indonezyjski": "Indonesian",
    "filipino": "Filipino", "kurdyjski": "Kurdish", "azerski": "Azerbaijani",
    "uzbecki": "Uzbek", "mołdawski": "Moldovan", "czeczeński": "Chechen",
    "Nieustalony": "Not stated", "nieustalony": "Not stated",
}
# The religion tree, cut where the branches partition the population. Each
# entry is (classification level, label as written) -> English name.
RELIGION = {
    (5, "katolicyzm"): "Catholic",
    (5, "chrześcijaństwo wschodnie (ortodoksyjne)"): "Orthodox",
    (5, "protestantyzm i tradycja protestancka"): "Protestant",
    (5, "nurt badaczy pisma świętego"): "Jehovah's Witnesses and Bible Students",
    (5, "inne chrześcijańskie"): "Other Christian",
    (4, "islam"): "Islam",
    (4, "judaizm"): "Judaism",
    (4, "buddyzm"): "Buddhism",
    (4, "hinduizm"): "Hinduism",
    (4, "pogaństwo - rekonstrukcjonizm i neopogaństwo"): "Pagan and neo-pagan",
    (4, "inne religie i wierzenia"): "Other religions",
    (3, "nienależący do żadnego wyznania"): "No religion",
    (2, "nieudzielający odpowiedzi na pytanie o wyznanie"): "Not stated",
    (2, "odmowa odpowiedzi"): "Not stated",
    (2, "odmawiający odpowiedzi na pytanie o wyznanie"): "Not stated",
    (2, "nie ustalono"): "Not stated",
}
NOTES = {
    "ethnicity": ("National-ethnic identification, NSP 2021. Multi-response: a person "
                  "may declare two identifications and every declaration is counted, "
                  "so the shares sum past 100% and each is the percentage of people who "
                  "named that identification, not a slice of a whole. Identifications "
                  "without an English name here are summed as Other."),
    "language": ("Language used at home, NSP 2021. Multi-response: a person may name "
                 "two languages and every answer is counted, so the shares sum past "
                 "100% and each is the percentage of people who use that language at "
                 "home. Languages without an English name here are summed as Other."),
    "religion": ("Religious affiliation, NSP 2021: one answer per person, and answering "
                 "was voluntary. The fifth of the population that declined is kept as "
                 "Not stated rather than renormalised away. Christianity is shown by "
                 "branch (GUS's level 5) and other religions by family (level 4)."),
}


def cell(value: Any) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def workbook(field: str):
    import openpyxl
    raw = http_get(FILES[field], binary=True, timeout=300)
    return openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)


def units_flat(field: str, level: str) -> dict[tuple[str, str], dict[str, Any]]:
    """Ethnicity and language sheets: one row per (unit, category).

    Columns for the voivodeship sheet are Województwo | code | category |
    count | pct; the powiat sheet has Powiat between Województwo and code.
    The unit's name appears once and is blank on its following rows.
    Returns {(voivodeship, unit name): {"__code__": code, "Ogółem": n, label: n}}.
    """
    ws = workbook(field)[SHEETS[field][level]]
    out: dict[tuple[str, str], dict[str, Any]] = {}
    voiv = name = None
    for row in ws.iter_rows(values_only=True):
        vals = [cell(v) for v in row]
        if level == "admin1":
            v, code, label, count = vals[0], vals[1], vals[2], vals[3]
            n = v
        else:
            v, n, code, label, count = vals[0], vals[1], vals[2], vals[3], vals[4]
        if v:
            voiv = v
        if n:
            name = n
        if not (voiv and name and code and label and re.fullmatch(r"\d+", code)):
            continue
        if not re.fullmatch(r"-?\d+", count):
            continue
        unit = out.setdefault((voiv, name), {"__code__": code})
        unit[label] = int(count)
    return out


def units_tree(level: str) -> dict[tuple[str, str], dict[str, Any]]:
    """The religion sheets: a classification tree, one row per node.

    Measured layouts. Voivodeship sheet: level | six label columns | count |
    three percentages; the level-1 row's label is the unit (POLSKA or a
    voivodeship). Powiat sheet: Województwo | Powiat | Kod powiatu | level |
    five label columns | count | percentages; the level-1 row's label is
    "Ogółem" and the unit is named in the first two columns, the voivodeship
    only on its first powiat. The cut through the tree is RELIGION above; the
    non-response arrives as two level-2 rows (declined, not established) and
    both are summed into Not stated.
    """
    ws = workbook("religion")[SHEETS["religion"][level]]
    out: dict[tuple[str, str], dict[str, Any]] = {}
    if level == "admin1":
        level_col, label_cols, count_col = 0, range(1, 7), 7
    else:
        level_col, label_cols, count_col = 3, range(4, 9), 9
    voiv = name = code = None
    for row in ws.iter_rows(values_only=True):
        vals = [cell(v) for v in row]
        if len(vals) <= count_col or not re.fullmatch(r"[1-7]", vals[level_col]):
            continue
        lvl = int(vals[level_col])
        label = next((vals[i] for i in label_cols if vals[i]), "")
        count = vals[count_col]
        if not label or not re.fullmatch(r"-?\d+", count):
            continue
        if lvl == 1:
            if level == "admin1":
                voiv = name = label
                code = None
            else:
                voiv = vals[0] or voiv
                name = vals[1]
                code = vals[2]
            out[(voiv, name)] = {"__code__": code or name, "__total__": int(count)}
            continue
        if not (voiv and name):
            continue
        english = RELIGION.get((lvl, label.lower()))
        if english is None:
            continue
        counts = out[(voiv, name)]
        counts[english] = counts.get(english, 0) + int(count)
    return out


def englished(field: str, counts: dict[str, int]) -> tuple[dict[str, int], int]:
    """Translate what has a name; sum the rest as Other. Returns (groups, total)."""
    table = {"ethnicity": ETHNICITY, "language": LANGUAGE}[field]
    total = counts.get("Ogółem", 0)
    out: dict[str, int] = {}
    other = 0
    for label, n in counts.items():
        if label.startswith("__") or label == "Ogółem" \
                or label.lower().startswith(("inna niż polska", "inny niż polski")):
            continue                      # the code, the total, the non-Polish subtotal
        english = table.get(label)
        if english is None:
            other += n
        else:
            out[english] = out.get(english, 0) + n
    if other:
        out["Other"] = other
    return out, total


def powiat_names(voiv: str, raw: str) -> tuple[str, list[str]]:
    """GUS's powiat label -> the record's name and the boundary file's spellings."""
    raw = raw.strip()
    city = re.sub(r"^(m\.\s*st\.|m\.st\.|m\.)\s*", "", raw, flags=re.I)
    if city != raw or raw[:1].isupper():
        return city, []                   # a city powiat, named for its city
    alias = COUNTY_ALIASES.get((voiv, raw))
    return f"powiat {raw}", [alias] if alias else []


def build(level: str) -> list[dict[str, Any]]:
    log(f"poland: {level}")
    fields = {
        "ethnicity": units_flat("ethnicity", level),
        "language": units_flat("language", level),
        "religion": units_tree(level),
    }
    for field, units in fields.items():
        log(f"  {field}: {len(units)} units")
    keys = sorted(set().union(*(set(u) for u in fields.values())))
    keys = [k for k in keys if k[1].upper() != "POLSKA"]
    if len(keys) != EXPECTED[level]:
        raise SystemExit(f"poland: {level} has {len(keys)} units, expected {EXPECTED[level]}: "
                         f"{[k[1] for k in keys][:20]}")
    records = []
    for voiv, raw in keys:
        code = next(str(fields[f][(voiv, raw)]["__code__"]) for f in fields
                    if (voiv, raw) in fields[f])
        if level == "admin1":
            name = VOIVODESHIPS.get(voiv.upper())
            if name is None:
                raise SystemExit(f"poland: unknown voivodeship {voiv!r}")
            aliases = [voiv.title(), voiv.lower()]
            parent, parent_name = "POL", None
            entity_id = f"POL-{code}"
        else:
            name, aliases = powiat_names(voiv, raw)
            parent_name = VOIVODESHIPS.get(voiv.upper())
            if parent_name is None:
                raise SystemExit(f"poland: unknown voivodeship {voiv!r} for powiat {raw!r}")
            parent = f"POL-{code[:2]}"
            entity_id = f"POL-{code}"
        values: dict[str, Any] = {}
        total = 0
        for field in ("ethnicity", "language"):
            counts = fields[field].get((voiv, raw))
            if not counts:
                values[field] = gap(NOT_AVAILABLE, f"{field} row missing from the GUS table")
                continue
            groups, total = englished(field, counts)
            values[field] = shares(groups, total=total) or gap(NOT_AVAILABLE)
            values[f"{field}_note"] = NOTES[field]
        rel = fields["religion"].get((voiv, raw))
        if rel:
            rtotal = rel["__total__"]
            groups = {k: v for k, v in rel.items() if not k.startswith("__")}
            summed = sum(groups.values())
            if abs(summed - rtotal) > rtotal * 0.005:
                raise SystemExit(f"poland: {raw} religion sums to {summed:,} against "
                                 f"{rtotal:,}; the cut through the tree is wrong")
            values["religion"] = shares(groups, total=rtotal) or gap(NOT_AVAILABLE)
            values["religion_note"] = NOTES["religion"]
            total = total or rtotal
        else:
            values["religion"] = gap(NOT_AVAILABLE, "religion row missing from the GUS table")
        records.append(record(
            entity_id, name, level=level, parent=parent, parent_name=parent_name,
            country="POL", aliases=aliases, codes={"teryt": code},
            population=measure(total, year=YEAR, source=SOURCE) if total else gap(NOT_AVAILABLE),
            sources=[{"field": "population/religion/ethnicity/language", "name": SOURCE,
                      "url": PAGE, "license": LICENCE}],
            **values,
        ))
    log(f"  {len(records)} records")
    return records


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--level", default="both", choices=["admin1", "admin2", "both"])
    args = ap.parse_args()
    if args.level in ("admin1", "both"):
        write_json(PROCESSED / "poland_voivodeship.json", build("admin1"))
    if args.level in ("admin2", "both"):
        write_json(PROCESSED / "poland_powiat.json", build("admin2"))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
