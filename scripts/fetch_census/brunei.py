#!/usr/bin/env python3
"""Brunei -- race and religion by district, population by mukim, BPP 2021.

The Department of Economic Planning and Statistics ran the sixth Population
and Housing Census (Banci Penduduk dan Perumahan, BPP) 2021 through 2021 and
reported it in *Report of the Population and Housing Census (BPP) 2021:
Demographic, Household and Housing Characteristics*, October 2022. The report
itself is 94 pages of narrative and charts; the figures are in its three
annexes, which DEPS also publishes as one workbook:

    https://deps.mofe.gov.bn/wp-content/uploads/2025/11/EXCEL-TABLE-A-C.xlsx

Two of its sheets are compositions:

* **A3** -- Population by Race, District and Sex, 2021.
* **A4** -- Population by Religion, District and Sex, 2021.

and one gives the level below:

* **C1** (on the sheet ``C1-C5``) -- Population by Mukim, Residential Status
  and Sex, 2021, for each of the four districts in turn.

**Race and religion stop at the district.** That is measured, not assumed.
Annex A's twelve tables cross race and religion with district, age and
residential status; Annex B's sixteen give population, households and occupied
living quarters by district, mukim and kampung; Annex C's ten give mukim and
kampung by residential status and age. No table in any of them puts race or
religion below the district, so the 38 mukims carry a head count and a stated
gap, and the gap names the tables that were read.

**Language is asked and not published.** The BPP 2021 questionnaire asks two
language questions -- E26, the languages a person can read and write, marked
*all that apply*, and E27, "the language mainly spoken at home", marked *one
only* -- and E27 is exactly a composition. None of the 38 annex tables reports
it, and the report's own Concepts and Definitions defines race, religion,
marital status, country of birth and nationality and says nothing of language.
So language here is ``not_available`` with that reason, never
``not_collected``: the question was asked, and the answer was not tabulated.

**What the categories mean.** Brunei's race question (E09) offers three
groups, and "Malay" is an administrative category rather than an ethnonym: the
report's own definition is that a Brunei Malay is "the persons belonging to
one of the following ethnic groups of the Malay race, namely Brunei, Tutong,
Belait, Kedayan, Dusun, Bisaya or Murut", and the questionnaire numbers those
seven under it. Chinese is its own group and "Others" is, again in the
report's words, "the rest of the population not included in the Malay and
Chinese racial groups" -- which in a country where 18.4% of the population are
temporary residents is largely foreign workers, and also the Malays of
Malaysia and Indonesia, the Ibans, and everyone else. The seven groups inside
Malay are published for the country only, never by district, so nothing here
splits Malay and nothing here renames it.

Religion (E10) offered Islam, Christianity, Buddhism, Hinduism and Others, and
the published table has four columns: the report states that "the other
religions, unstated faiths and no religious beliefs were grouped into
'Others'". That bucket welds a real answer to a non-answer, so it is carried
as **Other, none, or not stated** -- the project's existing label for exactly
that -- rather than as "Other religions", which would count the irreligious as
adherents of something.

**Self-checks.** The workbook's national column must reproduce the figures the
report's Executive Summary prints in prose -- 297,016 Malays, 42,132 Chinese,
101,567 Others; 362,035 Muslims, 29,462 Christians, 27,745 Buddhists, 21,473
Others; 440,715 people in all -- and the four districts must sum to each of
them. Each district's categories must sum to the district's own printed total,
and the mukims of a district must sum to the district. Any of those failing is
a ``SystemExit``: the sheet names are DEPS's and a re-issue could move a
column without changing one.

**Two names, and why they are not guesses.** geoBoundaries draws 38 mukims for
Brunei where the census counts 39. The differences are exactly two, and both
are forced rather than chosen:

* The census counts **Gadong A** and **Gadong B**; the boundary file draws one
  **Gadong**. Brunei Muara has 18 census mukims and 17 shapes, every other
  name matching outright, so the one shape is the two mukims. Their head
  counts are added and the record says so. Nothing else is summed, because
  nothing else needs to be.
* The census writes **Bokok**; the boundary file writes **Bunkok**. Temburong
  has five mukims in both, four of them identical (Amo, Bangar, Batu Apoi,
  Labu), so the fifth is one mukim under two spellings. It is declared as an
  alias, not matched by resemblance.

**The office moved its documents.** Every ``deps.mofe.gov.bn/DEPD Documents
Library/...`` link still printed on the department's own pages now answers 404;
the files live under ``wp-content/uploads/`` and were found through the site's
WordPress media API. The URLs below are the live ones, and the host answers an
ordinary verified client.

Usage:
    python -m scripts.fetch_census.brunei
    python -m scripts.fetch_census.brunei --probe
"""

from __future__ import annotations

import argparse
import io
import re
from typing import Any

from ._shared import (
    NOT_AVAILABLE, PROCESSED, dated, gap, http_get, log, measure, record,
    shares, write_json,
)

WORKBOOK = "https://deps.mofe.gov.bn/wp-content/uploads/2025/11/EXCEL-TABLE-A-C.xlsx"
REPORT = "https://deps.mofe.gov.bn/wp-content/uploads/2025/11/RPT-2.pdf"
ANNEX_A = "https://deps.mofe.gov.bn/wp-content/uploads/2025/11/ANNEX-A.pdf"
ANNEX_B = "https://deps.mofe.gov.bn/wp-content/uploads/2025/11/ANNEX-B.pdf"
ANNEX_C = "https://deps.mofe.gov.bn/wp-content/uploads/2025/11/ANNEX-C.pdf"
QUESTIONNAIRE = "https://deps.mofe.gov.bn/wp-content/uploads/2025/11/Q_BPP2021.pdf"
PAGE = "https://deps.mofe.gov.bn/population-social-statistics/"

DEPS = "Department of Economic Planning and Statistics, Brunei Darussalam"
CENSUS = ("Population and Housing Census (BPP) 2021, "
          f"{DEPS}")
YEAR = 2021

# The four districts, spelled as the workbook's column headers spell them.
DISTRICTS = ("Brunei Muara", "Belait", "Tutong", "Temburong")
# geoBoundaries writes the first one hyphenated. Declared rather than left to
# the matcher's hyphen rule, because a district that fails to join here is a
# bug: there are only four of them.
DISTRICT_ALIASES = {
    "Brunei Muara": ["Brunei-Muara", "Brunei and Muara", "Daerah Brunei Muara"],
}

# Table A3's row labels -> the group as this map writes it. The census's own
# English is "Malays" / "Chinese" / "Others"; "Malay" is the singular the rest
# of the map already uses for the same group (Singapore's CMIO reader makes
# the same change), and nothing else is altered.
RACE_LABELS = {"Malays": "Malay", "Chinese": "Chinese", "Others": "Others"}
# Table A4's, with the residual named for what DEPS says is inside it.
RELIGION_LABELS = {
    "Islam": "Islam",
    "Christianity": "Christianity",
    "Buddhism": "Buddhism",
    "Others": "Other, none, or not stated",
}
# The row that closes each table.
TOTAL_ROW = "Jumlah/Total"

# The 39 mukims, by district, spelled as table C1 spells them. Declared rather
# than discovered so that a renamed or dropped mukim is refused instead of
# quietly leaving a shape empty.
MUKIMS: dict[str, tuple[str, ...]] = {
    "Brunei Muara": (
        "Kianggeh", "Sungai Kedayan", "Saba", "Sungai Kebun",
        "Burong Pingai Ayer", "Peramu", "Tamoi", "Berakas A", "Berakas B",
        "Gadong A", "Gadong B", "Kota Batu", "Lumapas", "Kilanas",
        "Sengkurong", "Pangkalan Batu", "Mentiri", "Serasa"),
    "Belait": (
        "Kuala Belait", "Seria", "Liang", "Kuala Balai", "Labi",
        "Bukit Sawat", "Sukang", "Melilas"),
    "Tutong": (
        "Pekan Tutong", "Keriam", "Telisai", "Tanjong Maya", "Kiudang",
        "Lamunin", "Ukong", "Rambai"),
    "Temburong": ("Bangar", "Amo", "Labu", "Batu Apoi", "Bokok"),
}
# The boundary file draws one Gadong where the census counts two mukims; see
# the module docstring. The two head counts are added under the shape's name.
MERGED = {"Gadong": ("Gadong A", "Gadong B")}
# And writes Bokok's name differently.
MUKIM_ALIASES = {"Bokok": ["Bunkok"]}

# What the report's Executive Summary prints in prose, on printed pages 42 and
# 51-52. The workbook is a different document by the same office, so this is a
# cross-check and not a restatement: a column read one place to the left still
# sums to the total beside it, and would not survive this.
PUBLISHED_TOTAL = 440_715
PUBLISHED_RACE = {"Malay": 297_016, "Chinese": 42_132, "Others": 101_567}
PUBLISHED_RELIGION = {"Islam": 362_035, "Christianity": 29_462,
                      "Buddhism": 27_745, "Other, none, or not stated": 21_473}
# Table B1's district totals, printed again in chapter 4 of the report.
PUBLISHED_DISTRICT = {"Brunei Muara": 318_530, "Belait": 65_531,
                      "Tutong": 47_210, "Temburong": 9_444}

RACE_NOTE = (
    "Brunei's 2021 census (BPP 2021), Table A3 of its annexes: population by "
    "race, district and sex. \"Malay\" is the state's administrative category "
    "and covers the seven groups the report names as ethnic groups of the "
    "Malay race -- Brunei, Tutong, Belait, Kedayan, Dusun, Bisaya and Murut -- "
    "which are published for the country only and are not split here. "
    "\"Others\" is the report's own residual, everyone outside Malay and "
    "Chinese, and in a country where 18.4% of those counted were temporary "
    "residents it is largely foreign workers.")
RELIGION_NOTE = (
    "Brunei's 2021 census (BPP 2021), Table A4 of its annexes: population by "
    "religion, district and sex. Islam is the state religion. The census "
    "offered Islam, Christianity, Buddhism, Hinduism and Others and published "
    "four columns: the report states that other religions, unstated faiths "
    "and no religious belief were grouped into the last, which is why it is "
    "labelled for all three rather than as \"other religions\".")
LANGUAGE_NOTE = (
    "Not published. Brunei's 2021 census does ask language -- question E27, "
    "\"the language mainly spoken at home\", one answer per person, beside "
    "E26 on languages read and written -- and tabulates neither: none of the "
    "38 tables in Annexes A, B and C of the census report carries a language "
    "column at any level, and the report's Concepts and Definitions defines "
    "race, religion, marital status, country of birth and nationality and not "
    "language. Asked and not published, which is a different thing from not "
    "asked.")


def mukim_gap_note(people: int, merged: tuple[str, ...] | None) -> str:
    """Why a mukim has a head count and no composition."""
    counted = (f"The 2021 census counted {people:,} people in this mukim "
               f"(Table C1 of the census annexes)")
    if merged:
        counted = (f"The 2021 census counted {people:,} people here, the sum "
                   f"of its {' and '.join(merged)} (Table C1 of the census "
                   f"annexes)")
    return (f"{counted}, but Brunei publishes race and religion by district "
            f"only. Its mukim tables -- Annex B on population, households and "
            f"occupied living quarters, Annex C on residential status and age "
            f"-- carry no composition of any kind, and Annex A, which carries "
            f"race and religion, stops at the four districts.")


def sources(fields: str) -> list[dict[str, Any]]:
    return [
        {"field": fields, "name": f"{CENSUS}, annex tables A and C",
         "url": WORKBOOK, "year": YEAR},
        {"field": fields,
         "name": ("Report of the Population and Housing Census (BPP) 2021: "
                  "Demographic, Household and Housing Characteristics "
                  "(the national figures the workbook is checked against)"),
         "url": REPORT, "year": YEAR},
    ]


def slug(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "-", name).strip("-")


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def number(value: Any) -> int | None:
    """A count, or None where the cell is not one."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(round(value))
    body = text(value).replace(",", "").replace(" ", "")
    if not body or not re.fullmatch(r"-?\d+(\.0+)?", body):
        return None
    return int(round(float(body)))


def read_workbook(url: str = WORKBOOK) -> dict[str, list[list[Any]]]:
    """Every sheet of the annex workbook, as rows of raw cell values."""
    import openpyxl  # noqa: PLC0415
    blob = http_get(url, binary=True, timeout=300)
    book = openpyxl.load_workbook(io.BytesIO(blob), data_only=True,
                                  read_only=True)
    sheets = {name: [list(row) for row in book[name].iter_rows(values_only=True)]
              for name in book.sheetnames}
    book.close()
    return sheets


def district_columns(rows: list[list[Any]], sheet: str) -> dict[str, int]:
    """Which column each district's "Persons" figures are in.

    The header is a merged cell, so the district's name sits in the same
    column as the first of its three ``Orang / Lelaki / Perempuan`` columns,
    and "Jumlah" does the same for the national column. Found rather than
    counted, because the same tables are published as .xls and .xlsx with
    different amounts of padding between the blocks.
    """
    for row in rows:
        seen: dict[str, int] = {}
        for i, cell in enumerate(row):
            label = text(cell)
            if label and label not in seen:
                seen[label] = i
        if all(name in seen for name in DISTRICTS) and "Jumlah" in seen:
            return {"Jumlah": seen["Jumlah"],
                    **{name: seen[name] for name in DISTRICTS}}
    raise SystemExit(f"brunei: sheet {sheet!r} has no row naming all four "
                     f"districts and a total; DEPS changed the table")


def read_composition(rows: list[list[Any]], labels: dict[str, str], sheet: str
                     ) -> tuple[dict[str, dict[str, int]], dict[str, int]]:
    """One of tables A3 and A4 -> ``{place: {group: count}}`` and the totals.

    ``place`` is "Jumlah" for the country and the district's name otherwise.

    Every table is printed bilingually, the Malay caption on its own row above
    the English one that carries the figures, and for religion the two are the
    same word: A4 has "Islam" twice, once empty and once with the counts. So a
    row whose every figure is blank is the caption and is passed over, a row
    with some figures and not others is refused, and a second row of figures
    under one label is refused too. A label the table prints and this does not
    know is ignored on purpose *only* when it is a caption -- a fifth race or
    a split religion column arrives as an unknown label with figures, and the
    check against the report's own national counts is what catches it.
    """
    columns = district_columns(rows, sheet)
    counts: dict[str, dict[str, int]] = {place: {} for place in columns}
    totals: dict[str, int] = {}
    for row in rows:
        label = text(row[0] if row else None)
        if label != TOTAL_ROW and label not in labels:
            continue
        values = {place: (number(row[column]) if column < len(row) else None)
                  for place, column in columns.items()}
        if all(value is None for value in values.values()):
            continue                      # the Malay caption above the figures
        blank = sorted(p for p, v in values.items() if v is None)
        if blank:
            raise SystemExit(f"brunei: sheet {sheet!r}, row {label!r}: no "
                             f"count for {blank}")
        if label == TOTAL_ROW:
            if totals:
                raise SystemExit(f"brunei: sheet {sheet!r} has two total rows")
            totals = values                                   # type: ignore[assignment]
            continue
        group = labels[label]
        if group in counts["Jumlah"]:
            raise SystemExit(f"brunei: sheet {sheet!r} prints {label!r} twice")
        for place, value in values.items():
            counts[place][group] = value
    missing = {place: sorted(set(labels.values()) - set(comp))
               for place, comp in counts.items() if len(comp) != len(labels)}
    if missing or not totals:
        raise SystemExit(f"brunei: sheet {sheet!r} is incomplete: "
                         f"missing {missing}, totals {sorted(totals)}")
    return counts, totals


def check_composition(sheet: str, counts: dict[str, dict[str, int]],
                      totals: dict[str, int],
                      published: dict[str, int]) -> None:
    """The national column against the report, then every column against itself."""
    national = counts["Jumlah"]
    for group, expected in published.items():
        got = national.get(group)
        if got != expected:
            raise SystemExit(
                f"brunei: sheet {sheet!r} gives {group} {got:,} nationally "
                f"where the BPP 2021 report prints {expected:,}; this is not "
                f"the 2021 census table")
    if totals["Jumlah"] != PUBLISHED_TOTAL:
        raise SystemExit(f"brunei: sheet {sheet!r} totals {totals['Jumlah']:,}, "
                         f"the report's {PUBLISHED_TOTAL:,}")
    for place, comp in counts.items():
        summed = sum(comp.values())
        if summed != totals[place]:
            raise SystemExit(f"brunei: sheet {sheet!r}, {place}: the groups sum "
                             f"to {summed:,} against a printed total of "
                             f"{totals[place]:,}")
    for district in DISTRICTS:
        if totals[district] != PUBLISHED_DISTRICT[district]:
            raise SystemExit(
                f"brunei: sheet {sheet!r} gives {district} {totals[district]:,} "
                f"people, the census's {PUBLISHED_DISTRICT[district]:,}")
    for group in published:
        summed = sum(counts[d].get(group, 0) for d in DISTRICTS)
        if summed != national[group]:
            raise SystemExit(f"brunei: sheet {sheet!r}: the districts give "
                             f"{group} {summed:,}, the national column "
                             f"{national[group]:,}")
    log(f"  {sheet}: national column matches the report "
        + ", ".join(f"{g} {national[g]:,}" for g in published)
        + "; the four districts sum to it and to "
          f"{totals['Jumlah']:,} people")


def read_mukims(rows: list[list[Any]]) -> dict[str, dict[str, int]]:
    """Table C1 -> ``{district: {mukim: persons}}``.

    The sheet holds C1 and then C2 to C5, which list kampungs under the same
    mukim headings, so the walk is anchored: it starts at the first row whose
    first cell is the district's own name and takes the rows whose first cell
    is one of that district's declared mukims, in the order C1 prints them,
    stopping as soon as all of them are found. The head count is the row's
    first number, because it is the first of the twelve columns and that is
    true whatever padding the file carries.
    """
    out: dict[str, dict[str, int]] = {}
    for district, names in MUKIMS.items():
        start = next((i for i, row in enumerate(rows)
                      if text(row[0] if row else None) == district), None)
        if start is None:
            raise SystemExit(f"brunei: table C1 has no row for {district}")
        wanted, found = list(names), {}
        for row in rows[start + 1:]:
            label = text(row[0] if row else None)
            if label not in wanted:
                continue
            value = next((n for n in (number(c) for c in row[1:])
                          if n is not None), None)
            if value is None:
                raise SystemExit(f"brunei: table C1, {district}/{label}: "
                                 f"no head count on the row")
            found[label] = value
            wanted.remove(label)
            if not wanted:
                break
        if wanted:
            raise SystemExit(f"brunei: table C1 is missing {wanted} under "
                             f"{district}; DEPS renamed or dropped a mukim")
        summed = sum(found.values())
        if summed != PUBLISHED_DISTRICT[district]:
            raise SystemExit(f"brunei: {district}'s mukims sum to {summed:,} "
                             f"against the district's "
                             f"{PUBLISHED_DISTRICT[district]:,}")
        out[district] = found
    log(f"  C1: {sum(len(v) for v in out.values())} mukims read; each "
        f"district's mukims sum to its printed total")
    return out


def district_records(race: dict[str, dict[str, int]],
                     religion: dict[str, dict[str, int]],
                     totals: dict[str, int]) -> list[dict[str, Any]]:
    out = []
    for name in DISTRICTS:
        people = totals[name]
        ethnicity = shares(race[name], total=people)
        faith = shares(religion[name], total=people)
        out.append(record(
            f"BRN-{slug(name)}", name, level="admin1", parent="BRN",
            country="BRN", aliases=DISTRICT_ALIASES.get(name, []),
            population=measure(people, year=YEAR, source=CENSUS,
                               unit="persons"),
            ethnicity=ethnicity, ethnicity_year=dated(ethnicity, YEAR),
            ethnicity_note=RACE_NOTE,
            religion=faith, religion_year=dated(faith, YEAR),
            religion_note=RELIGION_NOTE,
            language=gap(NOT_AVAILABLE, LANGUAGE_NOTE),
            sources=sources("ethnicity/religion/population"),
        ))
    return out


def mukim_records(mukims: dict[str, dict[str, int]]) -> list[dict[str, Any]]:
    out = []
    for district, counts in mukims.items():
        merged_parts = {part for parts in MERGED.values() for part in parts}
        names = [n for n in counts if n not in merged_parts]
        names += [shape for shape, parts in MERGED.items()
                  if all(part in counts for part in parts)]
        for name in sorted(names):
            parts = MERGED.get(name)
            people = (sum(counts[p] for p in parts) if parts
                      else counts[name])
            note = mukim_gap_note(people, parts)
            out.append(record(
                f"BRN-{slug(district)}-{slug(name)}", name, level="admin2",
                parent=f"BRN-{slug(district)}", parent_name=district,
                parent_aliases=DISTRICT_ALIASES.get(district, []),
                country="BRN", aliases=MUKIM_ALIASES.get(name, []),
                population=measure(people, year=YEAR, source=CENSUS,
                                   unit="persons"),
                ethnicity=gap(NOT_AVAILABLE, note),
                religion=gap(NOT_AVAILABLE, note),
                language=gap(NOT_AVAILABLE, LANGUAGE_NOTE),
                sources=sources("population"),
            ))
    return out


def build(sheets: dict[str, list[list[Any]]]) -> list[dict[str, Any]]:
    for name in ("A3", "A4", "C1-C5"):
        if name not in sheets:
            raise SystemExit(f"brunei: the annex workbook has no sheet {name!r}; "
                             f"it holds {sorted(sheets)}")
    race, race_totals = read_composition(sheets["A3"], RACE_LABELS, "A3")
    check_composition("A3", race, race_totals, PUBLISHED_RACE)
    religion, faith_totals = read_composition(sheets["A4"], RELIGION_LABELS, "A4")
    check_composition("A4", religion, faith_totals, PUBLISHED_RELIGION)
    if race_totals != faith_totals:
        raise SystemExit("brunei: A3 and A4 disagree about how many people "
                         "each district holds")
    mukims = read_mukims(sheets["C1-C5"])
    return district_records(race, religion, race_totals) + mukim_records(mukims)


def probe() -> int:
    """What the workbook holds, sheet by sheet, and nothing written."""
    sheets = read_workbook()
    log(f"probe: {WORKBOOK}")
    for name, rows in sheets.items():
        width = max((len(r) for r in rows), default=0)
        log(f"  -- {name!r}: {len(rows)} rows x {width} cols")
        for row in rows[:6]:
            log("     " + " | ".join(text(c) for c in row[:14]))
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--probe", action="store_true",
                    help="print the annex workbook's sheets and write nothing")
    args = ap.parse_args()
    if args.probe:
        return probe()
    log("brunei: race and religion by district, population by mukim, BPP 2021")
    rows = build(read_workbook())
    districts = [r for r in rows if r["level"] == "admin1"]
    mukims = [r for r in rows if r["level"] == "admin2"]
    log(f"  {len(districts)} districts with race and religion; "
        f"{len(mukims)} mukims with a head count and a stated gap; "
        f"language not published at any level")
    write_json(PROCESSED / "brunei.json", rows)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
