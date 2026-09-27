#!/usr/bin/env python3
"""Probe the cross-national survey portals: what is open, what sits behind a login.

The European Social Survey, the European Values Study, Eurobarometer, ISSP and
Pew each publish through a portal of their own, several of them single-page
applications whose data calls are only visible in their scripts. Each argument
is one probe, ``mode:url`` with optional ``~~`` parts after it:

    head:URL                     status, final URL after redirects, type and size
    text:URL[~~N]                the first N characters of the body, tags stripped
    grep:URL~~REGEX[~~N]         up to N matches of REGEX in the raw body (with context)
    links:URL~~REGEX[~~N]        up to N links (href and text) matching REGEX
    js:URL~~REGEX[~~N]           the page's own scripts, each searched for REGEX
    post:URL~~B64JSON[~~N]       POST a JSON body (base64), print the first N characters
    zip:URL                      a zip archive's members
    xls:URL[~~ROWS[~~SHEETS]]    a workbook's sheets and first rows
    csv:URL[~~ROWS]              the first rows of a CSV (or gzipped CSV)
    cdx:URLPREFIX[~~N]           Internet Archive captures under a URL prefix
    pdf:URL~~REGEX[~~N]          lines of a PDF's text matching REGEX

The workflow splits the command on whitespace, so no part may hold a space: a
regex says ``\\s`` instead. Read-only: nothing is written, the log is the output.
No credential is sent anywhere and no login is attempted -- a 401, a 403 or a
redirect to a sign-in page is recorded as the answer.

Usage:
    python -m scripts.fetch_census.surveys_probe head:https://ess.sikt.no/
"""

from __future__ import annotations

import base64
import csv
import gzip
import html
import io
import json
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from common import USER_AGENT  # noqa: E402

TIMEOUT = 60
MODES = ("head", "text", "grep", "links", "js", "post", "zip", "xls", "csv", "cdx", "pdf")


def uri(url: str) -> str:
    return urllib.parse.quote(url, safe="%/:=&?~#+!$,;'@()*[]")


def fetch(url: str, data: bytes | None = None, headers: dict | None = None,
          timeout: int = TIMEOUT) -> tuple[int, str, dict, bytes]:
    """(status, final url, headers, body); an HTTP error is an answer, not a raise."""
    hdrs = {"User-Agent": USER_AGENT, "Accept": "*/*"}
    hdrs.update(headers or {})
    req = urllib.request.Request(uri(url), data=data, headers=hdrs)
    for wait in (5, 20, None):
        try:
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                body = resp.read()
                if resp.headers.get("Content-Encoding") == "gzip":
                    body = gzip.decompress(body)
                return resp.status, resp.geturl(), dict(resp.headers), body
        except urllib.error.HTTPError as exc:
            if exc.code in (429, 503) and wait is not None and "archive.org" in url:
                time.sleep(wait)
                continue
            try:
                body = exc.read() or b""
            except Exception:  # noqa: BLE001 - a body we cannot read is empty
                body = b""
            return exc.code, exc.geturl() or url, dict(exc.headers or {}), body
        except (urllib.error.URLError, TimeoutError, ConnectionError) as exc:
            if wait is None:
                return -1, url, {}, str(exc).encode()
            time.sleep(wait)
    return -1, url, {}, b""


def decode(body: bytes, headers: dict) -> str:
    m = re.search(r"charset=([\w-]+)", headers.get("Content-Type", "") or "")
    for enc in ([m.group(1)] if m else []) + ["utf-8", "cp1252", "latin-1"]:
        try:
            return body.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return body.decode("utf-8", "replace")


def strip(text: str) -> str:
    text = re.sub(r"(?is)<(script|style)\b.*?</\1>", " ", text)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    return re.sub(r"\s+", " ", html.unescape(text)).strip()


def say(line: str) -> None:
    print(line, flush=True)


def summary(url: str, status: int, final: str, headers: dict, body: bytes) -> None:
    moved = f" -> {final}" if final and final != uri(url) and final != url else ""
    extra = ""
    for h in ("Content-Disposition", "WWW-Authenticate", "Location"):
        if headers.get(h):
            extra += f" {h}={headers[h][:160]}"
    say(f"  HTTP {status}{moved} {headers.get('Content-Type', '?')} {len(body):,} bytes{extra}")


def show_row(row: list, width: int = 24, cols: int = 14) -> str:
    cells = [str(c).strip().replace("\n", " ")[:width] for c in row[:cols]]
    while cells and not cells[-1]:
        cells.pop()
    return " | ".join(cells)


def grep(text: str, rx: str, n: int, ctx: int = 80) -> None:
    seen: set[str] = set()
    for m in re.finditer(rx, text, re.I):
        a, b = max(0, m.start() - ctx), min(len(text), m.end() + ctx)
        snippet = re.sub(r"\s+", " ", text[a:b])
        key = m.group(0)
        if key in seen:
            continue
        seen.add(key)
        say(f"  ~ {snippet}")
        if len(seen) >= n:
            break
    if not seen:
        say("  (no match)")


def probe(arg: str) -> None:
    mode, _, rest = arg.partition(":")
    if mode not in MODES:
        say(f"?? unknown mode in {arg!r}")
        return
    parts = rest.split("~~")
    url, extra = parts[0], parts[1:]
    say(f"== {mode} {url} {' '.join(extra)}")
    if mode == "cdx":
        n = int(extra[0]) if extra else 60
        q = ("https://web.archive.org/cdx/search/cdx?url=" + urllib.parse.quote(url, safe="")
             + "&matchType=prefix&filter=statuscode:200&collapse=urlkey&limit=" + str(n * 4))
        status, _, headers, body = fetch(q, timeout=120)
        lines = decode(body, headers).splitlines()
        say(f"  HTTP {status}: {len(lines)} captures")
        for line in lines[:n]:
            f = line.split(" ")
            if len(f) >= 7:
                say(f"  {f[1]} {f[3]} {f[6]} {f[2]}")
        return
    if mode == "post":
        payload = base64.b64decode(extra[0])
        status, final, headers, body = fetch(url, data=payload,
                                             headers={"Content-Type": "application/json",
                                                      "Accept": "application/json"})
        summary(url, status, final, headers, body)
        say("  " + decode(body, headers)[:int(extra[1]) if len(extra) > 1 else 3000])
        return
    big = mode in ("zip", "xls", "csv", "pdf")
    status, final, headers, body = fetch(url, timeout=300 if big else TIMEOUT)
    summary(url, status, final, headers, body)
    if mode == "head":
        if status >= 400 or status < 0:
            say("  " + strip(decode(body, headers))[:300])
        else:
            text = decode(body, headers)
            title = re.search(r"(?is)<title[^>]*>(.*?)</title>", text)
            if title:
                say(f"  title: {strip(title.group(1))[:160]}")
            if re.search(r"(?i)log\s*in|sign\s*in|anmelden|register", text):
                grep(strip(text), r"(?:log\s*in|sign\s*in|anmelden|registr\w*)", 4, 60)
        return
    if status >= 400 or status < 0:
        say("  " + strip(decode(body, headers))[:400])
        return
    if mode == "text":
        n = int(extra[0]) if extra else 2000
        say("  " + strip(decode(body, headers))[:n])
    elif mode == "grep":
        grep(decode(body, headers), extra[0], int(extra[1]) if len(extra) > 1 else 20)
    elif mode == "links":
        rx = re.compile(extra[0], re.I) if extra and extra[0] else None
        n = int(extra[1]) if len(extra) > 1 else 60
        text = decode(body, headers)
        shown = 0
        for m in re.finditer(r"(?is)<a\b[^>]*href=[\"']([^\"']+)[\"'][^>]*>(.*?)</a>", text):
            href, label = html.unescape(m.group(1)), strip(m.group(2))
            if rx and not (rx.search(href) or rx.search(label)):
                continue
            say(f"  {urllib.parse.urljoin(final or url, href)}  [{label[:90]}]")
            shown += 1
            if shown >= n:
                break
        say(f"  ({shown} links shown)")
    elif mode == "js":
        text = decode(body, headers)
        rx = extra[0] if extra else r"https?://[^\"'\s]+"
        n = int(extra[1]) if len(extra) > 1 else 30
        srcs = re.findall(r"(?is)<script\b[^>]*src=[\"']([^\"']+)[\"']", text)
        say(f"  {len(srcs)} scripts; inline hits:")
        grep(text, rx, n)
        for src in srcs[:12]:
            full = urllib.parse.urljoin(final or url, html.unescape(src))
            s2, _, h2, b2 = fetch(full)
            say(f"  -- {full} HTTP {s2} {len(b2):,} bytes")
            if s2 == 200:
                grep(decode(b2, h2), rx, n)
    elif mode == "zip":
        try:
            zf = zipfile.ZipFile(io.BytesIO(body))
        except zipfile.BadZipFile:
            say("  not a zip: " + strip(decode(body[:2000], headers))[:300])
            return
        for info in zf.infolist()[:80]:
            say(f"  {info.file_size:>12,}  {info.filename}")
    elif mode == "xls":
        rows = int(extra[0]) if extra else 8
        sheets = int(extra[1]) if len(extra) > 1 else 3
        if body[:2] == b"PK":
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
            for ws in wb.worksheets[:sheets]:
                say(f"  sheet {ws.title!r} {ws.max_row}x{ws.max_column}")
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= rows:
                        break
                    say("    " + show_row(["" if v is None else v for v in row]))
        else:
            import xlrd
            wb = xlrd.open_workbook(file_contents=body)
            for sh in wb.sheets()[:sheets]:
                say(f"  sheet {sh.name!r} {sh.nrows}x{sh.ncols}")
                for r in range(min(rows, sh.nrows)):
                    say("    " + show_row(sh.row_values(r)))
    elif mode == "csv":
        rows = int(extra[0]) if extra else 8
        if body[:2] == b"\x1f\x8b":
            body = gzip.decompress(body)
        reader = csv.reader(io.StringIO(decode(body[:400000], headers)))
        for i, row in enumerate(reader):
            if i >= rows:
                break
            say("    " + show_row(row, 20, 20))
    elif mode == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(body))
        rx = re.compile(extra[0], re.I)
        n = int(extra[1]) if len(extra) > 1 else 40
        say(f"  {len(reader.pages)} pages")
        shown = 0
        for pno, page in enumerate(reader.pages, 1):
            for line in (page.extract_text() or "").splitlines():
                if rx.search(line):
                    say(f"  p{pno}: {line[:200]}")
                    shown += 1
                    if shown >= n:
                        return


def main(argv: list[str]) -> int:
    if not argv or argv[0] in ("-h", "--help"):
        say(__doc__ or "")
        return 0
    for arg in argv:
        try:
            probe(arg)
        except Exception as exc:  # noqa: BLE001 - one failed probe must not stop the rest
            say(f"  !! {type(exc).__name__}: {str(exc)[:300]}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
