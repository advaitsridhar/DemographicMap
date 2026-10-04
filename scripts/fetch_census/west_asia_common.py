"""What the readers for the Arab states and Israel share.

The map's units, under the boundary file's own labels (``common.as_drawn``);
OCHA's administrative gazetteers, which carry the P-codes several offices'
tables are keyed by; sex ratio and grouped median age computed one way; and
the run-stopping checks every reader makes before it writes anything.
"""

from __future__ import annotations

import io
import json
import re
import sys
import unicodedata
from pathlib import Path
from typing import Any

from ._shared import http_get, http_json, log, measure
from .binding import fold
from .cod_ps_age import grouped_median

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from common import as_drawn  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent.parent
SITE = ROOT / "site" / "data"
HDX_API = "https://data.humdata.org/api/3/action"


def units(iso3: str, level: str) -> list[dict[str, Any]]:
    """The map's units at one level, under the boundary file's labels."""
    path = SITE / level / f"{iso3}.units.json"
    return as_drawn(json.loads(path.read_text(encoding="utf-8")))


def parents_by_id(iso3: str) -> dict[str, str]:
    """The map's first-level id -> boundary-file label."""
    return {u["id"]: u["name"] for u in units(iso3, "admin1")}


def key(name: str) -> str:
    """A romanised Arabic place name reduced to what spellings agree on.

    Drops the article in all its assimilated forms (al-, el-, ad-, ash-,
    at-, az-, ...), apostrophes and hamzas, and folds the letters
    transliterations disagree on. Used only to *confirm* a binding made by
    code or to find a candidate a person then declares; never to bind alone.
    """
    text = unicodedata.normalize("NFKD", str(name or "")).encode("ascii", "ignore").decode()
    text = text.lower()
    text = re.sub(r"[`'‘’]", "", text)
    text = re.sub(r"\b(a[lnrtdsz]|as[h]?|ad[h]?|at[h]?|az|el|al)[- ]", " ", text)
    text = re.sub(r"[^a-z]", "", text)
    for a, b in (("kh", "h"), ("dh", "d"), ("th", "t"), ("sh", "s"), ("gh", "g"),
                 ("ou", "u"), ("ee", "i"), ("oo", "u"), ("aa", "a"), ("y", "i"),
                 ("q", "k"), ("w", "u"), ("e", "a"), ("o", "u")):
        text = text.replace(a, b)
    return re.sub(r"(.)\1+", r"\1", text)


def sex_ratio(men: float, women: float, *, year: int, source: str) -> dict[str, Any] | None:
    """Males per 100 females, to one decimal, as the European readers write it."""
    if not women:
        return None
    return measure(round(100 * men / women, 1), unit="males_per_100_females",
                   year=year, source=source)


def median_age(groups: list[tuple[int, int | None, float]], *, year: int,
               source: str) -> dict[str, Any] | None:
    """Median of five-year groups, interpolated in the group holding the middle person."""
    value = grouped_median(groups)
    if value is None:
        return None
    return measure(value, unit="years", year=year, source=source)


def hdx_resource(dataset: str, pattern: str) -> str:
    """The URL of the one resource of an HDX dataset whose name matches."""
    body = http_json(f"{HDX_API}/package_show?id={dataset}", cache=False)
    found = [r for r in body["result"].get("resources") or ()
             if re.search(pattern, str(r.get("name") or ""), re.I)]
    if len(found) != 1:
        raise SystemExit(f"{dataset}: {len(found)} resources match {pattern!r}: "
                         + ", ".join(str(r.get("name")) for r in found))
    return str(found[0]["url"])


def workbook(url: str) -> dict[str, list[list[Any]]]:
    """Every sheet of an .xlsx as rows of values."""
    import openpyxl
    blob = http_get(url, binary=True, cache=False)
    book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    return {name: [list(r) for r in book[name].iter_rows(values_only=True)]
            for name in book.sheetnames}


def gazetteer(sheets: dict[str, list[list[Any]]], level: int) -> list[dict[str, str]]:
    """OCHA's gazetteer rows at one admin level, keyed by lower-case header.

    OCHA's COD-AB workbooks name their sheets ``..._adm2`` or ``Admin2`` and
    their columns ``adm2_name``/``ADM2_EN`` and ``adm2_pcode``/``ADM2_PCODE``,
    depending on the vintage; both are read.
    """
    want = re.compile(rf"adm(?:in)?_?{level}(?!\d)", re.I)
    for name, rows in sheets.items():
        if not want.search(name.replace(" ", "")):
            continue
        header = [str(c or "").strip().lower() for c in rows[0]]
        out = []
        for row in rows[1:]:
            rec = {h: ("" if v is None else str(v).strip()) for h, v in zip(header, row) if h}
            if any(rec.values()):
                out.append(rec)
        return out
    raise SystemExit(f"no admin-{level} sheet among {sorted(sheets)}")


def field(row: dict[str, str], level: int, what: str) -> str:
    """A gazetteer row's admin-``level`` name or P-code, whatever the vintage calls it."""
    names = {"name": (f"adm{level}_name", f"adm{level}_en", f"admin{level}name_en",
                      f"adm{level}_name_en"),
             "pcode": (f"adm{level}_pcode", f"admin{level}pcode")}[what]
    for column in names:
        if row.get(column):
            return row[column]
    return ""


def bind_by_gazetteer(drawn: list[dict[str, Any]], parents: dict[str, str],
                      rows: list[dict[str, str]], level: int,
                      parent_alias: dict[str, str] | None = None,
                      unit_alias: dict[tuple[str, str], str] | None = None,
                      ) -> tuple[dict[str, str], list[str], list[str]]:
    """Each drawn unit's P-code, by its label among OCHA's units of its parent.

    The boundary file and OCHA's gazetteer spell a unit the same way when the
    one was cut from the other, which is the case this is for: a drawn label
    is bound only to the one gazetteer row of the same folded name under the
    same parent, and every unit left over on either side is returned.
    ``parent_alias`` maps the map's first-level label to OCHA's;
    ``unit_alias`` maps (OCHA parent, map label) to OCHA's label, declared.
    """
    parent_alias = parent_alias or {}
    unit_alias = unit_alias or {}
    index: dict[tuple[str, str], list[dict[str, str]]] = {}
    for row in rows:
        above = fold(field(row, level - 1, "name")) if level > 1 else ""
        index.setdefault((above, fold(field(row, level, "name"))), []).append(row)
    bound: dict[str, str] = {}
    left: list[str] = []
    for unit in drawn:
        parent = parents.get(unit["parent"], unit["parent"]) if level > 1 else ""
        parent = parent_alias.get(parent, parent)
        label = unit_alias.get((parent, unit["name"]), unit["name"])
        hits = index.get((fold(parent), fold(label)), [])
        if len(hits) == 1:
            bound[unit["id"]] = field(hits[0], level, "pcode")
        else:
            left.append(f"{unit['name']} ({parent}): {len(hits)} gazetteer rows")
    used = set(bound.values())
    spare = [f"{field(r, level, 'name')} ({field(r, level - 1, 'name')}, "
             f"{field(r, level, 'pcode')})" for r in rows
             if field(r, level, "pcode") not in used]
    if len(used) != len(bound):
        raise SystemExit("two drawn units bound to one gazetteer row")
    return bound, left, spare


AGE_CLOSED = re.compile(r"^([BMF])(\d{2})(\d{2})(_\d)?$")
AGE_OPEN = re.compile(r"^([BMF])(\d{2})PL(_\d)?$")


def uscb_table(sheets: dict[str, list[list[Any]]], name: str) -> tuple[list[dict[str, Any]],
                                                                        dict[str, str]]:
    """One sheet of a Census Bureau workbook as rows keyed by field name.

    The sheets carry two header rows -- field names, then the aliases a
    person reads -- and the same geography columns first; a sheet is found
    by its stripped, lower-cased name. Returns the rows and the
    name -> alias map.
    """
    from . import uscb
    wanted = name.strip().lower()
    actual = next((s for s in sheets if s.strip().lower() == wanted), None)
    if actual is None:
        raise SystemExit(f"no sheet {name!r}: {sorted(sheets)}")
    rows = sheets[actual]
    names, aliases = uscb.columns(rows)
    out = []
    for row in rows[2:]:
        rec = {n: v for n, v in zip(names, row) if n}
        if not any(v not in (None, "") for v in rec.values()):
            continue
        level = uscb.number(rec.get("ADM_LEVEL"))
        rec["_level"] = int(level) if level is not None else None
        out.append(rec)
    return out, {n: a for n, a in zip(names, aliases) if n}


def count(value: Any) -> float | None:
    """A cell as a count of people; the Bureau's negative sentinels are None."""
    from . import uscb
    return uscb.number(value)


def age_groups(row: dict[str, Any], sex: str = "B", suffix: str = "") -> list[tuple[int, int | None, float]]:
    """(from, to, people) for a Census Bureau row's five-year groups of one sex.

    ``sex`` is B (both), M or F; ``suffix`` selects one year's columns where
    a sheet carries several ("_7" for 2017). The groups must run from 0
    without a gap to an open-ended last group, or the run stops.
    """
    groups: list[tuple[int, int | None, float]] = []
    for name, value in row.items():
        if not isinstance(name, str):
            continue
        closed, opened = AGE_CLOSED.match(name), AGE_OPEN.match(name)
        m = closed or opened
        if not m or m.group(1) != sex:
            continue
        if ((closed.group(4) if closed else opened.group(3)) or "") != suffix:
            continue
        people = count(value)
        if people is None:
            raise SystemExit(f"{row.get('AREA_NAME')}: no count in {name}")
        if closed:
            groups.append((int(closed.group(2)), int(closed.group(3)), people))
        else:
            groups.append((int(opened.group(2)), None, people))
    groups.sort(key=lambda g: g[0])
    edge = 0
    for low, high, _n in groups:
        if low != edge:
            raise SystemExit(f"{row.get('AREA_NAME')}: age groups break at {edge}")
        edge = (high + 1) if high is not None else -1
    if not groups or groups[-1][1] is not None:
        raise SystemExit(f"{row.get('AREA_NAME')}: no open-ended last age group")
    return groups


def check(condition: bool, message: str) -> None:
    """A failed measurement stops the run; nothing is written."""
    if not condition:
        raise SystemExit(message)


def report(label: str, items: list[str], limit: int = 40) -> None:
    log(f"  {label}: {len(items)}" + (": " + "; ".join(items[:limit]) if items else ""))
