#!/usr/bin/env python3
"""The 2020 Russian census, Volume 5: ethnic composition and native language.

Two tables, both published per federal subject, both naming their own universe
in the sheet:

    Tom5_tab1  1. НАЦИОНАЛЬНЫЙ СОСТАВ НАСЕЛЕНИЯ      -> ethnicity
    Tom5_tab6  6. НАСЕЛЕНИЕ ПО РОДНОМУ ЯЗЫКУ         -> language

Not table 5. It is called ВЛАДЕНИЕ ЯЗЫКАМИ -- *proficiency*, which languages a
person knows -- and a person may know several, so its columns are independent
indicators rather than parts of one whole. Published as a composition it would
read past 100%, which is what Thailand's language file was refused for. The
native-language table is the one that partitions a population, and table 6 is
it.

**Fetched from the Internet Archive, and the citation says so.** rosstat.gov.ru
serves a valid certificate signed by the Russian Trusted Sub CA, an authority
operated by the Ministry of Digital Development that no ordinary trust store
carries -- Russian government sites moved to it after 2022. The site is not
refusing this client and this project will not disable verification to reach
it, so the files come from a public archive of the same URLs. That makes the
capture date part of the provenance: the map cites Rosstat's publication as
retrieved by the Internet Archive on the date below, not as fetched from
Rosstat today.

**Matched on ISO 3166-2, not on names.** The sheets are Cyrillic and the
boundary file is English, and norm() deliberately does not transliterate --
inventing a romanisation here would be this code guessing at a name rather
than reading one. Every subject therefore carries its ISO 3166-2 code, which
either matches a shape's code or does not, with no fuzzy step in between where
Chechnya can quietly wear Ingushetia's figures. One shape, Sakha, carries no
code in the boundary file and is aliased by name instead.
"""

from __future__ import annotations

import argparse
import io
import re
import time
from collections import Counter
from typing import Any

from ._shared import (
    NOT_AVAILABLE, PROCESSED, RAW, gap, http_get, log, record, shares,
    write_json,
)

# Where the two workbooks live once fetched. They are checked in, which is
# this project's existing answer for a source with no reliably fetchable URL --
# Nepal's report, Sri Lanka's tables and India's C-16 workbooks are all read
# from here. Russia qualifies twice over: its own host serves a certificate no
# ordinary trust store carries, and the archive standing in for it throttles
# hard enough to answer with an HTML page. A quarterly refresh should not
# depend on an archive being in a good mood, and 2.6 MB is a small price for
# a build that does the same thing every time.
STORE = RAW / "russia"

CAPTURE = "20250105165828"
# The "id_" suffix is the Wayback Machine's raw-content form, and it is not
# optional here. Without it the archive answers a request for a .xlsx with its
# playback wrapper -- 10 kB of "<!DOCTYPE html><title>Wayback Machine</title>"
# -- which is a valid page and a useless workbook. That was read here as rate
# limiting and written up as such four times before anyone looked at the body.
# It was never rate limiting. id_ returns 1.5 MB beginning "PK", first ask.
BASE = (f"https://web.archive.org/web/{CAPTURE}id_/"
        "https://rosstat.gov.ru/storage/mediabank/")
LANDING = ("https://rosstat.gov.ru/vpn/2020/"
           "Tom5_Nacionalnyj_sostav_i_vladenie_yazykami")

YEAR = 2021
SOURCE = ("Федеральная служба государственной статистики (Rosstat), "
          "Всероссийская перепись населения 2020 года, Том 5: Национальный "
          "состав и владение языками, retrieved via the Internet Archive "
          f"capture of {CAPTURE[:4]}-{CAPTURE[4:6]}-{CAPTURE[6:8]}")
LICENCE = "Rosstat open data"

# The sheet that names the whole country. Read as the control every other
# adapter here keeps, and never emitted: it is not a shape on this map.
COUNTRY_SHEET = "Российская Федерация"

# Sheets that exist and are deliberately not read, each for its own reason.
# Two are the *combined* forms of a subject the boundary file draws split, so
# reading them alongside their parts would count about four million people
# twice; two are territory this map already draws from Ukraine's own census.
SKIP = {
    "Архангельская область":
        "the form including Nenets AO, which is drawn separately",
    "Тюменская область":
        "the form including Khanty-Mansi and Yamalo-Nenets, drawn separately",
    "Республика Крым":
        "claimed by this file and by Ukraine's 2001 census, which this map "
        "already publishes on the same shapes",
    "г. Севастополь":
        "as Crimea, and geoBoundaries draws neither inside Russia",
}

# Sheet name -> ISO 3166-2. Written out rather than derived: a rule mapping
# Cyrillic to these codes would be a transliteration with extra steps, and
# every entry here is checkable against one published standard.
SUBJECTS = {
    "Белгородская область": "RU-BEL",
    "Брянская область": "RU-BRY",
    "Владимирская область": "RU-VLA",
    "Воронежская область": "RU-VOR",
    "Ивановская область": "RU-IVA",
    "Калужская область": "RU-KLU",
    "Костромская область": "RU-KOS",
    "Курская область": "RU-KRS",
    "Липецкая область": "RU-LIP",
    "Московская область": "RU-MOS",
    "Орловская область": "RU-ORL",
    "Рязанская область": "RU-RYA",
    "Смоленская область": "RU-SMO",
    "Тамбовская область": "RU-TAM",
    "Тверская область": "RU-TVE",
    "Тульская область": "RU-TUL",
    "Ярославская область": "RU-YAR",
    "г. Москва": "RU-MOW",
    "Республика Карелия": "RU-KR",
    "Республика Коми": "RU-KO",
    "Архангельская область без АО": "RU-ARK",
    "Ненецкий автономный округ": "RU-NEN",
    "Вологодская область": "RU-VLG",
    "Калининградская область": "RU-KGD",
    "Ленинградская область": "RU-LEN",
    "Мурманская область": "RU-MUR",
    "Новгородская область": "RU-NGR",
    "Псковская область": "RU-PSK",
    "г. Санкт-Петербург": "RU-SPE",
    "Республика Адыгея": "RU-AD",
    "Республика Калмыкия": "RU-KL",
    "Краснодарский край": "RU-KDA",
    "Астраханская область": "RU-AST",
    "Волгоградская область": "RU-VGG",
    "Ростовская область": "RU-ROS",
    "Республика Дагестан": "RU-DA",
    "Республика Ингушетия": "RU-IN",
    "Кабардино-Балкарская Республика": "RU-KB",
    "Карачаево-Черкесская Республика": "RU-KC",
    "РСО-Алания": "RU-SE",
    "Чеченская Республика": "RU-CE",
    "Ставропольский край": "RU-STA",
    "Республика Башкортостан": "RU-BA",
    "Республика Марий Эл": "RU-ME",
    "Республика Мордовия": "RU-MO",
    "Республика Татарстан": "RU-TA",
    "Удмуртская Республика": "RU-UD",
    "Чувашская Республика": "RU-CU",
    "Пермский край": "RU-PER",
    "Кировская область": "RU-KIR",
    "Нижегородская область": "RU-NIZ",
    "Оренбургская область": "RU-ORE",
    "Пензенская область": "RU-PNZ",
    "Самарская область": "RU-SAM",
    "Саратовская область": "RU-SAR",
    "Ульяновская область": "RU-ULY",
    "Курганская область": "RU-KGN",
    "Свердловская область": "RU-SVE",
    "Тюменская область без АО": "RU-TYU",
    "ХМАО": "RU-KHM",
    "ЯНАО": "RU-YAN",
    "Челябинская область": "RU-CHE",
    "Республика Алтай": "RU-AL",
    "Республика Тыва": "RU-TY",
    "Республика Хакасия": "RU-KK",
    "Алтайский край": "RU-ALT",
    "Красноярский край": "RU-KYA",
    "Иркутская область": "RU-IRK",
    "Кемеровская область - Кузбасс": "RU-KEM",
    "Новосибирская область": "RU-NVS",
    "Омская область": "RU-OMS",
    "Томская область": "RU-TOM",
    "Республика Бурятия": "RU-BU",
    "Республика Саха (Якутия)": "RU-SA",
    "Забайкальский край": "RU-ZAB",
    "Камчатский край": "RU-KAM",
    "Приморский край": "RU-PRI",
    "Хабаровский край": "RU-KHA",
    "Амурская область": "RU-AMU",
    "Магаданская область": "RU-MAG",
    "Сахалинская область": "RU-SAK",
    "Еврейская автономная область": "RU-YEV",
    "Чукотский автономный округ": "RU-CHU",
}

# The two tables do not spell every subject the same way. Table 1 writes ХМАО
# and ЯНАО; table 6 writes them out in full. Read under their own names those
# two sheets are simply unknown, and the two regions lose their language field
# without anything failing -- which is how this was nearly shipped.
ALSO_KNOWN_AS = {
    "Ханты-Мансийский АО - Югра": "ХМАО",
    "Ямало-Ненецкий АО": "ЯНАО",
}

# geoBoundaries carries an iso_3166_2 on 82 of Russia's 83 first-order shapes.
# Sakha is the exception, so it is the one subject matched by name.
BY_NAME = {"RU-SA": ("Sakha Republic",)}

# The row whose figure is everyone the question reached. Both tables name it,
# which is why neither needs a denominator inferred.
UNIVERSE = {
    "ethnicity": "Указавшие национальную принадлежность",
    "language": "Указавшие родной язык",
}

TABLE = {"ethnicity": "Tom5_tab1_VPN-2020.xlsx",
         "language": "Tom5_tab6_VPN-2020.xlsx"}

# How far the groups may sit from the universe the sheet publishes. The
# ethnicity table nests -- Авары is followed by Андийцы, Ахвахцы and Дидойцы,
# which are counted inside it -- so a reader that takes every row sums past
# the total. This is the guard that says so rather than publishing it.
TOLERANCE = 0.005


def workbook(filename: str) -> bytes:
    """The file, or a refusal that says what arrived instead.

    The check earns its place even now that BASE asks for raw content: the
    archive answers 200 either way, and an HTML page where a workbook was
    expected is a successful request for the wrong thing. http_get cannot see
    that -- from its side nothing failed -- so it belongs here, where something
    knows what was asked for. A ZIP starts "PK".

    What this guard could not do was explain itself. It refused four runs
    saying the archive was throttling, which was a guess made from a byte count
    and never checked against the body. The body said "Wayback Machine". The
    lesson kept here is that a guard which names the wrong cause is still
    better than none -- it stopped the bad file -- but that reading the
    evidence would have cost one probe instead of four runs.

    Uncached deliberately. http_get caches whatever it received, and a cached
    error page would fail every later run for a reason that had already gone
    away. Two files of about 2.6 MB once per refresh is the cheaper mistake.

    retries=1, arrived at by overshooting in both directions. http_get
    defaults to four, and four of those inside four of these multiply rather
    than add: a throttled fetch went twenty slow requests deep and ran
    seventeen minutes before it was killed. Zero then swung too far the other
    way, because a timeout is exactly the transient this layer should ride out
    and it aborted on the first one. One retry inside, four outside, and the
    outer loop owns the policy because it is the layer that can tell an error
    page from a dropped connection.
    """
    kept = STORE / filename
    if kept.exists():
        blob = kept.read_bytes()
        if blob[:2] == b"PK":
            log(f"  {filename}: {len(blob):,} bytes from {kept.parent}")
            return blob
        log(f"  {filename}: the checked-in copy is not a workbook; refetching")

    last = b""
    for attempt in range(4):
        blob = http_get(BASE + filename, binary=True, cache=False,
                        retries=1, timeout=90)
        if blob[:2] == b"PK":
            STORE.mkdir(parents=True, exist_ok=True)
            kept.write_bytes(blob)
            log(f"  {filename}: {len(blob):,} bytes, kept at {kept}")
            return blob
        last = blob
        wait = 15 * (attempt + 1)
        log(f"  {filename}: {len(blob):,} bytes beginning {blob[:12]!r}, "
            f"which is not a workbook; retrying in {wait}s")
        time.sleep(wait)
    raise SystemExit(
        f"RUS: {filename} came back as {len(last):,} bytes of {last[:60]!r} "
        f"four times over, and no copy is checked in at {kept}. Read those "
        f"bytes before concluding anything about why: the last time this "
        f"fired they were the archive's playback wrapper, and the URL wanted "
        f"the id_ suffix rather than the four rounds of backoff it got.")


def number(cell: Any) -> float | None:
    """A count, or None. Rosstat writes an em dash where a group is absent."""
    if cell is None:
        return None
    if isinstance(cell, bool):
        return None
    if isinstance(cell, (int, float)):
        return float(cell)
    text = str(cell).strip().replace(" ", "").replace(" ", "")
    if not text or text in {"-", "–", "—", "..."}:
        return None
    text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def label(cell: Any) -> str:
    """The group's published name.

    The census writes a group's self-designations in brackets after its name --
    "Башкиры (башкирцы, башкорт, мин, ...)". The name is the part before them;
    the list is a glossary, not part of what to print.
    """
    text = ("" if cell is None else str(cell)).strip()
    return re.split(r"\s*\(", text, 1)[0].strip()


def indent(cell: Any) -> int:
    """How far the sheet indents this label."""
    alignment = getattr(cell, "alignment", None)
    return int((alignment.indent or 0) if alignment else 0)


def parents(rows: list[tuple[str, float, int]]) -> list[tuple[str, float]]:
    """The groups that are not counted inside another group.

    Rosstat lists constituent peoples inside the group they belong to --
    Аварцы then Андийцы, Ахвахцы, Дидойцы; Татары then Кряшены -- and adding
    those to their parent counts the same people twice.

    Which rows are which is in the indent, and it took three tries to read it.
    Leading whitespace was wrong: it missed every real sub-group and fired on
    Русские in Воронежская область, dropping the largest group in a region of
    2.2 million. "Indented at all" was wrong the other way: Rosstat indents the
    whole list one level under "в том числе:", so in Чукотский автономный округ
    that read 94 of 95 rows as nested and kept 446 people out of 47,044.

    The level that means "a group" is therefore the commonest one, not zero and
    not the smallest -- a sheet has one row per ethnicity and only a handful of
    sub-groups, so the mode is the top of the list by construction.

    Rows at exactly that level, and no others. Deeper is inside something.
    Shallower is not a member of the list at all, and the row that proves it is
    "Лица, в переписных листах которых национальная принадлежность не указана"
    -- the people who stated nothing. Counting them took ХМАО to 135% of its
    own total, because the universe this reconciles against is Указавшие
    национальную принадлежность, those who *did* state. A residual outside the
    denominator cannot be one of its parts.
    """
    if not rows:
        return []
    depths = Counter(depth for _name, _value, depth in rows)
    top = depths.most_common(1)[0][0]
    return [(name, value) for name, value, depth in rows if depth == top]


def layout(sheet: Any, wanted: str) -> tuple[int, int] | None:
    """Which column holds the label and which the count, for this sheet.

    Not columns 0 and 1. Two sheets of the 88 -- Ингушетия and Красноярский
    край -- were saved with the pivot table's member keys still in the first
    column:

        [P16_Nationality].[Hierarchy].[Code02].&[3] | Аварцы | 63

    Reading position 0 there takes the key as the label and the label as the
    count, so the universe row is never recognised, no groups are collected,
    and the subject comes out empty. Empty is the dangerous outcome: it reads
    on the map as a census that did not ask, and it is the reader that did not
    look.

    So the columns are found rather than assumed, by locating the row that
    names the universe and taking the first numeric column to its right.
    """
    for row in sheet.iter_rows(max_row=40):
        for i, cell in enumerate(row):
            if label(cell.value).startswith(wanted):
                for j in range(i + 1, len(row)):
                    if number(row[j].value) is not None:
                        return i, j
    return None


def read(blob: bytes, field: str) -> dict[str, dict[str, Any]]:
    """{sheet name: {published, counts}} for every sheet in one table."""
    import openpyxl

    # Not read_only: the indent that says which rows are nested is a style,
    # and a read-only workbook does not carry styles. Two files of a few
    # megabytes are worth the memory to read them correctly.
    book = openpyxl.load_workbook(io.BytesIO(blob), data_only=True)
    wanted = UNIVERSE[field]
    out: dict[str, dict[str, Any]] = {}
    for sheet in book.sheetnames:
        key = ALSO_KNOWN_AS.get(sheet.strip(), sheet.strip())
        if key != COUNTRY_SHEET and key not in SUBJECTS and key not in SKIP:
            log(f"  unknown sheet, not read: {sheet!r}")
            continue
        where = layout(book[sheet], wanted)
        if where is None:
            raise SystemExit(
                f"RUS {field}: {sheet} has no row naming {wanted!r}. Every "
                f"other sheet in this file does, so this is a layout this "
                f"reader has not seen rather than a question the region was "
                f"not asked.")
        at_name, at_value = where
        published: float | None = None
        seen: list[tuple[str, float, int]] = []
        for row in book[sheet].iter_rows():
            if len(row) <= at_value:
                continue
            name = label(row[at_name].value)
            if not name:
                continue
            if published is None:
                if name.startswith(wanted):
                    published = number(row[at_value].value)
                continue
            if name.startswith("в том числе"):
                continue
            value = number(row[at_value].value)
            if value is None:
                continue
            seen.append((name, value, indent(row[at_name])))
        counts: dict[str, float] = {}
        for name, value in parents(seen):
            counts[name] = counts.get(name, 0.0) + value
        out[key] = {"published": published, "counts": counts,
                    "nested": len(seen) - len(parents(seen))}
    book.close()
    return out


def check(field: str, tables: dict[str, dict[str, Any]]) -> None:
    """Refuse a table this reader has misunderstood, and say which way."""
    worst = None
    for sheet, got in tables.items():
        published, counts = got["published"], got["counts"]
        if not published or not counts:
            continue
        ratio = sum(counts.values()) / published
        if worst is None or abs(ratio - 1) > abs(worst[1] - 1):
            worst = (sheet, ratio)
    if worst is None:
        raise SystemExit(f"RUS {field}: no sheet published a universe")
    sheet, ratio = worst
    off = abs(ratio - 1)
    got = tables[sheet]
    # The largest groups, because the first time this fired the whole story was
    # that Русские was missing from them and the ratio alone did not say so. A
    # sum that is wrong is usually wrong about something nameable.
    top = ", ".join(f"{name} {value:,.0f}" for name, value
                    in sorted(got["counts"].items(), key=lambda kv: -kv[1])[:3])
    log(f"  {field}: furthest from its own total is {sheet} at "
        f"{ratio:.4f} of {UNIVERSE[field]!r} ({got['published']:,.0f}); "
        f"largest {top}; {got['nested']} nested rows left to their parents")
    if off > TOLERANCE:
        raise SystemExit(
            f"RUS {field}: {sheet} sums to {ratio:.4f} of the "
            f"{got['published']:,.0f} the sheet publishes, past the "
            f"{TOLERANCE:.1%} this reader allows. Its largest groups came out "
            f"as {top} -- check whether one that belongs there is missing "
            f"before assuming the nesting is wrong, because that is what it "
            f"was last time.")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", default="russia_subject.json")
    args = ap.parse_args()

    fields: dict[str, dict[str, dict[str, Any]]] = {}
    for field, filename in TABLE.items():
        if fields:
            # A courtesy pause between two multi-megabyte requests to a free
            # public archive. Not a workaround for anything: the wrapper this
            # once seemed to be evidence of was a URL mistake, not a limit.
            time.sleep(5)
        blob = workbook(filename)
        log(f"{field}: {filename}, {len(blob):,} bytes")
        fields[field] = read(blob, field)
        check(field, fields[field])

    records = []
    for sheet, code in SUBJECTS.items():
        values: dict[str, Any] = {}
        for field in TABLE:
            got = fields[field].get(sheet)
            if not got or not got["counts"]:
                values[field] = gap(NOT_AVAILABLE)
                continue
            values[field] = shares(
                got["counts"], total=got["published"] or None)
        records.append(record(
            code, sheet, level="admin1", parent="RUS",
            country="RUS", iso_3166_2=code,
            aliases=list(BY_NAME.get(code, ())),
            **values,
            sources=[{"field": field, "name": SOURCE, "url": LANDING,
                      "license": LICENCE, "year": YEAR} for field in TABLE],
        ))

    # Every table, not just the first. Checking one of them let ХМАО and ЯНАО
    # through with ethnicity and no language, because table 6 spells their
    # names out where table 1 abbreviates: a subject present in one file and
    # absent from another is exactly the gap that looks like a country simply
    # not answering that question.
    for field in TABLE:
        # Present *and* read. A sheet that exists and yields nothing is the
        # same gap as a sheet that is absent, and it looks identical on the
        # map to a census that did not ask -- which is how Ингушетия and
        # Красноярский край nearly shipped empty.
        missing = [s for s in SUBJECTS
                   if not (fields[field].get(s) or {}).get("counts")]
        if missing:
            raise SystemExit(
                f"RUS {field}: {len(missing)} of {len(SUBJECTS)} configured "
                f"subjects came out empty from {TABLE[field]}: "
                + ", ".join(missing[:6])
                + ". A subject this file does not carry is a gap for that "
                  "field alone, which reads on the map as a question the "
                  "region was not asked.")
    path = PROCESSED / args.out
    write_json(path, records)
    log(f"  wrote {path} ({path.stat().st_size // 1024} kB)")
    log(f"  {len(records)} subjects, {len(SKIP)} sheets skipped by name")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
