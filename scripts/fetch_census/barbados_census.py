#!/usr/bin/env python3
"""Barbados: the 2021 census's tables by parish, and the 2010 census's parish populations.

The Barbados Statistical Service publishes the 2021 Population and Housing
Census's detailed tables as one workbook, *Census-2020-Tables.xlsx* (the round
is named for 2020; the census night was 1 August 2021):

* Table 01.02, population by parish, five-year age group and sex;
* Table 02.04, population by parish, sex and ethnic origin;
* Table 02.06, population by parish, sex and religion.

Their universe is what the BSS calls the tabulated resident population:
136,415 people, the residents whose records the census completed -- about half
the island, which the office's own estimate puts at some 269,000. So these are
the census's compositions of the people it tabulated, and every note says so.
The median age is interpolated within the five-year group holding the middle
person; the sex ratio is the table's men per thousand women.

No 2021 population is written: the tabulated counts are not the parishes'
populations. The population is the 2010 census's *estimated resident
population* by parish (Table C of the 2010 workbook, *Census-Tables-2010.xlsx*),
which includes the institutional population and the 18% undercount the BSS
estimated -- the parishes' last official count.

Checks, each of which stops the run: every parish's age groups, ethnic groups
and religions make its tabulated total, those totals agree across the three
tables, and the parishes make the island's row in each; the 2010 parishes make
the 2010 island total of 277,821.

Usage:
    python -m scripts.fetch_census.barbados_census
"""

from __future__ import annotations

import argparse
import io
import json
import re
from typing import Any

from ._shared import PROCESSED, http_get, log, measure, record, shares, write_json
from .binding import fold
from .cod_ps_age import grouped_median

OUT = PROCESSED / "barbados_census.json"
SITE = PROCESSED.parent.parent / "site" / "data"
URL_2021 = "https://stats.gov.bb/wp-content/uploads/2024/05/Census-2020-Tables.xlsx"
URL_2010 = "https://stats.gov.bb/wp-content/uploads/2021/03/Census-Tables-2010.xlsx"
PAGE = "https://stats.gov.bb/census/"
SOURCE_2021 = ("Barbados Statistical Service, 2021 Population and Housing Census, Tables "
               "01.02, 02.04 and 02.06 (tabulated resident population)")
SOURCE_2010 = ("Barbados Statistical Service, 2010 Population and Housing Census, Table C: "
               "Estimated Resident Population by Parish and Sex")
TABULATED = 136_415
RESIDENT_2010 = 277_821
PARISHES = 11

ETHNICITY = {
    "Black": "Black", "White": "White", "Asian": "Asian", "East Indian": "East Indian",
    "Middle Eastern": "Middle Eastern", "Mixed": "Mixed", "Other (Please Specify)": "Other",
    "Not Stated": "Not stated",
}
RELIGION = {
    "Adventist": "Seventh-day Adventist", "Anglican": "Anglican", "Baptist": "Baptist",
    "Church of God": "Church of God", "Methodist": "Methodist", "Moravian": "Moravian",
    "Nazarene": "Church of the Nazarene", "Other Pentecostal": "Pentecostal",
    "Roman Catholic": "Roman Catholic", "Salvation Army": "Salvation Army",
    "Wesleyan": "Wesleyan Holiness Church", "Brethren": "Brethren",
    "Jehovah's Witness": "Jehovah's Witnesses", "Mormon": "Mormon",
    "Other Christian": "Other Christian", "Baha'i": "Baha'i", "Hindu": "Hindu",
    "Islam": "Muslim", "Judaism": "Jewish", "Rastafarian": "Rastafarian",
    "Other Non-Christian": "Other religion", "No Religious Affiliation": "No religion",
    "Not Stated": "Not stated",
}
AGE = re.compile(r"^(\d+)\s*-\s*(\d+)$")
OPEN = re.compile(r"^(\d+)\s+and\s+over$", re.I)



def present(counts: dict[str, float]) -> list[dict]:
    """Shares of the categories anyone is counted in; a zero is no one, not a group."""
    return shares({k: v for k, v in counts.items() if v})

def text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("’", "'")).strip()


def count(value: Any, where: str) -> int:
    if value in (None, ""):
        return 0
    try:
        return int(round(float(value)))
    except (TypeError, ValueError):
        raise SystemExit(f"barbados_census: {where}: {value!r} is not a count")


def parish_key(name: str) -> str:
    """The BSS's "St. Michael" and the map's "Saint Michael" as one key."""
    return fold(re.sub(r"^St\.?\s*", "Saint ", text(name)))


def header(rows: list[list[Any]], first: str) -> tuple[int, list[str]]:
    """The row holding the category names (it begins, after two blanks, with "Total")."""
    for i, row in enumerate(rows):
        cells = [text(c) for c in row]
        if len(cells) > 2 and cells[2] == "Total" and first in cells:
            return i, cells
    raise SystemExit(f"barbados_census: no header row naming {first!r}")


def compositions(rows: list[list[Any]], labels: dict[str, str], rowtag: str,
                 what: str) -> dict[str, dict[str, Any]]:
    """{parish key (or "barbados"): {"name", "total", "counts"}} from a parish table.

    A parish's line is the one whose second cell is ``rowtag`` ("Total" for the
    ethnicity table, whose other lines are age groups; "Both Sexes" for the
    religion table, whose others are the sexes). A category not in ``labels``
    stops the run.
    """
    at, head = header(rows, next(iter(labels)))
    columns = {}
    for j, name in enumerate(head[3:], start=3):
        if not name:
            continue
        if name not in labels:
            raise SystemExit(f"barbados_census: {what} category {name!r} is not one this reads")
        columns[j] = labels[name]
    out, current = {}, None
    for row in rows[at + 1:]:
        cells = [text(c) for c in row] + [""] * 3
        if cells[0]:
            current = cells[0]
        if current is None or cells[1] != rowtag:
            continue
        where = f"{what}, {current}"
        total = count(row[2], where)
        counts: dict[str, int] = {}
        for j, label in columns.items():
            counts[label] = counts.get(label, 0) + count(row[j] if j < len(row) else None, where)
        if sum(counts.values()) != total:
            raise SystemExit(f"barbados_census: {where}: categories make "
                             f"{sum(counts.values()):,} of {total:,}")
        out[parish_key(current)] = {"name": current, "total": total, "counts": counts}
    return out


def ages(rows: list[list[Any]]) -> dict[str, dict[str, Any]]:
    """{parish key: {"total", "men", "women", "groups"}} from Table 01.02."""
    out, current = {}, None
    for row in rows:
        cells = [text(c) for c in row] + [""] * 5
        if cells[0] and cells[1] == "Total":
            current = parish_key(cells[0])
            out[current] = {"name": cells[0], "total": count(row[2], cells[0]),
                            "men": count(row[3], cells[0]), "women": count(row[4], cells[0]),
                            "groups": []}
            continue
        if current is None or not cells[1]:
            continue
        band = AGE.match(cells[1]) or OPEN.match(cells[1])
        if not band:
            continue
        low = int(band.group(1))
        high = int(band.group(2)) if band.re is AGE else None
        out[current]["groups"].append((low, high, count(row[2], cells[0] or current)))
    for key, unit in out.items():
        made = sum(g[2] for g in unit["groups"])
        if made != unit["total"] or unit["men"] + unit["women"] != unit["total"]:
            raise SystemExit(f"barbados_census: {unit['name']}'s age groups make {made:,} and "
                             f"its sexes {unit['men'] + unit['women']:,} of {unit['total']:,}")
    return out


def resident_2010(rows: list[list[Any]]) -> dict[str, int]:
    """{parish key: 2010 estimated resident population} from Table C."""
    out = {}
    for row in rows:
        cells = [text(c) for c in row] + [""] * 2
        if cells[0] and re.match(r"^\d", cells[1] or "x"):
            out[parish_key(cells[0])] = count(row[1], cells[0])
    island = out.pop("barbados", None)
    if island != RESIDENT_2010 or sum(out.values()) != RESIDENT_2010 or len(out) != PARISHES:
        raise SystemExit(f"barbados_census: Table C's {len(out)} parishes make "
                         f"{sum(out.values()):,}; the island row is {island}")
    return out


def check(tables: dict[str, dict[str, dict[str, Any]]]) -> None:
    """Each table's parishes make its island row, and the three agree parish by parish."""
    for what, table in tables.items():
        island = table.get("barbados")
        parts = {k: v for k, v in table.items() if k != "barbados"}
        if island is None or island["total"] != TABULATED or len(parts) != PARISHES:
            raise SystemExit(f"barbados_census: the {what} table's island row is "
                             f"{island and island['total']} with {len(parts)} parishes")
        if sum(v["total"] for v in parts.values()) != TABULATED:
            raise SystemExit(f"barbados_census: the {what} table's parishes do not make "
                             f"{TABULATED:,}")
    first, *others = tables.values()
    for key, unit in first.items():
        for other in others:
            if other.get(key, {}).get("total") != unit["total"]:
                raise SystemExit(f"barbados_census: {unit['name']}'s tabulated total differs "
                                 "between tables")


def main() -> int:
    argparse.ArgumentParser(description=__doc__).parse_args()
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(http_get(URL_2021, binary=True, cache=False)),
                                  read_only=True, data_only=True)

    def sheet(name: str) -> list[list[Any]]:
        return [list(r) for r in book[name].iter_rows(values_only=True)]

    age = ages(sheet("01.02"))
    ethnicity = compositions(sheet("02.04"), ETHNICITY, "Total", "ethnicity")
    religion = compositions(sheet("02.06"), RELIGION, "Both Sexes", "religion")
    check({"age": age, "ethnicity": ethnicity, "religion": religion})
    old = openpyxl.load_workbook(io.BytesIO(http_get(URL_2010, binary=True, cache=False)),
                                 read_only=True, data_only=True)
    resident = resident_2010([list(r) for r in old["Table C"].iter_rows(values_only=True)])
    log(f"  {PARISHES} parishes making the {TABULATED:,} tabulated in 2021 in each of three "
        f"tables, and the {RESIDENT_2010:,} estimated resident in 2010")

    admin1 = json.loads((SITE / "admin1" / "BRB.units.json").read_text())
    shapes = {fold(u["name"]): u for u in admin1}
    universe = (f"the 2021 census's tabulated resident population ({TABULATED:,} people, the "
                "residents whose census records were completed, about half the island)")
    records = []
    for key in sorted(k for k in age if k != "barbados"):
        shape = shapes.get(key)
        if shape is None:
            raise SystemExit(f"barbados_census: parish {age[key]['name']!r} has no polygon")
        a = age[key]
        records.append(record(
            f"BRB-BSS-{key}", shape["name"], level="admin1", parent="BRB", country="BRB",
            match_by="shape_id", shape_id=shape["id"],
            population=measure(resident[key], year=2010, source=SOURCE_2010),
            population_note=("The 2010 census's estimated resident population of the parish, "
                             "the institutional population and the estimated undercount "
                             "included. The 2021 census's parish tables count only the people "
                             "it tabulated and are not the parish's population."),
            median_age=measure(grouped_median(sorted(a["groups"], key=lambda g: g[0])),
                               unit="years", year=2021, source=SOURCE_2021),
            median_age_note=("Interpolated within the five-year age group holding the middle "
                             f"person, of {universe} (Table 01.02)."),
            sex_ratio=measure(round(1000 * a["men"] / a["women"]),
                              unit="males_per_1000_females", year=2021, source=SOURCE_2021),
            sex_ratio_note=f"Men per thousand women in {universe}.",
            ethnicity=present(ethnicity[key]["counts"]), ethnicity_year=2021,
            ethnicity_note=f"Ethnic origin, as asked of everyone in {universe} (Table 02.04).",
            religion=present(religion[key]["counts"]), religion_year=2021,
            religion_note=(f"Religion, as asked of everyone in {universe} (Table 02.06); "
                           "\"No Religious Affiliation\" is written as No religion."),
            sources=[{"field": "median_age/sex_ratio/ethnicity/religion", "name": SOURCE_2021,
                      "url": URL_2021, "year": 2021},
                     {"field": "population", "name": SOURCE_2010, "url": URL_2010,
                      "year": 2010}]))
    write_json(OUT, records)
    log(f"  wrote {OUT.name}: {len(records)} parishes")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
