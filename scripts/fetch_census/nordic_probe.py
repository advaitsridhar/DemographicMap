#!/usr/bin/env python3
"""Read-only probes of the Nordic and Baltic statistical offices' APIs.

The adapters for Sweden, Norway, Denmark, Finland, Iceland, Estonia, Latvia and
Lithuania each lean on a table layout -- which variables a table has, which
codes its geography holds, which years it covers. None of that is guessable
from here, and the sandbox that writes the adapters cannot reach a statistical
host, so this asks the runner and prints only what a decision needs: a table's
variables with their value counts and first and last codes, a folder's
listing, a search's hits, or a small query's first cells.

Each probe is named, and several run in one call:

    python -m scripts.fetch_census.nordic_probe --run swe_tables,nor_07459

Nothing is written; the output is the log.
"""

from __future__ import annotations

import argparse
import json
import re
import time
import urllib.error
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
from collections import Counter
from typing import Any

UA = "DemographicMap/1.0 (+https://github.com/advaitsridhar/DemographicMap) python-urllib"
TIMEOUT = 120

SCB = "https://api.scb.se/OV0104/v1/doris/en/ssd"
SSB = "https://data.ssb.no/api/v0/en/table"
STATBANK = "https://api.statbank.dk/v1"
STATFIN = "https://pxdata.stat.fi/PxWeb/api/v1/en/StatFin"
HAGSTOFA = "https://px.hagstofa.is/pxen/api/v1/en"
STAT_EE = "https://andmed.stat.ee/api/v1/en/stat"
CSB = "https://data.stat.gov.lv/api/v1/en/OSP_PUB"
OSP = "https://osp-rs.stat.gov.lt/rest_xml"


def fetch(url: str, payload: Any = None, accept: str = "application/json") -> bytes:
    data = json.dumps(payload).encode() if payload is not None else None
    headers = {"User-Agent": UA, "Accept": accept}
    if data is not None:
        headers["Content-Type"] = "application/json"
    delay = 3.0
    for attempt in range(4):
        try:
            req = urllib.request.Request(url, data=data, headers=headers)
            with urllib.request.urlopen(req, timeout=TIMEOUT) as resp:
                return resp.read()
        except urllib.error.HTTPError as exc:
            if exc.code == 429 and attempt < 3:
                time.sleep(delay)
                delay *= 2
                continue
            body = exc.read()[:300].decode("utf-8", "replace")
            raise RuntimeError(f"HTTP {exc.code}: {' '.join(body.split())}") from None
    raise RuntimeError("gave up after 429s")


def get_json(url: str, payload: Any = None) -> Any:
    raw = fetch(url, payload).decode("utf-8-sig", "replace")
    return json.loads(raw)


def short(values: list[str], texts: list[str], k: int = 6) -> str:
    pairs = [f"{v}={t}" for v, t in zip(values, texts)]
    if len(pairs) <= 2 * k:
        return "; ".join(pairs)
    return "; ".join(pairs[:k]) + " ... " + "; ".join(pairs[-k:])


def px_meta(url: str, grep: str | None = None, allvals: str | None = None) -> None:
    """A PxWeb table's variables: code, label, count, and first and last values."""
    meta = get_json(url)
    print(f"  title: {meta.get('title')}")
    for var in meta.get("variables", []):
        vals, texts = var.get("values", []), var.get("valueTexts", [])
        flags = "".join(f" [{k}]" for k in ("elimination", "time") if var.get(k))
        print(f"  var {var['code']!r} ({var.get('text')}){flags}: {len(vals)} values")
        print(f"      {short(vals, texts)}")
        lens = Counter(len(v) for v in vals)
        prefixes = Counter(re.match(r"[A-Za-z\-]*", v).group(0) for v in vals)
        if len(vals) > 12:
            print(f"      code lengths {dict(sorted(lens.items()))}; "
                  f"prefixes {dict(prefixes.most_common(8))}")
        if grep:
            hits = [f"{v}={t}" for v, t in zip(vals, texts) if re.search(grep, f"{v} {t}", re.I)]
            if hits:
                print(f"      /{grep}/: {len(hits)} -- " + "; ".join(hits[:40]))
        if allvals and re.fullmatch(allvals, var["code"]):
            print("      ALL: " + "; ".join(f"{v}={t}" for v, t in zip(vals, texts)))


def px_list(url: str, grep: str | None = None) -> None:
    items = get_json(url)
    print(f"  {len(items)} items")
    for item in items:
        line = f"{item.get('id')} [{item.get('type')}] {item.get('text')}"
        if not grep or re.search(grep, line, re.I):
            print(f"    {line}")


def px_search(url: str, limit: int = 40) -> None:
    hits = get_json(url)
    print(f"  {len(hits)} hits")
    for hit in hits[:limit]:
        print(f"    {hit.get('id')} | {hit.get('path')} | {hit.get('title')}")


def px_post(url: str, query: dict[str, Any], limit: int = 30) -> None:
    body = get_json(url, query)
    dims = body.get("dimension", {})
    ids = body.get("id", [])
    print(f"  dims {ids} sizes {body.get('size')}")
    for name in ids:
        cat = dims[name]["category"]
        index = cat["index"]
        order = sorted(index, key=index.get) if isinstance(index, dict) else list(index)
        labels = cat.get("label", {})
        print(f"    {name}: {short(order, [labels.get(c, c) for c in order], 5)}")
    values = body.get("value")
    if isinstance(values, dict):
        values = list(values.values())
    print(f"  {len(values)} values; first {values[:limit]}")
    nums = [v for v in values if isinstance(v, (int, float))]
    print(f"  sum of all {sum(nums):,.0f}")


def statbank_info(table: str, grep: str | None = None) -> None:
    info = get_json(f"{STATBANK}/tableinfo/{table}?lang=en&format=JSON")
    print(f"  {info.get('id')}: {info.get('text')} | unit {info.get('unit')} | "
          f"updated {info.get('updated')}")
    for var in info.get("variables", []):
        vals = var.get("values", [])
        ids, texts = [v["id"] for v in vals], [v["text"] for v in vals]
        print(f"  var {var['id']!r} ({var.get('text')}) elim={var.get('elimination')} "
              f"time={var.get('time')}: {len(vals)} values")
        print(f"      {short(ids, texts)}")
        if grep:
            hits = [f"{i}={t}" for i, t in zip(ids, texts) if re.search(grep, f"{i} {t}", re.I)]
            if hits:
                print(f"      /{grep}/: {len(hits)} -- " + "; ".join(hits[:40]))


def statbank_tables(grep: str) -> None:
    tables = get_json(f"{STATBANK}/tables?lang=en&format=JSON&includeInactive=true")
    print(f"  {len(tables)} tables")
    for t in tables:
        line = f"{t['id']} | {t['text']} | {t.get('firstPeriod')}-{t.get('latestPeriod')} | " \
               f"{','.join(t.get('variables', []))} | active={t.get('active')}"
        if re.search(grep, line, re.I):
            print(f"    {line}")


def sdmx_dataflows(grep: str) -> None:
    raw = fetch(f"{OSP}/dataflow/", accept="application/xml")
    root = ET.fromstring(raw)
    ns_name = "{http://www.sdmx.org/resources/sdmxml/schemas/v2_1/common}Name"
    flows = [el for el in root.iter() if el.tag.endswith("}Dataflow")]
    print(f"  {len(flows)} dataflows")
    shown = 0
    for flow in flows:
        names = {n.get("{http://www.w3.org/XML/1998/namespace}lang"): (n.text or "")
                 for n in flow.findall(ns_name)}
        line = f"{flow.get('id')} | {names.get('en', '')} | {names.get('lt', '')}"
        if re.search(grep, line, re.I):
            print(f"    {line}")
            shown += 1
            if shown > 400:
                print("    ...")
                break


def sdmx_series(flow: str, query: str = "", limit: int = 6) -> None:
    """An OSP dataflow's series: how many, their dimensions, and a few keys."""
    raw = fetch(f"{OSP}/data/{flow}{query}", accept="application/xml")
    root = ET.fromstring(raw)
    series = [el for el in root.iter() if el.tag.endswith("Series")]
    print(f"  {len(raw):,} bytes, {len(series)} series")
    if series and not any(kv.tag.endswith("}Value") for kv in series[0].iter()):
        # Structure-specific form: the key is the Series element's attributes
        # and each Obs carries TIME_PERIOD and OBS_VALUE.
        dims2: dict[str, Counter] = {}
        for s in series:
            for k, v in s.attrib.items():
                dims2.setdefault(k, Counter())[v] += 1
        for name, values in dims2.items():
            vals = list(values)
            print(f"    dim {name}: {len(vals)} values: {vals[:15]}{' ...' if len(vals) > 15 else ''}")
        for s in series[:limit]:
            obs = [(o.get("TIME_PERIOD"), o.get("OBS_VALUE")) for o in s if o.tag.endswith("Obs")]
            print(f"    {dict(s.attrib)} -> {obs[-3:]}")
        return
    if not series:
        print("  " + " ".join(raw[:1500].decode("utf-8", "replace").split()))
    dims: dict[str, Counter] = {}
    for s in series:
        for kv in s.iter():
            if kv.tag.endswith("}Value") and kv.get("id"):
                dims.setdefault(kv.get("id"), Counter())[kv.get("value")] += 1
    for name, values in dims.items():
        vals = list(values)
        print(f"    dim {name}: {len(vals)} values: {vals[:12]}{' ...' if len(vals) > 12 else ''}")
    for s in series[:limit]:
        key = {kv.get("id"): kv.get("value") for kv in s.iter()
               if kv.tag.endswith("}Value") and kv.get("id")}
        obs = []
        for o in s.iter():
            if o.tag.endswith("}Obs"):
                d = o.find(".//{*}ObsDimension")
                v = o.find(".//{*}ObsValue")
                obs.append((d.get("value") if d is not None else None,
                            v.get("value") if v is not None else None))
        print(f"    {key} -> {obs[-3:]}")


def fin_class(year: str) -> None:
    """Statistics Finland's municipality -> sub-region key for one year: the sub-regions."""
    url = (f"https://data.stat.fi/api/classifications/v2/correspondenceTables/"
           f"kunta_1_{year}0101%23seutukunta_1_{year}0101/maps?content=data&meta=max&lang=fi")
    body = get_json(url)
    print(f"  {len(body)} maps; first: {json.dumps(body[0], ensure_ascii=False)[:1500]}")
    targets = Counter(m.get("targetLocalId") for m in body)
    print(f"  {len(targets)} sub-regions")


def isl(fn, *a, **k):
    time.sleep(6)
    return fn(*a, **k)


def sdmx_obs(flow: str, query: str = "", limit: int = 8) -> None:
    """An OSP flow in the all-dimensions form: each Obs carries its own key."""
    raw = fetch(f"{OSP}/data/{flow}{query}", accept="application/xml")
    root = ET.fromstring(raw)
    obs = [el for el in root.iter() if el.tag.endswith("}Obs")]
    print(f"  {len(raw):,} bytes, {len(obs)} observations")
    dims: dict[str, Counter] = {}
    rows = []
    for o in obs:
        key = {}
        value = None
        for el in o.iter():
            if el.tag.endswith("}Value") and el.get("id"):
                key[el.get("id")] = el.get("value")
            if el.tag.endswith("}ObsValue"):
                value = el.get("value")
        for k, v in key.items():
            dims.setdefault(k, Counter())[v] += 1
        rows.append((key, value))
    for name, values in dims.items():
        vals = list(values)
        print(f"    dim {name}: {len(vals)} values: {vals[:40]}{' ...' if len(vals) > 40 else ''}")
    for key, value in rows[:limit]:
        print(f"    {key} = {value}")


def links(url: str, grep: str, limit: int = 80) -> None:
    raw = fetch(url, accept="text/html,*/*").decode("utf-8", "replace")
    hrefs = re.findall(r"""href=["']([^"'#]+)["'][^>]*>([^<]{0,120})""", raw, re.I)
    print(f"  {len(raw):,} characters, {len(hrefs)} links")
    shown = 0
    for href, label in hrefs:
        line = f"{' '.join(label.split())} -> {urllib.parse.urljoin(url, href)}"
        if re.search(grep, line, re.I):
            print(f"    {line}")
            shown += 1
            if shown >= limit:
                break


def sdmx_codes(path: str, grep: str | None = None, limit: int = 120) -> None:
    """The codelists an OSP structure refers to: id and names of each code."""
    raw = fetch(f"{OSP}/{path}", accept="application/xml")
    root = ET.fromstring(raw)
    lists = [el for el in root.iter() if el.tag.endswith("}Codelist")]
    print(f"  {len(raw):,} bytes, {len(lists)} codelists")
    for cl in lists:
        codes = [el for el in cl if el.tag.endswith("}Code")]
        print(f"    codelist {cl.get('id')}: {len(codes)} codes")
        shown = 0
        for code in codes:
            names = {n.get("{http://www.w3.org/XML/1998/namespace}lang"): n.text
                     for n in code if n.tag.endswith("}Name")}
            line = f"{code.get('id')} = {names.get('en')} | {names.get('lt')}"
            if not grep or re.search(grep, line, re.I):
                print(f"      {line}")
                shown += 1
                if shown >= limit:
                    break


def cdx(query: str, limit: int = 60) -> None:
    """What the Wayback Machine holds under a URL prefix."""
    url = ("https://web.archive.org/cdx/search/cdx?" + query
           + f"&output=json&limit={limit}&collapse=urlkey")
    rows = get_json(url)
    print(f"  {max(len(rows) - 1, 0)} captures")
    for row in rows[1:]:
        print(f"    {row[1]} {row[2]} {row[3] if len(row) > 3 else ''} {row[4] if len(row) > 4 else ''}")


def px_codes(url: str, var: str, pattern: str) -> None:
    """Every value of one variable whose code matches, as code=label, compactly."""
    meta = get_json(url)
    v = next(x for x in meta["variables"] if x["code"] == var)
    hits = [f"{c}={t}" for c, t in zip(v["values"], v["valueTexts"]) if re.fullmatch(pattern, c)]
    print(f"  {len(hits)} of {len(v['values'])}")
    for i in range(0, len(hits), 12):
        print("    " + "; ".join(hits[i:i + 12]))


def xlsx(url: str, rows: int = 30, width: int = 220) -> None:
    """A workbook's sheets, their sizes and first rows."""
    import io
    import openpyxl
    raw = fetch(url, accept="*/*")
    print(f"  {len(raw):,} bytes, starts {raw[:4]!r}")
    book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    for sheet in book.worksheets:
        print(f"  sheet {sheet.title!r}: {sheet.max_row} rows x {sheet.max_column} columns")
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= rows:
                break
            cells = ["" if c is None else str(c) for c in row]
            while cells and not cells[-1]:
                cells.pop()
            print(f"    {i + 1}: " + " | ".join(cells)[:width])


def osp_dims(flow: str, year: int = 2026) -> None:
    sdmx_obs(flow, f"?startPeriod={year}&endPeriod={year}", limit=2)


def xls(url: str, rows: int = 40, width: int = 220) -> None:
    """A legacy .xls workbook's sheets and first rows."""
    import xlrd
    raw = fetch(url, accept="*/*")
    print(f"  {len(raw):,} bytes, starts {raw[:4]!r}")
    book = xlrd.open_workbook(file_contents=raw)
    for sheet in book.sheets():
        print(f"  sheet {sheet.name!r}: {sheet.nrows} rows x {sheet.ncols} columns")
        for i in range(min(rows, sheet.nrows)):
            cells = [str(c) for c in sheet.row_values(i)]
            while cells and not cells[-1]:
                cells.pop()
            print(f"    {i + 1}: " + " | ".join(cells)[:width])


def text(url: str, grep: str | None = None, chars: int = 1500) -> None:
    raw = fetch(url, accept="*/*").decode("utf-8", "replace")
    print(f"  {len(raw):,} characters")
    if grep:
        hits = re.findall(grep, raw, re.I)
        print(f"  /{grep}/: {len(hits)} -- {hits[:40]}")
    else:
        print("  " + " ".join(raw[:chars].split()))


def px_walk(url: str, title: str | None = None, depth: int = 2, meta: str | None = None,
            pause: float = 0.4, grep: str | None = None, allvals: str | None = None) -> None:
    """A PxWeb folder, walked ``depth`` levels down: every table's id and title
    (those matching ``title``), and the variables of each whose id or title
    matches ``meta`` -- one call to see what a folder holds and how."""
    def walk(at: str, level: int) -> None:
        time.sleep(pause)
        try:
            items = get_json(at)
        except Exception as exc:                  # a folder's failure is its answer
            print(f"  {at}: FAILED {exc.__class__.__name__}: {str(exc)[:200]}")
            return
        for item in items:
            line = f"{item.get('id')} [{item.get('type')}] {item.get('text')}"
            if item.get("type") == "l":
                print(f"  {'  ' * (2 - level)}{line}")
                if level > 1:
                    walk(f"{at.rstrip('/')}/{item['id']}", level - 1)
            elif not title or re.search(title, line, re.I):
                print(f"  {'  ' * (2 - level)}{line}")
                if meta and re.search(meta, line, re.I):
                    time.sleep(pause)
                    try:
                        px_meta(f"{at.rstrip('/')}/{item['id']}", grep=grep, allvals=allvals)
                    except Exception as exc:
                        print(f"    meta FAILED {exc.__class__.__name__}: {str(exc)[:200]}")
    walk(url, depth)


def pdf(url: str, pages: str = "1-2", chars: int = 3000, grep: str | None = None) -> None:
    """A PDF's page count and the text of some pages (pdfplumber), or the lines
    matching ``grep`` on every page."""
    import io
    import pdfplumber
    raw = fetch(url, accept="*/*")
    print(f"  {len(raw):,} bytes, starts {raw[:5]!r}")
    with pdfplumber.open(io.BytesIO(raw)) as doc:
        print(f"  {len(doc.pages)} pages")
        lo, _, hi = pages.partition("-")
        wanted = range(int(lo) - 1, min(int(hi or lo), len(doc.pages)))
        for i, page in enumerate(doc.pages):
            text = page.extract_text() or ""
            if grep:
                for line in text.splitlines():
                    if re.search(grep, line, re.I):
                        print(f"    p{i + 1}: {line[:220]}")
            elif i in wanted:
                print(f"  --- page {i + 1}")
                print("    " + text[:chars].replace("\n", "\n    "))


def pdf_heads(urls: list[str], lines: int = 3) -> None:
    """For each URL: its HTTP outcome and, for a PDF, its page count and the
    first lines of page 1 -- which edition a fixed file name holds today."""
    import io
    import pdfplumber
    for url in urls:
        try:
            raw = fetch(url, accept="*/*")
        except Exception as exc:
            print(f"  {url.rsplit('/', 1)[-1][:110]}: {str(exc)[:90]}")
            continue
        if not raw.startswith(b"%PDF"):
            print(f"  {url.rsplit('/', 1)[-1][:110]}: {len(raw):,} bytes, not a PDF")
            continue
        with pdfplumber.open(io.BytesIO(raw)) as doc:
            head = (doc.pages[0].extract_text() or "").splitlines()[:lines]
            print(f"  {url.rsplit('/', 1)[-1][:110]}: {len(doc.pages)} pages -- "
                  + " | ".join(h[:120] for h in head))
        time.sleep(1.0)


SVK = "https://www.svenskakyrkan.se/filer/1374643/"


def klass_codes(url: str) -> None:
    body = get_json(url)
    codes = body.get("codes", [])
    print(f"  {len(codes)} codes")
    print("  " + "; ".join(f"{c['code']}={c['name']}" for c in codes))


def px_notes(url: str, query: dict[str, Any], grep: str | None = None, chars: int = 900) -> None:
    """A small PxWeb query answered as a .px file: its NOTE, VALUENOTE and SOURCE
    keywords, which the JSON metadata leaves out, and the data line."""
    q = dict(query)
    q["response"] = {"format": "px"}
    raw = fetch(url, q, accept="*/*")
    for enc in ("utf-8", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    print(f"  {len(raw):,} bytes")
    keys = r"(?:NOTEX?|VALUENOTEX?|CELLNOTEX?|SOURCE|CONTACT|INFOFILE|DESCRIPTION|DATANOTE|CONTENTS|DATA)"
    for m in re.finditer(keys + r"(?:\[[^\]]*\])?(?:\([^)]*\))?=(.*?);\s*(?=[A-Z-]+(?:\[|\(|=)|$)",
                         text, re.S):
        chunk = " ".join(m.group(0).replace('"\r\n"', "").replace('"\n"', "").split())
        if not grep or re.search(grep, chunk, re.I):
            print(f"    {chunk[:chars]}")


# ---------------------------------------------------------------------------
# The probes
# ---------------------------------------------------------------------------

PROBES: dict[str, Any] = {
    # Sweden
    "swe_tables": lambda: px_list(f"{SCB}/BE/BE0101/BE0101A"),
    "swe_befolkning": lambda: px_meta(f"{SCB}/BE/BE0101/BE0101A/BefolkningNy"),
    # Norway
    "nor_07459": lambda: px_meta(f"{SSB}/07459", grep=r"^(30|38|54|0710|0301|1804|5001) ",
                                 ),
    "nor_search_church": lambda: px_search(f"{SSB}/?query=church"),
    "nor_search_relig": lambda: px_search(f"{SSB}/?query=religious"),
    "nor_klass_2017": lambda: klass_codes(
        "https://data.ssb.no/api/klass/v1/classifications/131/codesAt.json?date=2017-06-01"),
    # Denmark
    "dnk_folk1a": lambda: statbank_info("FOLK1A"),
    "dnk_km1": lambda: statbank_info("KM1"),
    "dnk_tables": lambda: statbank_tables(r"church|relig|folkekirke|member|language|sprog"),
    # Finland
    "fin_vaerak": lambda: px_list(f"{STATFIN}/vaerak/", grep=r"relig|age|language|area|sub-region"),
    "fin_11re": lambda: px_meta(f"{STATFIN}/vaerak/11re.px", grep=r"^SK"),
    "fin_11rl": lambda: px_meta(f"{STATFIN}/vaerak/11rl.px", grep=r"^SK"),
    "fin_search_relig": lambda: px_search(f"{STATFIN}?query=religious"),
    # Iceland
    "isl_search_muni": lambda: px_search(f"{HAGSTOFA}/Ibuar?query=municipalities"),
    "isl_search_relig": lambda: px_search(f"{HAGSTOFA}/Samfelag?query=religious"),
    "isl_list": lambda: px_list(f"{HAGSTOFA}/Ibuar"),
    # Estonia
    "est_search_div": lambda: px_search(f"{STAT_EE}?query=administrative%20division"),
    "est_search_unit": lambda: px_search(f"{STAT_EE}?query=administrative%20unit"),
    "est_search_rl": lambda: px_search(f"{STAT_EE}?query=ethnic%20nationality"),
    # Latvia
    "lva_search_parish": lambda: px_search(f"{CSB}?query=parish"),
    "lva_search_terr": lambda: px_search(f"{CSB}?query=territorial%20unit"),
    "lva_search_census": lambda: px_search(f"{CSB}?query=ethnicity"),
    # Lithuania
    "ltu_flows": lambda: sdmx_dataflows(
        r"amži|age|tautyb|ethnic|kalb|langu|tikyb|relig|surašym|census"),
    # Round 18d: which edition the Church of Sweden's fixed file names hold
    # today, and the Finnish church's economic units in full.
    "r18d_swe_live": lambda: pdf_heads(
        [SVK + n for n in ("MedlemsutvecklingLKF.pdf", "MedlemsutvecklingLKF(1).pdf",
                           "MedlemsutvecklingLKF(2).pdf", "MedlemsutvecklingLKF(3).pdf",
                           "NyckeltalLKF.pdf", "NyckeltalLKF(1).pdf", "NyckeltalLKF(2).pdf")]
        + [SVK + "Medlemmar%20i%20Svenska%20kyrkan%20i%20forhallande%20till%20folkmangd%2031%20"
           f"december%20{y}%20per%20forsamling,%20kommun%20och%20lan%20samt%20riket%20(pdf).pdf"
           for y in range(2020, 2026)]
        + [SVK + f"Medlemsutveckling%20{y - 1}-{y},%20lan,%20kommun%20och%20forsamling%20samt%20"
           "riket%20(pdf).pdf" for y in range(2020, 2026)]),
    "r18d_fin_evl_units": lambda: xlsx(
        "https://web.archive.org/web/20220705145540id_/https://www.kirkontilastot.fi/tiedostot/"
        "J%C3%A4senm%C3%A4%C3%A4r%C3%A42020.xlsx", rows=400, width=140),
    # Round 18c: the Church of Sweden's membership against the population by
    # kommun, which its statistics folder has published as PDFs.
    "r18c_swe_cdx": lambda: cdx(
        "url=svenskakyrkan.se/filer/1374643/*&filter=original:.*(?:[Ff]olkm|[Kk]ommun|LKF|"
        "[Mm]edlemsutv|[Ff]orsamling).*", limit=300),
    "r18c_swe_page": lambda: text(
        "https://www.svenskakyrkan.se/statistik",
        grep=r"[^<>\"]{0,160}(?:folkm[aä]ngd|per kommun|kommun och l[aä]n|LKF)[^<>\"]{0,160}"),
    "r18c_swe_pdf2020": lambda: pdf(
        "https://web.archive.org/web/20220120151620id_/https://www.svenskakyrkan.se/filer/"
        "1374643/Medlemmar%20i%20Svenska%20kyrkan%20i%20forhallande%20till%20folkmangd%2031%20"
        "december%202020%20per%20forsamling,%20kommun%20och%20lan%20samt%20riket%20(pdf).pdf",
        pages="1-2", chars=2500),
    "r18c_swe_pdf2020_tail": lambda: pdf(
        "https://web.archive.org/web/20220120151620id_/https://www.svenskakyrkan.se/filer/"
        "1374643/Medlemmar%20i%20Svenska%20kyrkan%20i%20forhallande%20till%20folkmangd%2031%20"
        "december%202020%20per%20forsamling,%20kommun%20och%20lan%20samt%20riket%20(pdf).pdf",
        grep=r"(?:kommun|l[aä]n|riket|totalt|summa)"),
    "r18c_swe_lkf": lambda: pdf(
        "https://web.archive.org/web/20220401215811id_/https://www.svenskakyrkan.se/filer/"
        "1374643/MedlemsutvecklingLKF(1).pdf", pages="1-2", chars=2500),
    # Round 18b: what round 18 pointed at.
    "r18b_fin_vaerak": lambda: px_list(f"{STATFIN}/vaerak/"),
    "r18b_fin_11ru": lambda: px_meta(
        f"{STATFIN}/vaerak/11ru.px", grep=r"^(KU091|KU049|SK011|MK01|SSS) ",
        allvals=r"(?i)(?!alue|vuosi|timeperiod|sukupuoli|ikaryhma|contentscode).*"),
    "r18b_fin_159t": lambda: px_meta(
        f"{STATFIN}/vaerak/159t.px", grep=r"^(KU091|KU049|SK011|MK01|SSS) ",
        allvals=r"(?i)(?!alue|vuosi|timeperiod|sukupuoli|ikaryhma).*"),
    "r18b_nor_09817": lambda: px_meta(
        f"{SSB}/09817", grep=r"^(0710|0301|1804|5001|3005|1120|0101|K-0301) ",
        allvals=r"(?i)(?!region|tid).*"),
    "r18b_nor_11366": lambda: px_meta(
        f"{SSB}/11366", grep=r"^(0710|0301|1804|5001|3005|1120|0101) ",
        allvals=r"(?i)(?!region|tid).*"),
    "r18b_fin_evl_2019": lambda: xlsx(
        "https://web.archive.org/web/20220707222524id_/https://www.kirkontilastot.fi/tiedostot/"
        "J%C3%A4senm%C3%A4%C3%A4r%C3%A42019aluejako2020.xlsx", rows=25),
    "r18b_fin_evl_2019_live": lambda: xlsx(
        "https://www.kirkontilastot.fi/tiedostot/J%C3%A4senm%C3%A4%C3%A4r%C3%A42019aluejako2020.xlsx",
        rows=8),
    "r18b_fin_evl_2020": lambda: xlsx(
        "https://web.archive.org/web/20220705145540id_/https://www.kirkontilastot.fi/tiedostot/"
        "J%C3%A4senm%C3%A4%C3%A4r%C3%A42020.xlsx", rows=25),
    "r18b_fin_tableau1": lambda: text(
        "https://public.tableau.com/views/Jsentilasto2025/Tilastotaulukko.csv", chars=2500),
    "r18b_fin_tableau2": lambda: text(
        "https://public.tableau.com/views/Jsentilasto2025kirkkoonkuuluvuus/"
        "KirkkoonkuuluvuusTalousyksikt.csv", chars=2500),
    "r18b_swe_kyrkan_cdx": lambda: cdx(
        "url=svenskakyrkan.se/filer/*&filter=original:.*(?:[Kk]ommun|[Mm]edlem|[Tt]illh).*",
        limit=150),
    "r18b_isl_skra_old": lambda: xlsx(
        "https://web.archive.org/web/20240704090131id_/https://skra.is/library/Samnyttar-skrar-/"
        "Frettir/trufelagsskraning.xlsx", rows=12),
    # Round 18 (the gap round): nationality, country of birth and background
    # by the units the map draws, as the owner's decision of 19 September 2026
    # allows on the ethnicity field; the churches' own membership by
    # municipality; and the last places the 2011 Baltic censuses might reach.
    "r18_swe_walk": lambda: px_walk(
        f"{SCB}/BE/BE0101", title=r"born|birth|citizen|background|foreign|nationalit|country",
        depth=2, meta=r"region|kommun|county|municipal", grep=r"^(0114|0180|01|00) ",
        allvals=r"(?i)(fodel|medb|land|bakgr|utl).*"),
    "r18_swe_search": lambda: px_search(f"{SCB}?query=country%20of%20birth", limit=60),
    "r18_nor_search1": lambda: px_search(f"{SSB}/?query=country%20background", limit=60),
    "r18_nor_search2": lambda: px_search(f"{SSB}/?query=landbakgrunn", limit=60),
    "r18_nor_search3": lambda: px_search(f"{SSB}/?query=citizenship%20municipality", limit=40),
    "r18_dnk_folk1c": lambda: statbank_info("FOLK1C", grep=r"^(101|147|851|000|084)"),
    "r18_dnk_tables": lambda: statbank_tables(
        r"ancestry|origin|citizenship|country of birth|immigrant|descendant"),
    "r18_fin_11rt": lambda: px_meta(
        f"{STATFIN}/vaerak/11rt.px", grep=r"^(KU091|KU049|SK011|MK01|SSS) ",
        allvals=r"(?i)(?!alue|vuosi|timeperiod|sukupuoli|ikaryhma|contentscode).*"),
    "r18_fin_11rg": lambda: px_meta(
        f"{STATFIN}/vaerak/11rg.px", grep=r"^(KU091|KU049|SK011|MK01|SSS) ",
        allvals=r"(?i)(?!alue|vuosi|timeperiod|sukupuoli|ikaryhma|contentscode).*"),
    "r18_fin_11rp": lambda: px_meta(
        f"{STATFIN}/vaerak/11rp.px", grep=r"^(KU091|KU049|SK011|MK01|SSS) ",
        allvals=r"(?i)(?!alue|vuosi|timeperiod|sukupuoli|ikaryhma|contentscode).*"),
    "r18_isl_walk": lambda: px_walk(
        f"{HAGSTOFA}/Ibuar/mannfjoldi/3_bakgrunnur", depth=2, pause=7.0,
        meta=r"municipal|region|sveitarf|landshl"),
    "r18_swe_kyrkan": lambda: links(
        "https://www.svenskakyrkan.se/statistik",
        r"statistik|siffror|medlem|kommun|xls|filer|pdf|tillh", limit=120),
    "r18_fin_evl_311": lambda: text(
        "https://www.kirkontilastot.fi/viz.php?id=311",
        grep=r"(?:src|href|data-[a-z-]+|url|file)\s*[=:]\s*[\"'][^\"']{4,220}[\"']"),
    "r18_fin_evl_286": lambda: text(
        "https://www.kirkontilastot.fi/viz.php?id=286",
        grep=r"(?:src|href|data-[a-z-]+|url|file)\s*[=:]\s*[\"'][^\"']{4,220}[\"']"),
    "r18_fin_evl_files": lambda: cdx(
        "url=kirkontilastot.fi/*&filter=original:.*(?:xlsx|xls|csv|tiedostot).*", limit=150),
    "r18_isl_skra_cdx": lambda: cdx(
        "url=skra.is/library/*&filter=original:.*(?:[Tt]ru|[Ll]ifssk|[Tt]rufel).*", limit=150),
    "r18_est_smin_cdx": lambda: cdx(
        "url=siseministeerium.ee/*&filter=original:.*(?:[Rr]ahvus|[Rr]ahvastikuregist).*",
        limit=120),
    "r18_lva_od_parish": lambda: px_search(
        "https://data.stat.gov.lv/api/v1/en/OSP_OD?query=parish", limit=60),
    "r18_lva_od_lang": lambda: px_search(
        "https://data.stat.gov.lv/api/v1/en/OSP_OD?query=language", limit=60),
    # Round 17: retries of the archive's refusals, and SSB's own words on
    # where the members of a community outside the Church are counted.
    "r17_ltu_vilnius": lambda: xlsx(
        "https://web.archive.org/web/20220621221246id_/https://osp.stat.gov.lt/documents/10180/"
        "9601028/Vilnius_county_by_largest_ethnic_groups.xlsx", rows=45),
    "r17_ltu_vilnius_lt": lambda: xlsx(
        "https://web.archive.org/web/20220818192852id_/https://osp.stat.gov.lt/documents/10180/"
        "9601028/Gausiausiu_tautybiu_gyventojai_pagal_tautybe_Vilniaus_apskrityje.xlsx", rows=45),
    "r17_ltu_urban": lambda: xlsx(
        "https://web.archive.org/web/20221011003824id_/https://osp.stat.gov.lt/documents/10180/"
        "10367417/Urban_areas_population_by_largest_ethnic_group-EN.xlsx", rows=40),
    "r17_nor_trosamf": lambda: text(
        "https://www.ssb.no/kultur-og-fritid/religion-og-livssyn/statistikk/"
        "trus-og-livssynssamfunn-utanfor-den-norske-kyrkja",
        grep=r"[^<>]{0,300}(?:fylke|kommune|registrert|bustad|busett|bosted|adresse)[^<>]{0,300}"),
    "r17_nor_trosamf_en": lambda: text(
        "https://www.ssb.no/en/kultur-og-fritid/religion-og-livssyn/statistikk/"
        "religious-communities-and-life-stance-communities",
        grep=r"[^<>]{0,300}(?:county|municipalit|registered|resid|address)[^<>]{0,300}"),
    "r17_nor_vardok": lambda: text("https://www.ssb.no/a/metadata/conceptvariable/vardok/2337/nb",
                                   grep=r"[^<>]{0,400}(?:fylke|kommune|registrert|bosted)[^<>]{0,400}"),
    "r17_nor_kirke": lambda: text(
        "https://www.ssb.no/kultur-og-fritid/religion-og-livssyn/statistikk/den-norske-kyrkja",
        grep=r"[^<>]{0,300}(?:sokn|kommune|bustad|busett|registrert|Frøyland)[^<>]{0,300}"),
    "r17_fin_dvv_old": lambda: links(
        "https://web.archive.org/web/2021/https://dvv.fi/tilastot-ja-luettelot",
        r"uskon|rekisteritil|xlsx|tilast|kunn|v[äa]est"),
    "r17_fin_dvv_now": lambda: links("https://dvv.fi/tilastot-ja-luettelot",
                                     r"uskon|rekisteritil|xlsx|tilast|kunn|v[äa]est"),
    # Round 16: the checker's questions. How SSB's KOSTRA table places the
    # members of communities outside the Church; the 2021 census's ethnicity
    # and religion below the country in Lithuania; the register keepers of
    # Finland and Iceland; the universe of Latvia's 2011 home language.
    "r16_nor_12026_px": lambda: px_notes(f"{SSB}/12026", {"query": [
        {"code": "KOKkommuneregion0000", "selection": {"filter": "item", "values": ["1120", "1121"]}},
        {"code": "ContentsCode", "selection": {"filter": "item", "values": [
            "KOSmedlemmerdnk0000", "KOSmedltroslivs0000", "KOSpersoneralle0000"]}},
        {"code": "Tid", "selection": {"filter": "item", "values": ["2019", "2020"]}}]}),
    "r16_nor_12026_v2": lambda: text(
        "https://data.ssb.no/api/pxwebapi/v2/tables/12026/metadata?lang=en",
        grep=r'"notes?"\s*:\s*\[[^\]]{0,1500}'),
    "r16_nor_08531_px": lambda: px_notes(f"{SSB}/08531", {"query": [
        {"code": "Region", "selection": {"filter": "item", "values": ["03", "11"]}},
        {"code": "ReligionLivs", "selection": {"filter": "item", "values": ["999"]}},
        {"code": "ContentsCode", "selection": {"filter": "item", "values": ["Medlemmer"]}},
        {"code": "Tid", "selection": {"filter": "item", "values": ["2020"]}}]}),
    "r16_nor_klepp": lambda: px_post(f"{SSB}/12026", {"query": [
        {"code": "KOKkommuneregion0000", "selection": {"filter": "item", "values": [
            "1119", "1120", "1121", "1122", "0912", "4212"]}},
        {"code": "ContentsCode", "selection": {"filter": "item", "values": [
            "KOSmedlemmerdnk0000", "KOSmedltroslivs0000", "KOSpersoneralle0000"]}},
        {"code": "Tid", "selection": {"filter": "item", "values": ["2016", "2018", "2020"]}}],
        "response": {"format": "json-stat2"}}, limit=60),
    "r16_nor_12026_contents": lambda: px_meta(f"{SSB}/12026", allvals="ContentsCode"),
    "r16_nor_search_livssyn": lambda: px_search(f"{SSB}/?query=livssyn", limit=40),
    "r16_nor_search_trus": lambda: px_search(f"{SSB}/?query=religious%20communities", limit=40),
    "r16_isl_10200": lambda: isl(px_meta, f"{HAGSTOFA}/Samfelag/menning/5_trufelog/trufelogeldra/"
                                 "MAN10200.px", allvals=r"(?i)(?!Tr).*"),
    "r16_fin_cdx1": lambda: cdx("url=dvv.fi/documents/*&filter=original:.*(?:[Uu]skon|[Rr]eligi|"
                                "[Ss]eurakun).*", limit=100),
    "r16_fin_cdx2": lambda: cdx("url=dvv.fi/*&filter=original:.*(?:rekisteritilanne|tilastot).*",
                                limit=80),
    "r16_fin_cdx3": lambda: cdx("url=vrk.fi/*&filter=original:.*(?:[Uu]skon|[Rr]ekisteritil).*",
                                limit=80),
    "r16_fin_pxnet": lambda: links(
        "https://web.archive.org/web/2019/http://pxnet2.stat.fi/PXWeb/pxweb/en/StatFin/"
        "StatFin__vrm__vaerak/", r"relig|uskon|statfin_vaerak_pxt"),
    "r16_ltu_215": lambda: sdmx_obs("S3R167_M3010215_1", "?startPeriod=2021&endPeriod=2021",
                                    limit=6),
    "r16_ltu_162": lambda: sdmx_obs("S3R162_M3010215_2", "?startPeriod=2021&endPeriod=2021",
                                    limit=6),
    "r16_ltu_flows": lambda: sdmx_dataflows(r"M30102(?:1[4-9]|2[0-9])|tikyb|religin|gimtoj"),
    "r16_ltu_vilnius": lambda: xlsx(
        "https://web.archive.org/web/20220621221246id_/https://osp.stat.gov.lt/documents/10180/"
        "9601028/Vilnius_county_by_largest_ethnic_groups.xlsx", rows=45),
    "r16_ltu_urban": lambda: xlsx(
        "https://web.archive.org/web/20221011003824id_/https://osp.stat.gov.lt/documents/10180/"
        "10367417/Urban_areas_population_by_largest_ethnic_group-EN.xlsx", rows=30),
    "r16_lva_tsg1107_px": lambda: px_notes(
        "https://data.stat.gov.lv/api/v1/en/OSP_OD/tautassk/taut/tsk2011/TSG11-07.px", {"query": [
            {"code": "Teritoriālā vienība", "selection": {"filter": "item", "values": [
                "LV", "LV0010000"]}},
            {"code": "Dzimums", "selection": {"filter": "item", "values": ["T"]}},
            {"code": "Skaits, īpatsvars", "selection": {"filter": "item", "values": ["NUMB"]}},
            {"code": "Vecums (pilni gadi)", "selection": {"filter": "item", "values": ["TOTAL"]}}]}),
    "r16_lva_tsg1101_px": lambda: px_notes(
        "https://data.stat.gov.lv/api/v1/en/OSP_OD/tautassk/demogr/tsk2011/TSG11-01.px", {"query": [
            {"code": "Teritoriālā vienība", "selection": {"filter": "item", "values": [
                "LV", "LV0010000"]}},
            {"code": "Dzimums", "selection": {"filter": "item", "values": ["T"]}}]}),
    # Round 15
    "r15_est_rl222": lambda: px_meta(
        f"{STAT_EE}/rahvaloendus/rel2000/rahvus-emakeel-veerkeelte-oskus/RL222.PX",
        allvals=r"(?i)(elukoht|haldus|asustus|maakond|place|residence).*"),
    "r15_isl_skra_xlsx": lambda: xlsx(
        "https://www.skra.is/library/Samnyttar-skrar-/Frettir/20260910_Tru_lifskodunarfelog.xlsx",
        rows=30),
    "r15_fin_2017": lambda: links(
        "https://stat.fi/til/vaerak/2017/01/vaerak_2017_01_2018-10-01_tie_001_fi.html",
        r"tau_|xlsx|uskon|maakun|kunn"),
    "r15_fin_2019": lambda: links("https://stat.fi/til/vaerak/tau.html", r"uskon|relig"),
    # Round 14
    "r14_nor_08531_2020": lambda: px_post(f"{SSB}/08531", {"query": [
        {"code": "Region", "selection": {"filter": "item", "values": [
            "0", "03", "11", "15", "18", "30", "34", "38", "42", "46", "50", "54"]}},
        {"code": "ReligionLivs", "selection": {"filter": "all", "values": ["*"]}},
        {"code": "ContentsCode", "selection": {"filter": "item", "values": ["Medlemmer"]}},
        {"code": "Tid", "selection": {"filter": "item", "values": ["2020"]}}],
        "response": {"format": "json-stat2"}}, limit=80),
    "r14_est_2000eth": lambda: px_list(f"{STAT_EE}/rahvaloendus/rel2000/rahvus-emakeel-veerkeelte-oskus"),
    "r14_lva_tsg1101": lambda: px_meta(
        "https://data.stat.gov.lv/api/v1/en/OSP_OD/tautassk/demogr/tsk2011/TSG11-01.px",
        allvals=r"(?i).*(terit|vien|area).*"),
    "r14_fin_dvv_cdx": lambda: cdx("url=dvv.fi/documents/*&filter=original:.*(?:[Uu]skonto|[Rr]eligi).*",
                                   limit=80),
    "r14_fin_dvv_cdx2": lambda: cdx("url=dvv.fi/*&filter=original:.*(?:[Tt]ilasto|[Ss]tatisti).*",
                                    limit=60),
    "r14_fin_avain_meta": lambda: px_meta(
        "https://pxdata.stat.fi/PxWeb/api/v1/en/Kuntien_avainluvut/uusin/kuntien_avainluvut_viimeisin.px",
        allvals=r"(?i)(tiedot|contents.*|information)"),
    "r14_isl_skra_news": lambda: text(
        "https://www.skra.is/um-okkur/frettir/frett/2026/09/10/"
        "Skraning-i-tru-og-lifsskodunarfelog-fram-til-1.-september-2026/",
        grep=r"[^<>]{0,120}(?:sveitarf|xlsx|landshlut)[^<>]{0,120}"),
    "r14_isl_skra_tru": lambda: links("https://www.skra.is/folk/tru-og-lifsskodun/",
                                      r"xlsx|t[öo]lfr|fj[öo]ldi|sveitarf|frett"),
    # Round 13: the register keepers' own tables, older censuses by old units.
    "r13_isl_10001": lambda: isl(px_meta, f"{HAGSTOFA}/Samfelag/menning/5_trufelog/trufelog/"
                                 "MAN10001.px"),
    "r13_isl_eldra": lambda: isl(px_list, f"{HAGSTOFA}/Samfelag/menning/5_trufelog/trufelogeldra"),
    "r13_isl_10289": lambda: isl(px_meta, f"{HAGSTOFA}/Samfelag/menning/5_trufelog/trufelog/"
                                 "MAN10289.px"),
    "r13_isl_skra1": lambda: links("https://www.skra.is/um-okkur/tolfraedi/",
                                   r"tr[uú]|l[ií]fssk|xlsx|tolfr|sveitarf"),
    "r13_isl_skra2": lambda: links("https://www.skra.is/um-okkur/frettir/",
                                   r"tr[uú]f|l[ií]fssk|xlsx"),
    "r13_fin_dvv1": lambda: links("https://dvv.fi/vaestotietojarjestelman-rekisteritilanne",
                                  r"uskon|relig|xlsx|kunn"),
    "r13_fin_dvv2": lambda: links("https://dvv.fi/en/statistics",
                                  r"relig|xlsx|municip|statist"),
    "r13_fin_dvv3": lambda: links("https://dvv.fi/tilastot", r"uskon|xlsx|kunn|tilast"),
    "r13_fin_avain": lambda: px_list("https://pxdata.stat.fi/PxWeb/api/v1/en/Kuntien_avainluvut/uusin"),
    "r13_est_rel2000": lambda: px_list(f"{STAT_EE}/rahvaloendus/rel2000"),
    "r13_est_rl0430": lambda: px_meta(
        f"{STAT_EE}/rahvaloendus/rel2011/rahvastiku-demograafilised-ja-etno-kultuurilised-naitajad/"
        "rahvus-emakeel-ja-keelteoskus-murded/RL0430.PX", grep=r"OTH|COUNTY"),
    "r13_est_rl0442": lambda: px_meta(
        f"{STAT_EE}/rahvaloendus/rel2011/rahvastiku-demograafilised-ja-etno-kultuurilised-naitajad/"
        "rahvus-emakeel-ja-keelteoskus-murded/RL0442.PX", grep=r"OTH|COUNTY"),
    "r13_lva_taut2011": lambda: px_list("https://data.stat.gov.lv/api/v1/en/OSP_OD/tautassk/taut/tsk2011"),
    "r13_lva_demogr2011": lambda: px_list(
        "https://data.stat.gov.lv/api/v1/en/OSP_OD/tautassk/demogr/tsk2011"),
    "r13_lva_tsg1107": lambda: px_meta(
        "https://data.stat.gov.lv/api/v1/en/OSP_OD/tautassk/taut/tsk2011/TSG11-07.px",
        allvals="(?i).*(area|territ|region|county|valoda|language).*"),
    "r13_nor_08531": lambda: px_meta(f"{SSB}/08531", allvals="Region"),
    # Round 12: registers and surveys for the fields the censuses leave out.
    "r12_swe_kyrkan": lambda: links("https://www.svenskakyrkan.se/statistik",
                                    r"xlsx|xls|kommun|medlem|statistik"),
    "r12_fin_evl": lambda: links("https://www.kirkontilastot.fi/",
                                 r"xlsx|jasen|jäsen|kunta|kunn|tilast"),
    "r12_fin_evl2": lambda: links("https://evl.fi/tietoa-kirkosta/tilastotietoa/",
                                  r"xlsx|jasen|jäsen|kunta|kunn|tilast"),
    "r12_fin_search_uskonto": lambda: px_search(
        "https://pxdata.stat.fi/PxWeb/api/v1/fi/StatFin?query=uskonnollinen"),
    "r12_nor_search_sami": lambda: px_search(f"{SSB}/?query=sami%20language"),
    "r12_dnk_tables_lang": lambda: statbank_tables(r"sprog|language|tongue|dialect"),
    "r12_isl_search_lang": lambda: px_search(f"{HAGSTOFA}/Ibuar?query=language"),
    "r12_lva_search_relig": lambda: px_search(f"{CSB}?query=religious"),
    # Round 11
    "r11_ltu_eth_lt": lambda: xlsx("https://web.archive.org/web/20220722150228id_/https://osp.stat.gov.lt/documents/10180/9601028/Gyventojai_pagal_tautybe.xlsx", rows=45),
    "r11_ltu_rel_lt": lambda: xlsx("https://web.archive.org/web/20220818193000id_/https://osp.stat.gov.lt/documents/10180/9601028/Gyventojai_pagal_religine_bendruomene_0321.xlsx", rows=45),
    "r11_ltu_eth2011": lambda: xls("https://web.archive.org/web/20130929225433id_/http://osp.stat.gov.lt/documents/10180/217110/Gyventojai_pagal_tautybe_savivaldybese.xls/3b346c37-b28f-4dcc-9836-874b6ea951f7"),
    "r11_ltu_rel2011": lambda: xls("https://web.archive.org/web/20130929225123id_/http://osp.stat.gov.lt/documents/10180/217110/Gyv_religine_bendr_savivald.xls/b845994c-bcf6-4568-9b02-e849b55d6d37"),
    "r11_ltu_lang_tail": lambda: xlsx("https://web.archive.org/web/20221227012930id_/https://osp.stat.gov.lt/documents/10180/10367417/Population_by_mother_tongue_in_municipality-EN.xlsx/1c3c9ad4-5fa5-44b1-baa7-e6caef4b740b?version=1.0", rows=100),
    # Round 10
    "r10_ltu_eth": lambda: xlsx("https://web.archive.org/web/20221115072139id_/https://osp.stat.gov.lt/documents/10180/10367417/Population_by_ethnicity_1108.xlsx"),
    "r10_ltu_lang": lambda: xlsx("https://web.archive.org/web/20221227012930id_/https://osp.stat.gov.lt/documents/10180/10367417/Population_by_mother_tongue_in_municipality-EN.xlsx/1c3c9ad4-5fa5-44b1-baa7-e6caef4b740b?version=1.0"),
    "r10_ltu_rel": lambda: xlsx("https://web.archive.org/web/20221115072138id_/https://osp.stat.gov.lt/documents/10180/10367417/Population_by_religious_community_1108.xlsx"),
    "r10_ltu_202": lambda: osp_dims("S3R167_M3010202"),
    "r10_ltu_205": lambda: osp_dims("S3R167_M3010205"),
    "r10_ltu_206": lambda: osp_dims("S3R167_M3010206"),
    "r10_ltu_213": lambda: osp_dims("S3R167_M3010213"),
    "r10_ltu_214": lambda: osp_dims("S3R167_M3010214"),
    "r10_ltu_222": lambda: osp_dims("S3R167_M3010222"),
    "r10_ltu_224": lambda: osp_dims("S3R167_M3010224"),
    # Round 9
    "r9_lva_units_lv": lambda: px_codes("https://data.stat.gov.lv/api/v1/lv/OSP_PUB/POP/IR/IRD/IRD081",
                                        "AREA", r"LV00\d{5}"),
    "r9_lva_units_en": lambda: px_codes(f"{CSB}/POP/IR/IRD/IRD081", "AREA", r"LV00\d{5}"),
    "r9_lva_ird041_lv": lambda: px_codes("https://data.stat.gov.lv/api/v1/lv/OSP_PUB/POP/IR/IRD/IRD041",
                                         "AREA", r"LV.*"),
    "r9_lva_rig040": lambda: px_codes(f"{CSB}/POP/IR/IRE/RIG040", "AllAreaLV", r"LV00\d{5}"),
    # Round 8
    "r8_ltu_cdx_9601028": lambda: cdx("url=osp.stat.gov.lt/documents/10180/9601028/*", limit=200),
    "r8_ltu_cdx_10367417": lambda: cdx("url=osp.stat.gov.lt/documents/10180/10367417/*", limit=200),
    "r8_ltu_cdx_10439642": lambda: cdx("url=osp.stat.gov.lt/documents/10180/10439642/*", limit=200),
    "r8_ltu_cdx_savivald": lambda: cdx("url=osp.stat.gov.lt/documents/10180/*&filter=original:.*(?:savivald|kalb|tikyb|Tikyb|Kalb).*", limit=200),
    "r8_lva_taut": lambda: px_list("https://data.stat.gov.lv/api/v1/en/OSP_OD/tautassk/taut"),
    "r8_lva_demogr": lambda: px_list("https://data.stat.gov.lv/api/v1/en/OSP_OD/tautassk/demogr"),
    # Round 7
    "r7_ltu_cl1": lambda: sdmx_codes("codelist/LSD/savivaldybesRegdb/latest", limit=90),
    "r7_ltu_cl2": lambda: sdmx_codes("datastructure/LSD/M3010203/latest?references=children", limit=90),
    "r7_ltu_cl3": lambda: text(f"{OSP}/dataflow/LSD/S3R167_M3010203/latest?references=all", chars=1800),
    "r7_ltu_cdx": lambda: cdx("url=osp.stat.gov.lt/documents/10180/*&filter=original:.*(?:surasym|tautyb).*"),
    "r7_isl_trufelog": lambda: isl(px_list, f"{HAGSTOFA}/Samfelag/menning/5_trufelog/trufelog"),
    "r7_lva_tsk": lambda: px_list("https://data.stat.gov.lv/api/v1/en/OSP_OD/tautassk"),
    "r7_lva_tsk2021": lambda: px_search("https://data.stat.gov.lv/api/v1/en/OSP_OD?query=2021"),
    "r7_lva_ird081_units": lambda: px_meta(f"{CSB}/POP/IR/IRD/IRD081", grep=r"^LV00\d{5}"),
    "r7_est_rv0241_star": lambda: px_meta(
        f"{STAT_EE}/Lepetatud_tabelid/Rahvastik.Arhiiv/"
        "Rahvastikun%C3%A4itajad%20ja%20koosseis.%20Arhiiv/RV0241.PX", grep=r"\*|COUNTY"),
    "r7_est_rl0429_units": lambda: px_meta(
        f"{STAT_EE}/rahvaloendus/rel2011/rahvastiku-demograafilised-ja-etno-kultuurilised-naitajad/"
        "rahvus-emakeel-ja-keelteoskus-murded/RL0429.PX", grep=r"OTH|COUNTY"),
    "r7_est_rl0433_units": lambda: px_meta(
        f"{STAT_EE}/rahvaloendus/rel2011/rahvastiku-demograafilised-ja-etno-kultuurilised-naitajad/"
        "rahvus-emakeel-ja-keelteoskus-murded/RL0433.PX", allvals="Elukoht"),
    "r7_est_rl21434": lambda: px_meta(
        f"{STAT_EE}/rahvaloendus/rel2021/rahvastiku-demograafilised-ja-etno-kultuurilised-naitajad/"
        "rahvus-emakeel/RL21434.px", grep=r"COUNTY|county|maakond"),
    "r7_est_rl21452": lambda: px_meta(
        f"{STAT_EE}/rahvaloendus/rel2021/rahvastiku-demograafilised-ja-etno-kultuurilised-naitajad/"
        "usk/RL21452.px", grep=r"COUNTY|county|maakond"),
    "r7_est_rv0222u": lambda: px_meta(
        f"{STAT_EE}/rahvastik/rahvastikunaitajad-ja-koosseis/rahvaarv-ja-rahvastiku-koosseis/RV0222U.PX"),
    # Round 6
    "r6_ltu_codes": lambda: sdmx_codes("dataflow/LSD/S3R167_M3010203/latest?references=all",
                                       grep=r"savivald|apskr|^\d\d |LT0|^00 "),
    "r6_ltu_629more": lambda: sdmx_obs("S3R629_M3010217", "?startPeriod=2025&endPeriod=2026", limit=20),
    "r6_ltu_cdx1": lambda: cdx("url=osp.stat.gov.lt/documents/10180/*&filter=original:.*(?:surasym|tautyb|Surasym).*"),
    "r6_ltu_cdx2": lambda: cdx("url=osp.stat.gov.lt/gyventoju-ir-bustu-surasymai*"),
    "r6_ltu_datagov": lambda: text("https://data.gov.lt/datasets?q=sura%C5%A1ymas",
                                   grep=r"[^\"<>]{0,80}ura[sš]ym[^\"<>]{0,80}"),
    "r6_ltu_getdata": lambda: text("https://get.data.gov.lt/datasets/gov/lsd/:ns", chars=2000),
    "r6_fin_avain": lambda: px_list("https://pxdata.stat.fi/PxWeb/api/v1/en/Kuntien_avainluvut"),
    "r6_fin_avain_q": lambda: px_search("https://pxdata.stat.fi/PxWeb/api/v1/en/Kuntien_avainluvut?query=church"),
    "r6_isl_trufelog": lambda: isl(px_list, f"{HAGSTOFA}/Samfelag/menning/5_trufelog"),
    "r6_est_usk2021": lambda: px_list(
        f"{STAT_EE}/rahvaloendus/rel2021/rahvastiku-demograafilised-ja-etno-kultuurilised-naitajad/usk"),
    "r6_lva_od": lambda: px_list("https://data.stat.gov.lv/api/v1/en/OSP_OD"),
    "r6_lva_od_q": lambda: px_search("https://data.stat.gov.lv/api/v1/en/OSP_OD?query=language"),
    # Round 5
    "r5_ltu_203": lambda: sdmx_obs("S3R167_M3010203", "?startPeriod=2026&endPeriod=2026"),
    "r5_ltu_629": lambda: sdmx_obs("S3R629_M3010217", "?startPeriod=2026&endPeriod=2026"),
    "r5_ltu_216": lambda: sdmx_obs("S3R167_M3010216", "?startPeriod=2026&endPeriod=2026", limit=3),
    "r5_ltu_page1": lambda: links("https://osp.stat.gov.lt/gyventoju-ir-bustu-surasymai1",
                                  r"surasym|2021|xls|csv|duomen"),
    "r5_ltu_page2": lambda: links("https://osp.stat.gov.lt/2021-gyventoju-ir-bustu-surasymo-rezultatai",
                                  r"tautyb|kalb|tikyb|xls|csv|duomen"),
    "r5_ltu_datagov": lambda: text("https://data.gov.lt/datasets?q=surašymas", grep=r"[^\"<>]{0,80}surašym[^\"<>]{0,80}"),
    "r5_lva_dbs": lambda: px_list("https://data.stat.gov.lv/api/v1/en"),
    "r5_lva_ird081": lambda: px_meta(f"{CSB}/POP/IR/IRD/IRD081", allvals="AREA"),
    "r5_fin_evang": lambda: px_search(f"{STATFIN}?query=Evangelical"),
    "r5_fin_dbs": lambda: px_list("https://pxdata.stat.fi/PxWeb/api/v1/en"),
    "r5_fin_11ra_sk": lambda: px_meta(f"{STATFIN}/vaerak/11ra.px", grep=r"^SK"),
    "r5_isl_09000": lambda: isl(px_meta, f"{HAGSTOFA}/Ibuar/mannfjoldi/2_byggdir/x_eldraefni/MAN09000.px",
                               allvals=".*"),
    "r5_isl_menning": lambda: isl(px_list, f"{HAGSTOFA}/Samfelag/menning"),
    "r5_isl_02005": lambda: isl(px_meta, f"{HAGSTOFA}/Ibuar/mannfjoldi/2_byggdir/sveitarfelog/MAN02005.px"),
    "r5_est_rel2021": lambda: px_list(
        f"{STAT_EE}/rahvaloendus/rel2021/rahvastiku-demograafilised-ja-etno-kultuurilised-naitajad"),
    "r5_est_2021lang": lambda: px_list(
        f"{STAT_EE}/rahvaloendus/rel2021/rahvastiku-demograafilised-ja-etno-kultuurilised-naitajad/rahvus-emakeel"),
    # Round 4
    "r4_ltu_210": lambda: sdmx_series("S3R167_M3010210", "?startPeriod=2026&endPeriod=2026"),
    "r4_ltu_203": lambda: sdmx_series("S3R167_M3010203", "?startPeriod=2026&endPeriod=2026", limit=3),
    "r4_ltu_629": lambda: sdmx_series("S3R629_M3010217", "?startPeriod=2026&endPeriod=2026", limit=3),
    "r4_ltu_216": lambda: sdmx_series("S3R167_M3010216", "?startPeriod=2026&endPeriod=2026", limit=3),
    "r4_ltu_215": lambda: sdmx_series("S3R167_M3010215_1", "?startPeriod=2021&endPeriod=2026", limit=3),
    "r4_ltu_162": lambda: sdmx_series("S3R162_M3010215_2", "?startPeriod=2021&endPeriod=2026", limit=3),
    "r4_ltu_flows": lambda: sdmx_dataflows(r"gyventojų ir būstų|population and housing|surašymas|census"),
    "r4_lva_ird041": lambda: px_meta(f"{CSB}/POP/IR/IRD/IRD041"),
    "r4_lva_ird031": lambda: px_meta(f"{CSB}/POP/IR/IRD/IRD031"),
    "r4_lva_home": lambda: px_search(f"{CSB}?query=home"),
    "r4_lva_root": lambda: px_list(f"{CSB}"),
    "r4_lva_pop": lambda: px_list(f"{CSB}/POP"),
    "r4_fin_11ra": lambda: px_meta(f"{STATFIN}/vaerak/11ra.px"),
    "r4_fin_passiivi": lambda: px_search(
        "https://pxdata.stat.fi/PxWeb/api/v1/en/StatFin_Passiivi?query=religious%20community"),
    "r4_isl_old": lambda: isl(px_list, f"{HAGSTOFA}/Ibuar/mannfjoldi/2_byggdir/x_eldraefni"),
    "r4_isl_yfirlit": lambda: isl(px_list, f"{HAGSTOFA}/Ibuar/mannfjoldi/1_yfirlit"),
    "r4_isl_bak": lambda: isl(px_list, f"{HAGSTOFA}/Ibuar/mannfjoldi/3_bakgrunnur"),
    "r4_isl_root": lambda: isl(px_list, f"{HAGSTOFA}"),
    "r4_isl_samfelag": lambda: isl(px_list, f"{HAGSTOFA}/Samfelag"),
    "r4_est_rl0429": lambda: px_meta(
        f"{STAT_EE}/rahvaloendus/rel2011/rahvastiku-demograafilised-ja-etno-kultuurilised-naitajad/"
        "rahvus-emakeel-ja-keelteoskus-murded/RL0429.PX", allvals="Elukoht"),
    "r4_est_rl0433": lambda: px_meta(
        f"{STAT_EE}/rahvaloendus/rel2011/rahvastiku-demograafilised-ja-etno-kultuurilised-naitajad/"
        "rahvus-emakeel-ja-keelteoskus-murded/RL0433.PX"),
    "r4_est_rl0452": lambda: px_meta(
        f"{STAT_EE}/rahvaloendus/rel2011/rahvastiku-demograafilised-ja-etno-kultuurilised-naitajad/"
        "usk/RL0452.PX"),
    "r4_est_rl21429": lambda: px_meta(
        f"{STAT_EE}/rahvaloendus/rel2021/rahvastiku-demograafilised-ja-etno-kultuurilised-naitajad/"
        "rahvus-emakeel/RL21429.px"),
    "r4_est_rv0241_all": lambda: px_meta(
        f"{STAT_EE}/Lepetatud_tabelid/Rahvastik.Arhiiv/"
        "Rahvastikun%C3%A4itajad%20ja%20koosseis.%20Arhiiv/RV0241.PX", grep=r"\*"),
    # Round 3
    "r3_swe_ages": lambda: px_meta(f"{SCB}/BE/BE0101/BE0101A/BefolkningCKM", allvals="Alder"),
    "r3_nor_12026": lambda: px_post(f"{SSB}/12026", {"query": [
        {"code": "KOKkommuneregion0000", "selection": {"filter": "item", "values": ["0710", "0301", "0101", "0706"]}},
        {"code": "ContentsCode", "selection": {"filter": "item", "values": [
            "KOSmedlemmerdnk0000", "KOSmedltroslivs0000", "KOSpersoneralle0000"]}},
        {"code": "Tid", "selection": {"filter": "item", "values": ["2016", "2017", "2018"]}}],
        "response": {"format": "json-stat2"}}),
    "r3_nor_eka": lambda: px_meta(f"{SSB}/12026", grep=r"^EKA|^EAK"),
    "r3_fin_class2020": lambda: fin_class("2020"),
    "r3_fin_class2014": lambda: fin_class("2014"),
    "r3_fin_passiivi": lambda: px_search(
        "https://pxdata.stat.fi/PxWeb/api/v1/en/StatFin_Passiivi?query=religious"),
    "r3_isl_folder": lambda: isl(px_list, f"{HAGSTOFA}/Ibuar/mannfjoldi/2_byggdir/sveitarfelog"),
    "r3_isl_2byggdir": lambda: isl(px_list, f"{HAGSTOFA}/Ibuar/mannfjoldi/2_byggdir"),
    "r3_isl_mann": lambda: isl(px_list, f"{HAGSTOFA}/Ibuar/mannfjoldi"),
    "r3_isl_02001": lambda: isl(px_meta, f"{HAGSTOFA}/Ibuar/mannfjoldi/2_byggdir/sveitarfelog/MAN02001.px",
                               allvals="Sveitarf.*"),
    "r3_isl_relig": lambda: isl(px_search, f"{HAGSTOFA}/Ibuar?query=religious%20organizations"),
    "r3_est_ethfolder": lambda: px_list(
        f"{STAT_EE}/rahvaloendus/rel2011/rahvastiku-demograafilised-ja-etno-kultuurilised-naitajad/"
        "rahvus-emakeel-ja-keelteoskus-murded"),
    "r3_est_usk": lambda: px_list(
        f"{STAT_EE}/rahvaloendus/rel2011/rahvastiku-demograafilised-ja-etno-kultuurilised-naitajad/usk"),
    "r3_est_rv0241": lambda: px_meta(
        f"{STAT_EE}/Lepetatud_tabelid/Rahvastik.Arhiiv/"
        "Rahvastikun%C3%A4itajad%20ja%20koosseis.%20Arhiiv/RV0241.PX", allvals="Haldus.*|Elukoht.*|.*[Uu]nit.*"),
    "r3_est_rv0222": lambda: px_meta(
        f"{STAT_EE}/Lepetatud_tabelid/Rahvastik.Arhiiv/"
        "Rahvastikun%C3%A4itajad%20ja%20koosseis.%20Arhiiv/RV0222.PX"),
    "r3_lva_ird": lambda: px_list(f"{CSB}/POP/IR/IRD"),
    "r3_lva_lang": lambda: px_search(f"{CSB}?query=language"),
    "r3_lva_census": lambda: px_search(f"{CSB}?query=census%202021"),
    "r3_ltu_flows": lambda: sdmx_dataflows(
        r"tautyb|gimtoji|native language|religi|tikyb|mother tongue|ethnicity|surašymo"),
    "r3_ltu_s3r167_203": lambda: sdmx_series("S3R167_M3010203", "?startPeriod=2026&endPeriod=2026"),
    "r3_ltu_s3r629": lambda: sdmx_series("S3R629_M3010217", "?startPeriod=2026&endPeriod=2026"),
    "r3_ltu_s3r0155": lambda: sdmx_series("S3R0155_M3010203_1", "?startPeriod=2026&endPeriod=2026"),
    "r3_ltu_s3r167_216": lambda: sdmx_series("S3R167_M3010216", "?startPeriod=2026&endPeriod=2026"),
    "r3_ltu_s3r167_210": lambda: sdmx_series("S3R167_M3010210", "?startPeriod=2026&endPeriod=2026"),
    # Round 2
    "r2_swe_ckm": lambda: px_meta(f"{SCB}/BE/BE0101/BE0101A/BefolkningCKM"),
    "r2_nor_kostra": lambda: px_list(f"{SSB}/kf/kf04/kirke_kostra/SBMENU12107"),
    "r2_nor_kostra2": lambda: px_list(f"{SSB}/kf/kf04/kirke_kostra/SBMENU6460"),
    "r2_nor_trosamf": lambda: px_list(f"{SSB}/kf/kf04/trosamf"),
    "r2_nor_12026": lambda: px_meta(f"{SSB}/12026", grep=r"^(0710|1804|0301|K-0301) "),
    "r2_nor_06326": lambda: px_meta(f"{SSB}/06326"),
    "r2_nor_08531": lambda: px_meta(f"{SSB}/08531"),
    "r2_nor_search_members": lambda: px_search(f"{SSB}/?query=members%20church%20municipality"),
    "r2_dnk_km6": lambda: statbank_info("KM6"),
    "r2_fin_11rf": lambda: px_meta(f"{STATFIN}/vaerak/11rf.px", grep=r"^SK"),
    "r2_fin_11rx": lambda: px_meta(f"{STATFIN}/vaerak/11rx.px"),
    "r2_fin_11rm": lambda: px_meta(f"{STATFIN}/vaerak/11rm.px", grep=r"^SK"),
    "r2_fin_class": lambda: text(
        "https://data.stat.fi/api/classifications/v2/correspondenceTables/"
        "kunta_1_20250101%23seutukunta_1_20250101/maps?content=data&meta=max&lang=fi",
        chars=800),
    "r2_isl_muni": lambda: px_list(f"{HAGSTOFA}/Ibuar/mannfjoldi/2_byggdir/sveitarfelog"),
    "r2_isl_mann": lambda: px_list(f"{HAGSTOFA}/Ibuar/mannfjoldi"),
    "r2_isl_relig": lambda: px_search(f"{HAGSTOFA}/Ibuar?query=religious"),
    "r2_isl_02001": lambda: px_meta(f"{HAGSTOFA}/Ibuar/mannfjoldi/2_byggdir/sveitarfelog/MAN02001.px"),
    "r2_est_archive": lambda: px_list(
        f"{STAT_EE}/Lepetatud_tabelid/Rahvastik.Arhiiv/"
        "Rahvastikun%C3%A4itajad%20ja%20koosseis.%20Arhiiv"),
    "r2_est_rel2011": lambda: px_list(
        f"{STAT_EE}/rahvaloendus/rel2011/rahvastiku-demograafilised-ja-etno-kultuurilised-naitajad"),
    "r2_est_rl0428": lambda: px_meta(
        f"{STAT_EE}/rahvaloendus/rel2011/rahvastiku-demograafilised-ja-etno-kultuurilised-naitajad/"
        "rahvus-emakeel-ja-keelteoskus-murded/RL0428.PX"),
    "r2_est_search_age": lambda: px_search(f"{STAT_EE}?query=sex%20age%20administrative"),
    "r2_lva_ird081": lambda: px_meta(f"{CSB}/POP/IR/IRD/IRD081", grep=r"pag|Ain"),
    "r2_lva_rig010": lambda: px_meta(f"{CSB}/POP/IR/IRD/RIG010", grep=r"Ainaž"),
    "r2_lva_rig040": lambda: px_meta(f"{CSB}/POP/IR/IRE/RIG040", grep=r"Ainaž"),
    "r2_lva_search_age": lambda: px_search(f"{CSB}?query=single%20year"),
    "r2_ltu_flows": lambda: sdmx_dataflows(r"^S3R\d+_M301|surašym|census|Gyventojų surašymo"),
}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--run", required=True,
                    help="comma-separated probe names, or a prefix ending in '*'")
    args = ap.parse_args()
    names: list[str] = []
    for part in args.run.split(","):
        if part.endswith("*"):
            names += [n for n in PROBES if n.startswith(part[:-1])]
        else:
            names.append(part)
    for name in names:
        print(f"== {name}")
        try:
            PROBES[name]()
        except Exception as exc:        # a probe's failure is its answer
            print(f"  FAILED: {exc.__class__.__name__}: {exc}"[:600])
        time.sleep(1.0)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
