#!/usr/bin/env python3
"""Read-only probes for the CIS (Centro de Investigaciones Sociológicas): what
its site serves, and what one study's data file holds, before
``spain_cis`` relies on either.

Nothing is written but the log, and the log is kept short: the lines needed to
decide, never a whole page. A data file is read where it is downloaded (the
runner's HTTP cache, which git ignores) and only aggregates are printed --
variable names and labels, value labels, and frequency tables -- never a
respondent's row.

Subcommands:

* ``url URL... [--grep a,b] [--links a,b] [--chars N] [--around REGEX]`` --
  status, type and size of each URL, the lines of its text carrying a term,
  the links whose address or text carries a term, its first characters, or
  the raw page around a pattern (a script's handler, a form's fields).
* ``pdf URL... [--grep a,b] [--context N]`` -- the text of a PDF (a ficha
  técnica, a questionnaire), whole or the lines carrying a term.
* ``cdx PATTERN [--grep a,b] [--limit N]`` -- the Internet Archive's index of
  captured addresses under a prefix, to learn an official host's file naming.
* ``studies TITLE [--match REGEX] [--files REGEX]`` -- the CIS's studies
  whose title carries a phrase, newest first, with their published files,
  as Spain's open-data catalogue (datos.gob.es) lists them.
* ``xlsx URL [--find REGEX] [--sheets N] [--rows N]`` -- a tabulation
  workbook ("Resultados y cruces"): its sheets, the first rows of some, and
  the header and rows of every sheet where a term appears.
* ``data URL [--vars REGEX] [--labels A,B] [--freq A,B] [--by VAR]
  [--weight VAR]`` -- a study's data file (a zip holding an SPSS file, or the
  SPSS file itself): its members, its size, the variables whose name or label
  matches, the value labels of some, and frequencies, optionally split by
  another variable.

Usage:
    python -m scripts.fetch_census.spain_cis_probe url https://www.cis.es/ --links estudio
    python -m scripts.fetch_census.spain_cis_probe data URL --vars relig,lengua --freq CCAA
"""

from __future__ import annotations

import argparse
import html
import io
import json
import re
import urllib.error
import urllib.request
import zipfile
from collections import Counter, defaultdict
from typing import Any

from ._shared import RAW, log
from common import USER_AGENT  # noqa: E402  (on the path through _shared)

CACHE = RAW / "http_cache" / "cis"


def terms(value: str | None) -> list[str]:
    return [t.strip().lower() for t in (value or "").split(",") if t.strip()]


def get(url: str, timeout: int = 120) -> tuple[int, str, dict[str, str], bytes]:
    """Status, final address, headers and body -- an error's body included."""
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return resp.status, resp.geturl(), dict(resp.headers), resp.read()
    except urllib.error.HTTPError as err:
        return err.code, url, dict(err.headers or {}), err.read() or b""


def text_of(body: str) -> list[str]:
    body = re.sub(r"(?is)<(script|style)\b.*?</\1>", " ", body)
    body = re.sub(r"(?i)<br\s*/?>|</(p|div|tr|li|h\d|td|th|option)>", "\n", body)
    body = html.unescape(re.sub(r"<[^>]+>", " ", body))
    return [re.sub(r"\s+", " ", line).strip() for line in body.splitlines() if line.strip()]


def links_of(body: str) -> list[tuple[str, str]]:
    out = []
    for m in re.finditer(r"(?is)<a\b[^>]*?href\s*=\s*[\"']([^\"']+)[\"'][^>]*>(.*?)</a>", body):
        text = re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", m.group(2)))).strip()
        out.append((html.unescape(m.group(1)), text))
    return out


def cmd_url(args: argparse.Namespace) -> int:
    needles, wanted = terms(args.grep), terms(args.links)
    for url in args.url:
        log(f"== {url}")
        try:
            status, final, headers, raw = get(url, args.timeout)
        except Exception as err:                      # noqa: BLE001 -- reported
            log(f"  {type(err).__name__}: {str(err)[:300]}")
            continue
        kind = headers.get("Content-Type") or headers.get("content-type") or ""
        log(f"  HTTP {status}; {kind}; {len(raw):,} bytes; final {final}"
            + (f"; disposition {headers.get('Content-Disposition')}"
               if headers.get("Content-Disposition") else ""))
        if raw[:4] == b"%PDF" or raw[:2] == b"PK":
            log(f"  binary, starts {raw[:8]!r}")
            continue
        body = raw.decode("utf-8", errors="replace")
        if args.chars:
            log("  " + re.sub(r"\s+", " ", body[:args.chars]))
        if args.jsonld:
            def datasets(node: Any) -> list[dict[str, Any]]:
                if isinstance(node, list):
                    return [d for child in node for d in datasets(child)]
                if not isinstance(node, dict):
                    return []
                found = [node] if "distribution" in node else []
                return found + [d for key, child in node.items() if key != "distribution"
                                for d in datasets(child)]

            blocks = re.findall(r'(?is)<script[^>]*ld\+json[^>]*>(.*?)</script>', body)
            pairs = sorted(set(re.findall(
                r'"contentUrl"\s*:\s*"([^"]+)"\s*,\s*"@type"\s*:\s*"DataDownload"\s*,\s*"name"\s*:\s*"([^"]+)"',
                body)))
            log(f"  {len(blocks)} JSON-LD blocks; {len(pairs)} DataDownload entries in the page")
            for link, name in pairs:
                log(f"    {name}: {link}")
            for block in blocks:
                try:
                    data = json.loads(block, strict=False)
                except ValueError as err:
                    log(f"  JSON-LD unreadable: {err}")
                    continue
                for item in datasets(data):
                    log(f"  dataset {item.get('name')!r}; modified {item.get('dateModified')}; "
                        f"free {item.get('isAccessibleForFree')}")
                    for dist in item.get("distribution") or []:
                        log(f"    {dist.get('name')}: {dist.get('contentUrl')}")
        if args.around:
            spans = [m.start() for m in re.finditer(args.around, body, re.I)]
            log(f"  {len(spans)} raw matches of /{args.around}/")
            last = -10 ** 9
            shown = 0
            for at in spans:
                if at - last < args.context:
                    continue
                last = at
                piece = body[max(0, at - args.context):at + args.context]
                log("    ~ " + re.sub(r"\s+", " ", piece))
                shown += 1
                if shown >= args.limit:
                    break
        lines = text_of(body) if "<" in body[:3000] else body.splitlines()
        if needles:
            hits = [line for line in lines if any(n in line.lower() for n in needles)]
            log(f"  {len(lines)} lines, {len(hits)} with {needles}")
            for line in hits[:args.limit]:
                log(f"    | {line[:args.width]}")
        if wanted:
            found = [(h, t) for h, t in links_of(body)
                     if any(n in h.lower() or n in t.lower() for n in wanted)]
            log(f"  {len(found)} links with {wanted}")
            seen = set()
            for href, text in found:
                if href in seen:
                    continue
                seen.add(href)
                log(f"    -> {href[:args.width]}  [{text[:80]}]")
                if len(seen) >= args.limit:
                    break
    return 0


def cmd_cdx(args: argparse.Namespace) -> int:
    needles = terms(args.grep)
    url = ("https://web.archive.org/cdx/search/cdx?output=json&collapse=urlkey"
           f"&limit={args.scan}&url={args.pattern}")
    status, _, _, raw = get(url, args.timeout)
    log(f"== CDX {args.pattern}: HTTP {status}, {len(raw):,} bytes")
    try:
        rows = json.loads(raw.decode("utf-8", errors="replace") or "[]")
    except ValueError:
        log("  " + raw[:300].decode("utf-8", errors="replace"))
        return 0
    rows = rows[1:] if rows and rows[0] and rows[0][0] == "urlkey" else rows
    hits = [r for r in rows if not needles or any(n in r[2].lower() for n in needles)]
    log(f"  {len(rows)} captured addresses, {len(hits)} with {needles}")
    for r in hits[:args.limit]:
        log(f"    {r[1]} {r[4]} {r[3][:30]} {r[2][:args.width]}")
    return 0


def cmd_pdf(args: argparse.Namespace) -> int:
    """The text of a PDF (a ficha técnica, a questionnaire), whole or by term."""
    from pypdf import PdfReader                        # noqa: PLC0415 -- runner only
    needles = terms(args.grep)
    for url in args.url:
        log(f"== {url}")
        status, final, headers, raw = get(url, args.timeout)
        log(f"  HTTP {status}; {headers.get('Content-Type')}; {len(raw):,} bytes")
        if raw[:4] != b"%PDF":
            log("  not a PDF: " + re.sub(r"\s+", " ", raw[:200].decode("utf-8", "replace")))
            continue
        pages = PdfReader(io.BytesIO(raw)).pages
        lines = []
        for number, page in enumerate(pages, 1):
            for line in (page.extract_text() or "").splitlines():
                if line.strip():
                    lines.append((number, re.sub(r"\s+", " ", line).strip()))
        log(f"  {len(pages)} pages, {len(lines)} lines")
        if needles:
            keep = set()
            for i, (_, line) in enumerate(lines):
                if any(n in line.lower() for n in needles):
                    keep.update(range(max(0, i - args.context), min(len(lines), i + args.context + 1)))
            chosen = [lines[i] for i in sorted(keep)]
        else:
            chosen = lines
        for number, line in chosen[:args.limit]:
            log(f"    p{number}| {line[:args.width]}")
    return 0


def cmd_xlsx(args: argparse.Namespace) -> int:
    """A tabulation workbook: its sheets, their first rows, and where a term sits."""
    import openpyxl                                    # noqa: PLC0415 -- runner only
    raw = download(args.url, args.timeout)
    book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    names = book.sheetnames
    log(f"== {args.url}\n  {len(raw):,} bytes, {len(names)} sheets: "
        + " | ".join(names)[:args.width * 4])

    def rows_of(sheet: Any) -> list[list[Any]]:
        out = []
        for row in sheet.iter_rows(values_only=True):
            out.append(list(row[:args.cols]))
            if len(out) >= args.scan:
                break
        return out

    def show(rows: list[list[Any]], first: int, last: int) -> None:
        for i in range(max(0, first), min(len(rows), last)):
            cells = ["" if v is None else str(v).strip()[:args.cell] for v in rows[i]]
            while cells and not cells[-1]:
                cells.pop()
            if any(cells):
                log(f"    r{i + 1}: " + " | ".join(cells)[:args.width])

    for name in names[:args.sheets]:
        rows = rows_of(book[name])
        log(f"  -- sheet {name!r}: first {args.rows} rows")
        show(rows, 0, args.rows)
    if args.find:
        rx = re.compile(args.find, re.I)
        hits = 0
        for name in names:
            rows = rows_of(book[name])
            found = [i for i, row in enumerate(rows)
                     if any(v is not None and rx.search(str(v)) for v in row)]
            if not found:
                continue
            hits += 1
            log(f"  -- sheet {name!r}: {len(found)} rows match /{args.find}/; header and matches")
            show(rows, 0, args.header)
            for i in found[:args.limit]:
                show(rows, i, i + 1 + args.after)
            if hits >= args.max_sheets:
                break
    return 0


DATOS = "https://datos.gob.es/apidata/catalog/dataset/title/{title}.json?_pageSize={size}&_sort=-issued&_page={page}"
CIS_PUBLISHER = "EA0040772"


def value_of(field: Any) -> str:
    if isinstance(field, list):
        field = next((f for f in field if isinstance(f, dict) and f.get("_lang") == "es"),
                     field[0] if field else "")
    if isinstance(field, dict):
        return str(field.get("_value") or field.get("_about") or "")
    return str(field or "")


def cmd_studies(args: argparse.Namespace) -> int:
    """The CIS's studies whose title carries a phrase, newest first, as datos.gob.es lists them."""
    import urllib.parse
    shown = 0
    for page in range(args.pages):
        url = DATOS.format(title=urllib.parse.quote(args.title), size=args.size, page=page)
        status, _, _, raw = get(url, args.timeout)
        if status != 200:
            log(f"== {url}: HTTP {status}")
            break
        items = json.loads(raw.decode("utf-8", errors="replace"))["result"].get("items", [])
        log(f"== page {page}: {len(items)} datasets")
        if not items:
            break
        for item in items:
            if CIS_PUBLISHER not in str(item.get("publisher")):
                continue
            title = value_of(item.get("title"))
            if args.match and not re.search(args.match, title, re.I):
                continue
            keywords = item.get("keyword") or []
            keywords = [value_of(k) for k in (keywords if isinstance(keywords, list) else [keywords])]
            log(f"  {item.get('issued', '')[:16]} | {title[:150]} | {item.get('identifier', '')}")
            if args.keywords:
                log(f"      keywords: {', '.join(keywords)[:args.width]}")
            dists = item.get("distribution") or []
            for dist in dists if isinstance(dists, list) else [dists]:
                link = dist.get("accessURL") or dist.get("downloadURL") or ""
                if args.files and not re.search(args.files, str(link), re.I):
                    continue
                log(f"      -> {str(link)[:args.width]}")
            shown += 1
            if shown >= args.limit:
                return 0
    return 0


def download(url: str, timeout: int) -> bytes:
    import hashlib
    CACHE.mkdir(parents=True, exist_ok=True)
    path = CACHE / hashlib.sha256(url.encode()).hexdigest()[:24]
    if path.exists():
        return path.read_bytes()
    status, final, headers, raw = get(url, timeout)
    log(f"  HTTP {status}; {headers.get('Content-Type')}; {len(raw):,} bytes; final {final}")
    if status != 200:
        raise SystemExit(f"spain_cis_probe: {url} answered {status}: "
                         f"{raw[:300].decode('utf-8', errors='replace')}")
    path.write_bytes(raw)
    return raw


def open_data(raw: bytes) -> tuple[Any, Any, str]:
    """(frame, meta, member) for the SPSS file in a zip or the SPSS file itself."""
    import tempfile

    import pyreadstat                                  # noqa: PLC0415 -- runner only
    member = "(the file)"
    if raw[:2] == b"PK":
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            for info in zf.infolist():
                log(f"  member {info.filename} {info.file_size:,} bytes")
            savs = [i for i in zf.infolist() if i.filename.lower().endswith((".sav", ".zsav"))]
            if not savs:
                raise SystemExit("spain_cis_probe: no SPSS file in the zip")
            member = savs[0].filename
            raw = zf.read(savs[0])
    with tempfile.NamedTemporaryFile(suffix=".sav") as handle:
        handle.write(raw)
        handle.flush()
        try:
            frame, meta = pyreadstat.read_sav(handle.name)
        except Exception:                               # noqa: BLE001 -- encodings vary
            frame, meta = pyreadstat.read_sav(handle.name, encoding="latin1")
    return frame, meta, member


def label_of(meta: Any, var: str, value: Any) -> str:
    labels = meta.variable_value_labels.get(var, {})
    for key in (value, float(value) if isinstance(value, (int, float)) else value):
        if key in labels:
            return str(labels[key])
    return ""


def cmd_data(args: argparse.Namespace) -> int:
    log(f"== {args.url}")
    frame, meta, member = open_data(download(args.url, args.timeout))
    log(f"  {member}: {len(frame):,} rows, {len(frame.columns)} variables; "
        f"file label {meta.file_label!r}; encoding {getattr(meta, 'file_encoding', '')}")
    if args.min:
        var, _, value = args.min.partition("=")
        frame = frame[frame[var] >= float(value)]
        log(f"  {len(frame):,} rows with {var} >= {value}")
    names = dict(zip(meta.column_names, meta.column_labels))
    if args.vars:
        rx = re.compile(args.vars, re.I)
        hits = [(n, l) for n, l in names.items() if rx.search(n) or rx.search(str(l or ""))]
        log(f"  {len(hits)} variables match /{args.vars}/")
        for n, l in hits[:args.limit]:
            log(f"    {n}: {str(l)[:args.width]}")
    if args.list:
        log("  all variables: " + ", ".join(meta.column_names)[:args.list])
    for var in terms(args.labels):
        var = next((n for n in meta.column_names if n.lower() == var), var)
        labels = meta.variable_value_labels.get(var, {})
        log(f"  {var} ({names.get(var)!s:.100}) value labels: "
            + "; ".join(f"{k:g}={v}" if isinstance(k, float) else f"{k}={v}"
                        for k, v in list(labels.items())[:60]))
    weight = None
    if args.weight:
        weight = next((n for n in meta.column_names if n.lower() == args.weight.lower()), None)
        if weight is None:
            log(f"  no weight variable {args.weight}")
        else:
            w = frame[weight]
            log(f"  weight {weight}: min {w.min():.4f} max {w.max():.4f} mean {w.mean():.4f} "
                f"sum {w.sum():,.1f}; distinct {w.nunique()}")
    by = None
    if args.by:
        by = next((n for n in meta.column_names if n.lower() == args.by.lower()), None)
    def key(value: Any) -> Any:
        # A missing value is NaN, which equals nothing -- itself included --
        # so each one would count as a category of its own.
        return "NaN" if value != value else value

    for var in terms(args.freq):
        var = next((n for n in meta.column_names if n.lower() == var), None)
        if var is None:
            log(f"  no variable for --freq {args.freq}")
            continue
        if by is None:
            counts: Counter = Counter()
            weighted: Counter = Counter()
            for i, value in enumerate(frame[var]):
                value = key(value)
                counts[value] += 1
                if weight:
                    weighted[value] += float(frame[weight].iat[i])
            log(f"  {var} ({names.get(var)!s:.100}):")
            for value, n in sorted(counts.items(), key=lambda kv: str(kv[0])):
                log(f"    {value}: {n:,}" + (f" (w {weighted[value]:,.1f})" if weight else "")
                    + f"  {label_of(meta, var, value)[:60]}")
        else:
            table: dict[Any, Counter] = defaultdict(Counter)
            for a, b in zip(frame[by], frame[var]):
                table[key(a)][key(b)] += 1
            values = sorted({v for c in table.values() for v in c}, key=lambda v: str(v))
            log(f"  {var} by {by}: columns " + ", ".join(
                f"{v}={label_of(meta, var, v)[:18]}" for v in values))
            for group in sorted(table, key=lambda v: str(v)):
                row = table[group]
                log(f"    {group} {label_of(meta, by, group)[:22]:<22} n={sum(row.values()):>5}: "
                    + " ".join(f"{row.get(v, 0)}" for v in values))
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("url")
    p.add_argument("url", nargs="+")
    p.add_argument("--grep")
    p.add_argument("--links")
    p.add_argument("--chars", type=int, default=0)
    p.add_argument("--around", help="print the raw page around each match of this regex")
    p.add_argument("--jsonld", action="store_true",
                   help="print the files a page's schema.org Dataset lists")
    p.add_argument("--context", type=int, default=300)
    p.add_argument("--limit", type=int, default=40)
    p.add_argument("--width", type=int, default=200)
    p.add_argument("--timeout", type=int, default=90)
    p.set_defaults(func=cmd_url)
    p = sub.add_parser("pdf")
    p.add_argument("url", nargs="+")
    p.add_argument("--grep")
    p.add_argument("--context", type=int, default=0)
    p.add_argument("--limit", type=int, default=120)
    p.add_argument("--width", type=int, default=220)
    p.add_argument("--timeout", type=int, default=120)
    p.set_defaults(func=cmd_pdf)
    p = sub.add_parser("cdx")
    p.add_argument("pattern")
    p.add_argument("--grep")
    p.add_argument("--limit", type=int, default=60)
    p.add_argument("--scan", type=int, default=5000)
    p.add_argument("--width", type=int, default=160)
    p.add_argument("--timeout", type=int, default=120)
    p.set_defaults(func=cmd_cdx)
    p = sub.add_parser("studies")
    p.add_argument("title", help="a word or phrase of the title, e.g. Preelectoral")
    p.add_argument("--match", help="a regex the title must also match")
    p.add_argument("--files", help="show only distributions whose address matches this regex")
    p.add_argument("--keywords", action="store_true")
    p.add_argument("--pages", type=int, default=2)
    p.add_argument("--size", type=int, default=50)
    p.add_argument("--limit", type=int, default=40)
    p.add_argument("--width", type=int, default=220)
    p.add_argument("--timeout", type=int, default=120)
    p.set_defaults(func=cmd_studies)
    p = sub.add_parser("xlsx")
    p.add_argument("url")
    p.add_argument("--find")
    p.add_argument("--sheets", type=int, default=1)
    p.add_argument("--rows", type=int, default=12)
    p.add_argument("--cols", type=int, default=60)
    p.add_argument("--scan", type=int, default=3000)
    p.add_argument("--header", type=int, default=8)
    p.add_argument("--after", type=int, default=0)
    p.add_argument("--max-sheets", type=int, default=3)
    p.add_argument("--limit", type=int, default=20)
    p.add_argument("--cell", type=int, default=24)
    p.add_argument("--width", type=int, default=400)
    p.add_argument("--timeout", type=int, default=180)
    p.set_defaults(func=cmd_xlsx)
    p = sub.add_parser("data")
    p.add_argument("url")
    p.add_argument("--vars")
    p.add_argument("--labels")
    p.add_argument("--freq")
    p.add_argument("--by")
    p.add_argument("--weight")
    p.add_argument("--min", help="keep rows where VAR >= VALUE, written VAR=VALUE")
    p.add_argument("--list", type=int, default=0)
    p.add_argument("--limit", type=int, default=60)
    p.add_argument("--width", type=int, default=140)
    p.add_argument("--timeout", type=int, default=300)
    p.set_defaults(func=cmd_data)
    args = ap.parse_args()
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
