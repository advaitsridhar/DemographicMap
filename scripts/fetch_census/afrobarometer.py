#!/usr/bin/env python3
"""Afrobarometer Round 9: religion, ethnicity and language by region, 39 countries.

This is the first source on this map that is a **survey rather than a count**,
and everything below follows from that. 53,444 people were interviewed across
39 African countries; the questions are Q95 "What is your religion, if any?",
Q84a "What is your ethnic community, cultural group or tribe?" and Q2
"Language spoken in home", and every respondent carries the region they were
interviewed in. That is exactly the shape this map wants, for 36 countries
that have nothing at all.

Language was added on 20 September 2026, and the reason it was not here from
the start is worth recording: this module read the published *workbook* by
column position, and the five columns it took did not include Q2. Nothing
said the question was missing -- it simply was not extracted, and the
docstring then said the round did not ask it. The SPSS release, whose
variables are named, is now the source, and ``--extract`` reads it.

What it is not is a census. Afrobarometer is designed to be representative
**nationally**; a region is a sampling stratum, not an estimation domain. The
median region here holds 72 respondents and the thinnest holds 8, so a share
drawn from one carries an interval far wider than any census figure on this
map. Three rules keep that from being published as though it were precision:

* **Regions under MIN_SAMPLE respondents are dropped**, not estimated. 25 is
  the Demographic and Health Surveys' own threshold, and using the field's
  convention rather than inventing one keeps this comparable to how the source
  material is normally read. Botswana's 24 strata are first added into the
  ten districts they make up (POOLED), and the floor leaves 435 of the 505
  regions that gives.
* **A region between 25 and 49 is marked low precision** so the map can say so.
  DHS prints those in parentheses for the same reason.
* **"Not asked in the country" is honoured as a refusal to answer, not a zero.**
  Mauritania was never asked about religion, and Sudan, Tunisia and the
  Seychelles were never asked about ethnicity. Building a composition for them
  out of the remaining codes would invent a figure where a question does not
  exist, so those become ``not_collected`` -- the same status the EU countries
  that decline to ask about ethnicity already carry here.

Placed **first** in ADAPTER_FILES, which is the lowest authority: Ethiopia,
Mali and South Africa already have census figures and those must win. Because
merge_adapter works field by field, this still fills a field the census left
empty without touching one it filled.

The release has a row for every person interviewed, across 400-odd columns.
What is committed is not those rows but their counts (COUNTS): for each
region, question and answer, how many people gave it and their summed weight,
sorted by answer. With the labels beside them (LABELS), that is everything
this module reads. The sums are exact, which keeps every published share what
it was; it also means rows whose sums match, or add up to another row's, can
show that the same people gave both answers, so some people's three answers
can be matched up again from the file. They carry no names, and the release
itself prints each person's answers whole. ``--extract`` re-derives the counts
from the release, saved in data/raw/afrobarometer, where .gitignore keeps it
out of every commit.

Usage:
    python -m scripts.fetch_census.afrobarometer
    python -m scripts.fetch_census.afrobarometer --extract data/raw/afrobarometer/<release>.sav
    python -m scripts.fetch_census.afrobarometer \\
        --extract-xlsx data/raw/afrobarometer/R9.Merge_39ctry.xlsx
"""

from __future__ import annotations

import argparse
import csv
import gzip
import json
import re
from pathlib import Path
from typing import Any, Iterable

from ._shared import (NOT_AVAILABLE, NOT_COLLECTED, PROCESSED, RAW, gap, log,
                      record, shares, write_json)

HERE = RAW / "afrobarometer"
COUNTS = HERE / "r9_counts.csv.gz"
LABELS = HERE / "r9_labels.json"
OUT = "afrobarometer_region.json"
FIELDS = ("religion", "ethnicity", "language")

# Column positions in the published merge file, confirmed against the codebook.
COLUMNS = {"country": 1, "region": 9, "ethnicity": 297, "religion": 340, "weight": 396}

# DHS suppresses a cell below 25 unweighted cases and parenthesises 25-49.
MIN_SAMPLE = 25
LOW_PRECISION = 50

# Codes that are not an answer. 9994 is its own case and handled separately:
# the others are one person declining, 9994 is the country never being asked.
NOT_ANSWERED = {"-1", "9998", "9999"}
NOT_ASKED = "9994"
# A respondent who answered with their nationality, or said they do not think
# of themselves in those terms: the codebook's "(National identity) only, or
# 'doesn't think of self in those terms'". It is an answer and not a gap. It
# was written "No ethnic group", which is not what was said: 1,186 of
# Eswatini's 1,200 respondents carry the code, and its four regions read
# "No ethnic group" at 97 to 99.9%.
NATIONAL_ONLY = "9990"
NATIONAL_ONLY_LABEL = "National identity only"

# Afrobarometer's country code -> ISO3. Written out rather than matched on
# name: "Congo- Brazzaville" (the codebook's own spacing) is COG, the Republic
# of the Congo, and a name match that reached for COD instead would put one
# country's figures on the other's shapes without anything looking wrong.
ISO3 = {
    "2": "AGO", "3": "BEN", "4": "BWA", "5": "BFA", "6": "CPV", "7": "CMR",
    "8": "COG", "9": "CIV", "10": "SWZ", "11": "ETH", "12": "GAB", "13": "GMB",
    "14": "GHA", "15": "GIN", "16": "KEN", "17": "LSO", "18": "LBR", "19": "MDG",
    "20": "MWI", "21": "MLI", "22": "MRT", "23": "MUS", "24": "MAR", "25": "MOZ",
    "26": "NAM", "27": "NER", "28": "NGA", "29": "STP", "30": "SEN", "31": "SYC",
    "32": "SLE", "33": "ZAF", "34": "SDN", "35": "TZA", "36": "TGO", "37": "TUN",
    "38": "UGA", "39": "ZMB", "40": "ZWE",
}

SOURCE = "Afrobarometer Round 9 (2021-2023)"
UNIVERSE = ("Afrobarometer Round 9, a nationally representative sample of "
            "citizens of voting age. Regional shares are survey estimates, not "
            "census counts: they carry sampling error, and regions with fewer "
            "than 25 respondents are omitted rather than estimated.")

# Countries whose strata are added into the units the map draws before the
# 25-respondent floor is applied, with the word the note calls such a unit.
# Botswana is sampled in 24 strata over its ten districts: Francistown beside
# North East, Gaborone and Lobatse beside South East, Selibe Phikwe and five
# "Central" sub-districts in Central, and two to four in Kweneng, North-West,
# Southern and Kgalagadi. Read stratum by stratum, North-East District showed
# North East's 40 respondents for a polygon that holds Francistown too (56
# more), South-East District showed South East's 56 without Gaborone's 120,
# and Central, Kweneng, North-West, Southern and Kgalagadi showed nothing,
# though the survey interviewed 32 to 352 people in each. The strata are
# Round 8's, so its table says which district holds each
# (afrobarometer_r8.REGIONS), and their weighted answers and their
# respondents are added together as Round 8's are.
POOLED = {"BWA": "district"}


# The SPSS release's variables, which are named rather than counted from the
# left. The workbook route below is kept because it was how this was first
# read and the two agree exactly -- 1,010 of 1,010 country/region/field cells
# identical -- but the release is what should be read: it carries Q2, eight
# region labels the workbook's codebook omits, and 21 religion labels.
SAV_COLUMNS = {"country": "COUNTRY", "region": "REGION", "ethnicity": "Q84A",
               "religion": "Q95", "language": "Q2", "weight": "withinwt_hh"}
# withinwt_hh, not withinwt_ea: the household-level weight is the one in the
# workbook's column COLUMNS reads, checked row by row against both.
SAV_LABELS = {"COUNTRY": "country", "REGION": "region", "Q95": "religion",
              "Q84A": "ethnicity", "Q2": "language"}


# What each round commits in place of its respondents, and the three readers
# share: for every stratum, question and answer, the people who gave it and
# their summed weight, added in the order the respondents come. The rows are
# written sorted: in the order answers first appear, a stratum's first rows
# were its first respondent's answers, one person's record read back off the
# file.
def tally_answers(respondents: Iterable[tuple[tuple[str, ...], dict[str, str], float]]
                  ) -> dict[tuple[str, ...], list[Any]]:
    """Respondents and summed weight per stratum, question and answer."""
    counts: dict[tuple[str, ...], list[Any]] = {}
    for stratum, answers, weight in respondents:
        for question, value in answers.items():
            entry = counts.setdefault((*stratum, question, value), [0, 0.0])
            entry[0] += 1
            entry[1] += weight
    return counts


def write_counts(path: Path, strata: list[str],
                 counts: dict[tuple[str, ...], list[Any]]) -> int:
    """One row per stratum, question and answer, sorted; returns how many."""
    with gzip.open(path, "wt", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow([*strata, "question", "answer", "respondents", "weight"])
        for key, (n, weight) in sorted(counts.items()):
            writer.writerow([*key, n, repr(weight)])
    return len(counts)


def read_counts(path: Path) -> list[dict[str, str]]:
    with gzip.open(path, "rt", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def respondent(row: dict[str, Any]) -> tuple[tuple[str, str], dict[str, str], float] | None:
    """One interview as tally_answers takes it: its country and region, each
    question's answer as a code ("" for none), and its weight."""
    country, region = code(row["country"]), code(row["region"])
    if not country or not region:
        return None
    try:
        weight = float(row["weight"])
    except (TypeError, ValueError):
        weight = 1.0
    return ((country, region),
            {field: code(row[field]) or "" for field in FIELDS if field in row},
            weight)


def extract_sav(path: str) -> int:
    """Re-derive the committed counts and their labels from the SPSS release."""
    import pyreadstat                                # noqa: PLC0415 -- optional

    log(f"afrobarometer: reading {path}")
    frame, meta = pyreadstat.read_sav(path, usecols=list(SAV_COLUMNS.values()))
    labels = {name: {str(int(k)): v for k, v in
                     meta.variable_value_labels.get(name, {}).items()}
              for name in SAV_LABELS}
    HERE.mkdir(parents=True, exist_ok=True)
    LABELS.write_text(json.dumps(labels, ensure_ascii=False, indent=1,
                                 sort_keys=True), encoding="utf-8")
    log("  labels: " + ", ".join(f"{k} {len(v)}" for k, v in labels.items()))
    interviews = (respondent({name: "" if v != v else v
                              for name, v in zip(SAV_COLUMNS, values)})
                  for values in zip(*(frame[c] for c in SAV_COLUMNS.values())))
    rows = write_counts(COUNTS, ["country", "region"],
                        tally_answers(found for found in interviews if found))
    log(f"  wrote {COUNTS} ({rows:,} rows)")
    return rows


def load_labels() -> dict[str, dict[str, str]]:
    return json.loads(LABELS.read_text(encoding="utf-8"))


def code(value: Any) -> str | None:
    """Codes arrive as floats from the workbook; compare them as integers."""
    try:
        return str(int(float(value)))
    except (TypeError, ValueError):
        return None


# The codebook's "X only" means the respondent said X and named no sub-group,
# which reads as an odd label on a map. The plain name is what they said.
UNSPECIFIED = {"Christian only": "Christian", "Muslim only": "Muslim",
               "Sunni only": "Sunni"}

# The release's misspellings of a group's name. Q2's code 837 is "IIa", two
# capital I's, where Q84A's same code reads "Ila": Zambia's Ila, which the
# tree places and "IIa" was not. Q2 also writes Tanzania's Iraqw "Kiiraq",
# where Round 8 writes "Kiiraqw", and Ethiopia's Harari "Adergigna orHarari";
# Round 8's spelling of Harari is given the same form
# (afrobarometer_r8.SPELLING), so that a country's sum of the two rounds'
# regions names each language once.
LABEL_TYPOS = {"IIa": "Ila", "Kiiraq": "Kiiraqw",
               "Adergigna orHarari": "Adergigna or Harari"}

# Labels the release writes in French, by field, and Namibia's "Afrikaaner".
# Madagascar asks after "Malagasy officiel" and "Malagasy avec spécificité
# régionale", and Mauritius's ethnic list is written in French. "Hindou" and
# "Musulman" are two of the four communities Mauritius counts by religion and
# origin together, and are written as its constitution names them, "Hindu
# community" and "Muslim community": a bare "Muslim" is already the
# Factbook's name for Montenegro's Muslims, a people, and the two would be
# one group in the legend. "Euro-Mauricien" is Euro-Mauritian, not the
# narrower Franco-Mauritian the Factbook writes. Kept by field, because a
# religion label could be spelled like an ethnic one.
RELABEL: dict[str, dict[str, str]] = {
    "language": {
        "Malagasy officiel": "Official Malagasy",
        "Malagasy avec spécificité régionale": "Malagasy (regional variety)",
    },
    "ethnicity": {
        "Hindou": "Hindu community", "Musulman": "Muslim community",
        "Tamoul": "Tamil", "Telegou": "Telugu", "Chinois": "Chinese",
        "Afro-Mauricien (Creole)": "Afro-Mauritian (Creole)",
        "Euro-Mauricien (Blanc)": "Euro-Mauritian (White)",
        "Afrikaaner": "Afrikaner",
    },
}


def clean(label: str) -> str:
    """Drop the codebook's parenthetical gloss, keeping the group's name."""
    name = re.sub(r"\s*\((?:e\.g\.|i\.e\.|Do not).*$", "", label).strip().rstrip(",")
    name = UNSPECIFIED.get(name, name)
    return LABEL_TYPOS.get(name, name)


def named(field: str, value: str, table: dict[str, str]) -> str:
    """The group a code names, as it is written on the map."""
    if field == "ethnicity" and value == NATIONAL_ONLY:
        return NATIONAL_ONLY_LABEL
    name = clean(table.get(value, value))
    return RELABEL.get(field, {}).get(name, name)


def extract(path: str) -> int:
    """Re-derive the committed counts from the published workbook."""
    import openpyxl                                  # noqa: PLC0415 -- optional

    log(f"afrobarometer: reading {path}")
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book[book.sheetnames[0]]
    HERE.mkdir(parents=True, exist_ok=True)

    def interviews() -> Iterable[tuple[tuple[str, str], dict[str, str], float]]:
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row is None or len(row) <= COLUMNS["weight"]:
                continue
            found = respondent({name: row[i] for name, i in COLUMNS.items()})
            if found:
                yield found

    rows = write_counts(COUNTS, ["country", "region"], tally_answers(interviews()))
    book.close()
    log(f"  wrote {COUNTS} ({rows} rows)")
    return rows


def tabulate() -> tuple[dict[tuple[str, str], dict[str, dict[str, float]]],
                        dict[tuple[str, str], int], set[tuple[str, str]]]:
    """Weighted counts per country/region/field, plus the unweighted n."""
    tally: dict[tuple[str, str], dict[str, dict[str, float]]] = {}
    sample: dict[tuple[str, str], int] = {}
    never_asked: set[tuple[str, str]] = set()
    asked: set[tuple[str, str]] = set()
    for row in read_counts(COUNTS):
        country, region = row["country"], row["region"]
        field, value = row["question"], row["answer"]
        key = (country, region)
        cell = tally.setdefault(key, {name: {} for name in FIELDS})
        # Everyone answers every question once, if only by declining, so the
        # people counted under one question are the region's interviews.
        if field == FIELDS[0]:
            sample[key] = sample.get(key, 0) + int(row["respondents"])
        if value == NOT_ASKED:
            never_asked.add((country, field))
            continue
        asked.add((country, field))
        if not value or value in NOT_ANSWERED:
            continue
        cell[field][value] = cell[field].get(value, 0.0) + float(row["weight"])
    # A country counts as never asked only if no respondent in it was asked.
    return tally, sample, {pair for pair in never_asked if pair not in asked}


def units(tally: dict[tuple[str, str], dict[str, dict[str, float]]],
          sample: dict[tuple[str, str], int],
          labels: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    """What is published: each region as the survey samples it, and a POOLED
    country's units with their strata's weighted answers and respondents
    added together -- not their shares averaged.

    Listed in the order of the survey's codes; a pooled unit stands where its
    first stratum did.
    """
    # Round 8's module imports this one for its sampling rules, so its table
    # is read here rather than at the top.
    from .afrobarometer_r8 import REGIONS            # noqa: PLC0415
    out: list[dict[str, Any]] = []
    pooled: dict[tuple[str, str], dict[str, Any]] = {}
    for (country, region), cell in sorted(tally.items()):
        name = labels["REGION"].get(region)
        if not name:
            continue
        iso = ISO3.get(country)
        if not iso:
            raise SystemExit(f"afrobarometer: no ISO3 for country code {country}")
        if iso not in POOLED:
            out.append({"country": country, "iso": iso, "key": region, "name": name,
                        "cell": cell, "n": sample[(country, region)], "strata": [name]})
            continue
        shape = REGIONS.get((iso, name))
        if shape is None:
            # A unit summed without one of its strata would be short of it
            # and say nothing, so a stratum the table lacks stops the run.
            raise SystemExit(f"afrobarometer: {iso} stratum {name!r} is in no unit "
                             f"of afrobarometer_r8.REGIONS")
        unit = pooled.get((iso, shape))
        if unit is None:
            unit = pooled[(iso, shape)] = {
                "country": country, "iso": iso, "key": shape, "name": shape,
                "cell": {field: {} for field in cell}, "n": 0, "strata": []}
            out.append(unit)
        for field, counts in cell.items():
            into = unit["cell"][field]
            for value, weight in counts.items():
                into[value] = into.get(value, 0.0) + weight
        unit["n"] += sample[(country, region)]
        unit["strata"].append(name)
    return out


def listed(names: list[str]) -> str:
    """Names as a sentence lists them: "A and B", "A, B and C"."""
    names = sorted(names)
    return ", ".join(names[:-1]) + " and " + names[-1] if len(names) > 1 else names[0]


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--extract", metavar="SAV",
                    help="re-derive the committed counts and labels from the "
                         "SPSS release")
    ap.add_argument("--extract-xlsx", metavar="XLSX", dest="extract_xlsx",
                    help="the older route, reading the published workbook by "
                         "column position; writes no labels")
    ap.add_argument("--out", type=Path, default=None)
    args = ap.parse_args()

    if args.extract:
        extract_sav(args.extract)
        return 0
    if args.extract_xlsx:
        extract(args.extract_xlsx)
        return 0

    labels = load_labels()
    tally, sample, never_asked = tabulate()
    log(f"afrobarometer: {len(sample)} country/region cells, "
        f"{sum(sample.values())} interviews")
    for country, field in sorted(never_asked):
        log(f"  {labels['COUNTRY'].get(country, country)}: {field} not asked")
    # A question with no rows at all was not counted -- the workbook route has
    # no Q2 -- which is not the same as nobody answering it, so it is left out
    # rather than written as a gap in every region.
    counted = {row["question"] for row in read_counts(COUNTS)}
    for field in FIELDS:
        if field not in counted:
            log(f"  {field}: not in the counts, so no region carries it")

    records: list[dict[str, Any]] = []
    dropped = low = 0
    for unit in units(tally, sample, labels):
        country, iso, name, cell, n = (unit["country"], unit["iso"], unit["name"],
                                       unit["cell"], unit["n"])
        row_id = f"{iso}-AB9-{unit['key']}"
        if n < MIN_SAMPLE:
            # Not estimated, and the region says why rather than reading as
            # never surveyed: a reason the panel shows where nothing else has
            # a figure for the field (a count always stands in front of it).
            dropped += 1
            thin: dict[str, Any] = {}
            for field in FIELDS:
                if field not in counted:
                    continue
                if (country, field) in never_asked:
                    thin[field] = gap(NOT_COLLECTED)
                    thin[f"{field}_note"] = (
                        f"Afrobarometer did not ask about {field} in this country.")
                else:
                    thin[field] = gap(NOT_AVAILABLE, (
                        f"Afrobarometer's Round 9 survey (2021-2023) interviewed only {n} "
                        f"people in this region, too few to estimate {field} shares from, "
                        "so none are shown."))
            records.append(record(
                row_id, name, level="admin1", parent=iso, country=iso,
                **thin,
                sources=[{"field": "religion/ethnicity/language", "name": SOURCE,
                          "url": "https://www.afrobarometer.org/data/",
                          "license": "Afrobarometer data use policy"}]))
            continue
        fields: dict[str, Any] = {}
        for field, table in (("religion", labels["Q95"]),
                             ("ethnicity", labels["Q84A"]),
                             ("language", labels.get("Q2", {}))):
            if field not in counted or (field == "language" and not table):
                continue                     # counts or labels taken without Q2
            if (country, field) in never_asked:
                fields[field] = gap(NOT_COLLECTED)
                fields[f"{field}_note"] = (
                    f"Afrobarometer did not ask about {field} in this country.")
                continue
            # Added, not keyed: two codes can carry one label. Mozambique's
            # list has "Lomue" twice (550 and 559), and Zambézia's 65 weighted
            # answers under both read as the 7 under the second alone: 5.2%,
            # where the two make 34.6%, the region's largest language.
            counts: dict[str, float] = {}
            for value, weight in cell[field].items():
                group = named(field, value, table)
                counts[group] = counts.get(group, 0.0) + weight
            rows = shares(counts)
            fields[field] = rows or gap(NOT_AVAILABLE)
            if rows:
                note = f"{UNIVERSE} This region: {n} respondents."
                if len(unit["strata"]) > 1:
                    note += (f" The survey stratifies the {POOLED[iso]} as "
                             f"{listed(unit['strata'])}, and these shares are of "
                             "their respondents together.")
                if n < LOW_PRECISION:
                    note += (" Fewer than 50 respondents, so this share is"
                             " imprecise and should be read as indicative.")
                fields[f"{field}_note"] = note
        if n < LOW_PRECISION:
            low += 1
        records.append(record(
            row_id, name, level="admin1", parent=iso,
            country=iso,
            **fields,
            sources=[{"field": "religion/ethnicity/language", "name": SOURCE,
                      "url": "https://www.afrobarometer.org/data/",
                      "license": "Afrobarometer data use policy"}],
        ))
    log(f"  {len(records) - dropped} regions kept, {dropped} under {MIN_SAMPLE} "
        f"respondents, {low} marked low precision")
    write_json(args.out or PROCESSED / OUT, records)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
