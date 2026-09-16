#!/usr/bin/env python3
"""Report what a spreadsheet holds, before an adapter is written against it.

A statistical office's tables often arrive as a numbered series -- vol2_t1.xls,
vol2_t2.xls -- whose filenames say nothing about which one is the table you
want. Opening each in turn is the only way to find out, and doing that by
writing an adapter per guess is the slow way round.

This fetches a workbook, prints its sheet names and dimensions, and prints the
first rows of each sheet so the header structure is visible: how many rows to
skip, which column is the unit name, which are the shares. Nothing is written
and nothing is committed -- the output is the log.

Legacy ``.xls`` (BIFF) is read with xlrd, which since 2.0 reads that format and
no other; ``.xlsx`` with openpyxl. Both are in requirements.txt.

An Internet Archive capture is fetched raw with the ``id_`` modifier
(``/web/<timestamp>id_/<url>``), which returns the stored bytes rather than the
rewritten playback page -- without it a workbook comes back as HTML.

Usage:
    python scripts/probe_xls.py https://example.org/vol2_t1.xls
    python scripts/probe_xls.py --wayback 20190404144202 \
        http://www.recensamantromania.ro/wp-content/uploads/2015/05/vol2_t1.xls
    python scripts/probe_xls.py --rows 12 --sheets 3 URL...
"""

from __future__ import annotations

import argparse
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from common import http_get, log  # noqa: E402


def cell(value: object, width: int = 22) -> str:
    text = "" if value is None else str(value).strip()
    text = " ".join(text.split())
    return text[:width]


def rows_of(blob: bytes, name: str, args_start: int = 0) -> list[tuple[str, int, int, list[list[str]]]]:
    """(sheet name, rows, columns, first rows) for every sheet."""
    out: list[tuple[str, int, int, list[list[str]]]] = []
    if name.lower().endswith(".xlsx"):
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
        for sheet in book.worksheets:
            grid = [[cell(c) for c in row]
                    for row in sheet.iter_rows(max_row=args_start + 40, values_only=True)]
            out.append((sheet.title, sheet.max_row or 0, sheet.max_column or 0, grid))
        return out
    import xlrd
    book = xlrd.open_workbook(file_contents=blob)
    for sheet in book.sheets():
        grid = [[cell(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
                for r in range(min(sheet.nrows, args_start + 40))]
        out.append((sheet.name, sheet.nrows, sheet.ncols, grid))
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("url", nargs="+", help="one or more workbook URLs")
    ap.add_argument("--wayback", default=None, metavar="TIMESTAMP",
                    help="fetch each URL from the Internet Archive at this "
                         "capture, raw (the id_ modifier)")
    ap.add_argument("--rows", type=int, default=8, help="rows to print per sheet")
    ap.add_argument("--from", dest="start", type=int, default=0,
                    help="first row to print; the interesting part of a census "
                         "table is rarely the top of it")
    ap.add_argument("--sheets", type=int, default=6, help="sheets to print per workbook")
    ap.add_argument("--cols", type=int, default=12,
                    help="columns to print per row; a census table is wide and "
                         "the groups past the twelfth are the ones a first look misses")
    args = ap.parse_args()

    for url in args.url:
        target = (f"https://web.archive.org/web/{args.wayback}id_/{url}"
                  if args.wayback else url)
        log(f"{url}")
        try:
            blob = http_get(target, binary=True, cache=False, timeout=120)
        except Exception as exc:                      # noqa: BLE001 -- the log is the product
            log(f"  unreachable: {type(exc).__name__}: {exc}")
            continue
        log(f"  {len(blob):,} bytes")
        # A capture that played back as HTML is the id_ modifier missing, and
        # it must not be reported as an unreadable workbook.
        if blob[:15].lstrip().lower().startswith((b"<!doctype", b"<html")):
            log("  HTML, not a workbook -- the capture played back rewritten")
            continue
        try:
            sheets = rows_of(blob, url, args.start)
        except Exception as exc:                      # noqa: BLE001
            log(f"  cannot open: {type(exc).__name__}: {exc}")
            continue
        log(f"  {len(sheets)} sheet(s)")
        for name, nrows, ncols, grid in sheets[:max(0, args.sheets)]:
            log(f"  -- {name!r}: {nrows} rows x {ncols} cols")
            for row in grid[args.start:args.start + max(0, args.rows)]:
                if any(c for c in row):
                    log(f"     {row[:max(1, args.cols)]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
